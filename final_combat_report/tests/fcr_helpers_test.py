import sys
from datetime import datetime
from pathlib import Path

sys.path.append(str(Path(__file__).resolve().parent.parent))

import pytest
import pythoncom
import win32com.client
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, RGBColor

import helpers
from constants import UNIT_BATTALION, UNIT_COMPANY_ONE


@pytest.fixture
def doc():
    return Document()


# -------------------------
# get_signal_data
# -------------------------
def test_get_signal_data_missing_file_returns_empty_list(tmp_path):
    assert helpers.get_signal_data(str(tmp_path / "no_such.json")) == []


def test_get_signal_data_parses_jsonl_and_skips_blank_lines(tmp_path):
    path = tmp_path / "data.json"
    path.write_text('{"body": "a"}\n\n{"body": "b"}\n', encoding="utf-8")
    assert helpers.get_signal_data(str(path)) == [{"body": "a"}, {"body": "b"}]


def test_get_signal_data_invalid_json_line_returns_empty_list(tmp_path):
    path = tmp_path / "broken.json"
    path.write_text('{"body": "a"}\nне json\n', encoding="utf-8")
    assert helpers.get_signal_data(str(path)) == []


# -------------------------
# print_green / parse_any_date
# -------------------------
def test_print_green_outputs_text(capsys):
    helpers.print_green("готово")
    assert "готово" in capsys.readouterr().out


def test_parse_any_date_datetime_passthrough():
    dt = datetime(2026, 4, 21, 9, 0)
    assert helpers.parse_any_date(dt) is dt


def test_parse_any_date_parses_known_formats():
    assert helpers.parse_any_date("07:25 21.04.26") == datetime(2026, 4, 21, 7, 25)
    assert helpers.parse_any_date("21.04.2026") == datetime(2026, 4, 21)


def test_parse_any_date_unparsable_returns_min():
    assert helpers.parse_any_date("не дата") == datetime.min


# -------------------------
# clean_text
# -------------------------
def test_clean_text_normalizes_quotes_and_dashes():
    text = helpers.clean_text("“Газда” – «Шпонка» — test")
    assert text == '"Газда" - "Шпонка" - test'


def test_clean_text_collapses_whitespace_and_invisible_chars():
    text = helpers.clean_text("слово слово​  слово\t\nслово")
    assert text == "слово слово слово слово"


def test_clean_text_np_spacing():
    assert helpers.clean_text("н.п.Федорівське") == "н.п. Федорівське"


def test_clean_text_selo_spacing():
    assert helpers.clean_text("(Район с.Філія)") == "(Район с. Філія)"


def test_clean_text_selo_spacing_does_not_touch_lowercase_after_dot():
    # лукахед на ВЕЛИКУ літеру - "с." перед малою літерою (не назва села) не чіпаємо
    assert helpers.clean_text("щось с.іншого") == "щось с.іншого"


def test_clean_text_double_comma_with_space_collapses_to_comma_space():
    # сире "текст, ,(319-272)" (порожнє поле між двома комами) раніше через
    # " ," -> "," (прибирання пробілу ПЕРЕД комою) ставало "текст,,(319-272)"
    # - без пробілу зовсім; тепер лишається одна кома й пробіл
    assert helpers.clean_text("текст, ,(319-272)") == "текст, (319-272)"


def test_clean_text_literal_double_comma_collapses_to_comma_space():
    assert helpers.clean_text("текст,,(319-272)") == "текст, (319-272)"


def test_clean_text_quote_before_name_spacing():
    assert helpers.clean_text('ТЗ" Газда"') == 'ТЗ "Газда"'


def test_clean_text_closing_quote_mid_sentence_keeps_its_space():
    # регресія: стара регулярка розрізняла відкриваючу/закриваючу лапку за
    # сусіднім символом (літера з обох боків) - тому закриваючу лапку, після
    # якої одразу продовжувалось речення, теж хибно "виправляла" і ковтала
    # пробіл: 'Каспер" працював' ставало 'Каспер "працював'.
    text = helpers.clean_text('Екіпаж "Каспер" працював на логістику.')
    assert text == 'Екіпаж "Каспер" працював на логістику.'


def test_clean_text_alternating_quote_pairs_only_fix_odd_occurrences():
    # парність: 1-ше і 3-тє входження (відкриваючі) отримують пробіл перед і
    # без пробілу після; 2-ге і 4-те (закриваючі) лишаються без змін
    text = helpers.clean_text('ТЗ"Газда" звіт ТЗ"Шпонка" звіт')
    assert text == 'ТЗ "Газда" звіт ТЗ "Шпонка" звіт'


def test_clean_text_non_string_passthrough():
    assert helpers.clean_text(None) is None
    assert helpers.clean_text(5) == 5


def test_clean_text_known_typo_replacements():
    assert "знищено" in helpers.clean_text("знищенно ворога")
    assert "Лупиніс" in helpers.clean_text("Лупоніс 10 день")


# -------------------------
# parse_date_key / merge_and_sort_reports
# -------------------------
def test_parse_date_key_full_datetime():
    assert helpers.parse_date_key({"date": "14:05 21.04.2026"}) == datetime(2026, 4, 21, 14, 5)


def test_parse_date_key_missing_time_defaults_midnight():
    assert helpers.parse_date_key({"date": "21.04.2026"}) == datetime(2026, 4, 21, 0, 0)


def test_parse_date_key_strips_leading_none():
    assert helpers.parse_date_key({"date": "None 14:05 21.04.2026"}) == datetime(2026, 4, 21, 14, 5)


def test_parse_date_key_glued_time_and_date():
    assert helpers.parse_date_key({"date": "05:5813.07.26"}) == datetime(2026, 7, 13, 5, 58)


def test_parse_date_key_unparsable_returns_min():
    assert helpers.parse_date_key({"date": "не дата"}) == datetime.min
    assert helpers.parse_date_key({}) == datetime.min


def test_merge_and_sort_reports_sorts_and_strips_temp_key():
    a = [{"date": "10:00 01.05.2026", "text": "a"}]
    b = [{"Дата": "09:00 01.05.2026", "text": "b"}]
    result = helpers.merge_and_sort_reports(a, b)
    assert [r["text"] for r in result] == ["b", "a"]
    assert "date_obj" not in result[0]
    assert "date_obj" not in result[1]


# -------------------------
# convert_to_short_name / get_signature_officer
# -------------------------
def test_convert_to_short_name_basic():
    assert helpers.convert_to_short_name("Іванов Іван Іванович") == "Іван Іванов"


def test_convert_to_short_name_splits_glued_name_patronymic():
    assert helpers.convert_to_short_name("Петров ІванІванович") == "Іван Петров"


def test_convert_to_short_name_non_string_passthrough():
    assert helpers.convert_to_short_name(None) is None


def test_convert_to_short_name_invalid_raises():
    with pytest.raises(ValueError):
        helpers.convert_to_short_name("OnlyTwoWords")


def test_get_signature_officer_prefers_actual_rank_over_staff_rank():
    # реальні заголовки СПИСКУ: "Звання фактичне" (з великої), "звання за
    # штатом" (з малої) - саме так, несумісно, і є в реальному файлі
    rows = [{
        "Посада": "Командир батальйону", "П.І.Б.": "Іванов Іван Іванович",
        "Звання фактичне": "майор", "звання за штатом": "капітан",
    }]
    rank, name = helpers.get_signature_officer(rows, "Командир батальйону")
    assert rank == "майор"
    assert name == "Іван Іванов"


def test_get_signature_officer_falls_back_to_staff_rank():
    rows = [{
        "Посада": "Командир батальйону", "П.І.Б.": "Іванов Іван Іванович",
        "Звання фактичне": None, "звання за штатом": "капітан",
    }]
    rank, _ = helpers.get_signature_officer(rows, "Командир батальйону")
    assert rank == "капітан"


def test_get_signature_officer_not_found_returns_empty():
    assert helpers.get_signature_officer([], "Командир батальйону") == ("", "")


# -------------------------
# excel_col_to_index / get_rows / date_to_str / to_date
# -------------------------
def test_excel_col_to_index():
    assert helpers.excel_col_to_index('A') == 0
    assert helpers.excel_col_to_index('Z') == 25
    assert helpers.excel_col_to_index('AA') == 26
    assert helpers.excel_col_to_index('AD') == 29


def test_to_date_parsing_and_errors():
    assert helpers.to_date(datetime(2026, 4, 21)) == datetime(2026, 4, 21).date()
    assert helpers.to_date("2026-04-21") == datetime(2026, 4, 21).date()
    assert helpers.to_date("21.04.2026") == datetime(2026, 4, 21).date()
    with pytest.raises(ValueError):
        helpers.to_date(12345)


def test_date_to_str_minus_and_plus():
    assert helpers.date_to_str("21.04.2026", action='-', days=1) == "20.04.2026"
    assert helpers.date_to_str("21.04.2026", action='+', days=1) == "22.04.2026"


def test_get_rows_converts_timestamps():
    import pandas as pd

    df = pd.DataFrame({'A': [1], 'B': [pd.Timestamp('2026-04-21')], 'C': ['text']})
    rows = helpers.get_rows(df, ['A', 'B', 'C'])
    assert rows == [{'A': 1, 'B': '21.04.2026', 'C': 'text'}]


def test_get_rows_date_can_be_empty_with_nat():
    import pandas as pd

    df = pd.DataFrame([{"A": pd.NaT, "B": "x"}])
    rows = helpers.get_rows(df, ["A", "B"], date_can_be_empty=True)
    assert rows == [{"A": None, "B": "x"}]


# -------------------------
# add_paragraph_with_style / add_custom_heading
# -------------------------
def test_add_paragraph_with_style_basic_formatting(doc):
    paragraph = helpers.add_paragraph_with_style(
        doc, "Test text", font_name="Arial", font_size=16, bold=True,
        alignment=WD_ALIGN_PARAGRAPH.CENTER, first_line_indent=1.2, format_tabs=True,
    )
    run = paragraph.runs[0]
    assert run.text == "Test text"
    assert run.font.name == "Arial"
    assert run._element.rPr.rFonts.get(qn('w:eastAsia')) == "Arial"
    assert run.font.size.pt == 16
    assert run.font.bold is True
    assert run.font.color.rgb == RGBColor(0, 0, 0)
    assert paragraph.alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert paragraph.paragraph_format.first_line_indent.cm == pytest.approx(1.2, abs=0.01)

    tab_stops = paragraph.paragraph_format.tab_stops
    sec = doc.sections[0]
    width = Inches(sec.page_width.inches - (sec.left_margin.inches + sec.right_margin.inches))
    tabs = [t.position.inches for t in tab_stops]
    assert any(abs(t - width.inches) < 0.01 for t in tabs)


def test_add_paragraph_with_style_bold_word_list(doc):
    paragraph = helpers.add_paragraph_with_style(doc, "ОВТ – 0 од.", bold=["ОВТ"])
    bold_runs = [r for r in paragraph.runs if r.font.bold]
    non_bold_runs = [r for r in paragraph.runs if not r.font.bold]
    assert any(r.text == "ОВТ" for r in bold_runs)
    assert any("од." in r.text for r in non_bold_runs)


def test_add_paragraph_with_style_no_indent_by_default(doc):
    paragraph = helpers.add_paragraph_with_style(doc, "text")
    assert paragraph.paragraph_format.first_line_indent is None


def test_add_custom_heading(doc):
    heading = helpers.add_custom_heading(doc, "1. ВИСНОВКИ", heading=3, font_size=12)
    assert heading.style.name == "Heading 3"
    assert heading.runs[0].text == "1. ВИСНОВКИ"
    assert heading.runs[0].bold is True
    assert heading.runs[0].font.size.pt == 12


# -------------------------
# add_paragraph_with_parts / add_custom_heading_with_parts /
# mark_no_auto_highlight / highlight_changes_from_previous_report
# -------------------------
def test_add_paragraph_with_parts_highlights_only_marked_segments(doc):
    paragraph = helpers.add_paragraph_with_parts(doc, [("№", False), ("78", True), (".", False)])
    assert [r.text for r in paragraph.runs] == ["№", "78", "."]
    assert [bool(r.font.highlight_color) for r in paragraph.runs] == [False, True, False]


def test_add_paragraph_with_parts_bold_applies_to_all_runs(doc):
    paragraph = helpers.add_paragraph_with_parts(doc, [("a", False), ("b", True)], bold=True)
    assert all(r.bold for r in paragraph.runs)


def test_add_paragraph_with_parts_skips_empty_segments(doc):
    paragraph = helpers.add_paragraph_with_parts(doc, [("", False), ("text", False)])
    assert len(paragraph.runs) == 1


def test_add_paragraph_with_parts_applies_first_line_indent_when_positive(doc):
    paragraph = helpers.add_paragraph_with_parts(doc, [("text", False)], first_line_indent=1.2)
    assert paragraph.paragraph_format.first_line_indent.cm == pytest.approx(1.2, abs=0.01)


def test_add_custom_heading_with_parts_bold_and_heading_style(doc):
    paragraph = helpers.add_custom_heading_with_parts(doc, [("3.1.2 обстрілів: ", False), ("4", True)], heading=3)
    assert paragraph.style.name == "Heading 3"
    assert all(r.bold for r in paragraph.runs)
    assert [bool(r.font.highlight_color) for r in paragraph.runs] == [False, True]


def test_add_custom_heading_with_parts_skips_empty_segments(doc):
    paragraph = helpers.add_custom_heading_with_parts(doc, [("", False), ("текст", False)], heading=3)
    assert len(paragraph.runs) == 1


def test_parts_paragraph_is_auto_registered_for_exclusion(doc):
    paragraph = helpers.add_paragraph_with_parts(doc, [("text", False)])
    assert hasattr(doc, "_no_auto_highlight")
    assert paragraph._p in doc._no_auto_highlight


def test_mark_no_auto_highlight_survives_fresh_paragraph_wrappers(doc):
    # doc.paragraphs повертає нову Paragraph-обгортку щоразу - перевіряємо, що
    # виключення все одно спрацьовує для того самого абзацу через свіжу обгортку.
    paragraph = doc.add_paragraph("hello")
    helpers.mark_no_auto_highlight(doc, paragraph)

    fresh = doc.paragraphs[0]
    assert fresh is not paragraph
    assert fresh._p in doc._no_auto_highlight


def test_highlight_changes_no_previous_doc_is_noop(doc):
    doc.add_paragraph("будь-який текст")
    helpers.highlight_changes_from_previous_report(doc, None)
    assert not any(r.font.highlight_color for p in doc.paragraphs for r in p.runs)


def test_highlight_changes_identical_paragraph_not_highlighted(doc):
    doc.add_paragraph("той самий текст")
    previous = Document()
    previous.add_paragraph("той самий текст")

    helpers.highlight_changes_from_previous_report(doc, previous)

    assert not any(r.font.highlight_color for p in doc.paragraphs for r in p.runs)


def test_highlight_changes_new_paragraph_fully_highlighted(doc):
    doc.add_paragraph("новий абзац, якого не було вчора")
    previous = Document()
    previous.add_paragraph("зовсім інший текст")

    helpers.highlight_changes_from_previous_report(doc, previous)

    p = doc.paragraphs[0]
    assert all(r.font.highlight_color for r in p.runs)


def test_highlight_changes_respects_no_highlight_exclusion(doc):
    paragraph = helpers.add_paragraph_with_style(doc, "особового складу – 0, з них: ")
    helpers.mark_no_auto_highlight(doc, paragraph)
    previous = Document()
    previous.add_paragraph("особового складу – 1, з них: ")  # реальні дані вчора відрізняються

    helpers.highlight_changes_from_previous_report(doc, previous)

    assert not any(r.font.highlight_color for r in paragraph.runs)


def test_highlight_changes_skips_blank_paragraphs(doc):
    doc.add_paragraph("")
    previous = Document()

    helpers.highlight_changes_from_previous_report(doc, previous)

    assert not any(r.font.highlight_color for p in doc.paragraphs for r in p.runs)


def test_highlight_changes_table_cell_changed_vs_unchanged(doc):
    today_table = doc.add_table(rows=1, cols=2)
    today_table.rows[0].cells[0].text = "змінилось"
    today_table.rows[0].cells[1].text = "без змін"

    previous = Document()
    prev_table = previous.add_table(rows=1, cols=2)
    prev_table.rows[0].cells[0].text = "було інакше"
    prev_table.rows[0].cells[1].text = "без змін"

    helpers.highlight_changes_from_previous_report(doc, previous)

    changed_cell_runs = today_table.rows[0].cells[0].paragraphs[0].runs
    unchanged_cell_runs = today_table.rows[0].cells[1].paragraphs[0].runs
    assert all(r.font.highlight_color for r in changed_cell_runs)
    assert not any(r.font.highlight_color for r in unchanged_cell_runs)


def test_highlight_changes_table_cell_excluded_from_diff(doc):
    today_table = doc.add_table(rows=1, cols=1)
    cell_paragraph = today_table.rows[0].cells[0].paragraphs[0]
    cell_paragraph.text = "24.07.2026"
    helpers.mark_no_auto_highlight(doc, cell_paragraph)

    previous = Document()
    prev_table = previous.add_table(rows=1, cols=1)
    prev_table.rows[0].cells[0].text = "23.07.2026"

    helpers.highlight_changes_from_previous_report(doc, previous)

    assert not any(r.font.highlight_color for r in today_table.rows[0].cells[0].paragraphs[0].runs)


def test_highlight_changes_ignores_extra_rows_and_tables(doc):
    # Сьогоднішній документ має більше таблиць/рядків, ніж учорашній - зайве
    # просто ігнорується, без IndexError.
    doc.add_table(rows=2, cols=1)
    previous = Document()
    previous.add_table(rows=1, cols=1)

    helpers.highlight_changes_from_previous_report(doc, previous)  # не повинно кидати виняток


def test_highlight_changes_ignores_extra_table_beyond_previous_doc(doc):
    doc.add_table(rows=1, cols=1)
    doc.add_table(rows=1, cols=1)  # другої таблиці немає у вчорашньому документі
    previous = Document()
    previous.add_table(rows=1, cols=1)

    helpers.highlight_changes_from_previous_report(doc, previous)  # не повинно кидати виняток


def test_highlight_changes_ignores_extra_columns_beyond_previous_doc(doc):
    today_table = doc.add_table(rows=1, cols=2)
    today_table.rows[0].cells[1].text = "нова колонка"
    previous = Document()
    previous.add_table(rows=1, cols=1)

    helpers.highlight_changes_from_previous_report(doc, previous)  # не повинно кидати виняток


# -------------------------
# set_margins / set_line_spacing / add_text_for_header_and_footer_document
# -------------------------
def test_set_margins(doc):
    helpers.set_margins(doc, Inches(1), Inches(2), Inches(3), Inches(4))
    for s in doc.sections:
        assert s.top_margin == Inches(1)
        assert s.bottom_margin == Inches(2)
        assert s.left_margin == Inches(3)
        assert s.right_margin == Inches(4)


def test_set_line_spacing_applies_to_all_paragraphs(doc):
    doc.add_paragraph("перший")
    doc.add_paragraph("другий")
    helpers.set_line_spacing(doc, 12)
    for p in doc.paragraphs:
        assert p.paragraph_format.line_spacing.pt == 12


def test_add_text_for_header_and_footer_document(doc):
    helpers.add_text_for_header_and_footer_document(doc)
    sec = doc.sections[0]
    assert sec.different_first_page_header_footer is True
    assert sec.header.paragraphs[0].text.strip() == "ДЛЯ СЛУЖБОВОГО КОРИСТУВАННЯ"
    assert sec.footer.paragraphs[0].text.strip() == "ДЛЯ СЛУЖБОВОГО КОРИСТУВАННЯ"
    assert sec.first_page_footer.paragraphs[0].text.strip() == "ДЛЯ СЛУЖБОВОГО КОРИСТУВАННЯ"


# -------------------------
# paragraph_five_dot_one_table / paragraph_five_dot_two_table
# -------------------------
def test_paragraph_five_dot_one_table_structure_and_date_highlight(doc):
    unit_nospace = UNIT_COMPANY_ONE.replace(" ", "")
    units_data = {unit_nospace: {"безповоротні": 2, "санітарні_заг": 5}}
    helpers.paragraph_five_dot_one_table(doc, "24.07.2026", "25.07.2026", 18, units_data)

    table = doc.tables[0]
    assert table.rows[0].cells[0].text == "Підрозділ"
    assert "24.07.2026" in table.rows[0].cells[1].text
    assert "25.07.2026" in table.rows[0].cells[1].text

    date_runs = table.rows[0].cells[1].paragraphs[0].runs
    highlighted_texts = [r.text for r in date_runs if r.font.highlight_color]
    assert highlighted_texts == ["24.07.2026", "25.07.2026"]

    row_1rmp = next(r for r in table.rows if r.cells[0].text == unit_nospace)
    assert row_1rmp.cells[1].text == "2"
    assert row_1rmp.cells[6].text == "5"


def test_table_cell_parts_skips_empty_segments(doc):
    # cell.text = "" (усередині _table_cell_parts) лишає порожній перший run -
    # реальний доданий вміст завжди в останньому run.
    table = doc.add_table(rows=1, cols=1)
    para = helpers._table_cell_parts(table.rows[0].cells[0], [("", False), ("24.07.2026", True)])
    assert para.runs[-1].text == "24.07.2026"
    assert para.runs[-1].font.highlight_color


def test_paragraph_five_dot_one_table_empty_data_leaves_blank_cells(doc):
    helpers.paragraph_five_dot_one_table(doc, "24.07.2026", "25.07.2026", 18, None)
    table = doc.tables[0]
    row_1rmp = next(r for r in table.rows if r.cells[0].text == UNIT_COMPANY_ONE.replace(" ", ""))
    assert row_1rmp.cells[1].text == ""


def test_paragraph_five_dot_two_table_empty_data_leaves_blank_cells(doc):
    helpers.paragraph_five_dot_two_table(doc, "24.07.2026", "25.07.2026", 19, None)
    table = doc.tables[0]
    row_tankiv = next(r for r in table.rows if r.cells[0].text == "Танків")
    assert row_tankiv.cells[1].text == ""


def test_paragraph_five_dot_two_table_structure_and_date_highlight(doc):
    ovt_data = {"Танків": {"знищено": 1, "пошкоджено_заг": 3}}
    helpers.paragraph_five_dot_two_table(doc, "24.07.2026", "25.07.2026", 19, ovt_data)

    table = doc.tables[0]
    assert "24.07.2026" in table.rows[0].cells[1].text
    date_runs = table.rows[0].cells[1].paragraphs[0].runs
    highlighted_texts = [r.text for r in date_runs if r.font.highlight_color]
    assert highlighted_texts == ["24.07.2026", "25.07.2026"]

    row_tankiv = next(r for r in table.rows if r.cells[0].text == "Танків")
    assert row_tankiv.cells[1].text == "1"


def test_paragraph_five_dot_two_table_carries_forward_total_row(doc):
    ovt_summary_label = f"Всього ОВТ за {UNIT_BATTALION}"
    ovt_data = {ovt_summary_label: {"знищено": 4}}
    helpers.paragraph_five_dot_two_table(doc, "24.07.2026", "25.07.2026", 19, ovt_data)
    table = doc.tables[0]
    total_row = table.rows[-1]
    assert total_row.cells[0].text == ovt_summary_label
    assert total_row.cells[1].text == "4"


# -------------------------
# close_open_output_documents
# -------------------------
def test_close_open_output_documents_missing_directory_returns_early(tmp_path):
    # директорія не існує - виходимо одразу, до будь-яких win32com викликів
    helpers.close_open_output_documents(str(tmp_path / "does_not_exist"))


def test_close_open_output_documents_returns_when_pywin32_not_importable(tmp_path, monkeypatch):
    out = tmp_path / "output"
    out.mkdir()
    # None у sys.modules - сигнал CPython, що попередній імпорт провалився,
    # тож "import win32com.client" одразу підніме ImportError
    monkeypatch.setitem(sys.modules, "win32com.client", None)
    helpers.close_open_output_documents(str(out))  # не повинно впасти


def test_close_open_output_documents_returns_when_word_not_running(tmp_path, monkeypatch):
    out = tmp_path / "output"
    out.mkdir()
    monkeypatch.setattr(pythoncom, "CoInitialize", lambda: None)

    def _raise(name):
        raise Exception("Word.Application не запущено")

    monkeypatch.setattr(win32com.client, "GetActiveObject", _raise)
    helpers.close_open_output_documents(str(out))  # не повинно впасти


class _FakeWordDoc:
    def __init__(self, full_name, name):
        self.FullName = full_name
        self.Name = name
        self.closed_with = None

    def Close(self, SaveChanges=False):
        self.closed_with = SaveChanges


class _BrokenWordDoc:
    Name = "зламаний.docx"

    @property
    def FullName(self):
        raise RuntimeError("помилка COM при читанні шляху")


def test_close_open_output_documents_closes_matching_skips_unrelated_and_broken(tmp_path, monkeypatch, capsys):
    out = tmp_path / "output"
    out.mkdir()
    target_file = out / "report.docx"
    target_file.write_text("x")

    matching_doc = _FakeWordDoc(str(target_file), "report.docx")
    unrelated_doc = _FakeWordDoc(str(tmp_path / "other.docx"), "other.docx")
    broken_doc = _BrokenWordDoc()

    class _FakeWord:
        Documents = [matching_doc, unrelated_doc, broken_doc]

    monkeypatch.setattr(pythoncom, "CoInitialize", lambda: None)
    monkeypatch.setattr(pythoncom, "CoUninitialize", lambda: None)
    monkeypatch.setattr(win32com.client, "GetActiveObject", lambda name: _FakeWord())

    helpers.close_open_output_documents(str(out))

    assert matching_doc.closed_with is False
    assert unrelated_doc.closed_with is None
    assert "Закрито відкритий файл: report.docx" in capsys.readouterr().out


# -------------------------
# create_or_clear_output_directory
# -------------------------
def test_create_or_clear_output_directory(tmp_path):
    out = tmp_path / "output"
    out.mkdir()
    (out / "old.docx").write_text("stale")

    helpers.create_or_clear_output_directory(str(out))

    assert out.exists()
    assert list(out.iterdir()) == []


def test_create_or_clear_output_directory_creates_missing_dir(tmp_path):
    out = tmp_path / "does_not_exist_yet"
    helpers.create_or_clear_output_directory(str(out))
    assert out.exists() and out.is_dir()


def test_create_or_clear_output_directory_retries_on_permission_error_then_succeeds(tmp_path, monkeypatch):
    out = tmp_path / "output"
    out.mkdir()
    (out / "locked.docx").write_text("x")

    monkeypatch.setattr(helpers, "close_open_output_documents", lambda d: None)
    monkeypatch.setattr(helpers.time, "sleep", lambda s: None)

    calls = {"count": 0}
    real_rmtree = helpers.shutil.rmtree

    def fake_rmtree(path):
        calls["count"] += 1
        if calls["count"] < 3:
            raise PermissionError("locked")
        real_rmtree(path)

    monkeypatch.setattr(helpers.shutil, "rmtree", fake_rmtree)

    helpers.create_or_clear_output_directory(str(out))

    assert calls["count"] == 3
    assert out.exists()
    assert list(out.iterdir()) == []


def test_create_or_clear_output_directory_raises_after_all_retries_fail(tmp_path, monkeypatch):
    out = tmp_path / "output"
    out.mkdir()

    monkeypatch.setattr(helpers, "close_open_output_documents", lambda d: None)
    monkeypatch.setattr(helpers.time, "sleep", lambda s: None)

    def always_fails(path):
        raise PermissionError("still locked")

    monkeypatch.setattr(helpers.shutil, "rmtree", always_fails)

    with pytest.raises(PermissionError):
        helpers.create_or_clear_output_directory(str(out))
