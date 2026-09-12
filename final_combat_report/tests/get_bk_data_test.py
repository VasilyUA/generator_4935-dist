import sys
from pathlib import Path

sys.path.append(str(Path(__file__).resolve().parent.parent))

import openpyxl
import pytest

from get_normalized_data import get_bk_data as m

PREV = "24.07.2026"
TODAY = "25.07.2026"


@pytest.fixture(autouse=True)
def bk_dir(tmp_path, monkeypatch):
    # BK_UAV_FILE_PATH імпортований у модуль за значенням - патчимо саме ім'я
    # в get_bk_data, а не в constants, щоб не чіпати реальні resources/.
    monkeypatch.setattr(m, "BK_UAV_FILE_PATH", str(tmp_path / "БК ВБАК.xlsx"))
    return tmp_path


def _file_path(bk_dir):
    return bk_dir / "БК ВБАК.xlsx"


def _bk_file(bk_dir, rows, date=PREV, extra_qty_columns=None):
    """rows: список (id, synonyms, uav_type, item, qty) - значення qty йдуть у
    стовпець "кількість\\nза {date}". extra_qty_columns (опційно): ще стовпці
    кількості за іншими датами - {date: [qty, ...]} у тому ж порядку рядків."""
    wb = openpyxl.Workbook()
    ws = wb.active
    header = ["ID", "Синоніми", "Тип", "Назва", m._qty_header(date)]
    for extra_date in (extra_qty_columns or {}):
        header.append(m._qty_header(extra_date))
    ws.append(header)
    for i, row in enumerate(rows):
        line = list(row)
        for extra_date, values in (extra_qty_columns or {}).items():
            line.append(values[i])
        ws.append(line)
    wb.save(_file_path(bk_dir))


def _msg(text, spend):
    return {"text": text, "spend": dict(spend)}


# -------------------------
# read_bk_stock
# -------------------------
def test_read_bk_stock_missing_file_returns_empty():
    assert m.read_bk_stock(PREV) == []


def test_read_bk_stock_missing_date_column_returns_empty(bk_dir):
    _bk_file(bk_dir, [[1, "", "мавік", "МОА-400", 18]])
    assert m.read_bk_stock("01.01.2099") == []


def test_read_bk_stock_parses_rows_and_skips_blank_item(bk_dir):
    _bk_file(bk_dir, [
        [1, "", "мавік", "МОА-400", 18],
        [2, "", "", "", None],  # рядок без назви - має бути пропущений
    ])

    rows = m.read_bk_stock(PREV)

    assert len(rows) == 1
    assert rows[0]["item"] == "МОА-400"
    assert rows[0]["qty"] == 18


def test_read_bk_stock_picks_the_right_dated_column(bk_dir):
    _bk_file(
        bk_dir, [[1, "", "мавік", "МОА-400", 18]],
        date=PREV, extra_qty_columns={TODAY: [2]},
    )

    assert m.read_bk_stock(PREV)[0]["qty"] == 18
    assert m.read_bk_stock(TODAY)[0]["qty"] == 2


# -------------------------
# resolve_bk_synonyms
# -------------------------
def test_resolve_bk_synonyms_replaces_text_and_spend_when_stock_available(bk_dir):
    _bk_file(bk_dir, [[1, "свп осколок, уламок", "FPV", "HFS1300F", 39]])
    data = [_msg("Засіб: FPV. Витрата: свп осколок - 1 шт.", {"свп осколок": 1})]

    m.resolve_bk_synonyms(data, PREV, TODAY)

    assert data[0]["text"] == "Засіб: FPV. Витрата: HFS1300F - 1 шт."
    assert data[0]["spend"] == {"HFS1300F": 1}
    assert m.read_bk_stock(TODAY)[0]["qty"] == 38
    assert m.read_bk_stock(PREV)[0]["qty"] == 39  # учорашній стовпець не змінюється


def test_resolve_bk_synonyms_adds_a_new_column_not_a_new_file(bk_dir):
    _bk_file(bk_dir, [[1, "свп осколок", "FPV", "HFS1300F", 39]])
    data = [_msg("Засіб: FPV. Витрата: свп осколок - 1 шт.", {"свп осколок": 1})]

    m.resolve_bk_synonyms(data, PREV, TODAY)

    # той самий файл, обидва стовпці кількості в ньому
    wb = openpyxl.load_workbook(_file_path(bk_dir))
    header = [c.value for c in wb.active[1]]
    assert m._qty_header(PREV) in header
    assert m._qty_header(TODAY) in header
    assert len(list(_file_path(bk_dir).parent.glob("*.xlsx"))) == 1


def test_resolve_bk_synonyms_no_replacement_when_stock_is_zero(bk_dir):
    _bk_file(bk_dir, [[1, "свп осколок", "FPV", "HFS1300F", 0]])
    text = "Засіб: FPV. Витрата: свп осколок - 1 шт."
    data = [_msg(text, {"свп осколок": 1})]

    m.resolve_bk_synonyms(data, PREV, TODAY)

    assert data[0]["text"] == text
    assert data[0]["spend"] == {"свп осколок": 1}
    assert m.read_bk_stock(TODAY)[0]["qty"] == 0


def test_resolve_bk_synonyms_no_replacement_when_uav_type_not_mentioned(bk_dir):
    _bk_file(bk_dir, [[1, "свп осколок", "FPV", "HFS1300F", 39]])
    text = "Засіб: Vampire. Витрата: свп осколок - 1 шт."
    data = [_msg(text, {"свп осколок": 1})]

    m.resolve_bk_synonyms(data, PREV, TODAY)

    assert data[0]["text"] == text
    assert m.read_bk_stock(TODAY)[0]["qty"] == 39


def test_resolve_bk_synonyms_matches_compound_name_by_single_word(bk_dir):
    # реальний випадок: "СВП Уламок" не дорівнює жодному синоніму цілком, але
    # містить слово "уламок" - точний однослівний синонім з колонки B
    _bk_file(bk_dir, [[1, "свп осколок, уламок", "FPV", "HFS1300F", 39]])
    data = [_msg("Засіб: FPV. Витрата: СВП Уламок - 1 шт.", {"СВП Уламок": 1})]

    m.resolve_bk_synonyms(data, PREV, TODAY)

    assert data[0]["text"] == "Засіб: FPV. Витрата: HFS1300F - 1 шт."
    assert data[0]["spend"] == {"HFS1300F": 1}


def test_resolve_bk_synonyms_matches_informal_fragment_of_official_name(bk_dir):
    # реальний випадок: оператор пише неформальний/неповний фрагмент офіційної
    # назви (без слова "Боєприпас", з незакритою лапкою, в іншому порядку слів)
    # - жодного синоніма з колонки B немає, але значущі слова назви ("ко",
    # "пузатий", "змій") усі є в офіційній назві (колонка D)
    _bk_file(bk_dir, [[1, "щось інше", "FPV", 'Боєприпас КО 1,3 "Пузатий змій"', 44]])
    data = [_msg('Засіб: FPV. Витрата: КО "Пузатий змій 1,3 - 1 шт.', {'КО "Пузатий змій 1,3': 1})]

    m.resolve_bk_synonyms(data, PREV, TODAY)

    assert data[0]["spend"] == {'Боєприпас КО 1,3 "Пузатий змій"': 1}


def test_resolve_bk_synonyms_no_match_when_name_has_no_significant_tokens(bk_dir):
    # назва без жодного значущого слова (лише число) не повинна збігатися з
    # офіційною назвою просто через спільну цифру
    _bk_file(bk_dir, [[1, "щось інше", "FPV", 'Боєприпас КО 1,3 "Пузатий змій"', 44]])
    data = [_msg("Засіб: FPV. Витрата: 1,3 - 1 шт.", {"1,3": 1})]

    m.resolve_bk_synonyms(data, PREV, TODAY)

    assert data[0]["spend"] == {"1,3": 1}


def test_resolve_bk_synonyms_matches_name_without_variant_suffix(bk_dir):
    # реальний випадок: повідомлення каже "ОГБ", а синонім у таблиці - "ОГБ-1"
    # (з номером варіанта) - без номера теж має спрацювати
    _bk_file(bk_dir, [[1, "огб-1", "Vampire", "ОГ-Б1", 38]])
    data = [_msg("Засіб: Vampire. Витрата: ОГБ - 4 шт.", {"ОГБ": 4})]

    m.resolve_bk_synonyms(data, PREV, TODAY)

    assert data[0]["text"] == "Засіб: Vampire. Витрата: ОГ-Б1 - 4 шт."
    assert data[0]["spend"] == {"ОГ-Б1": 4}
    assert m.read_bk_stock(TODAY)[0]["qty"] == 34


def test_resolve_bk_synonyms_short_synonym_does_not_match_substring_of_other_word(bk_dir):
    # короткий синонім "ф" не повинен збігатися з довшим словом, що просто
    # містить цю літеру (наприклад "графік")
    _bk_file(bk_dir, [[1, "ф", "FPV", "Ф-1", 15]])
    data = [_msg("Засіб: FPV. Витрата: графік - 1 шт.", {"графік": 1})]

    m.resolve_bk_synonyms(data, PREV, TODAY)

    assert data[0]["spend"] == {"графік": 1}


def test_resolve_bk_synonyms_variant_suffix_stripping_does_not_shrink_below_safe_length(bk_dir):
    # "f-1" знявши суфікс дав би "f" (1 символ) - лишається лише точний
    # збіг "f-1" цілком, суфіксне знімання тут не застосовується
    _bk_file(bk_dir, [[1, "f-1", "FPV", "Ф-1", 15]])
    data = [_msg("Засіб: FPV. Витрата: fортеця - 1 шт.", {"fортеця": 1})]

    m.resolve_bk_synonyms(data, PREV, TODAY)

    assert data[0]["spend"] == {"fортеця": 1}


def test_resolve_bk_synonyms_picks_row_with_largest_remaining_stock(bk_dir):
    _bk_file(bk_dir, [
        [1, "мавік", "FPV", "GHO-1", 43],
        [2, "мавік", "FPV", "HGO-TA", 40],
        [3, "мавік", "FPV", "Ф-1", 15],
    ])
    data = [_msg("Засіб: FPV. Витрата: мавік - 1 шт.", {"мавік": 1})]

    m.resolve_bk_synonyms(data, PREV, TODAY)

    assert data[0]["spend"] == {"GHO-1": 1}
    by_item = {r["item"]: r["qty"] for r in m.read_bk_stock(TODAY)}
    assert by_item["GHO-1"] == 42
    assert by_item["HGO-TA"] == 40
    assert by_item["Ф-1"] == 15


def test_resolve_bk_synonyms_two_synonyms_same_message_decrement_sequentially(bk_dir):
    # обидва слова - синоніми з тієї самої групи B, обидва йдуть в один і той
    # самий (найбільший) рядок - разом -2 від залишку
    _bk_file(bk_dir, [[1, "свп осколок, уламок", "FPV", "HFS1300F", 39]])
    data = [_msg(
        "Засіб: FPV. Витрата: свп осколок - 1 шт, уламок - 1 шт.",
        {"свп осколок": 1, "уламок": 1},
    )]

    m.resolve_bk_synonyms(data, PREV, TODAY)

    assert data[0]["text"] == "Засіб: FPV. Витрата: HFS1300F - 1 шт, HFS1300F - 1 шт."
    assert data[0]["spend"] == {"HFS1300F": 2}
    assert m.read_bk_stock(TODAY)[0]["qty"] == 37


def test_resolve_bk_synonyms_stops_replacing_once_stock_hits_zero(bk_dir):
    # перше слово вичерпує залишок до 0 - друге (та сама група) вже не замінюється
    _bk_file(bk_dir, [[1, "свп осколок, уламок", "FPV", "HFS1300F", 1]])
    data = [_msg(
        "Засіб: FPV. Витрата: свп осколок - 1 шт, уламок - 1 шт.",
        {"свп осколок": 1, "уламок": 1},
    )]

    m.resolve_bk_synonyms(data, PREV, TODAY)

    assert "HFS1300F - 1 шт" in data[0]["text"]
    assert "уламок - 1 шт" in data[0]["text"]  # друге лишилось незамінене
    assert data[0]["spend"] == {"HFS1300F": 1, "уламок": 1}
    assert m.read_bk_stock(TODAY)[0]["qty"] == 0


def test_resolve_bk_synonyms_floors_at_zero_when_usage_exceeds_stock(bk_dir):
    _bk_file(bk_dir, [[1, "мавік", "FPV", "GHO-1", 1]])
    data = [_msg("Засіб: FPV. Витрата: мавік - 5шт.", {"мавік": 5})]

    m.resolve_bk_synonyms(data, PREV, TODAY)

    assert data[0]["spend"] == {"GHO-1": 5}
    assert m.read_bk_stock(TODAY)[0]["qty"] == 0


def test_resolve_bk_synonyms_ignores_messages_without_spend(bk_dir):
    _bk_file(bk_dir, [[1, "свп осколок", "FPV", "HFS1300F", 39]])
    data = [{"text": "просто повідомлення без витрати"}]

    m.resolve_bk_synonyms(data, PREV, TODAY)  # не повинно кидати виняток

    assert m.read_bk_stock(TODAY)[0]["qty"] == 39


def test_resolve_bk_synonyms_noop_when_file_missing(bk_dir):
    data = [_msg("Засіб: FPV. Витрата: свп осколок - 1 шт.", {"свп осколок": 1})]
    m.resolve_bk_synonyms(data, PREV, TODAY)
    assert data[0]["spend"] == {"свп осколок": 1}  # без змін
    assert not _file_path(bk_dir).exists()


def test_resolve_bk_synonyms_noop_when_previous_column_missing(bk_dir):
    _bk_file(bk_dir, [[1, "свп осколок", "FPV", "HFS1300F", 39]], date="01.01.2000")
    data = [_msg("Засіб: FPV. Витрата: свп осколок - 1 шт.", {"свп осколок": 1})]

    m.resolve_bk_synonyms(data, PREV, TODAY)

    assert data[0]["spend"] == {"свп осколок": 1}  # без змін - нема бази PREV
    assert m.read_bk_stock(TODAY) == []


def _bk_file_with_base(bk_dir, rows):
    """Файл БК ВБАК у ПОЧАТКОВОМУ стані - лише стовпець-БАЗА без дати
    (m._BASE_QTY_HEADER), як буває до першого запуску resolve_bk_synonyms()."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["ID", "Синоніми", "Тип", "Назва", m._BASE_QTY_HEADER])
    for row in rows:
        ws.append(list(row))
    wb.save(_file_path(bk_dir))


def test_resolve_bk_synonyms_bootstraps_from_dateless_base_column_on_first_run(bk_dir):
    # перший запуск узагалі - датованого стовпця за PREV ще нема ніде, лише
    # початковий стовпець-БАЗА без дати ("кількість.") - без фолбека на нього
    # функція назавжди лишалась би неактивною (нізвідки узяти перший датований)
    _bk_file_with_base(bk_dir, [[1, "свп осколок", "FPV", "HFS1300F", 39]])
    data = [_msg("Засіб: FPV. Витрата: свп осколок - 1 шт.", {"свп осколок": 1})]

    m.resolve_bk_synonyms(data, PREV, TODAY)

    assert data[0]["spend"] == {"HFS1300F": 1}
    assert m.read_bk_stock(TODAY)[0]["qty"] == 38


def test_resolve_bk_synonyms_prefers_dated_previous_column_over_base_when_both_exist(bk_dir):
    # якщо датований стовпець за PREV УЖЕ є - саме він база, а не стовпець-БАЗА
    # без дати (той уже застарів після першого реального запуску)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["ID", "Синоніми", "Тип", "Назва", m._BASE_QTY_HEADER, m._qty_header(PREV)])
    ws.append([1, "свп осколок", "FPV", "HFS1300F", 39, 10])
    wb.save(_file_path(bk_dir))
    data = [_msg("Засіб: FPV. Витрата: свп осколок - 1 шт.", {"свп осколок": 1})]

    m.resolve_bk_synonyms(data, PREV, TODAY)

    assert m.read_bk_stock(TODAY)[0]["qty"] == 9  # від PREV (10), не від бази (39)


# -------------------------
# resolve_bk_synonyms - BK_SYNONYM_CREW_OVERRIDES
# -------------------------
def test_resolve_bk_synonyms_crew_override_wins_over_largest_remaining_stock(bk_dir, monkeypatch):
    # без override "осколок" пішов би на HGO-TA (більший залишок), як у
    # test_resolve_bk_synonyms_picks_row_with_largest_remaining_stock; з
    # override для екіпажу "джміль" - примусово на ЗБ-2500, попри менший залишок
    monkeypatch.setattr(m, "BK_SYNONYM_CREW_OVERRIDES", {("джміль", "осколок"): "ЗБ-2500"})
    _bk_file(bk_dir, [
        [1, "свп осколок", "FPV, бамблбі", "ЗБ-2500", 26],
        [2, "свп осколок", "FPV, бамблбі", "HGO-TA", 40],
    ])
    data = [_msg("Засіб: бамблбі. Екіпаж: Джміль. Витрата: осколок - 2 шт.", {"осколок": 2})]

    m.resolve_bk_synonyms(data, PREV, TODAY)

    assert data[0]["spend"] == {"ЗБ-2500": 2}
    by_item = {r["item"]: r["qty"] for r in m.read_bk_stock(TODAY)}
    assert by_item["ЗБ-2500"] == 24
    assert by_item["HGO-TA"] == 40  # незачеплений


def test_resolve_bk_synonyms_crew_override_ignored_for_other_crews(bk_dir, monkeypatch):
    # override діє лише для вказаного екіпажу - для іншого екіпажу з тим самим
    # синонімом і далі звичайний резолв за найбільшим залишком
    monkeypatch.setattr(m, "BK_SYNONYM_CREW_OVERRIDES", {("джміль", "осколок"): "ЗБ-2500"})
    _bk_file(bk_dir, [
        [1, "осколок", "FPV, бамблбі", "ЗБ-2500", 26],
        [2, "осколок", "FPV, бамблбі", "HGO-TA", 40],
    ])
    data = [_msg("Засіб: бамблбі. Екіпаж: Каспер. Витрата: осколок - 2 шт.", {"осколок": 2})]

    m.resolve_bk_synonyms(data, PREV, TODAY)

    assert data[0]["spend"] == {"HGO-TA": 2}


def test_resolve_bk_synonyms_does_not_reprocess_already_resolved_display_name(bk_dir):
    # get_UAV_data.py вже сам зіставляє деякі неформальні назви на офіційні
    # (UAV_MUNITION_DISPLAY_NAMES, напр. "КО Пузатий змій" -> 'КО 1,3 "Пузатий
    # змій"') ДО того, як cost_data взагалі дійшло сюди. Без цього застереження
    # resolve_bk_synonyms() перезаписував би вже коректну назву на сирий запис
    # із колонки D таблиці (тут - з небажаним префіксом "Боєприпас").
    official = list(m.UAV_MUNITION_DISPLAY_NAMES.values())[0]
    _bk_file(bk_dir, [[1, "щось інше", "FPV, бамблбі", f"Боєприпас {official}", 25]])
    data = [_msg(f"Засіб: бамблбі. Витрата: {official} - 1 шт.", {official: 1})]

    m.resolve_bk_synonyms(data, PREV, TODAY)

    assert data[0]["spend"] == {official: 1}  # без змін
    assert m.read_bk_stock(TODAY)[0]["qty"] == 25  # залишок не списаний


def test_resolve_bk_synonyms_crew_override_null_explicitly_skips_without_generic_fallback(bk_dir, monkeypatch):
    # override з null - явно НЕ резолвити, і без фолбеку на звичайний пошук
    # (на відміну від "ціль поза залишком", де фолбек на генерик Є) - інакше
    # найбільший залишок серед ~40 однаково широких синонімів обрав би щось
    # для екіпажу/синоніма, для якого користувач явно попросив не чіпати.
    monkeypatch.setattr(m, "BK_SYNONYM_CREW_OVERRIDES", {("джміль", "фугас"): None})
    _bk_file(bk_dir, [[1, "фугас", "FPV, бамблбі", "HGO-TA", 40]])
    data = [_msg("Засіб: бамблбі. Екіпаж: Джміль. Витрата: фугас - 2 шт.", {"фугас": 2})]

    m.resolve_bk_synonyms(data, PREV, TODAY)

    assert data[0]["spend"] == {"фугас": 2}
    assert m.read_bk_stock(TODAY)[0]["qty"] == 40  # незачеплений


def test_resolve_bk_synonyms_crew_override_skipped_when_target_item_not_in_table(bk_dir, monkeypatch):
    # override вказує на офіційну назву, якої в таблиці взагалі немає (напр.
    # застаріле налаштування) - falls back до звичайного резолву
    monkeypatch.setattr(m, "BK_SYNONYM_CREW_OVERRIDES", {("джміль", "осколок"): "Немає такої позиції"})
    _bk_file(bk_dir, [[1, "осколок", "FPV, бамблбі", "HGO-TA", 40]])
    data = [_msg("Засіб: бамблбі. Екіпаж: Джміль. Витрата: осколок - 2 шт.", {"осколок": 2})]

    m.resolve_bk_synonyms(data, PREV, TODAY)

    assert data[0]["spend"] == {"HGO-TA": 2}


def test_resolve_bk_synonyms_crew_override_skipped_when_target_item_out_of_stock(bk_dir, monkeypatch):
    # цільова позиція override - без залишку: falls back до звичайного резолву
    monkeypatch.setattr(m, "BK_SYNONYM_CREW_OVERRIDES", {("джміль", "осколок"): "ЗБ-2500"})
    _bk_file(bk_dir, [
        [1, "осколок", "FPV, бамблбі", "ЗБ-2500", 0],
        [2, "осколок", "FPV, бамблбі", "HGO-TA", 40],
    ])
    data = [_msg("Засіб: бамблбі. Екіпаж: Джміль. Витрата: осколок - 2 шт.", {"осколок": 2})]

    m.resolve_bk_synonyms(data, PREV, TODAY)

    assert data[0]["spend"] == {"HGO-TA": 2}


def test_resolve_bk_synonyms_older_date_columns_are_untouched(bk_dir):
    _bk_file(
        bk_dir, [[1, "свп осколок", "FPV", "HFS1300F", 39]],
        date=PREV, extra_qty_columns={"23.07.2026": [50]},
    )
    data = [_msg("Засіб: FPV. Витрата: свп осколок - 1 шт.", {"свп осколок": 1})]

    m.resolve_bk_synonyms(data, PREV, TODAY)

    assert m.read_bk_stock("23.07.2026")[0]["qty"] == 50
    assert m.read_bk_stock(PREV)[0]["qty"] == 39
    assert m.read_bk_stock(TODAY)[0]["qty"] == 38


# -------------------------
# Повторна генерація того самого дня - стовпець TODAY перестворюється, а не дублюється
# -------------------------
def test_resolve_bk_synonyms_regeneration_recomputes_fresh_not_cumulatively(bk_dir):
    _bk_file(bk_dir, [[1, "мавік", "FPV", "GHO-1", 43]])
    first_run = [_msg("Засіб: FPV. Витрата: мавік - 1 шт.", {"мавік": 1})]
    m.resolve_bk_synonyms(first_run, PREV, TODAY)
    assert m.read_bk_stock(TODAY)[0]["qty"] == 42

    # повторна генерація того самого дня - cost_data знову "сира" (наче щойно
    # перечитана з Signal), той самий синонім
    second_run = [_msg("Засіб: FPV. Витрата: мавік - 1 шт.", {"мавік": 1})]
    m.resolve_bk_synonyms(second_run, PREV, TODAY)

    assert second_run[0]["spend"] == {"GHO-1": 1}
    assert m.read_bk_stock(TODAY)[0]["qty"] == 42  # те саме, а не 41 (не подвійне списання)
    assert m.read_bk_stock(PREV)[0]["qty"] == 43  # учорашній стовпець і далі незмінний


def test_resolve_bk_synonyms_regeneration_does_not_duplicate_todays_column(bk_dir):
    _bk_file(bk_dir, [[1, "мавік", "FPV", "GHO-1", 43]])
    for _ in range(3):
        m.resolve_bk_synonyms(
            [_msg("Засіб: FPV. Витрата: мавік - 1 шт.", {"мавік": 1})], PREV, TODAY,
        )

    wb = openpyxl.load_workbook(_file_path(bk_dir))
    header = [c.value for c in wb.active[1]]
    assert header.count(m._qty_header(TODAY)) == 1
    assert len(header) == 6  # ID, Синоніми, Тип, Назва, PREV, TODAY - без дублів
