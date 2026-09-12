import sys
from pathlib import Path

sys.path.append(str(Path(__file__).resolve().parent.parent))

import pytest
from docx import Document

import constants
import helpers
from get_normalized_data import get_previous_report_data as m


@pytest.fixture(autouse=True)
def resources_dir(tmp_path, monkeypatch):
    # RESOURCES_DIR імпортований у модуль за значенням - патчимо саме ім'я тут,
    # щоб тести не зверталися до реальної resources/ теки користувача.
    monkeypatch.setattr(m, "RESOURCES_DIR", str(tmp_path))
    return tmp_path


def test_to_int_digit_and_non_digit():
    assert m._to_int("5") == 5
    assert m._to_int("") == 0
    assert m._to_int(None) == 0
    assert m._to_int("не число") == 0


def test_get_previous_report_missing_file_returns_none():
    assert m.get_previous_report("01.01.2099") is None


def test_get_previous_report_returns_none_for_corrupted_docx(resources_dir):
    path = resources_dir / m.REPORT_FILE_NAME_TEMPLATE.format(date="24.07.2026")
    path.write_text("не насправді docx", encoding="utf-8")

    assert m.get_previous_report("24.07.2026") is None


def test_get_previous_report_reads_existing_docx(resources_dir):
    doc = Document()
    doc.add_paragraph("тестовий вміст")
    path = resources_dir / m.REPORT_FILE_NAME_TEMPLATE.format(date="24.07.2026")
    doc.save(path)

    result = m.get_previous_report("24.07.2026")

    assert result is not None
    assert result.paragraphs[0].text == "тестовий вміст"


def test_get_previous_report_number_extracts_from_header():
    doc = Document()
    doc.add_paragraph(
        f"ПІДСУМКОВЕ БОЙОВЕ ДОНЕСЕННЯ командира {constants.UNIT_BATTALION} {constants.UNIT_BRIGADE} №77"
    )
    assert m.get_previous_report_number(doc) == 77


def test_get_previous_report_number_none_for_missing_doc():
    assert m.get_previous_report_number(None) is None


def test_get_previous_report_number_none_when_not_found_in_first_paragraphs():
    doc = Document()
    for i in range(20):
        doc.add_paragraph("нерелевантний текст")
    doc.paragraphs[19].text = (
        f"{constants.UNIT_BATTALION} {constants.UNIT_BRIGADE} №99"
    )  # за межами перших 15 абзаців
    assert m.get_previous_report_number(doc) is None


# -------------------------
# get_previous_personnel_totals (таблиця 5.1)
# -------------------------
def test_get_previous_personnel_totals_reads_cumulative_columns():
    doc = Document()
    unit_nospace = constants.UNIT_COMPANY_ONE.replace(" ", "")
    units_data = {unit_nospace: {"безповоротні_заг": 2, "санітарні_заг": 1, "зниклі_заг": 0, "полонені_заг": 0}}
    helpers.paragraph_five_dot_one_table(doc, "23.07.2026", "24.07.2026", 18, units_data)

    result = m.get_previous_personnel_totals(doc)

    assert result[unit_nospace] == {"безповоротні_заг": 2, "санітарні_заг": 1, "зниклі_заг": 0, "полонені_заг": 0}


def test_get_previous_personnel_totals_skips_attached_units_header_row():
    doc = Document()
    helpers.paragraph_five_dot_one_table(doc, "23.07.2026", "24.07.2026", 18, {})
    result = m.get_previous_personnel_totals(doc)
    assert "Придані підрозділи" not in result


def test_get_previous_personnel_totals_skips_rows_with_too_few_cells():
    doc = Document()
    table = doc.add_table(rows=3, cols=3)  # замало колонок (потрібно >= 9)
    table.cell(2, 0).text = constants.UNIT_COMPANY_ONE.replace(" ", "")
    assert m.get_previous_personnel_totals(doc) == {}


def test_get_previous_personnel_totals_none_doc_returns_empty():
    assert m.get_previous_personnel_totals(None) == {}


def test_get_previous_personnel_totals_no_tables_returns_empty():
    assert m.get_previous_personnel_totals(Document()) == {}


# -------------------------
# get_previous_ovt_totals (таблиця 5.2, друга таблиця в документі)
# -------------------------
def test_get_previous_ovt_totals_reads_cumulative_columns():
    doc = Document()
    doc.add_table(rows=1, cols=1)  # заглушка замість таблиці 5.1 - ОВТ має бути tables[1]
    ovt_data = {"Танків": {"знищено_заг": 1, "пошкоджено_заг": 2, "всього_заг": 3}}
    helpers.paragraph_five_dot_two_table(doc, "23.07.2026", "24.07.2026", 19, ovt_data)

    result = m.get_previous_ovt_totals(doc)

    assert result["Танків"] == {"знищено_заг": 1, "пошкоджено_заг": 2, "всього_заг": 3}


def test_get_previous_ovt_totals_carries_forward_total_row():
    doc = Document()
    doc.add_table(rows=1, cols=1)
    ovt_summary_label = f"Всього ОВТ за {constants.UNIT_BATTALION}"
    ovt_data = {ovt_summary_label: {"знищено_заг": 4, "пошкоджено_заг": 0, "всього_заг": 4}}
    helpers.paragraph_five_dot_two_table(doc, "23.07.2026", "24.07.2026", 19, ovt_data)

    result = m.get_previous_ovt_totals(doc)

    assert result[ovt_summary_label]["знищено_заг"] == 4


def test_get_previous_ovt_totals_skips_rows_with_too_few_cells():
    doc = Document()
    doc.add_table(rows=1, cols=1)  # заглушка замість таблиці 5.1
    table = doc.add_table(rows=3, cols=3)  # замало колонок (потрібно >= 7)
    table.cell(2, 0).text = "Танків"
    assert m.get_previous_ovt_totals(doc) == {}


def test_get_previous_ovt_totals_skips_rows_with_empty_label():
    doc = Document()
    doc.add_table(rows=1, cols=1)  # заглушка замість таблиці 5.1
    table = doc.add_table(rows=3, cols=7)  # 7 колонок, як у реальній табл. 5.2
    for i in range(7):
        table.cell(2, i).text = "" if i == 0 else "1"
    assert m.get_previous_ovt_totals(doc) == {}


def test_get_previous_ovt_totals_requires_second_table():
    doc = Document()
    doc.add_table(rows=1, cols=1)  # лише одна таблиця в документі
    assert m.get_previous_ovt_totals(doc) == {}


def test_get_previous_ovt_totals_none_doc_returns_empty():
    assert m.get_previous_ovt_totals(None) == {}
