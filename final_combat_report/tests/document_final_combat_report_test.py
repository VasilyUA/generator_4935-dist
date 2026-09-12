import sys
from pathlib import Path

sys.path.append(str(Path(__file__).resolve().parent.parent))

import pytest
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

import document_final_combat_report as m
from constants import UNIT_COMPANY_ONE


@pytest.fixture
def doc():
    return Document()


# -------------------------
# P / H / P_parts / H_parts
# -------------------------
def test_P_defaults_to_1_2cm_first_line_indent(doc):
    paragraph = m.P(doc, "текст")
    assert paragraph.paragraph_format.first_line_indent.cm == pytest.approx(1.2, abs=0.01)


def test_P_no_highlight_registers_exclusion(doc):
    paragraph = m.P(doc, "текст", no_highlight=True)
    assert paragraph._p in doc._no_auto_highlight


def test_P_without_no_highlight_does_not_register_exclusion(doc):
    m.P(doc, "текст")
    assert not getattr(doc, "_no_auto_highlight", set())


def test_H_uses_heading_style_and_left_alignment(doc):
    paragraph = m.H(doc, "1. ВИСНОВКИ", heading=2)
    assert paragraph.style.name == "Heading 2"
    assert paragraph.alignment == WD_ALIGN_PARAGRAPH.LEFT


def test_H_no_highlight_registers_exclusion(doc):
    paragraph = m.H(doc, "1. ВИСНОВКИ", no_highlight=True)
    assert paragraph._p in doc._no_auto_highlight


def test_P_parts_defaults_to_1_2cm_indent_and_highlights_marked_segment(doc):
    paragraph = m.P_parts(doc, [("№", False), ("78", True)])
    assert paragraph.paragraph_format.first_line_indent.cm == pytest.approx(1.2, abs=0.01)
    assert [bool(r.font.highlight_color) for r in paragraph.runs] == [False, True]


def test_H_parts_uses_heading_style(doc):
    paragraph = m.H_parts(doc, [("3.1.2 обстрілів: ", False), ("4", True)], heading=3)
    assert paragraph.style.name == "Heading 3"


# -------------------------
# paragraph_shelling - лише ненульові лічильники підсвічуються
# -------------------------
def test_paragraph_shelling_highlights_only_nonzero_count(doc):
    m.paragraph_shelling(doc, ["09:00 артобстріл позицій"])

    heading = doc.paragraphs[0]
    assert heading.text == "3.1.2. Противник здійснив 1 обстрілів:"
    highlighted = [r.text for r in heading.runs if r.font.highlight_color]
    assert highlighted == ["1"]


def test_paragraph_shelling_zero_count_not_highlighted(doc):
    m.paragraph_shelling(doc, [])
    heading = doc.paragraphs[0]
    assert heading.text == "3.1.2. Противник здійснив 0 обстрілів:"
    assert not any(r.font.highlight_color for r in heading.runs)


def test_paragraph_shelling_shows_only_the_six_fixed_detail_categories():
    doc = Document()
    m.paragraph_shelling(doc, [])
    detail_labels = [
        "Ракетних ударів", "Авіаційних ударів", "Артилерійський обстріл",
        "Ударів дронів камікадзе", "Здійснення скидів з БпЛА", "Підрив на СВП",
    ]
    for label in detail_labels:
        assert any(p.text.startswith(f"{label} – 0:") for p in doc.paragraphs)
    # категорії поза фіксованою шісткою не отримують окремого блоку заголовка
    assert not any(p.text.startswith("Танковий –") for p in doc.paragraphs)


def test_paragraph_shelling_summary_line_matches_reference_format():
    # регресія: підсумковий рядок 3.1.2 має точно збігатись зі стилем
    # довідника - роздільники ";"/","/"." різні для різних полів, назви
    # полів рядковими літерами, і присутнє поле "БМП" (без "Мінометних")
    doc = Document()
    data_paintball = [
        "09:00 25.07.2026 артобстріл позицій",
        "10:00 25.07.2026 скид з mavic",
        "11:00 25.07.2026 скид з мавік",
        "12:00 25.07.2026 скид з mavic",
        "13:00 25.07.2026 удар fpv-дроном",
    ]
    m.paragraph_shelling(doc, data_paintball)
    summary_line = doc.paragraphs[1].text
    assert summary_line == (
        'авіаційних ударів - 0; артилерійських обстрілів - 1, РСЗВ – 0, гранатометів – 0, '
        'скиди з БпЛА — 3, удари дронів камікадзе (FPV, Молнія, FPV) – 1, БМП – 0, '
        'танковий – 0, ТОС – 0, БПЛА ЛАНЦЕТ (баражуючий боєприпас) — 0, снайпер - 0, '
        'стріл. зброя – 0; фосфор - 0; гранати - 0., підрив на СВП - 0.'
    )


# -------------------------
# generate_head_documents - лише номер і дата підсвічуються
# -------------------------
def test_generate_head_documents_highlights_only_number_and_date(doc):
    static_data = {"location_command_post": {"city": "НОВОПАВЛІВКА"}}
    m.generate_head_documents(doc, "25.07.2026", 18, static_data, 78)

    header = next(p for p in doc.paragraphs if "ПІДСУМКОВЕ" in p.text)
    highlighted = [r.text for r in header.runs if r.font.highlight_color]
    assert highlighted == ["78", "25.07.2026"]
    assert "№78" in header.text
    assert "18.00 25.07.2026" in header.text


# -------------------------
# open_generated_file - не падає, якщо startfile не зміг відкрити файл
# -------------------------
def test_open_generated_file_swallows_exception_when_startfile_fails(monkeypatch, capsys):
    def _raise(path):
        raise OSError("немає застосунку для відкриття")

    monkeypatch.setattr(m.os, "startfile", _raise)
    m.open_generated_file("some/path.docx")
    assert "Не вдалося автоматично відкрити файл" in capsys.readouterr().out


# -------------------------
# paragraph_combat_work_of_units - 3.2, кожне повідомлення окремим абзацом
# -------------------------
def test_paragraph_combat_work_of_units_renders_each_message(doc):
    unit_nospace = UNIT_COMPANY_ONE.replace(" ", "")
    m.paragraph_combat_work_of_units(doc, [{"text": f"09:00 25.07.2026 {unit_nospace} щось"}])
    assert any(p.text == f"09:00 25.07.2026 {unit_nospace} щось" for p in doc.paragraphs)


# -------------------------
# paragraph_progress_of_planned_tasks - 3.3, кожен запис окремим абзацом
# -------------------------
def test_paragraph_progress_of_planned_tasks_renders_each_entry(doc):
    m.paragraph_progress_of_planned_tasks(doc, [{"text": "виліт 1 доставлено", "date": "09:00 25.07.2026"}], [])
    assert any(p.text == "виліт 1 доставлено" for p in doc.paragraphs)


# -------------------------
# paragraph_artillery_task - 3.3.6, кожне повідомлення окремим абзацом
# -------------------------
def test_paragraph_artillery_task_renders_each_message(doc):
    m.paragraph_artillery_task(doc, [{"text": "13:35 ВГрМ Пежо витрата ОФ-843Б"}])
    assert any(p.text == "13:35 ВГрМ Пежо витрата ОФ-843Б" for p in doc.paragraphs)


# -------------------------
# paragraph_working_and_inspection_groups - регресія на відступ першого рядка
# -------------------------
def test_commission_entries_get_first_line_indent(doc):
    m.paragraph_working_and_inspection_groups(doc, ["09:49 25.07.2026 перевірка КСП"])
    entry = next(p for p in doc.paragraphs if p.text.startswith("09:49"))
    assert entry.paragraph_format.first_line_indent is not None
    assert entry.paragraph_format.first_line_indent.cm > 0


# -------------------------
# document_final_combat_report - наскрізний прогін без звернення до реальних resources/
# -------------------------
def test_document_final_combat_report_generates_valid_docx(tmp_path, monkeypatch):
    monkeypatch.setattr(m, "OUTPUT_DIR", str(tmp_path))
    monkeypatch.setattr(m, "get_static_data_json", lambda path: {
        "location_command_post": {"city": "НОВОПАВЛІВКА"},
        "paragraph_one": "тестове розвідувальне зведення",
        "paragraph_two": [f"{UNIT_COMPANY_ONE} утримує позиції"],
        "paragraph_seven": ["Немає"],
    })
    # get_previous_report - єдина точка, що читає resources/ - підміняємо, щоб
    # тест не залежав від реальних файлів користувача.
    monkeypatch.setattr(m, "get_previous_report", lambda date: None)
    monkeypatch.setattr(m.os, "startfile", lambda path: None)
    # resolve_bk_synonyms читає БК за selected_date/previous_date - з появою
    # фолбека на бездатний стовпець-БАЗУ (_find_base_qty_column) далека дата
    # в selected_date вже НЕ гарантує безпечний no-op проти РЕАЛЬНОГО файлу
    # resources/ - тепер підміняємо BK_UAV_FILE_PATH явно, як і в тесті нижче.
    import openpyxl
    from get_normalized_data import get_bk_data
    monkeypatch.setattr(get_bk_data, "BK_UAV_FILE_PATH", str(tmp_path / "БК ВБАК.xlsx"))

    data_uav = {"cost_data": [], "logistics_data": [], "loss_data": []}

    m.document_final_combat_report("31.12.2098", 18, [], [], data_uav, [], [], [])

    generated = list(tmp_path.glob("*.docx"))
    assert len(generated) == 1

    doc = Document(str(generated[0]))
    header = next(p for p in doc.paragraphs if "ПІДСУМКОВЕ" in p.text)
    assert "№1" in header.text  # немає попереднього звіту -> номер починається з 1
    assert any(p.text == "1. ВИСНОВКИ З ОЦІНКИ ПРОТИВНИКА" for p in doc.paragraphs)
    assert any(p.text == "7. ПРОБЛЕМНІ ПИТАННЯ" for p in doc.paragraphs)
    assert len(doc.tables) == 2


def test_document_final_combat_report_resolves_bk_synonyms_before_rendering(tmp_path, monkeypatch):
    # Регресія на те, що 3.3.5 (paragraph_combat_work_of_uav) і 6 (paragraph_six)
    # обидва бачать замінений (за таблицею БК ВБАК) текст/витрату - resolve_bk_synonyms
    # має відпрацювати ДО побудови обох розділів.
    import openpyxl
    from get_normalized_data import get_bk_data
    monkeypatch.setattr(get_bk_data, "BK_UAV_FILE_PATH", str(tmp_path / "БК ВБАК.xlsx"))
    # resolve_bk_synonyms читає базові залишки зі стовпця УЧОРАШНЬОЇ (previous_date) дати
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["ID", "Синоніми", "Тип", "Назва", get_bk_data._qty_header("30.12.2098")])
    ws.append([1, "мавік", "FPV", "GHO-1", 43])
    wb.save(str(tmp_path / "БК ВБАК.xlsx"))

    monkeypatch.setattr(m, "OUTPUT_DIR", str(tmp_path))
    monkeypatch.setattr(m, "get_static_data_json", lambda path: {
        "location_command_post": {"city": "НОВОПАВЛІВКА"},
        "paragraph_one": "тест", "paragraph_two": [], "paragraph_seven": [],
    })
    monkeypatch.setattr(m, "get_previous_report", lambda date: None)
    monkeypatch.setattr(m.os, "startfile", lambda path: None)

    data_uav = {
        "cost_data": [{"text": "Засіб: FPV. Витрата: мавік - 1 шт.", "spend": {"мавік": 1}}],
        "logistics_data": [],
        "loss_data": [],
    }

    m.document_final_combat_report("31.12.2098", 18, [], [], data_uav, [], [], [])

    generated = list(tmp_path.glob("Додаток*.docx"))
    doc = Document(str(generated[0]))
    texts = [p.text for p in doc.paragraphs]
    assert any("GHO-1 - 1 шт" in t for t in texts)  # 3.3.5 - текст повідомлення
    assert any(t.lower().startswith("вбпак:") and "gho-1" in t.lower() for t in texts)  # 6 - зведення

    rows = get_bk_data.read_bk_stock("31.12.2098")
    assert rows[0]["qty"] == 42  # списано у новий стовпець того самого файлу
