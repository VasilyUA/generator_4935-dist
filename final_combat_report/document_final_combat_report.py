from datetime import timedelta, datetime
import os

from docx import Document
from docx.shared import Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

from constants import OUTPUT_DIR, FILE_PATH_STATIC_DATA, REPORT_FILE_NAME_TEMPLATE, UNIT_BRIGADE, UNIT_COMPANY_ARTILLERY, UNIT_COMPANY_DSHR, UNIT_COMPANY_ONE, UNIT_COMPANY_RVP, UNIT_COMPANY_TWO, UNIT_COMPANY_RECONNAISSANCE, UNIT_BATTALION, UNIT_COMPANY_UAV, COMMANDER_POSADA, CHIEF_OF_STAFF_POSADA, COMMANDER_TITLE, CHIEF_OF_STAFF_TITLE
from get_normalized_data.get_paintball_data import classify_attacks
from helpers import clean_text, merge_and_sort_reports, paragraph_five_dot_one_table, paragraph_five_dot_two_table, highlight_changes_from_previous_report, parse_date_key, print_green, set_margins, set_line_spacing, add_text_for_header_and_footer_document, add_paragraph_with_style, add_custom_heading, add_paragraph_with_parts, add_custom_heading_with_parts, mark_no_auto_highlight, get_signature_officer
from get_normalized_data.get_UAV_data import calculate_uav_data
from get_normalized_data.get_basketball_data import calculate_basketball_data
from get_normalized_data.get_active_data import calculate
from get_normalized_data.get_previous_report_data import get_previous_report, get_previous_report_number, get_previous_personnel_totals, get_previous_ovt_totals
from get_normalized_data.get_bk_data import resolve_bk_synonyms
from get_data import get_static_data_json

# P/H - відступ 1.2см за замовчуванням; no_highlight=True виключає абзац із
# загального порівняння з учорашнім звітом (для статичних заголовків/заглушок)
def P(doc, text="", no_highlight=False, **kwargs):
    kwargs.setdefault("first_line_indent", 1.2)
    paragraph = add_paragraph_with_style(doc, text, **kwargs)
    return mark_no_auto_highlight(doc, paragraph) if no_highlight else paragraph

def H(doc, text, no_highlight=False, **kwargs):
    kwargs.setdefault("alignment", WD_ALIGN_PARAGRAPH.LEFT)
    kwargs.setdefault("first_line_indent", 1.2)
    paragraph = add_custom_heading(doc, text, **kwargs)
    return mark_no_auto_highlight(doc, paragraph) if no_highlight else paragraph

# P_parts/H_parts - абзац/заголовок із частин (text, highlight) для рядків, де
# підсвічувати треба лише конкретний токен (дату, номер, лічильник), не весь рядок
def P_parts(doc, parts, **kwargs):
    kwargs.setdefault("first_line_indent", 1.2)
    return add_paragraph_with_parts(doc, parts, **kwargs)

def H_parts(doc, parts, **kwargs):
    kwargs.setdefault("alignment", WD_ALIGN_PARAGRAPH.LEFT)
    kwargs.setdefault("first_line_indent", 1.2)
    return add_custom_heading_with_parts(doc, parts, **kwargs)

def blank(doc):
    return add_paragraph_with_style(doc, "")

def count_part(n):
    """(текст, підсвітити) - лічильник підсвічується лише коли не нульовий."""
    return (str(n), n != 0)

def static_subsection(doc, heading_text, heading_level, lines=()):
    """Підрозділ із повністю статичним вмістом (0-заглушки без джерела даних) -
    завжди виключений з підсвічування."""
    H(doc, heading_text, heading=heading_level, no_highlight=True)
    for line in lines:
        P(doc, line, no_highlight=True)
    blank(doc)

def document_final_combat_report(selected_date, hour_of_report, data_paintball, data_basketball, data_uav, data_active, data_commission, rows_with_data):
    doc = Document()
    static_data = get_static_data_json(FILE_PATH_STATIC_DATA)

    previous_date = (datetime.strptime(selected_date, "%d.%m.%Y") - timedelta(days=1)).strftime("%d.%m.%Y")
    previous_doc = get_previous_report(previous_date)
    report_number = (get_previous_report_number(previous_doc) or 0) + 1

    # синоніми з "Витрата:" (напр. "свп осколок") замінюємо на офіційну назву БК
    # з таблиці ще ДО побудови розділів 3.3.5/6, щоб обидва бачили вже змінений
    # текст/витрату. Базові залишки - з учорашньої (previous_date) таблиці, яка
    # сама не змінюється, тож повторна генерація сьогоднішнього звіту завжди
    # перераховує з нуля (а не списує ще раз поверх учорашнього прогону)
    resolve_bk_synonyms(data_uav['cost_data'], previous_date, selected_date)

    generate_head_documents(doc, selected_date, hour_of_report, static_data, report_number)
    paragraph_one(doc, static_data)
    paragraph_two(doc, static_data)
    paragraph_three(doc, selected_date, data_paintball, data_basketball, data_active, data_uav['cost_data'], data_uav['logistics_data'], data_uav['loss_data'], data_commission)
    paragraph_four(doc)
    paragraph_five(doc, selected_date, hour_of_report, previous_date, previous_doc)
    paragraph_six(doc, data_basketball, data_active, data_uav['cost_data'])
    paragraph_seven(doc, static_data)

    paragraph_footer(doc, rows_with_data)

    highlight_changes_from_previous_report(doc, previous_doc)

    safe_file_name = REPORT_FILE_NAME_TEMPLATE.format(date=selected_date)
    full_path = os.path.join(OUTPUT_DIR, safe_file_name)
    doc.save(full_path)
    print_green(f"- {safe_file_name}")

    open_generated_file(full_path)

def open_generated_file(full_path):
    try:
        os.startfile(full_path)
    except Exception as e:
        print(f"Не вдалося автоматично відкрити файл {full_path}: {e}")

def generate_head_documents(doc, selected_date, hour_of_report, static_data, report_number):
    set_margins(doc, Inches(0.786), Inches(0.786), Inches(1.18), Inches(0.395))
    set_line_spacing(doc, 1.0)  # встановлення міжрядкового інтервалу 1
    add_text_for_header_and_footer_document(doc)
    add_paragraph_with_style(doc, "Додаток 1", font_size=12, alignment=WD_ALIGN_PARAGRAPH.RIGHT)
    add_paragraph_with_style(doc, "ФОРМА № 5.18/КЦ", font_size=12, alignment=WD_ALIGN_PARAGRAPH.RIGHT)
    add_paragraph_with_style(doc, "Для службового користування", font_size=12, alignment=WD_ALIGN_PARAGRAPH.RIGHT)
    add_paragraph_with_style(doc, "Прим. № 1", font_size=12, alignment=WD_ALIGN_PARAGRAPH.RIGHT)
    blank(doc)
    blank(doc)
    P(doc, f"КОМАНДИРУ {UNIT_BRIGADE} (через командний пункт)")
    blank(doc)
    city = static_data.get('location_command_post', {}).get('city', '')
    P_parts(doc, [
        (f"ПІДСУМКОВЕ БОЙОВЕ ДОНЕСЕННЯ командира {UNIT_BATTALION} {UNIT_BRIGADE} №", False),
        (str(report_number), True),
        (f"\nКСП – {city} {hour_of_report}.00 ", False),
        (selected_date, True),
        (". Карта 25 000, видання 2024 року.", False),
    ])
    blank(doc)

# 0-заглушки без джерела реальних даних. bold: True - весь рядок, список слів - лише вони
_PARAGRAPH_ONE_LOSSES = [
    ("Орієнтовні втрати противника протягом доби (особовий склад, озброєння та військова техніка, матеріально-технічні засоби):", True),
    ("особового складу – 0, з них: ", False),
    (f"безповоротні – 0 з них: 0 — ({UNIT_COMPANY_UAV}).; 0 — ({UNIT_COMPANY_DSHR}).", False),
    (f"санітарні – 0 з них: 0 - ({UNIT_COMPANY_UAV}).", False),
    ("Полон – 0 ;", False),
    ("ОВТ – 0 од., з них: знищено – 0 од. з них: 0;", ['ОВТ']),
    ("пошкоджено – 0 од.:", False),
    ("танків – 0 од., з них: знищено – 0 од., пошкоджено– 0;", False),
    ("ББМ – 0 од., з них: знищено – 0 од., пошкоджено – 0 од.", False),
    ("РСЗВ – 0 од., з них: знищено – 0 од., пошкоджено – 0 од;", False),
    ("ГіМ – 0 од., з них: знищено – 0 од. пошкоджено – 0 од;", False),
    ("ПТ засобів – 0 од., з них: знищено – 0 од., пошкоджено – 0 од;", False),
    ("засоби ППО – 0 од., з них: знищено – 0 од., пошкоджено – 0 од. ", False),
    ("АТ – 0 од., з них: знищено – 0 од., пошкоджено – 0 од.;", False),
    ("спеціальна техніка – 0 од., з них: знищено – 0 од., пошкоджено – 0.", False),
    ("технічні засоби розвідки – 0 од., з них: знищено – 0 од., пошкоджено – 0 од.", False),
    ("Легкі транспортні засоби — 0 од., з них: знищено — од., пошкоджено — 0 од.", False),
    ("БпЛА – 0 од., з них: по типам: баражуючий боєприпас – 0 од., FPV-дрон – 0 од.", False),
    ("ПУ БпЛА – 0, з них: знищено – 0, пошкоджено – 0.", ['ПУ БпЛА']),
    ("пункти управління / укриття – 0 / 0", False),
    ("склади боєприпасів / ПММ – 0/ 0", False),
]

def paragraph_one(doc, static_data):
    H(doc, "1. ВИСНОВКИ З ОЦІНКИ ПРОТИВНИКА", no_highlight=True)
    blank(doc)
    P(doc, f"Згідно з розвідувальним зведенням {static_data.get('paragraph_one', '')}.")
    blank(doc)
    for text, bold in _PARAGRAPH_ONE_LOSSES:
        P(doc, text, bold=bold, no_highlight=True)
    blank(doc)

def paragraph_two(doc, static_data):
    H(doc, "2. ПОЛОЖЕННЯ ТА СТАН ПІДРОЗДІЛІВ НАШИХ ВІЙСЬК", no_highlight=True)
    blank(doc)
    for entry in static_data.get('paragraph_two', []):
        # запис або звичайний рядок, або [текст, список фраз для жирного
        # виділення] - як у довіднику (назви підрозділів, позивні, заголовки)
        text, bold_words = entry if isinstance(entry, list) else (entry, [])
        P(doc, text, bold=bold_words)

def paragraph_three(doc, selected_date, data_paintball, data_basketball, data_active, data_uav, logistics_data, data_uav_loss, data_commission):
    H(doc, "3. ХІД ВЕДЕННЯ БОЙОВИХ ДІЙ ТА ВИКОНАННЯ СПЛАНОВАНИХ ЗАВДАНЬ.", no_highlight=True)
    paragraph_assault_actions(doc) # штурмові дії 3.1.1
    paragraph_shelling(doc, data_paintball) # обстріли 3.1.2 /виконано
    paragraph_combat_work_of_units(doc, data_active) # бойова робота підрозділів батальйону 3.2 / виконано
    paragraph_progress_of_planned_tasks(doc, logistics_data, data_uav_loss) # хід виконання спланованих завдань 3.3
    paragraph_layered_mining(doc) # поширене мінування 3.3.1
    paragraph_restoring_previously_lost_position(doc) # відновлення раніше втрачене положення 3.3.2
    paragraph_fire_damage_aviation(doc) # вогневе ураження авіацією 3.3.3
    paragraph_air_defense_and_electronic_warfare(doc) # застосування ППО 3.3.4
    paragraph_combat_work_of_uav(doc, data_uav) # Застосування БпАК (дронів камікадзе та роботи НРК) 3.3.5
    paragraph_artillery_task(doc, data_basketball) # артилерійські завдання 3.3.6
    paragraph_electronic_warfare(doc) # силами та засобами РЕБ 3.3.7
    paragraph_working_and_inspection_groups(doc, data_commission) # робота комісій, робочих та інспекційних груп 3.4
    paragraph_emergency_events(doc) # надзвичайні події 3.5

def paragraph_assault_actions(doc):
    static_subsection(doc, "3.1.1. Противник 0 разів проводив штурмові (наступальні) дії:", 3)

def paragraph_shelling(doc, data_paintball):
    s = classify_attacks(data_paintball)
    total = s['всього']['count']

    H_parts(doc, [
        ("3.1.2. Противник здійснив ", False),
        count_part(total),
        (" обстрілів:", False),
    ], heading=3)

    # (назва, ключ, роздільник перед числом, роздільник ПІСЛЯ цього поля) - і
    # роздільники перед числом, і роздільники між полями різні за шаблоном
    # реального донесення (десь "; ", десь ", ", в кінці "., " перед останнім
    # полем і "." в самому кінці) - тому обидва задаються явно для кожного поля.
    summary_fields = [
        ("авіаційних ударів", 'авіаційних_ударів', " - ", "; "),
        ("артилерійських обстрілів", 'артилерійських_обстрілів', " - ", ", "),
        ("РСЗВ", 'рсзв', " – ", ", "),
        ("гранатометів", 'гранатомети', " – ", ", "),
        ("скиди з БпЛА", 'скиди_бпла', " — ", ", "),
        ("удари дронів камікадзе (FPV, Молнія, FPV)", 'удари_FPV_дронів', " – ", ", "),
        ("БМП", 'БМП', " – ", ", "),
        ("танковий", 'танковий', " – ", ", "),
        ("ТОС", 'ТОС', " – ", ", "),
        ("БПЛА ЛАНЦЕТ (баражуючий боєприпас)", 'БПЛА_Крило_ЛАНЦЕТ', " — ", ", "),
        ("снайпер", 'снайпер', " - ", ", "),
        ("стріл. зброя", 'стріл_зброя', " – ", "; "),
        ("фосфор", 'фосфор', " - ", "; "),
        ("гранати", 'гранати', " - ", "., "),
        ("підрив на СВП", 'Підриви', " - ", "."),
    ]
    # у підсумковому рядку - лише число підсвічується, і тільки якщо воно не нульове
    summary_parts = []
    for label, key, sep, trailing in summary_fields:
        summary_parts.append((f"{label}{sep}", False))
        summary_parts.append(count_part(s[key]['count']))
        summary_parts.append((trailing, False))
    P_parts(doc, summary_parts)

    # нижче завжди показуємо лише ці 6 категорій (навіть з нульовим значенням) - решта
    # рахується тільки в підсумковому рядку вище, без окремого блоку повідомлень
    detail_categories = [
        ("Ракетних ударів", s['ракетних_ударів']),
        ("Авіаційних ударів", s['авіаційних_ударів']),
        ("Артилерійський обстріл", s['артилерійських_обстрілів']),
        ("Ударів дронів камікадзе", s['удари_FPV_дронів']),
        ("Здійснення скидів з БпЛА", s['скиди_бпла']),
        ("Підрив на СВП", s['Підриви']),
    ]
    for label, cat in detail_categories:
        count = cat['count']
        P_parts(doc, [(f"{label} – ", False), count_part(count), (":", False)], bold=True)
        for msg in cat["messages"]:
            P(doc, f"{clean_text(msg)}")
        blank(doc)
        
def paragraph_combat_work_of_units(doc, data_active):
    H(doc, "3.2. Бойова робота підрозділів батальйону", heading=3, no_highlight=True)
    for msg in data_active:
        P(doc, f"{msg['text']}")
    blank(doc)

def paragraph_progress_of_planned_tasks(doc, logistics_data, data_uav_loss):
    H(doc, "3.3 Хід виконання спланованих завдань", heading=3, no_highlight=True)
    blank(doc)
    merged_data = merge_and_sort_reports(logistics_data, data_uav_loss)
    vylyoty = len([d for d in merged_data if 'виліт' in d.get('text', '').lower()])
    viyizdy = len([d for d in merged_data if 'виїз' in d.get('text', '').lower()])
    delivered = len([d for d in merged_data if 'доставлен' in d.get('text', '').lower()])
    P_parts(doc, [
        ("Всього ", False), count_part(vylyoty), (" вильотів, всього ", False),
        count_part(viyizdy), (" виїзд (з них ", False),
        count_part(delivered), (" доставлено)", False),
    ], bold=True)
    for d in merged_data:
        P(doc, clean_text(d.get("text", '')))
    blank(doc)

def paragraph_layered_mining(doc):
    static_subsection(
        doc,
        "3.3.1. Інформація щодо стану виконання заходів проведення суцільного ешелонованого мінування.",
        4,
        [
            "Перший рубіж",
            "обладнано (прокопано) траншей – 0 м;",
            "обладнано вогневих комірок – 0 од.;",
            "обладнано перекритих ділянок траншей – 0 од.;",
            "обладнано споруд для ведення вогню – 0 од.;",
            "обладнано споруд для ведення спостереження – 0 од.;",
            "обладнано бліндажів – 0 од.;",
            "Другий рубіж",
            "викопано протитанкового рову – 0 м;",
            "обладнано (прокопано) траншей – 0 м;",
            "перекрито ділянок траншей – 0 шт.;",
            "обшито траншей (геотекстиль/сітка) – 0 м;",
            "облаштовано брустверу – 0 м;",
            "виготовлено щитів для обшивки траншей (геотекстиль/сітка) – 0 шт.;",
            "накрито вогневих комірок гофрованою сталлю – 0 шт.;",
            "перекрито вогневих комірок дерев’яним накриттям – 0 шт.",
        ],
    )

def paragraph_restoring_previously_lost_position(doc):
    static_subsection(
        doc,
        "3.3.2. Виконання заходів відновлення раніше втрачене положення та боєздатності елементів бойового порядку.",
        4,
        ["Заходи відновлення боєздатності не проводяться."],
    )

def paragraph_fire_damage_aviation(doc):
    static_subsection(doc, "3.3.3. Авіацією УВ нанесено вогневе ураження по противнику:", 4,
                       ["Тактична, армійська авіація не застосовувалася."])

def paragraph_air_defense_and_electronic_warfare(doc):
    static_subsection(doc, "3.3.4. Застосування ППО:", 4)

def paragraph_electronic_warfare(doc):
    static_subsection(doc, "3.3.7 Силами та засобами РЕБ:", 3, [
        "здійснено протидію 0 БпЛА противника:",
        "зрив польотного завдання – 0 од.;",
        "посаджено – 0 од;",
    ])

def paragraph_working_and_inspection_groups(doc, data_commission):
    H(doc, "3.4. Робота комісій, робочих та інспекційних груп.", heading=3, no_highlight=True)
    blank(doc)
    for t in data_commission:
        P(doc, t)

def paragraph_emergency_events(doc):
    static_subsection(doc, "3.5. Надзвичайні події.", 3)

def paragraph_combat_work_of_uav(doc, data_uav):
    H(doc, "3.3.5 Застосування БпС (БпЛА, дронів камікадзе та роботи НРК):", heading=3, no_highlight=True)
    total = len(data_uav)
    hit = len([d for d in data_uav if ('уражен' in d.get('text', '').lower() or 'обстріл' in d.get('text', '').lower()) and 'не уражен' not in d.get('text', '').lower()])
    P_parts(doc, [
        ("Всього ", False), count_part(total), (" вильотів (з них ", False),
        count_part(hit), (" в ціль)", False),
    ], bold=True)
    for d in sorted(data_uav, key=parse_date_key):
        P(doc, clean_text(d.get("text", "")))

def paragraph_artillery_task(doc, data_basketball):
    H(doc, "3.3.6 Застосування артилерії.", heading=3, no_highlight=True)
    blank(doc)
    for d in data_basketball:
        P(doc, f"{clean_text(d.get('text', ''))}")
    blank(doc)

def paragraph_four(doc):
    H(doc, "4. ОСНОВНІ ЗАВДАННЯ, ЯКІ ПЛАНУЮТЬСЯ НА НАСТУПНУ ДОБУ.", no_highlight=True)
    blank(doc)
    P(doc, "")
    blank(doc)

def paragraph_five(doc, selected_date, hour_of_report, previous_date, previous_doc):
    H(doc, f"5. ВТРАТИ о/с підрозділів {UNIT_BATTALION}: ", no_highlight=True)
    blank(doc)
    paragraph_five_dot_one(doc, selected_date, hour_of_report, previous_date, previous_doc)
    paragraph_five_dot_two(doc, selected_date, previous_date, previous_doc)

def paragraph_five_dot_one(doc, selected_date, hour_of_report, previous_date, previous_doc):
    today = selected_date
    units_data = get_previous_personnel_totals(previous_doc)
    H(doc, "5.1. ВІДОМОСТІ ПРО БЕЗПОВОРОТНІ ТА САНІТАРНІ ВТРАТИ ОСОБОВОГО СКЛАДУ", heading=3, no_highlight=True)
    paragraph_five_dot_one_table(doc, previous_date, today, hour_of_report, units_data)
    blank(doc)
    blank(doc)

def paragraph_five_dot_two(doc, selected_date, previous_date, previous_doc):
    today = selected_date
    # звітний цикл ОВТ у реальних донесеннях фіксований - завжди 19:00-19:00,
    # незалежно від обраної hour_of_report (перевірено на попередніх звітах)
    ovt_data = get_previous_ovt_totals(previous_doc)
    H(doc, "5.2. ВТРАТИ ОВТ.", heading=3, no_highlight=True)
    paragraph_five_dot_two_table(doc, previous_date, today, 19, ovt_data)
    blank(doc)

def paragraph_six(doc, data_basketball, data_active, data_uav):
    H(doc, "6. ВИТРАТИ БОЄПРИПАСІВ, АВІАЦІЙНИХ ЗАСОБІВ УРАЖЕННЯ ТА АРТИЛЕРІЙСЬКИХ БОЄПРИПАСІВ", no_highlight=True)
    blank(doc)
    P(doc, f"ВБпАК: {calculate_uav_data(data_uav)}")
    P(doc, f"мінбатр: {calculate_basketball_data(data_basketball)} {calculate(data_active, UNIT_COMPANY_ARTILLERY)}")
    for unit in [UNIT_COMPANY_ONE, UNIT_COMPANY_TWO, UNIT_COMPANY_DSHR, UNIT_COMPANY_RVP, UNIT_COMPANY_RECONNAISSANCE]:
        P(doc, f"{unit}: {calculate(data_active, unit)}")
    blank(doc)

def paragraph_seven(doc, static_data):
    H(doc, "7. ПРОБЛЕМНІ ПИТАННЯ", no_highlight=True)
    for d in static_data.get('paragraph_seven', []):
        P(doc, f"{d}")
    blank(doc)
    blank(doc)

def paragraph_footer(doc, rows_with_data):
    signatories = [(COMMANDER_POSADA, COMMANDER_TITLE), (CHIEF_OF_STAFF_POSADA, CHIEF_OF_STAFF_TITLE)]
    blank(doc)
    for i, (posada, title) in enumerate(signatories):
        if i > 0:
            blank(doc)
        rank, name = get_signature_officer(rows_with_data, posada)
        add_paragraph_with_style(doc, title)
        add_paragraph_with_style(doc, f"{rank}\t{name}", format_tabs=True)
        