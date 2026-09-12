import os
import re
from docx import Document
from constants import UNIT_BATTALION, UNIT_BRIGADE, RESOURCES_DIR, REPORT_FILE_NAME_TEMPLATE

_REPORT_NUMBER_RE = re.compile(rf"{re.escape(UNIT_BATTALION)}\s+{re.escape(UNIT_BRIGADE)}\s*№\s*(\d+)")


def _report_file_path(date_str):
    return os.path.join(RESOURCES_DIR, REPORT_FILE_NAME_TEMPLATE.format(date=date_str))


def _to_int(text):
    text = (text or "").strip()
    return int(text) if text.isdigit() else 0


def get_previous_report(previous_date):
    """Повертає раніше згенерований документ (python-docx Document) за previous_date
    з папки resources/, або None якщо такого звіту немає."""
    path = _report_file_path(previous_date)
    if not os.path.isfile(path):
        return None
    try:
        return Document(path)
    except Exception:
        return None


def get_previous_report_number(previous_doc):
    """Номер вчорашнього ПБД зі шапки документа (напр. "...{UNIT_BRIGADE} №76" -> 76)."""
    if previous_doc is None:
        return None
    for p in previous_doc.paragraphs[:15]:
        m = _REPORT_NUMBER_RE.search(p.text)
        if m:
            return int(m.group(1))
    return None


def get_previous_personnel_totals(previous_doc):
    """Кумулятивні втрати о/с ("з початком включення") з таблиці 5.1 попереднього звіту."""
    if previous_doc is None or len(previous_doc.tables) < 1:
        return {}

    table = previous_doc.tables[0]
    keys = ["безповоротні_заг", "санітарні_заг", "зниклі_заг", "полонені_заг"]
    result = {}

    for row in table.rows[2:]:
        cells = [c.text.strip() for c in row.cells]
        if len(cells) < 9:
            continue
        label = cells[0]
        if not label or label == "Придані підрозділи":
            continue
        result[label] = {key: _to_int(cells[5 + i]) for i, key in enumerate(keys)}

    return result


def get_previous_ovt_totals(previous_doc):
    """Кумулятивні втрати ОВТ ("з початком включення") з таблиці 5.2 попереднього звіту,
    включно з підсумковим рядком "Всього ОВТ за {UNIT_BATTALION}" (переноситься як є, без перерахунку)."""
    if previous_doc is None or len(previous_doc.tables) < 2:
        return {}

    table = previous_doc.tables[1]
    keys = ["знищено_заг", "пошкоджено_заг", "всього_заг"]
    result = {}

    for row in table.rows[2:]:
        cells = [c.text.strip() for c in row.cells]
        if len(cells) < 7:
            continue
        label = cells[0]
        if not label:
            continue
        result[label] = {key: _to_int(cells[4 + i]) for i, key in enumerate(keys)}

    return result
