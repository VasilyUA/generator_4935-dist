import os
import re
import openpyxl

from constants import BK_UAV_FILE_PATH, BK_SYNONYM_CREW_OVERRIDES, UAV_MUNITION_DISPLAY_NAMES
from helpers import clean_text

# фіксовані колонки A-D; кількість - НЕ фіксована колонка, а окремий стовпець
# на кожну дату ("кількість\nза {дата}"), що додається праворуч від попередніх
_ID_COL, _SYNONYMS_COL, _UAV_TYPE_COL, _ITEM_COL = 1, 2, 3, 4

# нормалізовані офіційні назви, що get_UAV_data.py вже сам зіставив (напр.
# "КО Пузатий змій" -> 'КО 1,3 "Пузатий змій"') ДО того, як cost_data взагалі
# дійшло сюди - без цього застереження resolve_bk_synonyms() перезаписував би
# уже коректну назву на СИРИЙ запис із колонки D таблиці (з небажаним
# префіксом "Боєприпас" і зламаним регістром при подальшому форматуванні)

# заголовок стовпця-БАЗИ без дати ("кількість.") - реальний файл БК ВБАК.xlsx
# спочатку заповнюється ОДНИМ таким стовпцем (початкові залишки), і лише З
# ПЕРШИМ запуском resolve_bk_synonyms() з'являється перший датований стовпець.
# Без цього фолбека - якщо для previous_date ще нема датованого стовпця -
# функція мовчки нічого не робила б НАЗАВЖДИ (нізвідки узяти перший датований
# стовпець).
_BASE_QTY_HEADER = "кількість."


def _qty_header(date_str):
    return f"кількість\nза {date_str}"


def _normalize_item_name(name):
    # clean_text вирівнює лапки (“”/«» -> ") й пробіли - каталог БК і назви
    # з повідомлень (вже пропущені через clean_text) інакше не співпадуть
    return clean_text(str(name or "")).lower()


_ALREADY_DISPLAY_NAMES_NORM = {_normalize_item_name(v) for v in UAV_MUNITION_DISPLAY_NAMES.values()}


def _is_already_resolved_display_name(name_norm):
    return name_norm in _ALREADY_DISPLAY_NAMES_NORM


def _open_sheet():
    """Відкриває аркуш БК ВБАК. Повертає (workbook, worksheet) або (None, None),
    якщо файлу немає."""
    if not os.path.isfile(BK_UAV_FILE_PATH):
        return None, None
    wb = openpyxl.load_workbook(BK_UAV_FILE_PATH)
    return wb, wb.active


def _find_qty_column(ws, date_str):
    """Номер (1-based) стовпця "кількість\\nза date_str" у шапці, або None."""
    target = _qty_header(date_str).strip().lower()
    for cell in ws[1]:
        value = cell.value
        if isinstance(value, str) and value.strip().lower() == target:
            return cell.column
    return None


def _find_base_qty_column(ws):
    """Номер (1-based) стовпця-БАЗИ без дати (_BASE_QTY_HEADER), або None."""
    for cell in ws[1]:
        value = cell.value
        if isinstance(value, str) and value.strip().lower() == _BASE_QTY_HEADER:
            return cell.column
    return None


def _read_rows(ws, qty_col):
    rows = []
    for row_idx in range(2, ws.max_row + 1):
        item = ws.cell(row=row_idx, column=_ITEM_COL).value
        if not item:
            continue
        rows.append({
            "row_idx": row_idx,
            "id": ws.cell(row=row_idx, column=_ID_COL).value,
            "synonyms": ws.cell(row=row_idx, column=_SYNONYMS_COL).value,
            "uav_type": ws.cell(row=row_idx, column=_UAV_TYPE_COL).value,
            "item": item,
            "qty": int(ws.cell(row=row_idx, column=qty_col).value or 0),
        })
    return rows


def read_bk_stock(date_str):
    """Читає рядки БК ВБАК зі стовпцем кількості за date_str (для тестів й
    зовнішнього огляду). Кожен рядок - dict з id/synonyms/uav_type/item/qty.
    Якщо файлу чи стовпця за цю дату немає - []."""
    wb, ws = _open_sheet()
    if ws is None:
        return []
    col = _find_qty_column(ws, date_str)
    if col is None:
        return []
    return _read_rows(ws, col)


def _split_words(value):
    return [_normalize_item_name(w) for w in str(value or "").replace("\n", ",").split(",") if w.strip()]


_TRAILING_VARIANT_RE = re.compile(r"-\d+$")


def _strip_variant_suffix(word):
    """Прибирає суфікс варіанта на кшталт "-1" ("огб-1" -> "огб") - у
    повідомленні часто називають виріб без номера варіанта. Не чіпаємо, якщо
    залишок закороткий (2 символи й менше) - інакше короткі синоніми штибу
    "ф-1" звелися б до "ф", яке збігається з будь-яким словом на цю літеру."""
    stripped = _TRAILING_VARIANT_RE.sub("", word)
    return stripped if len(stripped) > 2 and stripped != word else None


def _name_matches_synonym(name_norm, synonyms):
    """Порівнює назву з витрати (name_norm) з переліком синонімів колонки B:
    багатослівний синонім ("свп осколок") - як підрядок-фраза, однослівний
    ("уламок", "ф") - як ОКРЕМЕ слово в назві (щоб короткі синоніми на кшталт
    "ф" не збігалися з будь-яким словом, що просто містить цю літеру), і
    додатково без суфікса варіанта ("огб-1" збігається і з "огб", і з "огб-1").
    Перше ловить складені назви на кшталт "СВП Уламок", де жоден синонім не
    дорівнює всій назві цілком, але одне з її слів - точний синонім."""
    name_words = name_norm.split()
    for syn in synonyms:
        if " " in syn:
            if syn in name_norm:
                return True
            continue
        if syn in name_words:
            return True
        stripped = _strip_variant_suffix(syn)
        if stripped is not None and stripped in name_words:
            return True
    return False


_PURE_NUMBER_RE = re.compile(r'^[\d,.\-]+$')


def _significant_tokens(name_norm):
    """Слова назви без коротких (1 символ) і суто числових токенів ("1,3") -
    саме вони мають однозначно вказувати на конкретний виріб."""
    tokens = name_norm.replace('"', ' ').split()
    return [t for t in tokens if len(t) > 1 and not _PURE_NUMBER_RE.match(t)]


def _matches_official_item(name_norm, item_norm):
    """Оператори іноді пишуть офіційну назву (колонка D) неформально й не
    повністю - без слова "Боєприпас", з незакритою лапкою, чи в іншому порядку
    слів (напр. "КО "Пузатий змій 1,3" замість "Боєприпас КО 1,3 «Пузатий
    змій»") - тому звіряємо не підрядком, а набором значущих слів: усі значущі
    слова з витрати мають зустрічатись десь у офіційній назві."""
    tokens = _significant_tokens(name_norm)
    if not tokens:
        return False
    item_tokens = set(item_norm.replace('"', ' ').split())
    return all(t in item_tokens for t in tokens)


def _find_replacement_row(rows, name_norm, text_lower):
    """Рядки, де є синонім name_norm (колонка B) АБО сама назва - неформальний/
    неповний варіант офіційної назви (колонка D), і десь у тексті повідомлення
    згадано тип БпЛА (колонка C), з-поміж них - з найбільшим залишком."""
    candidates = [
        row for row in rows
        if row["qty"] > 0
        and (
            _name_matches_synonym(name_norm, _split_words(row["synonyms"]))
            or _matches_official_item(name_norm, _normalize_item_name(row["item"]))
        )
        and any(t in text_lower for t in _split_words(row["uav_type"]))
    ]
    return max(candidates, key=lambda r: r["qty"]) if candidates else None


def _find_row_by_official_item(rows, item_name_norm):
    for row in rows:
        if _normalize_item_name(row["item"]) == item_name_norm:
            return row
    return None


# сентинел відрізняє "override сказав явно НЕ чіпати" (null у ресурс-файлі)
# від "жодного override не налаштовано" (None з _find_row_by_official_item) -
# обидва не мають опрацьовуватись однаково: перший НЕ повинен падати назад
# на звичайний пошук за залишком, другий - повинен.
_SKIP = object()


def _resolve_row_for_spend_item(rows, name_norm, text_lower):
    """BK_SYNONYM_CREW_OVERRIDES перевіряється ПЕРШИМ, ще до звичайного
    _find_replacement_row(): коли той самий синонім однаково збігається з
    десятками рядків (усі з ідентичним широким списком синонімів колонки B),
    вибір "найбільший залишок" не обов'язково дає виріб, який реально
    використовує САМЕ цей екіпаж (чи взагалі не мав би нічого міняти). Ключ
    overrides - (екіпаж, синонім); екіпаж шукається підрядком у тексті
    повідомлення. Повертає рядок таблиці, _SKIP (явно не резолвити), або None
    (нема ні override, ні звичайного збігу)."""
    for (crew, synonym), official_item in BK_SYNONYM_CREW_OVERRIDES.items():
        if crew not in text_lower or synonym not in name_norm:
            continue
        if official_item is None:
            return _SKIP
        row = _find_row_by_official_item(rows, _normalize_item_name(official_item))
        if row is not None and row["qty"] > 0:
            return row
        break  # override налаштовано, але ціль недоступна - звичайний пошук нижче

    return _find_replacement_row(rows, name_norm, text_lower)


def resolve_bk_synonyms(cost_data, previous_date, selected_date):
    """Для кожного повідомлення БпАК: якщо у витраті згадано синонім з колонки B
    ("synonyms") і десь у тексті повідомлення - тип БпЛА з колонки C ("uav_type"),
    замінює цей синонім у тексті й у витраті на офіційну назву БК з колонки D
    ("item") - береться рядок із найбільшим залишком серед тих, що підходять.
    Якщо серед відповідних рядків немає жодного із залишком > 0 - ні заміни
    (текст лишається як є). Мутує cost_data на місці.

    Файл БК ВБАК - один (без дати в імені); кожен день зберігається як окремий
    стовпець "кількість\\nза {дата}", що додається праворуч. Стовпець за
    previous_date - незмінна база для розрахунку. Стовпець за selected_date
    щоразу ПЕРЕСТВОРЮЄТЬСЯ заново (видаляється, якщо вже є, і рахується з нуля
    від previous_date) - тож повторна генерація того самого дня завжди дає той
    самий коректний результат, а не списує ще раз поверх учорашнього прогону.
    Якщо для previous_date ще нема датованого стовпця - базою стає стовпець-
    БАЗА без дати (_BASE_QTY_HEADER, як спочатку заповнений весь файл); це
    трапляється лише на САМОМУ ПЕРШОМУ запуску - далі завжди є датований
    стовпець за вчора. Якщо немає ні файлу, ні датованого стовпця, ні бази -
    нічого не робить."""
    wb, ws = _open_sheet()
    if ws is None:
        return
    prev_col = _find_qty_column(ws, previous_date) or _find_base_qty_column(ws)
    if prev_col is None:
        return
    rows = _read_rows(ws, prev_col)

    for d in cost_data:
        spend = d.get("spend")
        if not spend:
            continue
        text = d.get("text", "")
        text_lower = text.lower()
        new_spend = {}
        for name, qty in spend.items():
            name_norm = _normalize_item_name(name)
            if _is_already_resolved_display_name(name_norm):
                row = None
            else:
                row = _resolve_row_for_spend_item(rows, name_norm, text_lower)
            if row is None or row is _SKIP:
                new_spend[name] = new_spend.get(name, 0) + qty
                continue
            text = text.replace(f"{name} - {qty} шт", f"{row['item']} - {qty} шт")
            row["qty"] = max(0, row["qty"] - qty)
            new_spend[row["item"]] = new_spend.get(row["item"], 0) + qty
        d["text"] = text
        d["spend"] = new_spend

    today_col = _find_qty_column(ws, selected_date)
    if today_col is not None:
        ws.delete_cols(today_col)

    new_col = ws.max_column + 1
    ws.cell(row=1, column=new_col, value=_qty_header(selected_date))
    for row in rows:
        ws.cell(row=row["row_idx"], column=new_col, value=row["qty"])

    wb.save(BK_UAV_FILE_PATH)
