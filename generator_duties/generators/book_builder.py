"""Книга обов'язків - однакова за будовою для функціональних і посадових
(за отриманим зразком): обкладинка → ЗМІСТ → витяг зі Статуту → Положення →
блок на кожну унікальну посаду → підпис комбата → лист ознайомлення (усі
слоти штату по порядку) → підпис → напис про зшивання.

Оформлення власне (Times New Roman 14, поля як у зразку), а не скопійоване
зі зразка: у ньому десятки різних стилів і нумерацій з різних джерел.
Рядки ЗМІСТу - текст + поле PAGEREF на закладку заголовка: номери сторінок
рахує Word (generators/word_finalize.py)."""
import re
from dataclasses import dataclass, field

from docx import Document
from docx.enum.table import WD_ROW_HEIGHT_RULE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt

from constants import SECTION_DUTIES, SECTION_GENERAL, SECTION_HEADINGS, SECTIONS, SHEETS_MARKER, TO_FILL
from utils.docx_utils import (
    add_bookmark, add_page_break, add_paragraph_with_style, add_table, add_toc_entry, set_base_font, set_page,
    text_width,
)

INDENT_CM = 1.25
COVER_FONT_SIZE = 16
LIST_COLUMNS = ("№ з/п", "ВОС посади", "Посада", "Звання", "ПІБ", "Підпис в/с", "Примітка")
LIST_WIDTHS_CM = (1.1, 2.0, 4.1, 1.9, 3.3, 2.1, 2.5)
LIST_FONT_SIZE = 12
# Рядок листа ознайомлення не нижчий за це - щоб було де розписатися.
LIST_ROW_HEIGHT_CM = 0.8
STITCHING_TEXT = "Прошнуровано, пронумеровано та скріплено печаткою"
DATE_BLANK = "«____» ______________ 20____ року"
BLANK = "____________"
_ALREADY_LISTED_RE = re.compile(r"^([-–—•]|\d+[.)])")


@dataclass
class Signer:
    position: str = ""
    rank: str = ""
    name: str = ""


@dataclass
class BookInput:
    spec: dict
    year: int
    unit: dict
    signer: Signer
    positions: list
    duties: object
    extract: list = field(default_factory=list)
    regulation: list = field(default_factory=list)
    blank_rows: int = 0


@dataclass
class BookReview:
    to_fill: list = field(default_factory=list)
    missing: list = field(default_factory=list)
    without_vos: list = field(default_factory=list)


def _title_case_first_word(text):
    words = text.split(" ", 1)
    if words and words[0].isupper():
        words[0] = words[0][:1] + words[0][1:].lower()
    return " ".join(words)


def _specs_title(specs):
    """Заголовок вставленого документа для ЗМІСТу: перші центровані абзаци
    ("ПОЛОЖЕННЯ" + "про …") або перший абзац."""
    leading = []
    for spec in specs:
        if not spec.center:
            break
        leading.append(spec.text)
    return _title_case_first_word(" ".join(leading) if leading else specs[0].text)


def _bullet(text):
    return text if _ALREADY_LISTED_RE.match(text) else f"– {text}"


class _Bookmarks:
    def __init__(self):
        self.count = 0

    def next(self):
        self.count += 1
        return f"_RefDuty{self.count}", self.count


def _cover(doc, book):
    unit = book.unit
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    row = table.rows[0]
    row.height = Cm(25)
    row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
    cell = row.cells[0]
    cell.width = text_width(doc)
    lines = (
        ["", "Міністерство оборони України", unit.get("MILITARY_UNIT", "")]
        + [""] * 8
        + [(book.spec["cover_title"], True), (unit.get("BATTALION_GENITIVE", ""), True),
           (unit.get("MILITARY_UNIT_GENITIVE", ""), True)]
        + [""] * 12
        + [unit.get("SETTLEMENT", ""), f"{book.year} рік"]
    )
    for line in lines:
        text, bold = line if isinstance(line, tuple) else (line, False)
        add_paragraph_with_style(cell, text, font_size=COVER_FONT_SIZE, bold=bold, alignment=WD_ALIGN_PARAGRAPH.CENTER)


def _heading(doc, text, bookmarks, space_before=0):
    paragraph = add_paragraph_with_style(doc, text, bold=True, alignment=WD_ALIGN_PARAGRAPH.CENTER, heading_level=2,
                                         keep_with_next=True)
    paragraph.paragraph_format.space_before = Pt(space_before)
    name, bookmark_id = bookmarks
    add_bookmark(paragraph, name, bookmark_id)
    return paragraph


def _write_specs(doc, specs, bookmark):
    _heading(doc, specs[0].text, bookmark)
    for spec in specs[1:]:
        if spec.center:
            add_paragraph_with_style(doc, spec.text, bold=spec.bold, alignment=WD_ALIGN_PARAGRAPH.CENTER)
        else:
            text = _bullet(spec.text) if spec.listed else spec.text
            add_paragraph_with_style(doc, text, bold=spec.bold, first_line_indent_cm=INDENT_CM)


def _write_position(doc, position, duties, bookmark, review, first):
    _heading(doc, position.title, bookmark, space_before=0 if first else 12)
    if position.vos_codes:
        add_paragraph_with_style(doc, f"ВОС {', '.join(position.vos_codes)}", bold=True,
                                 alignment=WD_ALIGN_PARAGRAPH.CENTER, keep_with_next=True)
    else:
        review.without_vos.append(position.title)
    if not duties.has_position(position.key):
        review.missing.append(position.title)
    items_by_section = duties.for_position(position.key)
    for section in SECTIONS:
        items = items_by_section.get(section) or ([TO_FILL] if section == SECTION_DUTIES else [])
        if not items:
            continue
        add_paragraph_with_style(doc, SECTION_HEADINGS[section].format(position=position.title), bold=True,
                                 first_line_indent_cm=INDENT_CM, keep_with_next=True)
        for text in items:
            to_fill = text.strip().upper() == TO_FILL
            if to_fill:
                if position.title not in review.to_fill:
                    review.to_fill.append(position.title)
                text = f"{TO_FILL}: {section.lower()}"
            else:
                text = text.replace("{посада}", position.title)
                if section != SECTION_GENERAL:
                    text = _bullet(text)
            add_paragraph_with_style(doc, text, first_line_indent_cm=INDENT_CM, highlight=to_fill)


def _signature(doc, signer):
    add_paragraph_with_style(doc, "")
    add_paragraph_with_style(doc, signer.position or "Командир батальйону")
    add_paragraph_with_style(doc, f"{signer.rank or BLANK}\t{signer.name or BLANK}", tab_width=text_width(doc))


def _familiarization(doc, book, bookmark):
    _heading(doc, book.spec["list_title"], bookmark)
    add_paragraph_with_style(doc, "")
    slots = sorted((slot for position in book.positions for slot in position.slots), key=lambda slot: slot.row)
    rows = [(i, slot.vos, slot.display or slot.position, slot.rank, slot.pib, "", "")
            for i, slot in enumerate(slots, start=1)]
    # Порожні рядки - для в/с, що прийдуть пізніше (переміщення, донабір).
    rows += [(i,) for i in range(len(slots) + 1, len(slots) + book.blank_rows + 1)]
    add_table(doc, LIST_COLUMNS, rows, LIST_WIDTHS_CM, font_size=LIST_FONT_SIZE, min_row_height_cm=LIST_ROW_HEIGHT_CM)


def _stitching(doc, signer):
    add_paragraph_with_style(doc, STITCHING_TEXT, alignment=WD_ALIGN_PARAGRAPH.LEFT)
    add_paragraph_with_style(doc, SHEETS_MARKER, alignment=WD_ALIGN_PARAGRAPH.LEFT)
    _signature(doc, signer)
    add_paragraph_with_style(doc, DATE_BLANK, alignment=WD_ALIGN_PARAGRAPH.LEFT)


def build_book(book):
    """→ (Document, BookReview)."""
    doc = Document()
    set_page(doc, top=Cm(1.0), bottom=Cm(1.5), left=Cm(2.5), right=Cm(1.5))
    set_base_font(doc)
    review = BookReview()
    bookmarks = _Bookmarks()

    # Закладки й рядки ЗМІСТу - наперед, бо ЗМІСТ стоїть перед розділами.
    parts = []
    if book.extract:
        parts.append(("extract", book.extract, _specs_title(book.extract)))
    if book.regulation:
        parts.append(("regulation", book.regulation, _specs_title(book.regulation)))
    for position in book.positions:
        title = position.title + (f" ВОС {', '.join(position.vos_codes)}" if position.vos_codes else "")
        parts.append(("position", position, title))
    parts.append(("list", None, book.spec["list_title"]))
    marks = [bookmarks.next() for _ in parts]

    _cover(doc, book)
    add_page_break(doc)
    add_paragraph_with_style(doc, "ЗМІСТ", bold=True, alignment=WD_ALIGN_PARAGRAPH.CENTER)
    add_paragraph_with_style(doc, "")
    for number, ((_, _, title), (name, _)) in enumerate(zip(parts, marks), start=1):
        add_toc_entry(doc, f"{number}. {title}", name)

    first_position = True
    for (kind, payload, _), mark in zip(parts, marks):
        if kind in ("extract", "regulation"):
            add_page_break(doc)
            _write_specs(doc, payload, mark)
        elif kind == "position":
            if first_position:
                add_page_break(doc)
            _write_position(doc, payload, book.duties, mark, review, first_position)
            first_position = False
        else:
            _signature(doc, book.signer)
            add_page_break(doc)
            _familiarization(doc, book, mark)
            _signature(doc, book.signer)
    add_page_break(doc)
    _stitching(doc, book.signer)
    return doc, review
