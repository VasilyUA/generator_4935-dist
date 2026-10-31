"""Форматування python-docx. save_docx_safely/add_paragraph_with_style/
add_table - копії з generator_bn_and_map/utils/docx_utils.py (проєкти в цьому
репо не імпортують код один одного); закладки й поле PAGEREF - для ЗМІСТу."""
import os
import time
from typing import Union

from docx.enum.table import WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_COLOR_INDEX, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Emu, Pt, RGBColor

FONT_NAME = "Times New Roman"
FONT_SIZE = 14


def _close_matching_word_documents(matches):
    """Закриває (без збереження) кожен відкритий у Word документ, для якого
    matches(абсолютний_шлях) - True. Якщо Word не запущено або pywin32
    недоступний - тихо виходить."""
    try:
        import pythoncom
        import win32com.client
    except ImportError:
        return

    try:
        pythoncom.CoInitialize()
        try:
            word = win32com.client.GetActiveObject("Word.Application")
        except Exception:
            return
        for doc in list(word.Documents):
            try:
                full_name = os.path.abspath(doc.FullName)
            except Exception:
                continue
            if matches(full_name):
                doc.Close(SaveChanges=False)
    finally:
        pythoncom.CoUninitialize()


def save_docx_safely(doc, file_path, attempts=5, delay=0.5):
    """doc.save(), але спершу закриває цей самий файл у Word, якщо він там
    відкритий (попередня книга з такою самою назвою), і повторює спробу, якщо
    файл лишається заблокованим ще коротку мить (антивірус/OneDrive)."""
    abs_path = os.path.abspath(file_path)
    _close_matching_word_documents(lambda full_name: full_name == abs_path)
    for attempt in range(attempts):
        try:
            doc.save(file_path)
            return
        except PermissionError as e:
            if attempt == attempts - 1:
                raise PermissionError(
                    f"Не вдалось зберегти '{file_path}': файл заблоковано іншою програмою. "
                    f"Закрийте його (напр. Word) і спробуйте ще раз. ({e})"
                ) from e
            time.sleep(delay)


def set_page(doc, top, bottom, left, right):
    """A4, книжкова, поля як у зразку."""
    for section in doc.sections:
        section.page_height = Cm(29.7)
        section.page_width = Cm(21.0)
        section.top_margin = top
        section.bottom_margin = bottom
        section.left_margin = left
        section.right_margin = right


def text_width(doc):
    section = doc.sections[0]
    return Emu(section.page_width - section.left_margin - section.right_margin)


def set_base_font(doc, font_name=FONT_NAME, font_size=FONT_SIZE):
    """Normal і Heading 2 - Times New Roman, чорний: заголовки посад мають
    стиль Heading 2 лише для навігації Word, вигляд задається прямо в runs."""
    for style_name in ("Normal", "Heading 2"):
        style = doc.styles[style_name]
        style.font.name = font_name
        style.font.size = Pt(font_size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        rfonts = style.element.get_or_add_rPr().get_or_add_rFonts()
        for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
            rfonts.set(qn(attr), font_name)
        for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
            rfonts.attrib.pop(qn(attr), None)


def _set_paragraph_mark_font(paragraph, font_name, font_size):
    # Без цього порожній абзац (без жодного run-у з текстом) показує шрифт теми
    # замість заданого, бо Word бере шрифт кінцевого маркера абзацу (pPr/rPr).
    rPr = OxmlElement('w:rPr')
    rFonts = OxmlElement('w:rFonts')
    rFonts.set(qn('w:ascii'), font_name)
    rFonts.set(qn('w:hAnsi'), font_name)
    rFonts.set(qn('w:eastAsia'), font_name)
    rPr.append(rFonts)
    color = OxmlElement('w:color')
    color.set(qn('w:val'), '000000')
    rPr.append(color)
    sz = OxmlElement('w:sz')
    sz.set(qn('w:val'), str(int(font_size * 2)))
    rPr.append(sz)
    paragraph._p.get_or_add_pPr().append(rPr)


def style_run(run, font_name=FONT_NAME, font_size=FONT_SIZE, bold=False, italic=False):
    run.font.name = font_name
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), font_name)
    run.font.size = Pt(font_size)
    run.font.color.rgb = RGBColor(0, 0, 0)
    run.font.bold = bold
    run.font.italic = italic
    return run


def add_paragraph_with_style(container, text, font_name=FONT_NAME, font_size=FONT_SIZE, bold: Union[bool, list] = False,
                             alignment=WD_ALIGN_PARAGRAPH.JUSTIFY, first_line_indent_cm=0.0, tab_width=None,
                             italic=False, highlight=False, heading_level=None, left_indent_cm=None,
                             keep_with_next=False):
    """container - документ або клітинка таблиці. bold - True/False для всього
    абзацу або список фраз, які треба виділити жирним. tab_width - правий
    табулятор на цій відстані (рядки "звання<TAB>ПІБ" у підписах)."""
    paragraph = container.add_paragraph()
    if heading_level is not None:
        # Структура/навігація документа - як заголовок, але вигляд задається
        # прямими runs нижче.
        paragraph.style = f"Heading {heading_level}"
    paragraph.alignment = alignment
    paragraph.paragraph_format.line_spacing = 1.0
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.first_line_indent = Cm(first_line_indent_cm)
    paragraph.paragraph_format.keep_with_next = keep_with_next
    if left_indent_cm is not None:
        paragraph.paragraph_format.left_indent = Cm(left_indent_cm)
    if tab_width is not None:
        paragraph.paragraph_format.tab_stops.add_tab_stop(tab_width, WD_TAB_ALIGNMENT.RIGHT)

    def add_run(t, is_bold):
        run = style_run(paragraph.add_run(t), font_name, font_size, bold=is_bold, italic=italic)
        if highlight:
            run.font.highlight_color = WD_COLOR_INDEX.YELLOW
        return run

    if isinstance(bold, list):
        remaining = text
        while remaining:
            earliest_word = None
            earliest_pos = len(remaining)
            for word in bold:
                pos = remaining.find(word) if word else -1
                if pos != -1 and pos < earliest_pos:
                    earliest_pos = pos
                    earliest_word = word
            if earliest_word is None:
                add_run(remaining, False)
                break
            if earliest_pos > 0:
                add_run(remaining[:earliest_pos], False)
            add_run(earliest_word, True)
            remaining = remaining[earliest_pos + len(earliest_word):]
    elif text:
        add_run(text, bool(bold))

    _set_paragraph_mark_font(paragraph, font_name, font_size)
    return paragraph


def add_page_break(doc):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.add_run().add_break(WD_BREAK.PAGE)
    _set_paragraph_mark_font(paragraph, FONT_NAME, FONT_SIZE)
    return paragraph


def add_bookmark(paragraph, name, bookmark_id):
    """Закладка навколо всього тексту абзацу (ціль поля PAGEREF у ЗМІСТі)."""
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), str(bookmark_id))
    start.set(qn("w:name"), name)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), str(bookmark_id))
    p = paragraph._p
    first_run = p.find(qn("w:r"))
    if first_run is not None:
        first_run.addprevious(start)
    else:
        p.append(start)
    p.append(end)


def _field_run(paragraph, font_name, font_size, element=None, text=None):
    run = style_run(paragraph.add_run(text or ""), font_name, font_size)
    if element is not None:
        run._r.append(element)
    return run


def add_pageref_field(paragraph, bookmark_name, font_name=FONT_NAME, font_size=FONT_SIZE, placeholder="#"):
    """Поле PAGEREF <закладка> - номер сторінки, на якій стоїть закладка. Word
    рахує його під час оновлення полів (word_finalize або F9)."""
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    _field_run(paragraph, font_name, font_size, begin)
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = f" PAGEREF {bookmark_name} \\h "
    _field_run(paragraph, font_name, font_size, instr)
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    _field_run(paragraph, font_name, font_size, separate)
    _field_run(paragraph, font_name, font_size, text=placeholder)
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    _field_run(paragraph, font_name, font_size, end)


def add_toc_entry(doc, text, bookmark_name, font_name=FONT_NAME, font_size=FONT_SIZE):
    """Рядок ЗМІСТу: "текст ....... N" (правий табулятор з крапками + PAGEREF)."""
    paragraph = add_paragraph_with_style(doc, text, font_name, font_size)
    paragraph.paragraph_format.tab_stops.add_tab_stop(text_width(doc), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
    style_run(paragraph.add_run("\t"), font_name, font_size)
    add_pageref_field(paragraph, bookmark_name, font_name, font_size)
    return paragraph


def enable_update_fields_on_open(doc):
    """Word сам запропонує оновити поля (ЗМІСТ) під час відкриття - запасний
    варіант, коли word_finalize не зміг оновити їх через Word."""
    settings = doc.settings.element
    existing = settings.find(qn("w:updateFields"))
    if existing is None:
        existing = OxmlElement("w:updateFields")
        settings.append(existing)
    existing.set(qn("w:val"), "true")


def _repeat_as_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def add_table(doc, columns, rows, widths_cm, font_name=FONT_NAME, font_size=12, min_row_height_cm=None):
    """Таблиця 'Table Grid': рядок назв колонок (жирний, повторюється на
    кожній сторінці) і рядки даних (не нижчі за min_row_height_cm - місце для
    підпису від руки). Ширину визначає tblGrid (tbl.columns[i].width), а не
    лише cell.width - без цього Word ділить таблицю на рівні колонки."""
    n = len(columns)
    table = doc.add_table(rows=1 + len(rows), cols=n)
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = table.allow_autofit = False
    for i, width in enumerate(widths_cm[:n]):
        table.columns[i].width = Cm(width)
        for cell in table.columns[i].cells:
            cell.width = Cm(width)

    def fill(cell, text, bold=False, center=False):
        for line_index, line in enumerate(str(text).split("\n")):
            paragraph = cell.paragraphs[0] if line_index == 0 else cell.add_paragraph()
            if center:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            if line:
                style_run(paragraph.add_run(line), font_name, font_size, bold=bold)

    for i, column in enumerate(columns):
        fill(table.rows[0].cells[i], column, bold=True, center=True)
    _repeat_as_header(table.rows[0])
    for r, row in enumerate(rows, start=1):
        values = list(row) + [""] * (n - len(row))
        for c, value in enumerate(values[:n]):
            fill(table.rows[r].cells[c], value, center=(c == 0))
        if min_row_height_cm:
            table.rows[r].height = Cm(min_row_height_cm)
            table.rows[r].height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
    return table
