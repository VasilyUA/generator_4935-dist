from docx import Document

from constants import SECTION_DUTIES, SECTION_GENERAL, SECTION_RESPONSIBILITY, SECTION_RIGHTS
from content.positions import compare_key
from content.sample_reader import SampleBlock, find_block, parse_sample, read_sample
from duties_fixtures import STANDARD_RESPONSIBILITY, write_sample


def test_read_sample_takes_extract_until_regulation(tmp_path):
    sample = read_sample(write_sample(tmp_path / "Зразок.docx"))
    assert [spec.text for spec in sample.extract] == [
        "Витяг із Закону України Про Статут внутрішньої служби Збройних Сил України",
        "Загальні обов’язки військовослужбовців",
        "11. Військовослужбовець зобов’язаний бути дисциплінованим.",
    ]
    assert sample.extract[0].bold


def test_read_sample_finds_titles_from_heading_centered_lines_and_duties_marker(tmp_path):
    sample = read_sample(write_sample(tmp_path / "Зразок.docx"))
    assert [block.title for block in sample.blocks] == ["Командир батальйону", "Водій автомобільного відділення",
                                                        "Кухар їдальні"]


def test_read_sample_splits_sections_strips_bullets_and_stops_at_signature(tmp_path):
    sample = read_sample(write_sample(tmp_path / "Зразок.docx"))
    driver = sample.blocks[1]
    assert driver.sections[SECTION_GENERAL] == ["Водій підпорядковується командиру відділення."]
    assert driver.sections[SECTION_DUTIES] == ["утримувати автомобіль справним;"]
    assert driver.sections[SECTION_RIGHTS] == ["отримувати пальне;"]
    assert driver.sections[SECTION_RESPONSIBILITY] == ["за стан автомобіля."]
    cook = sample.blocks[2]
    assert SECTION_RESPONSIBILITY not in cook.sections
    assert all("військової частини" not in item for block in sample.blocks for items in block.sections.values()
               for item in items)


def test_read_sample_moves_items_present_in_most_blocks_to_common(tmp_path):
    sample = read_sample(write_sample(tmp_path / "Зразок.docx"))
    assert sample.common == {SECTION_RESPONSIBILITY: list(STANDARD_RESPONSIBILITY)}


def test_parse_sample_without_table_and_without_blocks():
    doc = Document()
    doc.add_paragraph("Текст без витягу")
    sample = parse_sample(doc)
    assert sample.extract == [] and sample.blocks == [] and sample.common == {}


def test_parse_sample_ignores_tables_before_blocks_and_untitled_blocks():
    doc = Document()
    doc.add_table(rows=1, cols=1)
    doc.add_paragraph().add_run("Витяг із статуту").bold = True
    doc.add_paragraph("ПОЛОЖЕННЯ")
    doc.add_table(rows=1, cols=1)
    doc.add_paragraph("Загальні положення")
    doc.add_paragraph("Текст без назви посади.")
    sample = parse_sample(doc)
    assert sample.blocks == []


def test_parse_sample_keeps_other_colon_lines_as_items_and_stops_at_table():
    doc = Document()
    doc.add_paragraph().add_run("Витяг із статуту").bold = True
    doc.add_paragraph("ПОЛОЖЕННЯ")
    doc.add_paragraph("Загальні положення")
    doc.add_paragraph("Кухар підпорядковується начальнику їдальні.")
    doc.add_paragraph("Кухар зобов’язаний:")
    doc.add_paragraph("знає та застосовує у діяльності:")
    doc.add_table(rows=1, cols=1)
    doc.add_paragraph("Загальні положення")
    sample = parse_sample(doc)
    assert [b.title for b in sample.blocks] == ["Кухар"]
    assert sample.blocks[0].sections[SECTION_DUTIES] == ["знає та застосовує у діяльності:"]


def test_parse_sample_title_with_vos_other_markers_and_unmarked_duties():
    doc = Document()
    doc.add_paragraph().add_run("Витяг із статуту").bold = True
    doc.add_paragraph("ПОЛОЖЕННЯ")
    heading = doc.add_paragraph()
    heading.add_run("Начальник лазні ВОС 1234567").bold = True
    heading.alignment = 1
    doc.add_paragraph("Загальні положення")
    doc.add_paragraph("Начальник лазні підпорядковується командиру взводу.")
    doc.add_paragraph("Начальник лазні повинен:")
    doc.add_paragraph("утримувати лазню;")
    doc.add_paragraph("Оператор", style="Heading 2")
    doc.add_paragraph("Загальні положення")
    doc.add_paragraph("Функції оператора:")
    doc.add_paragraph("обслуговувати обладнання;")
    doc.add_paragraph("Командир відділення", style="Heading 2")
    doc.add_paragraph("Загальні положення")
    doc.add_paragraph("Командир відділення підпорядковується командиру взводу.")
    doc.add_paragraph("знати особовий склад;")
    doc.add_paragraph("Кухар", style="Heading 2")
    doc.add_paragraph("Загальні положення")
    doc.add_paragraph("Кухар готує їжу.")
    doc.add_paragraph("Писар", style="Heading 2")
    doc.add_paragraph("Загальні положення")
    doc.add_paragraph("вести облік;")
    blocks = {block.title: block.sections for block in parse_sample(doc).blocks}
    assert blocks["Начальник лазні"][SECTION_DUTIES] == ["утримувати лазню;"]
    assert blocks["Оператор"][SECTION_DUTIES] == ["обслуговувати обладнання;"]
    assert blocks["Командир відділення"] == {SECTION_GENERAL: ["Командир відділення підпорядковується командиру взводу."],
                                             SECTION_DUTIES: ["знати особовий склад;"]}
    assert blocks["Кухар"] == {SECTION_GENERAL: ["Кухар готує їжу."]}
    assert blocks["Писар"] == {SECTION_DUTIES: ["вести облік;"]}


def test_parse_sample_drops_the_sample_battalion_type_from_texts():
    doc = Document()
    doc.add_paragraph().add_run("Витяг із статуту").bold = True
    doc.add_paragraph("ПОЛОЖЕННЯ")
    doc.add_paragraph("про батальйон матеріального забезпечення")
    doc.add_paragraph("Командир взводу", style="Heading 2")
    doc.add_paragraph("Загальні положення")
    doc.add_paragraph("Командир взводу матеріального забезпечення батальйону матеріального забезпечення "
                      "підпорядковується командирові Батальйону Матеріального Забезпечення.")
    sample = parse_sample(doc)
    assert sample.battalion_type == "матеріального забезпечення"
    assert sample.blocks[0].sections[SECTION_GENERAL] == [
        "Командир взводу матеріального забезпечення батальйону підпорядковується командирові Батальйону."
    ]


def test_find_block_matches_whole_word_prefix_and_prefers_longest():
    blocks = [SampleBlock("Водій"), SampleBlock("Водій автомобільного відділення"), SampleBlock("Кухар")]
    found = find_block(blocks, compare_key("Водій автомобільного відділення взводу"))
    assert found.title == "Водій автомобільного відділення"
    chief = [SampleBlock("Начальник штабу")]
    assert find_block(chief, compare_key("Начальник штабу-заступник командира батальйону")).title == "Начальник штабу"
    assert find_block(blocks, compare_key("Водій-електрик")) is None
    assert find_block(blocks, compare_key("Кухар")).title == "Кухар"
