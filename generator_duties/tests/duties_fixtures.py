"""Синтетичні вхідні файли для тестів generator_duties: штат (xlsx), зразок
(docx), сторінка Статуту (html). Лише вигадані дані: порядкові ПІБ, A0000."""
import openpyxl
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

from content.roster_reader import Slot
from content.positions import group_positions

SUFFIX = "батальйону A0000"
DATA_EXAMPLE = {
    "unit": {
        "MILITARY_UNIT": "Військова частина A0000",
        "MILITARY_UNIT_GENITIVE": "військової частини A0000",
        "BATTALION_GENITIVE": SUFFIX,
        "SETTLEMENT": "",
    },
    "roster": {"file": "", "sheet": "", "include_if_contains": [], "position_suffixes": [SUFFIX]},
    "signer": {"position": "", "rank": "", "name": ""},
    "duplex": False,
    "blank_rows": 20,
}
ROSTER_HEADER = ("№", "Посада", "ВОС", "Військове звання", "ПІБ")
ROSTER_ROWS = (
    (1, f"Командир {SUFFIX}", "1000001", "майор", "ПЕРШИЙ Перший Перший"),
    (2, f"Начальник штабу – заступник командира {SUFFIX}", "1000002", "капітан", "ДРУГИЙ Другий Другий"),
    (3, f"Командир 2 відділення 3 взводу {SUFFIX}", "1000003", "сержант", "ТРЕТІЙ Третій Третій"),
    (4, f"Стрілець 2 відділення 3 взводу {SUFFIX}", "1000004", "солдат", "ЧЕТВЕРТИЙ Четвертий Четвертий"),
    (5, f"Стрілець 3 відділення 3 взводу {SUFFIX}", "1000004", "", ""),
    (6, f"Водій господарчого відділення {SUFFIX}", "1000005", "солдат", "ШОСТИЙ Шостий Шостий"),
)


# Штат у будові повного штатно-посадового списку: кілька колонок "посада"
# (коротка, повна в називному, родовий), "підрозділ повністю", "П.І.Б.",
# рядки-розділювачі груп (лише коротка назва).
BATTALION = "батальйону A0001"
UNIT = "військової частини A0000"
FULL_HEADER = ("посадовий оклад", "категорія", "Посада зі всім", "Називний для посада 1", "Родовий для посада 1",
               "посада", "підрозділ повністю", "підрозділ", "Посада", "ВОС", "звання за штатом", "звання фактичне",
               "П.І.Б.")
FULL_ROWS = (
    (None, "У", None, None, None, None, None, None, "Управління", None, None, None, None),
    (1000, "оф", None, f"Командир {BATTALION} {UNIT}", f"Командира {BATTALION} {UNIT}", "Командира", BATTALION,
     "управління", "Командир батальйону", "1000001", "підполковник", "майор", "ПЕРШИЙ Перший Перший"),
    (900, "оф", None, f"Заступник командира {BATTALION} з озброєння {UNIT}",
     f"Заступника командира {BATTALION} з озброєння {UNIT}", "Заступника командира з озброєння", BATTALION,
     "управління", "Заступник командира батальйону з озброєння", "1000002", "майор", None, None),
    (None, "Р", None, None, None, None, None, None, "Рота", None, None, None, None),
    (500, "ряд", None, f"Стрілець 1 відділення 2 взводу роти {BATTALION} {UNIT}",
     f"Стрільця 1 відділення 2 взводу роти {BATTALION} {UNIT}", "Стрільця", f"1 відділення 2 взводу роти {BATTALION}",
     "рота", "Стрілець", "1000003", "солдат", "солдат", "ДРУГИЙ Другий Другий"),
    (500, "ряд", None, f"Стрілець 2 відділення 2 взводу роти {BATTALION} {UNIT}",
     f"Стрільця 2 відділення 2 взводу роти {BATTALION} {UNIT}", "Стрільця", f"2 відділення 2 взводу роти {BATTALION}",
     "рота", "Стрілець", "1000003", "солдат", None, None),
)


def write_roster(path, rows=ROSTER_ROWS, header=ROSTER_HEADER, title_rows=1, sheet_title="Штат"):
    workbook = openpyxl.Workbook()
    ws = workbook.active
    ws.title = sheet_title
    for _ in range(title_rows):
        ws.append(["Штатно-посадовий список"])
    ws.append(list(header))
    for row in rows:
        ws.append(list(row))
    workbook.save(path)
    return path


def make_positions(rows=ROSTER_ROWS):
    slots = [Slot(i + 3, row[1], row[3], row[4], row[2]) for i, row in enumerate(rows)]
    return group_positions(slots, [SUFFIX])


def _p(doc, text, bold=False, center=False, style=None, listed=False):
    paragraph = doc.add_paragraph(style=style)
    run = paragraph.add_run(text)
    run.bold = bold
    if center:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if listed:
        num_pr = paragraph._p.get_or_add_pPr().get_or_add_numPr()
        num_pr.get_or_add_numId().val = 1
        num_pr.get_or_add_ilvl().val = 0
    return paragraph


STANDARD_RESPONSIBILITY = (
    "за неналежне виконання своїх обов’язків;",
    "за завдання матеріальної шкоди.",
)


def write_sample(path, with_table=True):
    """Зразок у тій самій будові, що й отриманий: пам'ятка, ЗМІСТ, витяг,
    ПОЛОЖЕННЯ, блоки посад (з різними видами заголовків), підпис, таблиця."""
    doc = Document()
    _p(doc, "ВИМОГИ ДО ЗШИВАННЯ СПРАВИ")
    _p(doc, "1. Витяг із статуту внутрішньої служби…….3")
    _p(doc, "Витяг із Закону України Про Статут внутрішньої служби Збройних Сил України", bold=True)
    _p(doc, "Загальні обов’язки військовослужбовців", bold=True)
    _p(doc, "11. Військовослужбовець зобов’язаний бути дисциплінованим.")
    _p(doc, "ПОЛОЖЕННЯ", bold=True, center=True)
    _p(doc, "про батальйон", bold=True, center=True)
    _p(doc, "1. Призначення батальйону.", bold=True)
    # Блок 1: порожній заголовок - назва з рядка "… зобов’язаний:".
    _p(doc, "", style="Heading 2")
    _p(doc, "ВОС 0000000", bold=True, center=True)
    _p(doc, "Загальні положення", bold=True)
    _p(doc, "Командир батальйону підпорядковується командиру бригади.")
    _p(doc, "Командир батальйону зобов’язаний:", bold=True)
    _p(doc, "керувати батальйоном;", listed=True)
    _p(doc, "Командир батальйону несе відповідальність:", bold=True)
    for text in STANDARD_RESPONSIBILITY:
        _p(doc, text, listed=True)
    # Блок 2: заголовок Heading 2.
    _p(doc, "Водій автомобільного відділення", style="Heading 2")
    _p(doc, "ВОС 0000001", bold=True, center=True)
    _p(doc, "Загальні положення", bold=True)
    _p(doc, "Водій підпорядковується командиру відділення.")
    _p(doc, "Водій зобов’язаний:", bold=True)
    _p(doc, "- утримувати автомобіль справним;")
    _p(doc, "Водій має право:", bold=True)
    _p(doc, "- отримувати пальне;")
    _p(doc, "Водій несе відповідальність:", bold=True)
    for text in STANDARD_RESPONSIBILITY:
        _p(doc, text, listed=True)
    _p(doc, "за стан автомобіля.", listed=True)
    # Блок 3: центрований жирний заголовок у два абзаци.
    _p(doc, "Кухар", bold=True, center=True)
    _p(doc, "їдальні", bold=True, center=True)
    _p(doc, "Загальні положення", bold=True)
    _p(doc, "Кухар підпорядковується начальнику їдальні.")
    _p(doc, "Кухар зобов’язаний:", bold=True)
    _p(doc, "готувати їжу;", listed=True)
    _p(doc, "Кухар несе відповідальність:", bold=True)
    for text in STANDARD_RESPONSIBILITY:
        _p(doc, text, listed=True)
    _p(doc, "Командир батальйону військової частини A0000")
    _p(doc, "майор\tПерший ПЕРШИЙ")
    if with_table:
        doc.add_table(rows=2, cols=2)
        _p(doc, "Командир батальйону військової частини A0000")
    doc.save(path)
    return path


STATUTE_HTML = """<html><body>
<p class="rvps7">Відповідальність військовослужбовців</p>
<p class="rvps2">26. Військовослужбовці несуть дисциплінарну відповідальність.</p>
<p class="rvps2">{Стаття 26 із змінами}</p>
<p class="rvps2">27. Стягнення не звільняє від матеріальної відповідальності.</p>
<p class="rvps7">Розділ 3. Обов’язки посадових осіб</p>
<p class="rvps7">Командир батальйону (корабля 3 рангу)</p>
<p class="rvps2">101. Командир батальйону (корабля 3 рангу)* відповідає за бойову готовність.</p>
<p class="rvps8">__________ * Далі - командир батальйону.</p>
<p class="rvps2">Командир батальйону підпорядковується командирові бригади.</p>
<p class="rvps2">102. Командир батальйону зобов’язаний:</p>
<p class="rvps2">підтримувати бойову готовність;</p>
<p class="rvps2">керувати бойовою підготовкою.</p>
<p class="rvps7">Начальник штабу батальйону</p>
<p class="rvps2">103. Начальник штабу батальйону відповідає за управління.</p>
<p class="rvps2">104. Начальник штабу батальйону зобов’язаний:</p>
<p class="rvps2">організовувати управління;</p>
<p class="rvps2">105. Начальник штабу батальйону має право:</p>
<p class="rvps2">віддавати розпорядження.</p>
<p class="rvps7">Командир роти (корабля 4 рангу)</p>
<p class="rvps2">111. Командир роти відповідає за роту.</p>
<p class="rvps2">112. Командир роти зобов’язаний:</p>
<p class="rvps2">керувати ротою;</p>
<p class="rvps2">Крім того, командир окремої роти зобов’язаний:</p>
<p class="rvps2">вести господарство.</p>
<p class="rvps2">112 - 1 . Окрема стаття після переліку.</p>
<p class="rvps7">Командир відділення</p>
<p class="rvps2">125. Командир відділення відповідає за відділення.</p>
<p class="rvps2">126. Командир відділення зобов’язаний:</p>
<p class="rvps2">знати особовий склад;</p>
<p class="rvps7">Загальні обов’язки військовослужбовців рядового складу</p>
<p class="rvps2">127. Солдат відповідає за виконання обов’язків.</p>
<p class="rvps2">128. Солдат зобов’язаний:</p>
<p class="rvps2">знати свої обов’язки;</p>
</body></html>"""


def write_statute(path, html=STATUTE_HTML, encoding="utf-8"):
    with open(path, "w", encoding=encoding) as f:
        f.write(html)
    return path
