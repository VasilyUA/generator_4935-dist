import os
import time

import openpyxl
import pytest

from content.roster_reader import RosterError, detect_header, find_roster_file, read_roster
from duties_fixtures import BATTALION, FULL_HEADER, FULL_ROWS, ROSTER_ROWS, SUFFIX, write_roster


def test_read_roster_detects_header_below_title_rows(tmp_path):
    write_roster(tmp_path / "Штат.xlsx", title_rows=3)
    roster = read_roster(str(tmp_path))
    assert roster.sheet == "Штат"
    assert roster.headers == {"position": "Посада", "rank": "Військове звання", "pib": "ПІБ", "vos": "ВОС"}
    assert len(roster.slots) == len(ROSTER_ROWS)
    first = roster.slots[0]
    assert (first.row, first.rank, first.pib, first.vos) == (5, "майор", "ПЕРШИЙ Перший Перший", "1000001")
    assert roster.slots[4].vacant


def test_read_roster_skips_separator_rows_without_vos_and_person(tmp_path):
    rows = ((0, "Управління батальйону", "", "", ""),) + ROSTER_ROWS
    write_roster(tmp_path / "Штат.xlsx", rows=rows)
    assert len(read_roster(str(tmp_path)).slots) == len(ROSTER_ROWS)


def test_read_roster_keeps_rows_without_vos_when_there_is_no_vos_column(tmp_path):
    header = ("Посада", "Звання", "Прізвище, ім’я, по батькові")
    write_roster(tmp_path / "Штат.xlsx", rows=(("Стрілець", "", ""), ("", "", "")), header=header)
    roster = read_roster(str(tmp_path))
    assert [slot.position for slot in roster.slots] == ["Стрілець"]
    assert roster.slots[0].vos == ""


def test_read_roster_filters_rows_by_configured_text(tmp_path):
    rows = ROSTER_ROWS + ((7, "Стрілець відділення батальйону A0001", "1", "солдат", "СЬОМИЙ Сьомий Сьомий"),)
    write_roster(tmp_path / "Штат.xlsx", rows=rows)
    roster = read_roster(str(tmp_path), {"include_if_contains": [SUFFIX]})
    assert len(roster.slots) == len(ROSTER_ROWS)


def test_read_roster_skips_repeated_header_and_converts_numbers(tmp_path):
    rows = (("№", "Посада", "ВОС", "Військове звання", "ПІБ"), (1, "Стрілець", 1000004.0, "солдат", "ПЕРШИЙ Перший Перший"))
    write_roster(tmp_path / "Штат.xlsx", rows=rows)
    roster = read_roster(str(tmp_path))
    assert [(s.position, s.vos) for s in roster.slots] == [("Стрілець", "1000004")]


def test_read_roster_uses_configured_sheet_and_file(tmp_path):
    write_roster(tmp_path / "Інший.xlsx")
    path = write_roster(tmp_path / "Мій штат.xlsx", sheet_title="Основний")
    roster = read_roster(str(tmp_path), {"file": "Мій штат.xlsx", "sheet": "основний"})
    assert roster.path == str(path)
    with pytest.raises(RosterError, match="немає аркуша"):
        read_roster(str(tmp_path), {"file": "Мій штат.xlsx", "sheet": "Немає"})
    with pytest.raises(RosterError, match="не знайдено"):
        read_roster(str(tmp_path), {"file": "Немає.xlsx"})


def test_read_roster_tries_next_sheet_when_first_has_no_header(tmp_path):
    path = tmp_path / "Штат.xlsx"
    workbook = openpyxl.Workbook()
    workbook.active.append(["Довідка"])
    ws = workbook.create_sheet("Посади")
    ws.append(["Посада", "ПІБ"])
    ws.append(["Стрілець", "ПЕРШИЙ Перший Перший"])
    workbook.save(path)
    assert read_roster(str(tmp_path)).sheet == "Посади"


def test_read_roster_reports_missing_header(tmp_path):
    workbook = openpyxl.Workbook()
    workbook.active.append(["Довідка"])
    workbook.save(tmp_path / "Штат.xlsx")
    with pytest.raises(RosterError, match="заголовка"):
        read_roster(str(tmp_path))


def test_read_roster_explains_password_protected_file(tmp_path):
    (tmp_path / "Штат.xlsx").write_bytes(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\0" * 64)
    with pytest.raises(RosterError, match="паролем"):
        read_roster(str(tmp_path))


def test_read_roster_reports_broken_file(tmp_path):
    (tmp_path / "Штат.xlsx").write_bytes(b"not a workbook")
    with pytest.raises(RosterError, match="Не вдалось відкрити"):
        read_roster(str(tmp_path))


def test_find_roster_file_skips_duties_table_and_takes_newest(tmp_path):
    with pytest.raises(RosterError, match="немає файлу штату"):
        find_roster_file(str(tmp_path))
    write_roster(tmp_path / "Обов'язки.xlsx")
    write_roster(tmp_path / "~$Штат.xlsx")
    old = write_roster(tmp_path / "Старий.xlsx")
    new = write_roster(tmp_path / "Новий.xlsx")
    past = time.time() - 100
    os.utime(old, (past, past))
    assert find_roster_file(str(tmp_path)) == str(new)


def test_detect_header_prefers_actual_rank_over_staff_rank_and_skips_vos_position_column():
    header = ("ВОС посади", "Посада", "Штатне звання", "Військове звання", "ПІБ")
    row_index, columns = detect_header([("Назва",), header])
    assert row_index == 1
    assert columns == {"position": 1, "rank": 3, "pib": 4, "vos": 0}


def test_detect_header_returns_none_without_position_and_pib():
    assert detect_header([("Посада", "Звання")]) is None


def test_read_full_staff_list_picks_nominative_full_position_and_dotted_pib(tmp_path):
    write_roster(tmp_path / "ШТАТ.xlsx", rows=FULL_ROWS, header=FULL_HEADER, title_rows=0)
    roster = read_roster(str(tmp_path))
    assert roster.headers == {"position": "Називний для посада 1", "rank": "звання фактичне", "pib": "П.І.Б.",
                              "vos": "ВОС", "subunit": "підрозділ повністю", "group": "підрозділ"}
    assert [slot.position for slot in roster.slots] == [row[3] for row in FULL_ROWS if row[3]]
    assert [slot.vacant for slot in roster.slots] == [False, True, False, True]
    assert roster.slots[2].subunit == f"1 відділення 2 взводу роти {BATTALION}"
    assert [slot.group for slot in roster.slots] == ["управління", "управління", "рота", "рота"]


def test_detect_header_takes_position_column_with_most_distinct_values_and_full_subunit():
    rows = [
        ("Посада", "Посада повна", "підрозділ", "Підрозділ (повністю)", "ПІБ"),
        ("Стрілець", "Стрілець 1 відділення", "рота", "1 відділення роти", "ПЕРШИЙ Перший Перший"),
        ("Стрілець", "Стрілець 2 відділення", "рота", "2 відділення роти", ""),
    ]
    assert detect_header(rows)[1] == {"position": 1, "subunit": 3, "group": 2, "pib": 4}


def test_detect_header_searches_only_first_rows():
    rows = [("Довідка",)] * 3 + [("Посада", "ПІБ")]
    assert detect_header(rows, scan_rows=3) is None
    assert detect_header(rows, scan_rows=4)[0] == 3
