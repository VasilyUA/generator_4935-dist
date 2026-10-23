from docx import Document
from docx.enum.table import WD_ROW_HEIGHT_RULE
from docx.enum.text import WD_COLOR_INDEX
from docx.oxml.ns import qn

from constants import BOOKS, SECTION_DUTIES, SECTION_GENERAL, SECTION_RIGHTS, SHEET_FUNCTIONAL, SHEET_JOB, SHEETS_MARKER, TO_FILL
from content.docx_text import ParagraphSpec
from content.duties_table import DutiesTable, DutySet, PositionRow, SheetDuties, read_table
from content.positions import group_positions
from content.roster_reader import Slot
from content.sample_reader import read_sample
from content.sources import Sources
from content.statute_reader import parse_statute_html
from duties_fixtures import STATUTE_HTML, make_positions, write_sample
from generators.book_builder import BookInput, Signer, build_book
from generators.generate_books import blank_rows, generate_books, make_signer, review_lines
from generators.update_table import update_table
from generators.word_finalize import FinalizeError

UNIT = {"MILITARY_UNIT": "Військова частина A0000", "MILITARY_UNIT_GENITIVE": "військової частини A0000",
        "BATTALION_GENITIVE": "батальйону A0001", "SETTLEMENT": ""}


def _texts(doc):
    return [p.text for p in doc.paragraphs]


def _table(sheet, own, references=None):
    """own - {посада: {розділ: [пункти]}} - кожна посада зі своїм набором;
    references - {посада: назва набору} (набір може не існувати)."""
    table = DutiesTable()
    for title, sections in own.items():
        duty_set = DutySet(title)
        for section, texts in sections.items():
            for text in texts:
                duty_set.add(section, text)
        table.positions.append(PositionRow(title, {sheet: table.add_set(sheet, duty_set).name}))
    for title, name in (references or {}).items():
        table.positions.append(PositionRow(title, {sheet: name}))
    return table


def _book(positions, table, extract=(), regulation=(), blank=0):
    return BookInput(BOOKS[0], 2027, UNIT, Signer("Командир батальйону", "майор", "Перший ПЕРШИЙ"), positions,
                     SheetDuties(table, SHEET_FUNCTIONAL), list(extract), list(regulation), blank)


def test_build_book_structure_and_toc():
    positions = make_positions()
    table = _table(SHEET_FUNCTIONAL, {
        positions[0].title: {SECTION_DUTIES: ["керувати батальйоном;"],
                             SECTION_GENERAL: ["{посада} підпорядковується командиру бригади."]},
        positions[1].title: {SECTION_RIGHTS: ["– віддавати розпорядження"]},
    })
    extract = [ParagraphSpec("Витяг із Закону", bold=True), ParagraphSpec("Текст", listed=True),
               ParagraphSpec("По центру", center=True)]
    regulation = [ParagraphSpec("ПОЛОЖЕННЯ", True, True), ParagraphSpec("про батальйон", True, True),
                  ParagraphSpec("1. Призначення.")]
    doc, review = build_book(_book(positions, table, extract, regulation))
    texts = _texts(doc)
    toc = [t for t in texts if t[:2] in ("1.", "2.", "3.") and "\t" in t]
    assert toc[0].startswith("1. Витяг із Закону\t")
    assert toc[1].startswith("2. Положення про батальйон\t")
    assert toc[2].startswith("3. Командир батальйону ВОС 1000001\t")
    assert "– Текст" in texts
    assert "Командир батальйону підпорядковується командиру бригади." in texts
    assert "– керувати батальйоном;" in texts
    assert "– віддавати розпорядження" in texts
    assert "Начальник штабу – заступник командира батальйону має право:" in texts
    headings = [p.text for p in doc.paragraphs if p.style.name == "Heading 2"]
    assert headings == ["Витяг із Закону", "ПОЛОЖЕННЯ"] + [p.title for p in positions] + [BOOKS[0]["list_title"]]
    assert SHEETS_MARKER in texts
    assert texts.count("майор\tПерший ПЕРШИЙ") == 3
    bookmarks = doc.element.body.findall(".//" + qn("w:bookmarkStart"))
    assert len(bookmarks) == len(headings)
    assert len(doc.element.body.findall(".//" + qn("w:instrText"))) == len(headings)
    assert review.missing == [p.title for p in positions[2:]]
    assert review.to_fill == [p.title for p in positions[1:]]


def test_build_book_highlights_to_fill_and_lists_slots_then_blank_rows():
    positions = make_positions()
    table = _table(SHEET_FUNCTIONAL, {positions[0].title: {SECTION_DUTIES: [TO_FILL]}})
    doc, review = build_book(_book(positions, table, blank=3))
    highlighted = [p.text for p in doc.paragraphs if any(r.font.highlight_color == WD_COLOR_INDEX.YELLOW for r in p.runs)]
    assert highlighted[0] == f"{TO_FILL}: обов'язки"
    assert review.to_fill[0] == positions[0].title
    sheet = doc.tables[-1]
    assert [c.text for c in sheet.rows[0].cells] == ["№ з/п", "ВОС посади", "Посада", "Звання", "ПІБ", "Підпис в/с",
                                                     "Примітка"]
    names = [row.cells[4].text for row in sheet.rows[1:7]]
    assert names == ["ПЕРШИЙ Перший Перший", "ДРУГИЙ Другий Другий", "ТРЕТІЙ Третій Третій",
                     "ЧЕТВЕРТИЙ Четвертий Четвертий", "", "ШОСТИЙ Шостий Шостий"]
    assert [c.text for c in sheet.rows[1].cells[:4]] == ["1", "1000001", "Командир батальйону", "майор"]
    assert sheet.rows[3].cells[2].text == "Командир 2 відділення 3 взводу"
    assert len(sheet.rows) == 1 + 6 + 3
    assert [row.cells[0].text for row in sheet.rows[7:]] == ["7", "8", "9"]
    assert all(cell.text == "" for row in sheet.rows[7:] for cell in row.cells[1:])
    assert sheet.rows[8].height_rule == WD_ROW_HEIGHT_RULE.AT_LEAST


def test_build_book_without_vos_signer_or_positions():
    slots = [Slot(1, "Стрілець", "", "")]
    positions = group_positions(slots)
    book = BookInput(BOOKS[1], 2027, {}, Signer(), positions, SheetDuties(DutiesTable(), SHEET_JOB))
    doc, review = build_book(book)
    assert review.without_vos == ["Стрілець"]
    assert "Командир батальйону" in _texts(doc)
    assert "____________\t____________" in _texts(doc)
    empty = BookInput(BOOKS[1], 2027, {}, Signer(), [], SheetDuties(DutiesTable(), SHEET_JOB))
    doc, _ = build_book(empty)
    assert BOOKS[1]["list_title"] in _texts(doc)


def test_make_signer_from_staff_and_override():
    positions = make_positions()
    signer = make_signer(positions, {"unit": UNIT})
    assert signer == Signer("Командир батальйону A0000 військової частини A0000", "майор", "Перший ПЕРШИЙ")
    signer = make_signer(positions, {"signer": {"rank": "підполковник"}})
    assert signer.rank == "підполковник" and signer.name == "Перший ПЕРШИЙ"
    assert make_signer([], {}) == Signer()
    unit_in_title = group_positions([Slot(1, "Командир батальйону військової частини A0000", "майор", "ПЕРШИЙ Перший")],
                                    ["військової частини A0000"])
    assert make_signer(unit_in_title, {"unit": UNIT}).position == "Командир батальйону військової частини A0000"


def test_blank_rows_default_and_configured():
    assert blank_rows({}, 681) == 70
    assert blank_rows({}, 50) == 20
    assert blank_rows({"blank_rows": 0}, 681) == 0
    assert blank_rows({"blank_rows": 15}, 681) == 15
    assert blank_rows({"blank_rows": -1}, 681) == 70
    assert blank_rows({"blank_rows": "багато"}, 681) == 70
    assert blank_rows({"blank_rows": True}, 681) == 70


def _sources(tmp_path):
    return Sources(statute=parse_statute_html(STATUTE_HTML), sample=read_sample(write_sample(tmp_path / "s.docx")),
                   regulation=[ParagraphSpec("ПОЛОЖЕННЯ", True, True)])


def test_generate_books_writes_two_files_and_uses_finalize(tmp_path):
    sources = _sources(tmp_path)
    positions = make_positions()
    update_table(tmp_path / "t.xlsx", tmp_path / "b.xlsx", positions, sources)
    calls = []

    def fake_finalize(path, duplex):
        calls.append((path, duplex))
        return 10, 5

    results, signer = generate_books(positions, read_table(tmp_path / "t.xlsx"), sources,
                                     {"duplex": True, "blank_rows": 4}, 2027, str(tmp_path), finalize=fake_finalize)
    assert [r.sheet for r in results] == [SHEET_FUNCTIONAL, SHEET_JOB]
    assert [r.path for r in results] == [str(tmp_path / "Функціональні обов'язки 2027.docx"),
                                         str(tmp_path / "Посадові обов'язки 2027.docx")]
    assert all(call[1] for call in calls) and len(calls) == 2
    assert (results[0].pages, results[0].sheets, results[0].blank_rows) == (10, 5, 4)
    functional = Document(results[0].path)
    job = _texts(Document(results[1].path))
    assert "– керувати батальйоном;" in _texts(functional) and "– керувати батальйоном;" not in job
    assert "– підтримувати бойову готовність;" in job
    assert len(functional.tables[-1].rows) == 1 + 6 + 4
    assert review_lines(results, signer) == []


def test_generate_books_falls_back_without_word_and_reports(tmp_path):
    positions = make_positions()
    table = _table(SHEET_JOB, {"Застаріла посада": {SECTION_DUTIES: ["пункт"]}},
                   {positions[3].title: "Набір, якого немає"})
    table.unknown_sections = [(SHEET_JOB, "Солдат", "Покарання"), (SHEET_JOB, "Кухар", "")]

    def failing(path, duplex):
        raise FinalizeError("Word недоступний (немає)")

    results, signer = generate_books(positions, table, Sources(), {}, 2027, str(tmp_path), finalize=failing)
    job_doc = Document(results[1].path)
    assert any("аркушів" in t and "____" in t for t in _texts(job_doc))
    assert job_doc.settings.element.find(qn("w:updateFields")) is not None
    lines = review_lines(results, Signer(), ["попередження"])
    joined = "\n".join(lines)
    assert lines[0] == "- попередження"
    assert "Комбата не знайдено" in joined
    assert "Посади без набору обов'язків" in joined and "ЗАПОВНИТИ" in joined
    assert f"Посада «{positions[3].title}» посилається на набір «Набір, якого немає»" in joined
    job_missing = next(line for line in lines if line.startswith(f"- [{SHEET_JOB}] Посади без набору"))
    assert positions[3].title not in job_missing
    assert "Застаріла посада" in joined
    assert "Набір «Солдат»: невідомий розділ «Покарання»" in joined
    assert "Набір «Кухар»: пункт без розділу" in joined
    assert "Ctrl+A, F9" in joined
    assert "Без ВОС" not in joined


def test_review_lines_list_positions_without_vos(tmp_path):
    positions = group_positions([Slot(1, "Стрілець", "солдат", "ПЕРШИЙ Перший Перший")])
    table = _table(SHEET_FUNCTIONAL, {"Стрілець": {SECTION_DUTIES: ["пункт"]}})
    table.positions[0].sets[SHEET_JOB] = table.add_set(SHEET_JOB, table.set_for(SHEET_FUNCTIONAL, "Стрілець")).name
    results, signer = generate_books(positions, table, Sources(), {}, 2027, str(tmp_path),
                                     finalize=lambda path, duplex: (1, 1))
    lines = review_lines(results, signer)
    assert len(lines) == 3
    assert lines[1].startswith(f"- [{SHEET_FUNCTIONAL}] Без ВОС у штаті (1): Стрілець")
