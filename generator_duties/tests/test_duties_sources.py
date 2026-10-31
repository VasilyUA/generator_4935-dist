import os
import time

from docx import Document

from content.sources import find_resource, load_sources
from duties_fixtures import write_sample, write_statute


def test_load_sources_reads_all_files(tmp_path):
    write_statute(tmp_path / "Статут внутрішньої служби.htm")
    write_sample(tmp_path / "Зразки для Положень.docx")
    regulation = Document()
    regulation.add_paragraph("ПОЛОЖЕННЯ")
    regulation.add_paragraph("")
    regulation.save(tmp_path / "Положення про батальйон.docx")
    sources = load_sources(str(tmp_path))
    assert sources.warnings == []
    assert len(sources.used_files) == 3
    assert sources.statute and sources.sample.blocks
    assert [spec.text for spec in sources.regulation] == ["ПОЛОЖЕННЯ"]


def test_load_sources_warns_about_missing_files(tmp_path):
    sources = load_sources(str(tmp_path))
    assert sources.statute == [] and sources.sample is None and sources.regulation == []
    assert len(sources.warnings) == 3
    assert "zakon.rada.gov.ua" in sources.warnings[0]


def test_load_sources_warns_about_unreadable_files(tmp_path):
    (tmp_path / "Статут.htm").write_bytes(b"\x98")
    (tmp_path / "Зразок.docx").write_bytes(b"broken")
    (tmp_path / "Положення.docx").write_bytes(b"broken")
    sources = load_sources(str(tmp_path))
    assert len(sources.warnings) == 3
    assert all(w.startswith("Не вдалось прочитати") for w in sources.warnings)


def test_find_resource_takes_newest_and_skips_office_temp_files(tmp_path):
    old = tmp_path / "Зразок 1.docx"
    old.write_bytes(b"")
    new = tmp_path / "Зразок 2.docx"
    new.write_bytes(b"")
    (tmp_path / "~$Зразок.docx").write_bytes(b"")
    past = time.time() - 100
    os.utime(old, (past, past))
    assert find_resource(str(tmp_path), "Зраз*.docx") == str(new)
    assert find_resource(str(tmp_path), "Немає*.docx") is None
