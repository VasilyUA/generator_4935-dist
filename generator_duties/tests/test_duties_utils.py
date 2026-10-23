import os
import sys

import pytest
from docx import Document
from docx.oxml.ns import qn

from utils import docx_utils
from utils.docx_utils import add_bookmark, add_paragraph_with_style, add_toc_entry, enable_update_fields_on_open, save_docx_safely
from utils.names import signature_name
from utils.number_words import number_to_words, plural_form, sheets_text


@pytest.mark.parametrize("number, words", [
    (1, "один"), (2, "два"), (11, "одинадцять"), (21, "двадцять один"), (40, "сорок"), (100, "сто"),
    (125, "сто двадцять п’ять"), (212, "двісті дванадцять"), (999, "дев’ятсот дев’яносто дев’ять"),
    (1000, "одна тисяча"), (2001, "дві тисячі один"), (5000, "п’ять тисяч"),
])
def test_number_to_words(number, words):
    assert number_to_words(number) == words


def test_number_to_words_rejects_out_of_range():
    with pytest.raises(ValueError):
        number_to_words(0)


@pytest.mark.parametrize("number, form", [(1, "аркуш"), (21, "аркуш"), (11, "аркушів"), (3, "аркуші"),
                                          (13, "аркушів"), (24, "аркуші"), (5, "аркушів")])
def test_plural_form(number, form):
    assert plural_form(number, "аркуш", "аркуші", "аркушів") == form


def test_sheets_text():
    assert sheets_text(125) == "125 (сто двадцять п’ять) аркушів"


def test_signature_name():
    assert signature_name("ПЕРШИЙ Перший Перший") == "Перший ПЕРШИЙ"
    assert signature_name("Перший Перший") == "Перший Перший"
    assert signature_name("") == ""


def test_toc_entry_has_dotted_tab_and_pageref_field():
    doc = Document()
    paragraph = add_toc_entry(doc, "1. Розділ", "_RefDuty1")
    instr = paragraph._p.findall(".//" + qn("w:instrText"))
    assert instr[0].text.strip() == "PAGEREF _RefDuty1 \\h"
    tab = paragraph._p.find(".//" + qn("w:tab"))
    assert tab.get(qn("w:leader")) == "dot"


def test_bookmark_wraps_paragraph_text_or_empty_paragraph():
    doc = Document()
    paragraph = add_paragraph_with_style(doc, "Заголовок", bold=["Заго"])
    add_bookmark(paragraph, "_RefDuty1", 1)
    children = [child.tag for child in paragraph._p]
    assert children.index(qn("w:bookmarkStart")) < children.index(qn("w:r"))
    assert children[-1] == qn("w:bookmarkEnd")
    empty = doc.add_paragraph()
    add_bookmark(empty, "_RefDuty2", 2)
    assert empty._p[-1].tag == qn("w:bookmarkEnd")
    assert [r.text for r in paragraph.runs] == ["Заго", "ловок"]


def test_bold_phrases_with_text_after_last_phrase():
    doc = Document()
    paragraph = add_paragraph_with_style(doc, "а Б в", bold=["Б", ""], left_indent_cm=1, tab_width=docx_utils.Cm(5))
    assert [(r.text, r.bold) for r in paragraph.runs] == [("а ", False), ("Б", True), (" в", False)]


def test_enable_update_fields_on_open_is_idempotent():
    doc = Document()
    enable_update_fields_on_open(doc)
    enable_update_fields_on_open(doc)
    assert len(doc.settings.element.findall(qn("w:updateFields"))) == 1


def test_save_docx_safely_retries_then_explains(tmp_path, monkeypatch):
    monkeypatch.setitem(sys.modules, "win32com.client", None)
    calls = []

    class _Doc:
        def save(self, path):
            calls.append(path)
            if len(calls) < 2:
                raise PermissionError("locked")

    save_docx_safely(_Doc(), str(tmp_path / "a.docx"), delay=0)
    assert len(calls) == 2

    class _Locked:
        def save(self, path):
            raise PermissionError("locked")

    with pytest.raises(PermissionError, match="Закрийте"):
        save_docx_safely(_Locked(), str(tmp_path / "a.docx"), attempts=2, delay=0)


def test_close_matching_word_documents(monkeypatch, tmp_path):
    pytest.importorskip("win32com.client")
    import win32com.client
    target = os.path.abspath(tmp_path / "a.docx")
    closed = []

    class _Doc:
        def __init__(self, name):
            self.FullName = name

        def Close(self, SaveChanges):
            closed.append(self.FullName)

    class _Broken:
        @property
        def FullName(self):
            raise RuntimeError("gone")

    class _Word:
        Documents = [_Doc(target), _Doc(str(tmp_path / "b.docx")), _Broken()]

    monkeypatch.setattr(win32com.client, "GetActiveObject", lambda name: _Word())
    docx_utils._close_matching_word_documents(lambda name: name == target)
    assert closed == [target]

    def _raise(name):
        raise RuntimeError("no word")
    monkeypatch.setattr(win32com.client, "GetActiveObject", _raise)
    docx_utils._close_matching_word_documents(lambda name: True)
