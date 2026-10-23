import openpyxl
import pytest

from constants import (
    ALL_POSITIONS, SECTION_DUTIES, SECTION_GENERAL, SECTION_RESPONSIBILITY, SECTION_RIGHTS, SHEET_FUNCTIONAL, SHEET_JOB,
    SHEET_POSITIONS, TO_FILL,
)
from content.duties_table import (
    ALL_KEY, DutiesTable, DutySet, PositionRow, SheetDuties, TableError, canonical_section, name_from_sources,
    read_table, write_table,
)


def _set(name, **sections):
    duty_set = DutySet(name)
    for section, texts in sections.items():
        for text in texts:
            duty_set.add(section, text, "Джерело")
    return duty_set


def _table():
    table = DutiesTable()
    for sheet in (SHEET_FUNCTIONAL, SHEET_JOB):
        table.add_set(sheet, _set(ALL_POSITIONS, **{SECTION_RESPONSIBILITY: ["за неналежне виконання;"]}))
        table.add_set(sheet, _set("Солдат", **{SECTION_GENERAL: ["Солдат підпорядковується командиру."],
                                               SECTION_DUTIES: ["знати обов’язки;", "берегти зброю;"]}))
        table.add_set(sheet, _set("Кухар", **{SECTION_DUTIES: [TO_FILL]}))
    table.positions = [PositionRow("Стрілець", {SHEET_FUNCTIONAL: "Солдат", SHEET_JOB: "Солдат"}, 2, 1),
                       PositionRow("Кулеметник", {SHEET_FUNCTIONAL: "Солдат", SHEET_JOB: "Солдат"}, 1, 0),
                       PositionRow("Кухар", {SHEET_FUNCTIONAL: "Кухар", SHEET_JOB: "Кухар"}, 1, 1)]
    return table


def test_write_and_read_round_trip(tmp_path):
    path = tmp_path / "t.xlsx"
    write_table(path, _table())
    table = read_table(path)
    assert not table.legacy
    assert [(r.title, r.sets) for r in table.positions] == [(r.title, r.sets) for r in _table().positions]
    for sheet in (SHEET_FUNCTIONAL, SHEET_JOB):
        assert {k: s.texts() for k, s in table.sets[sheet].items()} == \
               {k: s.texts() for k, s in _table().sets[sheet].items()}
    assert table.set_for(SHEET_JOB, "солдат").sections[SECTION_DUTIES][0].source == "Джерело"


def test_written_sheets_name_each_set_and_section_once(tmp_path):
    path = tmp_path / "t.xlsx"
    write_table(path, _table())
    workbook = openpyxl.load_workbook(path)
    assert workbook.sheetnames == [SHEET_POSITIONS, SHEET_FUNCTIONAL, SHEET_JOB]
    positions = workbook[SHEET_POSITIONS]
    assert [c.value for c in positions[2]] == ["Стрілець", "Солдат", "Солдат", 2, 1]
    ws = workbook[SHEET_JOB]
    rows = [[c.value for c in row] for row in ws.iter_rows(min_row=2)]
    assert rows[0][:3] == [ALL_POSITIONS, SECTION_RESPONSIBILITY, "за неналежне виконання;"]
    assert rows[1] == [None, None, None, None]
    assert rows[2][:3] == ["Солдат", SECTION_GENERAL, "Солдат підпорядковується командиру."]
    assert rows[3][:3] == [None, SECTION_DUTIES, "знати обов’язки;"]
    assert rows[4][:3] == [None, None, "берегти зброю;"]
    assert ws.cell(2, 1).font.bold and ws.cell(4, 1).font.bold and not ws.cell(5, 1).font.bold
    kitchen = next(i for i, row in enumerate(ws.iter_rows(min_row=1), start=1) if row[0].value == "Кухар")
    assert ws.cell(kitchen, 3).fill.fgColor.rgb.endswith("FFFF00")


def test_written_sets_are_framed_and_merge_set_name_and_section(tmp_path):
    path = tmp_path / "t.xlsx"
    write_table(path, _table())
    workbook = openpyxl.load_workbook(path)
    ws = workbook[SHEET_JOB]
    # Рядки: 1 заголовок, 2 "Для всіх", 3 порожній, 4-6 "Солдат" (обов'язки - 5-6), 7 порожній, 8 "Кухар".
    assert {str(merged) for merged in ws.merged_cells.ranges} == {"A4:A6", "B5:B6"}
    assert ws.cell(4, 1).alignment.vertical == "center" and ws.cell(4, 1).alignment.horizontal == "center"
    for row, column in [(1, 1), (2, 4), (4, 3), (6, 4), (8, 1)]:
        border = ws.cell(row, column).border
        assert {border.left.style, border.right.style, border.top.style, border.bottom.style} == {"thin"}
    assert ws.cell(6, 1).border.bottom.style == "thin" and ws.cell(6, 1).border.left.style == "thin"
    assert ws.cell(3, 1).border.left.style is None and ws.cell(7, 3).border.top.style is None
    positions = workbook[SHEET_POSITIONS]
    assert positions.cell(1, 1).border.left.style == "thin"
    assert positions.cell(positions.max_row, 5).border.bottom.style == "thin"


def test_empty_set_is_written_and_read_back(tmp_path):
    table = DutiesTable()
    table.add_set(SHEET_JOB, DutySet("Порожній"))
    path = tmp_path / "t.xlsx"
    write_table(path, table)
    assert read_table(path).set_for(SHEET_JOB, "Порожній").texts() == {}


def test_fill_down_reading_tolerates_inserted_rows_repeated_sets_and_reports_problems(tmp_path):
    path = tmp_path / "t.xlsx"
    workbook = openpyxl.Workbook()
    workbook.active.title = SHEET_POSITIONS
    workbook.active.append(["Посада", "Функціональні обов'язки (набір)", "Посадові обов'язки (набір)"])
    workbook.active.append(["Стрілець", "", "Солдат"])
    ws = workbook.create_sheet(SHEET_JOB)
    ws.append(["Набір обов'язків", "Розділ", "Текст", "Джерело"])
    ws.append([None, None, "пункт до першого набору", None])
    ws.append(["Солдат", None, "пункт без розділу", None])
    ws.append([None, None, "ще пункт без розділу", None])
    ws.append([None, "Обов’язки", "знати обов’язки;", None])
    ws.append([None, None, "вставлений рядок;", None])
    ws.append([None, "Покарання", "щось", None])
    ws.append([None, None, "теж пропущено", None])
    ws.append(["Кухар", "Права", "отримувати продукти;", None])
    ws.append(["солдат", "Права", "звертатися до командира;", None])
    workbook.save(path)
    table = read_table(path)
    assert table.positions[0].sets == {SHEET_FUNCTIONAL: "", SHEET_JOB: "Солдат"}
    assert table.set_for(SHEET_JOB, "Солдат").texts() == {
        SECTION_DUTIES: ["знати обов’язки;", "вставлений рядок;"], SECTION_RIGHTS: ["звертатися до командира;"],
    }
    assert table.unknown_sections == [(SHEET_JOB, "", ""), (SHEET_JOB, "Солдат", ""), (SHEET_JOB, "Солдат", "Покарання")]
    assert table.sets[SHEET_FUNCTIONAL] == {}


def test_read_without_headers_uses_default_column_order(tmp_path):
    path = tmp_path / "t.xlsx"
    workbook = openpyxl.Workbook()
    workbook.active.title = SHEET_POSITIONS
    workbook.active.append(["Стрілець", "Солдат", "Солдат"])
    ws = workbook.create_sheet(SHEET_FUNCTIONAL)
    ws.append(["Солдат", "Обов'язки", "знати обов’язки;"])
    workbook.create_sheet(SHEET_JOB)
    workbook.save(path)
    table = read_table(path)
    assert table.positions[0].sets == {SHEET_FUNCTIONAL: "Солдат", SHEET_JOB: "Солдат"}
    assert table.set_for(SHEET_FUNCTIONAL, "Солдат").texts() == {SECTION_DUTIES: ["знати обов’язки;"]}
    assert table.sets[SHEET_JOB] == {}


def test_read_table_with_empty_positions_sheet(tmp_path):
    path = tmp_path / "t.xlsx"
    workbook = openpyxl.Workbook()
    workbook.active.title = SHEET_POSITIONS
    workbook.save(path)
    table = read_table(path)
    assert table.positions == [] and table.sets == {SHEET_FUNCTIONAL: {}, SHEET_JOB: {}}


def _legacy(path, sheets):
    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)
    for name, rows in sheets.items():
        ws = workbook.create_sheet(name)
        for row in rows:
            ws.append(list(row))
    workbook.save(path)
    return path


def test_legacy_table_is_converted_into_positions_and_shared_sets(tmp_path):
    header = ("Посада", "Розділ", "Текст", "Джерело")
    path = _legacy(tmp_path / "t.xlsx", {
        SHEET_JOB: [header,
                    (ALL_POSITIONS, "Відповідальність", "за неналежне виконання;", "Статут ВС, ст. 26, 27"),
                    ("Стрілець", "Обов'язки", "знати обов’язки;", "Статут ВС, ст. 127, 128"),
                    ("Кулеметник", "Обов'язки", "знати обов’язки;", "Статут ВС, ст. 127, 128"),
                    ("Водій", "Обов'язки", "берегти автомобіль;", "Зразок: Водій автомобільного відділення"),
                    ("Кухар", "Обов'язки", TO_FILL, TO_FILL),
                    ("Кухар", "Покарання", "щось", ""),
                    ("Писар", "Права", "", "")],
        SHEET_FUNCTIONAL: [("Стрілець", "Обов'язки", "пункт", "")],
    })
    table = read_table(path, {"Статут ВС, ст. 127, 128": "Солдат (Статут)"})
    assert table.legacy
    assert [r.title for r in table.positions] == ["Стрілець", "Кулеметник", "Водій", "Кухар", "Писар"]
    assert table.positions[0].sets == {SHEET_JOB: "Солдат (Статут)", SHEET_FUNCTIONAL: "Стрілець"}
    assert table.positions[1].sets == {SHEET_JOB: "Солдат (Статут)"}
    assert table.positions[2].sets[SHEET_JOB] == "Водій автомобільного відділення (зразок)"
    assert table.positions[3].sets[SHEET_JOB] == "Кухар"
    assert table.set_for(SHEET_JOB, ALL_POSITIONS).texts() == {SECTION_RESPONSIBILITY: ["за неналежне виконання;"]}
    assert table.unknown_sections == [(SHEET_JOB, "Кухар", "Покарання")]


def test_legacy_table_without_header_and_with_empty_sheet(tmp_path):
    path = _legacy(tmp_path / "t.xlsx", {SHEET_JOB: [("Стрілець", "Обов'язки", "пункт", "")], SHEET_FUNCTIONAL: []})
    table = read_table(path)
    assert table.positions[0].sets == {SHEET_JOB: "Стрілець"}


def test_add_set_reuses_identical_content_but_not_to_fill_and_makes_names_unique():
    table = DutiesTable()
    first = table.add_set(SHEET_JOB, _set("Солдат", **{SECTION_DUTIES: ["знати;"]}))
    same = DutySet("Інша назва")
    same.add(SECTION_DUTIES, "Знати", "інше джерело")
    assert table.add_set(SHEET_JOB, same) is first
    other = table.add_set(SHEET_JOB, _set("Солдат", **{SECTION_DUTIES: ["інше;"]}))
    assert other.name == "Солдат (2)"
    blank_one = table.add_set(SHEET_JOB, _set("Кухар", **{SECTION_DUTIES: [TO_FILL]}))
    blank_two = table.add_set(SHEET_JOB, _set("Кухар", **{SECTION_DUTIES: [TO_FILL]}))
    assert blank_one is not blank_two and blank_two.name == "Кухар (2)"


def test_add_set_never_merges_into_common_set():
    table = DutiesTable()
    table.add_set(SHEET_JOB, _set(ALL_POSITIONS, **{SECTION_DUTIES: ["знати;"]}))
    assert table.add_set(SHEET_JOB, _set("Солдат", **{SECTION_DUTIES: ["знати;"]})).name == "Солдат"


def test_duty_set_flags_and_users():
    assert DutySet("Порожній").is_to_fill and DutySet("Порожній").has_to_fill
    mixed = _set("Мішаний", **{SECTION_DUTIES: ["пункт", TO_FILL]})
    assert not mixed.is_to_fill and mixed.has_to_fill
    table = _table()
    assert [r.title for r in table.users(SHEET_JOB, table.set_for(SHEET_JOB, "Солдат"))] == ["Стрілець", "Кулеметник"]
    assert table.position("кухар").title == "Кухар" and table.position("писар") is None
    assert table.set_for(SHEET_JOB, "") is None


def test_name_from_sources():
    duty_set = DutySet("x")
    assert name_from_sources(duty_set, "Посада") == "Посада"
    duty_set.add(SECTION_GENERAL, TO_FILL, TO_FILL)
    duty_set.add(SECTION_GENERAL, "пункт", "")
    assert name_from_sources(duty_set, "Посада") == "Посада"
    duty_set.add(SECTION_DUTIES, "a", "Статут ВС, ст. 1")
    duty_set.add(SECTION_DUTIES, "b", "Статут ВС, ст. 1")
    duty_set.add(SECTION_RIGHTS, "c", "Зразок: Водій")
    duty_set.add(SECTION_RIGHTS, "d", "Інше")
    assert name_from_sources(duty_set, "Посада", {"Статут ВС, ст. 1": "Солдат (Статут)"}) == \
        "Солдат (Статут) + Водій (зразок) + Інше"


def test_canonical_section_accepts_case_and_apostrophe_variants():
    assert canonical_section(" обов’язки ") == SECTION_DUTIES
    assert canonical_section("ОБОВ'ЯЗКИ") == SECTION_DUTIES
    assert canonical_section("Покарання") is None


def test_sheet_duties_merges_common_set_before_own_set():
    table = _table()
    table.positions.append(PositionRow("Писар", {SHEET_JOB: "Немає такого"}))
    table.unknown_sections = [(SHEET_JOB, "Солдат", "Покарання"), (SHEET_FUNCTIONAL, "Кухар", "")]
    duties = SheetDuties(table, SHEET_JOB)
    assert duties.for_position("стрілець") == {
        SECTION_GENERAL: ["Солдат підпорядковується командиру."],
        SECTION_DUTIES: ["знати обов’язки;", "берегти зброю;"],
        SECTION_RESPONSIBILITY: ["за неналежне виконання;"],
    }
    assert duties.for_position("невідома") == {SECTION_RESPONSIBILITY: ["за неналежне виконання;"]}
    assert duties.has_position("кухар") and not duties.has_position("писар") and not duties.has_position("невідома")
    assert duties.missing_set_name("писар") == "Немає такого"
    assert duties.missing_set_name("стрілець") == "" and duties.missing_set_name("невідома") == ""
    assert duties.unknown_sections == [(SHEET_JOB, "Солдат", "Покарання")]
    del table.sets[SHEET_JOB][ALL_KEY]
    assert SheetDuties(table, SHEET_JOB).common == {}


def test_read_table_reports_broken_file(tmp_path):
    path = tmp_path / "t.xlsx"
    path.write_bytes(b"broken")
    with pytest.raises(TableError):
        read_table(path)


def test_write_table_reports_locked_file(tmp_path, monkeypatch):
    def _raise(self, path):
        raise PermissionError("locked")
    monkeypatch.setattr(openpyxl.Workbook, "save", _raise)
    with pytest.raises(TableError, match="закрийте"):
        write_table(tmp_path / "t.xlsx", DutiesTable())
