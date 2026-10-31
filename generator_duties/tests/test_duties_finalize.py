import pytest
from docx import Document

from constants import SHEETS_BLANK, SHEETS_MARKER
from generators import word_finalize
from generators.word_finalize import FinalizeError, finalize_with_word, finalize_without_word, sheets_for_pages


class _FakeFind:
    def __init__(self):
        self.executed = []

    def ClearFormatting(self):
        pass

    @property
    def Replacement(self):
        return self

    def Execute(self, *args):
        self.executed.append(args)


class _FakeDoc:
    def __init__(self, pages, fail=False):
        self.pages = pages
        self.fail = fail
        self.find = _FakeFind()
        self.saved = self.closed = False
        self.fields_updates = 0

    @property
    def Fields(self):
        return self

    def Update(self):
        self.fields_updates += 1

    def Repaginate(self):
        if self.fail:
            raise RuntimeError("COM error")

    def ComputeStatistics(self, kind):
        assert kind == word_finalize.WD_STATISTIC_PAGES
        return self.pages

    @property
    def Content(self):
        return self

    @property
    def Find(self):
        return self.find

    def Save(self):
        self.saved = True

    def Close(self, save):
        self.closed = True


class _FakeWord:
    def __init__(self, doc):
        self.doc = doc
        self.opened = []
        self.quit = False

    @property
    def Documents(self):
        return self

    def Open(self, path, AddToRecentFiles):
        self.opened.append(path)
        return self.doc

    def Quit(self):
        self.quit = True


def _opener(word, released):
    return lambda: (word, lambda: released.append(True))


def test_sheets_for_pages():
    assert sheets_for_pages(125, False) == 125
    assert sheets_for_pages(125, True) == 63


def test_finalize_with_word_updates_fields_counts_and_replaces_marker(tmp_path):
    doc = _FakeDoc(125)
    word = _FakeWord(doc)
    released = []
    assert finalize_with_word(tmp_path / "b.docx", duplex=True, open_word=_opener(word, released)) == (125, 63)
    replace = doc.find.executed[0]
    assert replace[0] == SHEETS_MARKER
    assert replace[9] == "63 (шістдесят три) аркуші"
    assert replace[10] == word_finalize.WD_REPLACE_ALL
    assert doc.fields_updates == 2 and doc.saved and doc.closed and word.quit and released


def test_finalize_with_word_wraps_com_errors(tmp_path):
    word = _FakeWord(_FakeDoc(1, fail=True))
    released = []
    with pytest.raises(FinalizeError, match="COM error"):
        finalize_with_word(tmp_path / "b.docx", open_word=_opener(word, released))
    assert word.doc.closed and released


def test_finalize_with_word_reports_missing_word(tmp_path):
    def _no_word():
        raise ImportError("no pywin32")
    with pytest.raises(FinalizeError, match="недоступний"):
        finalize_with_word(tmp_path / "b.docx", open_word=_no_word)


def test_finalize_with_word_ignores_errors_while_closing(tmp_path):
    doc = _FakeDoc(3)
    word = _FakeWord(doc)

    def _quit():
        raise RuntimeError("already closed")
    word.Quit = _quit
    assert finalize_with_word(tmp_path / "b.docx", open_word=_opener(word, [])) == (3, 3)


def test_open_word_starts_hidden_separate_instance(monkeypatch):
    pytest.importorskip("win32com.client")
    import pythoncom
    import win32com.client

    class _App:
        Visible = True
        DisplayAlerts = 1

    created = []
    monkeypatch.setattr(pythoncom, "CoInitialize", lambda: None)
    monkeypatch.setattr(win32com.client, "DispatchEx", lambda name: created.append(name) or _App())
    word, release = word_finalize._open_word()
    assert created == ["Word.Application"]
    assert word.Visible is False and word.DisplayAlerts == 0
    assert release is pythoncom.CoUninitialize


def test_finalize_without_word_replaces_marker_and_requests_field_update(tmp_path):
    path = tmp_path / "b.docx"
    doc = Document()
    doc.add_paragraph(SHEETS_MARKER)
    doc.add_paragraph("інший текст")
    doc.save(path)
    finalize_without_word(path)
    result = Document(path)
    assert result.paragraphs[0].text == f"{SHEETS_BLANK} аркушів"
    assert result.paragraphs[1].text == "інший текст"
