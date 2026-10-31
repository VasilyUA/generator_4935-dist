import openpyxl
import pytest

from constants import ALL_POSITIONS, SECTION_DUTIES, SHEET_FUNCTIONAL, SHEET_JOB, SHEET_POSITIONS, TO_FILL
from content.duties_table import DutySet, PositionRow, read_table, write_table
from content.sample_reader import read_sample
from content.sources import Sources
from content.statute_reader import parse_statute_html
from duties_fixtures import STATUTE_HTML, make_positions, write_sample
from generators.update_table import stale_positions, update_table

SOLDIER = "Загальні обов’язки військовослужбовців рядового складу (Статут)"


@pytest.fixture
def sources(tmp_path):
    return Sources(statute=parse_statute_html(STATUTE_HTML), sample=read_sample(write_sample(tmp_path / "s.docx")))


def _rows_with_gunner():
    from duties_fixtures import ROSTER_ROWS, SUFFIX
    return ROSTER_ROWS + ((7, f"Кулеметник 2 відділення 3 взводу {SUFFIX}", "1000006", "солдат", ""),)


def test_first_run_creates_positions_and_shared_sets(tmp_path, sources):
    path, backup = tmp_path / "t.xlsx", tmp_path / "b.xlsx"
    positions = make_positions(_rows_with_gunner())
    result = update_table(path, backup, positions, sources)
    assert result.created and not result.migrated and not backup.exists()
    assert result.added == [p.title for p in positions]
    table = read_table(path)
    assert [row.title for row in table.positions] == [p.title for p in positions]
    rifleman, gunner = table.positions[3], table.positions[5]
    assert rifleman.sets == gunner.sets == {SHEET_FUNCTIONAL: SOLDIER, SHEET_JOB: SOLDIER}
    counts = [cell.value for cell in openpyxl.load_workbook(path)[SHEET_POSITIONS][5][3:5]]
    assert counts == [2, 1]
    names = list(table.sets[SHEET_JOB])
    assert names[0] == "для всіх"
    assert len([key for key in names if key == SOLDIER.lower()]) == 1
    assert result.to_fill == {SHEET_FUNCTIONAL: [], SHEET_JOB: []}
    assert result.set_counts[SHEET_JOB] == len(table.sets[SHEET_JOB]) - 1


def test_second_run_keeps_user_edits_and_backs_up(tmp_path, sources):
    path, backup = tmp_path / "t.xlsx", tmp_path / "b.xlsx"
    positions = make_positions()
    update_table(path, backup, positions, sources)
    table = read_table(path)
    soldier = table.set_for(SHEET_JOB, SOLDIER)
    soldier.sections[SECTION_DUTIES][0].text = "мій текст"
    write_table(path, table)
    result = update_table(path, backup, positions, sources)
    assert not result.created and backup.exists()
    assert result.added == [] and result.refilled == []
    assert read_table(path).set_for(SHEET_JOB, SOLDIER).sections[SECTION_DUTIES][0].text == "мій текст"


def test_to_fill_sets_are_refilled_when_sources_appear(tmp_path, sources):
    path, backup = tmp_path / "t.xlsx", tmp_path / "b.xlsx"
    positions = make_positions()
    result = update_table(path, backup, positions, Sources())
    assert result.to_fill[SHEET_JOB] == [p.title for p in positions]
    assert len(read_table(path).sets[SHEET_JOB]) == len(positions) + 1
    result = update_table(path, backup, positions, sources)
    assert result.refilled == [p.title for p in positions]
    assert result.to_fill[SHEET_JOB] == []
    table = read_table(path)
    assert all(not duty_set.is_to_fill for duty_set in table.sets[SHEET_JOB].values())
    assert table.set_for(SHEET_JOB, ALL_POSITIONS).items()


def test_to_fill_sets_stay_when_sources_still_have_nothing(tmp_path):
    path, backup = tmp_path / "t.xlsx", tmp_path / "b.xlsx"
    positions = make_positions()
    update_table(path, backup, positions, Sources())
    result = update_table(path, backup, positions, Sources())
    assert result.refilled == []


def test_missing_set_reference_is_filled_again(tmp_path, sources):
    path, backup = tmp_path / "t.xlsx", tmp_path / "b.xlsx"
    positions = make_positions()
    update_table(path, backup, positions, sources)
    table = read_table(path)
    table.positions[0].sets[SHEET_JOB] = "Набір, якого немає"
    write_table(path, table)
    update_table(path, backup, positions, sources)
    assert read_table(path).positions[0].sets[SHEET_JOB] == "Командир батальйону (Статут)"


def test_shared_to_fill_set_is_not_refilled_for_one_position(tmp_path, sources):
    path, backup = tmp_path / "t.xlsx", tmp_path / "b.xlsx"
    positions = make_positions()
    update_table(path, backup, positions, sources)
    table = read_table(path)
    shared = table.add_set(SHEET_JOB, DutySet("Спільний"))
    shared.add(SECTION_DUTIES, TO_FILL, TO_FILL)
    table.positions[0].sets[SHEET_JOB] = table.positions[1].sets[SHEET_JOB] = "Спільний"
    write_table(path, table)
    result = update_table(path, backup, positions, sources)
    assert result.refilled == []
    assert read_table(path).positions[0].sets[SHEET_JOB] == "Спільний"


def test_stale_positions_are_removed_only_when_confirmed_with_their_sets(tmp_path, sources):
    path, backup = tmp_path / "t.xlsx", tmp_path / "b.xlsx"
    positions = make_positions()
    update_table(path, backup, positions, sources)
    table = read_table(path)
    table.add_set(SHEET_JOB, DutySet("Створений вручну"))
    write_table(path, table)
    driver_set = table.positions[-1].sets[SHEET_FUNCTIONAL]
    fewer = positions[:-1]
    assert stale_positions(read_table(path), fewer) == [positions[-1].title]
    result = update_table(path, backup, fewer, sources, remove_stale=False)
    assert result.stale_kept == [positions[-1].title]
    assert read_table(path).positions[-1].title == positions[-1].title
    result = update_table(path, backup, fewer, sources, remove_stale=True)
    assert result.removed == [positions[-1].title]
    table = read_table(path)
    assert positions[-1].title not in [row.title for row in table.positions]
    assert table.set_for(SHEET_JOB, "Створений вручну") is not None
    # Набір солдата потрібен іншим посадам - лишається.
    assert table.set_for(SHEET_FUNCTIONAL, driver_set) is not None


def test_set_used_only_by_removed_position_disappears(tmp_path, sources):
    path, backup = tmp_path / "t.xlsx", tmp_path / "b.xlsx"
    positions = make_positions()
    update_table(path, backup, positions, sources)
    commander_set = read_table(path).positions[0].sets[SHEET_JOB]
    update_table(path, backup, positions[1:], sources, remove_stale=True)
    assert read_table(path).set_for(SHEET_JOB, commander_set) is None


def test_existing_common_set_is_kept(tmp_path, sources):
    path, backup = tmp_path / "t.xlsx", tmp_path / "b.xlsx"
    common = DutySet(ALL_POSITIONS)
    common.add(SECTION_DUTIES, "спільний пункт")
    from content.duties_table import DutiesTable
    table = DutiesTable()
    table.add_set(SHEET_JOB, common)
    write_table(path, table)
    update_table(path, backup, make_positions(), sources)
    assert read_table(path).set_for(SHEET_JOB, ALL_POSITIONS).texts() == {SECTION_DUTIES: ["спільний пункт"]}


def test_legacy_table_is_migrated_and_renamed_titles_are_matched(tmp_path, sources):
    path, backup = tmp_path / "t.xlsx", tmp_path / "b.xlsx"
    positions = make_positions()
    workbook = openpyxl.Workbook()
    ws = workbook.active
    ws.title = SHEET_JOB
    ws.append(["Посада", "Розділ", "Текст", "Джерело"])
    ws.append(["Начальник штабу-заступник командира батальйону", "Обов'язки", TO_FILL, TO_FILL])
    ws.append(["Стрілець відділення взводу", "Обов'язки", "власний пункт", ""])
    workbook.save(path)
    result = update_table(path, backup, positions, sources)
    assert result.migrated
    table = read_table(path)
    assert SHEET_POSITIONS in openpyxl.load_workbook(path).sheetnames
    chief = table.position(positions[1].key)
    assert chief.title == positions[1].title
    assert chief.sets[SHEET_JOB] == "Начальник штабу батальйону (Статут)"
    assert positions[1].title in result.refilled
    rifleman = table.position(positions[3].key)
    assert table.set_for(SHEET_JOB, rifleman.sets[SHEET_JOB]).texts() == {SECTION_DUTIES: ["власний пункт"]}


def test_existing_row_keeps_set_and_new_position_is_added(tmp_path, sources):
    path, backup = tmp_path / "t.xlsx", tmp_path / "b.xlsx"
    positions = make_positions()
    update_table(path, backup, positions[:2], sources)
    table = read_table(path)
    assert [row.title for row in table.positions] == [p.title for p in positions[:2]]
    result = update_table(path, backup, positions, sources)
    assert result.added == [p.title for p in positions[2:]]
    assert isinstance(read_table(path).positions[0], PositionRow)
