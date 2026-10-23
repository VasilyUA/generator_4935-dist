"""Таблиця обов'язків resources/Обов'язки.xlsx - ЄДИНЕ місце текстів.

Аркуш "Посади": кожна посада штату один раз (у порядку штату) і назва набору
обов'язків для кожної книги. Аркуші "Функціональні" і "Посадові": кожен набір
один раз - назва лише в першому рядку набору, розділ - у першому рядку
розділу, далі пункти (рядок = один пункт). Посади з однаковими обов'язками
посилаються на один набір - текст не дублюється. Набір "Для всіх" додається
до кожної посади. Колонка "Джерело" - лише підказка для людини.

Таблицю старого формату (лише аркуші пунктів з колонкою "Посада") read_table
перетворює сама (legacy=True), щоб наступний запис був уже в новому форматі."""
import os
from dataclasses import dataclass, field

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from constants import (
    ALL_POSITIONS, COLUMN_POSITION, COLUMN_SECTION, COLUMN_SET, COLUMN_SOURCE, COLUMN_TEXT, POSITION_COLUMNS,
    SECTIONS, SET_COLUMN_BY_SHEET, SET_COLUMNS, SHEET_FUNCTIONAL, SHEET_JOB, SHEET_POSITIONS, TO_FILL,
)
from content.positions import compare_key

SHEET_NAMES = (SHEET_FUNCTIONAL, SHEET_JOB)
ALL_KEY = compare_key(ALL_POSITIONS)
_SECTION_BY_KEY = {compare_key(name): name for name in SECTIONS}
_POSITION_WIDTHS = (70, 55, 55, 12, 12)
_SET_WIDTHS = (45, 20, 110, 30)
_HEADER_FILL = PatternFill("solid", fgColor="D9D9D9")
_TO_FILL_FILL = PatternFill("solid", fgColor="FFFF00")
_THIN = Side(style="thin")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_LEGACY_COLUMNS = (COLUMN_POSITION, COLUMN_SECTION, COLUMN_TEXT, COLUMN_SOURCE)


class TableError(Exception):
    pass


def canonical_section(name):
    """"обов’язки" / "Обов'язки " → "Обов'язки"; невідомий розділ → None."""
    return _SECTION_BY_KEY.get(compare_key(name))


def is_to_fill(text):
    return text.strip().upper() == TO_FILL


@dataclass
class Item:
    text: str
    source: str = ""


@dataclass
class DutySet:
    name: str
    sections: dict = field(default_factory=dict)

    @property
    def key(self):
        return compare_key(self.name)

    def add(self, section, text, source=""):
        self.sections.setdefault(section, []).append(Item(text, source))

    def items(self):
        return [item for section in SECTIONS for item in self.sections.get(section, [])]

    @property
    def is_to_fill(self):
        """Порожній набір або лише "ЗАПОВНИТИ"."""
        return all(is_to_fill(item.text) for item in self.items())

    @property
    def has_to_fill(self):
        return not self.items() or any(is_to_fill(item.text) for item in self.items())

    def signature(self):
        """Зміст набору для порівняння (джерело й назва не важать)."""
        return tuple((section, tuple(compare_key(item.text) for item in self.sections.get(section, [])))
                     for section in SECTIONS if self.sections.get(section))

    def texts(self):
        return {section: [item.text for item in items] for section, items in self.sections.items() if items}


@dataclass
class PositionRow:
    title: str
    sets: dict = field(default_factory=dict)
    total: int = 0
    vacant: int = 0

    @property
    def key(self):
        return compare_key(self.title)


@dataclass
class DutiesTable:
    positions: list = field(default_factory=list)
    sets: dict = field(default_factory=lambda: {name: {} for name in SHEET_NAMES})
    unknown_sections: list = field(default_factory=list)
    legacy: bool = False

    def position(self, key):
        return next((row for row in self.positions if row.key == key), None)

    def set_for(self, sheet, name):
        return self.sets[sheet].get(compare_key(name or "")) if name else None

    def users(self, sheet, duty_set):
        return [row for row in self.positions if compare_key(row.sets.get(sheet, "")) == duty_set.key]

    def add_set(self, sheet, duty_set):
        """Додає набір; якщо такий самий зміст уже є - повертає наявний (текст
        не дублюється). Набори лише з "ЗАПОВНИТИ" не об'єднуються - кожна
        посада заповнює свій. Назва робиться унікальною: "Назва (2)"."""
        sets = self.sets[sheet]
        if not duty_set.is_to_fill:
            signature = duty_set.signature()
            for key, existing in sets.items():
                if key != ALL_KEY and existing.signature() == signature:
                    return existing
        base, number = duty_set.name, 1
        while duty_set.key in sets:
            number += 1
            duty_set.name = f"{base} ({number})"
        sets[duty_set.key] = duty_set
        return duty_set


def name_from_sources(duty_set, fallback, label_names=None):
    """Назва набору за джерелами його пунктів: "Командир відділення (Статут)",
    "Водій … (зразок)", кілька джерел - через " + "; без джерел - fallback."""
    label_names = label_names or {}
    names = []
    for item in duty_set.items():
        if not item.source or is_to_fill(item.text):
            continue
        if item.source in label_names:
            name = label_names[item.source]
        elif item.source.startswith("Зразок: "):
            name = f"{item.source[len('Зразок: '):]} (зразок)"
        else:
            name = item.source
        if name not in names:
            names.append(name)
    return " + ".join(names) if names else fallback


def _cell(value):
    return "" if value is None else str(value).strip()


def _columns(header, names):
    lowered = [_cell(value).lower() for value in header]
    found = {name: lowered.index(name.lower()) for name in names if name.lower() in lowered}
    return found


def _getter(values, columns):
    def get(name):
        index = columns.get(name)
        return _cell(values[index]) if index is not None and index < len(values) else ""
    return get


def _read_positions(ws):
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    columns = _columns(rows[0], POSITION_COLUMNS)
    if COLUMN_POSITION not in columns:
        columns = {name: i for i, name in enumerate(POSITION_COLUMNS)}
        rows = [None] + rows
    result = []
    for values in rows[1:]:
        get = _getter(values, columns)
        if get(COLUMN_POSITION):
            result.append(PositionRow(get(COLUMN_POSITION),
                                      {sheet: get(column) for sheet, column in SET_COLUMN_BY_SHEET.items()}))
    return result


def _read_sets(ws, sheet, unknown):
    sets = {}
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return sets
    columns = _columns(rows[0], SET_COLUMNS)
    if COLUMN_TEXT not in columns:
        columns = {name: i for i, name in enumerate(SET_COLUMNS)}
        rows = [None] + rows
    current, section, section_seen = None, None, False
    for values in rows[1:]:
        get = _getter(values, columns)
        if get(COLUMN_SET):
            current = sets.setdefault(compare_key(get(COLUMN_SET)), DutySet(get(COLUMN_SET)))
            section, section_seen = None, False
        if get(COLUMN_SECTION):
            section, section_seen = canonical_section(get(COLUMN_SECTION)), True
            if section is None:
                unknown.append((sheet, current.name if current else "", get(COLUMN_SECTION)))
        if not get(COLUMN_TEXT):
            continue
        if current is not None and section is not None:
            current.add(section, get(COLUMN_TEXT), get(COLUMN_SOURCE))
        elif not section_seen:
            # Пункт без розділу (чи без набору) - повідомляється один раз.
            unknown.append((sheet, current.name if current else "", ""))
            section_seen = True
    return sets


def _read_legacy_rows(ws):
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    columns = _columns(rows[0], _LEGACY_COLUMNS)
    if COLUMN_POSITION not in columns or COLUMN_TEXT not in columns:
        columns = {name: i for i, name in enumerate(_LEGACY_COLUMNS)}
        rows = [None] + rows
    result = []
    for values in rows[1:]:
        get = _getter(values, columns)
        if get(COLUMN_POSITION) or get(COLUMN_TEXT):
            result.append((get(COLUMN_POSITION), get(COLUMN_SECTION), get(COLUMN_TEXT), get(COLUMN_SOURCE)))
    return result


def _from_legacy(rows_by_sheet, label_names):
    """Старий формат (рядок = посада + пункт) → посади + набори без дублювання."""
    table = DutiesTable(legacy=True)
    rows_by_key = {}
    for sheet in SHEET_NAMES:
        groups = {}
        for position, section_name, text, source in rows_by_sheet.get(sheet, []):
            section = canonical_section(section_name)
            if section is None:
                table.unknown_sections.append((sheet, position, section_name))
                continue
            group = groups.setdefault(compare_key(position), DutySet(position))
            if text:
                group.add(section, text, source)
        common = groups.pop(ALL_KEY, None)
        if common is not None:
            common.name = ALL_POSITIONS
            table.sets[sheet][ALL_KEY] = common
        for key, content in groups.items():
            title = content.name
            content.name = name_from_sources(content, title, label_names)
            chosen = table.add_set(sheet, content)
            if key not in rows_by_key:
                rows_by_key[key] = PositionRow(title)
                table.positions.append(rows_by_key[key])
            rows_by_key[key].sets[sheet] = chosen.name
    return table


def read_table(path, label_names=None):
    """→ DutiesTable. label_names - назви наборів за джерелом (для перетворення
    старого формату), див. content.prefill.label_names."""
    try:
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except Exception as e:
        raise TableError(f"Не вдалось відкрити '{os.path.basename(path)}': {e}") from e
    try:
        by_name = {ws.title.strip().lower(): ws for ws in workbook.worksheets}
        if SHEET_POSITIONS.lower() not in by_name:
            return _from_legacy({name: _read_legacy_rows(by_name[name.lower()]) for name in SHEET_NAMES
                                 if name.lower() in by_name}, label_names)
        table = DutiesTable(positions=_read_positions(by_name[SHEET_POSITIONS.lower()]))
        for name in SHEET_NAMES:
            if name.lower() in by_name:
                table.sets[name] = _read_sets(by_name[name.lower()], name, table.unknown_sections)
        return table
    finally:
        workbook.close()


def _style_header(ws, widths):
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = _HEADER_FILL
        cell.alignment = Alignment(vertical="center", horizontal="center", wrap_text=True)
        cell.border = _BORDER
    for index, width in enumerate(widths, start=1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(index)].width = width
    ws.freeze_panes = "A2"


def _draw_borders(ws, rows, columns):
    for row in rows:
        for column in range(1, columns + 1):
            ws.cell(row, column).border = _BORDER


def _write_positions(ws, table):
    ws.append(list(POSITION_COLUMNS))
    for row in table.positions:
        ws.append([row.title] + [row.sets.get(sheet) or None for sheet in SHEET_NAMES] + [row.total, row.vacant])
    for cells in ws.iter_rows(min_row=2):
        for cell in cells:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    _draw_borders(ws, range(2, ws.max_row + 1), len(POSITION_COLUMNS))
    _style_header(ws, _POSITION_WIDTHS)
    ws.auto_filter.ref = ws.dimensions


def _write_sets(ws, sets):
    """Набір = блок рядків у рамці; назва набору й розділ об'єднані по
    вертикалі (значення лишається в першому рядку - read_table читає так
    само), між наборами - порожній рядок без рамки."""
    ws.append(list(SET_COLUMNS))
    blocks = []
    for duty_set in sets.values():
        set_rows, section_blocks = [], []
        for section in SECTIONS:
            section_rows = []
            for index, item in enumerate(duty_set.sections.get(section, [])):
                ws.append([None if set_rows else duty_set.name, section if index == 0 else None, item.text,
                           item.source or None])
                set_rows.append(ws.max_row)
                section_rows.append(ws.max_row)
                if is_to_fill(item.text):
                    ws.cell(ws.max_row, 3).fill = _TO_FILL_FILL
            section_blocks.append(section_rows)
        if not set_rows:
            ws.append([duty_set.name])
            set_rows.append(ws.max_row)
        blocks.append((set_rows, section_blocks))
        ws.append([])
    for cells in ws.iter_rows(min_row=2):
        for cell in cells:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    _style_header(ws, _SET_WIDTHS)
    for set_rows, section_blocks in blocks:
        name_cell = ws.cell(set_rows[0], 1)
        name_cell.font = Font(bold=True)
        name_cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        # Рамка - до об'єднання: merge_cells переносить її з першої клітинки на краї.
        _draw_borders(ws, set_rows, len(SET_COLUMNS))
        for rows, column in [(set_rows, 1)] + [(rows, 2) for rows in section_blocks]:
            if len(rows) > 1:
                ws.merge_cells(start_row=rows[0], start_column=column, end_row=rows[-1], end_column=column)


def write_table(path, table):
    """Файл перезаписується повністю: "Посади", потім аркуші наборів."""
    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)
    _write_positions(workbook.create_sheet(SHEET_POSITIONS), table)
    for name in SHEET_NAMES:
        _write_sets(workbook.create_sheet(name), table.sets[name])
    try:
        workbook.save(path)
    except PermissionError as e:
        raise TableError(f"Не вдалось записати '{os.path.basename(path)}' - закрийте його в Excel і спробуйте ще раз.") from e


class SheetDuties:
    """Аркуш таблиці для генерації книги: посада → "Для всіх" + її набір."""

    def __init__(self, table, sheet):
        self.sheet = sheet
        self.table = table
        common = table.sets[sheet].get(ALL_KEY)
        self.common = common.texts() if common else {}
        self.rows = {row.key: row for row in table.positions}
        self.unknown_sections = [entry for entry in table.unknown_sections if entry[0] == sheet]

    def _own_set(self, key):
        row = self.rows.get(key)
        return self.table.set_for(self.sheet, row.sets.get(self.sheet, "")) if row else None

    def has_position(self, key):
        return self._own_set(key) is not None

    def missing_set_name(self, key):
        """Назва набору, вказана для посади, якої немає на аркуші (інакше "")."""
        row = self.rows.get(key)
        name = row.sets.get(self.sheet, "") if row else ""
        return name if name and self._own_set(key) is None else ""

    def for_position(self, key):
        """{розділ: [пункти]} - спершу "Для всіх", потім набір посади."""
        own_set = self._own_set(key)
        own = own_set.texts() if own_set else {}
        return {section: self.common.get(section, []) + own.get(section, []) for section in SECTIONS
                if self.common.get(section) or own.get(section)}
