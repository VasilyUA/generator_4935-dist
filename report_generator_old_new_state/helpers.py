import os, time, re, shutil, pandas as pd

from datetime import datetime
from colorama import Fore, Style
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml.ns import qn
from docx.shared import Pt, Inches
from constants import RANKS, FULL_MILITARY_UNIT

#____________________________________________________    
def get_rows(data, columns):
    return [
        {col: row[col] for col in columns}
        for _, row in data[columns].iterrows()
    ]

def print_red(s):
    print(Fore.RED + s + Style.RESET_ALL)

def print_green(s):
    print(Fore.GREEN + s + Style.RESET_ALL)

def print_timeout(s, t):
    print('ROW:')
    time.sleep(t)
    print(s)

def set_line_spacing(doc, spacing):
    for paragraph in doc.paragraphs:
        paragraph.paragraph_format.line_spacing = Pt(spacing)
        

def create_or_clear_output_directory(directory):
    if os.path.exists(directory):
        shutil.rmtree(directory)
    os.makedirs(directory)

def add_paragraph_with_style(doc, text, font_name='Times New Roman', font_size=14, alignment=WD_ALIGN_PARAGRAPH.JUSTIFY, first_line_indent=0, format_tabs=False):
    paragraph = doc.add_paragraph()
    run = paragraph.add_run(text)
    run.font.name = font_name
    run._element.rPr.rFonts.set(qn('w:eastAsia'), font_name)
    run.font.size = Pt(font_size)

    paragraph.alignment = alignment
    paragraph.paragraph_format.line_spacing = 1.0  # Одинарний міжрядковий
    paragraph.paragraph_format.space_after = Pt(0)  # Взагалі нуль після абзацу

    if first_line_indent > 0:
        paragraph.paragraph_format.first_line_indent = Pt(first_line_indent)

    if format_tabs:
        width = Inches(doc.sections[0].page_width.inches - (doc.sections[0].left_margin.inches + doc.sections[0].right_margin.inches))
        tab_stops = paragraph.paragraph_format.tab_stops
        tab_stops.add_tab_stop(width, WD_TAB_ALIGNMENT.RIGHT)

    return paragraph


def excel_col_to_index(col):
    """Перетворює Excel-літеру (наприклад, 'A', 'Z', 'AA', 'AD') у індекс (0-відлік)."""
    col = col.upper()
    index = 0
    for c in col:
        index = index * 26 + (ord(c) - ord('A') + 1)
    return index - 1

def set_margins(doc, top, bottom, left, right):
    sections = doc.sections
    for section in sections:
        section.page_height = Inches(11.69)  # 297 мм
        section.page_width = Inches(8.27)    # 210 мм
        section.top_margin = top
        section.bottom_margin = bottom
        section.left_margin = left
        section.right_margin = right

def int_or_none(value):
    try:
        return int(value)
    except (ValueError, TypeError):
        return None


def convert_to_short_name(full_name, id_position=None):
    if not isinstance(full_name, str):
        return ''

    # 1. Виправляємо апостроф + пробіл -> зліплюємо назад
    cleaned = re.sub(r"(В|в)'[\s]+([а-яА-ЯІіЇїЄєҐґ])", r"\1'\2", full_name.strip())

    # 2. Прибираємо зайві пробіли
    cleaned = ' '.join(cleaned.split())

    # 3. Розбити на слова
    words = cleaned.split()

    # 4. Якщо менше ніж 3 — пробуємо розклеїти друге слово (ім’я+по батькові)
    if len(words) == 2:
        # Пробуємо вставити пробіл між ім’ям і по батькові
        possible_fix = re.sub(r'([а-яґєіїʼ’]+)([А-ЯІЇЄҐ])', r'\1 \2', words[1])
        fixed_words = [words[0]] + possible_fix.split()
        if len(fixed_words) == 3:
            words = fixed_words

    # 5. Перевірка
    if len(words) != 3:
        raise ValueError(f"❌ Некоректний ПІБ для {id_position}: \"{full_name}\" → після обробки: \"{' '.join(words)}\"")

    # 6. Повертаємо скорочене ім’я
    return f"{words[1]} {words[0]}"

def get_position_code(row):
    return {
        "subordinate_position_code_from": int_or_none(row['НОМЕР ПОСАДИ З ЯКОЇ ПЕРЕМІЩУЄТЬСЯ ОСОБА']), # A
        "subordinate_position_code_to": int_or_none(row['НОМЕР ПОСАДИ НА ЯКУ ПЕРЕМІЩУЄТЬСЯ ОСОБА']), # B
        "commander_position_code_in_old_state": int_or_none(row['КОМАНДИР В СТАРІЙ ШТАТІ ЯКИЙ КЛОПОЧЕ КОМАНДИРУ БАТАЛЬЙОНУ']), # C
        "commander_position_code_in_new_state": int_or_none(row['КОМАНДИР В НОВОМУ ШТАТІ ЯКИЙ КЛОПОЧЕ КОМАНДИРУ БАТАЛЬЙОНУ']), # D
        "higher_commander_position_code_in_old_state": int_or_none(row['КОМАНДИР БАТАЛЬЙОНУ В СТАРІЙ ШТАТІ']), # E
        "higher_commander_position_code_in_new_state": int_or_none(row['КОМАНДИР БАТАЛЬЙОНУ В НОВОМУ ШТАТІ']), # F
    }

def get_commander_row(position_codes, rows_personel_old, rows_personel_new, status_state, rows_tvo, rows_task, date):
    date_col = to_date(date)
    row = {}

    if status_state == "old":
        row = next((r for r in rows_personel_old if r["№ з.п."] == position_codes[f"commander_position_code_in_old_state"]), {})

    if status_state == "new":
        row = next((r for r in rows_personel_new if r["№ з.п."] == position_codes[f"commander_position_code_in_new_state"]), {})
    
    name = row.get("ПІБ", "") if row else ""
    if pd.isna(name) or name.strip() != "":
        tvo = next((
            r for r in rows_tvo
            if (
                row is not None
                and r['ПІБ'] == row.get('ПІБ')
                and to_date(r['Start']) <= to_date(date_col) <= to_date(r['End'])
                and r.get(f'{status_state} TVO Active', None) is True
            )
        ), {})

        tvo_code = tvo.get(f'№{status_state}', None)
        if tvo_code is not None:
            task_name = next((r for r in rows_task if r["ПІБ називний"] == tvo.get("ПІБ", "")), {})
            task_position = next((r for r in rows_task if r["Посада називний"] == tvo.get("ПОСАДА", "")), {})

            commander = {
                "ПІБ": row.get("ПІБ", ""),
                "Посада": tvo.get("ПОСАДА", ""),
                "ТВО": tvo.get(f'{status_state} TVO Active', False),
                "підрозділ повністю": row.get("підрозділ повністю", ""),
                'звання фактичне': row.get('звання фактичне', ''),
                'Посада давальний': task_position.get("Посада давальний", ""),
                'посада називний': task_position.get("посада називний", ""),
                'ПІБ родовий': task_name.get('ПІБ родовий', '')
            }

            return commander

    task_name = next((r for r in rows_task if r["ПІБ називний"] == row.get("ПІБ", "")), {})
    task_position = next((r for r in rows_task if r["Посада називний"] == row.get("Посада", "")), {})
    commander = {
            "ПІБ": row.get("ПІБ", ""),
            "Посада": row.get('Посада', ""),
            "підрозділ повністю": row.get("підрозділ повністю", ""),
            'звання фактичне': row.get('звання фактичне', ''),
            "ТВО":  False,
            'Посада давальний': task_position.get("Посада давальний", ""),
            'посада називний': task_position.get("посада називний", ""),
            'ПІБ родовий': task_name.get('ПІБ родовий', '')
        }
    
    return commander

def get_subordinate(position_codes, rows_personel_old, rows_personel_new, rows_task, rows_tvo, date):
    data_for_generation = {
        "old": {
            "personal_row": {}, 
            "commander_row": None, 
            "higher_commander_row": {}
        },
        "new": {
            "personal_row": {}, 
            "commander_row": None, 
            "higher_commander_row": {}
        },
        "task": {
            "personal_row": {}, 
            "higher_commander_row": {}
        }
    }

    data_for_generation["old"]["personal_row"] = next((r for r in rows_personel_old if r["№ з.п."] == position_codes["subordinate_position_code_from"]), {})
    data_for_generation["old"]["commander_row"] = get_commander_row(position_codes, rows_personel_old, rows_personel_new, "old", rows_tvo, rows_task, date)
    data_for_generation["old"]["higher_commander_row"] = next((r for r in rows_personel_old if r["№ з.п."] == position_codes["higher_commander_position_code_in_old_state"]), {})

    data_for_generation["new"]["personal_row"] = next((r for r in rows_personel_new if r["№ з.п."] == position_codes["subordinate_position_code_to"]), {})
    data_for_generation["new"]["commander_row"] = get_commander_row(position_codes, rows_personel_old, rows_personel_new, "new", rows_tvo, rows_task, date)
    data_for_generation["new"]["higher_commander_row"] = next((r for r in rows_personel_new if r["№ з.п."] == position_codes["higher_commander_position_code_in_new_state"]), {})
    
   
    data_for_generation["task"]["personal_row"] = next((r for r in rows_task if r["ПІБ називний"] == data_for_generation["new"]["personal_row"]["ПІБ"]), {})
    data_for_generation["task"]["higher_commander_row"] = next((r for r in rows_task if r["ПІБ називний"] == data_for_generation["new"]["higher_commander_row"]["ПІБ"]), {})


    return data_for_generation
    
from datetime import datetime, date

def to_date(x):
    if isinstance(x, datetime):
        return x.date()
    elif isinstance(x, date):  # якщо вже date
        return x
    elif isinstance(x, str):
        # список можливих форматів
        formats = [
            "%Y-%m-%d",     # 2025-08-22
            "%d.%m.%Y",     # 22.08.2025
            "%d/%m/%Y",     # 22/08/2025
            "%m/%d/%Y",     # 08/22/2025
            "%d-%m-%Y",     # 22-08-2025
            "%Y/%m/%d",     # 2025/08/22
            "%d %b %Y",     # 22 Aug 2025
            "%d %B %Y",     # 22 August 2025
        ]
        for fmt in formats:
            try:
                return datetime.strptime(x.strip(), fmt).date()
            except ValueError:
                continue
    raise ValueError(f"Невідомий формат або тип дати: {x}")

def find_higher_commander(rows_tvo, row, date):
    date_col = to_date(date)
 
    commander = next(
        (r for r in rows_tvo
         if r['ПОСАДА'] == row['Посада'] and r.get('ТВО', False) is True
         and to_date(r['Start']) <= date_col <= to_date(r['End'])),
        {}
    )

    if not commander:
        return row

    return commander

def get_tvo_position(rows_tvo, commander, date=datetime.now()):
    commander_tvo = find_higher_commander(rows_tvo, commander, date)
    is_tvo = commander_tvo.get('ТВО', False)

    role = "Тимчасово виконуючий обов'язки командира батальйону" if is_tvo else "Командир батальйону"
    return f"{role} {FULL_MILITARY_UNIT}"

def get_tvo_name(rows_tvo, commander, date=datetime.now()):
    commander_tvo = find_higher_commander(rows_tvo, commander, date)
    is_tvo = commander_tvo.get('ТВО', False)

    return f"{commander['звання фактичне']}\t{convert_to_short_name(commander_tvo['ПІБ']) if is_tvo else convert_to_short_name(commander['ПІБ'])}"


def get_text(data_for_generation):
    rank_old = RANKS.get(data_for_generation.get("old", {}).get("personal_row", {}).get("звання фактичне", ""), "")
    name_r = data_for_generation.get("task", {}).get("personal_row", {}).get("ПІБ родовий", "")
    Position = data_for_generation.get("old", {}).get("personal_row", {}).get("Посада", "").lower()
    Full_position = data_for_generation.get("old", {}).get("personal_row", {}).get("підрозділ повністю", "")
    vos = data_for_generation.get("old", {}).get("personal_row", {}).get("ВОС", "")
    rank_for_state = data_for_generation.get("old", {}).get("personal_row", {}).get("звання за штатом", "")
    number_old = int_or_none(data_for_generation.get("old", {}).get("personal_row", {}).get("№ з.п.", ""))

    old_text = f'{rank_old} {name_r}, {Position} {Full_position} {FULL_MILITARY_UNIT}, ВОС - {vos}, ШПК "{rank_for_state}" ({number_old});'

    
    Position_new = data_for_generation["new"]["personal_row"]["Посада"].lower()
    Full_position_new = data_for_generation["new"]["personal_row"]["підрозділ повністю"]
    vos_new = data_for_generation["new"]["personal_row"]["ВОС"]
    rank_for_state_new = data_for_generation["new"]["personal_row"]["звання за штатом"]
    number_new = int_or_none(data_for_generation["new"]["personal_row"]["№ з.п."])

    new_text = f'на посаду {Position_new} {Full_position_new} {FULL_MILITARY_UNIT}, ВОС - {vos_new}, ШПК "{rank_for_state_new}" ({number_new});'

    return f"{old_text} {new_text}"

from datetime import datetime, date

def input_date(prompt="Введіть дату (формат YYYY-MM-DD або DD.MM.YYYY): "):
    while True:
        user_input = input(prompt).strip()
        
        # Якщо пусто — підставляємо сьогодні
        if not user_input:
            today = date.today()
            print(f"Вибрано сьогоднішню дату: {today}")
            return today
        
        # Список можливих форматів
        formats = ["%Y-%m-%d", "%d.%m.%Y", "%d/%m/%Y", "%d-%m-%Y"]
        
        for fmt in formats:
            try:
                parsed_date = datetime.strptime(user_input, fmt).date()
                print(f"Вибрано дату: {parsed_date}")
                return parsed_date
            except ValueError:
                continue
        
        print("❌ Невірний формат. Спробуйте ще раз.")
