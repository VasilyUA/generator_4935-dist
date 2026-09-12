"""Генерує МІНІМАЛЬНИЙ приклад resources/СТАРИЙ_СПИСОК.xlsx +
resources/НОВИЙ_СПИСОК.xlsx + resources/ЗАДАЧІ.xlsm + mass_movement.xlsx
(у КОРЕНІ проєкту, не в resources/) - показує точну структуру колонок на
1 вигаданому переміщенні. Не чіпає файли, які вже існують.
Запуск: python resources_example/generate_sample.py (з кореня проєкту)."""
import json
import os

from openpyxl import Workbook

ROOT_DIR = os.path.join(os.path.dirname(__file__), "..")
RESOURCES_DIR = os.path.join(ROOT_DIR, "resources")
DATA_JSON_PATH = os.path.join(RESOURCES_DIR, "data.json")
OLD_PATH = os.path.join(RESOURCES_DIR, "СТАРИЙ_СПИСОК.xlsx")
NEW_PATH = os.path.join(RESOURCES_DIR, "НОВИЙ_СПИСОК.xlsx")
ZADACHI_PATH = os.path.join(RESOURCES_DIR, "ЗАДАЧІ.xlsm")
MASS_MOVEMENT_PATH = os.path.join(ROOT_DIR, "mass_movement.xlsx")

# Вигадані люди - НІКОЛИ не вставляйте сюди реальні ПІБ. Один умовний
# приклад: солдат переміщується з посади 1 (старий штат) на посаду 2 (новий
# штат), обидва рапорти клопоче/підписує один і той самий командир
# батальйону (позиція 99 у старому штаті, 98 - у новому).
SUBORDINATE_PIB = "ПЕРШИЙ Перший Першович"
COMMANDER_PIB = "ДРУГИЙ Другий Другович"


def _personel_row(letters_map, values):
    """letters_map: {літера: значення заголовка}. Повертає рядок довжиною до
    останньої потрібної літери, решту клітинок лишає порожніми."""
    max_col = max(_col_index(l) for l in letters_map)
    row = [""] * (max_col + 1)
    for letter, value in zip(letters_map, values):
        row[_col_index(letter)] = value
    return row


def _col_index(letter):
    idx = 0
    for ch in letter:
        idx = idx * 26 + (ord(ch.upper()) - ord("A") + 1)
    return idx - 1


def _build_personel_list(path, letters, sheet_name):
    """letters - словник {поле: літера}, у порядку підрозділ/№/посада/ВОС/звання за штатом/звання фактичне/ПІБ."""
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name
    header = [""] * (max(_col_index(l) for l in letters.values()) + 1)
    for field, letter in letters.items():
        header[_col_index(letter)] = field
    ws.append(header)
    for pos_id, pidrozdil, posada, pib in (
        (1 if path == OLD_PATH else 2, "Підрозділ 1", "Стрілець", SUBORDINATE_PIB),
        (99 if path == OLD_PATH else 98, "управління", "Командир батальйону", COMMANDER_PIB),
    ):
        row = [""] * len(header)
        row[_col_index(letters["підрозділ повністю"])] = pidrozdil
        row[_col_index(letters["№ з.п."])] = pos_id
        row[_col_index(letters["Посада"])] = posada
        row[_col_index(letters["ВОС"])] = "0000"
        row[_col_index(letters["звання за штатом"])] = "солдат" if posada == "Стрілець" else "майор"
        row[_col_index(letters["звання фактичне"])] = "солдат" if posada == "Стрілець" else "майор"
        row[_col_index(letters["ПІБ"])] = pib
        ws.append(row)
    wb.save(path)


def main():
    os.makedirs(RESOURCES_DIR, exist_ok=True)

    if not os.path.exists(DATA_JSON_PATH):
        with open(DATA_JSON_PATH, "w", encoding="utf-8") as f:
            json.dump({
                "unit": {
                    "PERSONEL_LIST_SHEET_NAME": "Штат",
                    "FULL_MILITARY_UNIT": "військової частини А0000",
                }
            }, f, ensure_ascii=False, indent=2)
        print(f"Створено {DATA_JSON_PATH} (заповніть реальними даними).")
    sheet_name = json.load(open(DATA_JSON_PATH, encoding="utf-8"))["unit"]["PERSONEL_LIST_SHEET_NAME"]

    if not os.path.exists(OLD_PATH):
        _build_personel_list(OLD_PATH, {
            "підрозділ повністю": "F", "№ з.п.": "J", "Посада": "K",
            "ВОС": "M", "звання за штатом": "N", "звання фактичне": "O", "ПІБ": "P",
        }, sheet_name)
        print(f"Створено {OLD_PATH} (аркуш '{sheet_name}').")
    else:
        print(f"{OLD_PATH} вже існує - пропускаю.")

    if not os.path.exists(NEW_PATH):
        _build_personel_list(NEW_PATH, {
            "підрозділ повністю": "P", "№ з.п.": "Q", "Посада": "R",
            "ВОС": "T", "звання за штатом": "U", "звання фактичне": "V", "ПІБ": "W",
        }, sheet_name)
        print(f"Створено {NEW_PATH} (аркуш '{sheet_name}').")
    else:
        print(f"{NEW_PATH} вже існує - пропускаю.")

    if not os.path.exists(ZADACHI_PATH):
        wb = Workbook()
        ws1 = wb.active
        ws1.title = "Аркуш1"
        ws1.append([
            "Посада називний", "посада називний", "ПІБ називний",
            "Посада родовий", "посада родовий", "ПІБ родовий",
            "Посада давальний", "посада давальний", "ПІБ давальний",
        ])
        for posada, pib in (("Стрілець", SUBORDINATE_PIB), ("Командир батальйону", COMMANDER_PIB)):
            ws1.append([posada, posada.lower(), pib, posada, posada.lower(), pib, posada, posada.lower(), pib])

        ws2 = wb.create_sheet("Аркуш2")  # ТВО
        ws2.append(["Підрозділ", "ПОСАДА", "Start", "End", "ЗВАННЯ", "ПІБ", "ТВО", "№old", "№new", "old TVO Active", "new TVO Active"])
        # Порожній (без рядків) - так само коректно, як і заповнений.

        wb.save(ZADACHI_PATH)
        print(f"Створено {ZADACHI_PATH} (Аркуш1 + порожній Аркуш2 'ТВО').")
    else:
        print(f"{ZADACHI_PATH} вже існує - пропускаю.")

    if not os.path.exists(MASS_MOVEMENT_PATH):
        wb = Workbook()
        ws = wb.active
        ws.title = "Аркуш1"
        ws.append([
            "НОМЕР ПОСАДИ З ЯКОЇ ПЕРЕМІЩУЄТЬСЯ ОСОБА", "НОМЕР ПОСАДИ НА ЯКУ ПЕРЕМІЩУЄТЬСЯ ОСОБА",
            "КОМАНДИР В СТАРІЙ ШТАТІ ЯКИЙ КЛОПОЧЕ КОМАНДИРУ БАТАЛЬЙОНУ",
            "КОМАНДИР В НОВОМУ ШТАТІ ЯКИЙ КЛОПОЧЕ КОМАНДИРУ БАТАЛЬЙОНУ",
            "КОМАНДИР БАТАЛЬЙОНУ В СТАРІЙ ШТАТІ", "КОМАНДИР БАТАЛЬЙОНУ В НОВОМУ ШТАТІ",
        ])
        ws.append([1, 2, 99, 98, 99, 98])
        wb.save(MASS_MOVEMENT_PATH)
        print(f"Створено {MASS_MOVEMENT_PATH}")
    else:
        print(f"{MASS_MOVEMENT_PATH} вже існує - пропускаю.")

    print("Замініть вигаданих людей на реальних.")


if __name__ == "__main__":
    main()
