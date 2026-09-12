import os

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches
from helpers import set_margins, set_line_spacing, add_paragraph_with_style
from constants import OUTPUT_DIR, RANKS, FULL_MILITARY_UNIT

from helpers import (
    get_text, 
    get_position_code,
    print_red, 
    add_paragraph_with_style,
    get_subordinate,
    set_margins,
    set_line_spacing,
    get_tvo_name,
    get_tvo_position
)
from validator import DataValidation

def change_position_report(rows_move, rows_personel_old, rows_personel_new, rows_task, rows_tvo, date):
    doc_main_mass_moved = Document()

    for index, row in enumerate(rows_move):
        position_codes = get_position_code(row)

        validator = DataValidation({
            "position_codes": position_codes,
            "rows_personel_old": rows_personel_old,
            "rows_personel_new": rows_personel_new,
            "rows_task": rows_task
        })
        is_valid = validator.is_valid()

        if not is_valid:
            print_red(f"Валідація не пройшла для даних: {position_codes}")
            continue

        # структурування даних
        data_for_generation = get_subordinate(position_codes, rows_personel_old, rows_personel_new, rows_task, rows_tvo, date)
        # print(data_for_generation)

        if index == 0:
            set_margins(doc_main_mass_moved, Inches(0.786), Inches(0.786), Inches(1.18), Inches(0.395))
            set_line_spacing(doc_main_mass_moved, 1.0)  # встановлення міжрядкового інтервалу 1
            add_paragraph_with_style(doc_main_mass_moved, f"Командиру {FULL_MILITARY_UNIT}", alignment=WD_ALIGN_PARAGRAPH.RIGHT)
            add_paragraph_with_style(doc_main_mass_moved, "")
            add_paragraph_with_style(doc_main_mass_moved, "")
            add_paragraph_with_style(doc_main_mass_moved, "")
            add_paragraph_with_style(doc_main_mass_moved, "РАПОРТ", alignment=WD_ALIGN_PARAGRAPH.CENTER)
            add_paragraph_with_style(doc_main_mass_moved, "")
            add_paragraph_with_style(doc_main_mass_moved, "")
            add_paragraph_with_style(doc_main_mass_moved, f"Прошу Вашого дозволу для підвищення бойової готовності 1 батальйону морської піхоти {FULL_MILITARY_UNIT} нижче поіменований особовий склад від займаних посад ЗВІЛЬНИТИ ТА ПРИЗНАЧИТИ:", first_line_indent=34)
            add_paragraph_with_style(doc_main_mass_moved, "")
        
     
        text = get_text(data_for_generation)
        add_paragraph_with_style(
            doc_main_mass_moved, 
            text,
            first_line_indent=34
        )

        if index + 1 == len(rows_move):
            add_paragraph_with_style(doc_main_mass_moved, "")
            add_paragraph_with_style(doc_main_mass_moved, "")
            add_paragraph_with_style(doc_main_mass_moved, get_tvo_position(rows_tvo, data_for_generation["new"]["higher_commander_row"]))
            add_paragraph_with_style(doc_main_mass_moved, get_tvo_name(rows_tvo, data_for_generation["new"]["higher_commander_row"]), format_tabs=True)
            filename = f"РАПОРТ НА ПЕРЕМІЩЕННЯ.docx"
            doc_main_mass_moved.save(os.path.join(OUTPUT_DIR, filename))
            return filename