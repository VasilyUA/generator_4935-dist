import os

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches
from helpers import (set_margins, set_line_spacing, add_paragraph_with_style, convert_to_short_name, get_tvo_name, get_tvo_position)
from constants import OUTPUT_DIR, RANKS, FULL_MILITARY_UNIT

def generate_accepted_position_report(data_for_generation, commander_position_code, rows_tvo):
    doc = Document()
    set_margins(doc, Inches(0.786), Inches(0.786), Inches(1.18), Inches(0.395))
    set_line_spacing(doc, 14)
    if commander_position_code:
        add_paragraph_with_style(doc, data_for_generation.get("new", {}).get("commander_row", {}).get('Посада давальний', '').capitalize(), first_line_indent=271)
    else:
        add_paragraph_with_style(doc, "Командиру батальйону", first_line_indent=271)

    add_paragraph_with_style(doc, "", font_name="Times New Roman", font_size=14)
    add_paragraph_with_style(doc, "", font_name="Times New Roman", font_size=14)
    add_paragraph_with_style(doc, "РАПОРТ", alignment=WD_ALIGN_PARAGRAPH.CENTER)
    add_paragraph_with_style(doc, "", font_name="Times New Roman", font_size=14)
    add_paragraph_with_style(doc, f"Дійсним доповідаю, що справи та посаду {data_for_generation["task"]["personal_row"]['посада родовий'].lower()} {data_for_generation["new"]["personal_row"]['підрозділ повністю']} {FULL_MILITARY_UNIT} здав.", first_line_indent=34)
    add_paragraph_with_style(doc, "Матеріальних цінностей за посадою не рахується.", first_line_indent=34)
    add_paragraph_with_style(doc, f"{data_for_generation["new"]["personal_row"]['Посада']} {data_for_generation["new"]["personal_row"]['підрозділ повністю']} {FULL_MILITARY_UNIT}")
    add_paragraph_with_style(doc, f"{data_for_generation["new"]["personal_row"]['звання фактичне']}\t{convert_to_short_name(data_for_generation["new"]["personal_row"]['ПІБ'])}", format_tabs=True)
    add_paragraph_with_style(doc, "", font_name="Times New Roman", font_size=14)
    add_paragraph_with_style(doc, "", font_name="Times New Roman", font_size=14)
    if commander_position_code:
        add_paragraph_with_style(doc, "Командиру батальйону", first_line_indent=251)
        add_paragraph_with_style(doc, "", font_name="Times New Roman", font_size=14)
        add_paragraph_with_style(doc, f"Доповідаю по суті рапорту {RANKS[data_for_generation["new"]["personal_row"]['звання фактичне']]} {convert_to_short_name(data_for_generation["task"]["personal_row"]["ПІБ родовий"])}", first_line_indent=34)
        add_paragraph_with_style(doc, "", font_name="Times New Roman", font_size=14)
        add_paragraph_with_style(doc, f"{data_for_generation["new"]["commander_row"]['посада називний']} {data_for_generation["new"]["commander_row"]['підрозділ повністю']} {FULL_MILITARY_UNIT}")
        add_paragraph_with_style(doc, f"{data_for_generation["new"]["commander_row"]['звання фактичне']}\t{convert_to_short_name(data_for_generation["new"]["commander_row"]['ПІБ'])}", format_tabs=True)
        add_paragraph_with_style(doc, "", font_name="Times New Roman", font_size=14)
        add_paragraph_with_style(doc, "", font_name="Times New Roman", font_size=14)
        add_paragraph_with_style(doc, f"Командиру {FULL_MILITARY_UNIT}", first_line_indent=251)
        add_paragraph_with_style(doc, "", font_name="Times New Roman", font_size=14)
        add_paragraph_with_style(doc, f"Доповідаю по суті рапорту {RANKS[data_for_generation["new"]["commander_row"]['звання фактичне']]} {convert_to_short_name(data_for_generation["new"]["commander_row"]["ПІБ родовий"])}", first_line_indent=34)
        add_paragraph_with_style(doc, "", font_name="Times New Roman", font_size=14)
        add_paragraph_with_style(doc, get_tvo_position(rows_tvo, data_for_generation["new"]["higher_commander_row"]))
        add_paragraph_with_style(doc, get_tvo_name(rows_tvo, data_for_generation["new"]["higher_commander_row"]), format_tabs=True)
    else:
        add_paragraph_with_style(doc, f"Командиру {FULL_MILITARY_UNIT}", first_line_indent=251)
        add_paragraph_with_style(doc, "", font_name="Times New Roman", font_size=14)
        add_paragraph_with_style(doc, f"Доповідаю по суті рапорту {RANKS[data_for_generation["new"]["personal_row"]['звання фактичне']]} {convert_to_short_name(data_for_generation["task"]["personal_row"]["ПІБ родовий"])}", first_line_indent=34)
        add_paragraph_with_style(doc, "", font_name="Times New Roman", font_size=14)
        add_paragraph_with_style(doc, get_tvo_position(rows_tvo, data_for_generation["new"]["higher_commander_row"]))
        add_paragraph_with_style(doc, get_tvo_name(rows_tvo, data_for_generation["new"]["higher_commander_row"]), format_tabs=True)
    doc_name_3 = os.path.join(OUTPUT_DIR, f"{data_for_generation["new"]["personal_row"]['ПІБ']} прийняв посаду.docx")
    doc.save(doc_name_3)

    return doc_name_3
