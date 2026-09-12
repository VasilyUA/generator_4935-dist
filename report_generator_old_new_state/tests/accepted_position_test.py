import sys
from pathlib import Path

sys.path.append(str(Path(__file__).resolve().parent.parent))

from docx import Document

import constants
from template import accepted_position as m


def _data_for_generation():
    return {
        "task": {"personal_row": {"посада родовий": "стрільця", "ПІБ родовий": "ЧЕТВЕРТИЙ Четвертий Четвертий"}},
        "new": {
            "personal_row": {
                "Посада": "Навідник", "підрозділ повністю": "2 рота",
                "звання фактичне": "сержант", "ПІБ": "ПЕРШИЙ Перший Перший",
            },
            "commander_row": {
                "посада називний": "командир роти", "підрозділ повністю": "1 рота",
                "звання фактичне": "капітан", "ПІБ": "ДРУГИЙ Другий Другий",
                "ПІБ родовий": "ДРУГИЙ Другий Другий",
            },
            "higher_commander_row": {
                "Посада": "Командир батальйону", "звання фактичне": "підполковник",
                "ПІБ": "ТРЕТІЙ Третій Третій",
            },
        },
    }


def test_generate_accepted_position_report_with_commander(tmp_path, monkeypatch):
    monkeypatch.setattr(m, "OUTPUT_DIR", str(tmp_path))

    doc_name = m.generate_accepted_position_report(_data_for_generation(), commander_position_code=7, rows_tvo=[])

    assert doc_name == str(tmp_path / "ПЕРШИЙ Перший Перший прийняв посаду.docx")
    text = "\n".join(p.text for p in Document(doc_name).paragraphs)
    assert "здав." in text
    assert constants.FULL_MILITARY_UNIT in text
    assert "Другий ДРУГИЙ" in text
    assert "Третій ТРЕТІЙ" in text


def test_generate_accepted_position_report_without_commander_falls_back(tmp_path, monkeypatch):
    monkeypatch.setattr(m, "OUTPUT_DIR", str(tmp_path))

    doc_name = m.generate_accepted_position_report(_data_for_generation(), commander_position_code=None, rows_tvo=[])

    text = "\n".join(p.text for p in Document(doc_name).paragraphs)
    assert f"Командиру {constants.FULL_MILITARY_UNIT}" in text
    assert "Другий ДРУГИЙ" not in text
    assert "Третій ТРЕТІЙ" in text
