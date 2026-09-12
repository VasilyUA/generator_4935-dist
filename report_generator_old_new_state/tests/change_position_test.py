import sys
from pathlib import Path

sys.path.append(str(Path(__file__).resolve().parent.parent))

from docx import Document

import constants
from template import change_position as m


def _row_move(from_id, to_id, commander_id, higher_id):
    return {
        "НОМЕР ПОСАДИ З ЯКОЇ ПЕРЕМІЩУЄТЬСЯ ОСОБА": from_id,
        "НОМЕР ПОСАДИ НА ЯКУ ПЕРЕМІЩУЄТЬСЯ ОСОБА": to_id,
        "КОМАНДИР В СТАРІЙ ШТАТІ ЯКИЙ КЛОПОЧЕ КОМАНДИРУ БАТАЛЬЙОНУ": commander_id,
        "КОМАНДИР В НОВОМУ ШТАТІ ЯКИЙ КЛОПОЧЕ КОМАНДИРУ БАТАЛЬЙОНУ": commander_id,
        "КОМАНДИР БАТАЛЬЙОНУ В СТАРІЙ ШТАТІ": higher_id,
        "КОМАНДИР БАТАЛЬЙОНУ В НОВОМУ ШТАТІ": higher_id,
    }


def _personel_row(pos_id, pib, posada="Стрілець", rank="сержант"):
    return {
        "№ з.п.": pos_id, "ПІБ": pib, "Посада": posada,
        "підрозділ повністю": "1 рота", "звання фактичне": rank,
        "звання за штатом": rank, "ВОС": "000000А/000",
    }


def _task_row(pos_id, pib):
    return {
        "ПІБ називний": pib, "Посада називний": "Навідник", "посада називний": "навідник",
        "посада родовий": "навідника", "ПІБ родовий": pib, "Посада давальний": "Навіднику",
    }


def _build_data(from_id, to_id, commander_id, higher_id, from_pib, to_pib, commander_pib, higher_pib):
    rows_personel_old = [
        _personel_row(from_id, from_pib),
        _personel_row(commander_id, commander_pib, posada="Командир роти", rank="капітан"),
        _personel_row(higher_id, higher_pib, posada="Командир батальйону", rank="підполковник"),
    ]
    rows_personel_new = [
        _personel_row(to_id, to_pib),
        _personel_row(commander_id, commander_pib, posada="Командир роти", rank="капітан"),
        _personel_row(higher_id, higher_pib, posada="Командир батальйону", rank="підполковник"),
    ]
    rows_task = [_task_row(to_id, to_pib)]
    return rows_personel_old, rows_personel_new, rows_task


def test_change_position_report_single_row_writes_header_and_returns_filename(tmp_path, monkeypatch):
    monkeypatch.setattr(m, "OUTPUT_DIR", str(tmp_path))
    rows_personel_old, rows_personel_new, rows_task = _build_data(
        1, 2, 3, 4, "ВОСЬМИЙ Восьмий Восьмий", "ДВАДЦЯТИЙ Двадцятий Двадцятий",
        "ДЕСЯТИЙ Десятий Десятий", "ОДИНАДЦЯТИЙ Одинадцятий Одинадцятий",
    )
    rows_move = [_row_move(1, 2, 3, 4)]

    filename = m.change_position_report(rows_move, rows_personel_old, rows_personel_new, rows_task, [], "15.08.2026")

    assert filename == "РАПОРТ НА ПЕРЕМІЩЕННЯ.docx"
    saved_path = tmp_path / filename
    assert saved_path.exists()

    text = "\n".join(p.text for p in Document(str(saved_path)).paragraphs)
    assert "ЗВІЛЬНИТИ ТА ПРИЗНАЧИТИ" in text
    assert constants.FULL_MILITARY_UNIT in text
    assert "Одинадцятий ОДИНАДЦЯТИЙ" in text


def test_change_position_report_multiple_rows_header_appears_once(tmp_path, monkeypatch):
    monkeypatch.setattr(m, "OUTPUT_DIR", str(tmp_path))
    old1, new1, task1 = _build_data(1, 2, 5, 6, "ДВАНАДЦЯТИЙ Дванадцятий Дванадцятий", "ТРИНАДЦЯТИЙ Тринадцятий Тринадцятий", "ЧОТИРНАДЦЯТИЙ Чотирнадцятий Чотирнадцятий", "ШІСТНАДЦЯТИЙ Шістнадцятий Шістнадцятий")
    old2, new2, task2 = _build_data(3, 4, 5, 6, "СІМНАДЦЯТИЙ Сімнадцятий Сімнадцятий", "ВІСІМНАДЦЯТИЙ Вісімнадцятий Вісімнадцятий", "ЧОТИРНАДЦЯТИЙ Чотирнадцятий Чотирнадцятий", "ШІСТНАДЦЯТИЙ Шістнадцятий Шістнадцятий")
    rows_personel_old = old1 + [r for r in old2 if r["№ з.п."] not in {5, 6}]
    rows_personel_new = new1 + [r for r in new2 if r["№ з.п."] not in {5, 6}]
    rows_task = task1 + task2
    rows_move = [_row_move(1, 2, 5, 6), _row_move(3, 4, 5, 6)]

    filename = m.change_position_report(rows_move, rows_personel_old, rows_personel_new, rows_task, [], "15.08.2026")

    text = "\n".join(p.text for p in Document(str(tmp_path / filename)).paragraphs)
    assert text.count("ЗВІЛЬНИТИ ТА ПРИЗНАЧИТИ") == 1
    assert "ТРИНАДЦЯТИЙ Тринадцятий Тринадцятий" in text
    assert "ВІСІМНАДЦЯТИЙ Вісімнадцятий Вісімнадцятий" in text


def test_change_position_report_skips_invalid_rows(tmp_path, monkeypatch):
    """DataValidation.is_valid() завжди True сьогодні (validator.py) - гілку
    "не пройшла валідація" перевіряємо примусово через monkeypatch."""
    import validator

    monkeypatch.setattr(m, "OUTPUT_DIR", str(tmp_path))
    monkeypatch.setattr(validator.DataValidation, "is_valid", lambda self: False)
    rows_personel_old, rows_personel_new, rows_task = _build_data(
        1, 2, 3, 4, "ДВАДЦЯТИЙ Двадцятий Двадцятий", "ДРУГИЙ Другий Другий", "ТРЕТІЙ Третій Третій", "ЧЕТВЕРТИЙ Четвертий Четвертий",
    )
    rows_move = [_row_move(1, 2, 3, 4)]

    filename = m.change_position_report(rows_move, rows_personel_old, rows_personel_new, rows_task, [], "15.08.2026")

    assert filename is None
    assert list(tmp_path.iterdir()) == []
