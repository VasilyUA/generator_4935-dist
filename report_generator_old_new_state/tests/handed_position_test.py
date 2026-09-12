import sys
from pathlib import Path

sys.path.append(str(Path(__file__).resolve().parent.parent))

from docx import Document

import constants
from template import handed_position as m


def _data_for_generation():
    return {
        "task": {"personal_row": {"посада родовий": "стрільця", "ПІБ родовий": "ЧЕТВЕРТИЙ Четвертий Четвертий"}},
        "old": {
            "personal_row": {
                "Посада": "Стрілець", "підрозділ повністю": "1 рота",
                "звання фактичне": "сержант", "ПІБ": "ЧЕТВЕРТИЙ Четвертий Четвертий",
            },
            "commander_row": {
                "посада називний": "командир роти", "підрозділ повністю": "1 рота",
                "звання фактичне": "капітан", "ПІБ": "ШОСТИЙ Шостий Шостий",
                "ПІБ родовий": "ШОСТИЙ Шостий Шостий",
            },
            "higher_commander_row": {
                "Посада": "Командир батальйону", "звання фактичне": "підполковник",
                "ПІБ": "СЬОМИЙ Сьомий Сьомий",
            },
        },
    }


def test_handed_position_report_with_commander(tmp_path, monkeypatch):
    monkeypatch.setattr(m, "OUTPUT_DIR", str(tmp_path))

    doc_name = m.handed_position_report(_data_for_generation(), commander_position_code=7, rows_tvo=[])

    assert doc_name == str(tmp_path / "ЧЕТВЕРТИЙ Четвертий Четвертий здав посаду.docx")
    text = "\n".join(p.text for p in Document(doc_name).paragraphs)
    assert "здав." in text
    assert constants.FULL_MILITARY_UNIT in text
    assert "Шостий ШОСТИЙ" in text
    assert "Сьомий СЬОМИЙ" in text


def test_handed_position_report_without_commander_falls_back(tmp_path, monkeypatch):
    monkeypatch.setattr(m, "OUTPUT_DIR", str(tmp_path))

    doc_name = m.handed_position_report(_data_for_generation(), commander_position_code=None, rows_tvo=[])

    text = "\n".join(p.text for p in Document(doc_name).paragraphs)
    assert f"Командиру {constants.FULL_MILITARY_UNIT}" in text
    assert "Шостий ШОСТИЙ" not in text
    assert "Сьомий СЬОМИЙ" in text
