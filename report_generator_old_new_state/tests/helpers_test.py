import sys
from datetime import date, datetime
from pathlib import Path

sys.path.append(str(Path(__file__).resolve().parent.parent))

import pandas as pd
import pytest
from colorama import Fore, Style
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

import helpers
import constants


# -------------------------
# get_rows
# -------------------------
def test_get_rows_selects_only_requested_columns():
    df = pd.DataFrame({"ПІБ": ["ПЕРШИЙ Перший Перший"], "Посада": ["Стрілець"], "extra": ["ігнор"]})
    result = helpers.get_rows(df, ["ПІБ", "Посада"])
    assert result == [{"ПІБ": "ПЕРШИЙ Перший Перший", "Посада": "Стрілець"}]


# -------------------------
# print_red / print_green / print_timeout
# -------------------------
def test_print_red_output(capsys):
    helpers.print_red("Помилка")
    assert capsys.readouterr().out == Fore.RED + "Помилка" + Style.RESET_ALL + "\n"


def test_print_green_output(capsys):
    helpers.print_green("Готово")
    assert capsys.readouterr().out == Fore.GREEN + "Готово" + Style.RESET_ALL + "\n"


def test_print_timeout_prints_row_then_message(capsys, monkeypatch):
    monkeypatch.setattr(helpers.time, "sleep", lambda s: None)
    helpers.print_timeout("текст", 0)
    out = capsys.readouterr().out
    assert "ROW:" in out
    assert "текст" in out


# -------------------------
# create_or_clear_output_directory
# -------------------------
def test_create_or_clear_output_directory_clears_existing(tmp_path):
    out = tmp_path / "output"
    out.mkdir()
    (out / "old.docx").write_text("stale")

    helpers.create_or_clear_output_directory(str(out))

    assert out.exists()
    assert list(out.iterdir()) == []


def test_create_or_clear_output_directory_creates_missing(tmp_path):
    out = tmp_path / "does_not_exist_yet"
    helpers.create_or_clear_output_directory(str(out))
    assert out.exists() and out.is_dir()


# -------------------------
# set_margins / set_line_spacing / add_paragraph_with_style
# -------------------------
def test_set_margins_sets_page_size_and_margins():
    from docx.shared import Inches

    doc = Document()
    helpers.set_margins(doc, Inches(0.5), Inches(0.5), Inches(1), Inches(0.5))
    section = doc.sections[0]
    assert round(section.page_width.inches, 2) == 8.27
    assert section.left_margin == Inches(1)


def test_set_line_spacing_applies_to_all_paragraphs():
    from docx.shared import Pt

    doc = Document()
    doc.add_paragraph("перший")
    helpers.set_line_spacing(doc, 14)
    assert doc.paragraphs[0].paragraph_format.line_spacing == Pt(14)


def test_add_paragraph_with_style_basic_text():
    doc = Document()
    paragraph = helpers.add_paragraph_with_style(doc, "текст", first_line_indent=34)
    assert paragraph.text == "текст"
    assert paragraph.alignment == WD_ALIGN_PARAGRAPH.JUSTIFY
    assert paragraph.paragraph_format.first_line_indent is not None


def test_add_paragraph_with_style_format_tabs_adds_tab_stop():
    doc = Document()
    paragraph = helpers.add_paragraph_with_style(doc, "текст\tзначення", format_tabs=True)
    assert len(paragraph.paragraph_format.tab_stops) == 1


# -------------------------
# excel_col_to_index
# -------------------------
def test_excel_col_to_index_single_and_double_letter():
    assert helpers.excel_col_to_index("A") == 0
    assert helpers.excel_col_to_index("AA") == 26


# -------------------------
# int_or_none
# -------------------------
def test_int_or_none_valid_and_invalid():
    assert helpers.int_or_none("5") == 5
    assert helpers.int_or_none("не число") is None


# -------------------------
# convert_to_short_name - ВІДМІННІСТЬ від report_generator: не-рядковий вхід
# повертає '' (не пропускає значення як є)
# -------------------------
def test_convert_to_short_name_three_words():
    assert helpers.convert_to_short_name("ДРУГИЙ Другий Другий") == "Другий ДРУГИЙ"


def test_convert_to_short_name_splices_two_glued_words():
    # ім'я+по-батькові зліплені без пробілу ("ТретійТретій") - розклеюються
    # за межею мала->велика літера
    assert helpers.convert_to_short_name("ТРЕТІЙ ТретійТретій") == "Третій ТРЕТІЙ"


def test_convert_to_short_name_non_string_returns_empty_string():
    assert helpers.convert_to_short_name(None) == ""
    assert helpers.convert_to_short_name(123) == ""


def test_convert_to_short_name_wrong_word_count_raises():
    with pytest.raises(ValueError, match="Некоректний ПІБ"):
        helpers.convert_to_short_name("ОдинСлово")


# -------------------------
# get_position_code
# -------------------------
def test_get_position_code_reads_all_six_columns():
    row = {
        "НОМЕР ПОСАДИ З ЯКОЇ ПЕРЕМІЩУЄТЬСЯ ОСОБА": "1",
        "НОМЕР ПОСАДИ НА ЯКУ ПЕРЕМІЩУЄТЬСЯ ОСОБА": "2",
        "КОМАНДИР В СТАРІЙ ШТАТІ ЯКИЙ КЛОПОЧЕ КОМАНДИРУ БАТАЛЬЙОНУ": "3",
        "КОМАНДИР В НОВОМУ ШТАТІ ЯКИЙ КЛОПОЧЕ КОМАНДИРУ БАТАЛЬЙОНУ": "4",
        "КОМАНДИР БАТАЛЬЙОНУ В СТАРІЙ ШТАТІ": "5",
        "КОМАНДИР БАТАЛЬЙОНУ В НОВОМУ ШТАТІ": "6",
    }
    codes = helpers.get_position_code(row)
    assert codes == {
        "subordinate_position_code_from": 1,
        "subordinate_position_code_to": 2,
        "commander_position_code_in_old_state": 3,
        "commander_position_code_in_new_state": 4,
        "higher_commander_position_code_in_old_state": 5,
        "higher_commander_position_code_in_new_state": 6,
    }


# -------------------------
# to_date - ширший список форматів, порівняно з report_generator
# -------------------------
def test_to_date_date_passthrough():
    d = date(2026, 8, 19)
    assert helpers.to_date(d) is d


def test_to_date_datetime_converts_to_date():
    assert helpers.to_date(datetime(2026, 8, 19)) == date(2026, 8, 19)


@pytest.mark.parametrize("raw,expected", [
    ("2026-08-19", date(2026, 8, 19)),
    ("19.08.2026", date(2026, 8, 19)),
    ("19/08/2026", date(2026, 8, 19)),
    ("08/19/2026", date(2026, 8, 19)),
    ("19-08-2026", date(2026, 8, 19)),
    ("2026/08/19", date(2026, 8, 19)),
])
def test_to_date_accepts_many_formats(raw, expected):
    assert helpers.to_date(raw) == expected


def test_to_date_unparseable_raises():
    with pytest.raises(ValueError):
        helpers.to_date("не дата")


def test_to_date_unsupported_type_raises():
    with pytest.raises(ValueError):
        helpers.to_date(123)


# -------------------------
# find_higher_commander
# -------------------------
def _tvo_row(pib, posada, start, end, active_key, active=True):
    return {"ПІБ": pib, "ПОСАДА": posada, "Start": start, "End": end, active_key: active}


def test_find_higher_commander_returns_tvo_when_active_and_in_range():
    row = {"Посада": "Командир батальйону"}
    rows_tvo = [_tvo_row("ТРЕТІЙ Третій Третій", "Командир батальйону", "01.08.2026", "31.08.2026", "ТВО")]

    result = helpers.find_higher_commander(rows_tvo, row, "15.08.2026")

    assert result["ПІБ"] == "ТРЕТІЙ Третій Третій"


def test_find_higher_commander_falls_back_to_row_when_no_match():
    row = {"Посада": "Командир батальйону"}
    result = helpers.find_higher_commander([], row, "15.08.2026")
    assert result is row


def test_find_higher_commander_ignores_inactive_tvo():
    row = {"Посада": "Командир батальйону"}
    rows_tvo = [_tvo_row("ЧЕТВЕРТИЙ Четвертий Четвертий", "Командир батальйону", "01.08.2026", "31.08.2026", "ТВО", active=False)]
    result = helpers.find_higher_commander(rows_tvo, row, "15.08.2026")
    assert result is row


# -------------------------
# get_tvo_position / get_tvo_name
# -------------------------
def test_get_tvo_position_not_tvo_uses_commander_title():
    commander = {"Посада": "Командир батальйону", "звання фактичне": "підполковник", "ПІБ": "ШОСТИЙ Шостий Шостий"}
    result = helpers.get_tvo_position([], commander, date="15.08.2026")
    assert result == f"Командир батальйону {constants.FULL_MILITARY_UNIT}"


def test_get_tvo_position_active_tvo_uses_acting_title():
    commander = {"Посада": "Командир батальйону", "звання фактичне": "підполковник", "ПІБ": "СЬОМИЙ Сьомий Сьомий"}
    rows_tvo = [_tvo_row("ВОСЬМИЙ Восьмий Восьмий", "Командир батальйону", "01.08.2026", "31.08.2026", "ТВО")]

    result = helpers.get_tvo_position(rows_tvo, commander, date="15.08.2026")

    assert result == f"Тимчасово виконуючий обов'язки командира батальйону {constants.FULL_MILITARY_UNIT}"


def test_get_tvo_name_not_tvo_uses_commander_own_pib():
    commander = {"Посада": "Командир батальйону", "звання фактичне": "підполковник", "ПІБ": "ШІСТНАДЦЯТИЙ Шістнадцятий Шістнадцятий"}
    result = helpers.get_tvo_name([], commander, date="15.08.2026")
    assert result == "підполковник\tШістнадцятий ШІСТНАДЦЯТИЙ"


def test_get_tvo_name_active_tvo_uses_substitute_pib():
    commander = {"Посада": "Командир батальйону", "звання фактичне": "підполковник", "ПІБ": "ДЕСЯТИЙ Десятий Десятий"}
    rows_tvo = [_tvo_row("ОДИНАДЦЯТИЙ Одинадцятий Одинадцятий", "Командир батальйону", "01.08.2026", "31.08.2026", "ТВО")]

    result = helpers.get_tvo_name(rows_tvo, commander, date="15.08.2026")

    assert result == "підполковник\tОдинадцятий ОДИНАДЦЯТИЙ"


# -------------------------
# get_commander_row
# -------------------------
def test_get_commander_row_old_state_no_tvo_returns_own_data():
    position_codes = {"commander_position_code_in_old_state": 3}
    rows_personel_old = [{"№ з.п.": 3, "ПІБ": "ДВАНАДЦЯТИЙ Дванадцятий Дванадцятий", "Посада": "Командир роти", "підрозділ повністю": "1 рота", "звання фактичне": "капітан"}]
    rows_task = [{"ПІБ називний": "ДВАНАДЦЯТИЙ Дванадцятий Дванадцятий", "Посада називний": "Командир роти", "Посада давальний": "Командиру роти", "посада називний": "командир роти"}]

    result = helpers.get_commander_row(position_codes, rows_personel_old, [], "old", [], rows_task, "15.08.2026")

    assert result["ПІБ"] == "ДВАНАДЦЯТИЙ Дванадцятий Дванадцятий"
    assert result["ТВО"] is False


def test_get_commander_row_old_state_active_tvo_overrides():
    position_codes = {"commander_position_code_in_old_state": 3}
    rows_personel_old = [{"№ з.п.": 3, "ПІБ": "ТРИНАДЦЯТИЙ Тринадцятий Тринадцятий", "Посада": "Командир роти", "підрозділ повністю": "1 рота", "звання фактичне": "капітан"}]
    rows_tvo = [{
        "ПІБ": "ТРИНАДЦЯТИЙ Тринадцятий Тринадцятий", "ПОСАДА": "Заступник командира роти",
        "Start": "01.08.2026", "End": "31.08.2026", "old TVO Active": True, "№old": 99,
    }]
    rows_task = [
        {"ПІБ називний": "ЧОТИРНАДЦЯТИЙ Чотирнадцятий Чотирнадцятий", "Посада називний": "Заступник командира роти", "Посада давальний": "Заступнику командира роти", "посада називний": "заступник командира роти", "ПІБ родовий": "Чотирнадцятого Чотирнадцятого"},
    ]

    result = helpers.get_commander_row(position_codes, rows_personel_old, [], "old", rows_tvo, rows_task, "15.08.2026")

    assert result["ТВО"] is True
    assert result["Посада"] == "Заступник командира роти"
    assert result["ПІБ"] == "ТРИНАДЦЯТИЙ Тринадцятий Тринадцятий"


def test_get_commander_row_new_state_no_match_returns_blank_defaults():
    position_codes = {"commander_position_code_in_new_state": 999}
    result = helpers.get_commander_row(position_codes, [], [], "new", [], [], "15.08.2026")
    assert result["ПІБ"] == ""
    assert result["ТВО"] is False


# -------------------------
# get_subordinate
# -------------------------
def test_get_subordinate_builds_full_structure():
    position_codes = {
        "subordinate_position_code_from": 1, "subordinate_position_code_to": 2,
        "commander_position_code_in_old_state": 3, "commander_position_code_in_new_state": 3,
        "higher_commander_position_code_in_old_state": 4, "higher_commander_position_code_in_new_state": 4,
    }
    rows_personel_old = [
        {"№ з.п.": 1, "ПІБ": "ШІСТНАДЦЯТИЙ Шістнадцятий Шістнадцятий", "Посада": "Стрілець", "підрозділ повністю": "1 рота", "звання фактичне": "сержант"},
        {"№ з.п.": 3, "ПІБ": "СІМНАДЦЯТИЙ Сімнадцятий Сімнадцятий", "Посада": "Командир роти", "підрозділ повністю": "1 рота", "звання фактичне": "капітан"},
        {"№ з.п.": 4, "ПІБ": "ВІСІМНАДЦЯТИЙ Вісімнадцятий Вісімнадцятий", "Посада": "Командир батальйону", "підрозділ повністю": "батальйон", "звання фактичне": "підполковник"},
    ]
    rows_personel_new = [
        {"№ з.п.": 2, "ПІБ": "ДВАДЦЯТИЙ Двадцятий Двадцятий", "Посада": "Навідник", "підрозділ повністю": "2 рота", "звання фактичне": "сержант"},
        {"№ з.п.": 3, "ПІБ": "СІМНАДЦЯТИЙ Сімнадцятий Сімнадцятий", "Посада": "Командир роти", "підрозділ повністю": "1 рота", "звання фактичне": "капітан"},
        {"№ з.п.": 4, "ПІБ": "ВІСІМНАДЦЯТИЙ Вісімнадцятий Вісімнадцятий", "Посада": "Командир батальйону", "підрозділ повністю": "батальйон", "звання фактичне": "підполковник"},
    ]
    rows_task = [
        {"ПІБ називний": "ДВАДЦЯТИЙ Двадцятий Двадцятий", "Посада називний": "Навідник", "ПІБ родовий": "Двадцятого Двадцятого", "посада родовий": "навідника"},
    ]

    result = helpers.get_subordinate(position_codes, rows_personel_old, rows_personel_new, rows_task, [], "15.08.2026")

    assert result["old"]["personal_row"]["ПІБ"] == "ШІСТНАДЦЯТИЙ Шістнадцятий Шістнадцятий"
    assert result["new"]["personal_row"]["ПІБ"] == "ДВАДЦЯТИЙ Двадцятий Двадцятий"
    assert result["old"]["higher_commander_row"]["ПІБ"] == "ВІСІМНАДЦЯТИЙ Вісімнадцятий Вісімнадцятий"
    assert result["task"]["personal_row"]["ПІБ називний"] == "ДВАДЦЯТИЙ Двадцятий Двадцятий"


# -------------------------
# get_text
# -------------------------
def test_get_text_builds_old_and_new_position_sentences():
    data_for_generation = {
        "old": {"personal_row": {
            "звання фактичне": "сержант", "Посада": "Стрілець", "підрозділ повністю": "1 рота",
            "ВОС": "000000А/000", "звання за штатом": "сержант", "№ з.п.": 1,
        }},
        "new": {"personal_row": {
            "Посада": "Навідник", "підрозділ повністю": "2 рота",
            "ВОС": "111111Б/111", "звання за штатом": "сержант", "№ з.п.": 2,
        }},
        "task": {"personal_row": {"ПІБ родовий": "ДВАДЦЯТИЙ Двадцятий Двадцятий родовий"}},
    }

    text = helpers.get_text(data_for_generation)

    assert "стрілець" in text
    assert "на посаду навідник" in text
    assert constants.FULL_MILITARY_UNIT in text
    assert "ДВАДЦЯТИЙ Двадцятий Двадцятий родовий" in text


# -------------------------
# input_date
# -------------------------
def test_input_date_blank_returns_today(monkeypatch, capsys):
    monkeypatch.setattr("builtins.input", lambda prompt="": "")
    result = helpers.input_date()
    assert result == date.today()


def test_input_date_parses_dotted_format(monkeypatch):
    monkeypatch.setattr("builtins.input", lambda prompt="": "19.08.2026")
    assert helpers.input_date() == date(2026, 8, 19)


def test_input_date_reprompts_until_valid(monkeypatch, capsys):
    answers = iter(["не дата", "19.08.2026"])
    monkeypatch.setattr("builtins.input", lambda prompt="": next(answers))
    result = helpers.input_date()
    assert result == date(2026, 8, 19)
    assert "Невірний формат" in capsys.readouterr().out
