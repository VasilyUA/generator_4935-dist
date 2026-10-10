import pytest
from openpyxl import Workbook

from bn_fixtures import write_rop_vop
from content.rop_vop_reader import classify, read_rop_vop


@pytest.mark.parametrize("text, expected", [
    ("Вогнева позиція \"АЛЬФА\"", ("ВП", "АЛЬФА")),
    ("Вогнева позиція мінометного розрахунку \"Альфа\" A0001", ("ВП", "АЛЬФА")),
    ("Вогнева засідка \"АЛЬФА\"", ("ВЗ", "АЛЬФА")),
    ("Позиція відділення \"АЛЬФА\"", ("ПВ", "АЛЬФА")),
    ("Спостережна позиція \"АЛЬФА\"", ("СП", "АЛЬФА")),
    ("СП \"Альфа\" A0001", ("СП", "АЛЬФА")),
    ("СПАР \"Альфа\" батарея", ("СПАР", "АЛЬФА")),
    ("ХАБ (Альфа)  A0001", ("ХАБ", "АЛЬФА")),
    ("Пункт БпЛА Mavic \"Альфа\"", ("ТЗ", "АЛЬФА")),
    ("Пункт зльоту бомбери \"Альфа\"", ("ТЗ", "АЛЬФА")),
    ("ПУ НРК \"Альфа\"", ("ПУ", "АЛЬФА")),
    ("РОП\n“АЛЬФА\"\nрота", ("РОП", "АЛЬФА")),
    ("ВОП «АЛЬФА»", ("ВОП", "АЛЬФА")),
    ("ЗКП A0001", ("ЗКП", "")),
    ("Всього за окремі позиції", (None, "")),
    ("Упр - 4 в/сл", (None, "")),
    ("", (None, "")),
])
def test_classify(text, expected):
    assert classify(text) == expected


def test_read_entries(tmp_path):
    table = read_rop_vop(str(write_rop_vop(tmp_path / "РОП_ВОП.xlsx")))
    entries = {(e.type, e.name): e for e in table.entries}
    assert entries[("БРО", "")].quantity == 50 and len(entries[("БРО", "")].mgrs) == 2
    assert entries[("РОП", "АЛЬФА")].quantity == 10
    assert entries[("КСП", "АЛЬФА")].quantity == 4
    assert entries[("ВОП", "БРАВО")].quantity == 6
    assert entries[("ВЗ", "ЧАРЛІ")].quantity == 3
    assert entries[("СП", "ДЕЛЬТА")].mgrs == ["31U DQ 40116 10116"]
    # "Окремі позиції": назва в колонці БРО, кількість - у колонці позицій.
    assert entries[("ВП", "ЕХО")].quantity == 4
    assert entries[("ХАБ", "ГОТЕЛІ")].quantity == 7
    assert table.zkp_mgrs == "31U DQ 40800 10800"
    assert table.warnings == []
    assert not any(e.quantity == 99 for e in table.entries)


def test_missing_headers(tmp_path):
    wb = Workbook()
    wb.active["A1"] = "довільна таблиця"
    path = tmp_path / "інше.xlsx"
    wb.save(path)
    table = read_rop_vop(str(path))
    assert table.entries == [] and len(table.warnings) == 1


def test_missing_coordinates_column(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws["C7"], ws["E7"], ws["H7"] = "БРО", "РОП", "ВОП"
    path = tmp_path / "без_координат.xlsx"
    wb.save(path)
    assert len(read_rop_vop(str(path)).warnings) == 1
