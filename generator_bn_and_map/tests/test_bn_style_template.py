import copy

import docx
import pytest
from docx.oxml.ns import qn

from bn_fixtures import tiny_png, write_style_reference
from generators.style_template import StyleTemplate, create_style_template, format_signature


@pytest.fixture
def style(tmp_path):
    reference = write_style_reference(tmp_path / "зразок.docx", tiny_png(tmp_path / "1.png"))
    path = tmp_path / "Стиль.docx"
    mapping = create_style_template(str(reference), str(path), {"tables": ["table:ack"]})
    return reference, path, mapping


def _ppr_xml(paragraph_element):
    ppr = paragraph_element.find(qn("w:pPr"))
    return ppr.xml if ppr is not None else ""


def test_mapping_shares_formats_and_names_table(style):
    _, _, mapping = style
    assert mapping[2] == mapping[3]                 # однаково відформатовані абзаци
    assert len({mapping[0], mapping[1], mapping[2], mapping[4], mapping[5]}) == 5
    assert mapping[6] == "table:ack"


def test_harvest_leaves_only_section_properties(style):
    _, path, mapping = style
    template = StyleTemplate(str(path))
    assert set(template.prototypes) == {v for v in mapping.values() if v.startswith("f")}
    assert list(template.tables) == ["table:ack"]
    body = list(template.new_document().element.body.iterchildren())
    assert [child.tag for child in body] == [qn("w:sectPr")]


def test_paragraph_copies_format_from_sample(style, tmp_path):
    reference, path, mapping = style
    template = StyleTemplate(str(path))
    doc = template.new_document()
    template.add_paragraph(doc, mapping[2], [("Новий ", None), ("жирний", True), (" текст", False)])
    template.add_paragraph(doc, mapping[4], [], page_break_after=True)
    out = tmp_path / "out.docx"
    doc.save(out)

    generated = docx.Document(out)
    sample = docx.Document(reference)
    paragraph = generated.paragraphs[0]
    assert paragraph.text == "Новий жирний текст"
    assert _ppr_xml(paragraph._p) == _ppr_xml(sample.paragraphs[2]._p)
    assert [r.font.size.pt for r in paragraph.runs] == [14, 14, 14]
    assert [r.bold for r in paragraph.runs] == [None, True, False]
    assert 'w:type="page"' in generated.paragraphs[1]._p.xml


def test_drawing_stays_with_its_own_format_only(style, tmp_path):
    _, path, mapping = style
    template = StyleTemplate(str(path))
    doc = template.new_document()
    template.add_paragraph(doc, mapping[0], [("Новий титул", None)])
    template.add_paragraph(doc, mapping[2], [("Текст", None)])
    paragraphs = doc.paragraphs
    assert "<w:drawing" in paragraphs[0]._p.xml and paragraphs[0].text == "Новий титул"
    assert "<w:drawing" not in paragraphs[1]._p.xml


def test_table_rows_titles_and_blanks(style, tmp_path):
    _, path, _ = style
    template = StyleTemplate(str(path))
    doc = template.new_document()
    template.add_table(doc, "table:ack", "АРКУШ\nдоведення", ["Посада", "Звання", "Підпис"],
                       [["Посада 1", "звання"], ["Посада 2", "звання", "x"]], blank_rows=2)
    table = doc.tables[0]
    assert len(table.rows) == 2 + 2 + 2
    assert [p.text for p in table.cell(0, 0).paragraphs] == ["АРКУШ", "доведення"]
    assert [c.text for c in table.rows[2].cells] == ["Посада 1", "звання", ""]
    assert [c.text for c in table.rows[3].cells] == ["Посада 2", "звання", "x"]
    assert all(c.text == "" for row in table.rows[4:] for c in row.cells)


def test_table_without_title_row(style, tmp_path):
    _, path, _ = style
    template = StyleTemplate(str(path))
    doc = template.new_document()
    template.add_table(doc, "table:ack", "", ["А", "Б", "В"], [["1", "2", "3"]], blank_rows=0)
    assert [c.text for c in doc.tables[0].rows[0].cells] == ["А", "А", "А"]


def test_signature_ignores_text_and_rsids(style):
    reference, _, _ = style
    paragraphs = docx.Document(reference).paragraphs
    other = copy.deepcopy(paragraphs[2]._p)
    other.set(qn("w:rsidR"), "00AB12CD")
    assert format_signature(other)[:2] == format_signature(paragraphs[3]._p)[:2]
