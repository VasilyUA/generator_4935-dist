import pytest
from docx.oxml.ns import qn

from bn_fixtures import tiny_png, write_style_reference
from content.position_refresh import PositionRefresher
from generators.document_builder import (
    ORIGIN_MANUAL, ORIGIN_REFRESHED, ORIGIN_SOURCE, ORIGIN_STANDARD, DocumentBuilder, apply_block_options, line_kind,
)
from generators.style_template import StyleTemplate, create_style_template


@pytest.fixture
def style(tmp_path):
    reference = write_style_reference(tmp_path / "зразок.docx", tiny_png(tmp_path / "1.png"))
    path = tmp_path / "Стиль.docx"
    mapping = create_style_template(str(reference), str(path), {"tables": ["table:ack"]})
    return StyleTemplate(str(path)), mapping


@pytest.mark.parametrize("text, kind", [
    ("", "empty"), ("MAVIK:", "header"), ("ВР № 1 (31U DQ 1 1)", "vr"), ("101   ЖС   31U DQ 48251 11932;", "target"),
    ("Сусід праворуч: A0002", "neighbour"), ("ліворуч: точка", "side"), ("- ВП «АЛЬФА» (31U DQ 40100 10100)", "position"),
    ("- евакуація", "dash"), ("Звичайний текст.", "body"), ("Опис з координатами (31U DQ 48251 11932):", "body"),
])
def test_line_kind(text, kind):
    assert line_kind(text) == kind


def test_only_and_nth():
    items = ["перший рядок", "другий рядок А", "інший", "другий рядок Б"]
    assert apply_block_options(items, {"only": ["другий"]}) == ["другий рядок А", "другий рядок Б"]
    assert apply_block_options(items, {"only": "другий", "nth": 1}) == ["другий рядок Б"]
    assert apply_block_options(items, {"only": "немає", "nth": 0}) == []


def test_style_mode_document(style):
    template, mapping = style
    title_fmt, heading_fmt, body_fmt, signature_fmt = mapping[0], mapping[1], mapping[2], mapping[5]
    spec = {
        "formats": {"body": body_fmt, "heading1": heading_fmt, "title": title_fmt},
        "title_page": [{"text": "Титул {bn_year}", "fmt": title_fmt, "page_break_after": True}],
        "sections": [{"heading": "РОЗДІЛ", "content": [
            "Статичний текст.",
            {"blank": 2, "fmt": body_fmt},
            {"block": "b", "fmt_by_kind": {"header": signature_fmt}},
            {"block": "s", "fmt_list": [signature_fmt, body_fmt]},
            {"block": "x", "fallback": ["Запасний текст."]},
            {"block": "b", "only": ["друг"], "nth": 0},
        ]}],
        "acknowledgement": {"title": "АРКУШ", "columns": ["Посада", "Звання", "Підпис"], "rows": [["П", "З"]], "blank_rows": 1},
    }
    blocks = {"b": ["Заголовок блоку:", "другий рядок"], "s": ["перший", "другий", "третій"]}
    builder = DocumentBuilder({"bn_year": "2026"}, blocks, style=template)
    doc = builder.build(spec)
    texts = [p.text for p in doc.paragraphs]
    assert texts[0] == "Титул 2026"
    assert sum('w:type="page"' in p._p.xml for p in doc.paragraphs) == 1
    assert "<w:drawing" in doc.paragraphs[0]._p.xml
    assert texts[1] == "РОЗДІЛ"                       # нумерацію дає стиль, не код
    heading_style = doc.paragraphs[1]._p.find(qn("w:pPr")).find(qn("w:pStyle"))
    assert heading_style is not None
    assert texts[2:5] == ["Статичний текст.", "", ""]
    assert texts[5:7] == ["Заголовок блоку:", "другий рядок"]
    assert doc.paragraphs[5].alignment == doc.paragraphs[7].alignment      # header -> формат підпису
    assert texts[7:10] == ["перший", "другий", "третій"]
    assert texts[10:12] == ["Запасний текст.", "другий рядок"]
    assert builder.fallback_blocks == ["x"]
    assert len(doc.tables[0].rows) == 4
    assert "РОЗДІЛ" in builder.log[1].text


def test_missing_block_highlighted_in_style_mode(style):
    template, mapping = style
    builder = DocumentBuilder({}, {}, style=template)
    doc = builder.build({"formats": {"body": mapping[2]}, "sections": [{"block": "немає"}]})
    assert doc.paragraphs[0].text == "‹немає даних: немає›"
    assert doc.paragraphs[0].runs[0].font.highlight_color is not None


def test_origins_and_automatic_percent():
    refresher = PositionRefresher([{"type": "ВП", "name": "АЛЬФА", "mgrs": "31U DQ 40100 10100"}])
    builder = DocumentBuilder(
        {"zkp": "31U DQ 40800 10800", "sig": "ГРІМ-101"}, {"src": ["Текст з джерела."]},
        refresher=refresher, known_names=["ГОРА 1"], source_mgrs={"31U DQ 40900 10900"},
    )
    builder.build({"sections": [
        "Незмінний текст.",
        "ЗКП ({zkp}) і сигнал ‘‘{sig}’’.",
        "ВП «АЛЬФА» (31U DQ 49999 19999).",
        "Рубіж (31U DQ 40900 10900) з БР.",
        "Рубіж штабу (31U DQ 45555 15555) за сигналом ‘‘ВІТЕР-45’’.",
        "КСП «ГОРА 1» без сигналу.",
        {"block": "src"},
    ]})
    origins = [entry.origin for entry in builder.log]
    assert origins == [ORIGIN_STANDARD, ORIGIN_REFRESHED, ORIGIN_REFRESHED, ORIGIN_STANDARD, ORIGIN_MANUAL,
                       ORIGIN_STANDARD, ORIGIN_SOURCE]
    manual = builder.manual_paragraphs()
    assert manual[0].staff_values == ("31U DQ 45555 15555", "‘‘ВІТЕР-45’’")
    staff_chars = len("31U DQ 45555 15555") + len("‘‘ВІТЕР-45’’")
    all_chars = sum(len(entry.text) for entry in builder.log)
    assert builder.automatic_percent() == round(100 - 100 * staff_chars / all_chars, 1)
    assert builder.refreshed_mentions == 1
    coverage = builder.coverage()
    assert set(coverage) == {ORIGIN_SOURCE, ORIGIN_REFRESHED, ORIGIN_STANDARD, ORIGIN_MANUAL}
    assert round(sum(coverage.values())) == 100


def test_empty_document_percentages():
    builder = DocumentBuilder({}, {})
    builder.build({})
    assert builder.automatic_percent() == 100.0
    assert builder.coverage()[ORIGIN_MANUAL] == 0
