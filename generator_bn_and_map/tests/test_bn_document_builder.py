from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX

from content.data_store import normalize_name
from generators.document_builder import DocumentBuilder, apply_block_options, find_stale_position_names


def _texts(doc):
    return [p.text for p in doc.paragraphs]


def _build(spec, context=None, blocks=None):
    builder = DocumentBuilder(context or {}, blocks or {})
    return builder, builder.build(spec)


def test_item_forms_and_formatting():
    spec = {"sections": [
        "звичайний",
        ["з жирним АЛЬФА", ["АЛЬФА"]],
        {"text": "по центру", "align": "center", "indent": False, "bold": True, "italic": True, "size": 12},
        {"text": "з лівим відступом", "left_indent": 2.5},
        {"text": "звання\tПІДПИС"},
        {"blank": 2},
        {"page_break": True},
        42,
    ]}
    _, doc = _build(spec)
    paragraphs = doc.paragraphs
    assert _texts(doc)[:5] == ["звичайний", "з жирним АЛЬФА", "по центру", "з лівим відступом", "звання\tПІДПИС"]
    assert paragraphs[0].alignment == WD_ALIGN_PARAGRAPH.JUSTIFY
    assert round(paragraphs[0].paragraph_format.first_line_indent.cm, 2) == 1.25
    assert [r.bold for r in paragraphs[1].runs] == [False, True]
    centered = paragraphs[2]
    assert centered.alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert centered.runs[0].bold and centered.runs[0].italic and centered.runs[0].font.size.pt == 12
    assert round(paragraphs[3].paragraph_format.left_indent.cm, 2) == 2.5
    assert len(paragraphs[4].paragraph_format.tab_stops) == 1
    assert _texts(doc)[5:7] == ["", ""]
    assert 'w:type="page"' in paragraphs[7]._p.xml


def test_heading_numbering_and_levels():
    spec = {"sections": [
        {"heading": "ПЕРШИЙ", "content": ["текст"]},
        {"heading": "ДРУГИЙ", "content": [
            {"heading": "ПІДРОЗДІЛ А", "content": []},
            {"heading": "ПІДРОЗДІЛ Б", "content": []},
            {"heading": "БЕЗ НОМЕРА", "number": False},
        ]},
    ]}
    _, doc = _build(spec)
    headings = [(p.style.name, p.text) for p in doc.paragraphs if p.style.name.startswith("Heading")]
    assert headings == [
        ("Heading 1", "1. ПЕРШИЙ"), ("Heading 1", "2. ДРУГИЙ"),
        ("Heading 2", "2.1. ПІДРОЗДІЛ А"), ("Heading 2", "2.2. ПІДРОЗДІЛ Б"), ("Heading 2", "БЕЗ НОМЕРА"),
    ]


def test_placeholders_known_and_unknown():
    builder, doc = _build({"header": ["№{bn_number} {невідомий} {bn_date}"]}, {"bn_number": "5", "bn_date": "01.02.2026"})
    assert _texts(doc) == ["№5 {невідомий} 01.02.2026"]
    assert builder.unknown_placeholders == {"невідомий"}


def test_none_context_value_is_unknown():
    builder, doc = _build({"header": ["{zkp_mgrs}"]}, {"zkp_mgrs": None})
    assert _texts(doc) == ["{zkp_mgrs}"] and builder.unknown_placeholders == {"zkp_mgrs"}


def test_blocks_with_prefix_and_options():
    blocks = {"b": ["перший", ["другий {не плейсхолдер}", ["другий"]], "пропустити це", "стоп тут", "після стопу"]}
    spec = {"sections": [{"block": "b", "prefix": "№{n} ", "skip": "Пропустити", "stop_before": ["стоп"]}]}
    builder, doc = _build(spec, {"n": "7"}, blocks)
    assert _texts(doc) == ["№7 перший", "другий {не плейсхолдер}"]
    assert builder.used_blocks == {"b": 2}
    assert builder.unknown_placeholders == set()


def test_missing_block_is_highlighted():
    builder, doc = _build({"sections": [{"block": "немає"}]})
    assert builder.missing_blocks == ["немає"]
    assert doc.paragraphs[0].runs[0].font.highlight_color == WD_COLOR_INDEX.YELLOW


def test_apply_block_options_without_options():
    assert apply_block_options(["а", "б"], {}) == ["а", "б"]


def test_title_page_header_footer_and_acknowledgement():
    spec = {
        "title_page": ["Титул"],
        "header": ["Шапка"],
        "signature": ["Підпис"],
        "acknowledgement": {"title": "АРКУШ\nдоведення", "rows": [["Посада 1", "звання", "", "ПЕРШИЙ Перший Перший"]], "blank_rows": 2},
        "closing": ["Кінець"],
    }
    _, doc = _build(spec)
    texts = _texts(doc)
    assert texts[0] == "Титул" and texts[2] == "Шапка"
    assert 'w:type="page"' in doc.paragraphs[1]._p.xml
    assert texts[-1] == "Кінець"
    table = doc.tables[0]
    assert len(table.rows) == 1 + 1 + 1 + 2
    assert table.cell(0, 0).text == "АРКУШ\nдоведення" and table.cell(0, 4).text == "АРКУШ\nдоведення"
    assert [c.text for c in table.rows[1].cells] == ["Посада", "Військове звання", "Підпис", "П.І.Б", "Дата ознайомлення"]
    assert [c.text for c in table.rows[2].cells] == ["Посада 1", "звання", "", "ПЕРШИЙ Перший Перший", ""]
    section = doc.sections[0]
    assert section.different_first_page_header_footer
    assert "PAGE" in section.header._element.xml
    assert section.footer.paragraphs[0].text == "ДЛЯ СЛУЖБОВОГО КОРИСТУВАННЯ"
    assert round(section.left_margin.cm, 1) == 3.0 and round(section.right_margin.cm, 1) == 1.0


def test_stale_position_names():
    spec = {
        "sections": [{"heading": "ВП «СТАРИЙ»", "content": [
            "Позиція ВП «АЛЬФА» і ХАБ \"НОВИЙ\" та ВП «СТАРИЙ»",
            {"text": "СП ‘‘ЗНИКЛИЙ’’ біля"},
            {"block": "x", "prefix": "ТЗ «ПРЕФІКС» "},
        ]}],
        "closing": ["кінець без назв"],
    }
    stale = find_stale_position_names(spec, ["АЛЬФА", "Новий", ""], normalize_name)
    assert stale == ["СТАРИЙ", "ЗНИКЛИЙ", "ПРЕФІКС"]
