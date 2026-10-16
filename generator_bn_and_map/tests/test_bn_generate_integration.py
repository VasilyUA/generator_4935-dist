import os
from datetime import date

import docx
import pytest

import generators.pipeline as pipeline
from bn_fixtures import BR_PARAGRAPHS, PBD_PARAGRAPHS, make_data, write_docx, write_rop_vop
from content.sources import find_file, find_latest_br, find_latest_source_folder
from generators.generate_bn import generate_bn, title_case
from generators.generate_map import generate_map


@pytest.fixture
def resources(tmp_path):
    res = tmp_path / "resources"
    res.mkdir()
    write_docx(res / "БР старе 1111.docx", [p.replace("1234", "1111").replace("19.08.2026", "01.08.2026") for p in BR_PARAGRAPHS
                                             if isinstance(p, str)])
    write_docx(res / "БР нове 1234.docx", BR_PARAGRAPHS)
    (res / "~$БР тимчасовий.docx").write_bytes(b"")
    old = res / "01.10.2026 стара"
    new = res / "09.10.2026 нова"
    for folder in (old, new):
        folder.mkdir()
    write_docx(new / "Додаток ПБД 09.10.2026.docx", PBD_PARAGRAPHS)
    write_rop_vop(new / "Додаток 2 (РОП_ВОП) 09.10.2026.xlsx")
    (res / "15.10.2026 це файл, не тека").write_text("")
    (res / "без дати").mkdir()
    return res


def test_sources(resources):
    folder, folder_date = find_latest_source_folder(str(resources))
    assert os.path.basename(folder) == "09.10.2026 нова" and folder_date == date(2026, 10, 9)
    assert os.path.basename(find_file(folder, ("пбд",), ".DOCX")) == "Додаток ПБД 09.10.2026.docx"
    assert find_file(folder, ("немає",), ".docx") is None
    assert find_file(None, ("ПБД",), ".docx") is None
    assert os.path.basename(find_latest_br(str(resources))) == "БР нове 1234.docx"
    assert find_latest_source_folder(str(resources / "немає")) == (None, None)


def test_find_latest_br_without_files(tmp_path):
    assert find_latest_br(str(tmp_path)) is None


def test_refresh_errors(tmp_path):
    with pytest.raises(pipeline.SourceError):
        pipeline.refresh_pbd_cache({}, str(tmp_path))
    (tmp_path / "09.10.2026 порожня").mkdir()
    with pytest.raises(pipeline.SourceError):
        pipeline.refresh_pbd_cache({}, str(tmp_path))


def test_refresh_without_rop_vop(tmp_path):
    folder = tmp_path / "09.10.2026 лише ПБД"
    folder.mkdir()
    write_docx(folder / "ПБД.docx", PBD_PARAGRAPHS)
    cache, warnings, used = pipeline.refresh_pbd_cache({}, str(tmp_path))
    assert len(used) == 1 and any("РОП/ВОП" in w for w in warnings)


def test_title_case():
    assert title_case("ПЕРШЕ ДРУГЕ") == "Перше Друге"
    assert title_case("") == ""


def test_full_generation(resources, tmp_path):
    data = make_data()
    cache, warnings, used = pipeline.refresh_pbd_cache(data, str(resources))
    assert len(used) == 2
    br = pipeline.load_latest_br(str(resources))
    assert br.number == "1234"

    output = tmp_path / "output"
    result = generate_bn(data, br, cache, date(2026, 8, 24), "06.00", "5", str(output))
    assert os.path.basename(result.path) == "Бойовий наказ КБ A0001 БАТ 24.08.2026.docx"
    doc = docx.Document(result.path)
    texts = [p.text for p in doc.paragraphs]
    assert "Тестове – 2026 рік" in texts
    assert "БОЙОВИЙ НАКАЗ КОМАНДИРА a0001 бат A0000 №5" in texts
    # За замовчуванням головне - БР: видання карти з шапки БР (2025), а не з ПБД (2024).
    assert "КСП – Тестове 06.00 24.08.2026. Карта 25 000, видання 2025р." in texts
    assert "1. ВИСНОВКИ З ОЦІНЮВАННЯ ПРОТИВНИКА" in texts
    assert "Відповідно до БР №1234 від 19.08.2026 A0001 продовжити ведення оборонного бою." in texts
    assert "Бойовий порядок мати в один ешелон." not in texts
    assert "Сусід праворуч: A0002 обороняє район." in texts
    assert "КСП (31U DQ 40500 10500), ЗКП (31U DQ 40800 10800), {невідомий}" in texts
    assert "‹немає даних: pbd:невідомий›" in texts
    assert result.used_blocks["pbd:rop:АЛЬФА:positions"] == 4
    assert any("{невідомий}" in w for w in result.warnings)
    assert any("ЗАСТАРІЛА" in w for w in result.warnings)
    assert any("pbd:невідомий" in w for w in result.warnings)

    map_result = generate_map(data, cache, str(output / "map"), str(tmp_path / "icons"))
    assert os.path.isfile(map_result.kmz_path) and os.path.isfile(map_result.csv_path)
    assert {p.name for p in map_result.points} >= {"ВЗ Чарлі", "КСП Альфа", "СП Індія", "ЗКП"}


def test_generate_bn_without_br_and_document(resources, tmp_path):
    data = make_data()
    cache, _, _ = pipeline.refresh_pbd_cache(data, str(resources))
    result = generate_bn(data, None, cache, date(2026, 8, 24), "06.00", "5", str(tmp_path))
    assert any("БР не знайдено" in w for w in result.warnings)
    with pytest.raises(ValueError):
        generate_bn({"unit": {}}, None, cache, date(2026, 8, 24), "06.00", "5", str(tmp_path))


def test_load_latest_br_without_files(tmp_path):
    assert pipeline.load_latest_br(str(tmp_path)) is None


def test_br_stage_signals_and_source_coordinates(resources):
    from generators.generate_bn import br_stage_signals, source_coordinates
    data = make_data()
    cache, _, _ = pipeline.refresh_pbd_cache(data, str(resources))
    br = pipeline.load_latest_br(str(resources))
    assert br_stage_signals(br) == {"1": "ГРІМ-101"}
    assert br_stage_signals(None) == {}
    coordinates = source_coordinates(br, cache)
    assert {"31U DQ 48000 11000", "31U DQ 40116 10116", "31U DQ 40800 10800", "31U DQ 40500 10500"} <= coordinates
    assert source_coordinates(None, {}) == set()


def test_generation_with_style_and_review(resources, tmp_path):
    from bn_fixtures import tiny_png, write_style_reference
    from generators.generate_bn import ORIGIN_TITLES, review_lines
    from generators.style_template import create_style_template

    reference = write_style_reference(tmp_path / "зразок.docx", tiny_png(tmp_path / "1.png"))
    style_path = tmp_path / "Стиль.docx"
    mapping = create_style_template(str(reference), str(style_path), {"tables": ["table:ack"]})
    data = make_data()
    data["bn"]["KSP_NAME"] = "ГОРА 1"
    document = data["bn"]["document"]
    document["formats"] = {"body": mapping[2], "heading1": mapping[1], "title": mapping[0]}
    document["sections"].append({"heading": "ПЛАН", "content": [
        "Район зосередження мати в районі – {reserve_area_БнГ};",
        "КСП «ГОРА 1» за координатами (31U DQ 41111 11111) за сигналом ‘‘{br_signal_1}’’.",
        "Штабний рубіж (31U DQ 45555 15555).",
        {"block": "pbd:signature"},
        {"block": "br_stage:1:right"},
    ]})
    cache, _, _ = pipeline.refresh_pbd_cache(data, str(resources))
    br = pipeline.load_latest_br(str(resources))
    result = generate_bn(data, br, cache, date(2026, 8, 24), "06.00", "5", str(tmp_path / "out"), str(style_path))

    texts = [p.text for p in docx.Document(result.path).paragraphs]
    assert "Район зосередження мати в районі – лісосмуга (31U DQ 40700 10700);" in texts
    assert "КСП «ГОРА 1» за координатами (31U DQ 40500 10500) за сигналом ‘‘ГРІМ-101’’." in texts
    assert "звання\tПІДПИС" in texts
    assert "праворуч: точка (31U DQ 48000 11000), відповідальний сусід." in texts
    assert "ПЛАН" in texts                               # заголовок без ручного номера
    assert result.staff_values >= 1 and 0 < result.automatic_percent < 100
    lines = review_lines(result)
    assert lines[0] == "БОЙОВИЙ НАКАЗ" and f"{result.automatic_percent}%" in lines[1]
    assert any(ORIGIN_TITLES["manual"] in line for line in lines)
    assert any("31U DQ 45555 15555" in line for line in lines)


def test_missing_style_file_warns(resources, tmp_path):
    data = make_data()
    cache, _, _ = pipeline.refresh_pbd_cache(data, str(resources))
    result = generate_bn(data, None, cache, date(2026, 8, 24), "06.00", "5", str(tmp_path), str(tmp_path / "немає.docx"))
    assert any("без стилю" in w for w in result.warnings)


def test_review_lines_without_manual():
    from generators.generate_bn import BnResult, review_lines
    lines = review_lines(BnResult("x.docx", coverage={"manual": 0}, automatic_percent=100.0))
    assert "100.0%" in lines[1] and not any("Штабні значення" in line for line in lines)


def _write_extra_sources(resources):
    from bn_fixtures import RR_PARAGRAPHS, write_signal_table
    write_docx(resources / "РР 77 від 17.08.2026.docx", RR_PARAGRAPHS)
    write_docx(resources / "ДОНЕСЕННЯ_ПОЗАТЕРМІНОВЕ 17.docx", [
        "ПОЗАТЕРМІНОВЕ БОЙОВЕ ДОНЕСЕННЯ командира A0001 №17 КСП – 10:00 11.08.2026.",
        "Вогнева позиція «СТАРА» перейменовано на «ЕХО»;",
    ])
    write_signal_table(resources / "Таблиця сигналів A0001.docx")
    delta = resources / "Дельта"
    delta.mkdir()
    (delta / "БП - дельта.csv").write_text(
        '"sidc","id","quantity","name","observation_datetime","reliability_credibility","staff_comments",'
        '"platform_type","direction","speed","coordinates"\n'
        '"10032500001401000000","vr-1","","ВР-1 БнГ","","","","","","",'
        '"LINESTRING (2.3000000 48.8000000, 2.3010000 48.8000000)"\n',
        encoding="utf-8",
    )


def test_generation_with_extra_sources(resources, tmp_path):
    from generators.generate_bn import review_lines
    from utils.mgrs import latlon_to_mgrs

    _write_extra_sources(resources)
    extras, used, problems = pipeline.load_extras(str(resources))
    assert problems == [] and len(used) == 4
    assert extras.rr.number == "77" and len(extras.renames) == 1 and extras.signals.rows and extras.delta.objects

    data = make_data()
    data["staff"] = {
        "вогневі_рубежі": {"вр-1": {"title": "ВР-1 БнГ", "text": "дорога ({1}), дорога ({2})",
                                     "mgrs": ["31U DQ 46000 16000", "31U DQ 46500 16000"]}},
        "точки": {"стабпункт": {"title": "Стабпункт", "mgrs": []}, "т-1": {"title": "Точка", "mgrs": ["31U DQ 47000 17000"]}},
    }
    data["bn"]["document"]["sections"].append({"heading": "ВБАК", "content": [
        "Райони особливої уваги мати:", {"block": "rr_stage:1:rou"},
        "Командиру БнГ A0001:",
        "ВР №1 - {staff:вогневі_рубежі:вр-1} сигнал висування – «ГРОЗА-501»",
        "За сигналом ‘‘ГРІМ-101’’ відійти.",
        "- стабпункт {staff:точки:стабпункт};",
        "Точка: {staff:точки:т-1}.",
        "ВП «СТАРА» за координатами (31U DQ 49999 19999).",
    ]})
    cache, _, _ = pipeline.refresh_pbd_cache(data, str(resources))
    br = pipeline.load_latest_br(str(resources))
    result = generate_bn(data, br, cache, date(2026, 8, 24), "06.00", "5", str(tmp_path / "out"), extras=extras)

    texts = [p.text for p in docx.Document(result.path).paragraphs]
    assert "РОУ №1/1 – ЧАРЛІ (31U DQ 40300 10300);" in texts
    first, second = (latlon_to_mgrs(48.8, 2.3), latlon_to_mgrs(48.8, 2.301))
    assert f"ВР №1 - дорога ({first}), дорога ({second}) сигнал висування – «ГРОЗА-501»" in texts
    assert "Точка: (31U DQ 47000 17000)." in texts
    assert not any("стабпункт" in text for text in texts)
    assert "ВП «ЕХО» за координатами (31U DQ 40300 10300)." in texts
    assert any("позатерміновим донесенням: ВП «СТАРА» -> ВП «ЕХО»" in w for w in result.warnings)
    assert not any("СТАРА" in w and "немає в ПБД" in w for w in result.warnings)
    assert result.priority_values["ksp_mgrs"] == ("31U DQ 40500 10500", "pbd")
    assert result.signals.path and os.path.isfile(result.signals.path)
    assert result.signal_stats["total"] == 2 and result.signal_stats["known"] == 1

    review = "\n".join(review_lines(result))
    assert "Координати КСП: ПБД / РОП_ВОП" in review and "Виправлена таблиця сигналів" in review
    assert "staff:вогневі_рубежі:вр-1 «ВР-1 БнГ»: 2 точок, джерело - Дельта" in review
    assert "staff:точки:стабпункт «Стабпункт»: координат немає" in review
    assert "- зі штабних рішень data_auto.json → staff" in review

    map_result = generate_map(data, cache, str(tmp_path / "out" / "map"), str(tmp_path / "icons"), extras.delta,
                              result.staff_entries + result.unused_staff + result.empty_staff)
    assert map_result.stats["staff"] == 1 and map_result.stats["carried"] == 1


def test_missing_command_post_coordinates_warn(resources, tmp_path):
    data = make_data()
    cache, _, _ = pipeline.refresh_pbd_cache(data, str(resources))
    cache.update({"ksp_mgrs": "", "zkp_mgrs": ""})
    result = generate_bn(data, None, cache, date(2026, 8, 24), "06.00", "5", str(tmp_path))
    assert any(w.startswith("Координати КСП: немає ні в БР, ні в ПБД") for w in result.warnings)
    assert any(w.startswith("Координати ЗКП: немає ні в БР, ні в ПБД") for w in result.warnings)


def test_load_extras_reports_broken_files(tmp_path):
    (tmp_path / "РР зламаний.docx").write_bytes(b"not a docx")
    extras, used, problems = pipeline.load_extras(str(tmp_path))
    assert extras.rr is None and used == [] and "РР зламаний.docx" in problems[0]
    empty, used, problems = pipeline.load_extras(str(tmp_path / "немає"))
    assert (empty.rr, empty.signals, empty.delta, empty.renames, used, problems) == (None, None, None, [], [], [])


def test_remove_backup_files(tmp_path):
    for name in ("data.json", "data.json.bak", "data.json.bak-2026-01-01", "інше.txt"):
        (tmp_path / name).write_text("{}", encoding="utf-8")
    (tmp_path / "тека.bak").mkdir()
    removed = pipeline.remove_backup_files(str(tmp_path))
    assert sorted(os.path.basename(p) for p in removed) == ["data.json.bak", "data.json.bak-2026-01-01"]
    assert set(os.listdir(tmp_path)) == {"data.json", "інше.txt", "тека.bak"}
