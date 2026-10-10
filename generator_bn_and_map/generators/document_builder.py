"""bn.document з data.json -> python-docx Document.

Формат item (як у сусідніх проєктах - final_combat_report paragraph_two):
  "текст"                              - звичайний абзац;
  ["текст", ["жирна фраза", …]]        - абзац з виділеними фразами;
  {"text", "bold", "fmt", "align", "indent", "left_indent", "italic", "size",
   "page_break_after"}                 - явне форматування; "fmt" - прототип
                                         абзацу зі "Стиль БН.docx" (точно як
                                         у зразку), решта - коли стилю немає;
  {"heading": "…", "content": [...], "fmt"} - заголовок (вкладеність = рівень);
                                         без стилю - автонумерація "N."/"N.M.",
                                         зі стилем - нумерація самого Word;
  {"block": "назва", "prefix", "stop_before", "skip", "only", "nth", "fmt", "fmt_by_kind", "fmt_list",
   "fallback"}                         - текст з БР/ПБД ("fallback" - items,
                                         якщо блок порожній);
  {"page_break": true}, {"blank": N, "fmt"}.
У статичних текстах підставляються {плейсхолдери} і оновлюються координати
названих позицій (content/position_refresh.py)."""
import re
from dataclasses import dataclass

from docx import Document
from docx.shared import Cm

from constants import CLASSIFICATION_MARKING
from content.docx_blocks import item_text, with_text
from utils.docx_utils import (
    ALIGNMENTS, FONT_SIZE, add_classification_header_footer, add_paragraph_with_style, add_table, set_margins,
)
from utils.mgrs import find_mgrs

PLACEHOLDER_RE = re.compile(r"\{(\w+)\}")
BODY_INDENT_CM = 1.25
_DEFAULT_ACK_COLUMNS = ["Посада", "Військове звання", "Підпис", "П.І.Б", "Дата ознайомлення"]
_DEFAULT_ACK_WIDTHS_CM = [3.97, 3.98, 3.11, 4.33, 3.46]
DEFAULT_ACK_TABLE = "table:ack"

# Умовні сигнали штабу на кшталт ‘‘НАЗВА-123’’ / «НАЗВА 45».
_SIGNAL_RE = re.compile(r"[«‘'\"“]{1,2}\s*(?P<name>[А-ЯІЇЄҐA-Z][А-ЯІЇЄҐA-Z'’]+\s*[-–]?\s*\d{2,4})\s*[»’'\"”]{1,2}")
_POSITION_START_RE = re.compile(r"^[-–—•\s]*(?:СПАР|КСП|ВОП|РОП|ХАБ|ВП|ПВ|СП|ВЗ|ТЗ|ПУ)(?!\w)")

ORIGIN_SOURCE = "source"        # блок з БР/ПБД
ORIGIN_REFRESHED = "refreshed"  # статичний текст з підставленими/оновленими даними
ORIGIN_STANDARD = "standard"    # статичний текст без ситуаційних даних
ORIGIN_MANUAL = "manual"        # статичний текст з координатами чи сигналами, яких немає в джерелах


def line_kind(text):
    """Вид рядка блоку - щоб підібрати йому формат зі зразка."""
    text = (text or "").strip()
    if not text:
        return "empty"
    if text.endswith(":") and len(text) <= 80 and not find_mgrs(text):
        return "header"
    if re.match(r"^ВР\s*№", text):
        return "vr"
    if re.match(r"^\d{2,4}\s{2,}", text):
        return "target"
    if text.startswith("Сусід"):
        return "neighbour"
    if re.match(r"^(праворуч|ліворуч)\b", text, re.I):
        return "side"
    if _POSITION_START_RE.match(text):
        return "position"
    if text[0] in "-–—•":
        return "dash"
    return "body"


def _as_list(value):
    if not value:
        return []
    return [value] if isinstance(value, str) else list(value)


def apply_block_options(items, options):
    """stop_before/skip/only - префікси абзаців (без урахування регістру): блок
    обрізається перед першим stop_before, абзаци зі skip відкидаються, з only
    лишаються лише ті, що з них починаються; nth - лише один абзац (за
    порядком) з того, що лишилось."""
    stops = [s.lower() for s in _as_list(options.get("stop_before"))]
    skips = [s.lower() for s in _as_list(options.get("skip"))]
    only = [s.lower() for s in _as_list(options.get("only"))]
    result = []
    for item in items:
        text = item_text(item).lower()
        if any(text.startswith(stop) for stop in stops):
            break
        if any(text.startswith(skip) for skip in skips):
            continue
        if only and not any(text.startswith(prefix) for prefix in only):
            continue
        result.append(item)
    nth = options.get("nth")
    if nth is not None:
        return result[nth:nth + 1]
    return result


def _segments(text, bold):
    """-> [(текст, жирний?)]; None - жирність як у прототипу абзацу."""
    if bold is True:
        return [(text, True)]
    if not bold:
        return [(text, None)]
    segments, remaining = [], text
    while remaining:
        positions = [(remaining.find(p), p) for p in bold if p and remaining.find(p) != -1]
        if not positions:
            segments.append((remaining, False))
            break
        pos, phrase = min(positions)
        if pos:
            segments.append((remaining[:pos], False))
        segments.append((phrase, True))
        remaining = remaining[pos + len(phrase):]
    return segments


@dataclass
class LoggedParagraph:
    origin: str
    path: str
    text: str
    staff_values: tuple = ()


def _signal_key(name):
    return re.sub(r"\W", "", name or "").upper()


class DocumentBuilder:
    def __init__(self, context, blocks, style=None, refresher=None, known_names=(), source_mgrs=()):
        """source_mgrs - координати з поточних БР/ПБД/РОП_ВОП: такі координати в
        статичному тексті вважаються підтвердженими джерелами."""
        self.context = context
        self.blocks = blocks
        self.style = style
        self.refresher = refresher
        self.known_names = {_signal_key(n) for n in known_names if n}
        self.source_mgrs = set(source_mgrs)
        self.unknown_placeholders = set()
        self.missing_blocks = []
        self.fallback_blocks = []
        self.used_blocks = {}
        self.formats = {}
        self.log = []
        self.refreshed_mentions = 0
        self.renamed = []
        self.unknown_mentions = []
        self._path = []
        self._render_mgrs = set()
        self._render_signals = set()

    # --- текст ----------------------------------------------------------------
    def render(self, text):
        def replace(m):
            key = m.group(1)
            if key in self.context and self.context[key] is not None:
                value = str(self.context[key])
                self._render_mgrs.update(c.normalized for c in find_mgrs(value) if c.valid)
                self._render_signals.add(_signal_key(value))
                return value
            self.unknown_placeholders.add(key)
            return m.group(0)

        return PLACEHOLDER_RE.sub(replace, str(text))

    def _render_static(self, raw):
        """Плейсхолдери + оновлення координат. -> (текст, походження, штабні
        значення - координати й сигнали, яких немає в джерелах)."""
        self._render_mgrs = set()
        self._render_signals = set()
        text = self.render(raw)
        substituted = text != str(raw)
        auto = set(self._render_mgrs)
        if self.refresher is not None:
            result = self.refresher.refresh(text)
            if result.text != text:
                substituted = True
            text = result.text
            auto |= result.auto_mgrs
            self.refreshed_mentions += result.refreshed
            self.renamed.extend(result.renamed)
            self.unknown_mentions.extend(result.unknown)
        staff = [c.raw for c in find_mgrs(text)
                 if c.valid and c.normalized not in auto and c.normalized not in self.source_mgrs]
        staff += [m.group(0) for m in _SIGNAL_RE.finditer(text)
                  if _signal_key(m.group("name")) not in self.known_names | self._render_signals]
        if staff:
            return text, ORIGIN_MANUAL, tuple(staff)
        return text, ORIGIN_REFRESHED if substituted else ORIGIN_STANDARD, ()

    def _log(self, origin, text, staff_values=()):
        if text and text.strip():
            self.log.append(LoggedParagraph(origin, " / ".join(self._path), text, tuple(staff_values)))

    # --- документ -------------------------------------------------------------
    def build(self, spec):
        self.formats = spec.get("formats") or {}
        if self.style is not None:
            doc = self.style.new_document()
        else:
            doc = Document()
            set_margins(doc, Cm(2), Cm(2), Cm(3), Cm(1))
            add_classification_header_footer(doc, CLASSIFICATION_MARKING)

        title_page = spec.get("title_page") or []
        if title_page:
            self.add_content(doc, title_page, default_align="center", default_indent=False, default_fmt="title")
            if not any(isinstance(e, dict) and e.get("page_break_after") for e in title_page):
                doc.add_page_break()
        self.add_content(doc, spec.get("header") or [], default_align="left", default_indent=False, default_fmt="header")
        self.add_content(doc, spec.get("sections") or [])
        self._path = []
        self.add_content(doc, spec.get("signature") or [], default_align="left", default_indent=False, default_fmt="signature")
        if spec.get("acknowledgement"):
            self.add_acknowledgement(doc, spec["acknowledgement"])
        self.add_content(doc, spec.get("closing") or [], default_align="left", default_indent=False, default_fmt="closing")
        return doc

    def add_content(self, doc, content, level=1, prefix="", default_align="justify", default_indent=True, default_fmt="body"):
        counter = 0
        for entry in content:
            if isinstance(entry, dict) and "heading" in entry:
                numbered = entry.get("number", True)
                heading_text = self.render(entry["heading"])
                fmt = entry.get("fmt") or self.formats.get(f"heading{level}")
                if self.style is not None and fmt and self.style.has(fmt):
                    # Нумерацію ("1.", "3.1.") дає сам стиль заголовка у шаблоні.
                    self.style.add_paragraph(doc, fmt, [(heading_text, None)])
                    number = prefix
                else:
                    if numbered:
                        counter += 1
                        number = f"{prefix}{counter}."
                        heading_text = f"{number} {heading_text}"
                    else:
                        number = prefix
                    add_paragraph_with_style(
                        doc, heading_text, bold=True,
                        alignment=ALIGNMENTS.get(entry.get("align", "justify"), ALIGNMENTS["justify"]),
                        first_line_indent_cm=self._indent(entry.get("indent", True)),
                        heading_level=min(level, 9),
                    )
                self._path = self._path[:level - 1] + [self.render(entry["heading"])[:60]]
                self._log(ORIGIN_STANDARD, heading_text)
                self.add_content(doc, entry.get("content") or [], level + 1, number if numbered else prefix)
            else:
                self.add_item(doc, entry, default_align=default_align, default_indent=default_indent, default_fmt=default_fmt)

    @staticmethod
    def _indent(value):
        if value is True:
            return BODY_INDENT_CM
        if value in (False, None):
            return 0.0
        return float(value)

    def _emit(self, doc, text, bold, options, default_align, default_indent, default_fmt):
        fmt = options.get("fmt") or self.formats.get(default_fmt) or self.formats.get("body")
        if self.style is not None and fmt and self.style.has(fmt):
            self.style.add_paragraph(doc, fmt, _segments(text, bold) if text else [],
                                     page_break_after=bool(options.get("page_break_after")))
            return
        add_paragraph_with_style(
            doc, text, bold=bold,
            alignment=ALIGNMENTS.get(options.get("align", default_align), ALIGNMENTS["justify"]),
            first_line_indent_cm=self._indent(options.get("indent", default_indent)),
            format_tabs="\t" in text,
            italic=bool(options.get("italic")),
            font_size=options.get("size", FONT_SIZE),
            left_indent_cm=options.get("left_indent"),
        )
        if options.get("page_break_after"):
            doc.add_page_break()

    def add_item(self, doc, item, default_align="justify", default_indent=True, default_fmt="body", render=True, fmt=None):
        options = {}
        if isinstance(item, str):
            text, bold = item, False
        elif isinstance(item, list):
            text = item[0] if item else ""
            bold = list(item[1]) if len(item) > 1 else False
        elif isinstance(item, dict):
            if "block" in item:
                self.add_block(doc, item)
                return
            if item.get("page_break"):
                doc.add_page_break()
                return
            if "blank" in item:
                for _ in range(int(item["blank"] or 0)):
                    self._emit(doc, "", False, item, default_align, default_indent, default_fmt)
                return
            text, bold, options = item.get("text", ""), item.get("bold", False), item
        else:
            return
        if fmt:
            options = {**options, "fmt": fmt}

        staff = ()
        if render:
            text, origin, staff = self._render_static(text)
            if isinstance(bold, list):
                bold = [self.render(phrase) for phrase in bold]
        else:
            origin = ORIGIN_SOURCE
        self._emit(doc, text, bold, options, default_align, default_indent, default_fmt)
        self._log(origin, text, staff)

    def add_block(self, doc, item):
        name = item["block"]
        items = apply_block_options(self.blocks.get(name) or [], item)
        if not items:
            fallback = item.get("fallback")
            if fallback:
                self.fallback_blocks.append(name)
                self.add_content(doc, fallback, len(self._path) + 1)
                return
            self.missing_blocks.append(name)
            if self.style is not None and self.formats.get("body") and self.style.has(self.formats["body"]):
                self.style.add_paragraph(doc, self.formats["body"], [(f"‹немає даних: {name}›", None)])
                self._highlight_last(doc)
            else:
                add_paragraph_with_style(doc, f"‹немає даних: {name}›", highlight=True, first_line_indent_cm=BODY_INDENT_CM)
            return
        self.used_blocks[name] = len(items)
        prefix = self.render(item.get("prefix", ""))
        by_kind = item.get("fmt_by_kind") or {}
        by_position = item.get("fmt_list") or []
        for index, block_item in enumerate(items):
            if index == 0 and prefix:
                block_item = with_text(block_item, prefix + item_text(block_item))
            fmt = (by_position[min(index, len(by_position) - 1)] if by_position else None) \
                or by_kind.get(line_kind(item_text(block_item))) or item.get("fmt")
            # Текст із джерел не містить плейсхолдерів - не обробляється, щоб
            # випадкові фігурні дужки в ньому не давали хибних попереджень.
            self.add_item(doc, block_item, render=False, fmt=fmt)

    @staticmethod
    def _highlight_last(doc):
        from docx.enum.text import WD_COLOR_INDEX
        paragraph = doc.paragraphs[-1]
        for run in paragraph.runs:
            run.font.highlight_color = WD_COLOR_INDEX.YELLOW

    def add_acknowledgement(self, doc, ack):
        columns = ack.get("columns") or _DEFAULT_ACK_COLUMNS
        rows = [[self.render(cell) for cell in row] for row in ack.get("rows", [])]
        blank_rows = int(ack.get("blank_rows", 0) or 0)
        title = self.render(ack.get("title", ""))
        table_id = ack.get("table", DEFAULT_ACK_TABLE)
        if self.style is not None and table_id in self.style.tables:
            self.style.add_table(doc, table_id, title, columns, rows, blank_rows)
        else:
            add_table(doc, title, columns, rows + [[""] * len(columns) for _ in range(blank_rows)],
                      ack.get("widths_cm") or _DEFAULT_ACK_WIDTHS_CM, font_size=ack.get("font_size", 12))
        self._log(ORIGIN_STANDARD, " ".join([title] + [" ".join(r) for r in rows]))

    # --- звіт -----------------------------------------------------------------
    def coverage(self):
        """Частка тексту (за кількістю символів) за походженням, %."""
        totals = {ORIGIN_SOURCE: 0, ORIGIN_REFRESHED: 0, ORIGIN_STANDARD: 0, ORIGIN_MANUAL: 0}
        for entry in self.log:
            totals[entry.origin] += len(entry.text)
        all_chars = sum(totals.values()) or 1
        return {origin: round(100 * chars / all_chars, 1) for origin, chars in totals.items()}

    def automatic_percent(self):
        """Частка тексту, яку пише програма: усе, крім самих штабних значень
        (координат і сигналів, яких немає в БР/ПБД/РОП_ВОП)."""
        all_chars = sum(len(entry.text) for entry in self.log) or 1
        staff_chars = sum(len(value) for entry in self.log for value in entry.staff_values)
        return round(100 - 100 * staff_chars / all_chars, 1)

    def manual_paragraphs(self):
        return [entry for entry in self.log if entry.origin == ORIGIN_MANUAL]


_POSITION_MENTION_RE = re.compile(
    r"\b(?:КСП|ВОП|РОП|ВП|ПВ|СП|ВЗ|ТЗ|ХАБ|ПУ|СПАР)\b[^«\"“„‘’'\n]{0,30}?[«\"“„‘’']{1,2}\s*([^«»\"“”„‘’'\n]{2,40}?)\s*[»\"”“’']{1,2}"
)


def _static_texts(content):
    for entry in content or []:
        if isinstance(entry, dict):
            if "heading" in entry:
                yield str(entry["heading"])
                yield from _static_texts(entry.get("content"))
            elif "text" in entry:
                yield str(entry["text"])
            elif "prefix" in entry:
                yield str(entry["prefix"])
            if "fallback" in entry:
                yield from _static_texts(entry.get("fallback"))
        elif isinstance(entry, (str, list)):
            yield item_text(entry)


def find_stale_position_names(spec, known_names, normalize):
    """Кодові назви позицій, згадані в СТАТИЧНИХ текстах bn.document, яких
    немає серед відомих (з ПБД/РОП_ВОП) - ймовірно застарілий текст."""
    known = {normalize(name) for name in known_names if name}

    def is_known(name):
        # "ГОРА" в тексті й "ГОРА 5" у назві КСП - та сама позиція.
        key = normalize(name)
        return key in known or (len(key) >= 4 and any(k.startswith(key) or key.startswith(k) for k in known if len(k) >= 4))

    stale = []
    for part in ("title_page", "header", "sections", "signature", "closing"):
        for text in _static_texts(spec.get(part)):
            for m in _POSITION_MENTION_RE.finditer(text):
                name = m.group(1).strip()
                if not is_known(name) and name.upper() not in stale:
                    stale.append(name.upper())
    return stale
