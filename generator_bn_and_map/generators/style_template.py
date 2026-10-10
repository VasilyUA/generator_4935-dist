"""Стиль БН, знятий з готового (еталонного) документа.

"Стиль БН.docx" (resources/, gitignored) - копія еталонного БН, у тілі якої
замість тексту лежать ПРОТОТИПИ: по одному абзацу на кожен набір
форматування (відступи, вирівнювання, інтервали, нумерація, шрифт першого
run-у, малюнок/рамка, якщо були) з текстом-міткою "⟦fNNN⟧", і таблиця
аркуша доведення після мітки "⟦table:ack⟧". Колонтитули, поля, стилі,
нумерація заголовків і налаштування документа лишаються від еталона - тож
згенерований БН виглядає так само, як зразок.

data.json["bn"]["document"] посилається на прототипи ключем "fmt"."""
import copy
import re

from docx import Document
from lxml import etree
from docx.oxml.ns import qn
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P
from docx.text.run import Run

MARKER_RE = re.compile(r"^⟦([\w:.\-]+)⟧$")
ACK_TABLE_ID = "table:ack"

_NON_TEXT_RUN_CHILDREN = {qn("w:drawing"), qn("w:pict"), qn("w:object")}
_DROPPED_TAGS = {
    qn("w:bookmarkStart"), qn("w:bookmarkEnd"), qn("w:proofErr"), qn("w:permStart"), qn("w:permEnd"),
    qn("w:del"), qn("w:moveFrom"), qn("w:commentRangeStart"), qn("w:commentRangeEnd"),
}
_UNWRAPPED_TAGS = {qn("w:ins"), qn("w:moveTo"), qn("w:smartTag"), qn("w:hyperlink")}


def _strip_rsids(element):
    for el in element.iter():
        for key in [k for k in el.attrib if "rsid" in k]:
            del el.attrib[key]


def _is_text_run(run):
    return not any(child.tag in _NON_TEXT_RUN_CHILDREN for child in run)


def _has_text(run):
    return run.find(qn("w:t")) is not None


def _flatten(paragraph):
    """Прибирає закладки/правки й розгортає обгортки (гіперпосилання, вставки),
    щоб усі runs стали прямими дітьми абзацу."""
    for el in list(paragraph.iter()):
        if el is paragraph or el.getparent() is None:
            continue
        if el.tag in _DROPPED_TAGS:
            el.getparent().remove(el)
    changed = True
    while changed:
        changed = False
        for el in list(paragraph.iter()):
            if el is not paragraph and el.tag in _UNWRAPPED_TAGS and el.getparent() is not None:
                parent = el.getparent()
                index = parent.index(el)
                for child in list(el):
                    if child.tag == qn("w:rPr"):
                        continue
                    parent.insert(index, child)
                    index += 1
                parent.remove(el)
                changed = True


def paragraph_skeleton(p_element):
    """Абзац -> (копія без текстових runs, rPr першого текстового run-у).
    Малюнки й рамки (w:drawing, w:pict) лишаються в копії."""
    p = copy.deepcopy(p_element)
    _strip_rsids(p)
    _flatten(p)
    first_rpr = None
    for run in list(p.iter(qn("w:r"))):
        if not _is_text_run(run):
            continue
        if first_rpr is None and _has_text(run):
            rpr = run.find(qn("w:rPr"))
            first_rpr = copy.deepcopy(rpr) if rpr is not None else None
        run.getparent().remove(run)
    return p, first_rpr


def format_signature(p_element):
    """Ключ набору форматування абзацу: pPr + rPr першого текстового run-у;
    абзаци з малюнком/рамкою отримують власний ключ (щоб малюнок не
    копіювався в інші абзаци)."""
    skeleton, rpr = paragraph_skeleton(p_element)
    ppr = skeleton.find(qn("w:pPr"))
    has_object = any(child.tag in _NON_TEXT_RUN_CHILDREN for child in skeleton.iter())
    return (
        etree.tostring(ppr) if ppr is not None else b"",
        etree.tostring(rpr) if rpr is not None else b"",
        id(p_element) if has_object else None,
    )


def _append_text_runs(p, rpr, segments):
    for text, bold in segments:
        if not text:
            continue
        r = p.makeelement(qn("w:r"), {})
        if rpr is not None:
            r.append(copy.deepcopy(rpr))
        p.append(r)
        run = Run(r, None)
        run.text = text
        if bold is not None:
            run.bold = bold


def _append_page_break(p):
    r = p.makeelement(qn("w:r"), {})
    br = r.makeelement(qn("w:br"), {qn("w:type"): "page"})
    r.append(br)
    p.append(r)


class StyleTemplate:
    def __init__(self, path):
        self.path = path
        self.document = Document(path)
        self.prototypes = {}
        self.tables = {}
        self._harvest()

    def _harvest(self):
        body = self.document.element.body
        pending_table = None
        for child in list(body.iterchildren()):
            if isinstance(child, CT_P):
                skeleton, rpr = paragraph_skeleton(child)
                text = "".join(t.text or "" for t in child.iter(qn("w:t"))).strip()
                m = MARKER_RE.match(text)
                if m:
                    if m.group(1).startswith("table:"):
                        pending_table = m.group(1)
                    else:
                        self.prototypes[m.group(1)] = (skeleton, rpr)
            elif isinstance(child, CT_Tbl) and pending_table:
                self.tables[pending_table] = copy.deepcopy(child)
                pending_table = None
            if child.tag != qn("w:sectPr"):
                body.remove(child)

    def has(self, fmt):
        return fmt in self.prototypes

    def new_document(self):
        """Порожній документ з усіма налаштуваннями еталона (викликати один раз на
        генерацію - повертає той самий об'єкт, з якого вже прибрано прототипи)."""
        return self.document

    def _insert(self, doc, element):
        body = doc.element.body
        sect = body.find(qn("w:sectPr"))
        if sect is not None:
            sect.addprevious(element)
        else:
            body.append(element)
        return element

    def add_paragraph(self, doc, fmt, segments=(), page_break_after=False):
        """segments - [(текст, bold або None)]; None - жирність як у прототипу."""
        skeleton, rpr = self.prototypes[fmt]
        p = copy.deepcopy(skeleton)
        _append_text_runs(p, rpr, segments)
        if page_break_after:
            _append_page_break(p)
        return self._insert(doc, p)

    def add_table(self, doc, table_id, title, columns, rows, blank_rows):
        tbl = copy.deepcopy(self.tables[table_id])
        _strip_rsids(tbl)
        trs = tbl.findall(qn("w:tr"))
        has_title = bool(title) and len(trs) >= 3
        header_index = 1 if has_title else 0
        data_prototype = trs[header_index + 1] if len(trs) > header_index + 1 else trs[header_index]
        blank_prototype = trs[-1]
        for tr in trs[header_index + 1:]:
            tbl.remove(tr)
        if has_title:
            _set_cell_lines(trs[0].findall(qn("w:tc"))[0], title.split("\n"))
        for tc, column in zip(trs[header_index].findall(qn("w:tc")), columns):
            _set_cell_lines(tc, str(column).split("\n"))
        for row in rows:
            tr = copy.deepcopy(data_prototype)
            for index, tc in enumerate(tr.findall(qn("w:tc"))):
                _set_cell_lines(tc, [str(row[index]) if index < len(row) else ""])
            tbl.append(tr)
        for _ in range(blank_rows):
            tr = copy.deepcopy(blank_prototype)
            for tc in tr.findall(qn("w:tc")):
                _set_cell_lines(tc, [""])
            tbl.append(tr)
        return self._insert(doc, tbl)


def _set_cell_lines(tc, lines):
    """Текст клітинки: кожен рядок - окремий абзац; форматування береться з
    наявних абзаців клітинки (зайві прибираються, бракуючі копіюють останній)."""
    paragraphs = tc.findall(qn("w:p"))
    skeletons = [paragraph_skeleton(p) for p in paragraphs] or []
    for p in paragraphs:
        tc.remove(p)
    for index, line in enumerate(lines):
        skeleton, rpr = skeletons[min(index, len(skeletons) - 1)] if skeletons else (tc.makeelement(qn("w:p"), {}), None)
        p = copy.deepcopy(skeleton)
        _append_text_runs(p, rpr, [(line, None)])
        tc.append(p)


def create_style_template(reference_path, output_path, assign_ids=None):
    """Еталонний БН -> "Стиль БН.docx". Повертає {індекс елемента тіла: fmt}.
    Однакові набори форматування отримують один fmt ("f001", "f002", ...);
    таблиці (у порядку появи) - мітки з assign_ids["tables"] або "table:N"."""
    doc = Document(reference_path)
    body = doc.element.body
    children = list(body.iterchildren())
    by_signature, mapping, prototypes, tables = {}, {}, [], []
    table_names = list((assign_ids or {}).get("tables", []))
    for index, child in enumerate(children):
        if isinstance(child, CT_P):
            signature = format_signature(child)
            if signature not in by_signature:
                fmt = f"f{len(by_signature) + 1:03d}"
                by_signature[signature] = fmt
                prototypes.append((fmt, child))
            mapping[index] = by_signature[signature]
        elif isinstance(child, CT_Tbl):
            name = table_names.pop(0) if table_names else f"table:{len(tables) + 1}"
            tables.append((name, child))
            mapping[index] = name
    sect = body.find(qn("w:sectPr"))
    for child in children:
        if child is not sect:
            body.remove(child)
    for fmt, original in prototypes:
        skeleton, rpr = paragraph_skeleton(original)
        _append_text_runs(skeleton, rpr, [(f"⟦{fmt}⟧", None)])
        sect.addprevious(skeleton)
    for name, original in tables:
        marker = body.makeelement(qn("w:p"), {})
        _append_text_runs(marker, None, [(f"⟦{name}⟧", None)])
        sect.addprevious(marker)
        table_copy = copy.deepcopy(original)
        _strip_rsids(table_copy)
        sect.addprevious(table_copy)
    doc.save(output_path)
    return mapping
