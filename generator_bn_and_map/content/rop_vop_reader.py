"""Додаток "Перелік … опорних пунктів" (РОП/ВОП, .xlsx) -> позиції з
кількістю особового складу і координатами, рядок ЗКП.

Колонки шукаються за текстом заголовків (рядок груп "БРО / РОП / ВОП /
позиції…" і під ним "назва / підрозділ / к-ть о/с"), а не за фіксованими
літерами - таблицю правлять вручну, колонки можуть зсуватись."""
import re
import warnings
from dataclasses import dataclass, field

from openpyxl import load_workbook

from utils.mgrs import find_mgrs

_HEADER_SCAN_ROWS = 15
_QUOTED_RE = re.compile(r"[«\"“„”'‘’]{1,2}\s*([^«»\"“”„'‘’]+?)\s*[»\"”“'’‘]{1,2}")
_PARENTHESIZED_RE = re.compile(r"\(\s*([^()]+?)\s*\)")

# (тип, перевірка) - порядок важливий: "СПАР" раніше за "СП", "Вогнева засідка"
# раніше за загальне "Вогнева позиція".
_TYPE_RULES = (
    ("БРО", lambda t: t.startswith("БРО")),
    ("РОП", lambda t: t.startswith("РОП")),
    ("КСП", lambda t: t.startswith("КСП")),
    ("ВОП", lambda t: t.startswith("ВОП")),
    ("ЗКП", lambda t: t.startswith("ЗКП")),
    ("ВЗ", lambda t: "вогнева засідка" in t.lower()),
    ("ВП", lambda t: "вогнева позиція" in t.lower()),
    ("ПВ", lambda t: "позиція відділення" in t.lower()),
    ("СП", lambda t: "спостережна позиція" in t.lower()),
    ("СПАР", lambda t: t.startswith("СПАР")),
    ("СП", lambda t: re.match(r"^СП\b", t) is not None),
    ("ХАБ", lambda t: t.startswith("ХАБ")),
    ("ПУ", lambda t: re.match(r"^ПУ\b", t) is not None),
    ("ТЗ", lambda t: re.match(r"^ТЗ\b", t) is not None or re.match(r"^Пункт\s+(БпЛА|зльоту)", t, re.I) is not None),
)


@dataclass
class RopVopEntry:
    type: str
    name: str
    quantity: int | None
    mgrs: list
    row: int


@dataclass
class RopVopTable:
    path: str
    entries: list = field(default_factory=list)
    zkp_mgrs: str = ""
    warnings: list = field(default_factory=list)


def _text(value):
    return "" if value is None else str(value).strip()


def classify(text):
    """Назва з клітинки -> (тип, назва) або (None, ""). Назва - з лапок, а якщо
    їх немає - з дужок ("ХАБ (Назва) …")."""
    text = text.strip()
    if not text or text.lower().startswith("всього"):
        return None, ""
    type_ = next((name for name, rule in _TYPE_RULES if rule(text)), None)
    if type_ is None:
        return None, ""
    m = _QUOTED_RE.search(text) or _PARENTHESIZED_RE.search(text)
    return type_, (m.group(1).strip().upper() if m else "")


def _to_int(value):
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return int(value)
    text = _text(value)
    return int(text) if text.isdigit() else None


def _detect_columns(ws):
    """-> {"coords": col, "levels": [(name_col, count_col), …] від найдрібнішого
    рівня (позиції) до найбільшого (БРО)}, або None."""
    group_row = None
    for row in ws.iter_rows(min_row=1, max_row=_HEADER_SCAN_ROWS):
        labels = {_text(c.value).upper(): c.column for c in row if _text(c.value)}
        if "РОП" in labels and "ВОП" in labels:
            group_row = (row[0].row, labels)
            break
    if group_row is None:
        return None
    row_index, labels = group_row
    group_starts = sorted(col for label, col in labels.items() if label in ("БРО", "РОП", "ВОП") or label.startswith("ПОЗИЦІЇ"))
    sub = {c.column: _text(c.value).lower() for c in ws[row_index + 1] if _text(c.value)}

    levels = []
    for i, start in enumerate(group_starts):
        end = group_starts[i + 1] if i + 1 < len(group_starts) else max(list(sub) + [start]) + 1
        in_range = {col: label for col, label in sub.items() if start <= col < end}
        name_col = next((col for col, label in sorted(in_range.items()) if label == "назва"), start)
        count_col = next((col for col, label in sorted(in_range.items()) if "к-ть" in label), None)
        levels.append((name_col, count_col))

    coords_col = None
    for row in ws.iter_rows(min_row=1, max_row=_HEADER_SCAN_ROWS):
        for c in row:
            if "місце дислокації" in _text(c.value).lower():
                coords_col = c.column
                break
        if coords_col:
            break
    if coords_col is None:
        return None
    return {"coords": coords_col, "levels": list(reversed(levels)), "data_start": row_index + 2}


def read_rop_vop(path):
    table = RopVopTable(path=path)
    with warnings.catch_warnings():
        warnings.simplefilter("ignore", UserWarning)
        ws = load_workbook(path, data_only=True).worksheets[0]
    columns = _detect_columns(ws)
    if columns is None:
        table.warnings.append("РОП/ВОП: не знайдено заголовки колонок (БРО/РОП/ВОП, «Місце дислокації…») - файл пропущено.")
        return table

    position_count_col = columns["levels"][0][1]
    for row in ws.iter_rows(min_row=columns["data_start"]):
        values = {c.column: c.value for c in row}
        for level_index, (name_col, count_col) in enumerate(columns["levels"]):
            type_, name = classify(_text(values.get(name_col)))
            if type_ is None:
                continue
            quantity = _to_int(values.get(count_col)) if count_col else None
            if quantity is None and position_count_col:
                # "Окремі позиції" записують назву в колонку БРО, а кількість -
                # у колонку позицій.
                quantity = _to_int(values.get(position_count_col))
            mgrs = [m.normalized for m in find_mgrs(_text(values.get(columns["coords"]))) if m.valid]
            table.entries.append(RopVopEntry(type_, name, quantity, mgrs, row[0].row))
            if type_ == "ЗКП" and mgrs and not table.zkp_mgrs:
                table.zkp_mgrs = mgrs[0]
            break
    return table
