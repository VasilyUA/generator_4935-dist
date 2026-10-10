"""Форматування python-docx - копія потрібних функцій з
generator_br_and_report_for_money/formatting/docx_utils.py (проєкти в цьому
репо не імпортують код один одного: той модуль під час імпорту тягне свій
constants.py з власним data.json)."""
import os
import time
from typing import Union

from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX, WD_TAB_ALIGNMENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Emu, Inches, Pt, RGBColor

FONT_NAME = "Times New Roman"
FONT_SIZE = 14

ALIGNMENTS = {
    "left": WD_ALIGN_PARAGRAPH.LEFT,
    "center": WD_ALIGN_PARAGRAPH.CENTER,
    "right": WD_ALIGN_PARAGRAPH.RIGHT,
    "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
}


def _close_matching_word_documents(matches):
    """Закриває (без збереження) кожен відкритий у Word документ, для якого
    matches(абсолютний_шлях) - True. Якщо Word не запущено або pywin32
    недоступний — тихо виходить."""
    try:
        import pythoncom
        import win32com.client
    except ImportError:
        return

    try:
        pythoncom.CoInitialize()
        try:
            word = win32com.client.GetActiveObject("Word.Application")
        except Exception:
            return
        for doc in list(word.Documents):
            try:
                full_name = os.path.abspath(doc.FullName)
            except Exception:
                continue
            if matches(full_name):
                doc.Close(SaveChanges=False)
    finally:
        pythoncom.CoUninitialize()


def save_docx_safely(doc, file_path, attempts=5, delay=0.5):
    """doc.save(), але спершу закриває цей самий файл у Word, якщо він там
    відкритий (напр. попередній БН з такою самою назвою), і повторює спробу,
    якщо файл лишається заблокованим ще коротку мить (антивірус/OneDrive)."""
    abs_path = os.path.abspath(file_path)
    _close_matching_word_documents(lambda full_name: full_name == abs_path)
    for attempt in range(attempts):
        try:
            doc.save(file_path)
            return
        except PermissionError as e:
            if attempt == attempts - 1:
                raise PermissionError(
                    f"Не вдалось зберегти '{file_path}': файл заблоковано іншою програмою. "
                    f"Закрийте його (напр. Word) і спробуйте ще раз. ({e})"
                ) from e
            time.sleep(delay)


def set_margins(doc, top, bottom, left, right):
    for section in doc.sections:
        section.page_height = Inches(11.69)  # 297 мм
        section.page_width = Inches(8.27)    # 210 мм
        section.top_margin = top
        section.bottom_margin = bottom
        section.left_margin = left
        section.right_margin = right


def _set_paragraph_mark_font(paragraph, font_name, font_size):
    # Без цього порожній абзац (без жодного run-у з текстом) показує шрифт теми
    # замість заданого, бо Word бере шрифт кінцевого маркера абзацу (pPr/rPr).
    rPr = OxmlElement('w:rPr')
    rFonts = OxmlElement('w:rFonts')
    rFonts.set(qn('w:ascii'), font_name)
    rFonts.set(qn('w:hAnsi'), font_name)
    rFonts.set(qn('w:eastAsia'), font_name)
    rPr.append(rFonts)
    color = OxmlElement('w:color')
    color.set(qn('w:val'), '000000')
    rPr.append(color)
    sz = OxmlElement('w:sz')
    sz.set(qn('w:val'), str(int(font_size * 2)))
    rPr.append(sz)
    paragraph._p.get_or_add_pPr().append(rPr)


def _style_run(run, font_name, font_size, bold=False, italic=False):
    run.font.name = font_name
    run._element.rPr.rFonts.set(qn('w:eastAsia'), font_name)
    run.font.size = Pt(font_size)
    run.font.color.rgb = RGBColor(0, 0, 0)
    run.font.bold = bold
    run.font.italic = italic
    return run


def add_paragraph_with_style(doc, text, font_name=FONT_NAME, font_size=FONT_SIZE, bold: Union[bool, list] = False,
                             alignment=WD_ALIGN_PARAGRAPH.JUSTIFY, first_line_indent_cm=0.0, format_tabs=False,
                             italic=False, highlight=False, heading_level=None, left_indent_cm=None):
    """bold - True/False для всього абзацу або список фраз, які треба виділити
    жирним (той самий формат, що й у сусідніх проєктах). format_tabs - правий
    табулятор на всю ширину тексту (рядки "звання<TAB>ПІБ" у підписах).
    left_indent_cm - відступ зліва (може бути від'ємним - на поле)."""
    paragraph = doc.add_paragraph()
    if heading_level is not None:
        # Структура/навігація документа - як заголовок, але вигляд задається
        # прямими runs нижче. keep_with_next стилю Heading вимикається, щоб
        # заголовок не тягнув за собою на нову сторінку довгий наступний текст.
        paragraph.style = f"Heading {heading_level}"
        paragraph.paragraph_format.keep_with_next = False
    paragraph.alignment = alignment
    paragraph.paragraph_format.line_spacing = 1.0
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.first_line_indent = Cm(first_line_indent_cm)
    if left_indent_cm is not None:
        paragraph.paragraph_format.left_indent = Cm(left_indent_cm)

    if format_tabs:
        section = doc.sections[0]
        width = Emu(section.page_width - section.left_margin - section.right_margin - Cm(left_indent_cm or 0))
        paragraph.paragraph_format.tab_stops.add_tab_stop(width, WD_TAB_ALIGNMENT.RIGHT)

    def add_run(t, is_bold):
        run = _style_run(paragraph.add_run(t), font_name, font_size, bold=is_bold, italic=italic)
        if highlight:
            run.font.highlight_color = WD_COLOR_INDEX.YELLOW
        return run

    if isinstance(bold, list):
        remaining = text
        while remaining:
            earliest_word = None
            earliest_pos = len(remaining)
            for word in bold:
                pos = remaining.find(word) if word else -1
                if pos != -1 and pos < earliest_pos:
                    earliest_pos = pos
                    earliest_word = word
            if earliest_word is None:
                add_run(remaining, False)
                break
            if earliest_pos > 0:
                add_run(remaining[:earliest_pos], False)
            add_run(earliest_word, True)
            remaining = remaining[earliest_pos + len(earliest_word):]
    elif text:
        add_run(text, bool(bold))

    _set_paragraph_mark_font(paragraph, font_name, font_size)
    return paragraph


def add_classification_header_footer(doc, marking, font_name=FONT_NAME, font_size=12):
    """Гриф у нижньому колонтитулі всіх сторінок, крім першої, і номер сторінки
    (поле PAGE) + гриф у верхньому - як в еталонному БН. Перша сторінка
    (титулка) має власний гриф у тексті, тож її колонтитули порожні."""
    section = doc.sections[0]
    section.different_first_page_header_footer = True
    section.first_page_header.is_linked_to_previous = False
    section.first_page_footer.is_linked_to_previous = False

    header_number = section.header.paragraphs[0]
    header_number.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _add_page_field(header_number, font_name, font_size)
    header_marking = section.header.add_paragraph()
    header_marking.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _style_run(header_marking.add_run(marking), font_name, font_size)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _style_run(footer.add_run(marking), font_name, font_size)


def _add_page_field(paragraph, font_name, font_size):
    def _new_run():
        return _style_run(paragraph.add_run(), font_name, font_size)

    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    _new_run()._r.append(fld_begin)

    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    _new_run()._r.append(instr_text)

    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    _new_run()._r.append(fld_end)


def add_table(doc, title, columns, rows, widths_cm, font_name=FONT_NAME, font_size=12):
    """Таблиця 'Table Grid' з необов'язковим рядком-заголовком (title на всю
    ширину, жирний, по центру), рядком назв колонок і рядками даних. Рядки,
    коротші за columns, доповнюються порожніми клітинками."""
    n = len(columns)
    title_rows = 1 if title else 0
    table = doc.add_table(rows=title_rows + 1 + len(rows), cols=n)
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = table.allow_autofit = False

    # Ширину визначає tblGrid (tbl.columns[i].width), а не лише cell.width -
    # без цього Word ділить таблицю на рівні колонки.
    for i, width in enumerate(widths_cm[:n]):
        table.columns[i].width = Cm(width)
        for cell in table.columns[i].cells:
            cell.width = Cm(width)

    def fill(cell, text, bold=False, center=False):
        for line_index, line in enumerate(str(text).split("\n")):
            paragraph = cell.paragraphs[0] if line_index == 0 else cell.add_paragraph()
            if center:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            if line:
                _style_run(paragraph.add_run(line), font_name, font_size, bold=bold)

    if title:
        merged = table.rows[0].cells[0].merge(table.rows[0].cells[n - 1]) if n > 1 else table.rows[0].cells[0]
        fill(merged, title, bold=True, center=True)
    for i, column in enumerate(columns):
        fill(table.rows[title_rows].cells[i], column, bold=True, center=True)
    for r, row in enumerate(rows, start=title_rows + 1):
        values = list(row) + [""] * (n - len(row))
        for c, value in enumerate(values[:n]):
            fill(table.rows[r].cells[c], value)
    return table
