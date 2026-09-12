import os, time, re, pandas as pd, shutil


from colorama import Fore, Style
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml.ns import qn
from docx.shared import Pt, Inches
from InquirerPy.prompts.checkbox import CheckboxPrompt
from InquirerPy.prompts.input import InputPrompt
from InquirerPy.prompts.number import NumberPrompt
from InquirerPy.prompts.list import ListPrompt
from datetime import datetime, timedelta
from constants import RANKS, VALUE_MASS_MOVEMENT, VALUE_NEW_POSITION, VALUE_RESIGNATION, VALUE_MILITARY_ASSAULT_COURSE, VALUE_SANITATION, VALUE_VACATION

# ================= ROUTES CONFIG =================
ROUTES = {
    VALUE_RESIGNATION: ['pos_mil', 'pos_commander', 'pos_higher_commander'],
    VALUE_NEW_POSITION: ['pos_mil', 'pos_commander', 'pos_higher_commander'],
    VALUE_MILITARY_ASSAULT_COURSE: ['pos_mil', 'pos_commander', 'pos_higher_commander'],
    VALUE_SANITATION: ['pos_mil', 'pos_higher_commander'],
    VALUE_VACATION: ['pos_mil', 'pos_commander', 'pos_higher_commander'],
}

# поля, які додаються окремо (не через ROUTES)
EXTRA_FIELDS = {
    VALUE_VACATION: ['part_vacation', 'selected_day', 'selected_month', 'selected_year', 'selected_days_for_vacation', 'region_for_vacation', 'district_for_vacation', 'settlement_for_vacation', 'street_for_vacation', 'house_number_for_vacation', 'my_number_phone', 'relative_for_military', 'relative_number_phone'],
}

# ================= FIELD HANDLERS =================
def int_prompt(message, required=True):
    return lambda: int_or_none(
        NumberPrompt(
            message=message,
            validate=lambda x: x.isdigit() if required else (x.isdigit() or x.strip() == ""),
            invalid_message="❌ Має бути число!" if required else "❌ Має бути число або пустим!"
        ).execute()
    )

FIELD_HANDLERS = {
    'pos_mil': int_prompt("Введіть номер посади військовослужбовця:"),
    'pos_commander': int_prompt("Введіть номер посади командира чи начальника (можна залишити порожнім):", required=False),
    'pos_higher_commander': int_prompt("Введіть номер посади командира батальйону:"),
    'part_vacation': lambda: ListPrompt(
        message="Оберіть частину відпустки:",
        choices=["першої", "другої", "третьої"]
    ).execute(),
    'selected_day': lambda:   InputPrompt(
        message="Введіть день початку відпустки (від 01 до 31):",
        validate=lambda x: x.isdigit() and 1 <= int(x) <= 31,
        invalid_message="❌ Має бути число від 1 до 31!"
    ).execute(),
   'selected_month': lambda: ListPrompt(
    message="Оберіть місяць початку відпустки:",
    choices=[*map(lambda i: [
        "січня","лютого","березня","квітня","травня","червня",
        "липня","серпня","вересня","жовтня","листопада","грудня"
    ][i % 12], range(datetime.now().month - 2, datetime.now().month + 1))]
).execute(),
    'selected_year': lambda: ListPrompt(
        message="Оберіть рік відпустки:",
        choices=get_years()
    ).execute(),
    'selected_days_for_vacation': lambda: ListPrompt(
        message="Оберіть кількість днів відпустки:",
        choices=['10 (десять)', '15 (п’ятнадцять)']
    ).execute(),
    'region_for_vacation': lambda: ListPrompt(
        message="Введіть область, де будете проводити відпустку:",
        choices=['Вінницька', 'Волинська', 'Дніпропетровська', 'Донецька', 'Житомирська', 'Закарпатська', 'Запорізька', 'Івано-Франківська', 'Київська', 'Кіровоградська', 'Луганська', 'Львівська', 'Миколаївська', 'Одеська', 'Полтавська', 'Рівненська', 'Сумська', 'Тернопільська', 'Харківська', 'Херсонська', 'Хмельницька', 'Черkaська', 'Чернівецька', 'Чернігівська']
    ).execute() + ' область',
    'district_for_vacation': lambda: InputPrompt(
        message="Введіть район, де будете проводити відпустку (Київський):",
        validate=lambda x: len(x.strip()) >= 3,
        invalid_message="❌ Вкажіть назву району!"
    ).execute() + ' район',
    'settlement_for_vacation': lambda: InputPrompt(
        message="Введіть населений пункт, де будете проводити відпустку (м. Київ):",
        validate=lambda x: len(x.strip()) >= 2 and re.search(r'\b(м\.|с\.|смт\.)\s?', x) is not None,
        invalid_message="❌ Вкажіть тип населеного пункту (м., с., смт.)"
    ).execute(),
    'street_for_vacation': lambda: 'вул. ' + InputPrompt(
        message="Введіть вулицю, де будете проводити відпустку (Хрещатик):",
        validate=lambda x: len(x.strip()) >= 3,
        invalid_message="❌ Вкажіть назву вулиці!"
    ).execute(),
    'house_number_for_vacation': lambda: 'буд. ' + InputPrompt(
        message="Введіть номер будинку, де будете проводити відпустку:",
        validate=lambda x: re.fullmatch(r'\d+([/\\\- ]\d+)?', x.strip()) is not None,
        invalid_message="❌ Вкажіть коректний номер (наприклад: 180, 180/1, 180-1)"
    ).execute(),
    'my_number_phone': lambda: InputPrompt(
        message="Введіть свій номер телефону (формат +380XXXXXXXXX):",
        validate=lambda x: re.match(r'^\+380\d{9}$', x) is not None,
        invalid_message="❌ Не вірний формат! Повинно бути +380XXXXXXXXX"
    ).execute(),
    'relative_for_military': lambda: ListPrompt(
        message="Оберіть близьку особу для резервного контакту з військовослужбовцем:",
        choices=['дружини', 'чоловіка', 'матері', 'батька', 'брата', 'сестри']
    ).execute(),
    'relative_number_phone': lambda: InputPrompt(
        message="Введіть номер телефону вибраної близької особи для резервного контакту з військовослужбовцем (формат +380XXXXXXXXX):",
        validate=lambda x: re.match(r'^\+380\d{9}$', x) is not None,
        invalid_message="❌ Не вірний формат! Повинно бути +380XXXXXXXXX"
    ).execute(),
}

# ================= MAIN =================

def get_value_data(rows_personel, rows_task):
    selected = CheckboxPrompt(
        message="Виберіть рапорт для генерації (Space → Enter):",
        choices=list(ROUTES.keys()) + [VALUE_MASS_MOVEMENT],
        validate=lambda x: len(x) > 0,
        invalid_message="❌ Оберіть хоча б одне!"
    ).execute()[0]

    if selected == VALUE_MASS_MOVEMENT:
        return selected, ()

    data = collect_input_data(selected)

    return selected, build_result(selected, rows_personel, rows_task, data)


# ================= COLLECT DATA =================

def collect_input_data(selected, data=None):
    data = data or {}

    for field in (*ROUTES.get(selected, []), *EXTRA_FIELDS.get(selected, [])):
        data[field] = FIELD_HANDLERS[field]()

    return data


# ================= RESULT BUILDER =================
def build_result(selected, rows_personel, rows_task, data):
    base = tuple(get_date_for_military(rows_personel, rows_task, data[pos]) for pos in ROUTES.get(selected, []))
    fields = EXTRA_FIELDS.get(selected, [])
    extra = ({f: data[f] for f in fields},)

    return (base + extra)


# ================= DATA PROCESSING =================
def get_date_for_military(rows_personel, rows_task, id_pos):
    personel_row = next((r for r in rows_personel if int_or_none(r['№']) == int_or_none(id_pos)), {})
    task_row = next((t for t in rows_task if int_or_none(t["№ посади"]) == int_or_none(id_pos)), {})

    rank_fact = personel_row.get('звання фактичне', '')
    ranks = RANKS.get(rank_fact, {})

    return {
        'position_id': personel_row.get('№', ''),
        'education': personel_row.get('ВОС', ''),
        'rank_state_nominative': personel_row.get('звання за штатом', ''),
        'rank_fact_nominative': rank_fact,
        'rank_fact_genitive': ranks.get('родовий', ''),
        'rank_fact_dative': ranks.get('давальний', ''),
        'position_full_nominative': task_row.get('посада називний', ''),
        'position_full_genitive': task_row.get('посада родовий', ''),
        'position_full_dative': task_row.get('посада давальний', ''),
        'position_short_dative': task_row.get('Посада давальний', ''),
        'name_nominative': task_row.get('ПІБ називний', ''),
        'name_genitive': task_row.get('ПІБ родовий', ''),
        'name_dative': task_row.get('ПІБ давальний', ''),
    }

def get_rows(transfer_data, columns_transfer, date_can_be_empty=False):
    if date_can_be_empty:
        return [
            {
                col: (date_to_str(val) if not pd.isna(val) else None)
                if isinstance(val := row[col], (pd.Timestamp, datetime)) else val
                for col in columns_transfer
            }
            for _, row in transfer_data[columns_transfer].iterrows()
        ]

    return [
        {
            col: date_to_str(row[col]) if isinstance(row[col], (pd.Timestamp, datetime)) else row[col]
            for col in columns_transfer
        }
        for _, row in transfer_data[columns_transfer].iterrows()
    ]

def date_to_str(date_str, action='-', days=0):
    return (to_date(date_str) + timedelta(days)).strftime('%d.%m.%Y') if action == '+' else (to_date(date_str) - timedelta(days)).strftime('%d.%m.%Y')

def to_date(x):
    if isinstance(x, datetime):
        return x.date()
    elif isinstance(x, str):
        for fmt in ("%Y-%m-%d", "%d.%m.%Y"):
            try:
                return datetime.strptime(x, fmt).date()
            except ValueError:
                continue
    raise ValueError(f"Невідомий формат або тип дати: {x}")

def read_excel_rows(file, sheet, letters):
    """Уніфікована функція для читання рядків з Excel."""
    df = pd.read_excel(file, sheet_name=sheet)
    col_indices = [excel_col_to_index(l) for l in letters]
    target_columns = [df.columns[i] for i in col_indices]
    return get_rows(df, target_columns)

def print_red(s):
    print(Fore.RED + s + Style.RESET_ALL)

def print_green(s):
    print(Fore.GREEN + s + Style.RESET_ALL)

def print_timeout(s, t):
    print('ROW:')
    time.sleep(t)
    print(s)

def set_line_spacing(doc, spacing):
    for paragraph in doc.paragraphs:
        paragraph.paragraph_format.line_spacing = Pt(spacing)
        
def create_or_clear_output_directory(directory):
    if os.path.exists(directory):
        shutil.rmtree(directory)
    os.makedirs(directory)

def add_paragraph_with_style(doc, text, font_name='Times New Roman', font_size=14, alignment=WD_ALIGN_PARAGRAPH.JUSTIFY, first_line_indent=0, space_after=True, format_tabs=False):
    paragraph = doc.add_paragraph()
    run = paragraph.add_run(text)
    run.font.name = font_name
    run._element.rPr.rFonts.set(qn('w:eastAsia'), font_name)
    run.font.size = Pt(font_size)

    paragraph.alignment = alignment
    paragraph.paragraph_format.line_spacing = 1.0  # Одинарний міжрядковий
    paragraph.paragraph_format.space_after = Pt(0)  # Взагалі нуль після абзацу

    if first_line_indent > 0:
        paragraph.paragraph_format.first_line_indent = Pt(first_line_indent)

    if format_tabs:
        width = Inches(doc.sections[0].page_width.inches - (doc.sections[0].left_margin.inches + doc.sections[0].right_margin.inches))
        tab_stops = paragraph.paragraph_format.tab_stops
        tab_stops.add_tab_stop(width, WD_TAB_ALIGNMENT.RIGHT)

    return paragraph

def excel_col_to_index(col):
    """Перетворює Excel-літеру (наприклад, 'A', 'Z', 'AA', 'AD') у індекс (0-відлік)."""
    col = col.upper()
    index = 0
    for c in col:
        index = index * 26 + (ord(c) - ord('A') + 1)
    return index - 1

def set_margins(doc, top, bottom, left, right):
    sections = doc.sections
    for section in sections:
        section.page_height = Inches(11.69)  # 297 мм
        section.page_width = Inches(8.27)    # 210 мм
        section.top_margin = top
        section.bottom_margin = bottom
        section.left_margin = left
        section.right_margin = right

def int_or_none(value):
    try:
        return int(value)
    except (ValueError, TypeError):
        return None


def convert_to_short_name(full_name, id_position=None):
    if not isinstance(full_name, str):
        return full_name

    # 1. Виправляємо апостроф + пробіл -> зліплюємо назад
    cleaned = re.sub(r"(В|в)'[\s]+([а-яА-ЯІіЇїЄєҐґ])", r"\1'\2", full_name.strip())

    # 2. Прибираємо зайві пробіли
    cleaned = ' '.join(cleaned.split())

    # 3. Розбити на слова
    words = cleaned.split()

    # 4. Якщо менше ніж 3 — пробуємо розклеїти друге слово (ім’я+по батькові)
    if len(words) == 2:
        # Пробуємо вставити пробіл між ім’ям і по батькові
        possible_fix = re.sub(r'([а-яґєіїʼ’]+)([А-ЯІЇЄҐ])', r'\1 \2', words[1])
        fixed_words = [words[0]] + possible_fix.split()
        if len(fixed_words) == 3:
            words = fixed_words

    # 5. Перевірка
    if len(words) != 3:
        raise ValueError(f"❌ Некоректний ПІБ для {id_position}: \"{full_name}\" → після обробки: \"{' '.join(words)}\"")

    # 6. Повертаємо скорочене ім’я
    return f"{words[1]} {words[0]}"

def is_empty(text):
    return not text or text.strip() == ""

def get_years():
    now = datetime.now()
    current = now.year

    # якщо до кінця року ≤ 1 місяць (грубо — грудень)
    if now.month == 12:
        years = [current - 1, current, current + 1]
    else:
        years = [current - 1, current]

    return [str(y) for y in years]