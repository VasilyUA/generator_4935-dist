import sys
from pathlib import Path

sys.path.append(str(Path(__file__).resolve().parent.parent))

import numpy as np
from docx import Document

import constants
from template import handed_position as m


def _soldier():
    return {
        "position_id": 5,
        "position_full_genitive": "Стрільця",
        "position_full_nominative": "Стрілець",
        "rank_fact_nominative": "сержант",
        "rank_fact_genitive": "сержанта",
        "name_nominative": "ЧЕТВЕРТИЙ Четвертий Четвертий",
        "name_genitive": "ЧЕТВЕРТИЙ Четвертий Четвертий",
    }


def _commander(position_id=7):
    return {
        "position_id": position_id,
        "position_short_dative": "Командиру роти",
        "position_full_nominative": "Командир роти",
        "rank_fact_nominative": "капітан",
        "rank_fact_genitive": "капітана",
        "name_nominative": "ШОСТИЙ Шостий Шостий",
        "name_genitive": "ШОСТИЙ Шостий Шостий",
    }


def _higher_commander():
    return {
        "position_full_nominative": "Командир батальйону",
        "rank_fact_nominative": "підполковник",
        "name_nominative": "СЬОМИЙ Сьомий Сьомий",
    }


def test_handed_position_report_with_commander(tmp_path, monkeypatch):
    monkeypatch.setattr(m, "OUTPUT_DIR", str(tmp_path))

    doc_name = m.handed_position_report(_soldier(), _commander(), _higher_commander())

    assert doc_name == str(tmp_path / "ЧЕТВЕРТИЙ Четвертий Четвертий здав посаду.docx")
    text = "\n".join(p.text for p in Document(doc_name).paragraphs)
    assert "здав" in text
    assert constants.FULL_MILITARY_UNIT in text
    assert "Шостий ШОСТИЙ" in text
    assert "Сьомий СЬОМИЙ" in text


def test_handed_position_report_without_commander_falls_back(tmp_path, monkeypatch):
    monkeypatch.setattr(m, "OUTPUT_DIR", str(tmp_path))
    commander = _commander(position_id=np.nan)

    doc_name = m.handed_position_report(_soldier(), commander, _higher_commander())

    text = "\n".join(p.text for p in Document(doc_name).paragraphs)
    assert "Командиру батальйону" in text
    assert "Шостий ШОСТИЙ" not in text
