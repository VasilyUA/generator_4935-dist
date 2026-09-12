import sys
from pathlib import Path

sys.path.append(str(Path(__file__).resolve().parent.parent))

from docx import Document

import constants
from template import change_position as m


def _soldier(name, position_id):
    return {
        "rank_fact_genitive": "сержанта",
        "name_genitive": name,
        "position_full_genitive": "стрільця",
        "education": "000000А/000",
        "rank_state": "сержант",
        "position_id": position_id,
    }


def _higher_commander():
    return {
        "position_full_nominative": "Командир батальйону",
        "rank_fact_nominative": "підполковник",
        "name_nominative": "ВІСІМНАДЦЯТИЙ Вісімнадцятий Вісімнадцятий",
    }


def test_change_position_report_single_row_writes_header_and_saves(tmp_path, monkeypatch):
    monkeypatch.setattr(m, "OUTPUT_DIR", str(tmp_path))
    doc = Document()

    result = m.change_position_report(
        doc, _soldier("ПЕРШИЙ Перший Перший", 5), _soldier("ДРУГИЙ Другий Другий", 6),
        _higher_commander(), index=1, len_row=1,
    )

    assert result == str(tmp_path / "РАПОРТ НА ПЕРЕМІЩЕННЯ.docx")
    text = "\n".join(p.text for p in Document(result).paragraphs)
    assert "ЗВІЛЬНИТИ ТА ПРИЗНАЧИТИ" in text
    assert constants.FULL_MILITARY_UNIT in text
    assert "ПЕРШИЙ Перший Перший" in text
    assert "(6)" in text  # position_id data_for_soldier_to - лише номер посади, не ім'я, потрапляє в текст
    assert "Вісімнадцятий ВІСІМНАДЦЯТИЙ" in text


def test_change_position_report_middle_row_does_not_save(tmp_path, monkeypatch):
    monkeypatch.setattr(m, "OUTPUT_DIR", str(tmp_path))
    doc = Document()

    result = m.change_position_report(
        doc, _soldier("ТРЕТІЙ Третій Третій", 5), _soldier("ЧЕТВЕРТИЙ Четвертий Четвертий", 6),
        _higher_commander(), index=1, len_row=2,
    )

    assert result is None
    assert list(tmp_path.iterdir()) == []


def test_change_position_report_header_only_on_first_row(tmp_path, monkeypatch):
    monkeypatch.setattr(m, "OUTPUT_DIR", str(tmp_path))
    doc = Document()

    m.change_position_report(
        doc, _soldier("ПЕРШИЙ Перший Перший", 5), _soldier("ДРУГИЙ Другий Другий", 6),
        _higher_commander(), index=1, len_row=2,
    )
    result = m.change_position_report(
        doc, _soldier("ТРЕТІЙ Третій Третій", 7), _soldier("ЧЕТВЕРТИЙ Четвертий Четвертий", 8),
        _higher_commander(), index=2, len_row=2,
    )

    text = "\n".join(p.text for p in Document(result).paragraphs)
    assert text.count("ЗВІЛЬНИТИ ТА ПРИЗНАЧИТИ") == 1
    assert "ПЕРШИЙ Перший Перший" in text
    assert "ТРЕТІЙ Третій Третій" in text
