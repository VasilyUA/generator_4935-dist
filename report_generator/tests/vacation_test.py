import sys
from pathlib import Path

sys.path.append(str(Path(__file__).resolve().parent.parent))

import numpy as np
from docx import Document

import constants
import template.vacation
# template/__init__.py робить "from .vacation import vacation", тож
# ОДНОЙМЕННИЙ атрибут пакета template.vacation переписується на функцію -
# дістаємо справжній модуль напряму із sys.modules, а не через атрибут пакета.
m = sys.modules["template.vacation"]


def _soldier():
    return {
        "position_full_nominative": "Стрілець",
        "rank_fact_nominative": "сержант",
        "rank_fact_genitive": "сержанта",
        "name_nominative": "ЧОТИРНАДЦЯТИЙ Чотирнадцятий Чотирнадцятий",
        "name_genitive": "ЧОТИРНАДЦЯТИЙ Чотирнадцятий Чотирнадцятий",
    }


def _commander(position_id=7):
    return {
        "position_id": position_id,
        "position_short_dative": "Командиру роти",
        "position_full_nominative": "Командир роти",
        "rank_fact_nominative": "капітан",
        "rank_fact_genitive": "капітана",
        "name_nominative": "ШІСТНАДЦЯТИЙ Шістнадцятий Шістнадцятий",
        "name_genitive": "ШІСТНАДЦЯТИЙ Шістнадцятий Шістнадцятий",
    }


def _higher_commander():
    return {
        "position_full_nominative": "Командир батальйону",
        "rank_fact_nominative": "підполковник",
        "name_nominative": "СІМНАДЦЯТИЙ Сімнадцятий Сімнадцятий",
    }


def _extra_fields():
    return {
        "part_vacation": "першої",
        "selected_day": "10",
        "selected_month": "серпня",
        "selected_year": "2026",
        "selected_days_for_vacation": "10 (десять)",
        "region_for_vacation": "Київська область",
        "district_for_vacation": "Київський район",
        "settlement_for_vacation": "м. Київ",
        "street_for_vacation": "вул. Хрещатик",
        "house_number_for_vacation": "буд. 1",
        "my_number_phone": "+380501234567",
        "relative_for_military": "дружини",
        "relative_number_phone": "+380671112233",
    }


def test_vacation_with_commander(tmp_path, monkeypatch):
    monkeypatch.setattr(m, "OUTPUT_DIR", str(tmp_path))

    doc_name = m.vacation(_soldier(), _commander(), _higher_commander(), _extra_fields())

    assert doc_name == str(tmp_path / "ЧОТИРНАДЦЯТИЙ Чотирнадцятий Чотирнадцятий здав посаду.docx")
    text = "\n".join(p.text for p in Document(doc_name).paragraphs)
    assert "щорічної основної відпустки" in text
    assert constants.FULL_MILITARY_UNIT in text
    assert "м. Київ" in text
    assert "Шістнадцятий ШІСТНАДЦЯТИЙ" in text


def test_vacation_without_commander_falls_back(tmp_path, monkeypatch):
    monkeypatch.setattr(m, "OUTPUT_DIR", str(tmp_path))
    commander = _commander(position_id=np.nan)

    doc_name = m.vacation(_soldier(), commander, _higher_commander(), _extra_fields())

    text = "\n".join(p.text for p in Document(doc_name).paragraphs)
    assert "Командиру батальйону" in text
    assert "ШІСТНАДЦЯТИЙ Шістнадцятий Шістнадцятий" not in text
