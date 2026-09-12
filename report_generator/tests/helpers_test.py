import sys
from datetime import datetime
from pathlib import Path

sys.path.append(str(Path(__file__).resolve().parent.parent))

import pandas as pd
import pytest
from colorama import Fore, Style
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from openpyxl import Workbook

import helpers
import constants
from conftest import FakePrompt, FakeCheckboxPrompt


def _write_xlsx(path, sheet, headers, rows):
    wb = Workbook()
    ws = wb.active
    ws.title = sheet
    ws.append(headers)
    for row in rows:
        ws.append(row)
    wb.save(path)
    return str(path)


# -------------------------
# int_or_none
# -------------------------
def test_int_or_none_valid_string():
    assert helpers.int_or_none("5") == 5


def test_int_or_none_invalid_string_returns_none():
    assert helpers.int_or_none("не число") is None


def test_int_or_none_none_input_returns_none():
    assert helpers.int_or_none(None) is None


# -------------------------
# excel_col_to_index
# -------------------------
def test_excel_col_to_index_single_letter():
    assert helpers.excel_col_to_index("A") == 0
    assert helpers.excel_col_to_index("Z") == 25


def test_excel_col_to_index_double_letter():
    assert helpers.excel_col_to_index("AA") == 26


def test_excel_col_to_index_lowercase():
    assert helpers.excel_col_to_index("a") == 0


# -------------------------
# to_date
# -------------------------
def test_to_date_datetime_passthrough():
    dt = datetime(2026, 8, 19)
    assert helpers.to_date(dt) == dt.date()


def test_to_date_iso_string():
    assert helpers.to_date("2026-08-19") == datetime(2026, 8, 19).date()


def test_to_date_dotted_string():
    assert helpers.to_date("19.08.2026") == datetime(2026, 8, 19).date()


def test_to_date_unparseable_raises():
    with pytest.raises(ValueError):
        helpers.to_date("не дата")


def test_to_date_unsupported_type_raises():
    with pytest.raises(ValueError):
        helpers.to_date(123)


# -------------------------
# date_to_str
# -------------------------
def test_date_to_str_subtracts_by_default():
    assert helpers.date_to_str("19.08.2026") == "19.08.2026"
    assert helpers.date_to_str("19.08.2026", days=1) == "18.08.2026"


def test_date_to_str_adds_when_action_plus():
    assert helpers.date_to_str("19.08.2026", action="+", days=1) == "20.08.2026"


# -------------------------
# is_empty
# -------------------------
def test_is_empty_none():
    assert helpers.is_empty(None) is True


def test_is_empty_blank_string():
    assert helpers.is_empty("   ") is True


def test_is_empty_non_blank_string():
    assert helpers.is_empty("текст") is False


# -------------------------
# get_years
# -------------------------
class _FixedDatetime(datetime):
    _fixed = datetime(2026, 8, 19)

    @classmethod
    def now(cls):
        return cls._fixed


def test_get_years_not_december_returns_two_years(monkeypatch):
    _FixedDatetime._fixed = datetime(2026, 8, 19)
    monkeypatch.setattr(helpers, "datetime", _FixedDatetime)
    assert helpers.get_years() == ["2025", "2026"]


def test_get_years_december_returns_three_years(monkeypatch):
    _FixedDatetime._fixed = datetime(2026, 12, 15)
    monkeypatch.setattr(helpers, "datetime", _FixedDatetime)
    assert helpers.get_years() == ["2025", "2026", "2027"]


# -------------------------
# convert_to_short_name
# -------------------------
def test_convert_to_short_name_three_words():
    assert helpers.convert_to_short_name("ПЕРШИЙ Перший Перший") == "Перший ПЕРШИЙ"


def test_convert_to_short_name_apostrophe_glued_back():
    # "В' ікторович" (апостроф + пробіл, типова помилка Excel-копіювання) ->
    # "В'ікторович" ДО розбиття на слова
    assert helpers.convert_to_short_name("ДРУГИЙ Другий В' ікторович") == "Другий ДРУГИЙ"


def test_convert_to_short_name_collapses_extra_spaces():
    assert helpers.convert_to_short_name("ТРЕТІЙ   Третій  Третій") == "Третій ТРЕТІЙ"


def test_convert_to_short_name_splices_two_glued_words():
    # ім'я+по-батькові зліплені без пробілу ("ЧетвертийЧетвертий") -
    # розклеюються за межею мала->велика літера
    assert helpers.convert_to_short_name("ЧЕТВЕРТИЙ ЧетвертийЧетвертий") == "Четвертий ЧЕТВЕРТИЙ"


def test_convert_to_short_name_wrong_word_count_raises():
    with pytest.raises(ValueError, match="Некоректний ПІБ"):
        helpers.convert_to_short_name("ОдинСлово", id_position=42)


def test_convert_to_short_name_non_string_passthrough():
    assert helpers.convert_to_short_name(None) is None
    assert helpers.convert_to_short_name(123) == 123


# -------------------------
# get_date_for_military
# -------------------------
def _personel_row(pos_id, rank_fact="сержант"):
    return {
        "№": pos_id,
        "ВОС": "000000А/000",
        "звання за штатом": rank_fact,
        "звання фактичне": rank_fact,
    }


def _task_row(pos_id, pib="ШОСТИЙ Шостий Шостий"):
    return {
        "№ посади": pos_id,
        "посада називний": "Стрілець",
        "посада родовий": "стрільця",
        "посада давальний": "стрільцю",
        "Посада давальний": "Стрільцю",
        "ПІБ називний": pib,
        "ПІБ родовий": f"{pib} родовий",
        "ПІБ давальний": f"{pib} давальний",
    }


def test_get_date_for_military_matches_by_position_id():
    rows_personel = [_personel_row(1)]
    rows_task = [_task_row(1)]

    result = helpers.get_date_for_military(rows_personel, rows_task, "1")

    assert result["position_id"] == 1
    assert result["rank_fact_nominative"] == "сержант"
    assert result["rank_fact_genitive"] == "сержанта"
    assert result["rank_fact_dative"] == "сержанту"
    assert result["name_nominative"] == "ШОСТИЙ Шостий Шостий"


def test_get_date_for_military_no_match_returns_blank_defaults():
    result = helpers.get_date_for_military([], [], "999")

    assert result["position_id"] == ""
    assert result["rank_fact_genitive"] == ""
    assert result["name_nominative"] == ""


def test_get_date_for_military_unknown_rank_blank_declensions():
    rows_personel = [_personel_row(1, rank_fact="невідоме звання")]
    rows_task = [_task_row(1)]

    result = helpers.get_date_for_military(rows_personel, rows_task, 1)

    assert result["rank_fact_nominative"] == "невідоме звання"
    assert result["rank_fact_genitive"] == ""


# -------------------------
# get_rows
# -------------------------
def test_get_rows_converts_timestamp_columns():
    df = pd.DataFrame({"ПІБ": ["СЬОМИЙ Сьомий Сьомий"], "Дата": [pd.Timestamp("2026-08-19")]})
    result = helpers.get_rows(df, ["ПІБ", "Дата"])
    assert result == [{"ПІБ": "СЬОМИЙ Сьомий Сьомий", "Дата": "19.08.2026"}]


def test_get_rows_date_can_be_empty_converts_nat_to_none():
    df = pd.DataFrame({"ПІБ": ["ВОСЬМИЙ Восьмий Восьмий"], "Дата": [pd.NaT]})
    result = helpers.get_rows(df, ["ПІБ", "Дата"], date_can_be_empty=True)
    assert result == [{"ПІБ": "ВОСЬМИЙ Восьмий Восьмий", "Дата": None}]


def test_get_rows_date_can_be_empty_still_converts_real_dates():
    df = pd.DataFrame({"Дата": [pd.Timestamp("2026-08-19")]})
    result = helpers.get_rows(df, ["Дата"], date_can_be_empty=True)
    assert result == [{"Дата": "19.08.2026"}]


# -------------------------
# read_excel_rows
# -------------------------
def test_read_excel_rows_reads_by_column_letters(tmp_path):
    path = _write_xlsx(
        tmp_path / "СПИСОК.xlsm", "Аркуш1",
        ["A", "B", "ПІБ", "D"],
        [[1, 2, "ДВАДЦЯТИЙ Двадцятий Двадцятий", 4]],
    )
    result = helpers.read_excel_rows(path, "Аркуш1", ["C"])
    assert result == [{"ПІБ": "ДВАДЦЯТИЙ Двадцятий Двадцятий"}]


# -------------------------
# print_red / print_green / print_timeout
# -------------------------
def test_print_red_output(capsys):
    helpers.print_red("Помилка")
    out = capsys.readouterr().out
    assert out == Fore.RED + "Помилка" + Style.RESET_ALL + "\n"


def test_print_green_output(capsys):
    helpers.print_green("Готово")
    out = capsys.readouterr().out
    assert out == Fore.GREEN + "Готово" + Style.RESET_ALL + "\n"


def test_print_timeout_prints_row_then_message(capsys, monkeypatch):
    monkeypatch.setattr(helpers.time, "sleep", lambda s: None)
    helpers.print_timeout("текст", 0)
    out = capsys.readouterr().out
    assert "ROW:" in out
    assert "текст" in out


# -------------------------
# create_or_clear_output_directory
# -------------------------
def test_create_or_clear_output_directory_clears_existing(tmp_path):
    out = tmp_path / "output"
    out.mkdir()
    (out / "old.docx").write_text("stale")

    helpers.create_or_clear_output_directory(str(out))

    assert out.exists()
    assert list(out.iterdir()) == []


def test_create_or_clear_output_directory_creates_missing(tmp_path):
    out = tmp_path / "does_not_exist_yet"
    helpers.create_or_clear_output_directory(str(out))
    assert out.exists() and out.is_dir()


# -------------------------
# set_margins / set_line_spacing / add_paragraph_with_style
# -------------------------
def test_set_margins_sets_page_size_and_margins():
    from docx.shared import Inches

    doc = Document()
    helpers.set_margins(doc, Inches(0.5), Inches(0.5), Inches(1), Inches(0.5))
    section = doc.sections[0]
    assert round(section.page_width.inches, 2) == 8.27
    assert round(section.page_height.inches, 2) == 11.69
    assert section.left_margin == Inches(1)


def test_set_line_spacing_applies_to_all_paragraphs():
    from docx.shared import Pt

    doc = Document()
    doc.add_paragraph("перший")
    doc.add_paragraph("другий")
    helpers.set_line_spacing(doc, 14)
    for paragraph in doc.paragraphs:
        assert paragraph.paragraph_format.line_spacing == Pt(14)


def test_add_paragraph_with_style_basic_text():
    doc = Document()
    paragraph = helpers.add_paragraph_with_style(doc, "текст", first_line_indent=34)
    assert paragraph.text == "текст"
    assert paragraph.alignment == WD_ALIGN_PARAGRAPH.JUSTIFY
    assert paragraph.paragraph_format.first_line_indent is not None


def test_add_paragraph_with_style_format_tabs_adds_tab_stop():
    doc = Document()
    paragraph = helpers.add_paragraph_with_style(doc, "текст\tзначення", format_tabs=True)
    assert len(paragraph.paragraph_format.tab_stops) == 1


# -------------------------
# collect_input_data / build_result / get_value_data (через FakePrompt)
# -------------------------
def test_collect_input_data_calls_handler_for_each_route_field(inquirer_inputs):
    inquirer_inputs.append("5")
    inquirer_inputs.append("")
    inquirer_inputs.append("7")

    data = helpers.collect_input_data(constants.VALUE_RESIGNATION)

    assert set(data.keys()) == {"pos_mil", "pos_commander", "pos_higher_commander"}
    assert data["pos_mil"] == 5


def test_collect_input_data_includes_extra_fields_for_vacation(inquirer_inputs):
    for _ in range(3 + len(helpers.EXTRA_FIELDS[constants.VALUE_VACATION])):
        inquirer_inputs.append("значення")

    data = helpers.collect_input_data(constants.VALUE_VACATION)

    for field in helpers.EXTRA_FIELDS[constants.VALUE_VACATION]:
        assert field in data


def test_build_result_shapes_base_and_extra_tuple():
    data = {"pos_mil": 1, "pos_commander": None, "pos_higher_commander": 3}
    rows_personel = [_personel_row(1), _personel_row(3)]
    rows_task = [_task_row(1), _task_row(3)]

    result = helpers.build_result(constants.VALUE_RESIGNATION, rows_personel, rows_task, data)

    assert len(result) == 4  # pos_mil + pos_commander + pos_higher_commander + extra dict
    assert result[0]["position_id"] == 1
    assert result[-1] == {}


def test_get_value_data_mass_movement_returns_empty_args(inquirer_inputs):
    inquirer_inputs.append([constants.VALUE_MASS_MOVEMENT])

    selected, args = helpers.get_value_data([], [])

    assert selected == constants.VALUE_MASS_MOVEMENT
    assert args == ()


def test_get_value_data_resignation_builds_full_result(inquirer_inputs):
    rows_personel = [_personel_row(1), _personel_row(2), _personel_row(3)]
    rows_task = [_task_row(1), _task_row(2), _task_row(3)]

    inquirer_inputs.append([constants.VALUE_RESIGNATION])
    inquirer_inputs.extend(["1", "2", "3"])

    selected, args = helpers.get_value_data(rows_personel, rows_task)

    assert selected == constants.VALUE_RESIGNATION
    assert len(args) == 4


# -------------------------
# FIELD_HANDLERS - валідаційна логіка полів відпустки
# -------------------------
def test_field_handlers_my_number_phone_validates_ukrainian_format(inquirer_inputs):
    inquirer_inputs.append("+380501234567")
    assert helpers.FIELD_HANDLERS["my_number_phone"]() == "+380501234567"
    validate = FakePrompt.last_kwargs["validate"]
    assert validate("+380501234567") is True
    assert validate("0501234567") is False


def test_field_handlers_house_number_accepts_slash_and_dash(inquirer_inputs):
    inquirer_inputs.append("180/1")
    assert helpers.FIELD_HANDLERS["house_number_for_vacation"]() == "буд. 180/1"
    validate = FakePrompt.last_kwargs["validate"]
    assert validate("180-1") is True
    assert validate("не число") is False


def test_field_handlers_region_appends_suffix(inquirer_inputs):
    inquirer_inputs.append("Київська")
    assert helpers.FIELD_HANDLERS["region_for_vacation"]() == "Київська область"


def test_field_handlers_district_appends_suffix_and_validates_length(inquirer_inputs):
    inquirer_inputs.append("Київський")
    assert helpers.FIELD_HANDLERS["district_for_vacation"]() == "Київський район"
    validate = FakePrompt.last_kwargs["validate"]
    assert validate("ки") is False


def test_field_handlers_settlement_requires_type_prefix(inquirer_inputs):
    inquirer_inputs.append("м. Київ")
    assert helpers.FIELD_HANDLERS["settlement_for_vacation"]() == "м. Київ"
    validate = FakePrompt.last_kwargs["validate"]
    assert validate("Київ") is False


def test_field_handlers_street_prepends_prefix(inquirer_inputs):
    inquirer_inputs.append("Хрещатик")
    assert helpers.FIELD_HANDLERS["street_for_vacation"]() == "вул. Хрещатик"


def test_field_handlers_relative_number_phone_same_validation_as_own(inquirer_inputs):
    inquirer_inputs.append("+380671112233")
    assert helpers.FIELD_HANDLERS["relative_number_phone"]() == "+380671112233"


def test_field_handlers_part_vacation_selects_from_list(inquirer_inputs):
    inquirer_inputs.append("першої")
    assert helpers.FIELD_HANDLERS["part_vacation"]() == "першої"


def test_field_handlers_selected_day_validates_range(inquirer_inputs):
    inquirer_inputs.append("15")
    assert helpers.FIELD_HANDLERS["selected_day"]() == "15"
    validate = FakePrompt.last_kwargs["validate"]
    assert validate("32") is False
    assert validate("0") is False


def test_field_handlers_selected_month_returns_a_choice(inquirer_inputs):
    inquirer_inputs.append("серпня")
    assert helpers.FIELD_HANDLERS["selected_month"]() == "серпня"


def test_field_handlers_selected_year_uses_get_years(inquirer_inputs):
    inquirer_inputs.append("2026")
    assert helpers.FIELD_HANDLERS["selected_year"]() == "2026"


def test_field_handlers_selected_days_for_vacation_choice(inquirer_inputs):
    inquirer_inputs.append("10 (десять)")
    assert helpers.FIELD_HANDLERS["selected_days_for_vacation"]() == "10 (десять)"


def test_int_prompt_required_true_validates_digits_only():
    prompt_fn = helpers.int_prompt("Введіть номер:")
    _BaseFakePromptQueueEmpty = FakePrompt._queue
    FakePrompt._queue = ["5"]
    try:
        assert prompt_fn() == 5
    finally:
        FakePrompt._queue = _BaseFakePromptQueueEmpty
    validate = FakePrompt.last_kwargs["validate"]
    assert validate("5") is True
    assert validate("абв") is False


def test_int_prompt_not_required_allows_blank():
    prompt_fn = helpers.int_prompt("Введіть номер:", required=False)
    FakePrompt._queue = [""]
    try:
        assert prompt_fn() is None
    finally:
        FakePrompt._queue = []
    validate = FakePrompt.last_kwargs["validate"]
    assert validate("") is True
    assert validate("абв") is False
