import sys
from pathlib import Path

sys.path.append(str(Path(__file__).resolve().parent.parent))

import numpy as np
from docx import Document

import constants
from template import military_assault_сourse as m


def _soldier():
    return {
        "position_id": 5,
        "position_full_genitive": "Стрільця",
        "rank_fact_nominative": "сержант",
        "rank_fact_genitive": "сержанта",
        "name_nominative": "ВОСЬМИЙ Восьмий Восьмий",
        "name_genitive": "ВОСЬМИЙ Восьмий Восьмий",
    }


def _commander(position_id=7):
    return {
        "position_id": position_id,
        "position_short_dative": "Командиру роти",
        "position_full_nominative": "Командир роти",
        "rank_fact_nominative": "капітан",
        "rank_fact_genitive": "капітана",
        "name_nominative": "ДЕСЯТИЙ Десятий Десятий",
        "name_genitive": "ДЕСЯТИЙ Десятий Десятий",
    }


def _higher_commander():
    return {
        "position_full_nominative": "Командир батальйону",
        "rank_fact_nominative": "підполковник",
        "name_nominative": "ОДИНАДЦЯТИЙ Одинадцятий Одинадцятий",
    }


def test_military_assault_course_with_commander(tmp_path, monkeypatch):
    monkeypatch.setattr(m, "OUTPUT_DIR", str(tmp_path))

    doc_name = m.military_assault_course(_soldier(), _commander(), _higher_commander())

    assert doc_name == str(tmp_path / "ВОСЬМИЙ Восьмий Восьмий смуга піхотинця.docx")
    text = "\n".join(p.text for p in Document(doc_name).paragraphs)
    assert "смуги морського піхотинця" in text
    assert constants.FULL_MILITARY_UNIT in text
    assert "Десятий ДЕСЯТИЙ" in text


def test_military_assault_course_without_commander_falls_back(tmp_path, monkeypatch):
    monkeypatch.setattr(m, "OUTPUT_DIR", str(tmp_path))
    commander = _commander(position_id=np.nan)

    doc_name = m.military_assault_course(_soldier(), commander, _higher_commander())

    text = "\n".join(p.text for p in Document(doc_name).paragraphs)
    assert "Командиру батальйону" in text
    assert "Десятий ДЕСЯТИЙ" not in text
