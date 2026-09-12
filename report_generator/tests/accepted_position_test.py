import sys
from pathlib import Path

sys.path.append(str(Path(__file__).resolve().parent.parent))

import numpy as np

import constants
from template import accepted_position as m


def _soldier():
    return {
        "position_id": 5,
        "position_full_genitive": "Стрільця",
        "position_full_nominative": "Стрілець",
        "rank_fact_nominative": "сержант",
        "rank_fact_genitive": "сержанта",
        "name_nominative": "ПЕРШИЙ Перший Перший",
        "name_genitive": "ПЕРШИЙ Перший Перший",
    }


def _commander(position_id=7):
    return {
        "position_id": position_id,
        "position_short_dative": "Командиру роти",
        "position_full_nominative": "Командир роти",
        "rank_fact_nominative": "капітан",
        "rank_fact_genitive": "капітана",
        "name_nominative": "ДРУГИЙ Другий Другий",
        "name_genitive": "ДРУГИЙ Другий Другий",
    }


def _higher_commander():
    return {
        "position_full_nominative": "Командир батальйону",
        "rank_fact_nominative": "підполковник",
        "name_nominative": "ТРЕТІЙ Третій Третій",
    }


def test_generate_accepted_position_report_with_commander(tmp_path, monkeypatch):
    monkeypatch.setattr(m, "OUTPUT_DIR", str(tmp_path))

    doc_name = m.generate_accepted_position_report(_soldier(), _commander(), _higher_commander())

    assert doc_name == str(tmp_path / "ПЕРШИЙ Перший Перший прийняв посаду.docx")
    assert Path(doc_name).exists()

    from docx import Document
    text = "\n".join(p.text for p in Document(doc_name).paragraphs)
    assert "прийняв" in text
    assert constants.FULL_MILITARY_UNIT in text
    assert "Другий ДРУГИЙ" in text
    assert "Третій ТРЕТІЙ" in text


def test_generate_accepted_position_report_without_commander_falls_back(tmp_path, monkeypatch):
    monkeypatch.setattr(m, "OUTPUT_DIR", str(tmp_path))
    commander = _commander(position_id=np.nan)

    doc_name = m.generate_accepted_position_report(_soldier(), commander, _higher_commander())

    from docx import Document
    text = "\n".join(p.text for p in Document(doc_name).paragraphs)
    assert "Командиру батальйону" in text
    assert "Другий ДРУГИЙ" not in text
