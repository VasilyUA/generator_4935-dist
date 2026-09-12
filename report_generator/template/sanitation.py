import pandas as pd, os

from datetime import datetime
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches
from helpers import convert_to_short_name, set_margins, set_line_spacing, add_paragraph_with_style, is_empty, print_green
from constants import OUTPUT_DIR, FULL_MILITARY_UNIT

def sanitation(data_for_soldier, data_for_higher_commander):
    doc = Document()
    set_margins(doc, Inches(0.786), Inches(0.786), Inches(1.18), Inches(0.395))
    set_line_spacing(doc, 14)  # встановлення міжрядкового інтервалу
    add_paragraph_with_style(doc, f"Командиру {FULL_MILITARY_UNIT}", first_line_indent=250)
    add_paragraph_with_style(doc, "")
    add_paragraph_with_style(doc, "")
    add_paragraph_with_style(doc, "")
    add_paragraph_with_style(doc, "")
    add_paragraph_with_style(doc, "")
    add_paragraph_with_style(doc, "")
    add_paragraph_with_style(doc, "РАПОРТ", alignment=WD_ALIGN_PARAGRAPH.CENTER)
    add_paragraph_with_style(doc, "")
    add_paragraph_with_style(doc, "")
    add_paragraph_with_style(doc, f"Прошу Вас виплатити мені {data_for_soldier.get('position_full_dative', '').lower()} {FULL_MILITARY_UNIT} {data_for_soldier.get('rank_fact_dative', '')} {data_for_soldier.get('name_dative', '')} грошову допомогу на оздоровлення за {datetime.now().year} рік, без надання відпустки, згідно розділу ХХІІІ Порядку виплати грошового забезпечення військовослужбовцям Збройних Сил України та деяким іншим особам, Затвердженого Наказом Міноборони України від 07.06.2018 №260 від «07» червня 2018 року «Про затвердження Порядку виплати грошового забезпечення військовослужбовцям Збройних Сил України та деяким іншим особам».", first_line_indent=34)
    add_paragraph_with_style(doc, "")
    add_paragraph_with_style(doc, "")
    add_paragraph_with_style(doc, f"{data_for_higher_commander.get('position_full_nominative', '')} {FULL_MILITARY_UNIT}", space_after=False)
    add_paragraph_with_style(doc, f"{data_for_higher_commander.get('rank_fact_nominative', '')}\t{convert_to_short_name(data_for_higher_commander.get('name_nominative', ''))}", format_tabs=True)
    doc_name = os.path.join(OUTPUT_DIR, f"{data_for_soldier.get('name_nominative')} на оздоровчі.docx")
    doc.save(doc_name)
    print_green(f"- {doc_name}")
    return doc_name