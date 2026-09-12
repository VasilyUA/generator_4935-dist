import os
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches
from helpers import set_margins, set_line_spacing, add_paragraph_with_style, convert_to_short_name, print_green
from constants import OUTPUT_DIR, FULL_MILITARY_UNIT

def change_position_report(doc, data_for_soldier_from, data_for_soldier_to, data_for_higher_commander, index, len_row):
        if index == 1:
            set_margins(doc, Inches(0.786), Inches(0.786), Inches(1.18), Inches(0.395))
            set_line_spacing(doc, 1.0)  # встановлення міжрядкового інтервалу 1
            add_paragraph_with_style(doc, f"Командиру {FULL_MILITARY_UNIT}", alignment=WD_ALIGN_PARAGRAPH.RIGHT)
            add_paragraph_with_style(doc, "")
            add_paragraph_with_style(doc, "")
            add_paragraph_with_style(doc, "")
            add_paragraph_with_style(doc, "РАПОРТ", alignment=WD_ALIGN_PARAGRAPH.CENTER)
            add_paragraph_with_style(doc, "")
            add_paragraph_with_style(doc, "")
            add_paragraph_with_style(doc, f"Прошу Вашого дозволу для підвищення бойової готовності 1 батальйону морської піхоти {FULL_MILITARY_UNIT} нижче поіменований особовий склад від займаних посад ЗВІЛЬНИТИ ТА ПРИЗНАЧИТИ:", first_line_indent=34)
            add_paragraph_with_style(doc, "")
            add_paragraph_with_style(doc, "")

        person = f"{data_for_soldier_from.get('rank_fact_genitive', '')} {data_for_soldier_from.get('name_genitive', '')},"
        from_position = f"{data_for_soldier_from.get('position_full_genitive', '').lower()} {FULL_MILITARY_UNIT}, ВОС - {data_for_soldier_from.get('education', '')}, ШПС \"{data_for_soldier_from.get('rank_state', '')}\" ({data_for_soldier_from.get('position_id', '')})"
        to_position = f"{data_for_soldier_to.get('position_full_genitive', '').lower()} {FULL_MILITARY_UNIT}, ВОС - {data_for_soldier_to.get('education', '')}, ШПС \"{data_for_soldier_to.get('rank_state', '')}\" ({data_for_soldier_to.get('position_id', '')})"
        add_paragraph_with_style(doc, f"{person} {from_position} - {to_position};", first_line_indent=34)

        print(len_row, index)

        if  len_row == index:
            add_paragraph_with_style(doc, "")
            add_paragraph_with_style(doc, "")
            add_paragraph_with_style(doc, f"{data_for_higher_commander.get('position_full_nominative', '')} {FULL_MILITARY_UNIT}", space_after=False)
            add_paragraph_with_style(doc, f"{data_for_higher_commander.get('rank_fact_nominative', '')}\t{convert_to_short_name(data_for_higher_commander.get('name_nominative', ''))}", format_tabs=True)

            doc_name = os.path.join(OUTPUT_DIR, f"РАПОРТ НА ПЕРЕМІЩЕННЯ.docx")
            doc.save(doc_name)

            print_green(f"- {doc_name}")
            return doc_name