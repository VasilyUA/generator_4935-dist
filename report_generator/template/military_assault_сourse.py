import pandas as pd, os

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches
from helpers import convert_to_short_name, set_margins, set_line_spacing, add_paragraph_with_style, is_empty, print_green
from constants import OUTPUT_DIR, FULL_MILITARY_UNIT

def military_assault_course(data_for_soldier, data_for_commander, data_for_higher_commander):
    doc = Document()
    set_margins(doc, Inches(0.786), Inches(0.786), Inches(1.18), Inches(0.395))
    set_line_spacing(doc, 14)  # встановлення міжрядкового інтервалу
    if pd.notna(data_for_commander.get('position_id')) and not is_empty(f"{data_for_commander.get('position_id')}"):
        add_paragraph_with_style(doc, data_for_commander.get('position_short_dative'), first_line_indent=250)
    else:
        add_paragraph_with_style(doc, "Командиру батальйону", first_line_indent=250)
    add_paragraph_with_style(doc, "")
    add_paragraph_with_style(doc, "")
    add_paragraph_with_style(doc, "")
    add_paragraph_with_style(doc, "РАПОРТ", alignment=WD_ALIGN_PARAGRAPH.CENTER)
    add_paragraph_with_style(doc, "")
    add_paragraph_with_style(doc, "")
    add_paragraph_with_style(doc, f"Прошу допустити мене {data_for_soldier.get('position_full_genitive', '').lower()} {FULL_MILITARY_UNIT} до проходження смуги морського піхотинця.", first_line_indent=34)
    add_paragraph_with_style(doc, f"{data_for_soldier.get('rank_fact_nominative', '')}\t{convert_to_short_name(data_for_soldier.get('name_nominative', ''))}", format_tabs=True)
    add_paragraph_with_style(doc, "")
    add_paragraph_with_style(doc, "")
    if pd.notna(data_for_commander.get('position_id')) and not is_empty(f"{data_for_commander.get('position_id')}"):
        add_paragraph_with_style(doc, "Командиру батальйону", first_line_indent=250)
        add_paragraph_with_style(doc, "")
        add_paragraph_with_style(doc, f"Доповідаю по суті рапорту {data_for_soldier.get('rank_fact_genitive', '')} {convert_to_short_name(data_for_soldier.get('name_genitive', ''))}", first_line_indent=34)
        add_paragraph_with_style(doc, "")
        add_paragraph_with_style(doc, f"{data_for_commander.get('position_full_nominative', '')} {FULL_MILITARY_UNIT}", space_after=False)
        add_paragraph_with_style(doc, f"{data_for_commander.get('rank_fact_nominative', '')}\t{convert_to_short_name(data_for_commander.get('name_nominative', ''))}", format_tabs=True)
        add_paragraph_with_style(doc, "")
        add_paragraph_with_style(doc, "")
        add_paragraph_with_style(doc, f"Командиру {FULL_MILITARY_UNIT}", first_line_indent=250)
        add_paragraph_with_style(doc, "")
        add_paragraph_with_style(doc, f"Доповідаю по суті рапорту {data_for_commander.get('rank_fact_genitive', '')} {convert_to_short_name(data_for_commander.get('name_genitive', ''))}", first_line_indent=34)
    else:
        add_paragraph_with_style(doc, f"Командиру {FULL_MILITARY_UNIT}", first_line_indent=250)
        add_paragraph_with_style(doc, "")
        add_paragraph_with_style(doc, f"Доповідаю по суті рапорту {data_for_soldier.get('rank_fact_genitive', '')} {convert_to_short_name(data_for_soldier.get('name_genitive', ''))}", first_line_indent=34)

    add_paragraph_with_style(doc, "")
    add_paragraph_with_style(doc, "")
    add_paragraph_with_style(doc, f"{data_for_higher_commander.get('position_full_nominative', '')} {FULL_MILITARY_UNIT}", space_after=False)
    add_paragraph_with_style(doc, f"{data_for_higher_commander.get('rank_fact_nominative', '')}\t{convert_to_short_name(data_for_higher_commander.get('name_nominative', ''))}", format_tabs=True)
    doc_name = os.path.join(OUTPUT_DIR, f"{data_for_soldier.get('name_nominative')} смуга піхотинця.docx")
    doc.save(doc_name)
    print_green(f"- {doc_name}")
    return doc_name
