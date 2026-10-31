"""Спільні операції з книгами Excel (openpyxl): відкриття, створення, збереження.

load_workbook стійкий до текстових значень користувацьких фільтрів. openpyxl
3.1.2 приймає в <customFilter val="..."> автофільтра лише число чи шаблон із
"*", а на будь-якому іншому тексті відмовляється відкривати книгу взагалі.
Саме так Excel зберігає звичайний фільтр "(Непусті)" (notEqual " ") чи
"дорівнює/не дорівнює <текст>". Реальний випадок: фільтр на колонці аркуша
"Табель" resources/ОБЛІК.xlsx - ValueError "Value must be either numerical or
a string containing a wildcard" ще до читання жодної клітинки. Такі значення
лишаються текстом, фільтр зберігається у файлі як був."""

import os

from openpyxl import Workbook, load_workbook
# Стаби openpyxl не описують дескриптор CustomFilterValueDescriptor, хоча
# він є в самій бібліотеці.
from openpyxl.worksheet.filters import CustomFilter, CustomFilterValueDescriptor  # pyright: ignore[reportAttributeAccessIssue]
from openpyxl.worksheet.worksheet import Worksheet

__all__ = ["empty_workbook", "load_workbook", "save_workbook", "workbook_with_sheet"]


class _TextTolerantFilterValue(CustomFilterValueDescriptor):
    def __set__(self, instance, value):
        if isinstance(value, str) and not self.pattern.match(value):
            instance.__dict__[self.name] = value
            return
        super().__set__(instance, value)


CustomFilter.val = _TextTolerantFilterValue(name="val")  # pyright: ignore[reportAttributeAccessIssue]


def workbook_with_sheet(title: str) -> tuple[Workbook, Worksheet]:
    """Нова книга з одним аркушем title."""
    wb = Workbook()
    ws = wb.worksheets[0]
    ws.title = title
    return wb, ws


def empty_workbook() -> Workbook:
    """Нова книга без аркушів - їх додають через create_sheet."""
    wb = Workbook()
    wb.remove(wb.worksheets[0])
    return wb


def save_workbook(wb: Workbook, path: str) -> None:
    """Зберігає книгу, за потреби створивши теку. Якщо файл відкритий в Excel,
    піднімає PermissionError з поясненням, що робити, а не сирий traceback."""
    os.makedirs(os.path.dirname(path) or ".", exist_ok=True)
    try:
        wb.save(path)
    except PermissionError as e:
        raise PermissionError(
            f"Не вдалось зберегти '{path}': файл зараз відкритий в іншій програмі "
            "(напр. Excel). Закрийте його та спробуйте ще раз."
        ) from e
