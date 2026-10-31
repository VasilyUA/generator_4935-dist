"""Звірка книг resources/addons між собою, без Табеля (окреме перше питання
index.py).

Кожна книга будує для людини власну шкалу (_paint_timeline лише з її записів,
за тими самими правилами, що й у звірці з Табелем). На датах, де щось
очікують обидві книги пари, статуси звіряються тими самими правилами
еквівалентності в обидва боки (_statuses_match: ППД=НОВ=БЗВП, 10 - статуси,
які ним оплачуються, "будь-який" тощо). День повернення чи виписки
(arrival_date будь-якої з книг) не звіряється - книги записують цю межу
по-різному. Запис, повністю вкладений у період іншої книги за
ADDON_NESTED_STATUSES (ШП посеред БЗВП), - теж не розбіжність.

Звіряються всі періоди, а не лише місяць Табеля: шкала людини триває від
найранішого запису в будь-якій книзі до last_date (сьогодні чи пізніший
відомий кінець періоду). Люди беруться з Табеля (resources/ОБЛІК.xlsx).

Звіт - окремий аркуш на кожен файл addons (назва аркуша - ім'я файлу), де
цей файл порівняно з усіма іншими. Одна таблиця з однією шапкою; рядок -
пара записів двох файлів, що суперечать одне одному. Назва файлу стоїть у
кожному рядку перед його аркушем, статусом і періодом: узагальнене "інший
файл" у шапці та розділені шапки користувач відхилив як незрозумілі.
Червоним - лише те, що не сходиться (статус, дата початку, дата кінця). Та
сама розбіжність є на аркушах обох файлів."""

from datetime import date, datetime, timedelta
from itertools import combinations
from typing import NamedTuple

from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter

from constants import ADDON_NESTED_STATUSES, DATE_NUMBER_FORMAT, PIB_COLUMN_NAME
from content.addon_mismatch_checker import (
    _paint_timeline,
    _roster_pidrozdil,
    _sheet_title,
    _statuses_match,
    status_options,
)
from content.addon_reader import ANY_STATUS, is_implausible_future_date
from content.excel_compat import empty_workbook, save_workbook
from content.oblik_timesheet import normalize_name
from content.payment_mismatch_checker import (
    _BORDER,
    _display_value,
    _FONT,
    _HEADER_ALIGNMENT,
    _HEADER_FILL,
    _HEADER_FONT,
    _MISMATCH_FILL,
)
from content.report_pib_column import pib_column_width

_ANY_STATUS_LABEL = "будь-який"


def _nested(outer, inner):
    """Чи може inner лежати всередині outer: статус inner дозволено
    ADDON_NESTED_STATUSES для статусу чи аркуша outer, і період inner (як
    записано у файлі) повністю в межах outer. Без кінця inner невідомо, чи він
    всередині, тож ні."""
    allowed = {status for key in (outer["status"], outer["source_sheet"])
               for status in ADDON_NESTED_STATUSES.get(key, ())}
    if not allowed & set(status_options(inner["status"])) or inner["end_date"] is None:
        return False
    return outer["start_date"] <= inner["start_date"] and (
        outer["end_date"] is None or inner["end_date"] <= outer["end_date"]
    )


def _agree(record_a, record_b, date_value):
    if date_value in (record_a.get("arrival_date"), record_b.get("arrival_date")):
        return True
    if _nested(record_a, record_b) or _nested(record_b, record_a):
        return True
    status_a, status_b = record_a["status"], record_b["status"]
    return _statuses_match(status_a, status_b) or _statuses_match(status_b, status_a)


def _last_date(period_records, today=None):
    """Кінець шкали - сьогодні або пізніший відомий кінець/початок періоду."""
    known = [r["end_date"] for r in period_records if r["end_date"]] + [r["start_date"] for r in period_records]
    return max([today or date.today()] + known)


def _days(start, end):
    return [(None, start + timedelta(days=offset)) for offset in range((end - start).days + 1)]


def _conflict_periods(conflict_days):
    """[(дата, запис книги A, запис книги B)] за зростанням дати -> періоди.
    Дні поспіль з тією самою парою записів - один період {"start", "end",
    "status_a", "status_b", "sheet_a", "sheet_b", "period_a", "period_b"};
    period_* - період самого запису, як він записаний у файлі: (початок,
    кінець), кінець None - у файлі не вказано."""
    periods = []
    for day, record_a, record_b in conflict_days:
        last = periods[-1] if periods else None
        if last and last["records"] == (record_a, record_b) and last["end"] + timedelta(days=1) == day:
            last["end"] = day
        else:
            periods.append({"records": (record_a, record_b), "start": day, "end": day})
    return [
        {"start": p["start"], "end": p["end"],
         "status_a": p["records"][0]["status"], "status_b": p["records"][1]["status"],
         "sheet_a": p["records"][0]["source_sheet"], "sheet_b": p["records"][1]["source_sheet"],
         "period_a": (p["records"][0]["start_date"], p["records"][0]["end_date"]),
         "period_b": (p["records"][1]["start_date"], p["records"][1]["end_date"])}
        for p in periods
    ]


def find_addon_conflicts(timesheet, period_records, source_files, last_date=None):
    """Пари книг (у порядку source_files), у яких для людини з Табеля є хоча б
    одна дата з різними статусами, за всі періоди. Повертає [{"files": (книга
    A, книга B), "roster_pib_raw", "roster_pidrozdil", "periods":
    _conflict_periods(...), "days": кількість розбіжних днів}]."""
    person_rows = timesheet.person_rows()
    pib_col = timesheet.label_columns[PIB_COLUMN_NAME]
    last_date = last_date or _last_date(period_records)

    by_person = {}
    for record in period_records:
        normalized = normalize_name(record["pib_raw"])
        if normalized in person_rows:
            by_person.setdefault(normalized, {}).setdefault(record["source_file"], []).append(record)

    conflicts = []
    for normalized, by_file in sorted(by_person.items()):
        present = [name for name in source_files if name in by_file]
        if len(present) < 2:
            continue
        first = min(r["start_date"] for name in present for r in by_file[name])
        days = _days(first, last_date)
        timelines = {name: _paint_timeline(by_file[name], days, last_date) for name in present}
        for file_a, file_b in combinations(present, 2):
            conflict_days = []
            for day in sorted(set(timelines[file_a]) & set(timelines[file_b])):
                record_a, record_b = timelines[file_a][day], timelines[file_b][day]
                if not _agree(record_a, record_b, day):
                    conflict_days.append((day, record_a, record_b))
            if conflict_days:
                row_idx = person_rows[normalized]
                conflicts.append({
                    "files": (file_a, file_b),
                    "roster_pib_raw": timesheet.ws.cell(row=row_idx, column=pib_col).value,
                    "roster_pidrozdil": _roster_pidrozdil(timesheet, row_idx),
                    "periods": _conflict_periods(conflict_days),
                    "days": len(conflict_days),
                })
    return conflicts


def _status_display(status):
    return _ANY_STATUS_LABEL if status == ANY_STATUS else _display_value(status)


def file_rows(conflicts, source_file):
    """Рядки аркуша файлу source_file: пари записів цього файлу ("my_*") та
    іншого ("other_*"), що суперечать одне одному, відсортовані за ПІБ,
    іншим файлом і датою. Звіт показує періоди самих записів, тож та сама пара
    кількома проміжками днів (між ними - день, що не звіряється, чи інший
    запис) дає один рядок: "days" - сума днів, "start"/"end" - перший і
    останній розбіжний день."""
    rows = {}
    for record in conflicts:
        file_a, file_b = record["files"]
        if source_file not in (file_a, file_b):
            continue
        mine, other = ("a", "b") if source_file == file_a else ("b", "a")
        for period in record["periods"]:
            row = {
                "roster_pidrozdil": record["roster_pidrozdil"], "roster_pib_raw": record["roster_pib_raw"],
                "other_file": file_b if mine == "a" else file_a,
                "my_sheet": period[f"sheet_{mine}"], "my_status": period[f"status_{mine}"],
                "my_period": period[f"period_{mine}"],
                "other_sheet": period[f"sheet_{other}"], "other_status": period[f"status_{other}"],
                "other_period": period[f"period_{other}"],
            }
            key = tuple(str(value) for value in row.values())
            days = (period["end"] - period["start"]).days + 1
            if key in rows:
                merged = rows[key]
                merged["start"], merged["end"] = min(merged["start"], period["start"]), max(merged["end"], period["end"])
                merged["days"] += days
            else:
                rows[key] = {**row, "start": period["start"], "end": period["end"], "days": days}
    return sorted(rows.values(), key=lambda r: (str(r["roster_pib_raw"]), r["other_file"], r["start"]))


class _Compared(NamedTuple):
    """Що звіряється між записами двох файлів: назва в поясненні й колонки
    цього та іншого файлу - вони підсвічуються червоним, коли різні."""
    name: str
    my_col: int
    other_col: int


_COMPARED = (_Compared("статус", 5, 10), _Compared("дата початку", 6, 11), _Compared("дата кінця", 7, 12))


def _differences(row):
    """Те з _COMPARED, що в записах двох файлів не сходиться."""
    (my_start, my_end), (other_start, other_end) = row["my_period"], row["other_period"]
    pairs = ((row["my_status"], row["other_status"]), (my_start, other_start), (my_end, other_end))
    return [compared for compared, (mine, other) in zip(_COMPARED, pairs) if mine != other]


def _explanation(row, source_file, today=None):
    """Текст "У чому розбіжність": що саме не сходиться (ті самі клітинки
    підсвічено червоним). Якщо дата котрогось із файлів далеко в майбутньому
    (is_implausible_future_date), додається попередження про ймовірну помилку
    в році - реальний випадок: виписка 2027 року замість 2026 давала
    "розбіжність" з відпусткою через місяць."""
    text = "Не сходиться: " + ", ".join(compared.name for compared in _differences(row))
    for file_name, period in ((source_file, row["my_period"]), (row["other_file"], row["other_period"])):
        for value in period:
            if is_implausible_future_date(value, today):
                text += (f". Увага: у «{file_name}» дата {value:%d.%m.%Y} - далеко в майбутньому, "
                         "ймовірно помилка в році (виправте у файлі)")
    return text


_TITLE_FONT = Font(name="Times New Roman", size=14, bold=True)
_NOTE_FONT = Font(name="Times New Roman", size=12, italic=True)
# Червоним - лише те, що не сходиться (той самий світло-червоний, що й у
# error_mis_statuses.xlsx); решта клітинок однакові, без кольорових груп
# колонок. Групи "цей файл" / "порівняно з файлом" / пояснення відділяє лише
# товща ліва межа.
_DIFFERENT_FONT = Font(name="Times New Roman", size=14, color="FF9C0006")
_IMPLAUSIBLE_DATE_FONT = Font(name="Times New Roman", size=14, bold=True, color="FFC00000")
_GROUP_START_BORDER = Border(left=Side(style="medium"), right=_BORDER.right, top=_BORDER.top, bottom=_BORDER.bottom)
_GROUP_START_COLUMNS = (3, 8, 13)
_LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)
_CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
_NO_CONFLICTS_TEXT = "Розбіжностей з іншими файлами немає."
_NO_END_TEXT = "не вказано"
_NOTE_TEXT = ("Червоним - те, що в записах двох файлів не сходиться. Дати - період запису так, як він "
              "записаний у самому файлі.")
_HEADER_ROW = 3
# Ширина колонки з датою: вужча показувала "#######" замість дати шрифтом
# звіту (Times New Roman 14).
_DATE_WIDTH = 16


# Колонок "розбіжні дні" немає: день повернення не звіряється, тож вони
# відрізнялися від однакових періодів обох файлів на день, і користувач
# сприймав це як розбіжність дат.
class _Column(NamedTuple):
    title: str
    width: int
    left_aligned: bool


_COLUMNS = (
    _Column("Підрозділ", 18, True), _Column("ПІБ", 34, True),
    _Column("Файл", 22, True), _Column("Аркуш", 20, True), _Column("Статус", 11, False),
    _Column("З", _DATE_WIDTH, False), _Column("По", _DATE_WIDTH, False),
    _Column("Порівняно з файлом", 26, True), _Column("Аркуш", 20, True), _Column("Статус", 11, False),
    _Column("З", _DATE_WIDTH, False), _Column("По", _DATE_WIDTH, False),
    _Column("У чому розбіжність", 45, True),
)


def _date_cell(value, missing=_NO_END_TEXT):
    return datetime.combine(value, datetime.min.time()) if value else missing


def _border(col_idx):
    return _GROUP_START_BORDER if col_idx in _GROUP_START_COLUMNS else _BORDER


def _write_file_sheet(ws, source_file, rows, other_files, today=None):
    ws.cell(row=1, column=1, value=(
        f"Файл «{source_file}» порівняно з: {', '.join(other_files) or '-'} "
        f"(розбіжностей: {len(rows)}, людей: {len({r['roster_pib_raw'] for r in rows})})"
    )).font = _TITLE_FONT
    ws.cell(row=2, column=1, value=_NOTE_TEXT).font = _NOTE_FONT
    for col_idx, column in enumerate(_COLUMNS, start=1):
        cell = ws.cell(row=_HEADER_ROW, column=col_idx, value=column.title)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.alignment = _HEADER_ALIGNMENT
        cell.border = _border(col_idx)
        ws.column_dimensions[get_column_letter(col_idx)].width = column.width
    ws.column_dimensions["B"].width = pib_column_width([row["roster_pib_raw"] for row in rows])  # ПІБ видно повністю
    ws.row_dimensions[_HEADER_ROW].height = 36
    ws.freeze_panes = ws.cell(row=_HEADER_ROW + 1, column=3)  # шапка й ПІБ лишаються на місці під час прокрутки

    if not rows:
        cell = ws.cell(row=_HEADER_ROW + 1, column=1, value=_NO_CONFLICTS_TEXT)
        ws.merge_cells(start_row=_HEADER_ROW + 1, start_column=1, end_row=_HEADER_ROW + 1, end_column=len(_COLUMNS))
        cell.font = _FONT
        cell.alignment = _CENTER
        return

    ws.auto_filter.ref = f"A{_HEADER_ROW}:{get_column_letter(len(_COLUMNS))}{_HEADER_ROW + len(rows)}"
    group_start, previous = _HEADER_ROW + 1, None
    for offset, data in enumerate(rows):
        row_idx = _HEADER_ROW + 1 + offset
        if data["roster_pib_raw"] != previous:
            group_start, previous = row_idx, data["roster_pib_raw"]
        (my_start, my_end), (other_start, other_end) = data["my_period"], data["other_period"]
        values = [
            _display_value(data["roster_pidrozdil"]), _display_value(data["roster_pib_raw"]),
            source_file, data["my_sheet"], _status_display(data["my_status"]),
            _date_cell(my_start), _date_cell(my_end),
            data["other_file"], data["other_sheet"], _status_display(data["other_status"]),
            _date_cell(other_start), _date_cell(other_end),
            _explanation(data, source_file, today),
        ]
        different = {col for compared in _differences(data) for col in (compared.my_col, compared.other_col)}
        for col_idx, (column, value) in enumerate(zip(_COLUMNS, values), start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            if isinstance(value, datetime) and is_implausible_future_date(value.date(), today):
                cell.font = _IMPLAUSIBLE_DATE_FONT
            else:
                cell.font = _DIFFERENT_FONT if col_idx in different else _FONT
            if col_idx in different:
                cell.fill = _MISMATCH_FILL
            cell.border = _border(col_idx)
            cell.alignment = _LEFT if column.left_aligned else _CENTER
            if isinstance(value, datetime):
                cell.number_format = DATE_NUMBER_FORMAT
        is_last_of_person = offset == len(rows) - 1 or rows[offset + 1]["roster_pib_raw"] != previous
        if is_last_of_person and row_idx > group_start:
            for col_idx in (1, 2):
                ws.merge_cells(start_row=group_start, start_column=col_idx, end_row=row_idx, end_column=col_idx)


def write_addon_conflicts_report(conflicts, output_path, source_files=(), today=None):
    """Пише звіт: аркуш на кожен файл addons (source_files - усі звірені
    файли), де цей файл порівняно з усіма іншими (file_rows). Аркуш файлу без
    розбіжностей містить лише напис _NO_CONFLICTS_TEXT. Якщо розбіжностей
    немає взагалі, файл не пишеться і повертається None. today (None -
    сьогодні) потрібен, щоб помітити дати, надто далекі в майбутньому."""
    if not conflicts:
        return None
    files = list(dict.fromkeys(list(source_files) + [name for record in conflicts for name in record["files"]]))

    wb = empty_workbook()
    used_titles = set()
    for source_file in files:
        others = [name for name in files if name != source_file]
        _write_file_sheet(wb.create_sheet(_sheet_title(source_file, used_titles)), source_file,
                          file_rows(conflicts, source_file), others, today)
    save_workbook(wb, output_path)
    return output_path
