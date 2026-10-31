import difflib
import functools
import re
from collections import Counter
from datetime import date, datetime, timedelta

from openpyxl.styles import PatternFill
from openpyxl.utils.exceptions import InvalidFileException

from constants import (
    DEFAULT_STATUS,
    LABEL_COLUMN_NAMES,
    PAYMENT_CATEGORY_100_EXCLUDE_NAMES,
    PAYMENT_CATEGORY_100_MAX_DAYS,
    PAYMENT_STATUS_CONVERSIONS,
    PAYMENT_STATUS_OVERRIDES,
    PAYMENT_STATUS_PRIORITY,
    PERSONEL_LIST_SHEET_NAME,
    PIB_COLUMN_NAME,
    STATUS_COLORS,
    STATUS_CONFLICT_OVERRIDES,
    STATUS_CONFLICT_PAID_BY_DEFAULT,
    STATUS_CONFLICT_PRIORITY_EXEMPT,
    month_limit_key,
)
from content.excel_compat import load_workbook, save_workbook


@functools.cache
def _fill_for_status(status):
    """PatternFill кольору STATUS_COLORS[status] або None, якщо для статусу
    кольору немає (тоді клітинка лишається без заливки - це не помилка)."""
    color = STATUS_COLORS.get(status)
    return PatternFill(start_color=color, end_color=color, fill_type="solid") if color else None


def _paint(cell, status):
    """Заливає клітинку кольором статусу; статус без кольору лишає її як є."""
    fill = _fill_for_status(status)
    if fill:
        cell.fill = fill


# Апостроф - будь-яке його написання; усередині слова разом із пробілами
# після нього ("В' ячеслав").
_APOSTROPHES = "'’ʼ`‘"
_APOSTROPHE_IN_WORD_RE = re.compile(rf"(?<=[^\W\d_])[{_APOSTROPHES}]\s*(?=[^\W\d_])")
_APOSTROPHE_RE = re.compile(rf"[{_APOSTROPHES}]")


def normalize_name(name):
    """Нормалізує ПІБ для зіставлення рапорту з ОБЛІК.xlsx: пробіли, регістр,
    апостроф. Це власна копія логіки generator_br_and_report_for_money
    (content/money_report_helpers.py, checker_accounting/checker.py) - проєкти
    не імпортують один одного. Апостроф (будь-якого виду, з пробілом після
    нього чи без, є він чи немає) для зіставлення не важить: в ОБЛІК було
    "В’ячеслав", у book61 "Вячеслав", і людину не зіставляло зовсім, тож і
    розбіжності ніде не було видно. Що написання різне, звіти показують
    окремо (spelled_differently)."""
    if not isinstance(name, str):
        return ""
    cleaned = _APOSTROPHE_RE.sub("", _APOSTROPHE_IN_WORD_RE.sub("", name.strip()))
    return " ".join(cleaned.split()).upper()


def spelled_differently(name_a, name_b):
    """Чи різні по суті два написання ПІБ однієї людини: не пробілами,
    регістром чи видом апострофа ("'" / "’"), а напр. є апостроф чи немає -
    звіти мають це показувати. Якщо з одного боку не рядок (напр. напис "Не
    знайдено в ОБЛІК.xlsx"), різні, коли різні самі значення."""
    if not (isinstance(name_a, str) and isinstance(name_b, str)):
        return name_a != name_b

    def spelling(name):
        unified = re.sub(rf"\s*[{_APOSTROPHES}]\s*", "'", name)
        return " ".join(unified.split()).upper()

    return spelling(name_a) != spelling(name_b)


def _strip_external_links(wb):
    """У resources/ОБЛІК.xlsx лишилось "осиротіле" зовнішнє посилання - залишок
    давно видаленої перевірки даних чи формули, що вказувала на старий файл
    resources/information_unit/*.xlsx. Ним ніщо в книзі не користується, але
    openpyxl не вміє правильно його перезберегти (зв'язок r:id лишається
    "rId1", а запис у externalLink1.xml.rels губиться), і Excel щоразу
    повідомляв про пошкоджений вміст згенерованого файлу ("Знайдено проблему
    з вмістом... Відновлені записи"). Тому перед збереженням посилання
    просто прибирається."""
    wb._external_links = []


class Timesheet:
    """Робоча книга ОБЛІК.xlsx (аркуш "Табель"). З неї читається лише
    структура заголовка: колонки-дати й іменовані колонки (за назвою, а не за
    позицією). Сама книга (wb/ws) лишається відкритою: apply_events редагує
    клітинки напряму, а save зберігає ту саму книгу, тож стилі, ширини
    колонок і об'єднані клітинки зберігаються самі собою (на відміну від
    checker_accounting/checker.py, який будує нову книгу й копіює стилі
    вручну)."""

    def __init__(self, file_path, sheet_name):
        self.file_path = file_path
        # Файл для перевірок користувач обирає діалогом (index.py), і буває,
        # що обрано книгу без аркуша "Табель" (напр. відомість) - тоді
        # зрозуміле повідомлення замість KeyError на весь скрипт.
        try:
            self.wb = load_workbook(file_path)
        except InvalidFileException as e:
            raise ValueError(f"Файл {file_path} - не книга Excel (.xlsx/.xlsm).") from e
        if sheet_name not in self.wb.sheetnames:
            raise ValueError(f"У файлі {file_path} немає аркуша '{sheet_name}' - оберіть файл ОБЛІК.")
        self.ws = self.wb[sheet_name]
        # {(row_idx, дата): статус}, який apply_payment_values замінив числом
        # (ППД -> 10). Число - лише вид оплати, тож звіти й повідомлення
        # командирам показують цей статус (рішення користувача).
        self.statuses_behind_numbers = {}

        header = [self.ws.cell(row=1, column=c).value for c in range(1, self.ws.max_column + 1)]

        self.label_columns = {
            value.strip().upper(): idx
            for idx, value in enumerate(header, start=1)
            if isinstance(value, str) and value.strip().upper() in LABEL_COLUMN_NAMES
        }
        if PIB_COLUMN_NAME not in self.label_columns:
            raise ValueError(f"У файлі {file_path} не знайдено колонку '{PIB_COLUMN_NAME}' у заголовку.")

        self.date_columns = sorted(
            (
                (idx, value.date() if isinstance(value, datetime) else value)
                for idx, value in enumerate(header, start=1)
                if isinstance(value, date)
            ),
            key=lambda pair: pair[1],
        )
        if not self.date_columns:
            raise ValueError(f"У файлі {file_path} не знайдено жодної колонки-дати в заголовку.")

    def person_rows(self):
        """{normalize_name(ПІБ): row_idx} - рядок з порожнім ПІБ пропускається."""
        pib_col = self.label_columns[PIB_COLUMN_NAME]
        rows = {}
        for row_idx in range(2, self.ws.max_row + 1):
            pib_raw = self.ws.cell(row=row_idx, column=pib_col).value
            if pib_raw:
                rows[normalize_name(pib_raw)] = row_idx
        return rows

    def save(self, output_path):
        _strip_external_links(self.wb)
        save_workbook(self.wb, output_path)


def export_for_money_project(source_path, target_path, target_sheet_name):
    """Готує вже згенерований "ОБЛІК для виплат" (source_path) для сусіднього
    проєкту generator_br_and_report_for_money (index.py пропонує це лише після
    повністю чистого прогону). Книга завантажується окремо, тож source_path не
    змінюється, а аркуш "Табель" перейменовується на target_sheet_name -
    constants.MONEY_PROJECT_SHEET_NAME, назву поточного місяця (напр.
    "СЕРПЕНЬ"), яку шукає той проєкт (інакше він бере запасний аркуш і читає
    не ті дані). Колонки (ПІДРОЗДІЛ/ПОСАДА/ЗВАННЯ/ПІБ, дати, ПІДСТАВИ/ДАТА В
    СТАТУС СПЕЦКОНТИНГЕНТУ) вже такі, як чекає той проєкт, тож перетворювати
    дані не треба."""
    wb = load_workbook(source_path)
    wb[PERSONEL_LIST_SHEET_NAME].title = target_sheet_name
    _strip_external_links(wb)
    save_workbook(wb, target_path)


# Поріг схожості (difflib.SequenceMatcher.ratio) для _typo_match, перевірений
# на реальному resources/ОБЛІК.xlsx (300+ ПІБ): справжня описка (переставлені
# літери в прізвищі, ratio ~0.96) дає рівно один збіг, а вигадане відсутнє ПІБ
# - жодного. Поріг достатньо високий, щоб не підказати іншу людину.
_TYPO_HINT_SIMILARITY_CUTOFF = 0.8


def _typo_match(timesheet, person_rows, pib_col, normalized):
    """(candidate_normalized, candidate_pib_raw), якщо в Табелі є рівно одне ПІБ,
    дуже схоже на normalized (difflib, поріг _TYPO_HINT_SIMILARITY_CUTOFF), але
    не тотожне. Зазвичай це описка в рапорті (напр. переставлені літери в
    прізвищі). apply_events застосовує подію до цього кандидата автоматично
    (рішення користувача), а не лише підказує. На відміну від
    _candidate_roster_person у payment_mismatch_checker, тут описка може бути
    й у прізвищі, тож "перші два слова" не підходять. (None, None) - збігів
    немає чи їх кілька: краще нічого не застосувати, ніж застосувати не до
    тієї людини.

    Прізвище (перше слово) має окремо пройти той самий поріг. Реальна помилка:
    дві різні людини з однаковими ім'ям і по батькові, але різними прізвищами
    (тест test_apply_events_does_not_auto_apply_when_similar_pib_has_a_different_surname)
    - спільна довга частина давала ratio усього ПІБ 0.89, вище за поріг, хоча
    прізвища зовсім різні (ratio ~0.5, а справжня описка - 0.9 і вище, напр.
    "ЧОТИРНАЦДЯТИЙ"/"ЧОТИРНАДЦЯТИЙ" - 0.92). Тоді це не описка, а інша
    людина."""
    matches = difflib.get_close_matches(normalized, person_rows.keys(), n=2, cutoff=_TYPO_HINT_SIMILARITY_CUTOFF)
    if len(matches) != 1:
        return None, None
    candidate_normalized = matches[0]
    target_surname = normalized.split(" ", 1)[0]
    candidate_surname = candidate_normalized.split(" ", 1)[0]
    surname_ratio = difflib.SequenceMatcher(None, target_surname, candidate_surname).ratio()
    if surname_ratio < _TYPO_HINT_SIMILARITY_CUTOFF:
        return None, None
    candidate_pib_raw = timesheet.ws.cell(row=person_rows[candidate_normalized], column=pib_col).value
    return candidate_normalized, candidate_pib_raw


def _automatic_conflict_winner(status_a, status_b):
    """Хто переможе, коли на одну дату два різні статуси, а готової пари в
    STATUS_CONFLICT_OVERRIDES немає: статус, що оплачується за замовчуванням
    (STATUS_CONFLICT_PAID_BY_DEFAULT), перемагає той, що ні. Реальний випадок:
    ВЛК почалося під час ВПСЗ - за ВЛК належить виплата, тож "ВЛК" перемагає
    "ВПСЗ".

    None (суперечність іде на ручну перевірку), якщо будь-який зі статусів - у
    STATUS_CONFLICT_PRIORITY_EXEMPT (адміністративний, дисциплінарний чи
    остаточний статус, для якого "хто оплачується" не має сенсу), або обидва
    оплачуються (напр. "ШП" і "ВЛК" - занадто неоднозначно), або жоден."""
    if status_a in STATUS_CONFLICT_PRIORITY_EXEMPT or status_b in STATUS_CONFLICT_PRIORITY_EXEMPT:
        return None
    a_paid = status_a in STATUS_CONFLICT_PAID_BY_DEFAULT
    b_paid = status_b in STATUS_CONFLICT_PAID_BY_DEFAULT
    if a_paid == b_paid:
        return None
    return status_a if a_paid else status_b


def apply_events(timesheet, events, report_dates):
    """Заповнює колонки-дати аркуша. Для кожної дати з рапортом (за
    зростанням) застосовуються події цієї дати (зміна статусу), а решта людей
    отримує статус з попередньої колонки-дати. Так само й для дат між
    рапортами: подій немає, тож статус просто переноситься.

    Статус, уже вписаний вручну в resources/ОБЛІК.xlsx, має вищий пріоритет,
    ніж подія з рапорту чи перенесений статус (рішення користувача, раніше
    було навпаки). Реальний випадок: людина фактично прибула в пункт постійної
    дислокації, а рапорт пише лише загальне "прибув в район виконання
    завдань", і жоден шаблон тексту цього не впізнав би - користувач знає
    краще. Це відбувається мовчки, без запису в unresolved: людина свідомо
    керує цим значенням. Такий "якір" переноситься на всі наступні колонки,
    доки не трапиться подія з рапорту чи інший вписаний вручну статус.

    Найперша колонка-дата (базовий стан напередодні першого рапорту) не
    змінюється, лише її заливка оновлюється за наявним текстом. Дати після
    останнього рапорту самі не заповнюються - жоден рапорт ще не підтвердив
    "без змін", тож порожні колонки там лишаються порожніми. Але вписаний
    вручну статус і там стає якорем і переноситься на наступні порожні
    колонки, до наступного такого статусу чи кінця табеля.

    Повертає [{"reason", ...}] - записи, про які треба знати:
    - ПІБ з рапорту немає в Табелі буквально, але є рівно одне дуже схоже
      написання (_typo_match, зазвичай описка в рапорті): подія однаково
      застосовується до цієї людини (рішення користувача), а запис лише
      повідомляє про це. Справді не застосовується, лише якщо схожих
      написань немає або їх кілька.
    - effective_date не збігається з жодною колонкою-датою (зазвичай описка
      в даті рапорту).
    - На одну людину й дату (можливо, з різних рапортів) прийшли різні
      статуси. Спершу: якщо одна з подій - прогноз (item["is_forecast"], дата
      відпустки лише обчислена як "наступний день" у
      _discharge_vlk_leave_events), перемагає справжня подія, хоч би в якому
      порядку вони йшли, і без запису (прогноз просто не справдився). Інакше,
      якщо пара є в STATUS_CONFLICT_OVERRIDES чи переможця визначає
      _automatic_conflict_winner ("хто оплачується"), застосовується він без
      запису. Інакше застосовується останній статус, а запис іде на ручну
      перевірку.

    Мовчки пропускаються (без запису):
    - effective_date раніша за базову колонку, а статус події той самий, що
      вже стоїть у базовій колонці (напр. "з 21.07 у відпустці", а табель
      починається з 31.07 і там уже ВПСЗ): подія лише підтверджує записане, а
      базова колонка не змінюється. Якщо ж статус інший (зазвичай описка в
      даті, "06.07" замість "06.08"), запис іде на ручну перевірку: сама
      рання дата ще не доказ, лише збіг з базовою колонкою знімає сумнів.
    - effective_date - рівно наступний день після останньої колонки-дати
      (зазвичай 1 число наступного місяця). Рапорт останнього дня місяця
      описує складену подію ("виписаний ... та проходить ВЛК ..., потребує
      відпустки ..." - _discharge_vlk_leave_events), де відпустка лише
      обчислена як "наступний день", а колонок наступного місяця ще немає.
      Перевірити це зараз неможливо, і це не помилка (описка дала б іншу
      дату). Підтвердження прийде з рапортом наступного місяця
      (_confirmed_health_leave_event чи звичайне перенесення статусу)."""
    person_rows = timesheet.person_rows()
    date_col_by_date = {d: idx for idx, d in timesheet.date_columns}
    ordered_dates = [d for _, d in timesheet.date_columns]
    pib_col = timesheet.label_columns[PIB_COLUMN_NAME]

    unresolved = []
    events_by_person_date = {}
    for event in events:
        normalized = normalize_name(event["pib_raw"])
        if normalized not in person_rows:
            candidate_normalized, candidate_pib_raw = _typo_match(timesheet, person_rows, pib_col, normalized)
            if candidate_normalized is None:
                unresolved.append({
                    "reason": f"ПІБ \"{event['pib_raw']}\" з рапорту не знайдено в {timesheet.file_path}.",
                    **event,
                })
                continue
            unresolved.append({
                "reason": (
                    f"ПІБ \"{event['pib_raw']}\" з рапорту не збігається буквально з роcтеровим написанням "
                    f"\"{candidate_pib_raw}\" - схоже на друкарську помилку в рапорті, статус застосовано "
                    "автоматично за схожістю написання (перевірте, чи це справді та сама людина)."
                ),
                **event,
            })
            normalized = candidate_normalized
        else:
            roster_pib_raw = timesheet.ws.cell(row=person_rows[normalized], column=pib_col).value
            if spelled_differently(event["pib_raw"], roster_pib_raw):
                unresolved.append({
                    "reason": (
                        f"ПІБ \"{event['pib_raw']}\" з рапорту написано інакше, ніж у роcтері "
                        f"\"{roster_pib_raw}\" (апостроф) - статус застосовано; перевірте, яке написання правильне."
                    ),
                    **event,
                })
        if event["effective_date"] not in date_col_by_date:
            if event["effective_date"] < ordered_dates[0]:
                baseline_col = timesheet.date_columns[0][0]
                baseline_value = timesheet.ws.cell(row=person_rows[normalized], column=baseline_col).value
                if baseline_value == event["status"]:
                    continue
            elif event["effective_date"] == ordered_dates[-1] + timedelta(days=1):
                continue
            unresolved.append({
                "reason": (
                    f"Дата \"{event['effective_date'].strftime('%d.%m.%Y')}\" для \"{event['pib_raw']}\" "
                    f"відсутня серед колонок-дат {timesheet.file_path}."
                ),
                **event,
            })
            continue

        key = (normalized, event["effective_date"])
        previous = events_by_person_date.get(key)
        if previous is not None and previous["status"] != event["status"]:
            previous_is_forecast = previous.get("is_forecast", False)
            event_is_forecast = event.get("is_forecast", False)
            if previous_is_forecast != event_is_forecast:
                # Прогнозована подія (напр. ВПСЗ "наступного дня" з
                # _discharge_vlk_leave_events, коли дату відпустки лише
                # обчислено як "день після ВЛК") ніколи не перемагає справжню
                # подію на ту саму дату з іншого рапорту - навіть якщо прогноз
                # прийшов пізніше. Реальний випадок: рапорт передбачав
                # відпустку наступного дня після ВЛК, а пізніший рапорт показав,
                # що людина того дня повернулась у ППД, а відпустка почалася на
                # кілька днів пізніше. Без запису на ручну перевірку - прогноз
                # просто не справдився.
                events_by_person_date[key] = event if previous_is_forecast else previous
                continue
            override_status = STATUS_CONFLICT_OVERRIDES.get(frozenset({previous["status"], event["status"]}))
            if override_status is None:
                override_status = _automatic_conflict_winner(previous["status"], event["status"])
            if override_status is not None:
                events_by_person_date[key] = event if event["status"] == override_status else previous
                continue
            unresolved.append({
                "reason": (
                    f"\"{event['pib_raw']}\" {event['effective_date'].strftime('%d.%m.%Y')}: суперечливі статуси "
                    f"в рапорті - \"{previous['status']}\" і \"{event['status']}\" (застосовано останній)."
                ),
                **event,
            })
        events_by_person_date[key] = event

    if not report_dates:
        return unresolved

    last_report_date = max(report_dates)
    fill_dates = [d for d in ordered_dates if d > ordered_dates[0] and d <= last_report_date]
    # Дати після останнього рапорту самі не заповнюються: жоден рапорт ще не
    # підтвердив для них "без змін". Але статус, уже вписаний там вручну в
    # resources/ОБЛІК.xlsx (реальний випадок - "БЗ" поза діапазоном
    # рапортів), переноситься на всі наступні порожні колонки, доки не
    # трапиться наступний вписаний статус чи не скінчаться колонки. Порожні
    # колонки до першого такого статусу лишаються порожніми - починати нема
    # з чого.
    tail_dates = [d for d in ordered_dates if d > last_report_date]

    for normalized, row_idx in person_rows.items():
        baseline_cell = timesheet.ws.cell(row=row_idx, column=timesheet.date_columns[0][0])
        current_status = baseline_cell.value
        # Значення базової колонки не змінюється - лише заливка, щоб колір
        # відповідав решті рядка.
        _paint(baseline_cell, current_status)
        deferred_status = None
        for date_value in fill_dates:
            event = events_by_person_date.get((normalized, date_value))
            existing_value = timesheet.ws.cell(row=row_idx, column=date_col_by_date[date_value]).value
            next_deferred_status = None
            if existing_value and event is not None and existing_value != event["status"]:
                # Статус, вписаний вручну в resources/ОБЛІК.xlsx, має перевагу
                # над подією з рапорту, навіть коли рапорт на ту саму дату каже
                # інше (людина фактично прибула в ППД, а рапорт пише лише
                # загальне "прибув в район виконання завдань"). Мовчки, без
                # запису на ручну перевірку: цим значенням людина керує сама.
                current_status = existing_value
                # Подія "після 17:00" (content/daily_report_reader.py) при цьому
                # не губиться: вписаний вручну статус описує сам день, а подія -
                # вечір, тож вона діє з наступного дня (якщо на нього немає
                # власної події чи вписаного статусу). Реальний випадок: 20.10
                # вручну "ППД", рапорт - "прибув в район виконання завдань 20.10
                # після 17.00", тож 21.10 має бути РВЗ, а не перенесений ППД.
                if event.get("after_hours"):
                    next_deferred_status = event["status"]
            elif event is not None:
                current_status = event["status"]
            else:
                # Події на цю дату немає, але в resources/ОБЛІК.xlsx уже вписано
                # інший статус (не порожньо й не той, що переноситься) - він
                # стає новим якорем, як і подія з рапорту. Реальний випадок:
                # "БЗ", вписаний вручну на дату всередині діапазону рапортів.
                if existing_value and existing_value != current_status:
                    current_status = existing_value
                elif not existing_value and deferred_status is not None:
                    current_status = deferred_status
            deferred_status = next_deferred_status
            cell = timesheet.ws.cell(row=row_idx, column=date_col_by_date[date_value], value=current_status)
            _paint(cell, current_status)

        tail_status = None
        for date_value in tail_dates:
            cell = timesheet.ws.cell(row=row_idx, column=date_col_by_date[date_value])
            if cell.value:
                tail_status = cell.value
            elif tail_status is not None:
                cell.value = tail_status
            if tail_status is not None:
                _paint(cell, tail_status)

    return unresolved


def check_payment_category_100_days(timesheet, exclude_names=PAYMENT_CATEGORY_100_EXCLUDE_NAMES):
    """Правило "категорія виплати 100 - не більше ніж
    PAYMENT_CATEGORY_100_MAX_DAYS днів поспіль у місяці" для кожної людини
    табеля: порушення - безперервна серія 100 у межах місяця, довша за число
    цього місяця (будь-який інший статус чи пропущена дата її перериває). ПІБ
    з exclude_names (через той самий normalize_name) правилу не підлягають.
    Повертає [{"reason", "pib_raw"}] - запис на кожну порушену серію; сам
    табель не змінюється."""
    excluded = {normalize_name(name) for name in exclude_names}
    pib_col = timesheet.label_columns[PIB_COLUMN_NAME]
    unresolved = []
    for normalized, row_idx in timesheet.person_rows().items():
        if normalized in excluded:
            continue
        pib_raw = timesheet.ws.cell(row=row_idx, column=pib_col).value
        for start, length in _category_100_runs(timesheet, row_idx):
            limit = PAYMENT_CATEGORY_100_MAX_DAYS[month_limit_key(start)]
            if length > limit:
                end = start + timedelta(days=length - 1)
                unresolved.append({
                    "reason": (
                        f"\"{pib_raw}\": категорія виплати 100 - {length} {_days_word(length)} поспіль "
                        f"({start.strftime('%d.%m.%Y')}-{end.strftime('%d.%m.%Y')}), дозволено не більше {limit}. "
                        "Якщо так має бути - додайте ПІБ до винятків EXCLUDE_NAMES (resources/data.json)."
                    ),
                    "pib_raw": pib_raw,
                })
    return unresolved


def _days_word(count):
    """1 день, 2 дні, 5 днів (21 день, 30 днів, 31 день)."""
    if count % 10 == 1 and count % 100 != 11:
        return "день"
    if count % 10 in (2, 3, 4) and count % 100 not in (12, 13, 14):
        return "дні"
    return "днів"


def _category_100_runs(timesheet, row_idx):
    """[(початкова дата, довжина)] - безперервні серії категорії 100 у рядку
    людини, кожна в межах одного місяця: серія рахується з 1-го числа, тож
    колонка останнього дня попереднього місяця (30.09 у жовтневому Табелі) до
    неї не входить. Серію перериває будь-яке інше значення чи пропущена дата
    (не сусідній день)."""
    runs = []
    run_start, run_len, prev_day = None, 0, None
    for col_idx, day in timesheet.date_columns:
        is_100 = _is_number(timesheet.ws.cell(row=row_idx, column=col_idx).value) and \
            timesheet.ws.cell(row=row_idx, column=col_idx).value == 100
        contiguous = prev_day is not None and (day - prev_day).days == 1 and day.month == prev_day.month
        if is_100 and run_len > 0 and contiguous:
            run_len += 1
        else:
            if run_len > 0:
                runs.append((run_start, run_len))
            run_start, run_len = (day, 1) if is_100 else (None, 0)
        prev_day = day
    if run_len > 0:
        runs.append((run_start, run_len))
    return runs


def _is_number(value):
    return isinstance(value, (int, float)) and not isinstance(value, bool)


def _normalize_status_text(value):
    """Верхній регістр і згорнуті пробіли - для пошуку статусу в
    PAYMENT_STATUS_CONVERSIONS/PAYMENT_STATUS_OVERRIDES без урахування регістру
    чи зайвих пробілів. Значення з information_unit вводять вручну (напр.
    "БПШП " із зайвим пробілом), і точний пошук мовчки не знаходив ключ - без
    жодного повідомлення. У клітинку однаково пишеться значення з
    оригінального словника: функція лише для пошуку ключа."""
    return " ".join(value.split()).upper() if isinstance(value, str) else value


def _normalized_lookup(mapping):
    """{нормалізований_ключ: (оригінальний_ключ, значення), ...}."""
    return {_normalize_status_text(key): (key, value) for key, value in mapping.items()}


_NORMALIZED_PAYMENT_STATUS_CONVERSIONS = _normalized_lookup(PAYMENT_STATUS_CONVERSIONS)
_NORMALIZED_PAYMENT_STATUS_OVERRIDES = {
    _normalize_status_text(base_status): _normalized_lookup(override_map)
    for base_status, override_map in PAYMENT_STATUS_OVERRIDES.items()
}
_NORMALIZED_PAYMENT_STATUS_PRIORITY = {
    frozenset(_normalize_status_text(status) for status in pair): winner
    for pair, winner in PAYMENT_STATUS_PRIORITY.items()
}

# Скільки різних ключів PAYMENT_STATUS_CONVERSIONS дають те саме значення (напр.
# "БЗ"/"полон"/"Інт" - усі в "100_БЗ"). Такі клітинки фарбуються за новим
# значенням (STATUS_COLORS["100_БЗ"]), а не за колишнім статусом, у якого
# свій колір, - інакше та сама "100_БЗ" мала б три різні кольори. Конвертації
# з єдиним джерелом (напр. "ППД" -> 10) фарбуються за колишнім статусом (як
# підтвердив користувач): там неоднозначності немає.
_CONVERSION_VALUE_COUNTS = Counter(PAYMENT_STATUS_CONVERSIONS.values())


def _resolve_value_priority_conflict(values):
    """Якщо серед values (після нормалізації) рівно два різні значення і ця
    пара - ключ у PAYMENT_STATUS_PRIORITY (словник веде сам користувач),
    повертає [переможець] замість values. Так суперечність між файлами
    information_unit для однієї людини й дати розв'язується автоматично, і
    переможець далі йде звичайним шляхом (PAYMENT_STATUS_OVERRIDES[DEFAULT_STATUS]),
    а не на ручну перевірку як "не єдине значення". Реальний випадок: "БЗ"
    (зниклий безвісти) має переважати "Розп" (розпорядження).

    values повертаються без змін, якщо різних значень не два (0-1 -
    суперечності немає, 3 і більше - неоднозначно) чи для пари немає запису в
    PAYMENT_STATUS_PRIORITY."""
    normalized_values = {_normalize_status_text(value) for value in values}
    if len(normalized_values) != 2:
        return values
    winner = _NORMALIZED_PAYMENT_STATUS_PRIORITY.get(frozenset(normalized_values))
    return [winner] if winner is not None else values


def _remember_status_behind_number(timesheet, row_idx, date_value, status, new_value):
    """Timesheet.statuses_behind_numbers: лише коли статус-текст стає числом."""
    if _is_number(new_value) and isinstance(status, str):
        timesheet.statuses_behind_numbers[(row_idx, date_value)] = status


def apply_payment_values(timesheet, payment_values):
    """Другий прохід табеля для виплат: замінює статуси на категорії виплати.

    1. Клітинка зі статусом з PAYMENT_STATUS_CONVERSIONS (напр. "ППД",
       "БПШП") одразу й без умов отримує фіксоване значення зі словника - той
       самий статус завжди дає той самий запис, хоч би що казав
       information_unit.

    2. Клітинка зі статусом з PAYMENT_STATUS_OVERRIDES (напр. "ВД",
       DEFAULT_STATUS): якщо payment_values для цих людини й дати дає рівно
       одне значення і воно є ключем мапи замінників цього статусу, статус
       замінюється на відповідне значення з мапи - не обов'язково той самий
       текст. Реальні випадки: "ВД" в ОБЛІК, файл каже "Навч" - стає "Навч";
       "РВЗ", файл каже "БЗВП" чи "Адап" - по суті та сама причина, що й
       "ППД" (категорія 10), тож стає числом 10; "РВЗ", файл каже "ЗВФН" чи
       "ЛМР" - стає цим текстом. Заливка - за новим значенням (кожна
       категорія 10/30/70/100/170 має свій колір). Для "ВД"-подібних статусів
       суперечність чи невідоме значення лишають статус як є, без запису:
       "ВД" сам собою уваги не потребує. Для DEFAULT_STATUS обробка йде далі.

    3. Клітинка з DEFAULT_STATUS ("РВЗ") або порожня (новоприбулу людину щойно
       додали до Табеля, і статусу ще немає ні базового, ні з рапорту - це те
       саме "ще не визначено"), якщо замінник вище не спрацював, отримує
       категорію виплати з payment_values ({(normalize_name(ПІБ), date):
       [значення, ...]}, read_payment_values). Спершу
       _resolve_value_priority_conflict: два суперечливі значення з пари в
       PAYMENT_STATUS_PRIORITY розв'язуються на користь переможця. Значення
       підставляється, лише якщо лишилось рівно одне число. Інша суперечність
       чи нечислове значення для "РВЗ" мовчки не застосовується, а йде на
       ручну перевірку.

    Для порожньої клітинки правило м'якше (рішення користувача): одне
    узгоджене значення information_unit, навіть нечислове й без запису серед
    замінників (напр. "ВП"/"СЗЧ" - і так зрозумілі статуси), записується
    напряму й мовчки: іншого джерела статусу в новоприбулої людини немає.
    Справжня суперечність (два різні значення на дату) лишає клітинку
    порожньою - теж мовчки. "Підрозділ ще не подав дані" (людини й дати
    немає в payment_values) пишеться на кожну дату лише для "РВЗ"; для
    порожньої клітинки це було б ~30 однакових записів на людину, тож
    замість них - один підсумковий запис наприкінці (див. нижче).

    Увесь пошук у PAYMENT_STATUS_CONVERSIONS/PAYMENT_STATUS_OVERRIDES - без
    урахування регістру й зайвих пробілів (_normalize_status_text): значення
    information_unit вводять вручну, і зайвий пробіл інакше мовчки ламав би
    заміну.

    Наприкінці: якщо в людини жодна колонка-дата так і не отримала значення
    (ні базового статусу в ОБЛІК.xlsx, ні події з рапорту, ні даних
    information_unit), - один запис на ручну перевірку: треба вписати
    базовий статус вручну. Якщо information_unit дав хоч одне значення, цього
    запису немає - саме так прибрано колишні хибні "усі колонки-дати порожні"
    для новоприбулих, чиї дані вже підставлено.

    Повертає [{"reason", ...}] - записи на ручну перевірку."""
    pib_col = timesheet.label_columns[PIB_COLUMN_NAME]
    unresolved = []

    for normalized, row_idx in timesheet.person_rows().items():
        pib_raw = timesheet.ws.cell(row=row_idx, column=pib_col).value
        has_any_value = False
        for col_idx, date_value in timesheet.date_columns:
            cell = timesheet.ws.cell(row=row_idx, column=col_idx)
            status = cell.value
            if status is not None:
                has_any_value = True
            normalized_status = _normalize_status_text(status)

            conversion_entry = _NORMALIZED_PAYMENT_STATUS_CONVERSIONS.get(normalized_status)
            if conversion_entry is not None:
                original_status, new_value = conversion_entry
                # Кілька різних статусів, що ведуть до того самого значення
                # (напр. "БЗ"/"полон"/"Інт" - усі в "100_БЗ"), фарбуються за
                # новим значенням, щоб однакова клітинка завжди мала той самий
                # колір. Статуси з єдиним джерелом (напр. "ППД" -> 10) - за
                # колишнім статусом (_CONVERSION_VALUE_COUNTS).
                color_key = new_value if _CONVERSION_VALUE_COUNTS[new_value] > 1 else original_status
                cell.value = new_value
                _paint(cell, color_key)
                _remember_status_behind_number(timesheet, row_idx, date_value, original_status, new_value)
                has_any_value = True
                continue

            values = _resolve_value_priority_conflict(payment_values.get((normalized, date_value)) or [])
            unique_values = set(values)

            # Порожня клітинка й одне узгоджене значення information_unit, яке
            # саме є ключем PAYMENT_STATUS_CONVERSIONS (напр. "ППД"/"БПШП"/
            # "ВПБП"/"БЗ"/"полон"/"Інт"), конвертується так само, як коли це
            # статус самої клітинки: "ВПБП" має стати "100_ВПБП" (як підтвердив
            # користувач), ніби Табель сам написав "ВПБП". Лише для порожньої
            # клітинки: "РВЗ" ніколи не буває ключем PAYMENT_STATUS_CONVERSIONS.
            if normalized_status is None and len(unique_values) == 1:
                conversion_entry = _NORMALIZED_PAYMENT_STATUS_CONVERSIONS.get(_normalize_status_text(values[0]))
                if conversion_entry is not None:
                    original_status, new_value = conversion_entry
                    color_key = new_value if _CONVERSION_VALUE_COUNTS[new_value] > 1 else original_status
                    cell.value = new_value
                    _paint(cell, color_key)
                    _remember_status_behind_number(timesheet, row_idx, date_value, original_status, new_value)
                    has_any_value = True
                    continue

            # Порожня клітинка шукає мапу замінників за DEFAULT_STATUS ("РВЗ"):
            # власного запису для None немає, а порожньо тут означає те саме "ще
            # не визначено". Без цього текстові замінники ("НОВ" -> "НОВ" тощо)
            # не спрацьовували б для новоприбулих, і кожна така дата хибно
            # падала б у "не єдине числове значення".
            override_map = _NORMALIZED_PAYMENT_STATUS_OVERRIDES.get(normalized_status or DEFAULT_STATUS)
            if override_map:
                # Порівнюються нормалізовані форми: два файли з "БПШП" і "бпшп" -
                # це та сама відповідь, а не суперечність.
                normalized_values = {_normalize_status_text(value) for value in values}
                if len(normalized_values) == 1:
                    override_entry = override_map.get(next(iter(normalized_values)))
                    if override_entry is not None:
                        original_value, new_status = override_entry
                        cell.value = new_status
                        _paint(cell, new_status)
                        _remember_status_behind_number(timesheet, row_idx, date_value, original_value, new_status)
                        has_any_value = True
                        continue

            # Порожня клітинка йде тим самим шляхом підстановки категорії, що й
            # DEFAULT_STATUS ("РВЗ"): новоприбулу людину додано до Табеля без
            # жодного статусу (у рапортах її немає, лише в information_unit),
            # тож information_unit може заповнити її так само.
            if normalized_status not in (DEFAULT_STATUS, None):
                continue

            date_text = date_value.strftime("%d.%m.%Y")
            if not values:
                # "Підрозділ ще не подав дані" має сенс лише для "РВЗ": людина
                # вже на обліку, просто категорії на цю дату ще немає. Для
                # порожньої клітинки (новоприбулий без жодних даних) запис на
                # кожну з ~30 дат не пишеться - замість нього один підсумковий
                # запис наприкінці, якщо жодна дата так нічого й не отримала.
                if normalized_status == DEFAULT_STATUS:
                    unresolved.append({
                        "reason": f"Немає категорії виплати для \"{pib_raw}\" на {date_text} - підрозділ ще не подав дані.",
                        "pib_raw": pib_raw, "report_date": date_value,
                    })
                continue

            if len(unique_values) > 1 or not _is_number(values[0]):
                if normalized_status is None:
                    # Порожня клітинка й одне узгоджене значення information_unit,
                    # яке не число і не підійшло під жодну мапу вище (напр.
                    # "ВП"/"СЗЧ" - і так зрозумілі статуси, просто без
                    # конвертації). Для новоприбулого без іншого джерела воно
                    # записується напряму й мовчки (рішення користувача).
                    # Справжня суперечність (два різні значення) лишає клітинку
                    # порожньою - теж мовчки, як і "немає даних" вище.
                    if len(unique_values) == 1:
                        cell.value = values[0]
                        has_any_value = True
                        _paint(cell, values[0])
                    continue

                unresolved.append({
                    "reason": (
                        f"\"{pib_raw}\" {date_text}: підрозділ подає {sorted(map(str, unique_values))} - "
                        f"не єдине числове значення категорії виплати, статус \"{normalized_status}\" лишено як є."
                    ),
                    "pib_raw": pib_raw, "report_date": date_value,
                })
                continue

            cell.value = values[0]
            has_any_value = True
            _paint(cell, values[0])

        # Людину щойно додали до Табеля без базового статусу, жоден рапорт її
        # не згадує, і information_unit теж не дав жодного значення: будувати
        # нема на чому, тож один явний запис на ручну перевірку. Якщо
        # information_unit дав хоч щось, запису немає - користувач просив, щоб
        # таких помилок у txt не було, коли значення підставилось.
        if not has_any_value:
            unresolved.append({
                "reason": (
                    f"\"{pib_raw}\": усі колонки-дати порожні - немає базового статусу в {timesheet.file_path} "
                    "і жодних даних від information_unit. Впишіть базовий статус уручну."
                ),
                "pib_raw": pib_raw, "report_date": None,
            })

    return unresolved
