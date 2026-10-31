import glob
import os
import re
from datetime import date, datetime
from typing import Any, NamedTuple
from zipfile import BadZipFile

import xlrd
from openpyxl.drawing.text import Font as _DrawingFont
from openpyxl.styles.fonts import Font as _StyleFont
from xlrd import xldate

from constants import DEFAULT_STATUS, INFORMATION_UNIT_SHEET_NAME, PIB_COLUMN_NAME
from content.excel_compat import load_workbook
from content.oblik_timesheet import normalize_name

# Деякі файли цієї теки мають вбудований малюнок чи штамп, у текстовому стилі
# якого атрибут "panose" порожній. openpyxl (Font.panose - HexBinary, шаблон
# "[0-9a-fA-F]+$", хоча б один символ) на цьому піднімає ValueError ("Unable
# to read workbook: ... invalid XML") ще до читання аркушів ТАБЕЛЬ/БЧС, і
# файл пропускався цілком, хоча малюнки нам не потрібні. Тож шаблон
# послаблено точково, у вже створеного дескриптора: змінювати
# HexBinary.pattern у класі марно, бо test_pattern компілюється один раз під
# час імпорту openpyxl.drawing.text.
_DrawingFont.panose.pattern = "[0-9a-fA-F]*$"
_DrawingFont.panose.test_pattern = re.compile(_DrawingFont.panose.pattern, re.VERBOSE)

# Схожий випадок з файлом БЧС одного з підрозділів: styles.xml має
# <font><family val="34"/>...</font> (Times New Roman), а openpyxl дозволяє
# Font.family лише 0..14 (офіційний діапазон OOXML) і падає з тією самою
# помилкою. Excel такий файл відкриває без нарікань, а людина з нього хибно
# потрапляла в "Немає категорії виплати ... - підрозділ ще не подав дані".
# Межу так само розширено у вже створеного дескриптора. Стаби openpyxl цього
# дескриптора не описують - звідси pyright: ignore.
_StyleFont.family.max = 255  # pyright: ignore[reportAttributeAccessIssue]

# .xls (старий двійковий формат) openpyxl не читає, тож для нього є окрема
# обгортка над xlrd (_load_workbook). Один файл цієї теки справді був .xls і
# мовчки пропускався.
_XLS_EXTENSION = ".xls"
_SUPPORTED_EXTENSIONS = (".xlsx", ".xlsm", _XLS_EXTENSION)

# Реальні .xls цієї теки - у cp1251 (кирилиця Windows), а xlrd визначає
# кодування неправильно, і без явного cp1251 назви аркушів і текст читаються
# як абракадабра.
_XLS_ENCODING = "cp1251"


def _is_office_lock_file(path):
    """Lock-файл Office ("~$Назва.xlsx"), що існує, поки книга відкрита в
    Excel. Розширення в нього те саме, тож без цієї перевірки він ішов би на
    читання і валив його, хоча це не книга, а маленька позначка."""
    return os.path.basename(path).startswith("~$")


class _XlsCell:
    """Клітинка з інтерфейсом openpyxl (.value) навколо вже перетвореного
    значення xlrd (_xls_cell_value)."""

    __slots__ = ("value",)

    def __init__(self, value):
        self.value = value


def _xls_cell_value(sheet, book, row, col):
    """Значення клітинки xlrd у тому вигляді, що й в openpyxl: дата - datetime
    (xlrd дає лише серійне число, а _tabel_date_row_and_columns/_parse_cell_date
    чекають дату), ціле число - int, а не float (інакше категорія виплати
    "100" показувалась би як "100.0"), порожня клітинка - None (не "" і не
    0.0)."""
    cell = sheet.cell(row, col)
    if cell.ctype == xlrd.XL_CELL_DATE:
        return xldate.xldate_as_datetime(cell.value, book.datemode)
    if cell.ctype == xlrd.XL_CELL_NUMBER:
        value = cell.value
        return int(value) if value == int(value) else value
    if cell.ctype == xlrd.XL_CELL_BOOLEAN:
        return bool(cell.value)
    if cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
        return None
    return cell.value or None


class _XlsSheetAdapter:
    """Аркуш xlrd (0-індексований sheet.cell(row, col)) з інтерфейсом аркуша
    openpyxl (1-індексований ws.cell(row=, column=).value, max_row,
    max_column, title), щоб усі читачі, написані під openpyxl, без змін
    працювали і з .xls."""

    def __init__(self, sheet, book):
        self._sheet = sheet
        self._book = book
        self.title = sheet.name
        self.max_row = sheet.nrows
        self.max_column = sheet.ncols

    def cell(self, row, column):
        r, c = row - 1, column - 1
        if r < 0 or c < 0 or r >= self.max_row or c >= self.max_column:
            return _XlsCell(None)
        return _XlsCell(_xls_cell_value(self._sheet, self._book, r, c))


class _XlsWorkbookAdapter:
    """Книга xlrd з інтерфейсом книги openpyxl - рівно стільки, скільки
    потрібно читачам: sheetnames, wb[title] і active."""

    def __init__(self, path):
        book = xlrd.open_workbook(path, encoding_override=_XLS_ENCODING)
        self._book = book
        self.sheetnames = book.sheet_names()

    def __getitem__(self, title):
        return _XlsSheetAdapter(self._book.sheet_by_name(title), self._book)

    @property
    def active(self):
        """Аркуш, вибраний у файлі (як wb.active в openpyxl), інакше перший.
        sheet_selected xlrd додає до аркуша під час читання - у стабах його немає."""
        sheets = self._book.sheets()
        selected = next((sheet for sheet in sheets if getattr(sheet, "sheet_selected", 0)), sheets[0])
        return _XlsSheetAdapter(selected, self._book)


def _load_workbook(path):
    """openpyxl (data_only=True) для .xlsx/.xlsm і _XlsWorkbookAdapter (xlrd)
    для .xls - з однаковим інтерфейсом (sheetnames, wb[title],
    ws.cell(row=, column=).value, ws.max_row/max_column/title), тож усі читачі
    працюють однаково з будь-яким форматом."""
    if os.path.splitext(path)[1].lower() == _XLS_EXTENSION:
        return _XlsWorkbookAdapter(path)
    return load_workbook(path, data_only=True)


def _safe_load_workbook(path):
    """_load_workbook(path), але пошкоджений файл не зупиняє весь прогін:
    колись одна книга з пошкодженим XML (той самий "panose") не давала
    прочитати жодного файлу теки. Порожній "panose" тепер виправлено ще на
    імпорті (див. вище), а цей try/except лишається для будь-яких інших
    помилок розбору, яких наперед не передбачиш. ValueError - у нього openpyxl
    загортає будь-яку помилку XML під час читання аркушів; KeyError і
    BadZipFile - пошкоджений архів .xlsx чи відсутня частина в ньому.

    Повертає (wb, None) або (None, причина). Причину викликач додає в skipped
    і читає далі."""
    try:
        return _load_workbook(path), None
    except (ValueError, KeyError, BadZipFile) as e:
        return None, f"Не вдалось прочитати файл {path}: {e}"

# Рядок з датами на аркуші ТАБЕЛЬ у реальних файлах різний. Здебільшого в
# рядку 1 - груповий заголовок "ДАТА", а самі дати - в рядку 2 (дані з рядка
# 3), але бувають файли з датами одразу в рядку 1, після ПІБ (дані з рядка 2).
# Тому рядок з датами шукається (_tabel_date_row_and_columns), а не
# фіксується: інакше другий варіант мовчки не читався б.
_TABEL_DATE_ROW_CANDIDATES = (1, 2)

# Аркуш "БЧС" - друге джерело в кожному файлі: знімок одного дня, а не
# колонка на кожен день, як у "ТАБЕЛЬ". Рядок 1 - назва підрозділу (можливо, з
# датою в довільному форматі), рядок 2 - заголовки ("ПРІЗВИЩЕ ім'я по
# батькові" замість "ПІБ", "БЧС" - статус, "Виплата додактової винагороди" -
# категорія виплати), дані - з рядка 3. У частині реальних файлів "ТАБЕЛЬ"
# містить лише застарілі дати (напр. увесь липень замість серпня), і тоді
# "БЧС" - єдине джерело актуальних даних.
_BCHS_SHEET_NAME = "бчс"
_BCHS_HEADER_ROW = 2
_DATE_IN_TEXT_RE = re.compile(r"(\d{2})\.(\d{2})\.(\d{2,4})")


def _find_sheet(wb, name):
    return next((wb[title] for title in wb.sheetnames if title.strip().lower() == name), None)


# Дата в рядку дат ТАБЕЛЬ буває текстом з пропущеною крапкою між місяцем і
# роком: "06.092026" замість "06.09.2026". strptime нижче такого не
# розпізнає, і вся колонка мовчки не вважалась датою, хоча виплати в ній є.
# Шаблон прив'язаний до всього рядка (^...$) і ловить лише цю форму.
_MISSING_DOT_DATE_RE = re.compile(r"^(?P<day>\d{2})\.(?P<month>\d{2})(?P<year>\d{4})$")


def _parse_cell_date(value):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str):
        stripped = value.strip()
        for fmt in ("%d.%m.%Y", "%d.%m.%y"):
            try:
                return datetime.strptime(stripped, fmt).date()
            except ValueError:
                continue
        match = _MISSING_DOT_DATE_RE.match(stripped)
        if match:
            try:
                return date(int(match.group("year")), int(match.group("month")), int(match.group("day")))
            except ValueError:
                return None
    return None


def _pib_column(ws):
    header = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column + 1)]
    return next((idx for idx, value in enumerate(header, start=1) if isinstance(value, str) and value.strip().upper() == PIB_COLUMN_NAME), None)


def _date_columns(ws, row):
    columns = ((c, _parse_cell_date(ws.cell(row=row, column=c).value)) for c in range(1, ws.max_column + 1))
    return [(c, parsed) for c, parsed in columns if parsed is not None]


def _tabel_date_row_and_columns(ws):
    """(date_row, date_cols): перший рядок з _TABEL_DATE_ROW_CANDIDATES (спершу
    1, потім 2), де є хоча б одна дата, і колонки-дати в ньому; (None, []) -
    дат немає. У файлах з груповим заголовком "ДАТА" рядок 1 дат не має, тож
    для них сам собою береться рядок 2."""
    for row in _TABEL_DATE_ROW_CANDIDATES:
        date_cols = _date_columns(ws, row)
        if date_cols:
            return row, date_cols
    return None, []


def _search_date(text):
    """Перша дата (DD.MM.YY чи DD.MM.YYYY) будь-де в тексті - назві файлу чи
    заголовку аркуша. На відміну від _parse_cell_date, текст може містити ще
    щось, крім дати."""
    match = _DATE_IN_TEXT_RE.search(text or "")
    if not match:
        return None
    day, month, year = match.groups()
    if len(year) == 2:
        year = f"20{year}"
    try:
        return date(int(year), int(month), int(day))
    except ValueError:
        return None


def _extract_bchs_date(path, ws):
    """Дата аркуша БЧС: спершу з назви файлу (як у рапортів), а якщо там її
    немає - із заголовка аркуша (рядок 1). Жодне з джерел не гарантоване:
    бувають назви без року ("...08.08_1_3.xlsx") і заголовки без дати."""
    return _search_date(os.path.basename(path)) or _search_date(str(ws.cell(row=1, column=1).value or ""))


def _bchs_columns(ws):
    """(pib_col, pidrozdil_col, posada_col, zvannya_col, status_col,
    payment_col) з рядка _BCHS_HEADER_ROW. Позиції в різних файлах різні, тож
    колонки шукаються за текстом заголовка: ПІБ - за "прізвищ" (тут заголовок
    "ПРІЗВИЩЕ ім'я по батькові"), Підрозділ/Посада/Звання і статус ("БЧС") -
    за точним збігом, виплата - за "виплат" (у реальних файлах - "Виплата
    додактової винагороди", саме з цією опискою)."""
    pib_col = pidrozdil_col = posada_col = zvannya_col = status_col = payment_col = None
    for c in range(1, ws.max_column + 1):
        value = ws.cell(row=_BCHS_HEADER_ROW, column=c).value
        if not isinstance(value, str):
            continue
        lowered = value.strip().lower()
        if "прізвищ" in lowered:
            pib_col = c
        elif lowered == "підрозділ":
            pidrozdil_col = c
        elif lowered == "посада":
            posada_col = c
        elif lowered == "звання":
            zvannya_col = c
        elif lowered == "бчс":
            status_col = c
        elif "виплат" in lowered:
            payment_col = c
    return pib_col, pidrozdil_col, posada_col, zvannya_col, status_col, payment_col


def _clean_cell_value(value):
    """Значення клітинки ТАБЕЛЬ/БЧС, готове до порівняння й збереження:
    - текст обрізається з країв: зайвий пробіл ("БПШП ") мовчки ламав збіг з
      PAYMENT_STATUS_CONVERSIONS/PAYMENT_STATUS_OVERRIDES, і статус лишався
      без змін без жодного повідомлення;
    - рядок лише з цифр (можливо, з "-" на початку) стає числом: категорія
      виплати "100", збережена як текст (формат "Текст" чи копіювання з іншої
      системи), інакше вважалась би нечисловою.
    Числа, дати й None - без змін."""
    if isinstance(value, str):
        value = value.strip()
        if value.lstrip("-").isdigit():
            return int(value)
    return value


def _tabel_label_columns(ws):
    """(pidrozdil_col, posada_col, zvannya_col) з рядка 1 аркуша ТАБЕЛЬ - за
    точним заголовком, як і _pib_column. Кожної з цих колонок у файлі може не
    бути (Підрозділ є лише в деяких файлах) - тоді None."""
    pidrozdil_col = posada_col = zvannya_col = None
    for c in range(1, ws.max_column + 1):
        value = ws.cell(row=1, column=c).value
        if not isinstance(value, str):
            continue
        upper = value.strip().upper()
        if upper == "ПІДРОЗДІЛ":
            pidrozdil_col = c
        elif upper == "ПОСАДА":
            posada_col = c
        elif upper == "ЗВАННЯ":
            zvannya_col = c
    return pidrozdil_col, posada_col, zvannya_col


class SnapshotEntry(NamedTuple):
    """Одне значення файлу information_unit для людини й дати разом із тим,
    звідки воно (назва файлу й дані людини в ньому) - для звіту розбіжностей."""
    file_name: str
    pib_raw: Any
    pidrozdil: Any
    posada: Any
    zvannya: Any
    value: Any


def _cell_value(ws, row, col):
    return ws.cell(row=row, column=col).value if col else None


def _read_tabel_sheet(wb, path, snapshot):
    """Аркуш ТАБЕЛЬ (колонка на кожен день) -> snapshot. Рядок із датами (а
    отже й перший рядок даних під ним) шукається за вмістом
    (_tabel_date_row_and_columns), бо реальні файли різняться. Повертає None,
    якщо прочитано, або причину, чому ні."""
    ws = _find_sheet(wb, INFORMATION_UNIT_SHEET_NAME)
    if ws is None:
        return f"У файлі {path} не знайдено аркуш ТАБЕЛЬ."

    pib_col = _pib_column(ws)
    date_row, date_cols = _tabel_date_row_and_columns(ws)
    if pib_col is None or date_row is None:
        return f"У файлі {path} (аркуш {ws.title}) не знайдено {'колонку ПІБ' if pib_col is None else 'жодної колонки-дати'}."

    file_name = os.path.basename(path)
    pidrozdil_col, posada_col, zvannya_col = _tabel_label_columns(ws)
    for row in range(date_row + 1, ws.max_row + 1):
        pib_raw = ws.cell(row=row, column=pib_col).value
        if not pib_raw:
            continue
        normalized = normalize_name(pib_raw)
        person = (file_name, pib_raw, _cell_value(ws, row, pidrozdil_col), _cell_value(ws, row, posada_col),
                  _cell_value(ws, row, zvannya_col))
        for col, day in date_cols:
            value = _clean_cell_value(ws.cell(row=row, column=col).value)
            if value not in (None, ""):
                snapshot.setdefault((normalized, day), []).append(SnapshotEntry(*person, value))
    return None


def _read_bchs_sheet(wb, path, snapshot):
    """Аркуш БЧС (знімок одного дня) -> snapshot. Статус "БЧС", відмінний від
    DEFAULT_STATUS ("РВЗ"), записується як є (як і текст у ТАБЕЛЬ:
    apply_payment_values розуміє це як "підрозділ каже інше"), інакше -
    число з колонки виплати, якщо воно є. Повертає None, якщо прочитано, або
    причину, чому ні."""
    ws = _find_sheet(wb, _BCHS_SHEET_NAME)
    if ws is None:
        return f"У файлі {path} не знайдено аркуш БЧС."

    day = _extract_bchs_date(path, ws)
    if day is None:
        return f"У файлі {path} (аркуш {ws.title}) не вдалось визначити дату (ні з назви файлу, ні із заголовка аркуша)."

    pib_col, pidrozdil_col, posada_col, zvannya_col, status_col, payment_col = _bchs_columns(ws)
    if pib_col is None:
        return f"У файлі {path} (аркуш {ws.title}) не знайдено колонку ПІБ."

    file_name = os.path.basename(path)
    for row in range(_BCHS_HEADER_ROW + 1, ws.max_row + 1):
        pib_raw = ws.cell(row=row, column=pib_col).value
        if not pib_raw:
            continue
        status = _clean_cell_value(_cell_value(ws, row, status_col))
        if isinstance(status, str) and status and status.upper() != DEFAULT_STATUS:
            value = status
        else:
            value = _clean_cell_value(_cell_value(ws, row, payment_col))
            if value in (None, ""):
                continue
        entry = SnapshotEntry(file_name, pib_raw, _cell_value(ws, row, pidrozdil_col), _cell_value(ws, row, posada_col),
                              _cell_value(ws, row, zvannya_col), value)
        snapshot.setdefault((normalize_name(pib_raw), day), []).append(entry)
    return None


def _skipped(reason):
    return {"reason": reason, "report_date": None}


def read_information_unit_snapshot(directory):
    """Читає всі файли directory. У кожному файлі два можливі джерела (назва
    аркуша - без урахування регістру): ТАБЕЛЬ (колонка на кожен день) і БЧС
    (знімок одного дня, дата - з назви файлу чи заголовка аркуша). ТАБЕЛЬ
    головний: якщо його прочитано, БЧС того самого файлу не читається - за
    вказівкою користувача, бо два аркуші одного файлу можуть суперечити один
    одному (напр. описка в ПІБ лише на одному), і читання обох лише множить
    розбіжності. БЧС - запасне джерело, коли ТАБЕЛЬ прочитати не вдалось.

    Повертає (snapshot, skipped):
    - snapshot - {(normalize_name(ПІБ), date): [SnapshotEntry, ...]}. Список,
      бо людина може бути в кількох файлах одразу (напр. придана іншому
      підрозділу) - що робити із суперечливими значеннями, вирішує
      apply_payment_values (content/oblik_timesheet.py).
    - skipped - [{"reason", "report_date"}] для файлів, з яких не вдалось
      прочитати ні ТАБЕЛЬ, ні БЧС (непідтримуваний формат, пошкоджений файл,
      немає потрібних аркушів чи колонок)."""
    snapshot = {}
    skipped = []
    for path in sorted(glob.glob(os.path.join(directory, "*"))):
        if _is_office_lock_file(path):
            continue
        if os.path.splitext(path)[1].lower() not in _SUPPORTED_EXTENSIONS:
            skipped.append(_skipped(f"Формат файлу не підтримується: {path}."))
            continue

        wb, load_reason = _safe_load_workbook(path)
        if wb is None:
            skipped.append(_skipped(load_reason))
            continue
        tabel_reason = _read_tabel_sheet(wb, path, snapshot)
        if tabel_reason is None:
            continue
        if _read_bchs_sheet(wb, path, snapshot) is not None:
            skipped.append(_skipped(tabel_reason))
    return snapshot, skipped


def snapshot_values(snapshot):
    """snapshot без атрибуції джерела: {(normalize_name(ПІБ), date): [значення,
    ...]} - у такому вигляді значення бере apply_payment_values. Значення - як
    у файлі: число - категорія виплати, текст - статус на кшталт "ВП"/"ШП"."""
    return {key: [entry.value for entry in entries] for key, entries in snapshot.items()}


def read_payment_values(directory):
    """(snapshot_values(snapshot), skipped) - див. read_information_unit_snapshot."""
    snapshot, skipped = read_information_unit_snapshot(directory)
    return snapshot_values(snapshot), skipped
