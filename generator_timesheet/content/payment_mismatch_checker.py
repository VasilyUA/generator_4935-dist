from datetime import datetime

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from constants import (
    DATE_NUMBER_FORMAT,
    DEFAULT_STATUS,
    PAYMENT_ABSENT_FROM_FILES_EXCLUDE_NAMES,
    PAYMENT_ABSENT_FROM_FILES_EXEMPT_STATUSES,
    PAYMENT_DUTY_CATEGORIES,
    PAYMENT_MISMATCH_COMPARABLE_STATUSES,
    PAYMENT_MISMATCH_LABEL_COLUMNS,
    PAYMENT_STATUS_CONVERSIONS,
    PAYMENT_STATUS_OVERRIDES,
    PAYMENT_STATUS_PRIORITY,
    PIB_COLUMN_NAME,
)
from content.bchs_subdivisions import file_subdivision
from content.commander_message import (
    NOT_IN_ROSTER_ISSUE,
    PAYMENTS_SOURCE,
    compose_message,
    date_issues,
    update_message_file,
)
from content.excel_compat import save_workbook, workbook_with_sheet
from content.oblik_timesheet import normalize_name, spelled_differently
from content.report_pib_column import PIB_ALIGNMENT, pib_column_width

# Підпис рядка "офіційного" джерела у звіті. На відміну від рядка з файлом
# information_unit, одного файлу-джерела тут немає: значення обчислено з усіх
# щоденних рапортів і вже застосованих категорій виплати.
_REPORT_SOURCE_LABEL = "ОБЛІК для виплат (за рапортами)"

# Значення сторони ОБЛІК для людини, якої в ОБЛІК для виплат.xlsx немає
# зовсім. Ставиться в кожну порівнювану колонку (ПІБ/Підрозділ/Посада/Звання/
# Статус), а не лише в дати: порожня клітинка без пояснення незрозуміла, а
# напис пояснює, чому там пусто.
_NOT_IN_ROSTER_LABEL = "Не знайдено в ОБЛІК.xlsx"

# Те саме для протилежного випадку - людина є в ОБЛІК, але жоден файл
# information_unit її не подає (_absent_from_files_records): і назва "файлу",
# і файлова сторона кожної колонки.
_NOT_IN_FILES_LABEL = "Немає в жодному файлі information_unit"

# Заповнювач для будь-якого іншого порожнього значення (напр. у Табелі немає
# колонки ПІДРОЗДІЛ чи файл не подає Посаду), щоб серед кольорових клітинок
# не було незрозуміло білих. Порівняння (_normalize_for_compare) працює з
# оригінальним значенням (None), заповнювач - лише для показу.
_EMPTY_PLACEHOLDER = "—"

# Стиль звіту: Times New Roman 14, зелена заливка там, де Табель і файл
# information_unit збігаються, червона - де ні (ті самі відтінки, що й
# STATUS_COLORS). Межі - на всіх клітинках, щоб таблицю легко було читати й
# без кольору. Без переносу тексту, натомість широкі колонки: перенесений
# текст у рядках висотою 20 накладався на сусідній рядок.
_FONT = Font(name="Times New Roman", size=14)
_MATCH_FILL = PatternFill(start_color="FFC6EFCE", end_color="FFC6EFCE", fill_type="solid")
_MISMATCH_FILL = PatternFill(start_color="FFFFC7CE", end_color="FFFFC7CE", fill_type="solid")
_CENTER = Alignment(horizontal="center", vertical="center", wrap_text=False)
_THIN_SIDE = Side(style="thin")
_BORDER = Border(left=_THIN_SIDE, right=_THIN_SIDE, top=_THIN_SIDE, bottom=_THIN_SIDE)

# Заголовок - такий самий, як у аркуша "Табель" в resources/ОБЛІК.xlsx
# (Bahnschrift Light SemiCondensed 14, жирний білий на темно-бірюзовому
# FF035C6B): користувач просив "вигляд як в ОБЛІК файлі".
_HEADER_FONT = Font(name="Bahnschrift Light SemiCondensed", size=14, bold=True, color="FFFFFFFF")
_HEADER_FILL = PatternFill(start_color="FF035C6B", end_color="FF035C6B", fill_type="solid")
_HEADER_ALIGNMENT = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _normalize_for_compare(value):
    """Зводить None і рядки до одного вигляду: пробіли згорнуті, краї обрізані,
    верхній регістр. ПІБ з різних джерел буває з подвійними пробілами
    всередині, і без цього така косметична різниця вважалась би
    розбіжністю."""
    if value is None:
        return ""
    if isinstance(value, str):
        return " ".join(value.split()).upper()
    return value


def _display_value(value):
    """Значення для клітинки - як є, лише None чи "" стає _EMPTY_PLACEHOLDER,
    щоб не було незрозуміло білих клітинок. На порівняння не впливає: воно
    працює з оригінальним значенням."""
    if value is None or value == "":
        return _EMPTY_PLACEHOLDER
    return value


def _statuses_behind_numbers(timesheet, row_idx, dates):
    """{дата: статус} - дати з dates, де число в ОБЛІК лише замінило статус
    (Timesheet.statuses_behind_numbers, ППД -> 10)."""
    behind = timesheet.statuses_behind_numbers
    return {date_value: behind[(row_idx, date_value)] for date_value in dates if (row_idx, date_value) in behind}


def _surname_firstname_key(normalized):
    """Перші два слова нормалізованого ПІБ (Прізвище Ім'я) без по батькові.
    Це точний збіг, а не нечіткий: решта навмисно ігнорується. Реальний
    випадок: один файл information_unit написав по батькові з опискою в одну-
    дві літери, normalize_name давав інший ключ, і людину не вдавалось
    зіставити з Табелем."""
    parts = normalized.split(" ")
    return " ".join(parts[:2]) if len(parts) >= 2 else normalized


def _candidate_roster_person(person_rows, file_normalized):
    """row_idx єдиного рядка Табеля з тими самими Прізвищем та Ім'ям, що й у
    file_normalized (_surname_firstname_key), коли file_normalized сам не є
    ключем у person_rows. Якщо кандидатів 0 чи кілька (однофамільці, тезки),
    повертає None: краще показати "не знайдено", ніж вгадати не ту людину."""
    prefix = _surname_firstname_key(file_normalized)
    candidates = [row_idx for normalized, row_idx in person_rows.items() if _surname_firstname_key(normalized) == prefix]
    return candidates[0] if len(candidates) == 1 else None


def _missing_from_roster_records(timesheet, person_rows, snapshot):
    """Записи для людей зі snapshot (information_unit), чий ПІБ не збігся з
    жодним рядком Табеля: людини в ОБЛІК для виплат.xlsx немає взагалі або
    ПІБ написано інакше (зазвичай - описка в по батькові). Без цього вони
    мовчки зникали б зі звіту, бо основний прохід find_mismatches іде по
    Табелю, а не по snapshot.

    Якщо в Табелі рівно один кандидат з тими самими Прізвищем та Ім'ям
    (_candidate_roster_person), сторона ОБЛІК показує його справжні ПІБ,
    Підрозділ, Посаду, Звання й поточний статус - щоб різницю в написанні було
    видно поруч. Інакше (0 чи кілька кандидатів) у кожній колонці сторони
    ОБЛІК - _NOT_IN_ROSTER_LABEL, а не порожня клітинка без пояснення."""
    pib_col = timesheet.label_columns[PIB_COLUMN_NAME]
    pidrozdil_col = timesheet.label_columns.get("ПІДРОЗДІЛ")
    posada_col = timesheet.label_columns.get("ПОСАДА")
    zvannya_col = timesheet.label_columns.get("ЗВАННЯ")
    date_col_by_date = {d: idx for idx, d in timesheet.date_columns}

    candidate_by_normalized = {}
    by_person_file = {}
    for (normalized, date_value), entries in snapshot.items():
        if normalized in person_rows:
            continue
        if normalized not in candidate_by_normalized:
            candidate_by_normalized[normalized] = _candidate_roster_person(person_rows, normalized)
        candidate_row = candidate_by_normalized[normalized]

        if candidate_row is not None and date_value in date_col_by_date:
            roster_value = timesheet.ws.cell(row=candidate_row, column=date_col_by_date[date_value]).value
        else:
            roster_value = _NOT_IN_ROSTER_LABEL

        for file_name, pib_raw, file_pidrozdil, file_posada, file_zvannya, value in entries:
            entry = by_person_file.setdefault(
                (normalized, file_name),
                {
                    "pib_raw": pib_raw, "pidrozdil": file_pidrozdil, "posada": file_posada, "zvannya": file_zvannya,
                    "dates": {}, "candidate_row": candidate_row,
                },
            )
            entry["dates"][date_value] = (roster_value, value)

    # Файли з однаковим вмістом зводяться до одного
    # (_dedupe_files_with_identical_content), окремо для кожної людини:
    # by_person_file змішує людей, тож спершу групуємо за normalized, щоб не
    # сплутати двох різних людей, чиї дані випадково збіглися.
    by_normalized = {}
    for (normalized, file_name), info in by_person_file.items():
        by_normalized.setdefault(normalized, {})[file_name] = info
    deduped_items = [
        ((normalized, file_name), info)
        for normalized, files in by_normalized.items()
        for file_name, info in _dedupe_files_with_identical_content(files).items()
    ]

    records = []
    for (_normalized, file_name), info in sorted(deduped_items, key=lambda item: (item[0][1], item[1]["pib_raw"] or "")):
        candidate_row = info["candidate_row"]
        if candidate_row is not None:
            roster_pib_raw = timesheet.ws.cell(row=candidate_row, column=pib_col).value
            roster_pidrozdil = timesheet.ws.cell(row=candidate_row, column=pidrozdil_col).value if pidrozdil_col else None
            roster_posada = timesheet.ws.cell(row=candidate_row, column=posada_col).value if posada_col else None
            roster_zvannya = timesheet.ws.cell(row=candidate_row, column=zvannya_col).value if zvannya_col else None
        else:
            roster_pib_raw = roster_pidrozdil = roster_posada = roster_zvannya = _NOT_IN_ROSTER_LABEL

        records.append({
            "file_name": file_name,
            "roster_pib_raw": roster_pib_raw,
            "roster_pidrozdil": roster_pidrozdil,
            "roster_posada": roster_posada,
            "roster_zvannya": roster_zvannya,
            "file_pib_raw": info["pib_raw"],
            "file_pidrozdil": info["pidrozdil"],
            "file_posada": info["posada"],
            "file_zvannya": info["zvannya"],
            "dates": info["dates"],
            "roster_statuses": _statuses_behind_numbers(timesheet, candidate_row, info["dates"]),
        })
    return records


def _absent_from_files_records(timesheet, person_rows, snapshot):
    """Дзеркало _missing_from_roster_records: люди, які є в ОБЛІК, але жоден
    файл information_unit їх не подає (людину не внесли до файлу БЧС її
    підрозділу). Не потрапляють: люди, чий останній статус - з
    PAYMENT_ABSENT_FROM_FILES_EXEMPT_STATUSES (вже не на обліку підрозділу:
    БЗ, переведені, загиблі...); люди з PAYMENT_ABSENT_FROM_FILES_EXCLUDE_NAMES
    (resources/data.json - напр. довготривале лікування); кандидати
    _candidate_roster_person для ПІБ з файлів, яких немає в ОБЛІК (їх уже
    показано тим записом як інше написання); і всі, якщо файлів немає
    зовсім. Файлова сторона - _NOT_IN_FILES_LABEL; дати - ті, що подають
    файли і на які в ОБЛІК є статус."""
    if not snapshot:
        return []
    reported = {normalized for normalized, _date_value in snapshot}
    file_dates = {date_value for _normalized, date_value in snapshot}
    spelled_differently = {
        _candidate_roster_person(person_rows, normalized) for normalized in reported if normalized not in person_rows
    }
    exempt = {_normalize_for_compare(status) for status in PAYMENT_ABSENT_FROM_FILES_EXEMPT_STATUSES}
    excluded_names = {normalize_name(name) for name in PAYMENT_ABSENT_FROM_FILES_EXCLUDE_NAMES}
    pib_col = timesheet.label_columns[PIB_COLUMN_NAME]
    label_cols = {name: timesheet.label_columns.get(name) for name in ("ПІДРОЗДІЛ", "ПОСАДА", "ЗВАННЯ")}

    records = []
    for normalized, row_idx in person_rows.items():
        if normalized in reported or normalized in excluded_names or row_idx in spelled_differently:
            continue
        values = [(d, timesheet.ws.cell(row=row_idx, column=c).value) for c, d in timesheet.date_columns]
        filled = [(d, value) for d, value in values if value is not None]
        if filled and _normalize_for_compare(filled[-1][1]) in exempt:
            continue
        roster = {name: timesheet.ws.cell(row=row_idx, column=col).value if col else None for name, col in label_cols.items()}
        records.append({
            "file_name": _NOT_IN_FILES_LABEL,
            "roster_pib_raw": timesheet.ws.cell(row=row_idx, column=pib_col).value,
            "roster_pidrozdil": roster["ПІДРОЗДІЛ"],
            "roster_posada": roster["ПОСАДА"],
            "roster_zvannya": roster["ЗВАННЯ"],
            "file_pib_raw": _NOT_IN_FILES_LABEL,
            "file_pidrozdil": _NOT_IN_FILES_LABEL,
            "file_posada": _NOT_IN_FILES_LABEL,
            "file_zvannya": _NOT_IN_FILES_LABEL,
            "dates": {d: (value, None) for d, value in filled if d in file_dates},
            "roster_statuses": _statuses_behind_numbers(timesheet, row_idx, file_dates),
        })
    return records


def _is_number(value):
    return isinstance(value, (int, float)) and not isinstance(value, bool)


_COMPARABLE_STATUSES = frozenset(_normalize_for_compare(status) for status in PAYMENT_MISMATCH_COMPARABLE_STATUSES)


def _is_comparable_status(value):
    """Порівнювати з information_unit має сенс, лише поки людина "на обліку"
    підрозділу для виплат: значення - число (вже застосована категорія
    виплати) або статус з PAYMENT_MISMATCH_COMPARABLE_STATUSES (за
    замовчуванням "РВЗ" і "ВД"). Для решти статусів (відпустка, шпиталь,
    СЗЧ, ВЛК, адміністративні) дані підрозділу цьому статусу не суперечать -
    людина або отримує виплату на загальних підставах, або дані застаріли, -
    тож порівняння лише давало б хибні розбіжності. Реальні випадки: людина
    в СЗЧ чи у відпустці за станом здоров'я не має потрапляти в
    error_mis_statuses.xlsx, хоч би що казав файл.

    Написання не важить: новоприбулій людині статус береться прямо з файлу
    БЧС, а там він буває й "РТГР" замість "РТГр"; точний збіг мовчки
    пропускав би таку клітинку."""
    return _normalize_for_compare(value) in _COMPARABLE_STATUSES or _is_number(value)


def _statuses_equivalent(a, b):
    """Чи еквівалентні дві сторони порівнюваної клітинки: нормалізовані
    значення однакові (_normalize_for_compare) або це два записи тієї самої
    ситуації через PAYMENT_STATUS_CONVERSIONS чи PAYMENT_STATUS_OVERRIDES (їх
    застосовує apply_payment_values ще до звірки). Реальні випадки:
    apply_payment_values уже перетворив базовий "ППД" на 10, а файл
    information_unit зі своїм словником і далі пише "ППД"; так само "РВЗ"
    замінено на 10 через "БЗВП"/"Адап", а файл пише "БЗВП"/"Адап". Без цього
    10 і текст ніколи б не збігались і масово потрапляли б у
    error_mis_statuses.xlsx.

    Еквівалентна й пара з PAYMENT_STATUS_PRIORITY (той самий словник, що й у
    _resolve_value_priority_conflict, його веде користувач), хоч би з якого
    боку прийшов кожен статус: якщо для пари визначено переможця, інший статус
    пари від підрозділу - не помилка, а підтвердження того самого факту з
    нижчим пріоритетом. Реальний випадок: "100_БЗ" у Табелі (перетворене "БЗ")
    і "Розп" у файлі - "БЗ" переважає, тож розбіжності немає."""
    if _normalize_for_compare(a) == _normalize_for_compare(b):
        return True
    pair = {_normalize_for_compare(a), _normalize_for_compare(b)}
    for status_text, converted in PAYMENT_STATUS_CONVERSIONS.items():
        if pair == {_normalize_for_compare(status_text), _normalize_for_compare(converted)}:
            return True
    for override_map in PAYMENT_STATUS_OVERRIDES.values():
        for info_unit_value, replacement in override_map.items():
            if pair == {_normalize_for_compare(info_unit_value), _normalize_for_compare(replacement)}:
                return True
    for status_pair in PAYMENT_STATUS_PRIORITY:
        if pair == {_normalize_for_compare(status) for status in status_pair}:
            return True
    return False


def _dedupe_files_with_identical_content(by_file):
    """Якщо два чи більше файли information_unit дають для людини цілком
    однакові дані (ПІБ, Підрозділ, Посада, Звання й ті самі значення за
    датами), це насправді один файл, що лежить двічі під трохи різними
    іменами (напр. зайвий пробіл у назві після копіювання). Лишається файл з
    першою за абеткою назвою - інакше той самий рядок розбіжності був би у
    звіті двічі."""
    seen_signatures = set()
    deduped = {}
    for file_name in sorted(by_file):
        info = by_file[file_name]
        signature = (
            info["pib_raw"], info["pidrozdil"], info["posada"], info["zvannya"],
            tuple(sorted(info["dates"].items())),
        )
        if signature in seen_signatures:
            continue
        seen_signatures.add(signature)
        deduped[file_name] = info
    return deduped


_SNAPSHOT_LABEL_INDEX = {"ПІДРОЗДІЛ": 2, "ПОСАДА": 3, "ЗВАННЯ": 4}
_INFO_LABEL_KEY = {"ПІДРОЗДІЛ": "pidrozdil", "ПОСАДА": "posada", "ЗВАННЯ": "zvannya"}


def _files_with_label(snapshot):
    """{колонка: {файли, де вона хоч раз заповнена}}. Файл, у якого колонка
    порожня для всіх (її там просто немає), не дає розбіжності "порожньо" в
    кожному своєму рядку."""
    present = {label: set() for label in PAYMENT_MISMATCH_LABEL_COLUMNS}
    for entries in snapshot.values():
        for entry in entries:
            for label, files in present.items():
                if entry[_SNAPSHOT_LABEL_INDEX[label]] not in (None, ""):
                    files.add(entry[0])
    return present


def _labels_differ(roster_labels, info, file_name, files_with_label):
    """Чи відрізняються посада чи звання (PAYMENT_MISMATCH_LABEL_COLUMNS) у
    файлі від ОБЛІК - те саме порівняння, що фарбує клітинку звіту червоним.
    Колонка, якої немає в Табелі чи у файлі, не порівнюється."""
    for label in PAYMENT_MISMATCH_LABEL_COLUMNS:
        if label not in roster_labels:
            continue
        file_value = info[_INFO_LABEL_KEY[label]]
        if file_value in (None, "") and file_name not in files_with_label[label]:
            continue
        if _normalize_for_compare(roster_labels[label]) != _normalize_for_compare(file_value):
            return True
    return False


def find_mismatches(timesheet, snapshot):
    """Для кожної людини й дати, де snapshot (read_information_unit_snapshot)
    має запис, звіряє остаточне значення в timesheet (ОБЛІК для виплат, уже
    після apply_payment_values) з кожним файлом окремо. Пара (final_value,
    file_value) зберігається в dates завжди - і коли значення збігаються, і
    коли final_value не порівнюється (напр. "ВПСЗ"/"СЗЧ"), - щоб рядок звіту
    показував усю історію файлу за датами, а не лише розбіжні дати з
    порожнечею між ними.

    Розбіжність (has_mismatch - чи потрапляє файл у результат) - лише коли
    водночас: final_value порівнюється (_is_comparable_status: статуси з
    PAYMENT_MISMATCH_COMPARABLE_STATUSES, за замовчуванням РВЗ/ВД, чи будь-яке
    число - напр. людина за рапортами вже "ВД", а підрозділ подає число, - це
    справжня розбіжність) і значення не еквівалентні (_statuses_equivalent:
    10 у ОБЛІК і "ППД" у файлі - одне й те саме). Статус, що не порівнюється
    (СЗЧ/ВПСЗ/ВЛК/...), сам розбіжності не дає, але якщо рядок уже є у звіті
    через іншу розбіжність, ця дата теж показується (зеленим чи червоним).

    У результат потрапляє й пара людина+файл, де посада чи звання у файлі
    відрізняються від ОБЛІК, навіть коли статуси збігаються (_labels_differ;
    реальний випадок - у файлі порожня посада й інше звання).

    Файли з цілком однаковими даними людини (файл випадково лежить двічі під
    різними іменами) зводяться до одного (_dedupe_files_with_identical_content).

    Окремо (_missing_from_roster_records) - люди, чий ПІБ не збігся з жодним
    рядком Табеля (немає в Табелі чи написано інакше, напр. описка в по
    батькові): інакше вони мовчки зникали б, бо основний прохід іде по Табелю.

    Повертає записи - по одному на пару людина+файл з хоча б однією
    розбіжністю: {"file_name", "roster_pib_raw", "roster_pidrozdil",
    "roster_posada", "roster_zvannya", "file_pib_raw", "file_pidrozdil",
    "file_posada", "file_zvannya", "dates": {date: (final_value, file_value)},
    "roster_statuses": {date: статус за числом з final_value}}."""
    pib_col = timesheet.label_columns[PIB_COLUMN_NAME]
    pidrozdil_col = timesheet.label_columns.get("ПІДРОЗДІЛ")
    posada_col = timesheet.label_columns.get("ПОСАДА")
    zvannya_col = timesheet.label_columns.get("ЗВАННЯ")

    person_rows = timesheet.person_rows()
    files_with_label = _files_with_label(snapshot)
    records = []
    for normalized, row_idx in person_rows.items():
        pib_raw = timesheet.ws.cell(row=row_idx, column=pib_col).value
        roster_pidrozdil = timesheet.ws.cell(row=row_idx, column=pidrozdil_col).value if pidrozdil_col else None
        roster_posada = timesheet.ws.cell(row=row_idx, column=posada_col).value if posada_col else None
        roster_zvannya = timesheet.ws.cell(row=row_idx, column=zvannya_col).value if zvannya_col else None
        roster_labels = {
            label: value for label, col, value in (
                ("ПІДРОЗДІЛ", pidrozdil_col, roster_pidrozdil),
                ("ПОСАДА", posada_col, roster_posada),
                ("ЗВАННЯ", zvannya_col, roster_zvannya),
            ) if col
        }

        by_file = {}
        for col_idx, date_value in timesheet.date_columns:
            final_value = timesheet.ws.cell(row=row_idx, column=col_idx).value
            for file_name, file_pib_raw, file_pidrozdil, file_posada, file_zvannya, value in snapshot.get((normalized, date_value), []):
                entry = by_file.setdefault(
                    file_name,
                    {"pib_raw": file_pib_raw, "pidrozdil": file_pidrozdil, "posada": file_posada, "zvannya": file_zvannya, "dates": {}, "has_mismatch": False},
                )
                entry["dates"][date_value] = (final_value, value)
                if _is_comparable_status(final_value) and not _statuses_equivalent(value, final_value):
                    entry["has_mismatch"] = True

        by_file = _dedupe_files_with_identical_content(by_file)
        for file_name, info in sorted(by_file.items()):
            # ПІБ, написане у файлі інакше (напр. з апострофом чи без), - теж
            # розбіжність, навіть коли все інше збігається: normalize_name
            # людину однаково зіставляє.
            if (not info["has_mismatch"] and not _labels_differ(roster_labels, info, file_name, files_with_label)
                    and not spelled_differently(pib_raw, info["pib_raw"])):
                continue
            records.append({
                "file_name": file_name,
                "roster_pib_raw": pib_raw,
                "roster_pidrozdil": roster_pidrozdil,
                "roster_posada": roster_posada,
                "roster_zvannya": roster_zvannya,
                "file_pib_raw": info["pib_raw"],
                "file_pidrozdil": info["pidrozdil"],
                "file_posada": info["posada"],
                "file_zvannya": info["zvannya"],
                "dates": info["dates"],
                "roster_statuses": _statuses_behind_numbers(timesheet, row_idx, info["dates"]),
            })

    records.extend(_missing_from_roster_records(timesheet, person_rows, snapshot))
    records.extend(_absent_from_files_records(timesheet, person_rows, snapshot))
    return records


def _write_compared_cell(ws, row_a, row_b, col_idx, roster_value, file_value, match=None):
    """Одна порівнювана колонка (ПІБ/Підрозділ/Посада/Звання/Статус): однакові
    значення Табеля й файлу - одна об'єднана зелена клітинка (_MATCH_FILL),
    різні - два окремі значення, по одному на рядок, обидва червоні
    (_MISMATCH_FILL). Порожнє значення показується як _EMPTY_PLACEHOLDER.
    "Збіг" за замовчуванням - за _normalize_for_compare; колонки-дати
    передають у match ширший критерій (_statuses_equivalent: напр. 10 і "ППД"
    - той самий статус, записаний по-різному), щоб еквівалентні значення не
    вважались розбіжністю."""
    is_match = match if match is not None else _normalize_for_compare(roster_value) == _normalize_for_compare(file_value)
    cell_a = ws.cell(row=row_a, column=col_idx, value=_display_value(roster_value))
    if is_match:
        ws.merge_cells(start_row=row_a, start_column=col_idx, end_row=row_b, end_column=col_idx)
        cell_a.fill = _MATCH_FILL
    else:
        cell_a.fill = _MISMATCH_FILL
        ws.cell(row=row_b, column=col_idx, value=_display_value(file_value)).fill = _MISMATCH_FILL


def _as_status(value):
    """Вид оплати людини на РВЗ (PAYMENT_DUTY_CATEGORIES) - це статус РВЗ."""
    return DEFAULT_STATUS if _is_number(value) and value in PAYMENT_DUTY_CATEGORIES else value


def _mismatch_display(record, date_value, final_value, file_value):
    """(сторона ОБЛІК, сторона файлу) розбіжної дати у звіті й повідомленні:
    статуси, а не види оплати (рішення користувача). Реальний випадок: у
    рапорті ППД, у файлі ФАКТИЧНО РВЗ з виплатою 100, а звіт показував "10" і
    "100", тож було незрозуміло, які це статуси. Тепер "ППД" (за 10 у ОБЛІК
    стоїть він, roster_statuses) і "РВЗ". Якщо обидві сторони - РВЗ, різниця
    саме у виді оплати (100 і 30), тож лишаються числа: "РВЗ" проти "РВЗ"
    нічого б не пояснив."""
    roster_value = record.get("roster_statuses", {}).get(date_value, final_value)
    roster_status, file_status = _as_status(roster_value), _as_status(file_value)
    if _normalize_for_compare(roster_status) == _normalize_for_compare(file_status):
        return roster_value, file_value
    return roster_status, file_status


def _write_manual_entry_cell(ws, row_a, row_b, col_idx):
    """Колонка для ручного заповнення (ПІДСТАВИ / ДАТА В СТАТУС
    СПЕЦКОНТИНГЕНТУ): даних для неї немає в жодному джерелі, тож вона
    лишається порожньою - об'єднана клітинка без заливки й без
    _EMPTY_PLACEHOLDER (людина заповнює її вручну, а "—" лише заважав би)."""
    ws.merge_cells(start_row=row_a, start_column=col_idx, end_row=row_b, end_column=col_idx)


# Порядок колонок - як в офіційному шаблоні звіту: Назва файлу, Підрозділ,
# Посада, Звання, ПІБ, дати і наприкінці дві колонки для ручного заповнення.
_LABEL_HEADERS = ["Назва файлу", "Підрозділ", "Посада", "Звання", "ПІБ"]
_LABEL_COLUMN_WIDTHS = (50, 26, 30, 24, 30)
_PIB_COLUMN = 5  # ширина - за найдовшим ПІБ (content/report_pib_column.py)
_DATE_COLUMN_WIDTH = 24
_MANUAL_ENTRY_HEADERS = ["ПІДСТАВИ", "ДАТА В СТАТУС СПЕЦКОНТИНГЕНТУ"]
_MANUAL_ENTRY_COLUMN_WIDTHS = (30, 30)


def write_mismatch_report(records, output_path):
    """Будує output_path (.xlsx) за структурою офіційного шаблону звіту: Назва
    файлу, Підрозділ, Посада, Звання, ПІБ, далі кожна дата - окрема колонка
    (заголовок - сама дата, як в ОБЛІК.xlsx), а наприкінці дві колонки для
    ручного заповнення (ПІДСТАВИ, ДАТА В СТАТУС СПЕЦКОНТИНГЕНТУ) - даних для
    них у проєкті немає, тож вони завжди порожні.

    На кожен запис (людина + файл з розбіжністю) - два рядки: перший -
    _REPORT_SOURCE_LABEL, другий - конкретний файл information_unit. Кілька
    розбіжних дат однієї пари - одна пара рядків з кількома заповненими
    колонками-датами (решта порожні, але з межами). Підрозділ, ПІБ, Посада,
    Звання й кожна дата - одна зелена об'єднана клітинка, якщо Табель і файл
    кажуть те саме, інакше два значення червоним. Підрозділ теж звіряється на
    прохання користувача, хоча файл часто законно каже інше (відрядження,
    прикріплення). "Назва файлу" лише позначає джерело рядка і не
    фарбується. Заголовок - як в ОБЛІК.xlsx (Bahnschrift Light SemiCondensed 14,
    жирний білий на темно-бірюзовому), дані - Times New Roman 14, тонкі межі на
    всіх клітинках, без переносу тексту (натомість широкі колонки).

    Повертає output_path або None, якщо records порожній: файл лише із
    заголовком не потрібен."""
    if not records:
        return None

    all_dates = sorted({date_value for record in records for date_value in record["dates"]})

    wb, ws = workbook_with_sheet("Розбіжності")

    date_headers = [datetime.combine(d, datetime.min.time()) for d in all_dates]
    headers = _LABEL_HEADERS + date_headers + _MANUAL_ENTRY_HEADERS
    for col_idx, value in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx, value=value)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.alignment = _HEADER_ALIGNMENT
        cell.border = _BORDER
        if isinstance(value, datetime):
            cell.number_format = DATE_NUMBER_FORMAT
    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 45

    date_col_by_date = {d: idx for idx, d in enumerate(all_dates, start=6)}
    manual_col_start = 6 + len(all_dates)

    current_row = 2
    for record in records:
        row_a, row_b = current_row, current_row + 1

        ws.cell(row=row_a, column=1, value=_REPORT_SOURCE_LABEL)
        ws.cell(row=row_b, column=1, value=record["file_name"])

        for col_idx, roster_value, file_value in (
            (2, record["roster_pidrozdil"], record["file_pidrozdil"]),
            (3, record["roster_posada"], record["file_posada"]),
            (4, record["roster_zvannya"], record["file_zvannya"]),
        ):
            _write_compared_cell(ws, row_a, row_b, col_idx, roster_value, file_value)
        _write_compared_cell(ws, row_a, row_b, 5, record["roster_pib_raw"], record["file_pib_raw"],
                             match=not spelled_differently(record["roster_pib_raw"], record["file_pib_raw"]))

        for date_value, (final_value, file_value) in record["dates"].items():
            is_match = _statuses_equivalent(final_value, file_value)
            roster_shown, file_shown = (final_value, file_value) if is_match else _mismatch_display(
                record, date_value, final_value, file_value,
            )
            _write_compared_cell(ws, row_a, row_b, date_col_by_date[date_value], roster_shown, file_shown, match=is_match)

        for col_idx in range(manual_col_start, manual_col_start + len(_MANUAL_ENTRY_HEADERS)):
            _write_manual_entry_cell(ws, row_a, row_b, col_idx)

        for row in (row_a, row_b):
            ws.row_dimensions[row].height = 20
            for col_idx in range(1, len(headers) + 1):
                cell = ws.cell(row=row, column=col_idx)
                cell.font = _FONT
                cell.alignment = PIB_ALIGNMENT if col_idx == _PIB_COLUMN else _CENTER
                cell.border = _BORDER

        current_row += 2

    for col_idx in range(1, len(headers) + 1):
        letter = get_column_letter(col_idx)
        if col_idx == _PIB_COLUMN:
            width = pib_column_width(
                [record["roster_pib_raw"] for record in records] + [record["file_pib_raw"] for record in records]
            )
        elif col_idx <= 5:
            width = _LABEL_COLUMN_WIDTHS[col_idx - 1]
        elif col_idx < manual_col_start:
            width = _DATE_COLUMN_WIDTH
        else:
            width = _MANUAL_ENTRY_COLUMN_WIDTHS[col_idx - manual_col_start]
        ws.column_dimensions[letter].width = width

    save_workbook(wb, output_path)
    return output_path


# Повідомлення командирам (content/commander_message.py): "як має бути" -
# значення з ОБЛІК.
_MESSAGE_LABELS = {"ПОСАДА": ("Посада", "posada"), "ЗВАННЯ": ("Звання", "zvannya")}
_NOT_IN_FILE_ISSUE = "немає у файлі БЧС підрозділу"


def _date_issues(record):
    """Розбіжні дати (ті, що у звіті червоні, - не _statuses_equivalent) ->
    проміжки зі значенням ОБЛІК (date_issues), як у звіті (_mismatch_display)."""
    return date_issues({
        date_value: _display_value(_mismatch_display(record, date_value, final_value, file_value)[0])
        for date_value, (final_value, file_value) in record["dates"].items()
        if not _statuses_equivalent(final_value, file_value)
    })


def _message_entry(record):
    """(підрозділ, ПІБ, [що виправити]) для одного запису звіту. Підрозділ -
    власник файлу (виправляє він); для "немає в жодному файлі" - з ОБЛІК."""
    if record["file_name"] == _NOT_IN_FILES_LABEL:
        return record["roster_pidrozdil"], record["roster_pib_raw"], [_NOT_IN_FILE_ISSUE]
    subdivision = (file_subdivision(record["file_name"]) or record["file_pidrozdil"]
                   or record["roster_pidrozdil"])
    if record["roster_pib_raw"] == _NOT_IN_ROSTER_LABEL:
        return subdivision, record["file_pib_raw"], [NOT_IN_ROSTER_ISSUE]
    pib, issues = record["roster_pib_raw"], []
    if spelled_differently(record["roster_pib_raw"], record["file_pib_raw"]):
        # Інше написання ПІБ у файлі (апостроф чи _missing_from_roster_records):
        # людина - як у файлі командира, а правильне написання - першим пунктом.
        pib = record["file_pib_raw"]
        issues.append(f"ПІБ: {record['roster_pib_raw']}")
    for label in PAYMENT_MISMATCH_LABEL_COLUMNS:
        title, key = _MESSAGE_LABELS[label]
        roster_value, file_value = record[f"roster_{key}"], record[f"file_{key}"]
        if _normalize_for_compare(roster_value) != _normalize_for_compare(file_value):
            issues.append(f"{title}: {_display_value(roster_value)}")
    return subdivision, pib, issues + _date_issues(record)


def mismatch_message(records):
    """Текст повідомлення командирам за записами find_mismatches (тими самими,
    з яких пишеться error_mis_statuses.xlsx); "" - нема про що писати."""
    return compose_message(_message_entry(record) for record in records)


def write_mismatch_message(records, output_path):
    """Додає ці записи в спільний файл повідомлення output_path
    (update_message_file); None - нема про що писати."""
    return update_message_file(output_path, PAYMENTS_SOURCE, [_message_entry(record) for record in records])
