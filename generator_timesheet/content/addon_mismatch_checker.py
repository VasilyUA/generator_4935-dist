"""Звірка щойно згенерованого ОБЛІК.xlsx (аркуш "Табель") з книгами
resources/addons: book61 (content/addon_reader.py, інструкція - аркуш "map"
у resources/ОБЛІК.xlsx), ЦВС (content/cvs_reader.py, аркуш "map" самої
книги) та журналом обліку БП. Чи відповідає статус у Табелі на кожну дату
тому, що випливає з цих книг: періоди відпусток, відряджень, лікування,
адміністративні накази, втрати.

Будує mis_match.xlsx у тому самому стилі, що й error_mis_statuses.xlsx
(write_mismatch_report у content/payment_mismatch_checker.py) - стилі
імпортуються звідти, а не копіюються."""

import difflib
import re
from bisect import bisect_left, bisect_right
from datetime import datetime, timedelta

from openpyxl.styles import PatternFill
from openpyxl.utils import get_column_letter

from constants import (
    ADDON_ARRIVAL_DAY_SHEETS,
    ADDON_ARRIVAL_DAY_STATUSES,
    ADDON_EQUIVALENT_STATUS_GROUPS,
    ADDON_NOT_RETURNED_AFTER_SZCH_STATUSES,
    ADDON_ONE_OF_STATUSES,
    ADDON_REQUIRED_BOOK61_SHEETS,
    ADDON_TABLE_PRIORITY_STATUSES,
    ADDON_TRACKED_STATUSES,
    DATE_NUMBER_FORMAT,
    PAYMENT_STATUS_CONVERSIONS,
    PAYMENT_STATUS_OVERRIDES,
    PIB_COLUMN_NAME,
    TRANSFER_STATUSES,
)
from content.addon_reader import ANY_STATUS, RETURNED_FROM_SZCH_STATUS
from content.excel_compat import empty_workbook, save_workbook
from content.oblik_timesheet import normalize_name, spelled_differently
from content.report_pib_column import PIB_ALIGNMENT, pib_column_width
from content.payment_mismatch_checker import (
    _BORDER,
    _CENTER,
    _display_value,
    _FONT,
    _HEADER_ALIGNMENT,
    _HEADER_FILL,
    _HEADER_FONT,
    _is_number,
    _MATCH_FILL,
    _MISMATCH_FILL,
    _normalize_for_compare,
)


def _is_payment_category_value(value):
    """Чи value - запис категорії виплати (число чи текст на кшталт
    "100_..."/"10_..." з PAYMENT_STATUS_CONVERSIONS/PAYMENT_STATUS_OVERRIDES),
    а не звичайний текстовий статус. У тих самих словниках є й тотожності
    ("Навч" -> "Навч", "ВПС" -> "ВПС", "Повер" -> "Повер": information_unit
    лише підтверджує точнішу причину, не змінюючи категорії). Їх треба
    розрізняти для _payment_converted_forms: категорія - це та сама ситуація,
    записана інакше, а тотожність - справді інша причина."""
    if _is_number(value):
        return True
    if isinstance(value, str) and "_" in value:
        return value.split("_", 1)[0].isdigit()
    return False


def _payment_converted_forms(status):
    """Усі значення, які apply_payment_values (content/oblik_timesheet.py) міг
    би записати в клітинку Табеля для дня, коли за рапортом статус - status:
    фіксована конвертація (PAYMENT_STATUS_CONVERSIONS[status]) або заміна за
    підтвердженням information_unit (PAYMENT_STATUS_OVERRIDES[status]), - але
    лише значення-категорії (_is_payment_category_value). Книга addons
    визначає факт незалежно, тож "100_БПШП" замість "БПШП" чи 10 замість "ВД" -
    та сама ситуація, записана інакше вже після неї, а не розбіжність.

    Реальні випадки в output/ОБЛІК.xlsx: book61 (аркуш "ШП") очікує "БПШП", а
    в Табелі "100_БПШП"; book61 (аркуш "ВД") очікує "ВД", а в Табелі 10; так
    само "ВПСЗ" (аркуш "ВП") і "100_ВПБП".

    "ВД" і текстове "Навч" навмисно не зводяться
    (PAYMENT_STATUS_OVERRIDES["ВД"]["Навч"] = "Навч" - тотожність, а не
    категорія): ВД і Навч - справжня різниця (_is_training_destination у
    content/daily_report_reader.py). "Навч" ніде не є базовим статусом
    (базові там лише DEFAULT_STATUS і "ВД"), тож для очікуваного "Навч" форм
    немає: число 10 у Табелі проти "Навч" у book61 лишається розбіжністю, а
    не гаситься через "10 ~ ВД"."""
    forms = set()
    converted = PAYMENT_STATUS_CONVERSIONS.get(status)
    if converted is not None and _is_payment_category_value(converted):
        forms.add(converted)
    # Категорія підрозділу, записана як є (100 -> 100: "Адап"/"РРез" уже
    # закінчились, constants._DUTY_CATEGORIES_AS_IS), - це кінець статусу, а не
    # інший запис того самого статусу, тож не форма.
    for info_unit_value, replacement in PAYMENT_STATUS_OVERRIDES.get(status, {}).items():
        if _is_payment_category_value(replacement) and info_unit_value != replacement:
            forms.add(replacement)
    # Статус, яким information_unit уточнює інший базовий статус, - теж форма:
    # "БЗВП"/"Адап"/"НОВ" у клітинці "РВЗ" записуються числом 10
    # (PAYMENT_STATUS_OVERRIDES[DEFAULT_STATUS]). Реальний випадок: журнал
    # обліку БП очікує "БЗВП", а Табель для виплат каже 10.
    for override_map in PAYMENT_STATUS_OVERRIDES.values():
        replacement = override_map.get(status)
        if replacement is not None and _is_payment_category_value(replacement):
            forms.add(replacement)
    return forms


def _paid_as(status):
    """Категорії виплати, якими статус оплачується за будь-яких обставин:
    фіксована конвертація (ППД -> 10) чи заміна, коли саме цей статус подає
    information_unit (БЗВП/Адап/НОВ/ВД -> 10). Без можливих уточнень базового
    статусу: РВЗ лише може стати 10."""
    forms = {PAYMENT_STATUS_CONVERSIONS.get(status)}
    forms.update(override_map.get(status) for override_map in PAYMENT_STATUS_OVERRIDES.values())
    return {_normalize_for_compare(form) for form in forms if form is not None and _is_payment_category_value(form)}


def status_options(status):
    """Статуси, будь-який з яких підходить (ADDON_ONE_OF_STATUSES), або сам status."""
    return ADDON_ONE_OF_STATUSES.get(status, (status,)) if isinstance(status, str) else (status,)


def _not_returned_after_szch(actual):
    """Чи статус Табеля каже, що людина після СЗЧ так і не повернулась на
    службу: ADDON_NOT_RETURNED_AFTER_SZCH_STATUSES чи їхня форма для виплат."""
    actual_normalized = _normalize_for_compare(actual)
    return any(
        actual_normalized in {_normalize_for_compare(form) for form in (status, *_payment_converted_forms(status))}
        for status in ADDON_NOT_RETURNED_AFTER_SZCH_STATUSES
    )


def _statuses_match(expected, actual):
    if expected == ANY_STATUS:
        return True
    if _normalize_for_compare(expected) == _normalize_for_compare(RETURNED_FROM_SZCH_STATUS):
        return not _not_returned_after_szch(actual)
    if len(status_options(expected)) > 1 or len(status_options(actual)) > 1:
        return any(_statuses_match(option_expected, option_actual)
                   for option_expected in status_options(expected) for option_actual in status_options(actual))
    expected_normalized = _normalize_for_compare(expected)
    actual_normalized = _normalize_for_compare(actual)
    if expected_normalized == actual_normalized:
        return True
    # Очікується сама категорія виплати (журнал обліку БП: "закордон" - 10):
    # збіг і для статусу в Табелі, що нею оплачується завжди (ППД/БЗВП/НОВ/
    # Адап/ВД -> 10), але не для "РВЗ", який лише може стати 10.
    if _is_number(expected) and expected_normalized in _paid_as(actual):
        return True
    for group in ADDON_EQUIVALENT_STATUS_GROUPS:
        normalized_group = {_normalize_for_compare(status) for status in group}
        if expected_normalized in normalized_group and actual_normalized in normalized_group:
            return True
    for expected_status, tabel_statuses in ADDON_TABLE_PRIORITY_STATUSES.items():
        if expected_normalized == _normalize_for_compare(expected_status) and any(
            actual_normalized == _normalize_for_compare(status) for status in tabel_statuses
        ):
            return True
    return any(actual_normalized == _normalize_for_compare(form) for form in _payment_converted_forms(expected))


# Прибуття з аркуша "Прибулі у ВД" (РРез -> ППД -> будь-який статус у день
# прийому посади, read_rrez_tvo_periods) на весь свій проміжок переважає
# будь-який паралельний період book61 (реальний випадок: ТВО і "ППД" з тієї
# самої дати, а "ППД" з аркуша ППД має тривати лише після прийому посади).
# Тому такі записи накладаються останніми, поверх решти, хоч би коли
# починались.
_PRIORITY_SOURCE_SHEETS = ("Прибулі у ВД",)


def _accepted_on_arrival_day(record, date_value, actual):
    """День прибуття з ADDON_ARRIVAL_DAY_SHEETS: "РВЗ" (30, 100) тут має
    перевагу (рішення користувача), тож значення з ADDON_ARRIVAL_DAY_STATUSES у
    Табелі цього дня - не розбіжність. Статус самого аркуша теж лишається
    правильним (_statuses_match)."""
    if record["source_sheet"] not in ADDON_ARRIVAL_DAY_SHEETS or record.get("arrival_date") != date_value:
        return False
    accepted = {str(_normalize_for_compare(status)) for status in ADDON_ARRIVAL_DAY_STATUSES}
    return str(_normalize_for_compare(actual)) in accepted


def _paint(timeline, record, start, end, dates):
    """dates - відсортовані дати шкали. Перебираються лише дні самого запису:
    у звірці книг між собою шкала охоплює роки, а не місяць."""
    for date_value in dates[bisect_left(dates, start):bisect_right(dates, end)]:
        timeline[date_value] = record


# Остаточні й адміністративні статуси не "повертаються" після пізніших подій:
# будь-який новіший період означає, що вони вже не діють (людина знову на
# службі чи запис book61 застарів). Реальний випадок: "Звіль" з Алф2 за 2025
# рік (відкритий період) "повертався" після ВЛК у 2026, хоча пізніші
# ШП/ВП/ВЛК прямо свідчать, що людина на службі. З "СЗЧ" так само: пізніша
# відпустка чи відрядження означає, що людина вже повернулась (СЗЧ з лютого
# 2026 "повертався" після ВП у жовтні).
_SUPERSEDED_BY_LATER_STATUSES = ("Звіль", *TRANSFER_STATUSES, "полон", "Загиб", "Смерть", "БЗ", "Розп", "СЗЧ")

# Відкритий період (кінця в книзі немає) решти статусів так само закінчується
# перед пізнішою подією. Реальний випадок (книга ЦВС): лікування без "Дата
# виписки" поверталось після пізніших ВЛК/ВП і тягнулось до кінця місяця.
# Виняток - ППД: туди людина повертається після лікування чи ТВО.
_RESUMING_OPEN_STATUSES = ("ППД",)


def _ends_at_next_event(record):
    if record["status"] in _SUPERSEDED_BY_LATER_STATUSES:
        return True
    if record["end_date"] is not None:
        return False
    # Відкритий запис "Прибулі у ВД" (посаду ще не прийнято) накладається
    # поверх решти, тож навіть "ППД" з нього закінчується перед пізнішою
    # подією. Реальний випадок: зарахування 08.10 без прийому посади, далі
    # відпустка 21.10 - інакше "ППД" перекрив би відпустку.
    return record["status"] not in _RESUMING_OPEN_STATUSES or record["source_sheet"] in _PRIORITY_SOURCE_SHEETS


def _known_end_of_open_period(record, records):
    """Кінець відкритого періоду, відомий з іншого запису того самого статусу
    (інша книга знає, коли цей епізод закінчився): найраніший кінець такого
    запису, що не закінчився до початку цього, або None. Реальний випадок:
    book61 "ВЛК" з 14.10 без "Кінець (ВЛК)", ЦВС - ВЛК 14.10-20.10 (документ
    "від 20.10.2026"). ВЛК тягнулась до кінця місяця, бо запис ЦВС починався
    того самого дня, а не пізніше, і _ends_at_next_event його не враховує."""
    if record["end_date"] is not None or record["status"] in _RESUMING_OPEN_STATUSES:
        return None
    known_ends = [
        other["end_date"] for other in records
        if other is not record and set(status_options(other["status"])) & set(status_options(record["status"]))
        and _ends_the_same_episode(record, other)
    ]
    return min(known_ends) if known_ends else None


def _ends_the_same_episode(record, other):
    """Чи кінець other - кінець того самого епізоду, що й відкритий record:
    other закінчився після початку record або того самого дня, але й
    почався того самого дня (ВЛК з документом того ж дня). Не підходять
    попередній епізод, що закінчився в день початку record (реальний випадок:
    виписка з однієї лікарні й госпіталізація до іншої 16.10 обривали нове
    лікування 16.10), і одноденний запис без документа (end_assumed: ЦВС ВЛК
    лише з датою проходження обривав ВЛК book61)."""
    end = other["end_date"]
    if end is None or other.get("end_assumed"):
        return False
    start = record["start_date"]
    return end > start or (end == start and other["start_date"] == start)


def _paint_timeline(records, date_columns, last_date):
    """{дата Табеля: запис, що її покриває} - одна послідовна шкала з усіх
    періодів людини в усіх книгах resources/addons. Записи накладаються за
    датою початку, пізніший поверх попереднього, тож кожен період фактично
    триває до дня перед початком наступного ("ППД -> ШП -> СЗЧ"). Якщо
    пізніший закінчився раніше, ніж за книгою мав закінчитись попередній,
    попередній після нього продовжується (напр. ППД після ТВО).

    Винятки (_ends_at_next_event): _SUPERSEDED_BY_LATER_STATUSES і відкриті
    періоди (end_date None, крім _RESUMING_OPEN_STATUSES) закінчуються за день
    до будь-якої пізнішої події й після неї не продовжуються. Відкритий
    період, чий кінець знає інша книга (_known_end_of_open_period),
    закінчується там само. Записи з _PRIORITY_SOURCE_SHEETS ("Прибулі у ВД")
    накладаються останніми; відкритий такий запис (навіть "ППД") теж
    закінчується перед пізнішою подією. Відкритий період без пізніших подій
    триває до останньої колонки-дати."""
    ordered = sorted(records, key=lambda r: (r["source_sheet"] in _PRIORITY_SOURCE_SHEETS, r["start_date"]))
    dates = sorted(date_value for _col_idx, date_value in date_columns)
    timeline = {}
    for record in ordered:
        start = record["start_date"]
        end = record["end_date"] or _known_end_of_open_period(record, records) or last_date
        if _ends_at_next_event(record):
            # Запис без документа (end_assumed: ЦВС ВЛК лише з датою
            # проходження) - не подія, після якої період закінчився: сам день
            # ВЛК однаково накладається поверх. Реальний випадок: БПШП з 15.10
            # без виписки, ЦВС ВЛК 28.10 без документа - 29.10 лишався без
            # жодного запису, хоча людина й далі на лікуванні.
            later_starts = [other["start_date"] for other in records
                            if other["start_date"] > start and not other.get("end_assumed")]
            if later_starts:
                end = min(end, min(later_starts) - timedelta(days=1))
        _paint(timeline, record, start, end, dates)
    return timeline


# Поріг схожості слова ПІБ з опискою - той самий, що й для описок у рапортах
# (content/oblik_timesheet.py, _TYPO_HINT_SIMILARITY_CUTOFF).
_NAME_TYPO_SIMILARITY = 0.8


def _misspelled_roster_name(addon_normalized, roster_words):
    """Табельне ПІБ тієї самої людини, записане в книзі з опискою: два слова
    з трьох однакові, а третє схоже (_NAME_TYPO_SIMILARITY), і такий кандидат
    рівно один. Реальний випадок: Табель - "Констянтин", book61 -
    "Костянтин", ЦВС - "Константин", прізвище й по батькові однакові. Інакше
    None. "Петро"/"Павло" чи "Олександр"/"Олексій" (різні люди з однаковим
    прізвищем і по батькові) поріг не проходять. roster_words - [(ПІБ Табеля,
    його слова)] лише з трьох слів."""
    words = addon_normalized.split()
    if len(words) != 3:
        return None
    candidates = []
    for roster_name, roster_name_words in roster_words:
        different = [i for i in range(3) if words[i] != roster_name_words[i]]
        if len(different) == 1 and difflib.SequenceMatcher(
                None, words[different[0]], roster_name_words[different[0]]).ratio() >= _NAME_TYPO_SIMILARITY:
            candidates.append(roster_name)
    return candidates[0] if len(candidates) == 1 else None


def _attach_misspelled_records(grouped, person_rows):
    """Записи книги з ПІБ, якого немає в Табелі, - людині Табеля, якщо це та
    сама людина з опискою в ПІБ (_misspelled_roster_name), а не просто хтось,
    кого в Табелі немає. Звіт тоді показує обидва написання ПІБ (як і для
    апострофа), бо record["pib_raw"] лишається написанням книги. Не
    приєднується, якщо в цій самій книзі вже є записи з точним табельним ПІБ:
    тоді книга знає людину Табеля, а схоже ПІБ - інша людина."""
    exact = {(name, record["source_file"]) for name, records in grouped.items() if name in person_rows
             for record in records}
    roster_words = [(name, name.split()) for name in person_rows if len(name.split()) == 3]
    attached = {name: list(records) for name, records in grouped.items()}
    for name, records in grouped.items():
        if name in person_rows:
            continue
        roster_name = _misspelled_roster_name(name, roster_words)
        if roster_name is None:
            continue
        moved = [record for record in records if (roster_name, record["source_file"]) not in exact]
        attached.setdefault(roster_name, []).extend(moved)
        attached[name] = [record for record in records if record not in moved]
    return attached


_NO_RECORD_LABEL = "немає запису на аркуші {sheet}"


def _status_targets(mapping, files_by_kind):
    """[(нормалізовані форми статусу - сам статус і його форми в ОБЛІК для
    виплат, назва книги, аркуш)] з mapping {статус: (тип книги, аркуш)} - лише
    для книг, які звірялись (files_by_kind {тип: назва})."""
    return [
        ({_normalize_for_compare(form) for form in (status, *_payment_converted_forms(status))},
         files_by_kind[kind], sheet)
        for status, (kind, sheet) in mapping.items() if kind in files_by_kind
    ]


def _target_for(actual, targets):
    """(назва книги, аркуш) першого з targets, до якого належить статус
    Табеля actual, або None."""
    normalized = _normalize_for_compare(actual)
    return next(((source_file, sheet) for forms, source_file, sheet in targets if normalized in forms), None)


def _covered_by_sheet(records, source_file, sheet, date_value):
    return any(
        record["source_file"] == source_file and record["source_sheet"] == sheet
        and record["start_date"] <= date_value and (record["end_date"] is None or date_value <= record["end_date"])
        for record in records
    )


def find_period_mismatches(timesheet, period_records, files_by_kind=None):
    """Для кожної людини з Табеля - одна послідовна шкала очікуваних статусів
    з усіх її записів addons (_paint_timeline), звірена з фактичним статусом
    Табеля на кожну дату. Реальний випадок: book61 "ППД" 13.08-04.10, "ШП"
    03.10-06.10, "СЗЧ" з 06.10 колись давали два рядки звіту з датами, що
    перекривались; тепер це один рядок: ППД до 02.10 -> ШП 03.10-05.10 -> СЗЧ
    з 06.10 (до наступного статусу, напр. "Розп").

    Людина, якої немає в Табелі, не звіряється взагалі, без винятків: Табель
    вирішує, кого звіряти, а книги лише кажуть, що очікується для цих людей.
    Колись люди з давніми "СЗЧ" і "ВД", яких давно немає в Табелі, хибно
    потрапляли в mis_match.xlsx.

    Шкала одна на людину, навіть коли записи йдуть з кількох аркушів чи
    файлів (правила _paint_timeline діють через усі книги). Звіт розділено за
    файлами addons: кожна дата належить файлу, чий запис визначив очікуваний
    статус, і на кожен файл - окремий запис людини (один рядок на аркуші
    цього файлу, без повторів ПІБ).

    Повертає [{"source_file", "sources" ("файл: аркуш" цього файлу в тому
    порядку, як вони змінюють одне одного на шкалі - для колонки "Джерело"),
    "roster_pidrozdil" (ПІДРОЗДІЛ з Табеля), "roster_pib_raw", "addon_pib_raw",
    "dates": {date: (actual, expected)}, ...}] - лише пари людина+файл, де є
    справжня розбіжність (_statuses_match - False), інше написання ПІБ чи
    відпустка без "Куди вибув".

    files_by_kind ({тип книги "book61"/"journal": назва серед звірених})
    вмикає й зворотну перевірку - для людей Табеля навіть без жодного запису в
    addons. Розбіжність "немає запису на аркуші ..." (на аркуші звіту
    відповідної книги) дає:
    - день зі статусом з ADDON_REQUIRED_BOOK61_SHEETS (відпустки), якого не
      покриває жоден запис потрібного аркуша book61 - навіть якщо цей день
      покриває інша книга (запис іншого аркуша book61 покаже звичайна звірка);
    - день зі статусом з ADDON_TRACKED_STATUSES, коли жодна книга нічого не
      каже (запис будь-якої книги звіряється звичайно, за всіма правилами)."""
    person_rows = timesheet.person_rows()
    pib_col = timesheet.label_columns[PIB_COLUMN_NAME]
    date_columns = timesheet.date_columns  # Timesheet гарантує хоча б одну колонку-дату
    last_date = date_columns[-1][1]

    grouped = {}
    for record in period_records:
        grouped.setdefault(normalize_name(record["pib_raw"]), []).append(record)
    grouped = _attach_misspelled_records(grouped, person_rows)

    results = []
    files_by_kind = files_by_kind or {}
    required_targets = _status_targets(
        {status: ("book61", sheet) for status, sheet in ADDON_REQUIRED_BOOK61_SHEETS.items()}, files_by_kind,
    )
    tracked_targets = _status_targets(ADDON_TRACKED_STATUSES, files_by_kind)
    people = set(grouped) | (set(person_rows) if files_by_kind else set())
    for normalized in sorted(people):
        records = grouped.get(normalized, [])
        row_idx = person_rows.get(normalized)
        if row_idx is None:
            continue
        timeline = _paint_timeline(records, date_columns, last_date)

        by_file = {}
        for col_idx, date_value in date_columns:
            record = timeline.get(date_value)
            actual = timesheet.ws.cell(row=row_idx, column=col_idx).value
            # Порожня клітинка - дата після останнього рапорту, яку apply_events
            # (content/oblik_timesheet.py) навмисно лишає порожньою. Даних ще
            # немає, тож і розбіжності немає.
            if actual is None:
                continue
            # Число (категорія виплати: 10 - це і ВД, і ППД, і БЗВП...) не каже,
            # запис якої книги мав би бути.
            missing = None
            required = None if _is_number(actual) else _target_for(actual, required_targets)
            if required:
                source_file, sheet = required
                if not (record and record["source_file"] == source_file) and not _covered_by_sheet(
                        records, source_file, sheet, date_value):
                    missing = required
            elif record is None and not _is_number(actual):
                missing = _target_for(actual, tracked_targets)
            if missing:
                source_file, sheet = missing
                entry = by_file.setdefault(source_file, {
                    "dates": {}, "sources": [],
                    "pib_raw": timesheet.ws.cell(row=row_idx, column=pib_col).value,
                })
                entry["dates"][date_value] = (actual, _NO_RECORD_LABEL.format(sheet=sheet))
                if f"{source_file}: {sheet}" not in entry["sources"]:
                    entry["sources"].append(f"{source_file}: {sheet}")
            if record is None:
                continue
            expected = actual if _accepted_on_arrival_day(record, date_value, actual) else record["status"]
            entry = by_file.setdefault(record["source_file"], {"dates": {}, "sources": [], "pib_raw": record["pib_raw"]})
            entry["dates"][date_value] = (actual, expected)
            if record.get("destination_missing"):  # ВП без "Куди вибув" (addon_reader._vp_destination_flags)
                entry.setdefault("no_destination", set()).add(date_value)
            source = f"{record['source_file']}: {record['source_sheet']}"
            if not entry["sources"] or entry["sources"][-1] != source:
                entry["sources"].append(source)

        # ПІБ, написане в книзі інакше (напр. з апострофом чи без), - теж
        # розбіжність: пара людина+книга потрапляє у звіт, навіть коли всі дати
        # збігаються (normalize_name людину зіставив).
        roster_pib_raw = timesheet.ws.cell(row=row_idx, column=pib_col).value
        for record in records:
            if spelled_differently(record["pib_raw"], roster_pib_raw):
                entry = by_file.setdefault(record["source_file"], {"dates": {}, "sources": []})
                entry["pib_raw"] = record["pib_raw"]
                source = f"{record['source_file']}: {record['source_sheet']}"
                if source not in entry["sources"]:
                    entry["sources"].append(source)

        for source_file, entry in by_file.items():
            if (all(_statuses_match(expected, actual) for actual, expected in entry["dates"].values())
                    and not spelled_differently(entry["pib_raw"], roster_pib_raw)
                    and not entry.get("no_destination")):
                continue
            results.append({
                "source_file": source_file,
                "sources": entry["sources"],
                "roster_pidrozdil": _roster_pidrozdil(timesheet, row_idx),
                "roster_pib_raw": timesheet.ws.cell(row=row_idx, column=pib_col).value,
                "addon_pib_raw": entry["pib_raw"],
                "dates": entry["dates"],
                # Дні відпустки без "Куди вибув" (аркуш "ВП") - звіт показує жовтим.
                "no_destination": sorted(entry.get("no_destination", ())),
                # Увесь рядок Табеля - для дат, які звіряються на аркуші іншої
                # книги чи не звіряються зовсім (звіт показує їх сірим).
                "tabel_values": {
                    date_value: timesheet.ws.cell(row=row_idx, column=col_idx).value
                    for col_idx, date_value in date_columns
                },
            })
    return results


def is_warning_only(record):
    """Запис find_period_mismatches - лише попередження, а не помилка (у
    терміналі жовтим): усі статуси збігаються, ПІБ написано так само, а
    єдине зауваження - відпустка без "Куди вибув" (no_destination)."""
    return (
        bool(record.get("no_destination"))
        and all(_statuses_match(expected, actual) for actual, expected in record["dates"].values())
        and not spelled_differently(record["roster_pib_raw"], record["addon_pib_raw"])
    )


def _roster_pidrozdil(timesheet, row_idx):
    """ПІДРОЗДІЛ людини з Табеля або None, якщо такої колонки немає."""
    pidrozdil_col = timesheet.label_columns.get("ПІДРОЗДІЛ")
    return timesheet.ws.cell(row=row_idx, column=pidrozdil_col).value if pidrozdil_col else None


def unmapped_people_in_roster(timesheet, unmapped_people):
    """Ті з read_unmapped_alf2_people (content/addon_reader.py), хто є в
    Табелі: як і в find_period_mismatches, кого немає в Табелі, той з
    addons не звіряється взагалі, навіть як застереження. Кожна людина - один
    раз (на Алф2 буває кілька рядків з тим самим ПІБ, а ПІБ у звіті не
    повторюються), разом з її ПІДРОЗДІЛОМ з Табеля ("roster_pidrozdil")."""
    person_rows = timesheet.person_rows()
    kept = {}
    for person in unmapped_people:
        normalized = normalize_name(person["pib_raw"])
        if normalized in person_rows and normalized not in kept:
            kept[normalized] = {**person, "roster_pidrozdil": _roster_pidrozdil(timesheet, person_rows[normalized])}
    return list(kept.values())


_UNMAPPED_WARNING_HEADERS = ["Підрозділ", "ПІБ (book61)", "Статус в Алф2 (book61)", "Примітка"]
_UNMAPPED_WARNING_COLUMN_WIDTHS = (26, 30, 24, 70)


_NO_DESTINATION_NOTE = "; ВП: не вказано «Куди вибув»"
# Нейтральна жовта заливка Excel - поряд із зеленою (збіг) і червоною
# (розбіжність статусу): так позначено відпустку без "Куди вибув".
_NO_DESTINATION_FILL = PatternFill(start_color="FFFFEB9C", end_color="FFFFEB9C", fill_type="solid")


def _source_label(record):
    """Колонка "Джерело": "book61: ППД -> book61: ШП -> book61: СЗЧ" (і
    примітка про відпустку без "Куди вибув", якщо є)."""
    label = " -> ".join(record["sources"])
    return label + _NO_DESTINATION_NOTE if record.get("no_destination") else label


# Той самий нейтральний сірий, що й "не звірялось" у звіті Відомості РОП/ВОП
# (content/rop_vop_statement.py): ні збіг (зелений), ні розбіжність (червоний).
_NOT_COMPARED_FILL = PatternFill(start_color="FFF2F2F2", end_color="FFF2F2F2", fill_type="solid")


def _write_period_mismatch_rows(ws, records, start_row, date_col_by_date):
    current_row = start_row
    for record in records:
        row_a, row_b = current_row, current_row + 1
        ws.cell(row=row_a, column=1, value="Табель (ОБЛІК.xlsx)")
        ws.cell(row=row_b, column=1, value=_source_label(record))

        ws.cell(row=row_a, column=2, value=_display_value(record["roster_pidrozdil"]))
        ws.merge_cells(start_row=row_a, start_column=2, end_row=row_b, end_column=2)

        roster_display = _display_value(record["roster_pib_raw"])
        addon_display = _display_value(record["addon_pib_raw"])
        cell_a = ws.cell(row=row_a, column=3, value=roster_display)
        if not spelled_differently(record["roster_pib_raw"], record["addon_pib_raw"]):
            ws.merge_cells(start_row=row_a, start_column=3, end_row=row_b, end_column=3)
            cell_a.fill = _MATCH_FILL
        else:
            cell_a.fill = _MISMATCH_FILL
            ws.cell(row=row_b, column=3, value=addon_display).fill = _MISMATCH_FILL

        for date_value, col_idx in date_col_by_date.items():
            if date_value not in record["dates"]:
                # Рядок без розривів: дата, яку звіряє інша книга (її аркуш) чи
                # жодна, показує значення Табеля сірим, як "не звірялось" у
                # звіті Відомості РОП/ВОП.
                cell = ws.cell(row=row_a, column=col_idx,
                               value=_display_value(record.get("tabel_values", {}).get(date_value)))
                ws.merge_cells(start_row=row_a, start_column=col_idx, end_row=row_b, end_column=col_idx)
                cell.fill = _NOT_COMPARED_FILL
                continue
            actual, expected = record["dates"][date_value]
            cell_actual = ws.cell(row=row_a, column=col_idx, value=_display_value(actual))
            if date_value in record.get("no_destination", ()) and _statuses_match(expected, actual):
                # Статус збігся, але у відпустки немає "Куди вибув" - лише жовтим,
                # без тексту (червона розбіжність статусу того ж дня має
                # перевагу).
                ws.merge_cells(start_row=row_a, start_column=col_idx, end_row=row_b, end_column=col_idx)
                cell_actual.fill = _NO_DESTINATION_FILL
            elif _statuses_match(expected, actual):
                ws.merge_cells(start_row=row_a, start_column=col_idx, end_row=row_b, end_column=col_idx)
                cell_actual.fill = _MATCH_FILL
            else:
                cell_actual.fill = _MISMATCH_FILL
                ws.cell(row=row_b, column=col_idx, value=_display_value(expected)).fill = _MISMATCH_FILL

        for row in (row_a, row_b):
            ws.row_dimensions[row].height = 20
        current_row += 2
    return current_row


_UNMAPPED_SHEET_TITLE = "Застереження"
_NO_MISMATCHES_TEXT = "Розбіжностей з цим файлом немає."
_SHEET_TITLE_FORBIDDEN_RE = re.compile(r"[\[\]:*?/\\]")
_SHEET_TITLE_MAX_LENGTH = 31


def _sheet_title(name, used):
    """Назва аркуша з імені файлу addons за правилами Excel: без []:*?/\\, не
    довше 31 символу й унікальна в книзі (used - уже зайняті, у нижньому
    регістрі)."""
    base = _SHEET_TITLE_FORBIDDEN_RE.sub("_", str(name)).strip() or "addon"
    title = base[:_SHEET_TITLE_MAX_LENGTH]
    number = 1
    while title.lower() in used:
        number += 1
        suffix = f" ({number})"
        title = base[:_SHEET_TITLE_MAX_LENGTH - len(suffix)] + suffix
    used.add(title.lower())
    return title


def write_addon_mismatch_report(period_mismatches, unmapped_people, output_path, source_files=()):
    """Будує output_path (mis_match.xlsx): окремий аркуш на кожен файл addons
    (назва аркуша - ім'я файлу, записи - за "source_file") з розбіжностями
    періодів у стилі error_mis_statuses.xlsx - по два рядки, зелена й червона
    заливка (_write_period_sheet). source_files - усі звірені файли: файл без
    розбіжностей теж отримує аркуш з написом _NO_MISMATCHES_TEXT, щоб було
    видно, що його звірено. Якщо є нерозпізнаний статус Алф2 ("самостріл"),
    додається окремий аркуш "Застереження" (текстом, а не за датами;
    read_unmapped_alf2_people).

    Повертає output_path або None, якщо показувати нічого (ні розбіжностей,
    ні попереджень) - як і write_mismatch_report."""
    if not period_mismatches and not unmapped_people:
        return None

    by_file = {source_file: [] for source_file in source_files}
    for record in period_mismatches:
        by_file.setdefault(record["source_file"], []).append(record)

    wb = empty_workbook()
    used_titles = set()
    for source_file, records in by_file.items():
        _write_period_sheet(wb.create_sheet(_sheet_title(source_file, used_titles)), records)

    if unmapped_people:
        _write_unmapped_sheet(wb.create_sheet(_sheet_title(_UNMAPPED_SHEET_TITLE, used_titles)), unmapped_people)

    save_workbook(wb, output_path)
    return output_path


def _write_period_sheet(ws, records):
    """Аркуш одного файлу addons: Джерело/Підрозділ/ПІБ і колонка на кожну
    дату його записів (_write_period_mismatch_rows); без записів - лише шапка
    й _NO_MISMATCHES_TEXT."""
    all_dates = sorted({date_value for record in records for date_value in record["dates"]})

    date_headers = [datetime.combine(d, datetime.min.time()) for d in all_dates]
    headers = ["Джерело", "Підрозділ", "ПІБ"] + date_headers
    for col_idx, value in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx, value=value)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.alignment = _HEADER_ALIGNMENT
        cell.border = _BORDER
        if isinstance(value, datetime):
            cell.number_format = DATE_NUMBER_FORMAT
    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 45

    date_col_by_date = {d: idx for idx, d in enumerate(all_dates, start=4)}
    if records:
        next_row = _write_period_mismatch_rows(ws, records, 2, date_col_by_date)
    else:
        ws.cell(row=2, column=1, value=_NO_MISMATCHES_TEXT)
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(headers))
        next_row = 3

    for row in range(2, next_row):
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=row, column=col_idx)
            cell.font = _FONT
            cell.alignment = PIB_ALIGNMENT if col_idx == 3 and records else _CENTER
            cell.border = _BORDER

    ws.column_dimensions["A"].width = max([30] + [len(_source_label(record)) + 2 for record in records])
    ws.column_dimensions["B"].width = 26
    ws.column_dimensions["C"].width = pib_column_width(
        [record["roster_pib_raw"] for record in records] + [record["addon_pib_raw"] for record in records]
    )
    for col_idx in range(4, len(headers) + 1):
        ws.column_dimensions[get_column_letter(col_idx)].width = 24


def _write_unmapped_sheet(warn_ws, unmapped_people):
    for col_idx, value in enumerate(_UNMAPPED_WARNING_HEADERS, start=1):
        cell = warn_ws.cell(row=1, column=col_idx, value=value)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.alignment = _HEADER_ALIGNMENT
        cell.border = _BORDER
    warn_ws.freeze_panes = "A2"
    warn_row = 2
    for person in unmapped_people:
        warn_ws.cell(row=warn_row, column=1, value=_display_value(person["roster_pidrozdil"]))
        warn_ws.cell(row=warn_row, column=2, value=person["pib_raw"])
        warn_ws.cell(row=warn_row, column=3, value=person["alf2_status"])
        warn_ws.cell(row=warn_row, column=4, value=(
            "Статус book61 (Алф2) не зіставлено з жодним статусом ОБЛІК у довіднику - перевірити вручну."
        ))
        warn_row += 1
    for row in range(2, warn_row):
        for col_idx in range(1, len(_UNMAPPED_WARNING_HEADERS) + 1):
            cell = warn_ws.cell(row=row, column=col_idx)
            cell.font = _FONT
            cell.alignment = _CENTER
            cell.border = _BORDER
    for col_idx, width in enumerate(_UNMAPPED_WARNING_COLUMN_WIDTHS, start=1):
        warn_ws.column_dimensions[get_column_letter(col_idx)].width = width
