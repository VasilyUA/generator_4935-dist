from datetime import date, datetime

import pytest
from openpyxl import Workbook, load_workbook

import content.oblik_timesheet as ot
import content.rop_vop_statement as rvs
from content.oblik_timesheet import normalize_name


def _write_workbook(path, header, rows):
    wb = Workbook()
    ws = wb.active
    ws.title = "Табель"
    for col_idx, value in enumerate(header, start=1):
        ws.cell(row=1, column=col_idx, value=value)
    for row_idx, row_values in enumerate(rows, start=2):
        for col_idx, value in enumerate(row_values, start=1):
            ws.cell(row=row_idx, column=col_idx, value=value)
    wb.save(str(path))
    return str(path)


_HEADER = [
    "ПОСАДА", "ЗВАННЯ", "ПІБ",
    datetime(2026, 7, 31), datetime(2026, 8, 1), datetime(2026, 8, 2), datetime(2026, 8, 3),
]


@pytest.fixture
def base_workbook(tmp_path):
    return _write_workbook(tmp_path / "ОБЛІК.xlsx", _HEADER, [
        ["Посада1", "сержант", "ПЕРШИЙ Перший Перший", 70, 70, 170, 30],
        ["Посада2", "матрос", "ДРУГИЙ Другий Другий", "РВЗ", 100, 100, None],
    ])


def test_build_statement_rows_marks_rop_and_vop_days(base_workbook):
    """70 -> "роп", 170 -> "воп", будь-яке інше значення (30/100/текст) -
    порожньо (немає запису в day_marks)."""
    sheet = ot.Timesheet(base_workbook, "Табель")

    records, days_in_month = rvs.build_statement_rows(sheet, 2026, 8)

    assert days_in_month == 31
    assert len(records) == 1
    record = records[0]
    assert record["pib_raw"] == "ПЕРШИЙ Перший Перший"
    assert record["posada"] == "Посада1"
    assert record["zvannya"] == "сержант"
    assert record["day_marks"] == {1: "роп", 2: "воп"}  # 01.08=70, 02.08=170, 03.08=30 (порожньо)
    assert record["rop_count"] == 1
    assert record["vop_count"] == 1


def test_build_statement_rows_excludes_person_without_any_rop_or_vop_day(base_workbook):
    """ДРУГИЙ - лише 100/None за серпень (жодного 70/170) - НЕ входить у
    результат, за прямою вказівкою користувача (той самий фільтр, що й у
    реальному зразку - resources/1бмпВОП-РОП_ЛИПЕНЬ_.xlsx)."""
    sheet = ot.Timesheet(base_workbook, "Табель")

    records, _days_in_month = rvs.build_statement_rows(sheet, 2026, 8)

    assert all(record["pib_raw"] != "ДРУГИЙ Другий Другий" for record in records)


def test_build_statement_rows_ignores_the_baseline_column_from_a_different_month(base_workbook):
    """Базова колонка (31.07 - тут теж 70 для ПЕРШОГО) НЕ входить у місяць
    "серпень" - лічильники/позначки рахують ЛИШЕ дні цільового місяця."""
    sheet = ot.Timesheet(base_workbook, "Табель")

    records, _days_in_month = rvs.build_statement_rows(sheet, 2026, 8)

    record = next(r for r in records if r["pib_raw"] == "ПЕРШИЙ Перший Перший")
    assert 31 not in record["day_marks"]  # 31.07 - не в серпні, не рахується як "31"


def test_build_statement_rows_returns_days_in_month_for_february(base_workbook):
    """calendar.monthrange - days_in_month враховує рік (лютий 2026 - НЕ
    високосний, 28 днів)."""
    sheet = ot.Timesheet(base_workbook, "Табель")

    _records, days_in_month = rvs.build_statement_rows(sheet, 2026, 2)

    assert days_in_month == 28


def test_write_statement_header_and_sheet_name_contain_month_and_year(tmp_path):
    records = [{
        "pib_raw": "ПЕРШИЙ Перший Перший", "posada": "Посада1", "zvannya": "сержант",
        "day_marks": {1: "роп", 2: "воп"}, "rop_count": 1, "vop_count": 1,
    }]
    output_path = tmp_path / "output" / "Відомість.xlsx"

    rvs.write_statement(records, 31, str(output_path), 2026, 8)

    wb = load_workbook(str(output_path))
    assert wb.sheetnames == ["СЕРПЕНЬ"]
    ws = wb["СЕРПЕНЬ"]
    assert "за СЕРПЕНЬ 2026 року" in ws.cell(row=4, column=1).value


def test_write_statement_row_shows_rop_and_vop_marks_and_counts(tmp_path):
    records = [{
        "pib_raw": "ПЕРШИЙ Перший Перший", "posada": "Посада1", "zvannya": "сержант",
        "day_marks": {1: "роп", 2: "воп"}, "rop_count": 1, "vop_count": 1,
    }]
    output_path = tmp_path / "output" / "Відомість.xlsx"

    rvs.write_statement(records, 31, str(output_path), 2026, 8)

    ws = load_workbook(str(output_path))["СЕРПЕНЬ"]
    header_row = 6
    data_row = 7
    assert ws.cell(row=data_row, column=1).value == 1  # № з/п
    assert ws.cell(row=data_row, column=2).value == "сержант"
    assert ws.cell(row=data_row, column=3).value == "Посада1"
    assert ws.cell(row=data_row, column=4).value == "ПЕРШИЙ Перший Перший"
    assert ws.cell(row=data_row, column=5).value == "роп"  # день 1
    assert ws.cell(row=data_row, column=6).value == "воп"  # день 2
    assert ws.cell(row=data_row, column=7).value is None   # день 3 - порожньо
    # 4 колонки-мітки + 31 день = 35; далі 5 лічильників (ВОП/ТО/ТМП/ТП/РОП).
    assert ws.cell(row=data_row, column=36).value == 1  # ВОП
    assert ws.cell(row=data_row, column=37).value == 0  # ТО - завжди 0
    assert ws.cell(row=data_row, column=38).value == 0  # ТМП - завжди 0
    assert ws.cell(row=data_row, column=39).value == 0  # ТП - завжди 0
    assert ws.cell(row=data_row, column=40).value == 1  # РОП
    assert ws.cell(row=header_row, column=36).value == "кількість\n днів на ВОП"


def test_write_statement_wraps_permission_error_with_a_friendly_message(tmp_path, monkeypatch):
    """Той самий принцип, що й Timesheet.save/export_for_money_project -
    цільовий файл зайнятий іншою програмою дає зрозуміле повідомлення, а не
    сирий traceback."""
    output_path = tmp_path / "output" / "Відомість.xlsx"

    def _raise(*args, **kwargs):
        raise PermissionError("[WinError 32] заблоковано іншою програмою")

    monkeypatch.setattr(Workbook, "save", _raise)

    with pytest.raises(PermissionError, match="спробуйте ще раз"):
        rvs.write_statement([], 31, str(output_path), 2026, 8)


def test_write_statement_with_no_records_still_writes_header_and_signature_block(tmp_path):
    """За прямою вказівкою користувача - формальний документ, завжди
    зберігається (на відміну від error_mis_statuses.xlsx), навіть якщо
    ЖОДНА людина не мала РОП/ВОП цього місяця - показує лише заголовок і
    підписний блок."""
    output_path = tmp_path / "output" / "Відомість.xlsx"

    rvs.write_statement([], 31, str(output_path), 2026, 8)

    ws = load_workbook(str(output_path))["СЕРПЕНЬ"]
    assert ws.cell(row=6, column=1).value == "№\nз/п"
    assert any(
        cell.value and "Начальник штабу" in str(cell.value)
        for row in ws.iter_rows()
        for cell in row
    )


def test_find_schedule_mismatches_no_record_when_all_scheduled_days_agree(base_workbook):
    """ПЕРШИЙ - 01.08=70 (роп), 02.08=170 (воп) обчислено - schedule каже те
    саме на ОБИДВА ці дні -> жодного запису розбіжностей."""
    sheet = ot.Timesheet(base_workbook, "Табель")
    schedule_marks = {
        (normalize_name("ПЕРШИЙ Перший Перший"), 1): "роп",
        (normalize_name("ПЕРШИЙ Перший Перший"), 2): "воп",
    }

    records = rvs.find_schedule_mismatches(sheet, 2026, 8, schedule_marks, schedule_people={})

    assert records == []


def test_find_schedule_mismatches_detects_schedule_value_with_no_computed_value(base_workbook):
    """03.08 - обчислено 30 (порожньо для Відомості), schedule каже "роп" -
    розбіжність, обчислена сторона - None. roster_*/schedule_* - Посада/
    Звання/ПІБ з КОЖНОЇ сторони окремо (за прямою вказівкою користувача -
    щоб звіт міг показати обидва боки, а не лише один спільний напис)."""
    sheet = ot.Timesheet(base_workbook, "Табель")
    schedule_marks = {(normalize_name("ПЕРШИЙ Перший Перший"), 3): "роп"}
    schedule_people = {
        normalize_name("ПЕРШИЙ Перший Перший"): {
            "pib_raw": "ПЕРШИЙ Перший Перший", "posada": "Навідник", "zvannya": "сержант", "pidrozdil": None,
        },
    }

    records = rvs.find_schedule_mismatches(sheet, 2026, 8, schedule_marks, schedule_people)

    assert len(records) == 1
    record = records[0]
    assert record["roster_pib_raw"] == "ПЕРШИЙ Перший Перший"
    assert record["roster_posada"] == "Посада1"
    assert record["roster_zvannya"] == "сержант"
    assert record["schedule_pib_raw"] == "ПЕРШИЙ Перший Перший"
    assert record["schedule_posada"] == "Навідник"
    assert record["schedule_zvannya"] == "сержант"
    assert record["pidrozdil"] is None
    assert record["dates"] == {date(2026, 8, 3): (None, "роп")}
    assert record["full_days"] == {date(2026, 8, 3): (None, "роп")}


def test_find_schedule_mismatches_pidrozdil_comes_from_roster_for_matched_person(tmp_path):
    """pidrozdil - ІНФОРМАЦІЙНЕ поле (НЕ звіряється роcтер/schedule): для
    ЗНАЙДЕНОЇ в роcтері людини - це роcтерове значення колонки ПІДРОЗДІЛ (не
    порівнюється зі schedule, бо schedule НЕ веде "рідну" підрозділову
    приналежність для КОЖНОЇ людини). recognized_subdivisions - СИНТЕТИЧНИЙ
    набір (а не RECOGNIZED_SUBDIVISIONS за замовчуванням, constants.py) -
    тест НЕ має залежати від реального, керованого користувачем вмісту
    generator_br_and_report_for_money/resources/data.json."""
    header = ["ПІДРОЗДІЛ", "ПОСАДА", "ЗВАННЯ", "ПІБ", datetime(2026, 8, 1)]
    roster_path = _write_workbook(tmp_path / "ОБЛІК.xlsx", header, [
        ["альфа", "Посада1", "сержант", "ПЕРШИЙ Перший Перший", "РВЗ"],
    ])
    sheet = ot.Timesheet(roster_path, "Табель")
    schedule_marks = {(normalize_name("ПЕРШИЙ Перший Перший"), 1): "роп"}

    records = rvs.find_schedule_mismatches(
        sheet, 2026, 8, schedule_marks, schedule_people={}, recognized_subdivisions={"альфа"},
    )

    assert len(records) == 1
    assert records[0]["pidrozdil"] == "альфа"


def test_find_schedule_mismatches_excludes_matched_person_with_unrecognized_pidrozdil(tmp_path):
    """За прямою вказівкою користувача - людина, чиє роcтерове ПІДРОЗДІЛ НЕ
    входить у recognized_subdivisions, взагалі НЕ порівнюється по днях,
    НАВІТЬ якщо є справжня розбіжність - виключається ЦІЛКОМ, а не просто
    показується інакше."""
    header = ["ПІДРОЗДІЛ", "ПОСАДА", "ЗВАННЯ", "ПІБ", datetime(2026, 8, 1)]
    roster_path = _write_workbook(tmp_path / "ОБЛІК.xlsx", header, [
        ["браво", "Посада1", "сержант", "ПЕРШИЙ Перший Перший", 70],
    ])
    sheet = ot.Timesheet(roster_path, "Табель")
    schedule_marks = {(normalize_name("ПЕРШИЙ Перший Перший"), 1): "воп"}  # роп(70) != воп - інакше булo б розбіжністю

    records = rvs.find_schedule_mismatches(
        sheet, 2026, 8, schedule_marks, schedule_people={}, recognized_subdivisions={"альфа"},
    )

    assert records == []


def test_find_schedule_mismatches_does_not_apply_filter_when_pidrozdil_column_absent(base_workbook):
    """Роcтер (base_workbook) узагалі НЕ має колонки ПІДРОЗДІЛ - фільтр НЕ
    застосовується, НАВІТЬ якщо recognized_subdivisions - порожній набір
    (немає даних - немає підстави виключати)."""
    sheet = ot.Timesheet(base_workbook, "Табель")
    schedule_marks = {(normalize_name("ПЕРШИЙ Перший Перший"), 1): "воп"}

    records = rvs.find_schedule_mismatches(
        sheet, 2026, 8, schedule_marks, schedule_people={}, recognized_subdivisions=frozenset(),
    )

    assert len(records) == 1


def test_find_schedule_mismatches_does_not_apply_filter_when_pidrozdil_cell_is_blank(tmp_path):
    """Колонка ПІДРОЗДІЛ Є, але ЦЯ клітинка порожня (None) - фільтр НЕ
    застосовується (те саме "відсутнє - не помилкове", що й для
    колонки, якої взагалі немає)."""
    header = ["ПІДРОЗДІЛ", "ПОСАДА", "ЗВАННЯ", "ПІБ", datetime(2026, 8, 1)]
    roster_path = _write_workbook(tmp_path / "ОБЛІК.xlsx", header, [
        [None, "Посада1", "сержант", "ПЕРШИЙ Перший Перший", "РВЗ"],
    ])
    sheet = ot.Timesheet(roster_path, "Табель")
    schedule_marks = {(normalize_name("ПЕРШИЙ Перший Перший"), 1): "роп"}

    records = rvs.find_schedule_mismatches(
        sheet, 2026, 8, schedule_marks, schedule_people={}, recognized_subdivisions={"альфа"},
    )

    assert len(records) == 1


def test_find_schedule_mismatches_excluded_person_does_not_appear_as_not_in_roster(tmp_path):
    """Людина, ВИКЛЮЧЕНА через нерозпізнаний ПІДРОЗДІЛ, УСЕ ОДНО рахується
    "знайденою" (matched_normalized) - інакше вона хибно з'явилась би вдруге
    як "Не знайдено в ОБЛІК.xlsx" у другому проході (для запису schedule без
    відповідника в роcтері)."""
    header = ["ПІДРОЗДІЛ", "ПОСАДА", "ЗВАННЯ", "ПІБ", datetime(2026, 8, 1)]
    roster_path = _write_workbook(tmp_path / "ОБЛІК.xlsx", header, [
        ["браво", "Посада1", "сержант", "ПЕРШИЙ Перший Перший", "РВЗ"],
    ])
    sheet = ot.Timesheet(roster_path, "Табель")
    schedule_marks = {(normalize_name("ПЕРШИЙ Перший Перший"), 1): "роп"}

    records = rvs.find_schedule_mismatches(
        sheet, 2026, 8, schedule_marks, schedule_people={}, recognized_subdivisions={"альфа"},
    )

    assert records == []


def test_find_schedule_mismatches_excludes_not_in_roster_person_with_unrecognized_pidrozdil(base_workbook):
    """Реальний випадок - людина не знайдена в роcтері, а дужкова дописка
    прикріплення зі schedule вказує на СТОРОННІЙ підрозділ (напр. "танкова
    рота" - НЕ рота/взвод/батарея ЦЬОГО батальйону) - її відсутність у
    роcтері ОЧІКУВАНА (вона й не має тут бути), а НЕ помилка, тож звірка НЕ
    повинна її показувати взагалі."""
    sheet = ot.Timesheet(base_workbook, "Табель")
    schedule_marks = {(normalize_name("ТРЕТІЙ Третій Третій"), 1): "роп"}
    schedule_people = {
        normalize_name("ТРЕТІЙ Третій Третій"): {
            "pib_raw": "ТРЕТІЙ Третій Третій", "posada": "Навідник", "zvannya": "сержант", "pidrozdil": "танкова рота",
        },
    }

    records = rvs.find_schedule_mismatches(
        sheet, 2026, 8, schedule_marks, schedule_people, recognized_subdivisions={"альфа"},
    )

    assert records == []


def test_find_schedule_mismatches_does_not_apply_filter_to_not_in_roster_person_without_pidrozdil(base_workbook):
    """Людина не знайдена в роcтері, і schedule НЕ дає ЖОДНОЇ дужкової
    дописки прикріплення (pidrozdil is None) - фільтр НЕ застосовується
    (немає даних - немає підстави виключати): це, найімовірніше, СПРАВЖНЯ
    помилка (typo в ПІБ чи людина дійсно відсутня в роcтері), а не хтось
    зі стороннього підрозділу, тож запис усе одно потрапляє в результат."""
    sheet = ot.Timesheet(base_workbook, "Табель")
    schedule_marks = {(normalize_name("ТРЕТІЙ Третій Третій"), 1): "роп"}
    schedule_people = {
        normalize_name("ТРЕТІЙ Третій Третій"): {
            "pib_raw": "ТРЕТІЙ Третій Третій", "posada": "Навідник", "zvannya": "сержант", "pidrozdil": None,
        },
    }

    records = rvs.find_schedule_mismatches(
        sheet, 2026, 8, schedule_marks, schedule_people, recognized_subdivisions={"альфа"},
    )

    assert len(records) == 1


def test_find_schedule_mismatches_full_days_includes_matches_alongside_mismatches(base_workbook):
    """full_days - ПОВНА мапа порівнюваних днів (де schedule непорожній),
    включно ЗІ ЗБІЖНИМИ - на відміну від "dates" (лише розбіжні). Потрібно
    write_schedule_mismatch_report, щоб показати збіжний день зеленим, а не
    порожньою клітинкою, навіть якщо в людини є лише ОДИН розбіжний день."""
    sheet = ot.Timesheet(base_workbook, "Табель")
    schedule_marks = {
        (normalize_name("ПЕРШИЙ Перший Перший"), 1): "роп",  # збігається (70=роп)
        (normalize_name("ПЕРШИЙ Перший Перший"), 2): "роп",  # розбіжність (170=воп, а не роп)
    }

    records = rvs.find_schedule_mismatches(sheet, 2026, 8, schedule_marks, schedule_people={})

    assert len(records) == 1
    record = records[0]
    assert record["dates"] == {date(2026, 8, 2): ("воп", "роп")}
    assert record["full_days"] == {
        date(2026, 8, 1): ("роп", "роп"),
        date(2026, 8, 2): ("воп", "роп"),
    }


def test_find_schedule_mismatches_ignores_day_where_schedule_has_no_entry(base_workbook):
    """За прямою вказівкою користувача - ПЕРШИЙ має обчислені "роп"/"воп" на
    01.08/02.08, але schedule_marks взагалі не має записів на ці дні (ще не
    заповнено в "офіційній" відомості) - день, ВІДСУТНІЙ у schedule, НЕ
    звіряється, навіть якщо обчислено щось непорожнє."""
    sheet = ot.Timesheet(base_workbook, "Табель")

    records = rvs.find_schedule_mismatches(sheet, 2026, 8, schedule_marks={}, schedule_people={})

    assert records == []


def test_find_schedule_mismatches_detects_disagreement_between_rop_and_vop(base_workbook):
    """01.08 обчислено 70 (роп), schedule каже "воп" - розбіжність з ОБОМА
    непорожніми, але РІЗНИМИ значеннями."""
    sheet = ot.Timesheet(base_workbook, "Табель")
    schedule_marks = {(normalize_name("ПЕРШИЙ Перший Перший"), 1): "воп"}

    records = rvs.find_schedule_mismatches(sheet, 2026, 8, schedule_marks, schedule_people={})

    assert len(records) == 1
    assert records[0]["dates"] == {date(2026, 8, 1): ("роп", "воп")}


@pytest.mark.parametrize("schedule_word", ["тмп", "тот"])
def test_find_schedule_mismatches_treats_vop_aliases_as_equivalent_to_computed_vop(base_workbook, schedule_word):
    """За прямою вказівкою користувача - "воп"/"тмп"/"тот" - РІЗНІ слова
    ТІЄЇ САМОЇ категорії виплати 170: 02.08 обчислено 170 (воп), schedule
    каже "тмп"/"тот" - НЕ розбіжність, хоча тексти буквально різні."""
    sheet = ot.Timesheet(base_workbook, "Табель")
    schedule_marks = {(normalize_name("ПЕРШИЙ Перший Перший"), 2): schedule_word}

    records = rvs.find_schedule_mismatches(sheet, 2026, 8, schedule_marks, schedule_people={})

    assert records == []


def test_find_schedule_mismatches_vop_alias_does_not_hide_a_genuine_mismatch_on_another_day(base_workbook):
    """Псевдонім ВОП покриває ЛИШЕ той день, де він застосовний - 02.08
    (воп/тмп - еквівалент) лишається без розбіжності, але 01.08 (роп/тмп -
    СПРАВЖНЯ розбіжність, "тмп" - НЕ псевдонім РОП) усе одно потрапляє в
    результат."""
    sheet = ot.Timesheet(base_workbook, "Табель")
    schedule_marks = {
        (normalize_name("ПЕРШИЙ Перший Перший"), 1): "тмп",
        (normalize_name("ПЕРШИЙ Перший Перший"), 2): "тмп",
    }

    records = rvs.find_schedule_mismatches(sheet, 2026, 8, schedule_marks, schedule_people={})

    assert len(records) == 1
    assert records[0]["dates"] == {date(2026, 8, 1): ("роп", "тмп")}
    assert records[0]["full_days"] == {
        date(2026, 8, 1): ("роп", "тмп"),
        date(2026, 8, 2): ("воп", "тмп"),
    }


def test_find_schedule_mismatches_computed_blank_and_schedule_vop_alias_is_still_a_mismatch(base_workbook):
    """Псевдонім ВОП НЕ поширюється на "взагалі нічого не обчислено" - 03.08
    обчислено 30 (порожньо), schedule каже "тот" - усе одно розбіжність
    (людина, за обчисленням, узагалі не на ВОП цього дня)."""
    sheet = ot.Timesheet(base_workbook, "Табель")
    schedule_marks = {(normalize_name("ПЕРШИЙ Перший Перший"), 3): "тот"}

    records = rvs.find_schedule_mismatches(sheet, 2026, 8, schedule_marks, schedule_people={})

    assert len(records) == 1
    assert records[0]["dates"] == {date(2026, 8, 3): (None, "тот")}


def test_find_schedule_mismatches_schedule_person_with_no_roster_match(base_workbook):
    """ТРЕТІЙ є в resources/schedule, але відсутній у роcтері (ОБЛІК.xlsx) -
    той самий принцип, що й _missing_from_roster_records
    (content/payment_mismatch_checker.py) - роcтерова сторона (Посада/
    Звання/ПІБ) отримує _NOT_IN_ROSTER_LABEL, а schedule-сторона - РЕАЛЬНІ
    дані з schedule_people (щоб було видно, ХТО саме ця людина), включно з
    pidrozdil - дужковою допискою прикріплення."""
    sheet = ot.Timesheet(base_workbook, "Табель")
    schedule_marks = {(normalize_name("ТРЕТІЙ Третій Третій"), 1): "роп"}
    schedule_people = {
        normalize_name("ТРЕТІЙ Третій Третій"): {
            "pib_raw": "ТРЕТІЙ Третій Третій", "posada": "Навідник", "zvannya": "сержант", "pidrozdil": "танкова рота",
        },
    }

    records = rvs.find_schedule_mismatches(
        sheet, 2026, 8, schedule_marks, schedule_people, recognized_subdivisions={"танкова рота"},
    )

    assert len(records) == 1
    record = records[0]
    assert record["roster_pib_raw"] == rvs._NOT_IN_ROSTER_LABEL
    assert record["roster_posada"] == rvs._NOT_IN_ROSTER_LABEL
    assert record["roster_zvannya"] == rvs._NOT_IN_ROSTER_LABEL
    assert record["schedule_pib_raw"] == "ТРЕТІЙ Третій Третій"
    assert record["schedule_posada"] == "Навідник"
    assert record["schedule_zvannya"] == "сержант"
    assert record["pidrozdil"] == "танкова рота"
    assert record["dates"] == {date(2026, 8, 1): (None, "роп")}
    assert record["full_days"] == record["dates"]


def test_find_schedule_mismatches_ignores_schedule_day_beyond_month_length_for_matched_person(base_workbook):
    """Захисний випадок - день у schedule_marks, що перевищує кількість днів
    ЦІЛЬОВОГО місяця (тут - лютий 2026, 28 днів), НЕ спричиняє помилку -
    просто пропускається, як і будь-який інший непорівнюваний день."""
    sheet = ot.Timesheet(base_workbook, "Табель")
    schedule_marks = {(normalize_name("ПЕРШИЙ Перший Перший"), 29): "роп"}

    records = rvs.find_schedule_mismatches(sheet, 2026, 2, schedule_marks, schedule_people={})

    assert records == []


def test_find_schedule_mismatches_ignores_schedule_day_beyond_month_length_for_unmatched_person(base_workbook):
    """Той самий захист - для запису БЕЗ відповідника в роcтері."""
    sheet = ot.Timesheet(base_workbook, "Табель")
    schedule_marks = {(normalize_name("ТРЕТІЙ Третій Третій"), 29): "роп"}

    records = rvs.find_schedule_mismatches(sheet, 2026, 2, schedule_marks, schedule_people={})

    assert records == []


def test_write_schedule_mismatch_report_wraps_permission_error_with_a_friendly_message(tmp_path, monkeypatch):
    """Той самий принцип, що й write_statement - цільовий файл зайнятий
    іншою програмою дає зрозуміле повідомлення, а не сирий traceback."""
    records = [{
        "roster_posada": "Посада1", "schedule_posada": "Посада1",
        "roster_zvannya": "сержант", "schedule_zvannya": "сержант",
        "roster_pib_raw": "ПЕРШИЙ Перший Перший", "schedule_pib_raw": "ПЕРШИЙ Перший Перший",
        "pidrozdil": None,
        "dates": {date(2026, 8, 1): ("роп", "воп")},
        "full_days": {date(2026, 8, 1): ("роп", "воп")},
    }]
    output_path = tmp_path / "output" / "Розбіжності.xlsx"

    def _raise(*args, **kwargs):
        raise PermissionError("[WinError 32] заблоковано іншою програмою")

    monkeypatch.setattr(Workbook, "save", _raise)

    with pytest.raises(PermissionError, match="спробуйте ще раз"):
        rvs.write_schedule_mismatch_report(records, str(output_path))


def test_write_schedule_mismatch_report_returns_none_for_empty_records(tmp_path):
    """Той самий принцип, що й write_mismatch_report - немає сенсу
    створювати файл лише із заголовком."""
    output_path = tmp_path / "output" / "Розбіжності.xlsx"

    assert rvs.write_schedule_mismatch_report([], str(output_path)) is None


def test_write_schedule_mismatch_report_writes_columns_rows_and_fills(tmp_path):
    """ПЕРШИЙ - знайдений у роcтері, roster/schedule збігаються по Посада/
    Звання/ПІБ (зелено) - Підрозділ показує роcтерове значення (без
    заливки/порівняння); 01.08 - день теж збігається (зелено), 02.08 -
    розбіжність (червоно, обидва значення)."""
    records = [{
        "roster_posada": "Посада1", "schedule_posada": "Посада1",
        "roster_zvannya": "сержант", "schedule_zvannya": "сержант",
        "roster_pib_raw": "ПЕРШИЙ Перший Перший", "schedule_pib_raw": "ПЕРШИЙ Перший Перший",
        "pidrozdil": "підрозділ 1",
        "dates": {
            date(2026, 8, 1): ("роп", "роп"),
            date(2026, 8, 2): ("роп", "воп"),
        },
        "full_days": {
            date(2026, 8, 1): ("роп", "роп"),
            date(2026, 8, 2): ("роп", "воп"),
        },
    }]
    output_path = tmp_path / "output" / "Розбіжності.xlsx"

    result_path = rvs.write_schedule_mismatch_report(records, str(output_path))

    assert result_path == str(output_path)
    ws = load_workbook(str(output_path))["Розбіжності"]
    assert ws.cell(row=1, column=1).value == "Джерело"
    assert ws.cell(row=1, column=2).value == "Підрозділ"
    assert ws.cell(row=1, column=3).value == "Посада"
    assert ws.cell(row=1, column=4).value == "Звання"
    assert ws.cell(row=1, column=5).value == "ПІБ"
    assert ws.cell(row=2, column=1).value == rvs._MISMATCH_COMPUTED_LABEL
    assert ws.cell(row=3, column=1).value == rvs._MISMATCH_SCHEDULE_LABEL
    assert ws.cell(row=2, column=2).value == "підрозділ 1"
    assert ws.cell(row=2, column=3).value == "Посада1"
    assert ws.cell(row=2, column=3).fill.start_color.rgb == "FFC6EFCE"
    assert ws.cell(row=2, column=4).value == "сержант"
    assert ws.cell(row=2, column=5).value == "ПЕРШИЙ Перший Перший"
    assert ws.cell(row=2, column=5).fill.start_color.rgb == "FFC6EFCE"
    # 01.08 - обидва боки збігаються ("роп"/"роп") - одне значення, зелена заливка.
    assert ws.cell(row=2, column=6).value == "роп"
    assert ws.cell(row=2, column=6).fill.start_color.rgb == "FFC6EFCE"
    # 02.08 - розбіжність ("роп"/"воп") - обидва значення показані, червона заливка.
    assert ws.cell(row=2, column=7).value == "роп"
    assert ws.cell(row=3, column=7).value == "воп"
    assert ws.cell(row=2, column=7).fill.start_color.rgb == "FFFFC7CE"
    assert ws.cell(row=3, column=7).fill.start_color.rgb == "FFFFC7CE"


def test_write_schedule_mismatch_report_shows_both_sides_when_label_fields_differ(tmp_path):
    """За прямою вказівкою користувача - людина, відсутня в роcтері, БІЛЬШЕ
    не показує ОДИН спільний напис "Не знайдено в ОБЛІК.xlsx" на 2
    об'єднані клітинки: роcтерова сторона (немає даних) і schedule-сторона
    (РЕАЛЬНІ Посада/Звання/ПІБ) показуються ОКРЕМО, червоним - щоб було
    видно, ХТО саме ця людина. Підрозділ - інформаційно (дужкова дописка
    прикріплення зі schedule), БЕЗ заливки/порівняння."""
    records = [{
        "roster_posada": rvs._NOT_IN_ROSTER_LABEL, "schedule_posada": "Навідник",
        "roster_zvannya": rvs._NOT_IN_ROSTER_LABEL, "schedule_zvannya": "сержант",
        "roster_pib_raw": rvs._NOT_IN_ROSTER_LABEL, "schedule_pib_raw": "ТРЕТІЙ Третій Третій",
        "pidrozdil": "танкова рота",
        "dates": {date(2026, 8, 1): (None, "роп")},
        "full_days": {date(2026, 8, 1): (None, "роп")},
    }]
    output_path = tmp_path / "output" / "Розбіжності.xlsx"

    rvs.write_schedule_mismatch_report(records, str(output_path))

    ws = load_workbook(str(output_path))["Розбіжності"]
    assert ws.cell(row=2, column=2).value == "танкова рота"
    assert ws.cell(row=2, column=3).value == rvs._NOT_IN_ROSTER_LABEL
    assert ws.cell(row=3, column=3).value == "Навідник"
    assert ws.cell(row=2, column=3).fill.start_color.rgb == "FFFFC7CE"
    assert ws.cell(row=3, column=3).fill.start_color.rgb == "FFFFC7CE"
    assert ws.cell(row=2, column=4).value == rvs._NOT_IN_ROSTER_LABEL
    assert ws.cell(row=3, column=4).value == "сержант"
    assert ws.cell(row=2, column=5).value == rvs._NOT_IN_ROSTER_LABEL
    assert ws.cell(row=3, column=5).value == "ТРЕТІЙ Третій Третій"
    assert ws.cell(row=2, column=5).fill.start_color.rgb == "FFFFC7CE"
    assert ws.cell(row=3, column=5).fill.start_color.rgb == "FFFFC7CE"


def test_write_schedule_mismatch_report_pidrozdil_placeholder_when_none(tmp_path):
    """Підрозділ - інформаційна колонка, тож "немає даних" ТЕЖ показується
    як "—" (за прямою вказівкою користувача - у таблиці немає порожніх
    клітинок), просто БЕЗ заливки (не звіряється)."""
    records = [{
        "roster_posada": "Посада1", "schedule_posada": "Посада1",
        "roster_zvannya": "сержант", "schedule_zvannya": "сержант",
        "roster_pib_raw": "ПЕРШИЙ Перший Перший", "schedule_pib_raw": "ПЕРШИЙ Перший Перший",
        "pidrozdil": None,
        "dates": {date(2026, 8, 3): (None, "роп")},
        "full_days": {date(2026, 8, 3): (None, "роп")},
    }]
    output_path = tmp_path / "output" / "Розбіжності.xlsx"

    rvs.write_schedule_mismatch_report(records, str(output_path))

    ws = load_workbook(str(output_path))["Розбіжності"]
    assert ws.cell(row=2, column=2).value == "—"
    assert ws.cell(row=2, column=6).value == "—"
    assert ws.cell(row=3, column=6).value == "роп"


def test_write_schedule_mismatch_report_shows_vop_alias_day_as_a_green_match(tmp_path):
    """За прямою вказівкою користувача - "воп"/"тмп" РІЗНІ слова тієї самої
    категорії 170: 02.08 - колонка звіту через розбіжність ДРУГОГО, але для
    ПЕРШОГО на цей день "воп"/"тмп" - еквівалент, тож зелена клітинка (не
    сіра/червона) зі значенням SCHEDULE, як подано ("тмп" - НЕ замінюється
    на "воп")."""
    records = [
        {
            "roster_posada": "Посада1", "schedule_posada": "Посада1",
            "roster_zvannya": "сержант", "schedule_zvannya": "сержант",
            "roster_pib_raw": "ПЕРШИЙ Перший Перший", "schedule_pib_raw": "ПЕРШИЙ Перший Перший",
            "pidrozdil": None,
            "dates": {date(2026, 8, 1): ("роп", "воп")},
            "full_days": {
                date(2026, 8, 1): ("роп", "воп"),
                date(2026, 8, 2): ("воп", "тмп"),
            },
        },
        {
            "roster_posada": "Посада2", "schedule_posada": "Посада2",
            "roster_zvannya": "матрос", "schedule_zvannya": "матрос",
            "roster_pib_raw": "ДРУГИЙ Другий Другий", "schedule_pib_raw": "ДРУГИЙ Другий Другий",
            "pidrozdil": None,
            "dates": {date(2026, 8, 2): ("роп", "щось інше")},
            "full_days": {date(2026, 8, 2): ("роп", "щось інше")},
        },
    ]
    output_path = tmp_path / "output" / "Розбіжності.xlsx"

    rvs.write_schedule_mismatch_report(records, str(output_path))

    ws = load_workbook(str(output_path))["Розбіжності"]
    # Колонки: 6 = 01.08, 7 = 02.08.
    assert ws.cell(row=2, column=7).value == "тмп"
    assert ws.cell(row=2, column=7).fill.start_color.rgb == "FFC6EFCE"


def test_write_schedule_mismatch_report_has_no_blank_cells_and_marks_not_compared_days_neutrally(tmp_path):
    """За прямою вказівкою користувача - у таблиці НЕМАЄ порожніх клітинок.
    Колонки звіту - 01.08/03.08 (розбіжні для ДРУГОГО) і 02.08 (розбіжний
    для ПЕРШОГО). ПЕРШИЙ збігається на 01.08 (зелений, зі значенням, хоча
    ЦЯ дата - колонка через розбіжність ДРУГОГО, а НЕ його власна) і взагалі
    не звірявся на 03.08 - нейтральна сіра клітинка з "—", а НЕ порожня і НЕ
    зелена/червона."""
    records = [
        {
            "roster_posada": "Посада1", "schedule_posada": "Посада1",
            "roster_zvannya": "сержант", "schedule_zvannya": "сержант",
            "roster_pib_raw": "ПЕРШИЙ Перший Перший", "schedule_pib_raw": "ПЕРШИЙ Перший Перший",
            "pidrozdil": None,
            "dates": {date(2026, 8, 2): ("воп", "роп")},
            "full_days": {
                date(2026, 8, 1): ("роп", "роп"),
                date(2026, 8, 2): ("воп", "роп"),
            },
        },
        {
            "roster_posada": "Посада2", "schedule_posada": "Посада2",
            "roster_zvannya": "матрос", "schedule_zvannya": "матрос",
            "roster_pib_raw": "ДРУГИЙ Другий Другий", "schedule_pib_raw": "ДРУГИЙ Другий Другий",
            "pidrozdil": None,
            "dates": {
                date(2026, 8, 1): (None, "воп"),
                date(2026, 8, 3): (None, "роп"),
            },
            "full_days": {
                date(2026, 8, 1): (None, "воп"),
                date(2026, 8, 3): (None, "роп"),
            },
        },
    ]
    output_path = tmp_path / "output" / "Розбіжності.xlsx"

    rvs.write_schedule_mismatch_report(records, str(output_path))

    ws = load_workbook(str(output_path))["Розбіжності"]
    # Колонки (сортовані дати): 6 = 01.08, 7 = 02.08, 8 = 03.08.
    # ПЕРШИЙ (рядки 2-3): 01.08 - збіг (зелений), 02.08 - розбіжність
    # (червоний), 03.08 - не звірявся взагалі (нейтральний сірий).
    assert ws.cell(row=2, column=6).value == "роп"
    assert ws.cell(row=2, column=6).fill.start_color.rgb == "FFC6EFCE"
    assert ws.cell(row=2, column=7).value == "воп"
    assert ws.cell(row=3, column=7).value == "роп"
    assert ws.cell(row=2, column=7).fill.start_color.rgb == "FFFFC7CE"
    assert ws.cell(row=2, column=8).value == "—"
    assert ws.cell(row=2, column=8).fill.start_color.rgb == "FFF2F2F2"
    # ДРУГИЙ (рядки 4-5): 01.08/03.08 - розбіжність (червоний), 02.08 - не
    # звірявся (нейтральний сірий).
    assert ws.cell(row=4, column=6).value == "—"
    assert ws.cell(row=5, column=6).value == "воп"
    assert ws.cell(row=4, column=6).fill.start_color.rgb == "FFFFC7CE"
    assert ws.cell(row=4, column=7).value == "—"
    assert ws.cell(row=4, column=7).fill.start_color.rgb == "FFF2F2F2"
    assert ws.cell(row=4, column=8).value == "—"
    assert ws.cell(row=5, column=8).value == "роп"
    assert ws.cell(row=4, column=8).fill.start_color.rgb == "FFFFC7CE"


def test_write_schedule_mismatch_report_rotates_date_headers_but_not_label_headers(tmp_path):
    """За прямою вказівкою користувача - заголовки дат розвернуті на 90°;
    заголовки міток (Джерело/Підрозділ/Посада/Звання/ПІБ) лишаються
    горизонтальними."""
    records = [{
        "roster_posada": "Посада1", "schedule_posada": "Посада1",
        "roster_zvannya": "сержант", "schedule_zvannya": "сержант",
        "roster_pib_raw": "ПЕРШИЙ Перший Перший", "schedule_pib_raw": "ПЕРШИЙ Перший Перший",
        "pidrozdil": None,
        "dates": {date(2026, 8, 1): ("роп", "воп")},
        "full_days": {date(2026, 8, 1): ("роп", "воп")},
    }]
    output_path = tmp_path / "output" / "Розбіжності.xlsx"

    rvs.write_schedule_mismatch_report(records, str(output_path))

    ws = load_workbook(str(output_path))["Розбіжності"]
    assert ws.cell(row=1, column=6).alignment.text_rotation == 90
    assert ws.cell(row=1, column=1).alignment.text_rotation in (0, None)
    assert ws.cell(row=1, column=2).alignment.text_rotation in (0, None)
    assert ws.cell(row=1, column=5).alignment.text_rotation in (0, None)
