"""Синтетичний resources/addons/book61.xlsm для тестів - рядок 1 заголовки,
рядок 2 суто числові індекси колонок (той самий квірк, що й у реальній
книзі), дані з рядка 3."""

from openpyxl import Workbook

VP_HEADER = ["№", "Підрозділ на час вибуття", "Звання на час вибуття", "П.І.Б.", "БЧС",
             "Дата вибуття", "Дата 08 ранку", "Дата прибуття", "Тип відпустки"]
VD_HEADER = ["№", "Підрозділ", "Звання", "П.І.Б.", "БЧС", "Дата вибуття", "Дата 08 ранку",
             "Дата прибуття", "Виплата"]
SHP_HEADER = ["№", "Підрозділ", "Звання", "П.І.Б.", "БЧС", "Дата вибуття", "Кінець (виписка)"]
VLK_HEADER = ["№", "Підрозділ", "Звання", "П.І.Б.", "БЧС", "Дата вибуття", "Кінець (ВЛК)"]
PPD_HEADER = ["№", "Підрозділ", "Звання", "П.І.Б.", "БЧС", "Дата прибуття на ППД", "Дата вибуття з ППД"]
ALF2_HEADER = ["Індекс особи", "П.І.Б", "СТАТУС", "з", "по", "Дата, наказ розпорядження, виключення"]
SZCH_HEADER = ["№", "Підрозділ", "Звання", "П.І.Б.", "БЧС", "Дата події", "Дата повернення з СЗЧ"]
VTRATY_HEADER = ["№", "Підрозділ", "Звання", "П.І.Б.", "БЧС", "Дата події"]
RREZ_HEADER = ["№", "Індекс посади до якої допустили до ТВО", "Підрозділ куда зарахували",
               "Посада на яку призначили", "Звання", "П.І.Б.", "Дата допущення до ТВО",
               "Дата зарахування в списки", "Дата прийому посади"]

# Усі аркуші, які відкриває content/addon_reader.py (read_all_period_records,
# read_unmapped_alf2_people) - відсутній аркуш дав би KeyError.
ALL_SHEET_HEADERS = {
    "ВП": VP_HEADER, "ВД": VD_HEADER, "ШП": SHP_HEADER, "ВЛК": VLK_HEADER, "ППД": PPD_HEADER,
    "Алф2": ALF2_HEADER, "СЗЧ": SZCH_HEADER, "Втрати": VTRATY_HEADER, "Прибулі у ВД": RREZ_HEADER,
}


def write_book61(path, sheets):
    """sheets: {назва аркуша: (заголовок, [рядок, ...])}."""
    wb = Workbook()
    wb.remove(wb.active)
    for sheet_name, (header, rows) in sheets.items():
        ws = wb.create_sheet(sheet_name)
        for col_idx, value in enumerate(header, start=1):
            ws.cell(row=1, column=col_idx, value=value)
        for col_idx in range(1, len(header) + 1):
            ws.cell(row=2, column=col_idx, value=col_idx)
        for row_idx, row_values in enumerate(rows, start=3):
            for col_idx, value in enumerate(row_values, start=1):
                ws.cell(row=row_idx, column=col_idx, value=value)
    wb.save(str(path))
    return str(path)


def write_full_book61(path, rows_by_sheet):
    """Книга з УСІМА аркушами ALL_SHEET_HEADERS - ті, яких немає в
    rows_by_sheet, лишаються порожніми (лише заголовок)."""
    return write_book61(path, {
        name: (header, rows_by_sheet.get(name, [])) for name, header in ALL_SHEET_HEADERS.items()
    })
