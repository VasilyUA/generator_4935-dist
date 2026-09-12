from datetime import date

import pytest
from docx import Document
from docx_fixtures import ARRIVAL_HEADER, DEPARTURE_HEADER, write_daily_report

import content.daily_report_reader as reader
from constants import _SPECIAL_GROUP_NAME, IGNORED_TABLE_HEADING_KEYWORDS, STATUS_HEADING_RULES


def test_report_date_from_filename():
    assert reader.report_date_from_filename("resources/report/01.08.2026 - щоденний рапорт.docx") == date(2026, 8, 1)
    assert reader.report_date_from_filename(r"C:\reports\08.08.2026 - щоденний рапорт .docx") == date(2026, 8, 8)


def test_report_date_from_filename_returns_none_without_a_date():
    assert reader.report_date_from_filename("щоденний рапорт.docx") is None


def test_latest_report_date_picks_the_maximum_date(tmp_path):
    write_daily_report(tmp_path / "01.08.2026 - щоденний рапорт.docx")
    write_daily_report(tmp_path / "31.08.2026 - щоденний рапорт.docx")
    write_daily_report(tmp_path / "15.08.2026 - щоденний рапорт.docx")

    assert reader.latest_report_date(str(tmp_path)) == date(2026, 8, 31)


def test_latest_report_date_ignores_office_lock_files(tmp_path):
    """Реальний випадок: користувач тримає рапорт СЬОГОДНІШНЬОГО дня
    відкритим у Word - lock-файл ("~$...") НЕ рахується найпізнішим,
    навіть якщо його ІМ'Я (якби воно розпізналось) вказувало б на пізнішу
    дату - той самий принцип, що й generators/generate_timesheet.py::
    _build_timesheet."""
    write_daily_report(tmp_path / "30.08.2026 - щоденний рапорт.docx")
    (tmp_path / "~$1.08.2026 - щоденний рапорт.docx").write_bytes(b"")

    assert reader.latest_report_date(str(tmp_path)) == date(2026, 8, 30)


def test_latest_report_date_ignores_files_without_a_parseable_date(tmp_path):
    write_daily_report(tmp_path / "01.08.2026 - щоденний рапорт.docx")
    write_daily_report(tmp_path / "щоденний рапорт без дати.docx")

    assert reader.latest_report_date(str(tmp_path)) == date(2026, 8, 1)


def test_latest_report_date_returns_none_for_empty_directory(tmp_path):
    assert reader.latest_report_date(str(tmp_path)) is None


def test_return_from_leave_resolves_to_default_status(tmp_path):
    path = write_daily_report(
        tmp_path / "01.08.2026 - щоденний рапорт.docx",
        groups=[(
            "ПРИБУЛИ до складу сил та засобів 9 армійського корпусу:",
            [("З відпустки:", ARRIVAL_HEADER, [["1", "сержант", "ПЕРШИЙ Перший Перший", "Посада", "01.08.2026", "Прибув"]])],
        )],
    )
    events, unresolved = reader.extract_status_events(path)

    assert unresolved == []
    assert events == [{
        "pib_raw": "ПЕРШИЙ Перший Перший", "status": "РВЗ", "effective_date": date(2026, 8, 1),
        "heading": "З відпустки:", "report_date": date(2026, 8, 1),
    }]


def test_return_from_leave_to_ppd_destination_resolves_to_ppd_status_end_to_end(tmp_path):
    """Реальний випадок, підтверджений користувачем: рядок під заголовком
    "З відпустки:" (звичайно - "РВЗ", тест вище) - АЛЕ колонка "Куди прибув"
    прямо каже "...до пункту постійної дислокації..." - людина повернулась
    САМЕ на постійну базу, а не в район виконання завдань - статус ЦЬОГО
    рядка МАЄ бути "ППД", а не "РВЗ"."""
    path = write_daily_report(
        tmp_path / "11.09.2026 - щоденний рапорт.docx",
        groups=[(
            "ПРИБУЛИ до складу сил та засобів 9 армійського корпусу:",
            [("З відпустки:", ARRIVAL_HEADER, [
                ["1", "старший сержант", "СІМДЕСЯТИЙ Сімдесятий Сімдесятий", "Головний сержант", "11.09.2026",
                 "Прибув до пункту постійної дислокації з щорічної відпустки"],
            ])],
        )],
    )
    events, unresolved = reader.extract_status_events(path)

    assert unresolved == []
    assert events == [{
        "pib_raw": "СІМДЕСЯТИЙ Сімдесятий Сімдесятий", "status": "ППД", "effective_date": date(2026, 9, 11),
        "heading": "З відпустки:", "report_date": date(2026, 9, 11),
    }]


def test_ppd_destination_override_requires_the_arrival_direction_not_just_the_words(tmp_path):
    """Реальний випадок (баг, знайдений при перевірці попереднього тесту
    проти реальних даних): рядок під заголовком "В РТГР:", чия колонка
    "Куди вибув" каже "Вибув З пункту постійної дислокації в РТГР" -
    людина ЗАЛИШАЄ ППД (напрямок "з", а не "до") - статус МАЄ лишитись
    "РТГр" за заголовком, а НЕ хибно перетворитись на "ППД" лише через
    збіг слів "постійної дислокації" в ІНШОМУ напрямку."""
    path = write_daily_report(
        tmp_path / "12.09.2026 - щоденний рапорт.docx",
        groups=[(
            "ВИБУЛИ зі складу сил та засобів 9 армійського корпусу:",
            [("В РТГР:", DEPARTURE_HEADER, [
                ["1", "старший сержант", "ВІСІМДЕСЯТИЙ Вісімдесятий Вісімдесятий", "Головний сержант", "12.09.2026",
                 "Вибув з пункту постійної дислокації в РТГР"],
            ])],
        )],
    )
    events, unresolved = reader.extract_status_events(path)

    assert unresolved == []
    assert events == [{
        "pib_raw": "ВІСІМДЕСЯТИЙ Вісімдесятий Вісімдесятий", "status": "РТГр", "effective_date": date(2026, 9, 12),
        "heading": "В РТГР:", "report_date": date(2026, 9, 12),
    }]


def test_return_from_ppd_abbreviated_heading_is_applied_end_to_end(tmp_path):
    """Реальний випадок: таблиця під заголовком "З ППД:" раніше потрапляла в
    unresolved як "невідомий заголовок" - тепер коректно розпізнається й
    застосовується як звичайна подія повернення (РВЗ)."""
    path = write_daily_report(
        tmp_path / "13.08.2026 - щоденний рапорт.docx",
        groups=[(
            "ПРИБУЛИ до складу сил та засобів 9 армійського корпусу:",
            [("З ППД:", ARRIVAL_HEADER, [["1", "сержант", "ПЕРШИЙ Перший Перший", "Посада", "13.08.2026", "Прибув"]])],
        )],
    )
    events, unresolved = reader.extract_status_events(path)

    assert unresolved == []
    assert events == [{
        "pib_raw": "ПЕРШИЙ Перший Перший", "status": "РВЗ", "effective_date": date(2026, 8, 13),
        "heading": "З ППД:", "report_date": date(2026, 8, 13),
    }]


def test_arrival_list_heading_applies_default_status_to_every_listed_person(tmp_path):
    """Реальний випадок: заголовок-СПИСОК "ПРИБУЛИ до пункту постійно
    дислокації військової частини ...:" (БЕЗ таблиці) - кожен наступний
    абзац "звання ПІБ" (той самий стиль "List Paragraph") стає подією РВЗ
    на дату рапорту, автоматично, без ручної перевірки."""
    path = write_daily_report(
        tmp_path / "14.08.2026 - щоденний рапорт.docx",
        arrival_list=(
            "ПРИБУЛИ до пункту постійно дислокації військової частини A0000:",
            ["солдат ПЕРШИЙ Перший Перший", "солдат ДРУГИЙ Другий Другий"],
        ),
    )
    events, unresolved = reader.extract_status_events(path)

    assert unresolved == []
    assert events == [
        {
            "pib_raw": "ПЕРШИЙ Перший Перший", "status": "РВЗ", "effective_date": date(2026, 8, 14),
            "heading": "ПРИБУЛИ до пункту постійно дислокації військової частини A0000:",
            "report_date": date(2026, 8, 14),
        },
        {
            "pib_raw": "ДРУГИЙ Другий Другий", "status": "РВЗ", "effective_date": date(2026, 8, 14),
            "heading": "ПРИБУЛИ до пункту постійно дислокації військової частини A0000:",
            "report_date": date(2026, 8, 14),
        },
    ]


def test_arrival_list_heading_accepts_the_adjectival_wording_too(tmp_path):
    """ARRIVAL_LIST_HEADING_KEYWORDS - ключові слова, не точна фраза -
    "постійної дислокації" (прикметниковий відмінок) має розпізнаватись так
    само, як реальний варіант "постійно дислокації"."""
    path = write_daily_report(
        tmp_path / "14.08.2026 - щоденний рапорт.docx",
        arrival_list=(
            "ПРИБУЛИ до пункту постійної дислокації військової частини A0000:",
            ["солдат ПЕРШИЙ Перший Перший"],
        ),
    )
    events, unresolved = reader.extract_status_events(path)

    assert unresolved == []
    assert events[0]["status"] == "РВЗ"


def test_arrival_list_section_ends_at_the_next_recognized_heading(tmp_path):
    """Реальний випадок: розділ "Поза межами..." одразу ПІСЛЯ списку
    "ПРИБУЛИ..." - заголовок цього розділу (сам НЕ схожий на ПІБ) має
    коректно завершити список прибулих і розпізнатись як звичайний перехід
    до розділу "Поза межами...", а не хибно "проковтнутись" як ще один
    рядок списку прибулих."""
    path = write_daily_report(
        tmp_path / "14.08.2026 - щоденний рапорт.docx",
        arrival_list=(
            "ПРИБУЛИ до пункту постійно дислокації військової частини A0000:",
            ["солдат ПЕРШИЙ Перший Перший"],
        ),
        freeform_paragraphs=["сержант ТРЕТІЙ Третій Третій госпіталізований."],
    )
    events, unresolved = reader.extract_status_events(path)

    assert unresolved == []
    assert events == [
        {
            "pib_raw": "ПЕРШИЙ Перший Перший", "status": "РВЗ", "effective_date": date(2026, 8, 14),
            "heading": "ПРИБУЛИ до пункту постійно дислокації військової частини A0000:",
            "report_date": date(2026, 8, 14),
        },
        {
            "pib_raw": "ТРЕТІЙ Третій Третій", "status": "ШП", "effective_date": date(2026, 8, 14),
            "heading": "поза межами", "report_date": date(2026, 8, 14),
        },
    ]


def test_departure_to_leave_resolves_to_vp_not_default():
    """"У відпустку" (відбуття) не повинно зловитись правилом "з відпустки"
    (повернення) - протилежні за змістом статуси."""
    assert reader._status_for_heading("У відпустку:") == "ВП"


def test_return_from_szch_takes_priority_over_bare_szch_keyword():
    """"з сзч" (повернення) має перевірятися РАНІШЕ за самостійне "сзч"
    (відбуття) - інакше коротший підрядок "сзч" зловив би обидва варіанти."""
    assert reader._status_for_heading("З СЗЧ:") == "РВЗ"
    assert reader._status_for_heading("СЗЧ:") == "СЗЧ"
    assert reader._status_for_heading("У СЗЧ:") == "СЗЧ"


def test_return_from_ppd_abbreviated_heading_resolves_to_default_status():
    """Реальний випадок: "З ППД:" (скорочення) - той самий сенс "повернення",
    що й "з пункту постійної дислокації" (повна назва), лише коротшим
    написанням - раніше не розпізнавався, лишався на ручну перевірку."""
    assert reader._status_for_heading("З ППД:") == "РВЗ"


def test_return_from_special_task_group_resolves_to_default_status():
    """Реальний випадок: "З [назва спецпризначеної групи]:" - той самий
    сенс "повернення", що й "з відрядженн.../з сзч" тощо - назва групи НЕ
    хардкодиться в тесті (resources/data.json, чутлива інформація)."""
    assert reader._status_for_heading(f"З {_SPECIAL_GROUP_NAME}:") == "РВЗ"


def test_departure_to_special_task_group_resolves_to_vd():
    """"У"/"До [група]:" (відбуття НА завдання групи) - той самий сенс, що
    й "у відрядженн" -> "ВД"."""
    assert reader._status_for_heading(f"У {_SPECIAL_GROUP_NAME}:") == "ВД"
    assert reader._status_for_heading(f"До {_SPECIAL_GROUP_NAME}:") == "ВД"


def test_wounded_evacuation_heading_resolves_to_shp():
    """Реальний заголовок пункту рапорту (раніше "Невідомий заголовок пункту
    рапорту" - на ручну перевірку) - той самий сенс, що й "госпіталізаці"
    (лікування), лише інше формулювання."""
    heading = "Отримали поранення (травму) та перебувають на етапі медичної евакуації до стабілізаційного пункту:"
    assert reader._status_for_heading(heading) == "ШП"


def test_irrecoverable_losses_heading_resolves_to_zagybel_end_to_end(tmp_path):
    """Реальний заголовок пункту рапорту, підтверджений користувачем -
    "Безповоротні втрати:" - ЗАВЖДИ "Загиб." (200), НЕЗАЛЕЖНО від тексту
    причини в останній (без назви) колонці рядка ("Смерть в медичному
    закладі" тут - НЕ те саме, що окремий статус "Смерть", який означає
    конкретно самогубство - лише ручний запис в ОБЛІК.xlsx може встановити
    ЦЕЙ, вужчий статус). Дата - з колонки "Дата смерті" (та сама
    "дата"-substring логіка _detect_pib_and_date_columns, що й "Дата
    вибуття"/"Дата прибуття")."""
    header = ["№", "Звання", "Прізвище ім'я \nпо батькові", "Посада", "Дата смерті", ""]
    path = write_daily_report(
        tmp_path / "03.09.2026 - щоденний рапорт.docx",
        groups=[(
            "ВИБУЛИ зі складу сил та засобів 9 армійського корпусу:",
            [("Безповоротні втрати:", header, [
                ["1", "молодший сержант", "ВІСІМНАДЦЯТИЙ Вісімнадцятий Вісімнадцятий", "Посада", "02.09.2026", "Смерть в медичному закладі"],
            ])],
        )],
    )
    events, unresolved = reader.extract_status_events(path)

    assert unresolved == []
    assert events == [{
        "pib_raw": "ВІСІМНАДЦЯТИЙ Вісімнадцятий Вісімнадцятий", "status": "Загиб.", "effective_date": date(2026, 9, 2),
        "heading": "Безповоротні втрати:", "report_date": date(2026, 9, 3),
    }]


@pytest.mark.parametrize("keyword, expected_status", list(STATUS_HEADING_RULES.items()))
def test_status_for_heading_resolves_every_configured_keyword(keyword, expected_status):
    """Кожен рядок STATUS_HEADING_RULES дійсно розпізнається зі своїм
    статусом, окремо від решти - регресійний тест на випадок, якщо хтось
    (напр. користувач, керуючи статусами напряму в constants.py) додасть
    новий рядок, що ненавмисно "затіняє" вже наявний (коротший підрядок
    перевіряється РАНІШЕ за довший, що його містить)."""
    assert reader._status_for_heading(f"{keyword.capitalize()}:") == expected_status


def test_status_heading_rules_is_a_dict():
    """STATUS_HEADING_RULES - dict (не tuple-of-tuples): дописати третій
    елемент до рядка (замість ОКРЕМОГО нового рядка "ключове слово": статус)
    структурно неможливо - типова помилка при ручному редагуванні цього
    списку статусів більше не валить розбір реального рапорту."""
    assert isinstance(STATUS_HEADING_RULES, dict)


def test_effective_date_from_cell_overrides_report_filename_date(tmp_path):
    """Комірка "Дата вибуття" зі значенням "після 17:00" з ПОПЕРЕДНЬОГО дня -
    подія має застосуватись до дати з КОМІРКИ, а не до дати самого рапорту
    (реальний випадок з resources/report/02.08.2026 - щоденний рапорт.docx)."""
    path = write_daily_report(
        tmp_path / "02.08.2026 - щоденний рапорт.docx",
        groups=[(
            "ВИБУЛИ зі складу сил та засобів 9 армійського корпусу:",
            [("Звільнення з військової служби:", DEPARTURE_HEADER, [
                ["1", "сержант", "ДРУГИЙ Другий Другий", "Посада", "01.08.2026 після 17:00", "Вибув"],
            ])],
        )],
    )
    events, unresolved = reader.extract_status_events(path)

    assert unresolved == []
    assert events[0]["status"] == "ЗВІЛЬНЕНИЙ"
    assert events[0]["effective_date"] == date(2026, 8, 1)
    assert events[0]["report_date"] == date(2026, 8, 2)


def test_unrecognized_subheading_is_reported_as_unresolved_not_dropped(tmp_path):
    path = write_daily_report(
        tmp_path / "01.08.2026 - щоденний рапорт.docx",
        groups=[(
            "ВИБУЛИ зі складу сил та засобів 9 армійського корпусу:",
            [("Якийсь новий, ще не описаний пункт:", DEPARTURE_HEADER, [
                ["1", "сержант", "ТРЕТІЙ Третій Третій", "Посада", "01.08.2026", "Вибув"],
            ])],
        )],
    )
    events, unresolved = reader.extract_status_events(path)

    assert events == []
    assert len(unresolved) == 1
    assert unresolved[0]["pib_raw"] == "ТРЕТІЙ Третій Третій"
    assert "Невідомий заголовок" in unresolved[0]["reason"]


def test_is_ignored_table_heading_recognizes_the_real_heading():
    """Реальний заголовок пункту рапорту, підтверджений користувачем -
    цей пункт МАЄ ігноруватись ПОВНІСТЮ (не отримувати статус, не
    потрапляти на ручну перевірку)."""
    heading = "Отримали поранення (травму) та на наступному етапі евакуації в медичний заклад:"
    assert reader._is_ignored_table_heading(heading)


def test_ignored_table_heading_produces_no_events_and_no_unresolved(tmp_path):
    """За прямою вказівкою користувача - пункт рапорту з заголовком із
    IGNORED_TABLE_HEADING_KEYWORDS (constants.py) МАЄ пропускатись
    ПОВНІСТЮ: НЕ отримує статус (на відміну від STATUS_HEADING_RULES -
    там КОЖЕН заголовок дає якийсь статус) і НЕ потрапляє в unresolved як
    "невідомий заголовок" (на відміну від справді незнайомого заголовка,
    тест вище) - рядки під ним просто зникають з обох файлів "Потребує
    ручної перевірки"."""
    path = write_daily_report(
        tmp_path / "01.08.2026 - щоденний рапорт.docx",
        groups=[(
            "ВИБУЛИ зі складу сил та засобів 9 армійського корпусу:",
            [("Отримали поранення (травму) та на наступному етапі евакуації в медичний заклад:", DEPARTURE_HEADER, [
                ["1", "сержант", "ЧЕТВЕРТИЙ Четвертий Четвертий", "Посада", "01.08.2026", "Вибув"],
            ])],
        )],
    )
    events, unresolved = reader.extract_status_events(path)

    assert events == []
    assert unresolved == []


def test_ignored_table_heading_keywords_is_a_tuple():
    """IGNORED_TABLE_HEADING_KEYWORDS - tuple рядків (не одинокий рядок,
    що python case б розбило по буквах при substring-переборі в
    _is_ignored_table_heading)."""
    assert isinstance(IGNORED_TABLE_HEADING_KEYWORDS, tuple)
    assert all(isinstance(keyword, str) for keyword in IGNORED_TABLE_HEADING_KEYWORDS)


def test_ignored_table_heading_also_matches_the_shorter_real_variant_without_nastupnomu(tmp_path):
    """Реальний заголовок пункту рапорту, підтверджений користувачем -
    "На етапі евакуації:" (БЕЗ "наступному" - коротше формулювання ТОГО
    САМОГО пункту, що й тест вище) - МАЄ ігноруватись УСЮДИ так само
    ПОВНІСТЮ (спільний підрядок "етапі евакуації" покриває ОБИДВА
    формулювання)."""
    assert reader._is_ignored_table_heading("На етапі евакуації:")

    path = write_daily_report(
        tmp_path / "06.09.2026 - щоденний рапорт.docx",
        groups=[(
            "ВИБУЛИ зі складу сил та засобів 9 армійського корпусу:",
            [("На етапі евакуації:", DEPARTURE_HEADER, [
                ["1", "старший матрос", "ДВАДЦЯТИЙ Двадцятий Двадцятий", "Посада", "06.09.2026", "Вибув"],
            ])],
        )],
    )
    events, unresolved = reader.extract_status_events(path)

    assert events == []
    assert unresolved == []


def test_rtgr_departure_heading_resolves_to_rtgr_status_end_to_end(tmp_path):
    """Реальний заголовок пункту рапорту, підтверджений користувачем -
    "В РТГР:" - статус "РТГр" на дату вибуття з колонки таблиці."""
    path = write_daily_report(
        tmp_path / "06.09.2026 - щоденний рапорт.docx",
        groups=[(
            "ВИБУЛИ зі складу сил та засобів 9 армійського корпусу:",
            [("В РТГР:", DEPARTURE_HEADER, [
                ["1", "старший сержант", "ШІСТНАДЦЯТИЙ Шістнадцятий Шістнадцятий", "Головний сержант", "06.09.2026", "Вибув з району виконання завдань в РТГР"],
            ])],
        )],
    )
    events, unresolved = reader.extract_status_events(path)

    assert unresolved == []
    assert events == [{
        "pib_raw": "ШІСТНАДЦЯТИЙ Шістнадцятий Шістнадцятий", "status": "РТГр", "effective_date": date(2026, 9, 6),
        "heading": "В РТГР:", "report_date": date(2026, 9, 6),
    }]


def test_medical_company_treatment_heading_resolves_to_lmr_status_end_to_end(tmp_path):
    """Реальний заголовок пункту рапорту, підтверджений користувачем -
    "Лікування в медичній роті військової частини ...:" - статус "ЛМР"
    (STATUS_MEANINGS - "Лікуються в медичній роті") на дату вибуття з
    колонки таблиці."""
    path = write_daily_report(
        tmp_path / "10.09.2026 - щоденний рапорт.docx",
        groups=[(
            "ВИБУЛИ зі складу сил та засобів 9 армійського корпусу:",
            [("Лікування в медичній роті військової частини A0000:", DEPARTURE_HEADER, [
                ["1", "старший матрос", "ТРИДЦЯТИЙ Тридцятий Тридцятий", "Водій-механік", "10.09.2026", "Вибув з району виконання завдань до медичної роти військової частини A0000"],
            ])],
        )],
    )
    events, unresolved = reader.extract_status_events(path)

    assert unresolved == []
    assert events == [{
        "pib_raw": "ТРИДЦЯТИЙ Тридцятий Тридцятий", "status": "ЛМР", "effective_date": date(2026, 9, 10),
        "heading": "Лікування в медичній роті військової частини A0000:", "report_date": date(2026, 9, 10),
    }]


def test_freeform_vpsz_phrase_without_date_falls_back_to_report_date(tmp_path):
    """Реальний випадок (resources/report/01.08.2026 - щоденний рапорт.docx,
    розділ "Поза межами"): "вибув у відпустку за станом здоров'я" - розпізнається
    як подія (статус "ВПСЗ"), а не лишається лише на ручну перевірку - навіть
    без дати БЕЗПОСЕРЕДНЬО перед зворотом (дата "30.07.2026" тут - дата рішення
    ВЛК, ПІСЛЯ звороту, а не дата самої відпустки - тож ігнорується, ефективна
    дата - дата рапорту)."""
    path = write_daily_report(
        tmp_path / "01.08.2026 - щоденний рапорт.docx",
        freeform_paragraphs=[
            "штаб-сержант ЧЕТВЕРТИЙ Четвертий Четвертий вибув у відпустку за станом здоров'я згідно рішенню ВЛК №123-ОД від 30.07.2026.",
        ],
    )
    events, unresolved = reader.extract_status_events(path)

    assert unresolved == []
    assert events == [{
        "pib_raw": "ЧЕТВЕРТИЙ Четвертий Четвертий", "status": "ВПСЗ", "effective_date": date(2026, 8, 1),
        "heading": "поза межами", "report_date": date(2026, 8, 1),
    }]


def test_freeform_vpsz_phrase_with_leading_date_uses_that_date(tmp_path):
    """Реальний випадок (06.08.2026 - щоденний рапорт.docx): дата ОДРАЗУ перед
    зворотом - ефективна дата саме вона, а не дата рапорту."""
    path = write_daily_report(
        tmp_path / "06.08.2026 - щоденний рапорт.docx",
        freeform_paragraphs=["сержант ШОСТИЙ Шостий Шостий 06.08.2026 року вибув у відпустку за станом здоров'я на 30 календарних днів."],
    )
    events, _unresolved = reader.extract_status_events(path)

    assert events == [{
        "pib_raw": "ШОСТИЙ Шостий Шостий", "status": "ВПСЗ", "effective_date": date(2026, 8, 6),
        "heading": "поза межами", "report_date": date(2026, 8, 6),
    }]


def test_freeform_vpsz_phrase_picks_date_right_before_phrase_not_an_earlier_unrelated_date(tmp_path):
    """Реальний випадок (08.08.2026 - щоденний рапорт.docx): речення з ДВОМА
    датами - ефективна дата саме та, що БЕЗПОСЕРЕДНЬО перед "вибув...", а не
    перша (пов'язана з іншою подією - випискою) дата речення."""
    path = write_daily_report(
        tmp_path / "08.08.2026 - щоденний рапорт.docx",
        freeform_paragraphs=[
            "матрос СЬОМИЙ Сьомий Сьомий 07.08.2026 виписаний з лікувального закладу та "
            "08.08.2026 року вибув у відпустку за станом здоров'я на 30 календарних днів.",
        ],
    )
    events, _unresolved = reader.extract_status_events(path)

    assert events == [{
        "pib_raw": "СЬОМИЙ Сьомий Сьомий", "status": "ВПСЗ", "effective_date": date(2026, 8, 8),
        "heading": "поза межами", "report_date": date(2026, 8, 8),
    }]


def test_freeform_szch_phrase_recognized(tmp_path):
    """Реальний випадок (resources/report/07.08.2026 - щоденний рапорт.docx):
    "вибув у СЗЧ" - розпізнається як подія (статус "СЗЧ"), як і "вибув у
    відпустку за станом здоров'я" (ВПСЗ). Дата в цьому реченні ("06.07.2026")
    - імовірна друкарська помилка (місяць замість "08") в самому рапорті,
    поза діапазоном колонок ОБЛІК.xlsx - але це вже турбота apply_events
    (див. test_generate_timesheet.py), тут лише перевіряється, що подія
    взагалі ВИТЯГУЄТЬСЯ з реченням саме з цією (написаною) датою."""
    path = write_daily_report(
        tmp_path / "07.08.2026 - щоденний рапорт.docx",
        freeform_paragraphs=["матрос ВОСЬМИЙ Восьмий Восьмий під час проходження ВЛК 06.07.2026 вибув у СЗЧ."],
    )
    events, unresolved = reader.extract_status_events(path)

    assert unresolved == []
    assert events == [{
        "pib_raw": "ВОСЬМИЙ Восьмий Восьмий", "status": "СЗЧ", "effective_date": date(2026, 7, 6),
        "heading": "поза межами", "report_date": date(2026, 8, 7),
    }]


def test_freeform_szch_phrase_with_plausible_date_uses_that_date(tmp_path):
    path = write_daily_report(
        tmp_path / "07.08.2026 - щоденний рапорт.docx",
        freeform_paragraphs=["матрос ДЕСЯТИЙ Десятий Десятий 07.08.2026 вибув у СЗЧ."],
    )
    events, _unresolved = reader.extract_status_events(path)

    assert events == [{
        "pib_raw": "ДЕСЯТИЙ Десятий Десятий", "status": "СЗЧ", "effective_date": date(2026, 8, 7),
        "heading": "поза межами", "report_date": date(2026, 8, 7),
    }]


def test_freeform_hospitalization_transfer_resolves_to_shp_even_though_status_does_not_change(tmp_path):
    """Реальний випадок (04.08.2026 - щоденний рапорт.docx): виписаний з
    ОДНОГО закладу й ТОГО Ж дня госпіталізований до ІНШОГО - статус
    "перезаходить" у ШП (був ШП, лишається ШП), тож раніше не було жодної
    події для застосування, і речення даремно потрапляло на ручну перевірку
    (за словами користувача - "не треба вказувати в доповіді", коли статус не
    змінюється). Тепер розпізнається як подія (статус "ШП" незалежно від
    БАЗОВОГО статусу - навіть якщо він вже "ШП", застосування безпечне)."""
    path = write_daily_report(
        tmp_path / "04.08.2026 - щоденний рапорт.docx",
        freeform_paragraphs=[
            "1. молодший сержант ОДИНАДЦЯТИЙ Одинадцятий Одинадцятий 04.08.2026 виписаний із військово-медичного "
            "центру Південного регіону м. Одеса та 04.08.2026 госпіталізований до військової частини A0001.",
        ],
    )
    events, unresolved = reader.extract_status_events(path)

    assert unresolved == []
    assert events == [{
        "pib_raw": "ОДИНАДЦЯТИЙ Одинадцятий Одинадцятий", "status": "ШП", "effective_date": date(2026, 8, 4),
        "heading": "поза межами", "report_date": date(2026, 8, 4),
    }]


def test_freeform_department_transfer_resolves_to_shp_even_though_status_does_not_change(tmp_path):
    """Реальний випадок: "переведений з [відділення] ... до [відділення]"
    (переведення МІЖ відділеннями в межах лікування, а не до іншого
    підрозділу) - статус "перезаходить" у ШП (був ШП, лишається ШП), тож
    речення даремно потрапляло на ручну перевірку. Тепер розпізнається як
    подія (статус "ШП" незалежно від БАЗОВОГО статусу - той самий "безпечний
    no-op" принцип, що й госпіталізація до іншого закладу)."""
    path = write_daily_report(
        tmp_path / "12.08.2026 - щоденний рапорт.docx",
        freeform_paragraphs=[
            "майор ЧОТИРНАДЦЯТИЙ Чотирнадцятий Чотирнадцятий 12.08.2026 переведений з неврологічного відділення "
            "№2 медичного закладу B0001 до нейрореабілітаційного відділення медичного закладу B0001.",
        ],
    )
    events, unresolved = reader.extract_status_events(path)

    assert unresolved == []
    assert events == [{
        "pib_raw": "ЧОТИРНАДЦЯТИЙ Чотирнадцятий Чотирнадцятий", "status": "ШП", "effective_date": date(2026, 8, 12),
        "heading": "поза межами", "report_date": date(2026, 8, 12),
    }]


def test_freeform_transfer_to_another_department_resolves_to_shp(tmp_path):
    """Реальний випадок (баг, підтверджений користувачем): "переведений в
    інше відділення [заклад]" - КОРОТШЕ формулювання ТОГО САМОГО факту, що й
    "переведений з [відділення] ... до [відділення]" вище - називає ЛИШЕ
    відділення ПРИЗНАЧЕННЯ (без "з", без відділення-джерела) - раніше не
    розпізнавалось (жоден наявний зворот цього не покривав), даремно
    потрапляло на ручну перевірку. Статус лишається "ШП" (той самий
    "безпечний no-op" принцип)."""
    path = write_daily_report(
        tmp_path / "26.08.2026 - щоденний рапорт.docx",
        freeform_paragraphs=[
            "26.08.2026 матрос ДВАДЦЯТЬШОСТИЙ Двадцятьшостий Двадцятьшостий переведений в інше відділення "
            "«Приклад» медичного закладу B0001.",
        ],
    )
    events, unresolved = reader.extract_status_events(path)

    assert unresolved == []
    assert events == [{
        "pib_raw": "ДВАДЦЯТЬШОСТИЙ Двадцятьшостий Двадцятьшостий", "status": "ШП", "effective_date": date(2026, 8, 26),
        "heading": "поза межами", "report_date": date(2026, 8, 26),
    }]


def test_freeform_hospitalization_without_a_leading_date_falls_back_to_report_date(tmp_path):
    """Реальний випадок (06.08.2026 - щоденний рапорт.docx): госпіталізація
    БЕЗ дати в самому реченні (людина була у відпустці за станом здоров'я,
    госпіталізована без вказаної дати) - ефективна дата - дата рапорту."""
    path = write_daily_report(
        tmp_path / "06.08.2026 - щоденний рапорт.docx",
        freeform_paragraphs=[
            "старший матрос ДВАНАДЦЯТИЙ Дванадцятий Дванадцятий під час відпустки за станом здоров'я "
            "госпіталізований до КНП «Центральна міська лікарня» м. Гайворон.",
        ],
    )
    events, _unresolved = reader.extract_status_events(path)

    assert events == [{
        "pib_raw": "ДВАНАДЦЯТИЙ Дванадцятий Дванадцятий", "status": "ШП", "effective_date": date(2026, 8, 6),
        "heading": "поза межами", "report_date": date(2026, 8, 6),
    }]


def test_freeform_evacuation_transfer_resolves_to_shp(tmp_path):
    """Реальний випадок (08.08.2026 - щоденний рапорт.docx): "переводиться на
    наступний етап медичної евакуації" - інше формулювання того самого "ще в
    медичній системі" сценарію, що й "госпіталізован" - теж "ШП", незалежно
    від БАЗОВОГО статусу."""
    path = write_daily_report(
        tmp_path / "08.08.2026 - щоденний рапорт.docx",
        freeform_paragraphs=[
            "матрос ТРИНАДЦЯТИЙ Тринадцятий Тринадцятий 08.08.2026 року виписаний із лікувального закладу "
            "КНП «Приклад» та переводиться на наступний етап медичної евакуації.",
        ],
    )
    events, unresolved = reader.extract_status_events(path)

    assert unresolved == []
    assert events == [{
        "pib_raw": "ТРИНАДЦЯТИЙ Тринадцятий Тринадцятий", "status": "ШП", "effective_date": date(2026, 8, 8),
        "heading": "поза межами", "report_date": date(2026, 8, 8),
    }]


def test_freeform_treatment_stage_transfer_resolves_to_shp(tmp_path):
    """Реальний випадок: "виписаний ... та на наступному етапі лікування до
    [заклад]" - виписка з ОДНОГО закладу, але ЗРАЗУ на наступний етап
    ЛІКУВАННЯ (не медичної евакуації, як у сусідньому тесті вище) в ІНШОМУ -
    той самий "ще лікується" сценарій, теж "ШП", незалежно від БАЗОВОГО
    статусу. ОБИДВА ключових слова ("виписан" І "наступний етап лікування")
    мають бути присутні - навмисно вужче, ніж просто "наступний етап
    лікування" саме по собі."""
    path = write_daily_report(
        tmp_path / "17.08.2026 - щоденний рапорт.docx",
        freeform_paragraphs=[
            "старший матрос ШІСТНАДЦЯТИЙ Шістнадцятий Шістнадцятий виписаний 17.08.2026 із КНП «Приклад» "
            "м. Приклад та на наступному етапі лікування до КНП «Обласний приклад».",
        ],
    )
    events, unresolved = reader.extract_status_events(path)

    assert unresolved == []
    assert events == [{
        "pib_raw": "ШІСТНАДЦЯТИЙ Шістнадцятий Шістнадцятий", "status": "ШП", "effective_date": date(2026, 8, 17),
        "heading": "поза межами", "report_date": date(2026, 8, 17),
    }]


def test_freeform_treatment_stage_mention_without_discharge_stays_unresolved(tmp_path):
    """"наступний етап лікування" САМ ПО СОБІ (без "виписаний" перед ним) -
    НЕ розпізнається, лишається на ручну перевірку - навмисно вужчий
    критерій (ОБИДВА ключових слова), щоб не хапати випадкові згадки поза
    контекстом виписки з ОДНОГО закладу в ІНШИЙ."""
    path = write_daily_report(
        tmp_path / "17.08.2026 - щоденний рапорт.docx",
        freeform_paragraphs=[
            "старший матрос СІМНАДЦЯТИЙ Сімнадцятий Сімнадцятий переходить на наступний етап лікування "
            "до КНП «Приклад».",
        ],
    )
    events, unresolved = reader.extract_status_events(path)

    assert events == []
    assert len(unresolved) == 1


def test_freeform_discharge_with_dated_vlk_certificate_produces_vlk_then_leave(tmp_path):
    """Реальний випадок (05.08.2026 - щоденний рапорт.docx): "виписаний ...
    згідно довідки ВЛК №... від DATE потребує відпустки за станом здоров'я"
    - ДВІ події з ОДНОГО речення: сам день довідки ВЛК - статус "ВЛК",
    наступний день - "ВПСЗ" (речення прямо згадує "за станом здоров'я",
    тож відпустка починається саме як ВПСЗ, а не просто ВП)."""
    path = write_daily_report(
        tmp_path / "05.08.2026 - щоденний рапорт.docx",
        freeform_paragraphs=[
            "старший матрос ЧОТИРНАДЦЯТИЙ Чотирнадцятий Чотирнадцятий 04.08.2026 року виписаний з лікувального закладу "
            "згідно довідки ВЛК №2026-0804-1047-3411-0 від 04.08.2026 потребує відпустки за станом "
            "здоров'я на 30 календарних днів.",
        ],
    )
    events, unresolved = reader.extract_status_events(path)

    assert unresolved == []
    assert events == [
        {
            "pib_raw": "ЧОТИРНАДЦЯТИЙ Чотирнадцятий Чотирнадцятий", "status": "ВЛК", "effective_date": date(2026, 8, 4),
            "heading": "поза межами", "report_date": date(2026, 8, 5),
        },
        {
            "pib_raw": "ЧОТИРНАДЦЯТИЙ Чотирнадцятий Чотирнадцятий", "status": "ВПСЗ", "effective_date": date(2026, 8, 5),
            "heading": "поза межами", "report_date": date(2026, 8, 5),
        },
    ]


def test_freeform_needs_leave_with_spelled_out_vlk_commission_decision_produces_vlk_then_leave(tmp_path):
    """Реальний випадок, підтверджений користувачем: "... вибув для
    проходження військово-лікарської комісії ... Відповідно до рішення
    військово-лікарської комісії від ДАТА потребує відпустки ... за станом
    здоров'я по пораненню." - ПОВНА назва "військово-лікарської комісії" (а
    не абревіатура "ВЛК") із "від ДАТА" - той САМИЙ сенс, що й "ВЛК ... від
    ДАТА" (_VLK_DATE_RE, date_after3): ДВІ події - сам день рішення ВЛК -
    статус "ВЛК", наступний день - "ВПСЗ"."""
    path = write_daily_report(
        tmp_path / "31.08.2026 - щоденний рапорт.docx",
        freeform_paragraphs=[
            "31.08.2026 штаб-сержант ДВАНАДЦЯТИЙ Дванадцятий Дванадцятий закінчив відпустку за станом "
            "здоров'я та вибув для проходження військово-лікарської комісії згідно з направленням. "
            "Відповідно до рішення військово-лікарської комісії від 31.08.2026 потребує відпустки на 30 "
            "календарних діб за станом здоров'я по пораненню.",
        ],
    )
    events, unresolved = reader.extract_status_events(path)

    assert unresolved == []
    assert events == [
        {
            "pib_raw": "ДВАНАДЦЯТИЙ Дванадцятий Дванадцятий", "status": "ВЛК", "effective_date": date(2026, 8, 31),
            "heading": "поза межами", "report_date": date(2026, 8, 31),
        },
        {
            "pib_raw": "ДВАНАДЦЯТИЙ Дванадцятий Дванадцятий", "status": "ВПСЗ", "effective_date": date(2026, 9, 1),
            "heading": "поза межами", "report_date": date(2026, 8, 31),
        },
    ]


def test_freeform_needs_leave_for_wound_treatment_resolves_to_vpbp_not_vp(tmp_path):
    """Реальний випадок (баг, підтверджений користувачем): "закінчив
    проходження військово-лікарської комісії ..., відповідно до висновку
    військово-лікарської комісії від ДАТА потребує відпустки ДЛЯ
    ЛІКУВАННЯ ПІСЛЯ ПОРАНЕННЯ на 30 календарних днів." - речення НЕ згадує
    "за станом здоров'я" (тест вище, дає "ВПСЗ"), тож наступний день хибно
    отримував загальний статус "ВП", хоча причина відпустки прямо названа -
    "ВПБП" ("Відпустка для лікування після тяжкого поранення",
    STATUS_MEANINGS) - точніший, вже підтримуваний статус."""
    path = write_daily_report(
        tmp_path / "11.09.2026 - щоденний рапорт.docx",
        freeform_paragraphs=[
            "старший лейтенант СОТИЙ Сотий Сотий 11.09.2026 закінчив проходження військово-лікарської "
            "комісії у військовій частині A0000, відповідно до висновку військово-лікарської комісії від "
            "11.09.2026 року потребує відпустки для лікування після поранення на 30 календарних днів.",
        ],
    )
    events, unresolved = reader.extract_status_events(path)

    assert unresolved == []
    assert events == [
        {
            "pib_raw": "СОТИЙ Сотий Сотий", "status": "ВЛК", "effective_date": date(2026, 9, 11),
            "heading": "поза межами", "report_date": date(2026, 9, 11),
        },
        {
            "pib_raw": "СОТИЙ Сотий Сотий", "status": "ВПБП", "effective_date": date(2026, 9, 12),
            "heading": "поза межами", "report_date": date(2026, 9, 11),
        },
    ]


def test_freeform_needs_leave_mentioning_both_health_and_wound_treatment_resolves_to_vpsz(tmp_path):
    """Реальний випадок (баг, підтверджений користувачем): речення МОЖЕ
    згадувати ОБИДВІ фрази одразу - "... під час відпустки ЗА СТАНОМ
    ЗДОРОВ'Я по пораненню, ... потребує відпустки ДЛЯ ЛІКУВАННЯ ПІСЛЯ
    ПОРАНЕННЯ на 30 днів." - "за станом здоров'я" МАЄ перемагати (дає
    "ВПСЗ", а не "ВПБП" - тест вище): це той самий статус, що його ПІЗНІШЕ
    підтверджує окреме речення наступного рапорту ("вибув у відпустку за
    станом здоров'я") - без цього пріоритету "ВПБП" з ЦЬОГО рапорту й
    "ВПСЗ" з наступного хибно позначались би як суперечність, хоча це той
    самий факт відпустки."""
    path = write_daily_report(
        tmp_path / "08.09.2026 - щоденний рапорт.docx",
        freeform_paragraphs=[
            "штаб-сержант СОТИЙ Сотий Сотий 08.09.2026 закінчив проходження "
            "військово-лікарської комісії у військовій частині A0000 під час відпустки за станом "
            "здоров'я по пораненню, відповідно до висновку військово-лікарської комісії від 08.09.2026 "
            "року потребує відпустки для лікування після поранення на 30 календарних днів.",
        ],
    )
    events, unresolved = reader.extract_status_events(path)

    assert unresolved == []
    assert events == [
        {
            "pib_raw": "СОТИЙ Сотий Сотий", "status": "ВЛК", "effective_date": date(2026, 9, 8),
            "heading": "поза межами", "report_date": date(2026, 9, 8),
        },
        {
            "pib_raw": "СОТИЙ Сотий Сотий", "status": "ВПСЗ", "effective_date": date(2026, 9, 9),
            "heading": "поза межами", "report_date": date(2026, 9, 8),
        },
    ]


def test_freeform_evacuation_transfer_without_the_word_medical_resolves_to_shp(tmp_path):
    """Реальний випадок: "направлений на наступний етап евакуації" - БЕЗ
    слова "медичної" (на відміну від тесту вище) - той самий сенс, коротшим
    формулюванням - теж "ШП", незалежно від БАЗОВОГО статусу."""
    path = write_daily_report(
        tmp_path / "22.08.2026 - щоденний рапорт.docx",
        freeform_paragraphs=[
            "22.08.2026 молодший сержант СЬОМИЙ Сьомий Сьомий виписаний із КНП «Приклад» ДОР та "
            "направлений на наступний етап евакуації до лікувального закладу.",
        ],
    )
    events, unresolved = reader.extract_status_events(path)

    assert unresolved == []
    assert events == [{
        "pib_raw": "СЬОМИЙ Сьомий Сьомий", "status": "ШП", "effective_date": date(2026, 8, 22),
        "heading": "поза межами", "report_date": date(2026, 8, 22),
    }]


def test_freeform_evacuation_transfer_accepts_case_variant_and_typo_too(tmp_path):
    """Реальний випадок, підтверджений користувачем: "...та перебуває на
    НАСТУПНОМУ ЕТАП евакуації." - "наступному" (не "наступний" - інший
    відмінок) і "етап" (не "етапі" - ймовірно, описка) - той самий сенс,
    той самий "ШП"."""
    path = write_daily_report(
        tmp_path / "02.09.2026 - щоденний рапорт.docx",
        freeform_paragraphs=[
            "молодший сержант ДЕСЯТИЙ Десятий Десятий 02.09.2026 виписаний із КНП «Приклад» та "
            "перебуває на наступному етап евакуації.",
        ],
    )
    events, unresolved = reader.extract_status_events(path)

    assert unresolved == []
    assert events == [{
        "pib_raw": "ДЕСЯТИЙ Десятий Десятий", "status": "ШП", "effective_date": date(2026, 9, 2),
        "heading": "поза межами", "report_date": date(2026, 9, 2),
    }]


def test_freeform_hospitalization_transfer_accepts_the_missing_letter_typo(tmp_path):
    """Реальний випадок, підтверджений користувачем: "...та ДАТА
    ГОСПІТАЛІЗОАНИЙ військовою частиною ..." - друкарська помилка, пропущена
    буква "в" (замість "госпіталізоВАНИЙ") - той самий сенс, той самий
    "ШП"."""
    path = write_daily_report(
        tmp_path / "03.09.2026 - щоденний рапорт.docx",
        freeform_paragraphs=[
            "молодший сержант СІМНАДЦЯТИЙ Сімнадцятий Сімнадцятий 02.09.2026 виписаний з "
            "медичного закладу B0001 та 03.09.2026 госпіталізоаний військовою частиною A0001.",
        ],
    )
    events, unresolved = reader.extract_status_events(path)

    assert unresolved == []
    assert events == [{
        "pib_raw": "СІМНАДЦЯТИЙ Сімнадцятий Сімнадцятий", "status": "ШП", "effective_date": date(2026, 9, 3),
        "heading": "поза межами", "report_date": date(2026, 9, 3),
    }]


def test_freeform_evacuation_transfer_accepts_hospitalization_wording_too(tmp_path):
    """Реальний випадок, підтверджений користувачем: "...та перебуває на
    наступному етап ГОСПІТАЛІЗАЦІЇ." - інше слово, ніж "евакуації" (тест
    вище), той самий "ще лікується, ще не на волі" сенс, той самий "ШП"."""
    path = write_daily_report(
        tmp_path / "04.09.2026 - щоденний рапорт.docx",
        freeform_paragraphs=[
            "матрос ШІСТНАДЦЯТИЙ Шістнадцятий Шістнадцятий 04.09.2026 виписаний із медичного "
            "закладу B0001 та перебуває на наступному етап госпіталізації.",
        ],
    )
    events, unresolved = reader.extract_status_events(path)

    assert unresolved == []
    assert events == [{
        "pib_raw": "ШІСТНАДЦЯТИЙ Шістнадцятий Шістнадцятий", "status": "ШП", "effective_date": date(2026, 9, 4),
        "heading": "поза межами", "report_date": date(2026, 9, 4),
    }]


def test_freeform_completed_vlk_produces_a_same_day_health_leave_event(tmp_path):
    """Реальний випадок (баг, підтверджений користувачем): "закінчив
    проходження військово-лікарської комісії ДАТА, потребує відпустки за
    станом здоров'я ..." - ЩЕ ОДНЕ формулювання дати БІЛЯ згадки ВЛК (на
    відміну від "довідки ВЛК №... від ДАТА"/"ДАТА за висновком ВЛК" вище) -
    людина ЗАВЕРШИЛА саму комісію цього дня (а не лише посилається на
    ДОКУМЕНТ, датований цим днем) - тож, на відміну від тих двох форм, ТУТ
    - ОДНА подія: відпустка починається ТОГО САМОГО дня (ВПСЗ - речення
    прямо згадує "за станом здоров'я"), без окремого "дня ВЛК" (комісію вже
    пройдено)."""
    path = write_daily_report(
        tmp_path / "21.08.2026 - щоденний рапорт.docx",
        freeform_paragraphs=[
            "молодший сержант ШОСТИЙ Шостий Шостий закінчив проходження військово-лікарської комісії "
            "21.08.2026, потребує відпустки за станом здоров'я на 30 календарних діб у зв'язку з пораненням.",
        ],
    )
    events, unresolved = reader.extract_status_events(path)

    assert unresolved == []
    assert events == [{
        "pib_raw": "ШОСТИЙ Шостий Шостий", "status": "ВПСЗ", "effective_date": date(2026, 8, 21),
        "heading": "поза межами", "report_date": date(2026, 8, 21),
    }]


def test_freeform_discharge_without_a_dated_vlk_reference_produces_shp_then_leave(tmp_path):
    """Реальний випадок (07.08.2026 - щоденний рапорт.docx): "виписаний ...
    DATE та згідно висновку ВЛК потребує відпустки для лікування" - БЕЗ дати
    БІЛЯ самої згадки ВЛК (лише "згідно висновку", без "від DATE") - день
    виписки лишається "ШП" (ще не покинув медичну систему), а не "ВЛК"."""
    path = write_daily_report(
        tmp_path / "07.08.2026 - щоденний рапорт.docx",
        freeform_paragraphs=[
            "молодший сержант ШІСТНАДЦЯТИЙ Шістнадцятий Шістнадцятий виписаний з лікувального закладу "
            "07.08.2026 та згідно висновку ВЛК потребує відпустки для лікування на 30 календарних діб",
        ],
    )
    events, unresolved = reader.extract_status_events(path)

    assert unresolved == []
    assert events == [
        {
            "pib_raw": "ШІСТНАДЦЯТИЙ Шістнадцятий Шістнадцятий", "status": "ШП", "effective_date": date(2026, 8, 7),
            "heading": "поза межами", "report_date": date(2026, 8, 7),
        },
        {
            "pib_raw": "ШІСТНАДЦЯТИЙ Шістнадцятий Шістнадцятий", "status": "ВП", "effective_date": date(2026, 8, 8),
            "heading": "поза межами", "report_date": date(2026, 8, 7),
        },
    ]


def test_freeform_date_before_vlk_conclusion_produces_vlk_then_leave(tmp_path):
    """Реальний випадок (07.08.2026 - щоденний рапорт.docx): "під час
    відпустки за станом здоров'я DATE за висновком ВЛК потребує відпустки для
    лікування" - дата ПЕРЕД згадкою ВЛК (а не після, як "довідки ВЛК №... від
    DATE") - так само розпізнається: сам день ВЛК лишається "ВЛК" (виняток),
    а відпустка наступного дня - "ВПСЗ" (речення прямо згадує "за станом
    здоров'я", тож це не просто ВП). Подвійний пробіл між "потребує" й
    "відпустки" - як у реальному тексті - теж має спрацювати."""
    path = write_daily_report(
        tmp_path / "07.08.2026 - щоденний рапорт.docx",
        freeform_paragraphs=[
            "матрос СІМНАДЦЯТИЙ Сімнадцятий Сімнадцятий під час відпустки за станом здоров'я 07.08.2026 "
            "за висновком ВЛК потребує  відпустки для лікування на 30 календарних діб",
        ],
    )
    events, unresolved = reader.extract_status_events(path)

    assert unresolved == []
    assert events == [
        {
            "pib_raw": "СІМНАДЦЯТИЙ Сімнадцятий Сімнадцятий", "status": "ВЛК", "effective_date": date(2026, 8, 7),
            "heading": "поза межами", "report_date": date(2026, 8, 7),
        },
        {
            "pib_raw": "СІМНАДЦЯТИЙ Сімнадцятий Сімнадцятий", "status": "ВПСЗ", "effective_date": date(2026, 8, 8),
            "heading": "поза межами", "report_date": date(2026, 8, 7),
        },
    ]


def test_freeform_discharge_mention_without_needs_leave_phrase_is_not_auto_applied(tmp_path):
    """Без "потребує відпустки" - НІЧОГО не застосовується, навіть якщо
    речення згадує і виписку, і ВЛК: самих по собі цих згадок недостатньо -
    надто неоднозначно, куди саме поділась людина після виписки."""
    path = write_daily_report(
        tmp_path / "07.08.2026 - щоденний рапорт.docx",
        freeform_paragraphs=[
            "матрос ВІСІМНАДЦЯТИЙ Вісімнадцятий Вісімнадцятий 07.08.2026 виписаний з лікувального закладу "
            "згідно довідки ВЛК №123 від 07.08.2026 повернувся до району виконання завдань.",
        ],
    )
    events, unresolved = reader.extract_status_events(path)

    assert events == []
    assert len(unresolved) == 1
    assert unresolved[0]["pib_raw"] == "ВІСІМНАДЦЯТИЙ Вісімнадцятий Вісімнадцятий"


def test_freeform_needs_leave_without_discharge_or_vlk_date_is_not_auto_applied(tmp_path):
    """"потребує відпустки" є, але немає ЖОДНОЇ з дат (ні біля ВЛК, ні біля
    "виписаний") - немає якоря для дня самої події, тож нічого не
    застосовується, лишається на ручну перевірку."""
    path = write_daily_report(
        tmp_path / "07.08.2026 - щоденний рапорт.docx",
        freeform_paragraphs=["матрос ВІСІМНАДЦЯТИЙ Вісімнадцятий Вісімнадцятий потребує відпустки за станом здоров'я."],
    )
    events, unresolved = reader.extract_status_events(path)

    assert events == []
    assert len(unresolved) == 1
    assert unresolved[0]["pib_raw"] == "ВІСІМНАДЦЯТИЙ Вісімнадцятий Вісімнадцятий"


def test_freeform_discharge_with_vlk_in_progress_and_date_before_name_produces_vlk_then_leave(tmp_path):
    """Реальний випадок, підтверджений користувачем: "ДАТА звання ПІБ
    виписаний з [заклад] ... проходить ВЛК, потребує відпустку за станом
    здоров'я ... після поранення." - "проходить ВЛК" (теперішній час, а не
    "закінчив"/"довідки ВЛК №... від ДАТА") НЕ підходить під _VLK_DATE_RE, і
    дата стоїть ПЕРЕД ПІБ (на початку абзацу), а не біля самого "виписан" -
    раніше не мав якоря для дня події взагалі (лишалось на ручну перевірку,
    БАЗОВИЙ статус мовчки не змінювався). Явна згадка "проходить ВЛК" - день
    виписки "ВЛК" (а не "ШП"), відпустка наступного дня - "ВПСЗ" (речення
    прямо згадує "за станом здоров'я")."""
    path = write_daily_report(
        tmp_path / "28.08.2026 - щоденний рапорт.docx",
        freeform_paragraphs=[
            "28.08.2026 матрос ЧЕТВЕРТИЙ Четвертий Четвертий виписаний з КНП «Приклад» "
            "проходить ВЛК, потребує відпустку за станом здоров'я на 30 календарних діб після поранення.",
        ],
    )
    events, unresolved = reader.extract_status_events(path)

    assert unresolved == []
    assert events == [
        {
            "pib_raw": "ЧЕТВЕРТИЙ Четвертий Четвертий", "status": "ВЛК", "effective_date": date(2026, 8, 28),
            "heading": "поза межами", "report_date": date(2026, 8, 28),
        },
        {
            "pib_raw": "ЧЕТВЕРТИЙ Четвертий Четвертий", "status": "ВПСЗ", "effective_date": date(2026, 8, 29),
            "heading": "поза межами", "report_date": date(2026, 8, 28),
        },
    ]


def test_freeform_needs_leave_accepts_genitive_case_too(tmp_path):
    """"потребує відпустки" (родовий відмінок) - той самий сценарій, що й
    "потребує відпустку" (знахідний) вище - реальні рапорти трапляються з
    ОБОМА відмінками, _NEEDS_LEAVE_RE має розпізнавати обидва."""
    path = write_daily_report(
        tmp_path / "28.08.2026 - щоденний рапорт.docx",
        freeform_paragraphs=[
            "28.08.2026 матрос ВОСЬМИЙ Восьмий Восьмий виписаний з КНП «Приклад» проходить ВЛК, "
            "потребує відпустки за станом здоров'я на 30 календарних діб після поранення.",
        ],
    )
    events, unresolved = reader.extract_status_events(path)

    assert unresolved == []
    assert events == [
        {
            "pib_raw": "ВОСЬМИЙ Восьмий Восьмий", "status": "ВЛК", "effective_date": date(2026, 8, 28),
            "heading": "поза межами", "report_date": date(2026, 8, 28),
        },
        {
            "pib_raw": "ВОСЬМИЙ Восьмий Восьмий", "status": "ВПСЗ", "effective_date": date(2026, 8, 29),
            "heading": "поза межами", "report_date": date(2026, 8, 28),
        },
    ]


def test_freeform_vlk_in_progress_without_any_date_before_name_is_not_auto_applied(tmp_path):
    """Згадка "проходить ВЛК" Є, але дати ПЕРЕД ПІБ немає взагалі (абзац
    починається одразу зі звання й ПІБ, без дати), і "виписан" - задалеко
    від дати для _DISCHARGE_DATE_RE - немає якоря для дня події, лишається
    на ручну перевірку."""
    path = write_daily_report(
        tmp_path / "28.08.2026 - щоденний рапорт.docx",
        freeform_paragraphs=[
            "матрос ДРУГИЙ Другий Другий виписаний із КНП «Приклад» проходить ВЛК, "
            "потребує відпустки за станом здоров'я на 30 календарних діб.",
        ],
    )
    events, unresolved = reader.extract_status_events(path)

    assert events == []
    assert len(unresolved) == 1
    assert unresolved[0]["pib_raw"] == "ДРУГИЙ Другий Другий"


def test_freeform_discharge_with_date_before_name_but_no_vlk_in_progress_mention_is_not_auto_applied(tmp_path):
    """Дата ПЕРЕД ПІБ Є (а не біля "виписан" - тут вона задалеко для
    _DISCHARGE_DATE_RE), і речення НЕ згадує "проходить ВЛК" - немає якоря
    для дня події взагалі (ні через _DISCHARGE_DATE_RE, ні через фолбек
    ВЛК-в-процесі, який вимагає САМЕ цю згадку), лишається на ручну
    перевірку, як і раніше."""
    path = write_daily_report(
        tmp_path / "28.08.2026 - щоденний рапорт.docx",
        freeform_paragraphs=[
            "28.08.2026 матрос ШОСТИЙ Шостий Шостий виписаний із КНП «Приклад», "
            "потребує відпустки за станом здоров'я на 30 календарних діб.",
        ],
    )
    events, unresolved = reader.extract_status_events(path)

    assert events == []
    assert len(unresolved) == 1
    assert unresolved[0]["pib_raw"] == "ШОСТИЙ Шостий Шостий"


def test_freeform_discharge_date_near_discharge_word_with_vlk_in_progress_produces_vlk_not_shp(tmp_path):
    """Дата БІЛЯ самого "виписан" (_DISCHARGE_DATE_RE, а не перед ПІБ) - і
    речення ТАКОЖ прямо каже "проходить ВЛК" - день виписки МАЄ стати "ВЛК",
    а НЕ "ШП" (на відміну від test_freeform_discharge_without_a_dated_vlk_
    reference_produces_shp_then_leave вище, де ЦІЄЇ згадки немає)."""
    path = write_daily_report(
        tmp_path / "28.08.2026 - щоденний рапорт.docx",
        freeform_paragraphs=[
            "матрос ДЕСЯТИЙ Десятий Десятий 28.08.2026 виписаний із КНП «Приклад» та проходить ВЛК, "
            "потребує відпустки за станом здоров'я на 30 календарних діб.",
        ],
    )
    events, unresolved = reader.extract_status_events(path)

    assert unresolved == []
    assert events == [
        {
            "pib_raw": "ДЕСЯТИЙ Десятий Десятий", "status": "ВЛК", "effective_date": date(2026, 8, 28),
            "heading": "поза межами", "report_date": date(2026, 8, 28),
        },
        {
            "pib_raw": "ДЕСЯТИЙ Десятий Десятий", "status": "ВПСЗ", "effective_date": date(2026, 8, 29),
            "heading": "поза межами", "report_date": date(2026, 8, 28),
        },
    ]


def test_freeform_discharge_then_awol_produces_single_szch_event_on_discharge_date(tmp_path):
    """Реальний випадок (22.08.2026 - щоденний рапорт.docx): "ДАТА виписаний
    із [заклад]. Після виписки до району виконання завдань не прибув, місце
    перебування не встановлено, самовільно залишив військову частину (СЗЧ)."
    - ДВА окремих речення (крапка МІЖ ними, а не один зворот, як для
    "вибув у СЗЧ" вище), тож раніше НІ freeform-патерни, НІ
    _discharge_vlk_leave_events (немає "потребує відпустки") цього не бачили,
    і речення даремно потрапляло на ручну перевірку. Тепер розпізнається як
    ОДНА подія - статус "СЗЧ" саме на дату ВИПИСКИ (а не на дату самого
    рапорту)."""
    path = write_daily_report(
        tmp_path / "22.08.2026 - щоденний рапорт.docx",
        freeform_paragraphs=[
            "17.08.2026 матрос ДВАДЦЯТЬПЕРШИЙ Двадцятьперший Двадцятьперший виписаний із Військово-медичного "
            "клінічного центру Центрального регіону. Після виписки до району виконання завдань не прибув, "
            "місце перебування не встановлено, самовільно залишив військову частину (СЗЧ).",
        ],
    )
    events, unresolved = reader.extract_status_events(path)

    assert unresolved == []
    assert events == [{
        "pib_raw": "ДВАДЦЯТЬПЕРШИЙ Двадцятьперший Двадцятьперший", "status": "СЗЧ", "effective_date": date(2026, 8, 17),
        "heading": "поза межами", "report_date": date(2026, 8, 22),
    }]


def test_freeform_discharge_then_awol_falls_back_to_date_next_to_the_discharge_word(tmp_path):
    """Той самий зворот, але з ІНШИМ порядком слів - дата стоїть ЩІЛЬНО біля
    "виписаний" (ПІСЛЯ ПІБ), а не перед ПІБ на початку абзацу - перед ПІБ
    дати взагалі немає, тож функція має впасти на запасний варіант
    (_DISCHARGE_DATE_RE, той самий, що й _discharge_vlk_leave_events)."""
    path = write_daily_report(
        tmp_path / "22.08.2026 - щоденний рапорт.docx",
        freeform_paragraphs=[
            "матрос ДВАДЦЯТЬДРУГИЙ Двадцятьдругий Двадцятьдругий 17.08.2026 виписаний із Військово-медичного "
            "клінічного центру Центрального регіону. Після виписки до району виконання завдань не прибув, "
            "місце перебування не встановлено, самовільно залишив військову частину (СЗЧ).",
        ],
    )
    events, unresolved = reader.extract_status_events(path)

    assert unresolved == []
    assert events == [{
        "pib_raw": "ДВАДЦЯТЬДРУГИЙ Двадцятьдругий Двадцятьдругий", "status": "СЗЧ", "effective_date": date(2026, 8, 17),
        "heading": "поза межами", "report_date": date(2026, 8, 22),
    }]


def test_freeform_discharge_then_awol_without_any_date_is_not_auto_applied(tmp_path):
    """AWOL-фраза є, але дати немає ЗОВСІМ (ні перед ПІБ, ні біля "виписан")
    - немає якоря для дня самої події, тож нічого не застосовується,
    лишається на ручну перевірку (той самий принцип, що й
    test_freeform_needs_leave_without_discharge_or_vlk_date_is_not_auto_applied)."""
    path = write_daily_report(
        tmp_path / "22.08.2026 - щоденний рапорт.docx",
        freeform_paragraphs=[
            "матрос ДВАДЦЯТЬТРЕТІЙ Двадцятьтретій Двадцятьтретій виписаний із лікувального закладу. "
            "Після виписки не прибув, самовільно залишив військову частину (СЗЧ).",
        ],
    )
    events, unresolved = reader.extract_status_events(path)

    assert events == []
    assert len(unresolved) == 1
    assert unresolved[0]["pib_raw"] == "ДВАДЦЯТЬТРЕТІЙ Двадцятьтретій Двадцятьтретій"


def test_freeform_discharge_without_awol_phrase_is_not_auto_applied_as_szch(tmp_path):
    """Виписка сама по собі (без фрази "самовільно залишив військову
    частину") - НЕ повинна ставати подією "СЗЧ" лише за згадкою виписки -
    надто неоднозначно (людина могла повернутись у частину, піти на ВЛК
    тощо - див. інші тести discharge вище); лишається на ручну перевірку."""
    path = write_daily_report(
        tmp_path / "22.08.2026 - щоденний рапорт.docx",
        freeform_paragraphs=[
            "17.08.2026 матрос ДВАДЦЯТЬПЕРШИЙ Двадцятьперший Двадцятьперший виписаний із лікувального закладу "
            "та повернувся до району виконання завдань.",
        ],
    )
    events, unresolved = reader.extract_status_events(path)

    assert events == []
    assert len(unresolved) == 1
    assert unresolved[0]["pib_raw"] == "ДВАДЦЯТЬПЕРШИЙ Двадцятьперший Двадцятьперший"


def test_signature_block_with_no_spacing_style_is_ignored_entirely(tmp_path):
    """Реальний випадок: рядок підпису ("Тимчасово виконуючий обов'язки
    командира ...", стиль "No Spacing") і звання+ПІБ під ним не стосуються
    жодного військовослужбовця зі складу - не повинні потрапляти в
    unresolved як "текст без розпізнаного ПІБ", і все ПІСЛЯ підпису теж
    ігнорується (у справжніх рапортах підпис завжди останній блок)."""
    path = write_daily_report(
        tmp_path / "01.08.2026 - щоденний рапорт.docx",
        freeform_paragraphs=["якийсь текст без розпізнаваного ПІБ."],
        signature=(f"Тимчасово виконуючий обов'язки командира {reader.SIGNATORY_UNIT_NAMES[0]}", "майор\tДесятий Десятий"),
    )
    events, unresolved = reader.extract_status_events(path)

    assert events == []
    assert len(unresolved) == 1  # лише "якийсь текст..." - НІЧОГО про сам підпис
    assert "розпізнаваного" not in unresolved[0]["reason"]  # підпис не додав СВІЙ запис


def test_signature_block_detected_by_content_even_without_no_spacing_style(tmp_path):
    """Той самий підпис, але БЕЗ стилю "No Spacing" (напр. інший шаблон
    рапорту) - все одно розпізнається за вмістом: роль підписанта
    (SIGNATORY_ROLE_KEYWORDS) + назва підрозділу (resources/data.json - тест
    НЕ хардкодить справжню назву підрозділу, а бере її з конфігурації)."""
    path = write_daily_report(
        tmp_path / "01.08.2026 - щоденний рапорт.docx",
        freeform_paragraphs=[f"{reader.SIGNATORY_ROLE_KEYWORDS[0].capitalize()} {reader.SIGNATORY_UNIT_NAMES[0]}"],
    )
    events, unresolved = reader.extract_status_events(path)

    assert events == []
    assert unresolved == []


def test_own_unit_mention_without_signatory_role_keyword_is_not_treated_as_signature(tmp_path):
    """Згадка ВЛАСНОГО підрозділу в реченні про людину (без слова ролі
    підписанта, напр. "командир") - НЕ повинна хибно зникати як "підпис"."""
    path = write_daily_report(
        tmp_path / "01.08.2026 - щоденний рапорт.docx",
        freeform_paragraphs=[f"матрос ДВАДЦЯТИЙ Двадцятий Двадцятий прибув до району виконання завдань {reader.SIGNATORY_UNIT_NAMES[0]}."],
    )
    events, unresolved = reader.extract_status_events(path)

    assert events == []
    assert len(unresolved) == 1
    assert unresolved[0]["pib_raw"] == "ДВАДЦЯТИЙ Двадцятий Двадцятий"


def test_report_addressee_line_at_top_never_triggers_signature_detection(tmp_path):
    """"Командиру [назва підрозділу]" (адресат, ПЕРШИЙ рядок рапорту, з
    docx_fixtures.write_daily_report) - написаний схоже на підпис (той самий
    корінь "командир" і назва підрозділу), але "Командиру" (давальний
    відмінок) - явний виняток у _is_signatory_paragraph, тож увесь звичайний
    вміст рапорту (і після цього рядка) обробляється як завжди."""
    path = write_daily_report(
        tmp_path / "01.08.2026 - щоденний рапорт.docx",
        groups=[(
            "ПРИБУЛИ до складу сил та засобів 9 армійського корпусу:",
            [("З відпустки:", ARRIVAL_HEADER, [["1", "сержант", "ПЕРШИЙ Перший Перший", "Посада", "01.08.2026", "Прибув"]])],
        )],
    )
    events, unresolved = reader.extract_status_events(path)

    assert unresolved == []
    assert events == [{
        "pib_raw": "ПЕРШИЙ Перший Перший", "status": "РВЗ", "effective_date": date(2026, 8, 1),
        "heading": "З відпустки:", "report_date": date(2026, 8, 1),
    }]


def test_table_after_signature_is_ignored(tmp_path):
    """Підпис - завжди ОСТАННІЙ блок у справжніх рапортах, але навіть якби
    після нього трапилась таблиця (напр. пошкоджений/нестандартний файл) -
    вона так само повністю ігнорується, а не обробляється як звичайна."""
    path = write_daily_report(
        tmp_path / "01.08.2026 - щоденний рапорт.docx",
        signature=(f"Тимчасово виконуючий обов'язки командира {reader.SIGNATORY_UNIT_NAMES[0]}", "майор\tДесятий Десятий"),
    )
    doc = Document(path)
    doc.add_paragraph("З відпустки:")
    table = doc.add_table(rows=2, cols=len(ARRIVAL_HEADER))
    for col_idx, text in enumerate(ARRIVAL_HEADER):
        table.rows[0].cells[col_idx].text = text
    table.rows[1].cells[2].text = "ПЕРШИЙ Перший Перший"
    table.rows[1].cells[4].text = "01.08.2026"
    doc.save(path)

    events, unresolved = reader.extract_status_events(path)

    assert events == []
    assert unresolved == []


def test_is_bare_status_heading_matches_keyword_with_trailing_colon():
    assert reader._is_bare_status_heading("СЗЧ:") is True
    assert reader._is_bare_status_heading("сзч:") is True
    assert reader._is_bare_status_heading("  СЗЧ:  ") is True


def test_is_bare_status_heading_rejects_a_real_sentence_that_merely_mentions_the_keyword():
    """substring-збіг тут НЕ повинен спрацьовувати (на відміну від
    _status_for_heading) - інакше СПРАВЖНЄ речення про конкретну людину, яке
    лише ЗГАДУЄ ключове слово, мовчки зникло б замість піти на ручну
    перевірку."""
    assert reader._is_bare_status_heading("ПЕРШИЙ Перший Перший переведений у підрозділ СЗЧ.") is False


def test_freeform_bare_subheading_without_a_name_is_not_sent_for_manual_review(tmp_path):
    """Реальний випадок: розділ "Поза межами..." має міні-підзаголовок
    "СЗЧ:" ПЕРЕД реченням про конкретну людину - сам підзаголовок ("СЗЧ:",
    без ПІБ) не описує подію, тож не повинен потрапляти в unresolved окремо
    (лише РЕЧЕННЯ про людину, якщо воно саме не розпізналось, обробляється
    як завжди)."""
    path = write_daily_report(
        tmp_path / "07.08.2026 - щоденний рапорт.docx",
        freeform_paragraphs=[
            "СЗЧ:",
            "матрос ДЕСЯТИЙ Десятий Десятий 07.08.2026 вибув у СЗЧ.",
        ],
    )
    events, unresolved = reader.extract_status_events(path)

    assert unresolved == []
    assert events == [{
        "pib_raw": "ДЕСЯТИЙ Десятий Десятий", "status": "СЗЧ", "effective_date": date(2026, 8, 7),
        "heading": "поза межами", "report_date": date(2026, 8, 7),
    }]


def test_confirmed_health_leave_sentence_produces_vpsz_event(tmp_path):
    """Реальний випадок: НАСТУПНОГО дня після рапорту "потребує відпустки за
    висновком ВЛК ..." (без прямої причини - _discharge_vlk_leave_events дає
    "ВП") - ОКРЕМИЙ рапорт підтверджує причину ОКРЕМИМ реченням "З ДАТА
    відпустка за станом здоров'я висновок ВЛК №... від ...:" - розпізнається
    як ОКРЕМА подія "ВПСЗ" на дату, вказану в самому реченні (не дата
    рапорту)."""
    path = write_daily_report(
        tmp_path / "11.08.2026 - щоденний рапорт.docx",
        freeform_paragraphs=[
            "старший матрос ДВАДЦЯТИЙ Двадцятий Двадцятий з 11.08.2026 відпустка за станом "
            "здоров'я висновок ВЛК №2026-0810-1200-0001-0 від 10.08.2026:",
        ],
    )
    events, unresolved = reader.extract_status_events(path)

    assert unresolved == []
    assert events == [{
        "pib_raw": "ДВАДЦЯТИЙ Двадцятий Двадцятий", "status": "ВПСЗ", "effective_date": date(2026, 8, 11),
        "heading": "поза межами", "report_date": date(2026, 8, 11),
    }]


def test_confirmed_health_leave_sentence_does_not_interfere_with_needs_leave_sentence(tmp_path):
    """"потребує відпустки за висновком ВЛК ..." (БЕЗ "за станом здоров'я" в
    ЦЬОМУ реченні) - і далі йде звичайним шляхом _discharge_vlk_leave_events
    ("ВЛК" тоді "ВП", не "ВПСЗ") - нова "підтверджувальна" перевірка НЕ
    повинна хибно спрацювати на ЦЕ речення."""
    path = write_daily_report(
        tmp_path / "10.08.2026 - щоденний рапорт.docx",
        freeform_paragraphs=[
            "старший матрос ДВАДЦЯТИЙ Двадцятий Двадцятий потребує відпустки за висновком ВЛК "
            "№2026-0810-1200-0001-0 від 10.08.2026.",
        ],
    )
    events, unresolved = reader.extract_status_events(path)

    assert unresolved == []
    assert events == [
        {
            "pib_raw": "ДВАДЦЯТИЙ Двадцятий Двадцятий", "status": "ВЛК", "effective_date": date(2026, 8, 10),
            "heading": "поза межами", "report_date": date(2026, 8, 10),
        },
        {
            "pib_raw": "ДВАДЦЯТИЙ Двадцятий Двадцятий", "status": "ВП", "effective_date": date(2026, 8, 11),
            "heading": "поза межами", "report_date": date(2026, 8, 10),
        },
    ]


def test_confirmed_health_leave_sentence_with_phrase_before_date_produces_vpsz_event(tmp_path):
    """Реальний випадок: підтверджувальне речення з ІНШИМ порядком слів -
    "відпустка за станом здоров'я З ДАТА" (фраза, потім дата - а не "З ДАТА
    відпустка...", як у test_confirmed_health_leave_sentence_produces_vpsz_
    event) - той самий результат: подія "ВПСЗ" на дату з речення."""
    path = write_daily_report(
        tmp_path / "14.08.2026 - щоденний рапорт.docx",
        freeform_paragraphs=[
            "матрос ТРИНАДЦЯТИЙ Тринадцятий Тринадцятий відпустка за станом здоров'я з 14.08.2026 на 30 календарних днів;",
        ],
    )
    events, unresolved = reader.extract_status_events(path)

    assert unresolved == []
    assert events == [{
        "pib_raw": "ТРИНАДЦЯТИЙ Тринадцятий Тринадцятий", "status": "ВПСЗ", "effective_date": date(2026, 8, 14),
        "heading": "поза межами", "report_date": date(2026, 8, 14),
    }]


def test_completed_vlk_then_dated_leave_sentence_produces_vpsz_event(tmp_path):
    """Реальний випадок (23.08.2026 - щоденний рапорт.docx): НАСТУПНИЙ рапорт
    переказує ТОЙ САМИЙ факт, що й "закінчив ВЛК ДАТА, потребує відпустки за
    станом здоров'я..." (_discharge_vlk_leave_events, date_after2 форма) з
    ІНШОГО, попереднього рапорту - але ІНШИМИ словами: "ДАТА1 звання ПІБ
    завершив проходження військово-лікарської комісії. З ДАТА2 направляється
    у відпустку для лікування ..." - БЕЗ "потребує" і БЕЗ "за станом
    здоров'я" (жоден наявний механізм цього не бачив, лишалось на ручну
    перевірку). Розпізнається як ОКРЕМА подія "ВПСЗ" на ДАТУ2, вказану ПРЯМО
    в реченні (не дата рапорту)."""
    path = write_daily_report(
        tmp_path / "23.08.2026 - щоденний рапорт.docx",
        freeform_paragraphs=[
            "21.08.2026 молодший сержант ДВАДЦЯТЬПЕРШИЙ Двадцятьперший Двадцятьперший завершив проходження "
            "військово-лікарської комісії. З 22.08.2026 направляється у відпустку для лікування у зв'язку з "
            "пораненням строком на 30 календарних діб.",
        ],
    )
    events, unresolved = reader.extract_status_events(path)

    assert unresolved == []
    assert events == [{
        "pib_raw": "ДВАДЦЯТЬПЕРШИЙ Двадцятьперший Двадцятьперший", "status": "ВПСЗ", "effective_date": date(2026, 8, 22),
        "heading": "поза межами", "report_date": date(2026, 8, 23),
    }]


def test_completed_vlk_then_left_on_leave_sentence_produces_vpsz_event(tmp_path):
    """Реальний випадок, підтверджений користувачем: ТОЙ САМИЙ "завершив/
    закінчив проходження ВЛК. З ДАТА2 направляється у відпустку" факт, що й
    у тесті вище, АЛЕ ІНШИМ дієсловом для самої відпустки - "ДАТА2 ВИБУВ у
    відпустку" (а не "з ДАТА2 направляється"), БЕЗ прийменника "з", дата
    ПЕРЕД дієсловом, і КОМА (а не крапка) між частинами речення.
    _DATED_TREATMENT_LEAVE_RE раніше не бачив цю форму (лише "направля-").
    Як і в тесті вище - ЛИШЕ "ВПСЗ" на дату відпустки, БЕЗ окремого дня
    "ВЛК" (комісію вже завершено, той самий "no anchor" принцип)."""
    path = write_daily_report(
        tmp_path / "01.09.2026 - щоденний рапорт.docx",
        freeform_paragraphs=[
            "31.08.2026 штаб-сержант ДВАНАДЦЯТИЙ Дванадцятий Дванадцятий закінчив проходження  військово-"
            "лікарської комісії, 01.09.2026 вибув у відпустку на 30 календарних діб за станом здоров'я по "
            "пораненню.",
        ],
    )
    events, unresolved = reader.extract_status_events(path)

    assert unresolved == []
    assert events == [{
        "pib_raw": "ДВАНАДЦЯТИЙ Дванадцятий Дванадцятий", "status": "ВПСЗ", "effective_date": date(2026, 9, 1),
        "heading": "поза межами", "report_date": date(2026, 9, 1),
    }]


def test_discharge_then_dated_leave_sentence_produces_vpsz_event(tmp_path):
    """Реальний випадок (баг, підтверджений користувачем): ТОЙ САМИЙ "з ДАТА
    направляється у відпустку для лікування ..." зворот, що й
    test_completed_vlk_then_dated_leave_sentence_produces_vpsz_event, але
    медичний контекст - "виписаний із [заклад]" (з лікарні), а НЕ "завершив
    ВЛК" - ОДНЕ речення (кома, а не крапка МІЖ ними, на відміну від
    _discharge_then_awol_event вище). Раніше жоден механізм цього не бачив
    (_discharge_vlk_leave_events вимагає "потребує відпустки", тут -
    "направляється"; _DISCHARGE_DATE_RE не знаходить дату щільно біля
    "виписан" - дата стоїть ПЕРЕД ПІБ, а "виписан" - вже після переліку
    установи). Тепер розпізнається як ОКРЕМА подія "ВПСЗ" на дату, вказану
    ПРЯМО в реченні для самої відпустки (не дата рапорту, не дата виписки)."""
    path = write_daily_report(
        tmp_path / "25.08.2026 - щоденний рапорт.docx",
        freeform_paragraphs=[
            "24.08.2026 матрос ДВАДЦЯТЬЧЕТВЕРТИЙ Двадцятьчетвертий Двадцятьчетвертий виписаний із КНП «Приклад», "
            "з 24.08.2026 направляється у відпустку для лікування у зв'язку з пораненням строком на 30 "
            "календарних діб.",
        ],
    )
    events, unresolved = reader.extract_status_events(path)

    assert unresolved == []
    assert events == [{
        "pib_raw": "ДВАДЦЯТЬЧЕТВЕРТИЙ Двадцятьчетвертий Двадцятьчетвертий", "status": "ВПСЗ", "effective_date": date(2026, 8, 24),
        "heading": "поза межами", "report_date": date(2026, 8, 25),
    }]


def test_discharge_then_dated_leave_with_vlk_in_progress_adds_vlk_anchor_event(tmp_path):
    """Реальний випадок, підтверджений користувачем: ТОЙ САМИЙ "виписаний ...
    з ДАТА направляється у відпустку ..." зворот, що й
    test_discharge_then_dated_leave_sentence_produces_vpsz_event, АЛЕ
    речення ТАКОЖ прямо каже "приступив до проходження ВЛК" - на відміну
    від ТОГО тесту (лише "ВПСЗ" на дату відпустки, людина вважається "вже
    вийшла з медичного процесу"), тут ДОДАТКОВО з'являється подія "ВЛК" на
    дату виписки/початку комісії (_vlk_in_progress_anchor_event) - людина
    ЩОЙНО РОЗПОЧАЛА комісію, день виписки МАЄ бути "ВЛК", а не лишатись
    попереднім статусом."""
    path = write_daily_report(
        tmp_path / "25.08.2026 - щоденний рапорт.docx",
        freeform_paragraphs=[
            "24.08.2026 матрос ЧЕТВЕРТИЙ Четвертий Четвертий виписаний із КНП «Приклад» та приступив до "
            "проходження ВЛК, з 25.08.2026 направляється у відпустку для лікування у зв'язку з пораненням "
            "строком на 30 календарних діб.",
        ],
    )
    events, unresolved = reader.extract_status_events(path)

    assert unresolved == []
    assert events == [
        {
            "pib_raw": "ЧЕТВЕРТИЙ Четвертий Четвертий", "status": "ВЛК", "effective_date": date(2026, 8, 24),
            "heading": "поза межами", "report_date": date(2026, 8, 25),
        },
        {
            "pib_raw": "ЧЕТВЕРТИЙ Четвертий Четвертий", "status": "ВПСЗ", "effective_date": date(2026, 8, 25),
            "heading": "поза межами", "report_date": date(2026, 8, 25),
        },
    ]


def test_discharge_then_dated_leave_with_vlk_in_progress_but_same_day_skips_anchor_event(tmp_path):
    """ТОЙ САМИЙ "проходить"/"приступив до проходження ВЛК" зворот, що й
    вище, АЛЕ дата виписки й дата відпустки в реченні - ОДНА й ТА САМА
    (обидві "24.08.2026", як у test_discharge_then_dated_leave_sentence_
    produces_vpsz_event) - ДОДАТКОВА подія "ВЛК" НЕ додається (та сама дата
    не може бути ОДНОЧАСНО "ВЛК" і "ВПСЗ" - суперечність), лишається лише
    "ВПСЗ", як і без згадки ВЛК."""
    path = write_daily_report(
        tmp_path / "25.08.2026 - щоденний рапорт.docx",
        freeform_paragraphs=[
            "24.08.2026 матрос ДВАДЦЯТЬТРЕТІЙ Двадцятьтретій Двадцятьтретій виписаний із КНП «Приклад» та "
            "приступив до проходження ВЛК, з 24.08.2026 направляється у відпустку для лікування у зв'язку з "
            "пораненням строком на 30 календарних діб.",
        ],
    )
    events, unresolved = reader.extract_status_events(path)

    assert unresolved == []
    assert events == [{
        "pib_raw": "ДВАДЦЯТЬТРЕТІЙ Двадцятьтретій Двадцятьтретій", "status": "ВПСЗ", "effective_date": date(2026, 8, 24),
        "heading": "поза межами", "report_date": date(2026, 8, 25),
    }]


def test_confirmed_health_leave_with_vlk_in_progress_adds_vlk_anchor_event(tmp_path):
    """Реальний випадок, підтверджений користувачем: "ПІБ ДАТА1 виписаний ...
    та проходить ВЛК, з ДАТА2 відпустка за станом здоров'я ..." - БЕЗ
    "потребує"/"направляється" (голий іменник "відпустка" - _CONFIRMED_
    HEALTH_LEAVE_RE, а не _NEEDS_LEAVE_RE/_DATED_TREATMENT_LEAVE_RE) - той
    самий "проходить ВЛК" фолбек ДОДАЄ подію "ВЛК" на ДАТУ1 (тут - через
    _DISCHARGE_DATE_RE, дата ЩІЛЬНО біля "виписан"), а НЕ лише "ВПСЗ" на
    ДАТУ2."""
    path = write_daily_report(
        tmp_path / "15.08.2026 - щоденний рапорт.docx",
        freeform_paragraphs=[
            "матрос ШОСТИЙ Шостий Шостий 14.08.2026 виписаний з КНП «Приклад» та проходить ВЛК, "
            "з 15.08.2026 відпустка за станом здоров'я на 30 календарних днів.",
        ],
    )
    events, unresolved = reader.extract_status_events(path)

    assert unresolved == []
    assert events == [
        {
            "pib_raw": "ШОСТИЙ Шостий Шостий", "status": "ВЛК", "effective_date": date(2026, 8, 14),
            "heading": "поза межами", "report_date": date(2026, 8, 15),
        },
        {
            "pib_raw": "ШОСТИЙ Шостий Шостий", "status": "ВПСЗ", "effective_date": date(2026, 8, 15),
            "heading": "поза межами", "report_date": date(2026, 8, 15),
        },
    ]


def test_confirmed_health_leave_with_vlk_in_progress_but_no_anchor_date_skips_anchor_event(tmp_path):
    """"проходить ВЛК" згадано, але немає ЖОДНОЇ дати ні біля "виписан"
    (заклад назвний задовго - понад 40 символів до дати відпустки), ні
    перед самим ПІБ (абзац починається одразу зі звання й ПІБ) - немає
    якоря для дня ВЛК, лишається лише "ВПСЗ" на дату відпустки, як і без
    згадки ВЛК узагалі."""
    path = write_daily_report(
        tmp_path / "20.08.2026 - щоденний рапорт.docx",
        freeform_paragraphs=[
            "матрос ОДИНАДЦЯТИЙ Одинадцятий Одинадцятий виписаний із дуже довгої назви медичного закладу для "
            "гарантованого перевищення сорока символів та проходить ВЛК, з 20.08.2026 відпустка за станом "
            "здоров'я на 30 днів.",
        ],
    )
    events, unresolved = reader.extract_status_events(path)

    assert unresolved == []
    assert events == [{
        "pib_raw": "ОДИНАДЦЯТИЙ Одинадцятий Одинадцятий", "status": "ВПСЗ", "effective_date": date(2026, 8, 20),
        "heading": "поза межами", "report_date": date(2026, 8, 20),
    }]


def test_completed_vlk_mention_without_dated_leave_phrase_is_not_auto_applied(tmp_path):
    """Згадка завершення ВЛК є, але БЕЗ "З ДАТА направляється у відпустку" -
    надто неоднозначно (могло бути ІНШЕ продовження), тож нічого не
    застосовується, лишається на ручну перевірку."""
    path = write_daily_report(
        tmp_path / "23.08.2026 - щоденний рапорт.docx",
        freeform_paragraphs=[
            "21.08.2026 молодший сержант ДВАДЦЯТЬПЕРШИЙ Двадцятьперший Двадцятьперший завершив проходження "
            "військово-лікарської комісії.",
        ],
    )
    events, unresolved = reader.extract_status_events(path)

    assert events == []
    assert len(unresolved) == 1
    assert unresolved[0]["pib_raw"] == "ДВАДЦЯТЬПЕРШИЙ Двадцятьперший Двадцятьперший"


def test_dated_leave_phrase_without_medical_context_mention_is_not_auto_applied(tmp_path):
    """"З ДАТА направляється у відпустку ..." є, але БЕЗ медичного контексту
    (ні завершення ВЛК, ні "виписаний") - надто неоднозначно (причина
    відпустки невідома, могла бути звичайна, немедична відпустка), тож
    нічого не застосовується, лишається на ручну перевірку."""
    path = write_daily_report(
        tmp_path / "23.08.2026 - щоденний рапорт.docx",
        freeform_paragraphs=[
            "молодший сержант ДВАДЦЯТЬПЕРШИЙ Двадцятьперший Двадцятьперший. "
            "З 22.08.2026 направляється у відпустку для лікування строком на 30 календарних діб.",
        ],
    )
    events, unresolved = reader.extract_status_events(path)

    assert events == []
    assert len(unresolved) == 1
    assert unresolved[0]["pib_raw"] == "ДВАДЦЯТЬПЕРШИЙ Двадцятьперший Двадцятьперший"


def test_health_leave_then_vlk_sentence_produces_vpsz_and_vlk_events(tmp_path):
    """Реальний випадок: "з ДАТА1 у відпустці за станом здоров'я та ДАТА2 на
    проходження військово-лікарської комісії" - ОДНЕ речення, ДВІ послідовні
    події з ОБОМА датами явно вказаними: "ВПСЗ" з ДАТА1, "ВЛК" з ДАТА2
    (комісія визначає подальший статус) - на відміну від
    _CONFIRMED_HEALTH_LEAVE_RE (дата йде БЕЗПОСЕРЕДНЬО перед фразою), тут між
    датою й "відпустці" є слово "у"."""
    path = write_daily_report(
        tmp_path / "20.08.2026 - щоденний рапорт.docx",
        freeform_paragraphs=[
            "молодший сержант ВІСІМНАДЦЯТИЙ Вісімнадцятий Вісімнадцятий з 21.07.2026 у відпустці за станом "
            "здоров'я та 20.08.2026 на проходження військово-лікарської комісії.",
        ],
    )
    events, unresolved = reader.extract_status_events(path)

    assert unresolved == []
    assert events == [
        {
            "pib_raw": "ВІСІМНАДЦЯТИЙ Вісімнадцятий Вісімнадцятий", "status": "ВПСЗ", "effective_date": date(2026, 7, 21),
            "heading": "поза межами", "report_date": date(2026, 8, 20),
        },
        {
            "pib_raw": "ВІСІМНАДЦЯТИЙ Вісімнадцятий Вісімнадцятий", "status": "ВЛК", "effective_date": date(2026, 8, 20),
            "heading": "поза межами", "report_date": date(2026, 8, 20),
        },
    ]


def test_health_leave_then_vlk_sentence_with_end_date_range_and_line_break(tmp_path):
    """Реальний випадок: "з ДАТА1\\nпо ДАТА_кінця відпустці за станом
    здоров'я, з ДАТА2 року направлений на проходження військово-лікарської
    комісії" - переніс рядка ВСЕРЕДИНІ речення (між "з ДАТА1" і "по
    ДАТА_кінця") і ДОДАТКОВА "по ДАТА_кінця" (кінець діапазону відпустки,
    сама по собі НЕ використовується - forward-fill і так продовжує "ВПСЗ" аж
    до наступної події ("ВЛК" з ДАТА2)), і "року направлений" (два слова, не
    одне) між ДАТА2 і "на проходження" - _HEALTH_LEAVE_DATED_RE/
    _VLK_REFERRAL_DATED_RE мають ПРАВИЛЬНО вибрати дату ПОЧАТКУ кожної події
    (ДАТА1/ДАТА2), а не найближчу дату в реченні взагалі (ДАТА_кінця)."""
    path = write_daily_report(
        tmp_path / "20.08.2026 - щоденний рапорт.docx",
        freeform_paragraphs=[
            "молодший сержант ДВАДЦЯТЬОДИН Двадцятьодин Двадцятьодин з 21.07.2026\n"
            "по 19.08.2026 відпустці за станом здоров'я, з 20.08.2026 року направлений на "
            "проходження військово-лікарської комісії.",
        ],
    )
    events, unresolved = reader.extract_status_events(path)

    assert unresolved == []
    assert events == [
        {
            "pib_raw": "ДВАДЦЯТЬОДИН Двадцятьодин Двадцятьодин", "status": "ВПСЗ", "effective_date": date(2026, 7, 21),
            "heading": "поза межами", "report_date": date(2026, 8, 20),
        },
        {
            "pib_raw": "ДВАДЦЯТЬОДИН Двадцятьодин Двадцятьодин", "status": "ВЛК", "effective_date": date(2026, 8, 20),
            "heading": "поза межами", "report_date": date(2026, 8, 20),
        },
    ]


def test_health_leave_mention_without_vlk_referral_stays_unresolved(tmp_path):
    """"з ДАТА у відпустці за станом здоров'я" САМА ПО СОБІ (без ДАТОВАНОЇ
    згадки ВЛК поруч) - НЕ розпізнається цим механізмом (ОБИДВІ фрази мають
    бути присутні РАЗОМ) - лишається на ручну перевірку, як і раніше."""
    path = write_daily_report(
        tmp_path / "20.08.2026 - щоденний рапорт.docx",
        freeform_paragraphs=[
            "молодший сержант ДВАДЦЯТИЙ Двадцятий Двадцятий з 21.07.2026 у відпустці за станом здоров'я.",
        ],
    )
    events, unresolved = reader.extract_status_events(path)

    assert events == []
    assert len(unresolved) == 1


def test_freeform_continued_business_trip_resolves_to_vd_even_though_status_does_not_change(tmp_path):
    """Реальний випадок: "продовжено відрядження ... від ДАТА" - відрядження
    й далі триває (був "ВД", лишається "ВД") - раніше даремно потрапляло на
    ручну перевірку. Тепер розпізнається як подія (статус "ВД" незалежно
    від БАЗОВОГО статусу - той самий "безпечний no-op" принцип, що й
    госпіталізація до іншого закладу)."""
    path = write_daily_report(
        tmp_path / "14.08.2026 - щоденний рапорт.docx",
        freeform_paragraphs=[
            "матрос ЧОТИРНАДЦЯТИЙ Чотирнадцятий Чотирнадцятий продовжено відрядження на 30 календарних діб від 14.08.2026",
        ],
    )
    events, unresolved = reader.extract_status_events(path)

    assert unresolved == []
    assert events == [{
        "pib_raw": "ЧОТИРНАДЦЯТИЙ Чотирнадцятий Чотирнадцятий", "status": "ВД", "effective_date": date(2026, 8, 14),
        "heading": "поза межами", "report_date": date(2026, 8, 14),
    }]


def test_freeform_rehabilitation_continuation_resolves_to_shp_even_though_status_does_not_change(tmp_path):
    """Реальний випадок: виписаний з ОДНОГО закладу та "від ДАТА проходить
    реабілітацію" в ІНШОМУ - людина й далі лікується (був "ШП", лишається
    "ШП") - раніше даремно потрапляло на ручну перевірку. Тепер
    розпізнається як подія (статус "ШП" незалежно від БАЗОВОГО статусу)."""
    path = write_daily_report(
        tmp_path / "15.08.2026 - щоденний рапорт.docx",
        freeform_paragraphs=[
            "старший матрос ШІСТНАДЦЯТИЙ Шістнадцятий Шістнадцятий 13.08.2026 виписаний з медичного закладу B0001 "
            "та від 15.08.2026 проходить реабілітацію у медичному закладі B0001.",
        ],
    )
    events, unresolved = reader.extract_status_events(path)

    assert unresolved == []
    assert events == [{
        "pib_raw": "ШІСТНАДЦЯТИЙ Шістнадцятий Шістнадцятий", "status": "ШП", "effective_date": date(2026, 8, 15),
        "heading": "поза межами", "report_date": date(2026, 8, 15),
    }]


def test_freeform_rehabilitation_continuation_accepts_z_preposition_too(tmp_path):
    """Реальний випадок, підтверджений користувачем: "...та З 02.09.2026
    проходить реабілітацію..." - прийменник "з" (а не лише "від", тест
    вище) - той самий сенс, той самий "ШП"."""
    path = write_daily_report(
        tmp_path / "02.09.2026 - щоденний рапорт.docx",
        freeform_paragraphs=[
            "старший матрос СЬОМИЙ Сьомий Сьомий виписаний 02.09.2026 із медичного закладу B0001 та "
            "з 02.09.2026 проходить реабілітацію у медичному закладі B0001.",
        ],
    )
    events, unresolved = reader.extract_status_events(path)

    assert unresolved == []
    assert events == [{
        "pib_raw": "СЬОМИЙ Сьомий Сьомий", "status": "ШП", "effective_date": date(2026, 9, 2),
        "heading": "поза межами", "report_date": date(2026, 9, 2),
    }]


def test_freeform_started_vlk_mention_resolves_to_vlk_status(tmp_path):
    """Реальний випадок, підтверджений користувачем: "ДАТА почав проходження
    військово-лікарської комісії у військовій частині ..." - людина ЩОЙНО
    розпочала комісію, БЕЗ жодної згадки виписки чи потреби відпустки в
    тому ж реченні (складніший _discharge_vlk_leave_events на таке НЕ
    зреагував би - гейт "потребує відпустки" відсутній) - просте одноподієве
    правило FREEFORM_STATUS_PATTERNS."""
    path = write_daily_report(
        tmp_path / "04.09.2026 - щоденний рапорт.docx",
        freeform_paragraphs=[
            "сержант ВОСЬМИЙ Восьмий Восьмий 04.09.2026 почав проходження військово-лікарської "
            "комісії у військовій частині A0001",
        ],
    )
    events, unresolved = reader.extract_status_events(path)

    assert unresolved == []
    assert events == [{
        "pib_raw": "ВОСЬМИЙ Восьмий Восьмий", "status": "ВЛК", "effective_date": date(2026, 9, 4),
        "heading": "поза межами", "report_date": date(2026, 9, 4),
    }]


def test_freeform_passed_vlk_mention_without_hyphen_resolves_to_vlk_status(tmp_path):
    """Реальний випадок, підтверджений користувачем: "виписаний ДАТА із
    [заклад] та з ДАТА пройшов військово лікарську комісію ..." - ІНША
    граматична форма ("пройшов", БЕЗ проміжного іменника "проходження", і
    "військово лікарську" - БЕЗ дефіса, двома окремими словами), той самий
    сенс/статус "ВЛК", що й "почав проходження військово-лікарської
    комісії" (тест вище). Сама "виписаний" частина речення НЕ відповідає
    жодному іншому freeform-звороту сама по собі, тож НЕ заважає цьому."""
    path = write_daily_report(
        tmp_path / "02.09.2026 - щоденний рапорт.docx",
        freeform_paragraphs=[
            "старший матрос ОДИНАДЦЯТИЙ Одинадцятий Одинадцятий виписаний 02.09.2026 із медичного закладу "
            "B0001 та з 02.09.2026 пройшов військово лікарську комісію у медичному закладі B0001.",
        ],
    )
    events, unresolved = reader.extract_status_events(path)

    assert unresolved == []
    assert events == [{
        "pib_raw": "ОДИНАДЦЯТИЙ Одинадцятий Одинадцятий", "status": "ВЛК", "effective_date": date(2026, 9, 2),
        "heading": "поза межами", "report_date": date(2026, 9, 2),
    }]


def test_freeform_family_circumstances_leave_resolves_to_vps_status(tmp_path):
    """Реальний випадок, підтверджений користувачем: "з ДАТА відпустка за
    сімейними обставинами." - без жодного дієслова ("вибув"/"перебуває") -
    саме речення лише НАЗИВАЄ статус "ВПС" ("Відпустка за сімейними
    обставинами" - окремий, вже підтримуваний статус, відмінний від "ВП"/
    "ВПСЗ")."""
    path = write_daily_report(
        tmp_path / "02.09.2026 - щоденний рапорт.docx",
        freeform_paragraphs=[
            "матрос ДЕСЯТИЙ Десятий Десятий з 02.09.2026 відпустка за сімейними обставинами.",
        ],
    )
    events, unresolved = reader.extract_status_events(path)

    assert unresolved == []
    assert events == [{
        "pib_raw": "ДЕСЯТИЙ Десятий Десятий", "status": "ВПС", "effective_date": date(2026, 9, 2),
        "heading": "поза межами", "report_date": date(2026, 9, 2),
    }]


def test_freeform_rtgr_departure_with_a_named_group_resolves_to_rtgr_status(tmp_path):
    """За прямою вказівкою користувача - НА МАЙБУТНЄ: той самий факт, що й
    заголовок таблиці "В РТГР:" (STATUS_HEADING_RULES), може колись
    з'явитись і ВІЛЬНИМ ТЕКСТОМ - "вибув до ртгр [назва угрупування]" -
    статус "РТГр" незалежно від назви угрупування після (навмисно не
    хардкодиться - див. коментар у constants.py)."""
    path = write_daily_report(
        tmp_path / "11.09.2026 - щоденний рапорт.docx",
        freeform_paragraphs=[
            "молодший сержант СОРОКОВИЙ Сороковий Сороковий 11.09.2026 вибув до ртгр «Приклад».",
        ],
    )
    events, unresolved = reader.extract_status_events(path)

    assert unresolved == []
    assert events == [{
        "pib_raw": "СОРОКОВИЙ Сороковий Сороковий", "status": "РТГр", "effective_date": date(2026, 9, 11),
        "heading": "поза межами", "report_date": date(2026, 9, 11),
    }]


def test_freeform_bare_rtgr_mention_resolves_to_rtgr_status(tmp_path):
    """Той самий "ртгр" підрядок, коротше формулювання - БЕЗ "вибув до"
    (напр. "перебуває в РТГр") - той самий статус "РТГр"."""
    path = write_daily_report(
        tmp_path / "11.09.2026 - щоденний рапорт.docx",
        freeform_paragraphs=[
            "старший сержант ШІСТДЕСЯТИЙ Шістдесятий Шістдесятий 11.09.2026 перебуває в РТГр.",
        ],
    )
    events, unresolved = reader.extract_status_events(path)

    assert unresolved == []
    assert events == [{
        "pib_raw": "ШІСТДЕСЯТИЙ Шістдесятий Шістдесятий", "status": "РТГр", "effective_date": date(2026, 9, 11),
        "heading": "поза межами", "report_date": date(2026, 9, 11),
    }]


def test_freeform_paragraph_without_a_recognizable_name_still_kept_for_review(tmp_path):
    path = write_daily_report(
        tmp_path / "01.08.2026 - щоденний рапорт.docx",
        freeform_paragraphs=["якийсь текст без розпізнаваного ПІБ."],
    )
    events, unresolved = reader.extract_status_events(path)

    assert events == []
    assert len(unresolved) == 1
    assert unresolved[0]["pib_raw"] is None


def test_freeform_paragraph_with_name_but_no_matching_phrase_still_kept_for_review(tmp_path):
    path = write_daily_report(
        tmp_path / "01.08.2026 - щоденний рапорт.docx",
        freeform_paragraphs=["матрос ДЕСЯТИЙ Десятий Десятий переведений в інший підрозділ без подробиць."],
    )
    events, unresolved = reader.extract_status_events(path)

    assert events == []
    assert len(unresolved) == 1
    assert unresolved[0]["pib_raw"] == "ДЕСЯТИЙ Десятий Десятий"
    assert "Поза межами" in unresolved[0]["reason"]
    assert "raw_text" in unresolved[0]


def test_blank_paragraph_is_skipped(tmp_path):
    path = write_daily_report(tmp_path / "01.08.2026 - щоденний рапорт.docx", freeform_paragraphs=["", "текст без ПІБ."])
    events, unresolved = reader.extract_status_events(path)

    assert events == []
    assert len(unresolved) == 1  # лише непорожній абзац


def test_table_row_with_empty_pib_is_skipped(tmp_path):
    path = write_daily_report(
        tmp_path / "01.08.2026 - щоденний рапорт.docx",
        groups=[(
            "ПРИБУЛИ до складу сил та засобів 9 армійського корпусу:",
            [("З відпустки:", ARRIVAL_HEADER, [
                ["1", "сержант", "", "Посада", "01.08.2026", "Прибув"],
                ["2", "сержант", "ПЕРШИЙ Перший Перший", "Посада", "01.08.2026", "Прибув"],
            ])],
        )],
    )
    events, unresolved = reader.extract_status_events(path)

    assert unresolved == []
    assert len(events) == 1
    assert events[0]["pib_raw"] == "ПЕРШИЙ Перший Перший"


def test_table_without_recognizable_pib_column_is_reported_as_unresolved(tmp_path):
    path = write_daily_report(
        tmp_path / "01.08.2026 - щоденний рапорт.docx",
        groups=[(
            "ПРИБУЛИ до складу сил та засобів 9 армійського корпусу:",
            [("З відпустки:", ["№", "Звання", "ПІБ???", "Посада", "Дата", "Куди"], [["1", "сержант", "ПЕРШИЙ Перший Перший", "Посада", "01.08.2026", "Прибув"]])],
        )],
    )
    events, unresolved = reader.extract_status_events(path)

    assert events == []
    assert len(unresolved) == 1
    assert "колонку ПІБ" in unresolved[0]["reason"]


def test_multiple_groups_and_rows_in_one_report(tmp_path):
    path = write_daily_report(
        tmp_path / "05.08.2026 - щоденний рапорт.docx",
        groups=[
            (
                "ПРИБУЛИ до складу сил та засобів 9 армійського корпусу:",
                [("З відрядження:", ARRIVAL_HEADER, [
                    ["1", "старший сержант", "ПЕРШИЙ Перший Перший", "Посада", "05.08.2026", "Прибув"],
                    ["2", "молодший сержант", "ДРУГИЙ Другий Другий", "Посада", "05.08.2026", "Прибув"],
                ])],
            ),
            (
                "ВИБУЛИ зі складу сил та засобів 9 армійського корпусу:",
                [("У відпустку:", DEPARTURE_HEADER, [
                    ["1", "матрос", "ТРЕТІЙ Третій Третій", "Посада", "05.08.2026", "Вибув"],
                ])],
            ),
        ],
    )
    events, unresolved = reader.extract_status_events(path)

    assert unresolved == []
    statuses = {event["pib_raw"]: event["status"] for event in events}
    assert statuses == {
        "ПЕРШИЙ Перший Перший": "РВЗ", "ДРУГИЙ Другий Другий": "РВЗ", "ТРЕТІЙ Третій Третій": "ВП",
    }
