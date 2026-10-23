from datetime import date, datetime

import pytest
from book61_fixtures import write_full_book61
from cvs_fixtures import VLK_SHEET, write_cvs
from openpyxl import Workbook, load_workbook

import content.addon_mismatch_checker as amc
import content.cvs_reader as cr
import content.training_journal_reader as tjr
from content.addon_reader import load_addon_workbook
from generators.check_addon_mismatches import check_addon_mismatches_file


@pytest.fixture(autouse=True)
def _own_unit(monkeypatch):
    monkeypatch.setattr(cr, "OWN_MILITARY_UNIT_CODES", frozenset({"0000"}))


def write_journal(path, rows_by_sheet):
    """Синтетичний журнал обліку БП: аркуш -> рядки під заголовком "ПІБ"/"З"/"До"(+ ще колонка)."""
    wb = Workbook()
    wb.remove(wb.active)
    for sheet_name, (header, rows) in rows_by_sheet.items():
        ws = wb.create_sheet(sheet_name)
        ws.append(header)
        for row in rows:
            ws.append(row)
    wb.save(str(path))
    return str(path)


def test_reads_every_journal_sheet_as_its_status_period(tmp_path):
    """За прямою вказівкою користувача: БЗВП у власній в/ч (колонка "НЦ") -
    "БЗВП" (10), у навчальному центрі - "Навч"; фахова підготовка за кордоном -
    категорія 10, в НЦ - "Навч"; адаптація - "Адап"."""
    path = write_journal(tmp_path / "journal.xlsx", {
        "БЗВП": (["ПІБ", "З", "До", "НЦ"], [
            ["ПЕРШИЙ Перший Перший", datetime(2026, 10, 1), datetime(2026, 10, 20), "A0000"],
            ["ДРУГИЙ Другий Другий", datetime(2026, 10, 5), None, "198"],
            [None, datetime(2026, 10, 5), None, None],
        ]),
        "Адаптація": (["ПІБ", "З", "До"], [["ТРЕТІЙ Третій Третій", "з 03.10.2026", datetime(2026, 10, 9)]]),
        "Фахова підготовка ": (["ПІБ", "З", "До", "Фах", "НЦ"], [
            ["ЧЕТВЕРТИЙ Четвертий Четвертий", datetime(2026, 10, 2), datetime(2026, 10, 29), "Водій", "НЦ"],
            ["ШОСТИЙ Шостий Шостий", datetime(2026, 10, 2), datetime(2026, 10, 9), "Країна1", " Закордон "],
        ]),
        "Інше": (["ПІБ", "З", "До"], [["П'ЯТИЙ П'ятий П'ятий", datetime(2026, 10, 1), None]]),
    })

    records = tjr.read_all_period_records(load_addon_workbook(path), "journal")

    assert [(r["pib_raw"], r["status"], r["start_date"], r["end_date"], r["source_sheet"], r["arrival_date"],
             r["source_file"]) for r in records] == [
        ("ПЕРШИЙ Перший Перший", "БЗВП", date(2026, 10, 1), date(2026, 10, 20), "БЗВП", date(2026, 10, 20), "journal"),
        ("ДРУГИЙ Другий Другий", "Навч", date(2026, 10, 5), None, "БЗВП", None, "journal"),
        ("ТРЕТІЙ Третій Третій", "Адап", date(2026, 10, 3), date(2026, 10, 9), "Адаптація", date(2026, 10, 9), "journal"),
        ("ЧЕТВЕРТИЙ Четвертий Четвертий", "Навч", date(2026, 10, 2), date(2026, 10, 29), "Фахова підготовка",
         date(2026, 10, 29), "journal"),
        ("ШОСТИЙ Шостий Шостий", 10, date(2026, 10, 2), date(2026, 10, 9), "Фахова підготовка",
         date(2026, 10, 9), "journal"),
    ]


def test_an_expected_payment_category_matches_every_status_paid_by_it():
    """Фахова підготовка за кордоном - "можна залишати 10": збіг і для 10, і
    для статусу, що оплачується категорією 10; "Навч"/100/РВЗ - розбіжність."""
    for actual in (10, "ППД", "БЗВП", "НОВ", "Адап", "ВД"):
        assert amc._statuses_match(10, actual) is True, actual
    for actual in ("Навч", 100, "РВЗ", "10_МЕД"):
        assert amc._statuses_match(10, actual) is False, actual


def test_journal_is_recognized_by_its_sheets_and_columns_only(tmp_path):
    journal = write_journal(tmp_path / "journal.xlsx", {"Адаптація": (["ПІБ", "З", "До"], [])})
    assert tjr.is_training_journal_workbook(load_addon_workbook(journal))
    wrong_columns = write_journal(tmp_path / "other.xlsx", {"БЗВП": (["ПІБ", "Дата"], [])})
    assert not tjr.is_training_journal_workbook(load_addon_workbook(wrong_columns))
    assert not tjr.is_training_journal_workbook(load_addon_workbook(write_cvs(tmp_path / "cvs.xlsx", {VLK_SHEET: []})))
    assert not tjr.is_training_journal_workbook(load_addon_workbook(write_full_book61(tmp_path / "book61.xlsm", {})))


def test_training_statuses_match_their_payment_category_and_identical_statuses():
    """Табель для виплат пише "БЗВП"/"Адап"/"НОВ" числом 10; "НОВ" і "БЗВП"
    ідентичні (обидва = ППД); "Навч" числом 10 не стає - це й далі розбіжність."""
    for expected in ("БЗВП", "Адап", "НОВ"):
        assert amc._statuses_match(expected, 10) is True
    assert amc._statuses_match("БЗВП", "НОВ") is True
    assert amc._statuses_match("Адап", "Адап.") is True
    assert amc._statuses_match("БЗВП", "РВЗ") is False
    assert amc._statuses_match("Навч", 10) is False


def test_check_addon_mismatches_file_reports_the_journal_on_its_own_sheet(tmp_path):
    header = ["ПОСАДА", "ЗВАННЯ", "ПІБ"] + [datetime(2026, 10, day) for day in range(1, 5)]
    oblik_wb = Workbook()
    ws = oblik_wb.active
    ws.title = "Табель"
    ws.append(header)
    ws.append(["Посада1", "сержант", "ПЕРШИЙ Перший Перший", 10, "НОВ", "РВЗ", 100])
    oblik = str(tmp_path / "ОБЛІК.xlsx")
    oblik_wb.save(oblik)
    journal = write_journal(tmp_path / "Журнал.xlsx", {
        "БЗВП": (["ПІБ", "З", "До", "НЦ"], [["ПЕРШИЙ Перший Перший", datetime(2026, 10, 1), datetime(2026, 10, 4), "A0000"]]),
    })
    report = str(tmp_path / "mis_match.xlsx")

    result = check_addon_mismatches_file(oblik, [journal], report_file_name=report)

    assert result.checked_files == [journal]
    bad = {d.day: (actual, expected) for record in result.period_mismatches
           for d, (actual, expected) in record["dates"].items() if not amc._statuses_match(expected, actual)}
    assert bad == {3: ("РВЗ", "БЗВП")}  # 04.10 - день повернення ("До"), 100 - не розбіжність
    assert load_workbook(report).sheetnames == ["Журнал"]
