import pytest
from openpyxl import Workbook, load_workbook

from content.excel_compat import empty_workbook, save_workbook, workbook_with_sheet


def test_workbook_with_sheet_has_exactly_one_sheet_with_the_given_title():
    wb, ws = workbook_with_sheet("Розбіжності")

    assert wb.sheetnames == ["Розбіжності"]
    assert ws is wb["Розбіжності"]


def test_empty_workbook_has_no_sheets():
    assert empty_workbook().sheetnames == []


def test_save_workbook_creates_the_missing_folder(tmp_path):
    wb, ws = workbook_with_sheet("Аркуш")
    ws.cell(row=1, column=1, value="значення")
    path = tmp_path / "нова тека" / "книга.xlsx"

    save_workbook(wb, str(path))

    assert load_workbook(str(path))["Аркуш"].cell(row=1, column=1).value == "значення"


def test_save_workbook_explains_a_file_locked_by_another_program(tmp_path, monkeypatch):
    """Файл відкритий в Excel - зрозуміле повідомлення замість сирого traceback."""
    def _locked(*args, **kwargs):
        raise PermissionError("[WinError 32] заблоковано іншою програмою")

    monkeypatch.setattr(Workbook, "save", _locked)

    with pytest.raises(PermissionError, match="Закрийте його та спробуйте ще раз"):
        save_workbook(Workbook(), str(tmp_path / "книга.xlsx"))
