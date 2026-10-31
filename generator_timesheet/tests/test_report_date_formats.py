"""Клітинки-дати у звітах проєкту - формат, який однаково читають Excel та
інші програми перегляду (constants.DATE_NUMBER_FORMAT). Реальний випадок,
підтверджений користувачем: файл, перекинутий на інший пристрій, показував
30.09.2026 як "273.09.2026" - там "DD" означає день року, а "mm" - хвилини."""

from datetime import date, datetime

import pytest
from openpyxl import load_workbook

import content.addon_mismatch_checker as amc
import content.payment_mismatch_checker as pmc
import content.rop_vop_statement as rvs

_PERSON = "ПЕРШИЙ Перший Перший"


def _payment_report(path):
    pmc.write_mismatch_report([{
        "file_name": "файл1.xlsx", "roster_pib_raw": _PERSON, "file_pib_raw": _PERSON,
        "roster_pidrozdil": "Підрозділ1", "roster_posada": "Посада1", "roster_zvannya": "сержант",
        "file_pidrozdil": "Підрозділ1", "file_posada": "Посада1", "file_zvannya": "сержант",
        "dates": {date(2026, 9, 30): ("РВЗ", 100)},
    }], path)
    return "Розбіжності"


def _addon_report(path):
    amc.write_addon_mismatch_report([{
        "source_file": "book61", "sources": ["book61: ППД"], "roster_pidrozdil": "Підрозділ1",
        "roster_pib_raw": _PERSON, "addon_pib_raw": _PERSON,
        "dates": {date(2026, 9, 30): ("РВЗ", "ППД")},
    }], [], path)
    return "book61"


def _schedule_report(path):
    rvs.write_schedule_mismatch_report([{
        "roster_posada": "Посада1", "schedule_posada": "Посада1",
        "roster_zvannya": "сержант", "schedule_zvannya": "сержант",
        "roster_pib_raw": _PERSON, "schedule_pib_raw": _PERSON, "pidrozdil": None,
        "dates": {date(2026, 9, 30): ("роп", "воп")},
        "full_days": {date(2026, 9, 30): ("роп", "воп")},
    }], path)
    return "Розбіжності"


@pytest.mark.parametrize("write_report", [_payment_report, _addon_report, _schedule_report])
def test_report_date_headers_use_only_day_month_year_tokens_in_portable_case(tmp_path, write_report):
    path = str(tmp_path / "report.xlsx")
    sheet_name = write_report(path)

    ws = load_workbook(path)[sheet_name]
    formats = [cell.number_format for cell in ws[1] if isinstance(cell.value, datetime)]
    assert formats
    for number_format in formats:
        assert {ch for ch in number_format if ch.isalpha()} == {"d", "M", "y"}, number_format
