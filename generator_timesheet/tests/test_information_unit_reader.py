from datetime import date, datetime

import xlwt
from openpyxl import Workbook

import content.information_unit_reader as iur
from content.oblik_timesheet import normalize_name

_XLS_DATE_STYLE = xlwt.easyxf(num_format_str="DD.MM.YYYY")


def _write_info_unit_xls_file(path, rows, dates):
    """Синтетичний .xls (старий бінарний формат, xlwt) аркуш ТАБЕЛЬ - дати
    ОДРАЗУ в рядку 1 (той самий структурний варіант, що й у реальному .xls
    файлі цієї теки), дані з рядка 2."""
    wb = xlwt.Workbook(encoding="utf-8")
    ws = wb.add_sheet("ТАБЕЛЬ")
    for col_idx, text in enumerate(["Посада", "Звання", "ПІБ"]):
        ws.write(0, col_idx, text)
    for idx, value in enumerate(dates, start=3):
        ws.write(0, idx, value, _XLS_DATE_STYLE)
    for row_idx, row_values in enumerate(rows, start=1):
        for col_idx, value in enumerate(row_values):
            ws.write(row_idx, col_idx, value)
    wb.save(str(path))
    return str(path)


def _write_bchs_xls_file(path, title_row1, header_row2, rows):
    """Синтетичний .xls аркуш БЧС - та сама структура, що й _write_bchs_file
    (openpyxl), лише записана через xlwt (старий бінарний формат)."""
    wb = xlwt.Workbook(encoding="utf-8")
    ws = wb.add_sheet("БЧС")
    ws.write(0, 0, title_row1)
    for col_idx, text in enumerate(header_row2):
        ws.write(1, col_idx, text)
    for row_idx, row_values in enumerate(rows, start=2):
        for col_idx, value in enumerate(row_values):
            ws.write(row_idx, col_idx, value)
    wb.save(str(path))
    return str(path)


def _write_info_unit_file(path, sheet_name, rows, header_row2):
    """rows - [[Посада, Звання, ПІБ, значення...], ...]; header_row2 - значення
    дат (рядок 2, як у справжніх файлах - рядок 1 має лише групуючий "ДАТА")."""
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name
    for col_idx, text in enumerate(["Посада", "Звання", "ПІБ", "ДАТА"], start=1):
        ws.cell(row=1, column=col_idx, value=text)
    for idx, value in enumerate(header_row2, start=4):
        ws.cell(row=2, column=idx, value=value)
    for row_idx, row_values in enumerate(rows, start=3):
        for col_idx, value in enumerate(row_values, start=1):
            ws.cell(row=row_idx, column=col_idx, value=value)
    wb.save(str(path))
    return str(path)


def _write_bchs_file(path, title_row1, header_row2, rows, sheet_name="БЧС"):
    """Реальна структура аркуша БЧС: рядок 1 - назва підрозділу (можливо, з
    датою), рядок 2 - заголовки колонок (позиція НЕ фіксована), дані - з
    рядка 3."""
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name
    ws.cell(row=1, column=1, value=title_row1)
    for col_idx, text in enumerate(header_row2, start=1):
        ws.cell(row=2, column=col_idx, value=text)
    for row_idx, row_values in enumerate(rows, start=3):
        for col_idx, value in enumerate(row_values, start=1):
            ws.cell(row=row_idx, column=col_idx, value=value)
    wb.save(str(path))
    return str(path)


_BCHS_HEADER = ["Посада", "Звання", "ПРІЗВИЩЕ ім'я по батькові", "БЧС", "Виплата додактової винагороди"]


def _write_info_unit_file_dates_in_row1(path, sheet_name, rows, dates):
    """Реальний структурний варіант: дати ОДРАЗУ в рядку 1 (одразу після ПІБ,
    без окремого групового заголовка "ДАТА" в рядку 1 + самих дат у рядку 2)
    - дані тоді з рядка 2, а не з рядка 3, як у _write_info_unit_file."""
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name
    for col_idx, text in enumerate(["Посада", "Звання", "ПІБ"], start=1):
        ws.cell(row=1, column=col_idx, value=text)
    for idx, value in enumerate(dates, start=4):
        ws.cell(row=1, column=idx, value=value)
    for row_idx, row_values in enumerate(rows, start=2):
        for col_idx, value in enumerate(row_values, start=1):
            ws.cell(row=row_idx, column=col_idx, value=value)
    wb.save(str(path))
    return str(path)


def test_parse_cell_date_accepts_a_plain_date_object():
    assert iur._parse_cell_date(date(2026, 8, 1)) == date(2026, 8, 1)


def test_parse_cell_date_returns_none_for_unparseable_or_wrong_type():
    assert iur._parse_cell_date("не дата") is None
    assert iur._parse_cell_date(None) is None
    assert iur._parse_cell_date(100) is None


def test_read_payment_values_happy_path(tmp_path):
    _write_info_unit_file(
        tmp_path / "1 рота.xlsx", "ТАБЕЛЬ",
        rows=[["Посада1", "сержант", "ПЕРШИЙ Перший Перший", 100, 30]],
        header_row2=[datetime(2026, 8, 1), datetime(2026, 8, 2)],
    )
    values, skipped = iur.read_payment_values(str(tmp_path))

    assert skipped == []
    assert values[(normalize_name("ПЕРШИЙ Перший Перший"), date(2026, 8, 1))] == [100]
    assert values[(normalize_name("ПЕРШИЙ Перший Перший"), date(2026, 8, 2))] == [30]


def test_clean_cell_value_strips_whitespace_from_text():
    """Реальний випадок: значення з зайвим пробілом (напр. "БПШП " - вручну
    введений текст) мовчки "губило" збіг з PAYMENT_STATUS_CONVERSIONS/
    PAYMENT_STATUS_OVERRIDES (content/oblik_timesheet.py) - без обрізання
    пробілів тут ця розбіжність ніколи б не виявилась."""
    assert iur._clean_cell_value(" БПШП ") == "БПШП"


def test_clean_cell_value_coerces_a_purely_numeric_string_to_int():
    """Реальний ризик: категорія виплати "100", збережена як ТЕКСТ (напр.
    клітинка відформатована як "Текст" чи скопійована з іншої системи) -
    без цього перетворення apply_payment_values (_is_number) сприймав би її
    як "нечислове значення", хоча вона однозначно являє собою число."""
    assert iur._clean_cell_value("100") == 100
    assert iur._clean_cell_value(" 30 ") == 30
    assert iur._clean_cell_value("-5") == -5


def test_clean_cell_value_leaves_non_numeric_text_and_other_types_unchanged():
    assert iur._clean_cell_value("БПШП") == "БПШП"
    assert iur._clean_cell_value(100) == 100
    assert iur._clean_cell_value(None) is None
    assert iur._clean_cell_value(date(2026, 8, 1)) == date(2026, 8, 1)


def test_read_payment_values_strips_whitespace_from_tabel_cell(tmp_path):
    """Реальний випадок: клітинка ТАБЕЛЬ із зайвим пробілом навколо статусу
    - _read_tabel_sheet (на відміну від _read_bchs_sheet, яка ЗАВЖДИ це
    робила) РАНІШЕ не обрізала пробіли взагалі."""
    _write_info_unit_file(
        tmp_path / "1 рота.xlsx", "ТАБЕЛЬ",
        rows=[["Посада1", "сержант", "ПЕРШИЙ Перший Перший", " БПШП "]],
        header_row2=[datetime(2026, 8, 1)],
    )
    values, skipped = iur.read_payment_values(str(tmp_path))

    assert skipped == []
    assert values[(normalize_name("ПЕРШИЙ Перший Перший"), date(2026, 8, 1))] == ["БПШП"]


def test_read_payment_values_coerces_numeric_string_in_tabel_cell(tmp_path):
    """Реальний ризик: категорія виплати збережена як ТЕКСТ "100", а не
    число 100, у клітинці ТАБЕЛЬ."""
    _write_info_unit_file(
        tmp_path / "1 рота.xlsx", "ТАБЕЛЬ",
        rows=[["Посада1", "сержант", "ПЕРШИЙ Перший Перший", "100"]],
        header_row2=[datetime(2026, 8, 1)],
    )
    values, skipped = iur.read_payment_values(str(tmp_path))

    assert skipped == []
    assert values[(normalize_name("ПЕРШИЙ Перший Перший"), date(2026, 8, 1))] == [100]


def test_read_payment_values_reads_tabel_sheet_with_dates_directly_in_row1(tmp_path):
    """Реальний випадок: аркуш ТАБЕЛЬ, де дати йдуть ОДРАЗУ в рядку 1 (одразу
    після ПІБ, без окремого групового заголовка "ДАТА" + рядка дат) - дані
    тоді з рядка 2, а не з рядка 3. Раніше такий файл мовчки не читався
    взагалі (жодної колонки-дати не знаходилось, бо код завжди дивився лише
    в рядок 2)."""
    _write_info_unit_file_dates_in_row1(
        tmp_path / "підрозділ 1.xlsx", "ТАБЕЛЬ",
        rows=[["Командир роти", "старший лейтенант", "ПЕРШИЙ Перший Перший", 100, 100, 100]],
        dates=[datetime(2026, 8, 1), datetime(2026, 8, 2), datetime(2026, 8, 3)],
    )
    values, skipped = iur.read_payment_values(str(tmp_path))

    assert skipped == []
    assert values[(normalize_name("ПЕРШИЙ Перший Перший"), date(2026, 8, 1))] == [100]
    assert values[(normalize_name("ПЕРШИЙ Перший Перший"), date(2026, 8, 2))] == [100]
    assert values[(normalize_name("ПЕРШИЙ Перший Перший"), date(2026, 8, 3))] == [100]


def test_parse_cell_date_accepts_the_missing_dot_typo():
    """Реальний випадок, підтверджений користувачем: клітинка-дата аркуша
    ТАБЕЛЬ була записана як текст "06.092026" - пропущена крапка МІЖ
    місяцем і роком (мало бути "06.09.2026") - без цього ВСЯ колонка цього
    дня мовчки не вважалась колонкою-датою взагалі, хоча значення виплат у
    ній були."""
    assert iur._parse_cell_date("06.092026") == date(2026, 9, 6)


def test_parse_cell_date_rejects_an_invalid_date_matching_the_missing_dot_shape():
    """Той самий "день.місяцьРІК" вигляд, що й тест вище, але з неможливим
    місяцем (13) - НЕ має падати з винятком, лише повернути None, як і
    будь-яке інше нерозпізнане значення."""
    assert iur._parse_cell_date("06.132026") is None


def test_read_payment_values_reads_a_column_with_the_missing_dot_typo(tmp_path):
    """Той самий реальний випадок наскрізно - read_payment_values МАЄ
    прочитати значення виплати навіть під датою-колонкою з цією
    друкарською помилкою, а не пропускати всю колонку."""
    _write_info_unit_file_dates_in_row1(
        tmp_path / "підрозділ 1.xlsx", "ТАБЕЛЬ",
        rows=[["Командир роти", "старший лейтенант", "ПЕРШИЙ Перший Перший", 100, 100]],
        dates=[datetime(2026, 9, 5), "06.092026"],
    )
    values, skipped = iur.read_payment_values(str(tmp_path))

    assert skipped == []
    assert values[(normalize_name("ПЕРШИЙ Перший Перший"), date(2026, 9, 6))] == [100]


def test_read_payment_values_parses_string_dates_both_year_formats(tmp_path):
    _write_info_unit_file(
        tmp_path / "рота.xlsx", "ТАБЕЛЬ",
        rows=[["Посада1", "сержант", "ПЕРШИЙ Перший Перший", 100, 30]],
        header_row2=["01.08.26", "02.08.2026"],
    )
    values, skipped = iur.read_payment_values(str(tmp_path))

    assert skipped == []
    assert values[(normalize_name("ПЕРШИЙ Перший Перший"), date(2026, 8, 1))] == [100]
    assert values[(normalize_name("ПЕРШИЙ Перший Перший"), date(2026, 8, 2))] == [30]


def test_read_payment_values_is_case_insensitive_about_sheet_name(tmp_path):
    _write_info_unit_file(
        tmp_path / "рота.xlsx", "Табель",
        rows=[["Посада1", "сержант", "ПЕРШИЙ Перший Перший", 100]],
        header_row2=[datetime(2026, 8, 1)],
    )
    values, skipped = iur.read_payment_values(str(tmp_path))

    assert skipped == []
    assert values[(normalize_name("ПЕРШИЙ Перший Перший"), date(2026, 8, 1))] == [100]


def test_read_payment_values_combines_multiple_files_for_the_same_person_and_date(tmp_path):
    _write_info_unit_file(
        tmp_path / "рота1.xlsx", "ТАБЕЛЬ",
        rows=[["Посада1", "сержант", "ПЕРШИЙ Перший Перший", 100]],
        header_row2=[datetime(2026, 8, 1)],
    )
    _write_info_unit_file(
        tmp_path / "рота2.xlsx", "ТАБЕЛЬ",
        rows=[["Посада1", "сержант", "ПЕРШИЙ Перший Перший", 30]],
        header_row2=[datetime(2026, 8, 1)],
    )
    values, _skipped = iur.read_payment_values(str(tmp_path))

    assert sorted(values[(normalize_name("ПЕРШИЙ Перший Перший"), date(2026, 8, 1))]) == [30, 100]


def test_read_payment_values_skips_empty_pib_and_blank_cells(tmp_path):
    _write_info_unit_file(
        tmp_path / "рота.xlsx", "ТАБЕЛЬ",
        rows=[
            ["Посада1", "сержант", "", 100, 100],
            ["Посада2", "матрос", "ДРУГИЙ Другий Другий", None, ""],
        ],
        header_row2=[datetime(2026, 8, 1), datetime(2026, 8, 2)],
    )
    values, skipped = iur.read_payment_values(str(tmp_path))

    assert skipped == []
    assert values == {}


def test_read_payment_values_skips_unsupported_extension(tmp_path):
    (tmp_path / "стара.doc").write_text("не таблиця взагалі")

    values, skipped = iur.read_payment_values(str(tmp_path))

    assert values == {}
    assert len(skipped) == 1
    assert "не підтримується" in skipped[0]["reason"]


def test_safe_load_workbook_returns_the_workbook_on_success(tmp_path):
    path = tmp_path / "рота.xlsx"
    Workbook().save(str(path))

    wb, reason = iur._safe_load_workbook(str(path))

    assert wb is not None
    assert reason is None


def test_safe_load_workbook_reports_a_reason_for_a_corrupted_zip(tmp_path):
    """Реальний випадок: файл цієї теки має розширення .xlsx, але вміст
    пошкоджений настільки, що це навіть не дійсний zip-архів (openpyxl сам
    - zip-контейнер) - BadZipFile, а не звичайне читання."""
    path = tmp_path / "пошкоджена.xlsx"
    path.write_bytes(b"not a real xlsx file at all, just garbage bytes")

    wb, reason = iur._safe_load_workbook(str(path))

    assert wb is None
    assert "пошкоджена.xlsx" in reason


def test_module_import_tolerates_embedded_drawings_with_empty_panose():
    """Реальний випадок (файл цієї теки за 04.09.2026): вбудований малюнок/
    штамп, чий текстовий стиль записує ПОРОЖНЄ значення атрибута "panose" -
    БЕЗ патча (сам ІМПОРТ content.information_unit_reader) openpyxl.
    drawing.text.Font(panose="") підняв би ValueError "Value does not match
    pattern" (Font.panose - HexBinary, вимагає ХОЧА Б ОДИН hex-символ) - САМЕ
    ЦЕ й зупиняло читання ВСЬОГО файлу (openpyxl.reader.excel.
    read_worksheets -> find_images -> SpreadsheetDrawing.from_tree ->
    Font.__init__), хоча самі малюнки нас не цікавлять узагалі - лише дані
    аркушів ТАБЕЛЬ/БЧС."""
    from openpyxl.drawing.text import Font

    Font(typeface="Arial", panose="")


def test_safe_load_workbook_reports_a_reason_for_invalid_internal_xml(tmp_path, monkeypatch):
    """Реальний випадок (баг, підтверджений користувачем): файл цієї теки МАВ
    дійсний zip-контейнер, але один із внутрішніх XML-фрагментів (вбудований
    малюнок/шрифт зі значенням "panose", що не відповідає власній перевірці
    формату openpyxl) не проходить розбір - openpyxl сам явно згортає ЦЕ в
    ValueError ("Unable to read workbook: ... invalid XML") - раніше ОДНА
    така книга зупиняла ВЕСЬ прогін (непіймана помилка з traceback)."""
    path = tmp_path / "панозе.xlsx"
    path.write_text("непошкоджений сам по собі шлях - помилку піднімає підмінений _load_workbook")

    def _raise_value_error(_path):
        raise ValueError("Unable to read workbook: could not read worksheets. This is most probably because the workbook source files contain some invalid XML.")

    monkeypatch.setattr(iur, "_load_workbook", _raise_value_error)

    wb, reason = iur._safe_load_workbook(str(path))

    assert wb is None
    assert "панозе.xlsx" in reason


def test_read_payment_values_skips_a_corrupted_file_and_still_reads_the_rest(tmp_path):
    """Реальний випадок (баг, підтверджений користувачем): ОДИН пошкоджений
    файл цієї теки НЕ повинен зупиняти прогін для решти, справних файлів -
    записується в skipped із причиною, і читання ПРОДОВЖУЄТЬСЯ."""
    (tmp_path / "пошкоджена.xlsx").write_bytes(b"garbage, not a real workbook")
    _write_info_unit_file(
        tmp_path / "рота.xlsx", "Табель",
        [["Посада1", "сержант", "ПЕРШИЙ Перший Перший", 100]],
        header_row2=[date(2026, 8, 1)],
    )

    values, skipped = iur.read_payment_values(str(tmp_path))

    assert values[(normalize_name("ПЕРШИЙ Перший Перший"), date(2026, 8, 1))] == [100]
    assert len(skipped) == 1
    assert "пошкоджена.xlsx" in skipped[0]["reason"]


def test_read_payment_values_skips_office_lock_file(tmp_path):
    """Реальний випадок: файл цієї теки ВІДКРИТИЙ в Excel - lock-файл
    "~$Назва.xlsx" (ТЕ САМЕ розширення .xlsx) не повинен спричиняти
    PermissionError/крах читання - тихо пропускається, БЕЗ запису в skipped
    (це не "непідтримуваний формат", а взагалі не справжня книга)."""
    (tmp_path / "~$рота.xlsx").write_text("lock-файл, не справжня книга")

    values, skipped = iur.read_payment_values(str(tmp_path))

    assert values == {}
    assert skipped == []


def test_read_payment_values_reads_xls_tabel_sheet(tmp_path):
    """Реальний випадок: файл цієї теки був у старому бінарному форматі
    (.xls, не .xlsx) - раніше мовчки потрапляв у skipped ("формат файлу не
    підтримується"), і дані підрозділу взагалі не враховувались - тепер
    читається так само, як .xlsx (через xlrd, content/information_unit_reader
    _load_workbook/_XlsWorkbookAdapter)."""
    _write_info_unit_xls_file(
        tmp_path / "1 рота.xls",
        rows=[["Посада1", "сержант", "ПЕРШИЙ Перший Перший", 100, 30]],
        dates=[datetime(2026, 8, 1), datetime(2026, 8, 2)],
    )
    values, skipped = iur.read_payment_values(str(tmp_path))

    assert skipped == []
    assert values[(normalize_name("ПЕРШИЙ Перший Перший"), date(2026, 8, 1))] == [100]
    assert values[(normalize_name("ПЕРШИЙ Перший Перший"), date(2026, 8, 2))] == [30]


def test_read_payment_values_reads_xls_bchs_sheet_when_tabel_unavailable(tmp_path):
    """.xls файл БЕЗ аркуша ТАБЕЛЬ (лише БЧС) - той самий fallback, що й для
    .xlsx (read_payment_values_falls_back_to_bchs_when_tabel_unreadable)."""
    _write_bchs_xls_file(
        tmp_path / "рота.xls",
        title_row1="ВЗВОД ЗВ'ЯЗКУ 08.08.2026",
        header_row2=_BCHS_HEADER,
        rows=[["Посада1", "матрос", "ПЕРШИЙ Перший Перший", "РВЗ", 100]],
    )
    values, skipped = iur.read_payment_values(str(tmp_path))

    assert skipped == []
    assert values[(normalize_name("ПЕРШИЙ Перший Перший"), date(2026, 8, 8))] == [100]


def test_load_workbook_dispatches_to_openpyxl_for_xlsx(tmp_path):
    path = _write_info_unit_file(tmp_path / "1 рота.xlsx", "ТАБЕЛЬ", rows=[], header_row2=[])
    wb = iur._load_workbook(path)
    assert "ТАБЕЛЬ" in wb.sheetnames


def test_load_workbook_dispatches_to_xlrd_adapter_for_xls(tmp_path):
    path = _write_info_unit_xls_file(tmp_path / "1 рота.xls", rows=[], dates=[])
    wb = iur._load_workbook(path)
    assert isinstance(wb, iur._XlsWorkbookAdapter)
    assert "ТАБЕЛЬ" in wb.sheetnames


def test_is_office_lock_file_matches_tilde_dollar_prefix():
    assert iur._is_office_lock_file("resources/information_unit/~$рота.xlsx") is True
    assert iur._is_office_lock_file("resources/information_unit/рота.xlsx") is False


def test_xls_adapter_converts_boolean_cell_to_python_bool(tmp_path):
    wb = xlwt.Workbook()
    wb.add_sheet("Аркуш1").write(0, 0, True)
    path = tmp_path / "bool.xls"
    wb.save(str(path))

    adapter = iur._load_workbook(str(path))
    ws = adapter["Аркуш1"]

    assert ws.cell(row=1, column=1).value is True


def test_xls_adapter_returns_none_for_a_gap_cell_within_the_used_range(tmp_path):
    """Клітинка "в дірці" (порожня, але МІЖ двома заповненими в тому самому
    рядку - xlrd дає їй ctype EMPTY, значення '') - має читатись як None
    (той самий вигляд, що й в openpyxl), а не порожній рядок."""
    wb = xlwt.Workbook()
    ws_raw = wb.add_sheet("Аркуш1")
    ws_raw.write(0, 0, "a")
    ws_raw.write(0, 2, "c")
    path = tmp_path / "gap.xls"
    wb.save(str(path))

    adapter = iur._load_workbook(str(path))
    ws = adapter["Аркуш1"]

    assert ws.cell(row=1, column=2).value is None


def test_xls_adapter_returns_none_for_out_of_bounds_cell(tmp_path):
    """Той самий вигляд, що й openpyxl - ws.cell() за межами фактичного
    діапазону аркуша не кидає помилку, а дає порожню клітинку (реальний
    випадок: _bchs_columns читає ФІКСОВАНИЙ рядок _BCHS_HEADER_ROW, який
    може виявитись за межами короткого аркуша)."""
    wb = xlwt.Workbook()
    wb.add_sheet("Аркуш1").write(0, 0, "єдиний рядок")
    path = tmp_path / "short.xls"
    wb.save(str(path))

    adapter = iur._load_workbook(str(path))
    ws = adapter["Аркуш1"]

    assert ws.cell(row=5, column=1).value is None
    assert ws.cell(row=1, column=5).value is None


def test_read_payment_values_skips_file_without_tabel_sheet(tmp_path):
    wb = Workbook()
    wb.active.title = "БЧС"
    wb.save(str(tmp_path / "рота.xlsx"))

    values, skipped = iur.read_payment_values(str(tmp_path))

    assert values == {}
    assert len(skipped) == 1
    assert "не знайдено аркуш ТАБЕЛЬ" in skipped[0]["reason"]


def test_read_payment_values_skips_file_without_pib_column(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "ТАБЕЛЬ"
    ws.cell(row=1, column=1, value="Щось")
    ws.cell(row=2, column=2, value=datetime(2026, 8, 1))
    wb.save(str(tmp_path / "рота.xlsx"))

    values, skipped = iur.read_payment_values(str(tmp_path))

    assert values == {}
    assert len(skipped) == 1
    assert "колонку ПІБ" in skipped[0]["reason"]


def test_read_payment_values_skips_file_without_date_columns(tmp_path):
    _write_info_unit_file(
        tmp_path / "рота.xlsx", "ТАБЕЛЬ",
        rows=[["Посада1", "сержант", "ПЕРШИЙ Перший Перший", "не дата"]],
        header_row2=["не дата"],
    )
    values, skipped = iur.read_payment_values(str(tmp_path))

    assert values == {}
    assert len(skipped) == 1
    assert "жодної колонки-дати" in skipped[0]["reason"]


def test_read_payment_values_reads_bchs_sheet_with_date_from_filename(tmp_path):
    """Реальний випадок: аркуш ТАБЕЛЬ у деяких файлах resources/information_unit
    містить лише ЗАСТАРІЛІ дати (напр. увесь липень), тож аркуш БЧС (знімок
    ОДНОГО конкретного дня) лишається ЄДИНИМ джерелом актуальних даних."""
    _write_bchs_file(
        tmp_path / "БЧС ВЗ 08.08.2026.xlsx",
        title_row1="ВЗВОД ЗВ'ЯЗКУ 08.08.2026",
        header_row2=_BCHS_HEADER,
        rows=[["Посада1", "матрос", "ПЕРШИЙ Перший Перший", "РВЗ", 100]],
    )
    values, skipped = iur.read_payment_values(str(tmp_path))

    assert skipped == []
    assert values[(normalize_name("ПЕРШИЙ Перший Перший"), date(2026, 8, 8))] == [100]


def test_read_payment_values_reads_bchs_sheet_with_two_digit_year_in_filename(tmp_path):
    """Реальні файли часто мають дворічний рік у назві (напр. "08.08.26")."""
    _write_bchs_file(
        tmp_path / "МП 08.08.26.xlsx",
        title_row1="МІНОМЕТНИЙ ПІДРОЗДІЛ",
        header_row2=_BCHS_HEADER,
        rows=[["Посада1", "матрос", "ПЕРШИЙ Перший Перший", "РВЗ", 100]],
    )
    values, skipped = iur.read_payment_values(str(tmp_path))

    assert skipped == []
    assert values[(normalize_name("ПЕРШИЙ Перший Перший"), date(2026, 8, 8))] == [100]


def test_read_payment_values_reads_bchs_sheet_with_date_only_in_title(tmp_path):
    """Ім'я файлу без валідної дати (немає року поряд з день.місяць) -
    запасний варіант - дата із заголовка (рядок 1) самого аркуша БЧС."""
    _write_bchs_file(
        tmp_path / "БЧС РВ.xlsx",
        title_row1="РОЗВІДУВАЛЬНИЙ ВЗВОД (08.08.2026)",
        header_row2=_BCHS_HEADER,
        rows=[["Посада1", "матрос", "ПЕРШИЙ Перший Перший", "РВЗ", 30]],
    )
    values, skipped = iur.read_payment_values(str(tmp_path))

    assert skipped == []
    assert values[(normalize_name("ПЕРШИЙ Перший Перший"), date(2026, 8, 8))] == [30]


def test_read_payment_values_detects_bchs_columns_regardless_of_position(tmp_path):
    """Позиція колонок ПРІЗВИЩЕ/БЧС/Виплата РІЗНИТЬСЯ між реальними файлами -
    визначаються за текстом заголовка, а не фіксованою позицією."""
    _write_bchs_file(
        tmp_path / "БЧС ІСВ 08.08.2026.xlsx",
        title_row1="ІНЖЕНЕРНО-САПЕРНИЙ ВЗВОД 08.08.2026",
        header_row2=["№", "ПРІЗВИЩЕ ім'я по батькові", None, "Виплата додактової винагороди", "Посада", "Звання", "БЧС"],
        rows=[[1, "ПЕРШИЙ Перший Перший", None, 100, "Посада1", "матрос", "РВЗ"]],
    )
    values, skipped = iur.read_payment_values(str(tmp_path))

    assert skipped == []
    assert values[(normalize_name("ПЕРШИЙ Перший Перший"), date(2026, 8, 8))] == [100]


def test_read_payment_values_records_non_default_bchs_status_as_text(tmp_path):
    """Статус БЧС, ІНШИЙ за DEFAULT_STATUS ("РВЗ") - записується як текст (той
    самий підхід, що й нечислове значення в ТАБЕЛЬ): apply_payment_values
    трактує це як "підрозділ каже щось ІНШЕ", а не мовчки підставляє число."""
    _write_bchs_file(
        tmp_path / "БЧС ВЗ 08.08.2026.xlsx",
        title_row1="ВЗВОД ЗВ'ЯЗКУ 08.08.2026",
        header_row2=_BCHS_HEADER,
        rows=[["Посада1", "матрос", "ДРУГИЙ Другий Другий", "ВД", 100]],
    )
    values, skipped = iur.read_payment_values(str(tmp_path))

    assert skipped == []
    assert values[(normalize_name("ДРУГИЙ Другий Другий"), date(2026, 8, 8))] == ["ВД"]


def test_read_payment_values_prefers_tabel_and_ignores_bchs_in_the_same_file(tmp_path):
    """ТАБЕЛЬ - ПРІОРИТЕТНЕ джерело (за прямою вказівкою користувача): якщо
    файл має ОБИДВА аркуші, БЧС того самого файлу взагалі НЕ читається -
    навіть якщо в ньому є дані на дату, якої немає в ТАБЕЛЬ (два аркуші
    ОДНОГО файлу можуть суперечити один одному, напр. через друкарську
    помилку в ПІБ на одному з них - читання лише ОДНОГО джерела уникає
    подвоєння розбіжностей)."""
    wb = Workbook()
    ws_tabel = wb.active
    ws_tabel.title = "ТАБЕЛЬ"
    for col_idx, text in enumerate(["Посада", "Звання", "ПІБ", "ДАТА"], start=1):
        ws_tabel.cell(row=1, column=col_idx, value=text)
    ws_tabel.cell(row=2, column=4, value=datetime(2026, 7, 1))
    ws_tabel.cell(row=3, column=1, value="Посада1")
    ws_tabel.cell(row=3, column=2, value="сержант")
    ws_tabel.cell(row=3, column=3, value="ПЕРШИЙ Перший Перший")
    ws_tabel.cell(row=3, column=4, value=100)

    ws_bchs = wb.create_sheet("БЧС")
    ws_bchs.cell(row=1, column=1, value="ВЗВОД ЗВ'ЯЗКУ 08.08.2026")
    for col_idx, text in enumerate(_BCHS_HEADER, start=1):
        ws_bchs.cell(row=2, column=col_idx, value=text)
    ws_bchs.cell(row=3, column=1, value="Посада1")
    ws_bchs.cell(row=3, column=2, value="сержант")
    ws_bchs.cell(row=3, column=3, value="ПЕРШИЙ Перший Перший")
    ws_bchs.cell(row=3, column=4, value="РВЗ")
    ws_bchs.cell(row=3, column=5, value=30)

    wb.save(str(tmp_path / "змішаний.xlsx"))

    values, skipped = iur.read_payment_values(str(tmp_path))

    assert skipped == []
    assert values[(normalize_name("ПЕРШИЙ Перший Перший"), date(2026, 7, 1))] == [100]
    assert (normalize_name("ПЕРШИЙ Перший Перший"), date(2026, 8, 8)) not in values


def test_read_payment_values_falls_back_to_bchs_when_tabel_unreadable(tmp_path):
    """ТАБЕЛЬ - пріоритетне, але НЕ єдине джерело: якщо аркуш ТАБЕЛЬ у
    файлі взагалі не вдалось прочитати (тут - його просто немає), БЧС того
    самого файлу все одно читається як запасний варіант."""
    wb = Workbook()
    ws_bchs = wb.active
    ws_bchs.title = "БЧС"
    ws_bchs.cell(row=1, column=1, value="ВЗВОД ЗВ'ЯЗКУ 08.08.2026")
    for col_idx, text in enumerate(_BCHS_HEADER, start=1):
        ws_bchs.cell(row=2, column=col_idx, value=text)
    ws_bchs.cell(row=3, column=1, value="Посада1")
    ws_bchs.cell(row=3, column=2, value="сержант")
    ws_bchs.cell(row=3, column=3, value="ПЕРШИЙ Перший Перший")
    ws_bchs.cell(row=3, column=4, value="РВЗ")
    ws_bchs.cell(row=3, column=5, value=30)

    wb.save(str(tmp_path / "лише_бчс.xlsx"))

    values, skipped = iur.read_payment_values(str(tmp_path))

    assert skipped == []
    assert values[(normalize_name("ПЕРШИЙ Перший Перший"), date(2026, 8, 8))] == [30]


def test_read_payment_values_skips_file_with_neither_source_readable(tmp_path):
    """Файл лише з аркушем БЧС, у якого немає жодної дати (ні в назві файлу,
    ні в заголовку) - НЕМАЄ аркуша ТАБЕЛЬ узагалі, тож ОБИДВА джерела
    падають - файл лишається пропущеним, а не мовчки без жодного результату."""
    _write_bchs_file(
        tmp_path / "БЧС без дати.xlsx",
        title_row1="ЯКИЙСЬ ВЗВОД",
        header_row2=_BCHS_HEADER,
        rows=[["Посада1", "матрос", "ПЕРШИЙ Перший Перший", "РВЗ", 100]],
    )
    values, skipped = iur.read_payment_values(str(tmp_path))

    assert values == {}
    assert len(skipped) == 1


def test_read_bchs_sheet_reports_when_no_date_found(tmp_path):
    path = _write_bchs_file(
        tmp_path / "БЧС без дати.xlsx",
        title_row1="ЯКИЙСЬ ВЗВОД",
        header_row2=_BCHS_HEADER,
        rows=[["Посада1", "матрос", "ПЕРШИЙ Перший Перший", "РВЗ", 100]],
    )
    from openpyxl import load_workbook
    wb = load_workbook(path, data_only=True)

    reason = iur._read_bchs_sheet(wb, path, {})

    assert reason is not None
    assert "не вдалось визначити дату" in reason


def test_read_bchs_sheet_reports_when_no_pib_column_found(tmp_path):
    path = _write_bchs_file(
        tmp_path / "БЧС ВЗ 08.08.2026.xlsx",
        title_row1="ВЗВОД ЗВ'ЯЗКУ 08.08.2026",
        header_row2=["Посада", "Звання", "БЧС", "Виплата додактової винагороди"],
        rows=[["Посада1", "матрос", "РВЗ", 100]],
    )
    from openpyxl import load_workbook
    wb = load_workbook(path, data_only=True)

    reason = iur._read_bchs_sheet(wb, path, {})

    assert reason is not None
    assert "колонку ПІБ" in reason


def test_read_bchs_sheet_reports_when_sheet_missing(tmp_path):
    _write_info_unit_file(
        tmp_path / "лише_табель.xlsx", "ТАБЕЛЬ",
        rows=[["Посада1", "сержант", "ПЕРШИЙ Перший Перший", 100]],
        header_row2=[datetime(2026, 8, 1)],
    )
    from openpyxl import load_workbook
    wb = load_workbook(str(tmp_path / "лише_табель.xlsx"), data_only=True)

    reason = iur._read_bchs_sheet(wb, "лише_табель.xlsx", {})

    assert reason is not None
    assert "не знайдено аркуш БЧС" in reason


def test_search_date_returns_none_for_an_invalid_date():
    assert iur._search_date("32.13.2026") is None
    assert iur._search_date("нічого схожого на дату") is None


def test_read_bchs_sheet_skips_rows_with_empty_pib(tmp_path):
    """Реальний випадок: аркуш БЧС може мати рядок-підзаголовок (напр. для
    об'єднаної групи колонок "МІСЦЕЗНАХОДЖЕННЯ") одразу під заголовками,
    порожній у колонці ПІБ - такий рядок пропускається, а не падає."""
    _write_bchs_file(
        tmp_path / "БЧС ВЗ 08.08.2026.xlsx",
        title_row1="ВЗВОД ЗВ'ЯЗКУ 08.08.2026",
        header_row2=_BCHS_HEADER,
        rows=[
            [None, None, None, None, None],
            ["Посада1", "матрос", "ПЕРШИЙ Перший Перший", "РВЗ", 100],
        ],
    )
    values, skipped = iur.read_payment_values(str(tmp_path))

    assert skipped == []
    assert values[(normalize_name("ПЕРШИЙ Перший Перший"), date(2026, 8, 8))] == [100]


def test_bchs_columns_handles_missing_status_or_payment_column(tmp_path):
    """status_col/payment_col можуть бути відсутні незалежно один від одного -
    не мають зривати читання решти колонок (просто ця колонка не заповнюється)."""
    path = _write_bchs_file(
        tmp_path / "БЧС без виплати.xlsx",
        title_row1="ВЗВОД ЗВ'ЯЗКУ 08.08.2026",
        header_row2=["Посада", "Звання", "ПРІЗВИЩЕ ім'я по батькові", "БЧС"],
        rows=[["Посада1", "матрос", "ПЕРШИЙ Перший Перший", "РВЗ"]],
    )
    values, skipped = iur.read_payment_values(str(tmp_path))

    assert skipped == []
    assert values == {}  # РВЗ, але немає колонки виплати - нічого не записано


def test_tabel_date_row_and_columns_prefers_row1_when_dates_are_there():
    wb = Workbook()
    ws = wb.active
    for col_idx, text in enumerate(["Посада", "Звання", "ПІБ"], start=1):
        ws.cell(row=1, column=col_idx, value=text)
    ws.cell(row=1, column=4, value=datetime(2026, 8, 1))

    date_row, date_cols = iur._tabel_date_row_and_columns(ws)

    assert date_row == 1
    assert date_cols == [(4, date(2026, 8, 1))]


def test_tabel_date_row_and_columns_falls_back_to_row2_old_format():
    wb = Workbook()
    ws = wb.active
    for col_idx, text in enumerate(["Посада", "Звання", "ПІБ", "ДАТА"], start=1):
        ws.cell(row=1, column=col_idx, value=text)
    ws.cell(row=2, column=4, value=datetime(2026, 8, 1))

    date_row, date_cols = iur._tabel_date_row_and_columns(ws)

    assert date_row == 2
    assert date_cols == [(4, date(2026, 8, 1))]


def test_tabel_date_row_and_columns_returns_none_when_no_dates_anywhere():
    wb = Workbook()
    ws = wb.active
    ws.cell(row=1, column=1, value="Щось")

    date_row, date_cols = iur._tabel_date_row_and_columns(ws)

    assert date_row is None
    assert date_cols == []
