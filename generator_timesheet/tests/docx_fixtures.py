from docx import Document

ARRIVAL_HEADER = ["№", "Звання", "Прізвище ім'я \nпо батькові", "Посада", "Дата прибуття", "Куди прибув"]
DEPARTURE_HEADER = ["№", "Звання", "Прізвище ім'я \nпо батькові", "Посада", "Дата вибуття", "Куди вибув"]


def _add_table(doc, header, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(header))
    for col_idx, text in enumerate(header):
        table.rows[0].cells[col_idx].text = text
    for row_idx, values in enumerate(rows, start=1):
        for col_idx, value in enumerate(values):
            table.rows[row_idx].cells[col_idx].text = value


def write_daily_report(path, groups=(), arrival_list=None, freeform_paragraphs=(), signature=None):
    """Мінімальний, але структурно вірний .docx щоденного рапорту - та сама
    послідовність блоків, що й у справжніх рапортах (resources/report/*.docx):
    "List Paragraph" абзац групи (ПРИБУЛИ/ВИБУЛИ), під ним - "Normal" абзац
    підзаголовка й одразу таблиця людей.

    groups - [(group_heading, [(subheading, header, [row_values, ...]), ...]), ...].
    arrival_list - (heading_text, [name_line, ...]) чи None - заголовок-СПИСОК
    "ПРИБУЛИ до пункту постійно(ї) дислокації ...:" (реальний випадок, БЕЗ
    таблиці): і заголовок, і КОЖЕН рядок імені - той самий стиль "List
    Paragraph" (на відміну від розділу "Поза межами..." нижче, де рядки
    людей - звичайна "Normal" проза).
    freeform_paragraphs - рядки розділу "Поза межами..." (додається в кінці,
    якщо непорожній - як реальний останній розділ рапорту).
    signature - (позиція, звання_і_піб) чи None (без підпису) - як реальний
    підпис наприкінці рапорту ("No Spacing" стиль для позиції, "Normal" -
    для звання+ПІБ)."""
    doc = Document()
    doc.add_paragraph("Командиру військової частини A0000")
    doc.add_paragraph("РАПОРТ")
    doc.add_paragraph("Дійсним доповідаю, що станом на 17 годину 00 хвилин ... відбулися зміни, а саме:")

    for group_heading, subsections in groups:
        doc.add_paragraph(group_heading, style="List Paragraph")
        for subheading, header, rows in subsections:
            doc.add_paragraph(subheading)
            _add_table(doc, header, rows)

    if arrival_list:
        heading_text, name_lines = arrival_list
        doc.add_paragraph(heading_text, style="List Paragraph")
        for name_line in name_lines:
            doc.add_paragraph(name_line, style="List Paragraph")

    if freeform_paragraphs:
        doc.add_paragraph("Поза межами складу сил та засобів 9 армійського корпусу:", style="List Paragraph")
        for text in freeform_paragraphs:
            doc.add_paragraph(text)

    if signature:
        position_line, rank_and_name_line = signature
        doc.add_paragraph(position_line, style="No Spacing")
        doc.add_paragraph(rank_and_name_line)

    doc.save(str(path))
    return str(path)
