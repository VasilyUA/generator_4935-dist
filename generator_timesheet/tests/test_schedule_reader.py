from openpyxl import Workbook

from content.oblik_timesheet import normalize_name
from content.schedule_reader import latest_schedule_file, read_schedule_marks

# Реальний файл має 8 зайвих колонок перед видимими даними (A..H) - ПІБ/дні
# починаються з колонки I (9) - тут узято те саме зміщення, щоб перевірити,
# що читання шукає колонки ЗА ТЕКСТОМ заголовка, а не фіксованою позицією.
_JUNK_COLUMNS = 8
_PIB_COL = _JUNK_COLUMNS + 1
_POSADA_COL = _PIB_COL - 2
_ZVANNYA_COL = _PIB_COL - 1
_HEADER_ROW = 6


def _write_schedule_workbook(path, days, rows, with_posada_zvannya=False):
    """days - список номерів днів-колонок (у порядку колонок); rows - список
    (pib, {день: значення}) або (pib, {день: значення}, posada, zvannya),
    якщо with_posada_zvannya=True - рядок з pib=None позначає підписний
    блок (порожній ПІБ), на якому read_schedule_marks має зупинитись."""
    wb = Workbook()
    ws = wb.active
    day_cols = {day: _PIB_COL + 1 + i for i, day in enumerate(days)}

    if with_posada_zvannya:
        ws.cell(row=_HEADER_ROW, column=_POSADA_COL, value="Посада")
        ws.cell(row=_HEADER_ROW, column=_ZVANNYA_COL, value="Військове звання")
    ws.cell(row=_HEADER_ROW, column=_PIB_COL, value="Прізвище,\nвласне ім'я")
    for day, col in day_cols.items():
        ws.cell(row=_HEADER_ROW, column=col, value=f"{day:02d}")

    for offset, row_values in enumerate(rows, start=1):
        pib, day_values = row_values[0], row_values[1]
        row = _HEADER_ROW + offset
        ws.cell(row=row, column=_PIB_COL, value=pib)
        for day, value in day_values.items():
            ws.cell(row=row, column=day_cols[day], value=value)
        if with_posada_zvannya and len(row_values) > 2:
            posada, zvannya = row_values[2], row_values[3]
            ws.cell(row=row, column=_POSADA_COL, value=posada)
            ws.cell(row=row, column=_ZVANNYA_COL, value=zvannya)

    wb.save(str(path))
    return str(path)


def test_latest_schedule_file_picks_the_latest_date_for_the_target_month(tmp_path):
    older = tmp_path / "1бмпВОП-РОП(СЕРПЕНЬ)15.08.2026.xlsx"
    newer = tmp_path / "1бмпВОП-РОП(СЕРПЕНЬ)29.08.2026.xlsx"
    older.touch()
    newer.touch()

    result = latest_schedule_file(str(tmp_path), 2026, 8)

    assert result == str(newer)


def test_latest_schedule_file_ignores_office_lock_files(tmp_path):
    real_file = tmp_path / "1бмпВОП-РОП(СЕРПЕНЬ)29.08.2026.xlsx"
    lock_file = tmp_path / "~$1бмпВОП-РОП(СЕРПЕНЬ)30.08.2026.xlsx"
    real_file.touch()
    lock_file.touch()

    result = latest_schedule_file(str(tmp_path), 2026, 8)

    assert result == str(real_file)


def test_latest_schedule_file_ignores_files_from_a_different_month_or_year(tmp_path):
    (tmp_path / "1бмпВОП-РОП(ЛИПЕНЬ)31.07.2026.xlsx").touch()
    (tmp_path / "1бмпВОП-РОП(СЕРПЕНЬ)01.08.2025.xlsx").touch()

    result = latest_schedule_file(str(tmp_path), 2026, 8)

    assert result is None


def test_latest_schedule_file_ignores_files_without_a_parseable_date(tmp_path):
    (tmp_path / "1бмпВОП-РОП(СЕРПЕНЬ).xlsx").touch()

    result = latest_schedule_file(str(tmp_path), 2026, 8)

    assert result is None


def test_latest_schedule_file_returns_none_for_empty_directory(tmp_path):
    assert latest_schedule_file(str(tmp_path), 2026, 8) is None


def test_latest_schedule_file_treats_filename_date_as_one_day_after_the_actual_date(tmp_path):
    """Реальний випадок, підтверджений користувачем: "1бмпВОП-РОП(СЕРПЕНЬ)
    01.09.2026.xlsx" - аркуш усередині каже "Серпень" (відомість ЗА
    31.08.2026, подана ЗРАНКУ 01.09.2026) - дата в ІМЕНІ файлу на ОДИН день
    ПІЗНІША за фактичну дату відомості, тож ЦЕЙ файл МАЄ рахуватись
    серпневим (2026, 8), а НЕ вересневим (2026, 9)."""
    path = tmp_path / "1бмпВОП-РОП(СЕРПЕНЬ) 01.09.2026.xlsx"
    path.touch()

    assert latest_schedule_file(str(tmp_path), 2026, 8) == str(path)
    assert latest_schedule_file(str(tmp_path), 2026, 9) is None


def test_latest_schedule_file_picks_the_latest_actual_date_across_a_month_boundary(tmp_path):
    """Той самий "-1 день" застосовується при виборі НАЙПІЗНІШОГО файлу
    серед КІЛЬКОХ - "31.08.2026" (фактично 30.08) і "01.09.2026" (фактично
    31.08) обидва належать серпню, і другий - пізніший."""
    earlier = tmp_path / "1бмпВОП-РОП(СЕРПЕНЬ) 31.08.2026.xlsx"
    later = tmp_path / "1бмпВОП-РОП(СЕРПЕНЬ) 01.09.2026.xlsx"
    earlier.touch()
    later.touch()

    result = latest_schedule_file(str(tmp_path), 2026, 8)

    assert result == str(later)


def test_read_schedule_marks_finds_columns_despite_header_offset(tmp_path):
    path = _write_schedule_workbook(
        tmp_path / "schedule.xlsx",
        days=[1, 2],
        rows=[("ПЕРШИЙ Перший Перший", {1: "роп", 2: "воп"})],
    )

    marks, _people = read_schedule_marks(path)

    assert marks == {
        (normalize_name("ПЕРШИЙ Перший Перший"), 1): "роп",
        (normalize_name("ПЕРШИЙ Перший Перший"), 2): "воп",
    }


def test_read_schedule_marks_strips_attachment_suffix_before_normalizing(tmp_path):
    """Реальний випадок - ПІБ із допискою прикріплення в дужках наприкінці
    ("... (танкова рота)") - дописка НЕ входить у ключ, інакше зіставлення з
    роcтером провалилось би."""
    path = _write_schedule_workbook(
        tmp_path / "schedule.xlsx",
        days=[1],
        rows=[("ПЕРШИЙ Перший Перший (танкова рота)", {1: "роп"})],
    )

    marks, people = read_schedule_marks(path)

    assert marks == {(normalize_name("ПЕРШИЙ Перший Перший"), 1): "роп"}
    normalized = normalize_name("ПЕРШИЙ Перший Перший")
    assert people[normalized]["pib_raw"] == "ПЕРШИЙ Перший Перший"
    assert people[normalized]["pidrozdil"] == "танкова рота"


def test_read_schedule_marks_skips_blank_day_cells(tmp_path):
    """Порожня клітинка-день - "ще не заповнено", НЕ трактується як "роп"/
    "воп" чи будь-яке інше твердження - взагалі не потрапляє в результат."""
    path = _write_schedule_workbook(
        tmp_path / "schedule.xlsx",
        days=[1, 2],
        rows=[("ПЕРШИЙ Перший Перший", {1: "роп"})],  # день 2 лишається порожнім
    )

    marks, _people = read_schedule_marks(path)

    assert marks == {(normalize_name("ПЕРШИЙ Перший Перший"), 1): "роп"}


def test_read_schedule_marks_stops_at_the_first_blank_pib(tmp_path):
    """Підписний блок нижче ростеру НЕ має порожнього рядка-роздільника перед
    собою в реальних файлах - зупинка на ПЕРШОМУ порожньому ПІБ, а не на
    фіксованій кількості рядків."""
    path = _write_schedule_workbook(
        tmp_path / "schedule.xlsx",
        days=[1],
        rows=[
            ("ПЕРШИЙ Перший Перший", {1: "роп"}),
            (None, {}),
            ("ДРУГИЙ Другий Другий", {1: "воп"}),
        ],
    )

    marks, people = read_schedule_marks(path)

    assert marks == {(normalize_name("ПЕРШИЙ Перший Перший"), 1): "роп"}
    assert normalize_name("ДРУГИЙ Другий Другий") not in people


def test_read_schedule_marks_returns_empty_marks_and_people_when_header_not_found(tmp_path):
    path = tmp_path / "empty.xlsx"
    Workbook().save(str(path))

    assert read_schedule_marks(str(path)) == ({}, {})


def test_read_schedule_marks_captures_posada_and_zvannya_when_present(tmp_path):
    path = _write_schedule_workbook(
        tmp_path / "schedule.xlsx",
        days=[1],
        rows=[("ПЕРШИЙ Перший Перший", {1: "роп"}, "Навідник", "сержант")],
        with_posada_zvannya=True,
    )

    _marks, people = read_schedule_marks(path)

    normalized = normalize_name("ПЕРШИЙ Перший Перший")
    assert people[normalized]["posada"] == "Навідник"
    assert people[normalized]["zvannya"] == "сержант"
    assert people[normalized]["pidrozdil"] is None


def test_read_schedule_marks_posada_and_zvannya_are_none_when_columns_absent(tmp_path):
    path = _write_schedule_workbook(
        tmp_path / "schedule.xlsx",
        days=[1],
        rows=[("ПЕРШИЙ Перший Перший", {1: "роп"})],
    )

    _marks, people = read_schedule_marks(path)

    normalized = normalize_name("ПЕРШИЙ Перший Перший")
    assert people[normalized]["posada"] is None
    assert people[normalized]["zvannya"] is None
