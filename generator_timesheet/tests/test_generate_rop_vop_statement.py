import os
from datetime import date, datetime

from docx_fixtures import write_daily_report
from openpyxl import Workbook, load_workbook

import generators.generate_rop_vop_statement as generate_rop_vop_statement_module
from generators.generate_rop_vop_statement import generate_rop_vop_statement

_HEADER = ["ПОСАДА", "ЗВАННЯ", "ПІБ", datetime(2026, 7, 31), datetime(2026, 8, 1), datetime(2026, 8, 2)]


def _write_roster(path, rows):
    wb = Workbook()
    ws = wb.active
    ws.title = "Табель"
    for col_idx, value in enumerate(_HEADER, start=1):
        ws.cell(row=1, column=col_idx, value=value)
    for row_idx, row_values in enumerate(rows, start=2):
        for col_idx, value in enumerate(row_values, start=1):
            ws.cell(row=row_idx, column=col_idx, value=value)
    wb.save(str(path))
    return str(path)


def _write_info_unit_file(path, rows, dates):
    wb = Workbook()
    ws = wb.active
    ws.title = "ТАБЕЛЬ"
    for col_idx, text in enumerate(["Посада", "Звання", "ПІБ", "ДАТА"], start=1):
        ws.cell(row=1, column=col_idx, value=text)
    for idx, value in enumerate(dates, start=4):
        ws.cell(row=2, column=idx, value=value)
    for row_idx, row_values in enumerate(rows, start=3):
        for col_idx, value in enumerate(row_values, start=1):
            ws.cell(row=row_idx, column=col_idx, value=value)
    wb.save(str(path))
    return str(path)


def test_generate_rop_vop_statement_end_to_end(tmp_path):
    """Реальний випадок: підрозділ подає 70 (РОП) для ПЕРШОГО - той самий
    пайплайн, що й generate_payments_timesheet (apply_payment_values),
    ПОВНІСТЮ незалежний виклик. ДРУГИЙ (лише 100) - НЕ входить у відомість."""
    report_dir = tmp_path / "report"
    report_dir.mkdir()
    write_daily_report(report_dir / "01.08.2026 - щоденний рапорт.docx")
    write_daily_report(report_dir / "02.08.2026 - щоденний рапорт.docx")

    roster_path = _write_roster(tmp_path / "ОБЛІК.xlsx", [
        ["Посада1", "сержант", "ПЕРШИЙ Перший Перший", "РВЗ", None, None],
        ["Посада2", "матрос", "ДРУГИЙ Другий Другий", "РВЗ", None, None],
    ])
    info_unit_dir = tmp_path / "information_unit"
    info_unit_dir.mkdir()
    _write_info_unit_file(
        info_unit_dir / "1 рота.xlsx",
        rows=[
            ["Посада1", "сержант", "ПЕРШИЙ Перший Перший", 70, 70],
            ["Посада2", "матрос", "ДРУГИЙ Другий Другий", 100, 100],
        ],
        dates=[datetime(2026, 8, 1), datetime(2026, 8, 2)],
    )
    output_dir = tmp_path / "output"

    output_path, people_count, mismatch_report_path = generate_rop_vop_statement(
        report_dir=str(report_dir),
        personel_list_file_name=roster_path,
        personel_list_sheet_name="Табель",
        information_unit_dir=str(info_unit_dir),
        schedule_dir=str(tmp_path / "schedule"),
        output_dir=str(output_dir),
        year=2026,
        month=8,
    )

    assert mismatch_report_path is None
    assert people_count == 1
    assert output_path.endswith("1бмпВОП-РОП_СЕРПЕНЬ_.xlsx")
    ws = load_workbook(output_path)["СЕРПЕНЬ"]
    assert ws.cell(row=7, column=4).value == "ПЕРШИЙ Перший Перший"
    assert ws.cell(row=7, column=5).value == "роп"  # 01.08
    assert ws.cell(row=7, column=6).value == "роп"  # 02.08
    assert all(
        ws.cell(row=r, column=4).value != "ДРУГИЙ Другий Другий"
        for r in range(7, ws.max_row + 1)
    )


def test_generate_rop_vop_statement_uses_explicit_output_file_name_when_given(tmp_path):
    """За прямою вказівкою користувача - пункт меню "Згенерувати все"
    (index.py) передає ФІКСОВАНУ назву ("Відомість 170_70.xlsx"), а не
    динамічну (підрозділ+місяць) - output_file_name, ЯКЩО переданий явно,
    МАЄ використовуватись НАПРЯМУ, без жодної зміни."""
    report_dir = tmp_path / "report"
    report_dir.mkdir()
    write_daily_report(report_dir / "01.08.2026 - щоденний рапорт.docx")

    roster_path = _write_roster(tmp_path / "ОБЛІК.xlsx", [
        ["Посада1", "сержант", "ПЕРШИЙ Перший Перший", "РВЗ", None, None],
    ])
    info_unit_dir = tmp_path / "information_unit"
    info_unit_dir.mkdir()
    _write_info_unit_file(
        info_unit_dir / "1 рота.xlsx",
        rows=[["Посада1", "сержант", "ПЕРШИЙ Перший Перший", 70, 70]],
        dates=[datetime(2026, 8, 1), datetime(2026, 8, 2)],
    )
    output_dir = tmp_path / "output"
    explicit_output_file_name = str(output_dir / "Відомість 170_70.xlsx")

    output_path, people_count, mismatch_report_path = generate_rop_vop_statement(
        report_dir=str(report_dir),
        personel_list_file_name=roster_path,
        personel_list_sheet_name="Табель",
        information_unit_dir=str(info_unit_dir),
        schedule_dir=str(tmp_path / "schedule"),
        output_dir=str(output_dir),
        output_file_name=explicit_output_file_name,
        year=2026,
        month=8,
    )

    assert mismatch_report_path is None
    assert people_count == 1
    assert output_path == explicit_output_file_name


def test_generate_rop_vop_statement_saves_a_file_even_when_nobody_qualifies(tmp_path):
    """За прямою вказівкою користувача - формальний документ, ЗАВЖДИ
    зберігається (на відміну від error_mis_statuses.xlsx), навіть якщо
    жодна людина не має РОП/ВОП цього місяця."""
    report_dir = tmp_path / "report"
    report_dir.mkdir()
    write_daily_report(report_dir / "01.08.2026 - щоденний рапорт.docx")

    roster_path = _write_roster(tmp_path / "ОБЛІК.xlsx", [
        ["Посада1", "сержант", "ПЕРШИЙ Перший Перший", "РВЗ", None, None],
    ])
    info_unit_dir = tmp_path / "information_unit"
    info_unit_dir.mkdir()
    _write_info_unit_file(
        info_unit_dir / "1 рота.xlsx",
        rows=[["Посада1", "сержант", "ПЕРШИЙ Перший Перший", 100, 100]],
        dates=[datetime(2026, 8, 1), datetime(2026, 8, 2)],
    )
    output_dir = tmp_path / "output"

    output_path, people_count, mismatch_report_path = generate_rop_vop_statement(
        report_dir=str(report_dir),
        personel_list_file_name=roster_path,
        personel_list_sheet_name="Табель",
        information_unit_dir=str(info_unit_dir),
        schedule_dir=str(tmp_path / "schedule"),
        output_dir=str(output_dir),
        year=2026,
        month=8,
    )

    assert mismatch_report_path is None
    assert people_count == 0
    assert output_dir.joinpath("1бмпВОП-РОП_СЕРПЕНЬ_.xlsx").exists()


# Той самий зсув колонок (8 зайвих колонок перед видимими даними), що й
# реальний resources/schedule/*.xlsx - _write_schedule_file нижче будує файл
# ЦІЄЇ САМОЇ структури (читання шукає колонки за текстом заголовка, а не
# фіксованою позицією - content/schedule_reader.py).
def _write_schedule_file(path, day_values):
    """day_values - {день: "роп"/"воп"} для ОДНОГО ПЕРШИЙ Перший Перший."""
    wb = Workbook()
    ws = wb.active
    junk_columns = 8
    pib_col = junk_columns + 1
    days = sorted(day_values)
    day_cols = {day: pib_col + 1 + i for i, day in enumerate(days)}

    ws.cell(row=6, column=pib_col, value="Прізвище,\nвласне ім'я")
    for day, col in day_cols.items():
        ws.cell(row=6, column=col, value=f"{day:02d}")
    ws.cell(row=7, column=pib_col, value="ПЕРШИЙ Перший Перший")
    for day, value in day_values.items():
        ws.cell(row=7, column=day_cols[day], value=value)

    wb.save(str(path))
    return str(path)


def test_generate_rop_vop_statement_blocks_output_when_schedule_disagrees(tmp_path):
    """За прямою вказівкою користувача - обчислено 70/70 (роп/роп) для
    01.08/02.08, а "офіційна" відомість (resources/schedule) каже "воп" на
    01.08 - розбіжність БЛОКУЄ фінальний файл: замість нього зберігається
    ЛИШЕ звіт розбіжностей."""
    report_dir = tmp_path / "report"
    report_dir.mkdir()
    write_daily_report(report_dir / "01.08.2026 - щоденний рапорт.docx")
    write_daily_report(report_dir / "02.08.2026 - щоденний рапорт.docx")

    roster_path = _write_roster(tmp_path / "ОБЛІК.xlsx", [
        ["Посада1", "сержант", "ПЕРШИЙ Перший Перший", "РВЗ", None, None],
    ])
    info_unit_dir = tmp_path / "information_unit"
    info_unit_dir.mkdir()
    _write_info_unit_file(
        info_unit_dir / "1 рота.xlsx",
        rows=[["Посада1", "сержант", "ПЕРШИЙ Перший Перший", 70, 70]],
        dates=[datetime(2026, 8, 1), datetime(2026, 8, 2)],
    )
    schedule_dir = tmp_path / "schedule"
    schedule_dir.mkdir()
    _write_schedule_file(schedule_dir / "ВІДОМІСТЬ29.08.2026.xlsx", {1: "воп"})
    output_dir = tmp_path / "output"
    mismatch_report_file_name = str(output_dir / "error_mis_statuses_Відомість.xlsx")

    output_path, people_count, mismatch_report_path = generate_rop_vop_statement(
        report_dir=str(report_dir),
        personel_list_file_name=roster_path,
        personel_list_sheet_name="Табель",
        information_unit_dir=str(info_unit_dir),
        schedule_dir=str(schedule_dir),
        output_dir=str(output_dir),
        mismatch_report_file_name=mismatch_report_file_name,
        year=2026,
        month=8,
    )

    assert output_path is None
    assert people_count is None
    assert mismatch_report_path == mismatch_report_file_name
    assert os.path.exists(mismatch_report_file_name)
    assert not output_dir.joinpath("1бмпВОП-РОП_СЕРПЕНЬ_.xlsx").exists()


class _FixedToday:
    """Підмінює модульне ім'я "date" (from datetime import date) у
    generators/generate_rop_vop_statement.py - модуль викликає ЛИШЕ
    date.today() (жодної іншої date-функціональності), тож заглушка
    потребує ЛИШЕ .today()."""
    def __init__(self, fixed):
        self._fixed = fixed

    def today(self):
        return self._fixed


def test_generate_rop_vop_statement_proceeds_when_schedule_agrees(tmp_path):
    """Ті самі дані, що й у тесті вище, але "офіційна" відомість
    ПІДТВЕРДЖУЄ обчислене (роп на 01.08) - жодної розбіжності, звичайна
    генерація відбувається як завжди."""
    report_dir = tmp_path / "report"
    report_dir.mkdir()
    write_daily_report(report_dir / "01.08.2026 - щоденний рапорт.docx")
    write_daily_report(report_dir / "02.08.2026 - щоденний рапорт.docx")

    roster_path = _write_roster(tmp_path / "ОБЛІК.xlsx", [
        ["Посада1", "сержант", "ПЕРШИЙ Перший Перший", "РВЗ", None, None],
    ])
    info_unit_dir = tmp_path / "information_unit"
    info_unit_dir.mkdir()
    _write_info_unit_file(
        info_unit_dir / "1 рота.xlsx",
        rows=[["Посада1", "сержант", "ПЕРШИЙ Перший Перший", 70, 70]],
        dates=[datetime(2026, 8, 1), datetime(2026, 8, 2)],
    )
    schedule_dir = tmp_path / "schedule"
    schedule_dir.mkdir()
    _write_schedule_file(schedule_dir / "ВІДОМІСТЬ29.08.2026.xlsx", {1: "роп"})
    output_dir = tmp_path / "output"

    output_path, people_count, mismatch_report_path = generate_rop_vop_statement(
        report_dir=str(report_dir),
        personel_list_file_name=roster_path,
        personel_list_sheet_name="Табель",
        information_unit_dir=str(info_unit_dir),
        schedule_dir=str(schedule_dir),
        output_dir=str(output_dir),
        year=2026,
        month=8,
    )

    assert mismatch_report_path is None
    assert people_count == 1
    assert output_path.endswith("1бмпВОП-РОП_СЕРПЕНЬ_.xlsx")


def test_generate_rop_vop_statement_defaults_to_latest_report_month_not_todays_date(tmp_path, monkeypatch):
    """Реальний випадок, підтверджений користувачем: сьогодні вже ВЕРЕСЕНЬ
    (system clock), але report_dir/information_unit/schedule ще НЕ мають
    жодного вересневого файлу (типова затримка подачі на початку місяця) -
    year/month=None МАЄ орієнтуватись на місяць НАЙПІЗНІШОГО рапорту, що
    РЕАЛЬНО є (тут - СЕРПЕНЬ), а НЕ на date.today() (ВЕРЕСЕНЬ) - інакше
    генерація дала б ПОРОЖНІЙ результат ("жодна людина не має днів
    РОП/ВОП"), хоча серпневі дані цілком готові."""
    monkeypatch.setattr(generate_rop_vop_statement_module, "date", _FixedToday(date(2026, 9, 3)))

    report_dir = tmp_path / "report"
    report_dir.mkdir()
    write_daily_report(report_dir / "01.08.2026 - щоденний рапорт.docx")
    write_daily_report(report_dir / "02.08.2026 - щоденний рапорт.docx")

    roster_path = _write_roster(tmp_path / "ОБЛІК.xlsx", [
        ["Посада1", "сержант", "ПЕРШИЙ Перший Перший", "РВЗ", None, None],
    ])
    info_unit_dir = tmp_path / "information_unit"
    info_unit_dir.mkdir()
    _write_info_unit_file(
        info_unit_dir / "1 рота.xlsx",
        rows=[["Посада1", "сержант", "ПЕРШИЙ Перший Перший", 70, 70]],
        dates=[datetime(2026, 8, 1), datetime(2026, 8, 2)],
    )
    output_dir = tmp_path / "output"

    output_path, people_count, mismatch_report_path = generate_rop_vop_statement(
        report_dir=str(report_dir),
        personel_list_file_name=roster_path,
        personel_list_sheet_name="Табель",
        information_unit_dir=str(info_unit_dir),
        schedule_dir=str(tmp_path / "schedule"),
        output_dir=str(output_dir),
    )

    assert mismatch_report_path is None
    assert people_count == 1
    assert output_path.endswith("1бмпВОП-РОП_СЕРПЕНЬ_.xlsx")


def test_generate_rop_vop_statement_falls_back_to_today_when_no_report_has_a_parseable_date(tmp_path, monkeypatch):
    """report_dir має ЛИШЕ файл(и) БЕЗ розпізнаваної дати - latest_report_
    date повертає None, тож year/month падає на date.today(), як і раніше."""
    monkeypatch.setattr(generate_rop_vop_statement_module, "date", _FixedToday(date(2026, 8, 1)))

    report_dir = tmp_path / "report"
    report_dir.mkdir()
    write_daily_report(report_dir / "щоденний рапорт без дати.docx")

    roster_path = _write_roster(tmp_path / "ОБЛІК.xlsx", [
        ["Посада1", "сержант", "ПЕРШИЙ Перший Перший", "РВЗ", None, None],
    ])
    info_unit_dir = tmp_path / "information_unit"
    info_unit_dir.mkdir()
    output_dir = tmp_path / "output"

    output_path, _people_count, _mismatch_report_path = generate_rop_vop_statement(
        report_dir=str(report_dir),
        personel_list_file_name=roster_path,
        personel_list_sheet_name="Табель",
        information_unit_dir=str(info_unit_dir),
        schedule_dir=str(tmp_path / "schedule"),
        output_dir=str(output_dir),
    )

    assert output_path.endswith("1бмпВОП-РОП_СЕРПЕНЬ_.xlsx")
