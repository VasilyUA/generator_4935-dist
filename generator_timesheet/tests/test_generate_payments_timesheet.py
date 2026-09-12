from datetime import date, datetime

from docx_fixtures import ARRIVAL_HEADER, write_daily_report
from openpyxl import Workbook, load_workbook

from content.oblik_timesheet import normalize_name
from generators.generate_payments_timesheet import _drop_unresolved_already_in_mismatch_report, generate_payments_timesheet

_HEADER = ["ПОСАДА", "ЗВАННЯ", "ПІБ", datetime(2026, 7, 31), datetime(2026, 8, 1), datetime(2026, 8, 2)]


def _write_roster(path, rows, header=None):
    wb = Workbook()
    ws = wb.active
    ws.title = "Табель"
    for col_idx, value in enumerate(header or _HEADER, start=1):
        ws.cell(row=1, column=col_idx, value=value)
    for row_idx, row_values in enumerate(rows, start=2):
        for col_idx, value in enumerate(row_values, start=1):
            ws.cell(row=row_idx, column=col_idx, value=value)
    wb.save(str(path))
    return str(path)


def _write_info_unit_file(path, rows, dates):
    wb = Workbook()
    ws = wb.active
    ws.title = "ТАБЕЛЬ"
    for col_idx, text in enumerate(["Посада", "Звання", "ПІБ", "ДАТА"], start=1):
        ws.cell(row=1, column=col_idx, value=text)
    for idx, value in enumerate(dates, start=4):
        ws.cell(row=2, column=idx, value=value)
    for row_idx, row_values in enumerate(rows, start=3):
        for col_idx, value in enumerate(row_values, start=1):
            ws.cell(row=row_idx, column=col_idx, value=value)
    wb.save(str(path))
    return str(path)


def _write_info_unit_file_dates_in_row1(path, rows, dates):
    """Реальний структурний варіант: дати ОДРАЗУ в рядку 1 (без окремого
    групового заголовка "ДАТА" + рядка дат) - дані з рядка 2, а не з рядка 3."""
    wb = Workbook()
    ws = wb.active
    ws.title = "ТАБЕЛЬ"
    for col_idx, text in enumerate(["Посада", "Звання", "ПІБ"], start=1):
        ws.cell(row=1, column=col_idx, value=text)
    for idx, value in enumerate(dates, start=4):
        ws.cell(row=1, column=idx, value=value)
    for row_idx, row_values in enumerate(rows, start=2):
        for col_idx, value in enumerate(row_values, start=1):
            ws.cell(row=row_idx, column=col_idx, value=value)
    wb.save(str(path))
    return str(path)


def test_generate_payments_timesheet_replaces_rvz_with_payment_category(tmp_path):
    report_dir = tmp_path / "report"
    report_dir.mkdir()
    write_daily_report(
        report_dir / "01.08.2026 - щоденний рапорт.docx",
        groups=[(
            "ПРИБУЛИ до складу сил та засобів 9 армійського корпусу:",
            [("З відпустки:", ARRIVAL_HEADER, [["1", "матрос", "ПЕРШИЙ Перший Перший", "Посада", "01.08.2026", "Прибув"]])],
        )],
    )
    roster_path = _write_roster(tmp_path / "ОБЛІК.xlsx", [
        ["Посада1", "сержант", "ПЕРШИЙ Перший Перший", "ВП", None, None],
        ["Посада2", "матрос", "ДРУГИЙ Другий Другий", "РВЗ", None, None],
    ])
    info_unit_dir = tmp_path / "information_unit"
    info_unit_dir.mkdir()
    _write_info_unit_file(
        info_unit_dir / "1 рота.xlsx",
        rows=[["Посада1", "сержант", "ПЕРШИЙ Перший Перший", 100, 100]],
        dates=[datetime(2026, 8, 1), datetime(2026, 8, 2)],
    )
    output_path = tmp_path / "output" / "ОБЛІК для виплат.xlsx"
    review_path = tmp_path / "output" / "review.txt"
    mismatch_report_path = tmp_path / "output" / "error_mis_statuses.xlsx"

    result_path, unresolved, _mismatch_open_path = generate_payments_timesheet(
        report_dir=str(report_dir),
        personel_list_file_name=roster_path,
        personel_list_sheet_name="Табель",
        information_unit_dir=str(info_unit_dir),
        output_file_name=str(output_path),
        review_file_name=str(review_path),
        mismatch_report_file_name=str(mismatch_report_path),
    )

    assert result_path == str(output_path)
    saved = load_workbook(str(output_path))["Табель"]
    # ПЕРШИЙ: З відпустки на 01.08 -> РВЗ -> замінено на 100 (подає 1 рота); 02.08 продовжує 100.
    assert [saved.cell(row=2, column=c).value for c in range(4, 6)] == ["ВП", 100]
    # ДРУГИЙ: РВЗ від самого початку, але 1 рота нічого про нього не подає - лишається на ручну перевірку.
    assert saved.cell(row=3, column=4).value == "РВЗ"
    assert any("ДРУГИЙ" in item.get("pib_raw", "") for item in unresolved)
    assert review_path.exists()
    # error_mis_statuses.xlsx: ПЕРШИЙ повністю збігається з "1 рота" (100==100),
    # а про ДРУГОГО цей файл узагалі нічого не подає - розбіжностей немає, тож
    # файл НЕ створюється (за прямою вказівкою користувача - немає сенсу
    # писати файл лише із заголовком).
    assert not mismatch_report_path.exists()


def test_generate_payments_timesheet_creates_neither_review_nor_mismatch_file_when_fully_clean(tmp_path):
    """За прямою вказівкою користувача - коли прогін ПОВНІСТЮ чистий (жодного
    unresolved-запису, жодної розбіжності), НІ Потребує_ручної_перевірки_
    виплати.txt (на відміну від generate_timesheet, де review-файл пишеться
    завжди - skip_review_if_empty=True лише тут), НІ error_mis_statuses.xlsx
    НЕ створюються взагалі - немає сенсу писати файли, коли розповідати
    нема про що."""
    report_dir = tmp_path / "report"
    report_dir.mkdir()
    write_daily_report(report_dir / "01.08.2026 - щоденний рапорт.docx")

    roster_path = _write_roster(
        tmp_path / "ОБЛІК.xlsx",
        [["Посада1", "сержант", "ПЕРШИЙ Перший Перший", "РВЗ"]],
        header=["ПОСАДА", "ЗВАННЯ", "ПІБ", datetime(2026, 8, 1)],
    )
    info_unit_dir = tmp_path / "information_unit"
    info_unit_dir.mkdir()
    _write_info_unit_file(
        info_unit_dir / "1 рота.xlsx",
        rows=[["Посада1", "сержант", "ПЕРШИЙ Перший Перший", 100]],
        dates=[datetime(2026, 8, 1)],
    )

    review_path = tmp_path / "output" / "review.txt"
    mismatch_report_path = tmp_path / "output" / "error_mis_statuses.xlsx"

    _output_path, unresolved, mismatch_open_path = generate_payments_timesheet(
        report_dir=str(report_dir),
        personel_list_file_name=roster_path,
        personel_list_sheet_name="Табель",
        information_unit_dir=str(info_unit_dir),
        output_file_name=str(tmp_path / "output" / "ОБЛІК.xlsx"),
        review_file_name=str(review_path),
        mismatch_report_file_name=str(mismatch_report_path),
    )

    assert unresolved == []
    assert mismatch_open_path is None
    assert not review_path.exists()
    assert not mismatch_report_path.exists()


def _write_bchs_file(path, title_row1, rows):
    wb = Workbook()
    ws = wb.active
    ws.title = "БЧС"
    ws.cell(row=1, column=1, value=title_row1)
    header = ["Посада", "Звання", "ПРІЗВИЩЕ ім'я по батькові", "БЧС", "Виплата додактової винагороди"]
    for col_idx, text in enumerate(header, start=1):
        ws.cell(row=2, column=col_idx, value=text)
    for row_idx, row_values in enumerate(rows, start=3):
        for col_idx, value in enumerate(row_values, start=1):
            ws.cell(row=row_idx, column=col_idx, value=value)
    wb.save(str(path))
    return str(path)


def test_generate_payments_timesheet_replaces_rvz_using_bchs_sheet_only_source(tmp_path):
    """Регресійний тест на реальний випадок: людина отримує "РВЗ" у
    результаті, хоча ЄДИНЕ джерело даних про неї - файл information_unit,
    де ЛИШЕ аркуш БЧС (без жодного релевантного запису в ТАБЕЛЬ) прямо каже
    "100" для потрібної дати - число має підставитись, а не лишитись "РВЗ"."""
    report_dir = tmp_path / "report"
    report_dir.mkdir()
    write_daily_report(report_dir / "01.08.2026 - щоденний рапорт.docx")

    roster_path = _write_roster(tmp_path / "ОБЛІК.xlsx", [
        ["Посада1", "сержант", "ПЕРШИЙ Перший Перший", "РВЗ", None, None],
    ])

    info_unit_dir = tmp_path / "information_unit"
    info_unit_dir.mkdir()
    _write_bchs_file(
        info_unit_dir / "БЧС ВЗ 01.08.2026.xlsx",
        title_row1="ВЗВОД ЗВ'ЯЗКУ 01.08.2026",
        rows=[["Посада1", "сержант", "ПЕРШИЙ Перший Перший", "РВЗ", 100]],
    )

    output_path = tmp_path / "output" / "ОБЛІК для виплат.xlsx"

    _result_path, unresolved, _mismatch_open_path = generate_payments_timesheet(
        report_dir=str(report_dir),
        personel_list_file_name=roster_path,
        personel_list_sheet_name="Табель",
        information_unit_dir=str(info_unit_dir),
        output_file_name=str(output_path),
        review_file_name=str(tmp_path / "output" / "review.txt"),
        mismatch_report_file_name=str(tmp_path / "output" / "error_mis_statuses.xlsx"),
    )

    # 31.07 (базова колонка) - у файлі information_unit немає про неї даних,
    # лишається "РВЗ" на ручну перевірку - це ОКРЕМЕ, очікуване питання, не
    # те, що перевіряє цей тест.
    assert not any("01.08.2026" in item["reason"] for item in unresolved)
    saved = load_workbook(str(output_path))["Табель"]
    assert saved.cell(row=2, column=5).value == 100  # 01.08.2026 - замінено на 100, а не лишилось "РВЗ"


def test_generate_payments_timesheet_replaces_rvz_when_tabel_sheet_has_dates_in_row1(tmp_path):
    """Регресійний тест на реальний випадок: аркуш ТАБЕЛЬ файлу information_unit
    має дати ОДРАЗУ в рядку 1 (без окремого рядка-групи "ДАТА" в рядку 1 +
    самих дат у рядку 2, як у решти файлів) - раніше такий файл мовчки НЕ
    читався взагалі (жодної колонки-дати не знаходилось), тож людина
    лишалась "РВЗ" на кожен день, хоча файл прямо каже "100"."""
    header = ["ПОСАДА", "ЗВАННЯ", "ПІБ", datetime(2026, 7, 31), datetime(2026, 8, 1), datetime(2026, 8, 2), datetime(2026, 8, 3)]
    wb = Workbook()
    ws = wb.active
    ws.title = "Табель"
    for col_idx, value in enumerate(header, start=1):
        ws.cell(row=1, column=col_idx, value=value)
    ws.cell(row=2, column=1, value="Командир роти")
    ws.cell(row=2, column=2, value="старший лейтенант")
    ws.cell(row=2, column=3, value="ПЕРШИЙ Перший Перший")
    ws.cell(row=2, column=4, value="РВЗ")
    roster_path = tmp_path / "ОБЛІК.xlsx"
    wb.save(str(roster_path))

    report_dir = tmp_path / "report"
    report_dir.mkdir()
    write_daily_report(report_dir / "03.08.2026 - щоденний рапорт.docx")

    info_unit_dir = tmp_path / "information_unit"
    info_unit_dir.mkdir()
    _write_info_unit_file_dates_in_row1(
        info_unit_dir / "підрозділ 1.xlsx",
        rows=[["Командир роти", "старший лейтенант", "ПЕРШИЙ Перший Перший", 100, 100, 100]],
        dates=[datetime(2026, 8, 1), datetime(2026, 8, 2), datetime(2026, 8, 3)],
    )

    output_path = tmp_path / "output" / "ОБЛІК для виплат.xlsx"

    generate_payments_timesheet(
        report_dir=str(report_dir),
        personel_list_file_name=str(roster_path),
        personel_list_sheet_name="Табель",
        information_unit_dir=str(info_unit_dir),
        output_file_name=str(output_path),
        review_file_name=str(tmp_path / "output" / "review.txt"),
        mismatch_report_file_name=str(tmp_path / "output" / "error_mis_statuses.xlsx"),
    )

    saved = load_workbook(str(output_path))["Табель"]
    # 01.08, 02.08, 03.08 - усі три дні замінено на 100, а не лишились "РВЗ".
    assert [saved.cell(row=2, column=c).value for c in range(5, 8)] == [100, 100, 100]


def test_generate_payments_timesheet_writes_mismatch_report_after_saving_payments_file(tmp_path):
    """Реальний випадок: файл information_unit каже "ВД" (не число) для
    людини - apply_payment_values НЕ підставляє нечислове значення (лишає
    "РВЗ"), тож звіт розбіжностей (сформований ОДРАЗУ ПІСЛЯ ОБЛІК для виплат,
    звіряючи саме ЦЕЙ згенерований результат) має явно показати "РВЗ" і "ВД"
    поряд, з назвою файлу-джерела."""
    report_dir = tmp_path / "report"
    report_dir.mkdir()
    write_daily_report(report_dir / "01.08.2026 - щоденний рапорт.docx")

    roster_path = _write_roster(tmp_path / "ОБЛІК.xlsx", [
        ["Посада1", "сержант", "ПЕРШИЙ Перший Перший", "РВЗ", None, None],
    ])

    info_unit_dir = tmp_path / "information_unit"
    info_unit_dir.mkdir()
    _write_bchs_file(
        info_unit_dir / "БЧС ВЗ 01.08.2026.xlsx",
        title_row1="ВЗВОД ЗВ'ЯЗКУ 01.08.2026",
        rows=[["Посада1", "сержант", "ПЕРШИЙ Перший Перший", "ВД", 100]],
    )

    mismatch_report_path = tmp_path / "output" / "error_mis_statuses.xlsx"

    _output_path, _unresolved, mismatch_open_path = generate_payments_timesheet(
        report_dir=str(report_dir),
        personel_list_file_name=roster_path,
        personel_list_sheet_name="Табель",
        information_unit_dir=str(info_unit_dir),
        output_file_name=str(tmp_path / "output" / "ОБЛІК для виплат.xlsx"),
        review_file_name=str(tmp_path / "output" / "review.txt"),
        mismatch_report_file_name=str(mismatch_report_path),
    )

    # Є хоч один запис - третій елемент повертає шлях до звіту (index.py
    # відкриває його автоматично), а не None.
    assert mismatch_open_path == str(mismatch_report_path)
    saved = load_workbook(str(mismatch_report_path))["Розбіжності"]
    assert saved.cell(row=2, column=1).value == "ОБЛІК для виплат (за рапортами)"
    assert saved.cell(row=3, column=1).value == "БЧС ВЗ 01.08.2026.xlsx"
    assert saved.cell(row=2, column=5).value == "ПЕРШИЙ Перший Перший"


def _mismatch_record(**overrides):
    base = {
        "file_name": "файл1.xlsx", "roster_pib_raw": "ПЕРШИЙ Перший Перший", "file_pib_raw": "ПЕРШИЙ Перший Перший",
        "roster_pidrozdil": None, "roster_posada": "Посада1", "roster_zvannya": "сержант",
        "file_pidrozdil": None, "file_posada": "Посада1", "file_zvannya": "сержант",
        "dates": {date(2026, 8, 1): ("РВЗ", "ВЛК")},
    }
    base.update(overrides)
    return base


def test_drop_unresolved_already_in_mismatch_report_removes_the_duplicate():
    """За прямою вказівкою користувача - unresolved-запис apply_payment_values
    для (людина, дата), яка ВЖЕ показана як розбіжність у
    error_mis_statuses.xlsx (records), - дублює той самий факт у ДВОХ файлах
    одразу, тож прибирається, лишаючи ЛИШЕ error_mis_statuses.xlsx."""
    unresolved = [{
        "reason": "\"ПЕРШИЙ Перший Перший\" 01.08.2026: підрозділ подає ['ВЛК'] - не єдине числове значення категорії виплати, статус \"РВЗ\" лишено як є.",
        "pib_raw": "ПЕРШИЙ Перший Перший", "report_date": date(2026, 8, 1),
    }]
    records = [_mismatch_record()]

    assert _drop_unresolved_already_in_mismatch_report(unresolved, records) == []


def test_drop_unresolved_already_in_mismatch_report_keeps_entries_without_a_matching_record():
    """Запис БЕЗ відповідника в records (напр. "Немає категорії виплати" -
    підрозділ узагалі нічого не подав, або ІНША дата/людина) НЕ прибирається -
    для НЬОГО error_mis_statuses.xlsx нічого не показує (немає даних для
    порівняння), тож лишається єдиним джерелом інформації."""
    unresolved = [
        {
            "reason": "Немає категорії виплати для \"ДРУГИЙ Другий Другий\" на 01.08.2026 - підрозділ ще не подав дані.",
            "pib_raw": "ДРУГИЙ Другий Другий", "report_date": date(2026, 8, 1),
        },
        {
            "reason": "...",
            "pib_raw": "ПЕРШИЙ Перший Перший", "report_date": date(2026, 8, 2),  # ІНША дата, не 01.08
        },
    ]
    records = [_mismatch_record()]

    assert _drop_unresolved_already_in_mismatch_report(unresolved, records) == unresolved


def test_generate_payments_timesheet_does_not_duplicate_non_numeric_mismatch_in_review_file(tmp_path):
    """Реальний випадок: файл information_unit каже "ВЛК" (не число) для
    людини - apply_payment_values лишає "РВЗ" і, БЕЗ цього виправлення, додав
    би унерозв'язаний запис "не єдине числове значення" в Потребує_ручної_
    перевірки_виплати.txt - хоча ТА САМА розбіжність ВЖЕ показана в
    error_mis_statuses.xlsx (ВЛК - "порівнюваний" статус). За прямою
    вказівкою користувача - НЕ дублюється: unresolved НЕ містить цей запис,
    error_mis_statuses.xlsx і далі показує розбіжність."""
    report_dir = tmp_path / "report"
    report_dir.mkdir()
    write_daily_report(report_dir / "01.08.2026 - щоденний рапорт.docx")

    roster_path = _write_roster(tmp_path / "ОБЛІК.xlsx", [
        ["Посада1", "сержант", "ВОСЬМИЙ Восьмий Восьмий", "РВЗ", None, None],
    ])

    info_unit_dir = tmp_path / "information_unit"
    info_unit_dir.mkdir()
    _write_bchs_file(
        info_unit_dir / "БЧС ВЗ 01.08.2026.xlsx",
        title_row1="ВЗВОД ЗВ'ЯЗКУ 01.08.2026",
        rows=[["Посада1", "сержант", "ВОСЬМИЙ Восьмий Восьмий", "ВЛК", 100]],
    )

    mismatch_report_path = tmp_path / "output" / "error_mis_statuses.xlsx"

    _output_path, unresolved, mismatch_open_path = generate_payments_timesheet(
        report_dir=str(report_dir),
        personel_list_file_name=roster_path,
        personel_list_sheet_name="Табель",
        information_unit_dir=str(info_unit_dir),
        output_file_name=str(tmp_path / "output" / "ОБЛІК для виплат.xlsx"),
        review_file_name=str(tmp_path / "output" / "review.txt"),
        mismatch_report_file_name=str(mismatch_report_path),
    )

    # Дубльований запис (01.08 - "не єдине числове значення", ВЖЕ показано в
    # error_mis_statuses.xlsx) прибрано, АЛЕ 31.07 (базова колонка, "Немає
    # категорії виплати" - інша, НЕпов'язана причина, для НЕЇ немає розбіжності
    # у файлі, отже НЕ прибирається) - і далі лишається.
    assert not any("не єдине числове значення" in item["reason"] for item in unresolved)
    assert any("Немає категорії виплати" in item["reason"] for item in unresolved)
    assert mismatch_open_path == str(mismatch_report_path)
    saved = load_workbook(str(mismatch_report_path))["Розбіжності"]
    assert saved.cell(row=3, column=1).value == "БЧС ВЗ 01.08.2026.xlsx"


def test_generate_payments_timesheet_mismatch_report_includes_pidrozdil_column(tmp_path):
    """Роcтер із колонкою ПІДРОЗДІЛ - звіт розбіжностей має показати її поряд
    із Посадою/Званням, звірену з файлом information_unit (за прямою
    вказівкою користувача) - файл цього тесту не подає Підрозділ узагалі
    (None), тож роcтерове значення "підрозділ 1" - розбіжність, НЕ об'єднана
    клітинка."""
    report_dir = tmp_path / "report"
    report_dir.mkdir()
    write_daily_report(report_dir / "01.08.2026 - щоденний рапорт.docx")

    roster_path = _write_roster(
        tmp_path / "ОБЛІК.xlsx",
        [["підрозділ 1", "Посада1", "сержант", "ПЕРШИЙ Перший Перший", "РВЗ", None]],
        header=["ПІДРОЗДІЛ", "ПОСАДА", "ЗВАННЯ", "ПІБ", datetime(2026, 7, 31), datetime(2026, 8, 1)],
    )

    info_unit_dir = tmp_path / "information_unit"
    info_unit_dir.mkdir()
    _write_bchs_file(
        info_unit_dir / "БЧС ВЗ 01.08.2026.xlsx",
        title_row1="ВЗВОД ЗВ'ЯЗКУ 01.08.2026",
        rows=[["Посада1", "сержант", "ПЕРШИЙ Перший Перший", "ВД", 100]],
    )

    mismatch_report_path = tmp_path / "output" / "error_mis_statuses.xlsx"

    generate_payments_timesheet(
        report_dir=str(report_dir),
        personel_list_file_name=roster_path,
        personel_list_sheet_name="Табель",
        information_unit_dir=str(info_unit_dir),
        output_file_name=str(tmp_path / "output" / "ОБЛІК для виплат.xlsx"),
        review_file_name=str(tmp_path / "output" / "review.txt"),
        mismatch_report_file_name=str(mismatch_report_path),
    )

    saved = load_workbook(str(mismatch_report_path))["Розбіжності"]
    assert saved.cell(row=1, column=2).value == "Підрозділ"
    assert saved.cell(row=2, column=2).value == "підрозділ 1"
    merged = {str(rng) for rng in saved.merged_cells.ranges}
    assert "B2:B3" not in merged  # Підрозділ - файл None, роcтер "підрозділ 1" - розбіжність


def test_generate_payments_timesheet_mismatch_report_includes_people_missing_from_roster(tmp_path):
    """Реальний випадок: людина є у файлі information_unit, але ВЗАГАЛІ
    відсутня в роcтері (ОБЛІК.xlsx) - раніше мовчки не потрапляла в
    error_mis_statuses.xlsx узагалі (звіт порівнював лише людей з ростера)."""
    report_dir = tmp_path / "report"
    report_dir.mkdir()
    write_daily_report(report_dir / "01.08.2026 - щоденний рапорт.docx")

    roster_path = _write_roster(tmp_path / "ОБЛІК.xlsx", [
        ["Посада1", "сержант", "ПЕРШИЙ Перший Перший", "РВЗ", None],
    ])

    info_unit_dir = tmp_path / "information_unit"
    info_unit_dir.mkdir()
    _write_bchs_file(
        info_unit_dir / "БЧС ВЗ 01.08.2026.xlsx",
        title_row1="ВЗВОД ЗВ'ЯЗКУ 01.08.2026",
        rows=[
            ["Посада1", "сержант", "ПЕРШИЙ Перший Перший", "РВЗ", 100],
            ["Посада3", "матрос", "ТРЕТІЙ Третій Третій", "РВЗ", 100],
        ],
    )

    mismatch_report_path = tmp_path / "output" / "error_mis_statuses.xlsx"

    generate_payments_timesheet(
        report_dir=str(report_dir),
        personel_list_file_name=roster_path,
        personel_list_sheet_name="Табель",
        information_unit_dir=str(info_unit_dir),
        output_file_name=str(tmp_path / "output" / "ОБЛІК для виплат.xlsx"),
        review_file_name=str(tmp_path / "output" / "review.txt"),
        mismatch_report_file_name=str(mismatch_report_path),
    )

    saved = load_workbook(str(mismatch_report_path))["Розбіжності"]
    # Немає роcтерового рядка для ТРЕТІЙ - ПІБ НЕ об'єднана, ім'я показане
    # лише на рядку файлу (row_b); роcтерова сторона (row_a) показує
    # "Не знайдено в ОБЛІК.xlsx", а не порожню клітинку.
    pib_values = [saved.cell(row=r, column=5).value for r in range(2, saved.max_row + 1)]
    assert "ТРЕТІЙ Третій Третій" in pib_values
    row = next(r for r in range(2, saved.max_row + 1) if saved.cell(row=r, column=5).value == "ТРЕТІЙ Третій Третій")
    assert saved.cell(row=row, column=1).value == "БЧС ВЗ 01.08.2026.xlsx"
    assert saved.cell(row=row - 1, column=1).value == "ОБЛІК для виплат (за рапортами)"
    assert saved.cell(row=row - 1, column=5).value == "Не знайдено в ОБЛІК.xlsx"
    # Дата (01.08) - ОКРЕМА колонка (заголовок - сама дата), як в ОБЛІК.xlsx.
    date_col = next(c for c in range(6, saved.max_column + 1) if saved.cell(row=1, column=c).value == datetime(2026, 8, 1))
    assert saved.cell(row=row - 1, column=date_col).value == "Не знайдено в ОБЛІК.xlsx"
    assert saved.cell(row=row, column=date_col).value == 100


def test_generate_payments_timesheet_mismatch_report_flags_patronymic_spelling_difference(tmp_path):
    """Реальний випадок: роcтер (ОБЛІК.xlsx) і файл information_unit
    подають ТЕ САМЕ прізвище+ім'я, але РІЗНЕ по-батькові (друкарська
    помилка) - раніше це взагалі не потрапляло в error_mis_statuses.xlsx як
    порівнянна розбіжність ПІБ (normalize_name різний -> дані файлу не
    зіставлялись з роcтером), тепер має показати РЕАЛЬНЕ роcтерове ПІБ поряд
    із файловим написанням, а не мовчати чи показувати лише загальний
    напис "не знайдено"."""
    report_dir = tmp_path / "report"
    report_dir.mkdir()
    write_daily_report(report_dir / "01.08.2026 - щоденний рапорт.docx")

    roster_path = _write_roster(tmp_path / "ОБЛІК.xlsx", [
        ["Посада1", "сержант", "ПЕРШИЙ Перший Перший", "РВЗ", None],
    ])

    info_unit_dir = tmp_path / "information_unit"
    info_unit_dir.mkdir()
    _write_bchs_file(
        info_unit_dir / "БЧС ВЗ 01.08.2026.xlsx",
        title_row1="ВЗВОД ЗВ'ЯЗКУ 01.08.2026",
        # "Перекший" замість роcтерового "Перший" - та сама людина, друкарська
        # помилка в по-батькові (реальний випадок з подібною одно-двобуквеною різницею).
        rows=[["Посада1", "сержант", "ПЕРШИЙ Перший Перекший", "РВЗ", 100]],
    )

    mismatch_report_path = tmp_path / "output" / "error_mis_statuses.xlsx"

    generate_payments_timesheet(
        report_dir=str(report_dir),
        personel_list_file_name=roster_path,
        personel_list_sheet_name="Табель",
        information_unit_dir=str(info_unit_dir),
        output_file_name=str(tmp_path / "output" / "ОБЛІК для виплат.xlsx"),
        review_file_name=str(tmp_path / "output" / "review.txt"),
        mismatch_report_file_name=str(mismatch_report_path),
    )

    saved = load_workbook(str(mismatch_report_path))["Розбіжності"]
    assert saved.cell(row=2, column=1).value == "ОБЛІК для виплат (за рапортами)"
    assert saved.cell(row=3, column=1).value == "БЧС ВЗ 01.08.2026.xlsx"
    # РЕАЛЬНЕ роcтерове написання - НЕ загальний напис "не знайдено" -
    # показане поряд із файловим, щоб розбіжність у ПІБ було видно.
    assert saved.cell(row=2, column=5).value == "ПЕРШИЙ Перший Перший"
    assert saved.cell(row=3, column=5).value == "ПЕРШИЙ Перший Перекший"
    merged = {str(rng) for rng in saved.merged_cells.ranges}
    assert "E2:E3" not in merged  # ПІБ відрізняється - НЕ об'єднана
    assert saved.cell(row=2, column=5).fill.fgColor.rgb == "FFFFC7CE"
    assert saved.cell(row=3, column=5).fill.fgColor.rgb == "FFFFC7CE"


def test_generate_payments_timesheet_does_not_create_mismatch_report_when_everything_matches(tmp_path):
    """За прямою вказівкою користувача - немає сенсу створювати
    error_mis_statuses.xlsx лише із заголовком, коли розповідати нема про що."""
    report_dir = tmp_path / "report"
    report_dir.mkdir()
    write_daily_report(report_dir / "01.08.2026 - щоденний рапорт.docx")

    roster_path = _write_roster(tmp_path / "ОБЛІК.xlsx", [
        ["Посада1", "сержант", "ПЕРШИЙ Перший Перший", "РВЗ", None, None],
    ])
    info_unit_dir = tmp_path / "information_unit"
    info_unit_dir.mkdir()

    mismatch_report_path = tmp_path / "output" / "error_mis_statuses.xlsx"

    _result_path, unresolved, mismatch_open_path = generate_payments_timesheet(
        report_dir=str(report_dir),
        personel_list_file_name=roster_path,
        personel_list_sheet_name="Табель",
        information_unit_dir=str(info_unit_dir),
        output_file_name=str(tmp_path / "output" / "ОБЛІК для виплат.xlsx"),
        review_file_name=str(tmp_path / "output" / "review.txt"),
        mismatch_report_file_name=str(mismatch_report_path),
    )

    assert unresolved != []  # ПЕРШИЙ лишається "РВЗ" без жодних даних про виплату
    assert mismatch_open_path is None  # немає жодного запису - нема сенсу відкривати звіт автоматично
    assert not mismatch_report_path.exists()


def test_generate_payments_timesheet_uses_the_same_output_file_name_as_generate_timesheet():
    """За прямою вказівкою користувача - PAYMENTS_OUTPUT_FILE_NAME НАВМИСНО
    дорівнює OUTPUT_FILE_NAME (обидва - "ОБЛІК.xlsx"): обидва прогони завжди
    окремі (output/ очищається перед кожним, index.py), колізії імені
    немає."""
    from constants import OUTPUT_FILE_NAME, PAYMENT_MISMATCH_REPORT_FILE_NAME, PAYMENTS_OUTPUT_FILE_NAME
    assert PAYMENTS_OUTPUT_FILE_NAME == OUTPUT_FILE_NAME
    assert PAYMENT_MISMATCH_REPORT_FILE_NAME not in (OUTPUT_FILE_NAME, PAYMENTS_OUTPUT_FILE_NAME)
