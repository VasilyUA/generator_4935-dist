import os
from collections import Counter
from datetime import date, datetime

import pytest
from book61_fixtures import write_full_book61
from cvs_fixtures import MEDICAL_SHEET, VLK_SHEET, write_cvs
from openpyxl import Workbook, load_workbook

import content.addon_mismatch_checker as amc
from content.oblik_timesheet import Timesheet, normalize_name
from generators.check_addon_mismatches import addon_files, check_addon_mismatches_file

_HEADER = ["ПОСАДА", "ЗВАННЯ", "ПІБ", datetime(2026, 7, 1), datetime(2026, 7, 2), datetime(2026, 7, 3)]
_WEEK_HEADER = ["ПОСАДА", "ЗВАННЯ", "ПІБ"] + [datetime(2026, 10, day) for day in range(1, 8)]


def _write_roster(path, rows, header=_HEADER):
    wb = Workbook()
    ws = wb.worksheets[0]
    ws.title = "Табель"
    for col_idx, value in enumerate(header, start=1):
        ws.cell(row=1, column=col_idx, value=value)
    for row_idx, row_values in enumerate(rows, start=2):
        for col_idx, value in enumerate(row_values, start=1):
            ws.cell(row=row_idx, column=col_idx, value=value)
    wb.save(str(path))
    return str(path)


def _period(pib_raw, status, start, end, source_sheet="ВП", source_file="book61"):
    return {
        "pib_raw": pib_raw, "status": status, "start_date": start, "end_date": end,
        "source_sheet": source_sheet, "source_file": source_file,
    }


# --- _statuses_match ---------------------------------------------------------

def test_statuses_match_exact():
    assert amc._statuses_match("ВП", "ВП") is True
    assert amc._statuses_match("ВП", "ВЛК") is False


def test_statuses_match_treats_bpshp_and_vpbp_as_equivalent():
    """Рішення користувача: book61 НЕ розрізняє "БПШП"/"ВПБП"
    на рівні даних (content/addon_reader.py::read_shp_family_periods)."""
    assert amc._statuses_match("БПШП", "ВПБП") is True
    assert amc._statuses_match("ВПБП", "БПШП") is True


def test_statuses_match_treats_ppd_and_nov_as_equivalent():
    """Рішення користувача: "ППД" і "НОВ" ідентичні статуси,
    це не розбіжність (незалежно від регістру написання в Табелі)."""
    assert amc._statuses_match("ППД", "НОВ") is True
    assert amc._statuses_match("НОВ", "ППД") is True
    assert amc._statuses_match("ППД", "Нов") is True
    assert amc._statuses_match("СЗЧ", "НОВ") is False


def test_statuses_match_treats_zagyb_and_smert_as_equivalent():
    """Рішення користувача: "Смерть" і "Загиб" не дають
    розбіжності між собою (реальний випадок: book61 "Загиб", Табель
    "Смерть")."""
    assert amc._statuses_match("Загиб", "Смерть") is True
    assert amc._statuses_match("Смерть", "Загиб") is True
    assert amc._statuses_match("Загиб", "РВЗ") is False


def test_statuses_match_treats_ppd_and_bzvp_as_equivalent():
    """Рішення користувача: "ППД" і "БЗВП" ідентичні статуси,
    це не розбіжність."""
    assert amc._statuses_match("ППД", "БЗВП") is True
    assert amc._statuses_match("БЗВП", "ППД") is True
    assert amc._statuses_match("ППД", "Бзвп") is True


def test_statuses_match_prefers_tabel_10_med_over_book61_vpsz():
    """Рішення користувача: book61 "ВПСЗ", Табель "10_МЕД":
    пріоритетніший 10_МЕД (за нього платиться), це не розбіжність. Пріоритет
    односторонній - це не еквівалентність статусів."""
    assert amc._statuses_match("ВПСЗ", "10_МЕД") is True
    assert amc._statuses_match("ВПСЗ", "10_мед") is True
    assert amc._statuses_match("10_МЕД", "ВПСЗ") is False
    assert amc._statuses_match("ВПСЗ", "РВЗ") is False


@pytest.mark.parametrize("expected", ["БЗВП", "Адап"])
def test_statuses_match_prefers_tabel_rrez_over_an_expected_bzvp_or_adaptation(expected):
    """Рішення користувача: "РРез" пріоритетніший за "БЗВП" і
    "Адап": журнал очікує БЗВП/Адап, Табель - РРез (будь-яке написання) - не
    розбіжність; навпаки (очікується РРез, у Табелі БЗВП/Адап) - розбіжність."""
    assert amc._statuses_match(expected, "РРез") is True
    assert amc._statuses_match(expected, "Ррез") is True
    assert amc._statuses_match("РРез", expected) is False


def test_find_period_mismatches_skips_vpsz_when_tabel_has_10_med(tmp_path):
    path = _write_roster(tmp_path / "ОБЛІК.xlsx", [
        ["Посада1", "сержант", "ПЕРШИЙ Перший Перший", "10_МЕД", "10_МЕД", "10_МЕД"],
    ])
    timesheet = Timesheet(path, "Табель")
    records = [_period("ПЕРШИЙ Перший Перший", "ВПСЗ", date(2026, 7, 1), None, source_sheet="ВП")]

    assert amc.find_period_mismatches(timesheet, records) == []


def test_find_period_mismatches_skips_a_person_whose_only_difference_is_ppd_vs_nov(tmp_path):
    path = _write_roster(tmp_path / "ОБЛІК.xlsx", [
        ["Посада1", "сержант", "ПЕРШИЙ Перший Перший", "НОВ", "НОВ", "ППД"],
    ])
    timesheet = Timesheet(path, "Табель")
    records = [_period("ПЕРШИЙ Перший Перший", "ППД", date(2026, 7, 1), None, source_sheet="ППД")]

    assert amc.find_period_mismatches(timesheet, records) == []


def test_write_addon_mismatch_report_shows_the_tabel_value_for_an_equivalent_status(tmp_path):
    """Рішення користувача: пріоритет Табеля: для ППД/НОВ
    клітинка - ОДНА (об'єднана), зелена, зі значенням із Табеля."""
    period_mismatches = [{
        "source_file": "book61",
        "sources": ["book61: ППД", "book61: СЗЧ"], "roster_pidrozdil": "Підрозділ1",
        "roster_pib_raw": "ПЕРШИЙ Перший Перший", "addon_pib_raw": "ПЕРШИЙ Перший Перший",
        "dates": {date(2026, 10, 1): ("НОВ", "ППД"), date(2026, 10, 2): ("РВЗ", "СЗЧ")},
    }]
    output_path = str(tmp_path / "mis_match.xlsx")

    amc.write_addon_mismatch_report(period_mismatches, [], output_path)

    ws = load_workbook(output_path)["book61"]
    assert ws.cell(row=2, column=4).value == "НОВ"
    assert "D2:D3" in {str(merged) for merged in ws.merged_cells.ranges}
    assert ws.cell(row=2, column=4).fill.start_color.rgb == amc._MATCH_FILL.start_color.rgb


def test_statuses_match_treats_none_as_distinct_from_any_status():
    assert amc._statuses_match("ВП", None) is False


def test_statuses_match_treats_payment_converted_forms_as_equivalent():
    """Реальні випадки на output/ОБЛІК.xlsx: apply_payment_values (content/
    oblik_timesheet.py) може замінити ПОДІЄВИЙ статус на фіксовану
    категорію виплати (constants.py: PAYMENT_STATUS_CONVERSIONS/
    PAYMENT_STATUS_OVERRIDES) - book61 звіряється зі своїм ВЛАСНИМ,
    незалежним визначенням факту (до будь-якої конвертації), тож це й
    далі ТА САМА ситуація, а не розбіжність ("БПШП"/"100_БПШП",
    "ВД"/10)."""
    assert amc._statuses_match("БПШП", "100_БПШП") is True
    assert amc._statuses_match("ВД", 10) is True
    assert amc._statuses_match("ВПСЗ", "100_ВПБП") is True


def test_statuses_match_does_not_treat_navch_as_equivalent_to_10():
    """Застереження користувача: "10" еквівалентний "ВД", АЛЕ якщо book61
    очікує саме "Навч" (не "ВД") - число 10 в Табелі й далі МАЄ лишатись
    розбіжністю, а не гаситись через ширше правило "10 ~ ВД" вище."""
    assert amc._statuses_match("Навч", 10) is False


# --- find_period_mismatches ---------------------------------------------------

def test_find_period_mismatches_skips_a_matching_period(tmp_path):
    path = _write_roster(tmp_path / "ОБЛІК.xlsx", [
        ["Посада1", "сержант", "ПЕРШИЙ Перший Перший", "ВП", "ВП", "РВЗ"],
    ])
    timesheet = Timesheet(path, "Табель")
    records = [_period("ПЕРШИЙ Перший Перший", "ВП", date(2026, 7, 1), date(2026, 7, 2))]

    assert amc.find_period_mismatches(timesheet, records) == []


def test_find_period_mismatches_reports_a_genuine_mismatch(tmp_path):
    path = _write_roster(tmp_path / "ОБЛІК.xlsx", [
        ["Посада1", "сержант", "ПЕРШИЙ Перший Перший", "РВЗ", "РВЗ", "РВЗ"],
    ])
    timesheet = Timesheet(path, "Табель")
    records = [_period("ПЕРШИЙ Перший Перший", "ВП", date(2026, 7, 1), date(2026, 7, 2))]

    results = amc.find_period_mismatches(timesheet, records)

    assert len(results) == 1
    assert results[0]["sources"] == ["book61: ВП"]
    assert results[0]["roster_pib_raw"] == "ПЕРШИЙ Перший Перший"
    assert results[0]["dates"] == {
        date(2026, 7, 1): ("РВЗ", "ВП"),
        date(2026, 7, 2): ("РВЗ", "ВП"),
    }


def test_find_period_mismatches_ignores_not_yet_filled_in_cells(tmp_path):
    """Реальний випадок (виявлено на book61.xlsm): дата ПІСЛЯ останнього
    наявного рапорту - apply_events (content/oblik_timesheet.py) навмисно
    лишає такі колонки порожніми - це НЕ розбіжність, ще просто немає
    даних, щоб судити."""
    path = _write_roster(tmp_path / "ОБЛІК.xlsx", [
        ["Посада1", "сержант", "ПЕРШИЙ Перший Перший", "ВП", None, None],
    ])
    timesheet = Timesheet(path, "Табель")
    records = [_period("ПЕРШИЙ Перший Перший", "ВП", date(2026, 7, 1), date(2026, 7, 3))]

    assert amc.find_period_mismatches(timesheet, records) == []


def test_find_period_mismatches_never_includes_a_person_missing_from_the_roster(tmp_path):
    """Рішення користувача: БЕЗ винятків: кого немає в Табелі,
    той не звіряється з жодним аркушем book61 (реальні випадки з аркушів
    "ВД" і "СЗЧ" хибно потрапляли в mis_match.xlsx)."""
    path = _write_roster(tmp_path / "ОБЛІК.xlsx", [
        ["Посада1", "сержант", "ІНШИЙ Інший Інший", "РВЗ", "РВЗ", "РВЗ"],
    ])
    timesheet = Timesheet(path, "Табель")
    records = [
        _period("ПЕРШИЙ Перший Перший", "Навч", date(2026, 7, 1), date(2026, 7, 2), source_sheet="ВД"),
        _period("ДРУГИЙ Другий Другий", "ППД", date(2026, 7, 1), None, source_sheet="ППД"),
        _period("ТРЕТІЙ Третій Третій", "СЗЧ", date(2026, 7, 1), None, source_sheet="СЗЧ"),
        _period("ЧЕТВЕРТИЙ Четвертий Четвертий", "ПРВДБ", date(2023, 10, 31), None, source_sheet="Алф2"),
    ]

    assert amc.find_period_mismatches(timesheet, records) == []


def test_find_period_mismatches_builds_one_sequential_timeline_per_person(tmp_path):
    """Рішення користувача (реальний випадок): book61 "ППД" до
    04.10, "ШП" 03.10-06.10, "СЗЧ" з 06.10 - ОДИН рядок звіту, кожен період
    триває до дня перед початком наступного: ППД до 02.10 -> ШП
    03.10-05.10 -> СЗЧ з 06.10."""
    path = _write_roster(tmp_path / "ОБЛІК.xlsx", [
        ["Посада1", "сержант", "ПЕРШИЙ Перший Перший", "ППД", "ППД", "ШП", "ШП", "ШП", "СЗЧ", "РВЗ"],
    ], header=_WEEK_HEADER)
    timesheet = Timesheet(path, "Табель")
    records = [
        _period("ПЕРШИЙ Перший Перший", "СЗЧ", date(2026, 10, 6), None, source_sheet="СЗЧ"),
        _period("ПЕРШИЙ Перший Перший", "ППД", date(2026, 8, 13), date(2026, 10, 4), source_sheet="ППД"),
        _period("ПЕРШИЙ Перший Перший", "ШП", date(2026, 10, 3), date(2026, 10, 6), source_sheet="ШП"),
    ]

    results = amc.find_period_mismatches(timesheet, records)

    assert len(results) == 1
    assert results[0]["sources"] == ["book61: ППД", "book61: ШП", "book61: СЗЧ"]
    assert {d.day: expected for d, (_actual, expected) in results[0]["dates"].items()} == {
        1: "ППД", 2: "ППД", 3: "ШП", 4: "ШП", 5: "ШП", 6: "СЗЧ", 7: "СЗЧ",
    }


def test_find_period_mismatches_does_not_revive_an_old_open_status_after_newer_periods(tmp_path):
    """Реальний випадок, підтверджений користувачем: "Звіль" з Алф2 за
    2025 рік (відкритий) не може діяти в 2026, якщо після нього вже є
    пізніші періоди (ВЛК, що закінчився 03.10) - новіша подія закриває
    давню, і після кінця ВЛК "Звіль" НЕ повертається."""
    path = _write_roster(tmp_path / "ОБЛІК.xlsx", [
        ["Посада1", "сержант", "ПЕРШИЙ Перший Перший", "ВЛК", "ВЛК", "ВЛК", "РВЗ", "РВЗ", "РВЗ", "РВЗ"],
    ], header=_WEEK_HEADER)
    timesheet = Timesheet(path, "Табель")
    records = [
        _period("ПЕРШИЙ Перший Перший", "Звіль", date(2025, 11, 15), None, source_sheet="Алф2"),
        _period("ПЕРШИЙ Перший Перший", "ВЛК", date(2026, 9, 28), date(2026, 10, 3), source_sheet="ВЛК"),
    ]

    assert amc.find_period_mismatches(timesheet, records) == []


def test_find_period_mismatches_does_not_revive_an_open_szch_after_a_later_period(tmp_path):
    """Реальний випадок, підтверджений користувачем: СЗЧ з лютого (без дати
    повернення в book61) "повертався" після ВП, що закінчився 03.10, - і
    04.10 хибно очікувалось СЗЧ. Пізніша подія закриває СЗЧ назавжди."""
    path = _write_roster(tmp_path / "ОБЛІК.xlsx", [
        ["Посада1", "сержант", "ПЕРШИЙ Перший Перший", "ВП", "ВП", "ВП", "РВЗ", "РВЗ", "РВЗ", "РВЗ"],
    ], header=_WEEK_HEADER)
    timesheet = Timesheet(path, "Табель")
    records = [
        _period("ПЕРШИЙ Перший Перший", "СЗЧ", date(2026, 2, 28), None, source_sheet="СЗЧ"),
        _period("ПЕРШИЙ Перший Перший", "ВП", date(2026, 9, 17), date(2026, 10, 3), source_sheet="ВП"),
    ]

    assert amc.find_period_mismatches(timesheet, records) == []


def test_find_period_mismatches_ends_a_final_status_at_the_next_event_even_if_it_is_later_resumable(tmp_path):
    """"Розп"/"Звіль"/... - остаточні статуси: пізніший ВП закриває їх
    назавжди, навіть коли сам ВП короткий і після нього нічого немає."""
    path = _write_roster(tmp_path / "ОБЛІК.xlsx", [
        ["Посада1", "сержант", "ПЕРШИЙ Перший Перший", "Розп", "ВП", "ВП", "РВЗ", "РВЗ", "РВЗ", "РВЗ"],
    ], header=_WEEK_HEADER)
    timesheet = Timesheet(path, "Табель")
    records = [
        _period("ПЕРШИЙ Перший Перший", "Розп", date(2026, 9, 1), None, source_sheet="Алф2 (Розп)"),
        _period("ПЕРШИЙ Перший Перший", "ВП", date(2026, 10, 2), date(2026, 10, 3), source_sheet="ВП"),
    ]

    assert amc.find_period_mismatches(timesheet, records) == []


def test_find_period_mismatches_ends_an_open_period_at_the_next_event_except_ppd(tmp_path):
    """Реальний випадок (книга ЦВС): лікування без "Дата виписки" не
    "повертається" після пізнішої ВЛК до кінця місяця; відкритий ППД -
    повертається (людина знову на ППД після лікування)."""
    path = _write_roster(tmp_path / "ОБЛІК.xlsx", [
        ["Посада1", "сержант", "ПЕРШИЙ Перший Перший", "ШП", "ВЛК", "100", "100", "100", "100", "100"],
        ["Посада2", "матрос", "ДРУГИЙ Другий Другий", "ППД", "ШП", "ППД", "ППД", "ППД", "ППД", "ППД"],
    ], header=_WEEK_HEADER)
    timesheet = Timesheet(path, "Табель")
    records = [
        _period("ПЕРШИЙ Перший Перший", "ШП", date(2026, 9, 20), None, source_sheet="СОМАТИКА-ПОРАНЕНІ"),
        _period("ПЕРШИЙ Перший Перший", "ВЛК", date(2026, 10, 2), date(2026, 10, 2), source_sheet="ВЛК"),
        _period("ДРУГИЙ Другий Другий", "ППД", date(2026, 9, 1), None, source_sheet="ППД"),
        _period("ДРУГИЙ Другий Другий", "ШП", date(2026, 10, 2), date(2026, 10, 2), source_sheet="ШП"),
    ]

    assert amc.find_period_mismatches(timesheet, records) == []


def test_find_period_mismatches_ends_an_open_period_where_another_book_ends_the_same_status(tmp_path):
    """Реальний випадок, підтверджений користувачем: book61 "ВЛК" з 14.10
    без "Кінець (ВЛК)", ЦВС - ВЛК 14.10-20.10 (той самий день початку) -
    з 21.10 ВЛК не очікується. Так само, коли запис іншої книги почався
    раніше, але закінчився вже після початку відкритого періоду."""
    path = _write_roster(tmp_path / "ОБЛІК.xlsx", [
        ["Посада1", "сержант", "ПЕРШИЙ Перший Перший", "РВЗ", "ВЛК", "ВЛК", "ВЛК", "100", "100", "100"],
        ["Посада2", "матрос", "ДРУГИЙ Другий Другий", "ШП", "ШП", "ШП", "ШП", "100", "100", "100"],
    ], header=_WEEK_HEADER)
    timesheet = Timesheet(path, "Табель")
    records = [
        _period("ПЕРШИЙ Перший Перший", "ВЛК", date(2026, 10, 2), None, source_sheet="ВЛК"),
        _period("ПЕРШИЙ Перший Перший", "ВЛК", date(2026, 10, 2), date(2026, 10, 4), source_sheet="ВЛК", source_file="cvs"),
        _period("ДРУГИЙ Другий Другий", "ШП", date(2026, 10, 2), None, source_sheet="ШП"),
        _period("ДРУГИЙ Другий Другий", "ШП", date(2026, 10, 1), date(2026, 10, 4), source_sheet="СОМАТИКА-ПОРАНЕНІ",
                source_file="cvs"),
    ]

    assert amc.find_period_mismatches(timesheet, records) == []


def test_find_period_mismatches_resumes_a_regular_period_after_an_interruption(tmp_path):
    """Звичайний (не остаточний) період, перерваний коротшим пізнішим, -
    продовжується після нього до власного кінця за book61."""
    path = _write_roster(tmp_path / "ОБЛІК.xlsx", [
        ["Посада1", "сержант", "ПЕРШИЙ Перший Перший", "ВП", "ВП", "ВЛК", "ВЛК", "ВП", "РВЗ", "РВЗ"],
    ], header=_WEEK_HEADER)
    timesheet = Timesheet(path, "Табель")
    records = [
        _period("ПЕРШИЙ Перший Перший", "ВП", date(2026, 10, 1), date(2026, 10, 6), source_sheet="ВП"),
        _period("ПЕРШИЙ Перший Перший", "ВЛК", date(2026, 10, 3), date(2026, 10, 4), source_sheet="ВЛК"),
    ]

    results = amc.find_period_mismatches(timesheet, records)

    assert len(results) == 1
    assert {d.day: expected for d, (_actual, expected) in results[0]["dates"].items()} == {
        1: "ВП", 2: "ВП", 3: "ВЛК", 4: "ВЛК", 5: "ВП", 6: "ВП",
    }


def test_find_period_mismatches_accepts_the_whole_arrival_sequence_without_errors(tmp_path):
    """Рішення користувача ("Прибулі у ВД"): до зарахування в
    списки - "РРез", від зарахування - лише "ППД" (у Табелі - 10), у день
    прийому посади - будь-який статус (навіть коли паралельний аркуш "ППД"
    каже інакше), далі - знову аркуш "ППД". Жодної розбіжності."""
    path = _write_roster(tmp_path / "ОБЛІК.xlsx", [
        ["Посада1", "сержант", "ПЕРШИЙ Перший Перший", "Ррез", "Ррез", "Ррез", 10, 10, 30, 10],
    ], header=_WEEK_HEADER)
    timesheet = Timesheet(path, "Табель")
    records = [
        _period("ПЕРШИЙ Перший Перший", "ППД", date(2026, 9, 22), None, source_sheet="ППД"),
        _period("ПЕРШИЙ Перший Перший", "РРез", date(2026, 9, 22), date(2026, 10, 3), source_sheet="Прибулі у ВД"),
        _period("ПЕРШИЙ Перший Перший", "ППД", date(2026, 10, 4), date(2026, 10, 5), source_sheet="Прибулі у ВД"),
        _period("ПЕРШИЙ Перший Перший", amc.ANY_STATUS, date(2026, 10, 6), date(2026, 10, 6),
                source_sheet="Прибулі у ВД"),
    ]

    assert amc.find_period_mismatches(timesheet, records) == []


def test_find_period_mismatches_prefers_rvz_on_the_arrival_day(tmp_path):
    """Рішення користувача: у день прибуття з ШП/ВЛК/ВД
    пріоритетніший "РВЗ" (у Табелі для виплат - 30, 100): це НЕ
    розбіжність; сам статус аркуша в цей день - теж ні."""
    path = _write_roster(tmp_path / "ОБЛІК.xlsx", [
        ["Посада1", "сержант", "ПЕРШИЙ Перший Перший", "ШП", "ШП", 30, "РВЗ", "РВЗ", "РВЗ", "РВЗ"],
        ["Посада2", "матрос", "ДРУГИЙ Другий Другий", "ВЛК", "ВЛК", "Рвз", "РВЗ", "РВЗ", "РВЗ", "РВЗ"],
        ["Посада3", "матрос", "ТРЕТІЙ Третій Третій", "ВД", 100, "РВЗ", "РВЗ", "РВЗ", "РВЗ", "РВЗ"],
        ["Посада4", "матрос", "ЧЕТВЕРТИЙ Четвертий Четвертий", "ШП", "ШП", "ШП", "РВЗ", "РВЗ", "РВЗ", "РВЗ"],
    ], header=_WEEK_HEADER)
    timesheet = Timesheet(path, "Табель")
    records = [
        {**_period("ПЕРШИЙ Перший Перший", "ШП", date(2026, 9, 28), date(2026, 10, 3), source_sheet="ШП"),
         "arrival_date": date(2026, 10, 3)},
        {**_period("ДРУГИЙ Другий Другий", "ВЛК", date(2026, 9, 28), date(2026, 10, 3), source_sheet="ВЛК"),
         "arrival_date": date(2026, 10, 3)},
        {**_period("ТРЕТІЙ Третій Третій", "ВД", date(2026, 9, 28), date(2026, 10, 2), source_sheet="ВД"),
         "arrival_date": date(2026, 10, 2)},
        {**_period("ЧЕТВЕРТИЙ Четвертий Четвертий", "ШП", date(2026, 9, 28), date(2026, 10, 3), source_sheet="ШП"),
         "arrival_date": date(2026, 10, 3)},
    ]

    assert amc.find_period_mismatches(timesheet, records) == []


def test_find_period_mismatches_still_flags_another_status_on_the_arrival_day(tmp_path):
    """Не "РВЗ"/30/100 і не сам статус аркуша (напр. СЗЧ) у день прибуття -
    і далі розбіжність; так само "РВЗ" ДО дня прибуття."""
    path = _write_roster(tmp_path / "ОБЛІК.xlsx", [
        ["Посада1", "сержант", "ПЕРШИЙ Перший Перший", "ШП", 30, "СЗЧ", "РВЗ", "РВЗ", "РВЗ", "РВЗ"],
    ], header=_WEEK_HEADER)
    timesheet = Timesheet(path, "Табель")
    records = [{**_period("ПЕРШИЙ Перший Перший", "ШП", date(2026, 9, 28), date(2026, 10, 3), source_sheet="ШП"),
                "arrival_date": date(2026, 10, 3)}]

    results = amc.find_period_mismatches(timesheet, records)

    assert len(results) == 1
    flagged = {d.day for d, (actual, expected) in results[0]["dates"].items() if not amc._statuses_match(expected, actual)}
    assert flagged == {2, 3}


def test_find_period_mismatches_accepts_any_tabel_status_when_vd_says_not_departed(tmp_path):
    """book61 "ВД", "Виплата" - "не вибув": статус бере Табель, розбіжності
    немає, хоч би що там стояло."""
    path = _write_roster(tmp_path / "ОБЛІК.xlsx", [
        ["Посада1", "сержант", "ПЕРШИЙ Перший Перший", "РВЗ", 30, 100, "ППД", "РВЗ", "РВЗ", "РВЗ"],
    ], header=_WEEK_HEADER)
    timesheet = Timesheet(path, "Табель")
    records = [_period("ПЕРШИЙ Перший Перший", amc.ANY_STATUS, date(2026, 10, 1), date(2026, 10, 4), source_sheet="ВД")]

    assert amc.find_period_mismatches(timesheet, records) == []


def test_statuses_match_accepts_anything_for_the_any_status():
    assert amc._statuses_match(amc.ANY_STATUS, "30") is True
    assert amc._statuses_match(amc.ANY_STATUS, 100) is True


def test_open_arrival_period_without_a_position_is_checked_and_ends_at_the_next_event(tmp_path):
    """Рішення користувача (реальний випадок): ТВО, зарахування, посаду ще
    не прийнято - "РРез" до зарахування, далі відкритий "ППД"; "РРез" у
    Табелі в день зарахування - розбіжність. Відкритий "ППД" накладається
    поверх решти, тож закінчується перед пізнішою подією (відпустка)."""
    path = _write_roster(tmp_path / "ОБЛІК.xlsx", [
        ["Посада1", "матрос", "ПЕРШИЙ Перший Перший", "РРез", "РРез", "РРез", "РРез", "РВЗ", "РВЗ", None],
        ["Посада2", "матрос", "ДРУГИЙ Другий Другий", "ППД", "ППД", "ППД", "ВПС", "ППД", "ППД", "ППД"],
    ], header=_WEEK_HEADER)
    records = [
        _period("ПЕРШИЙ Перший Перший", "РРез", date(2026, 10, 1), date(2026, 10, 3), source_sheet="Прибулі у ВД"),
        _period("ПЕРШИЙ Перший Перший", "ППД", date(2026, 10, 4), None, source_sheet="Прибулі у ВД"),
        _period("ДРУГИЙ Другий Другий", "ППД", date(2026, 10, 1), None, source_sheet="Прибулі у ВД"),
        _period("ДРУГИЙ Другий Другий", "ППД", date(2026, 10, 1), date(2026, 10, 3), source_sheet="ППД"),
        _period("ДРУГИЙ Другий Другий", "ВПС", date(2026, 10, 4), date(2026, 10, 4), source_sheet="ВП"),
        _period("ДРУГИЙ Другий Другий", "ППД", date(2026, 10, 5), None, source_sheet="ППД"),
    ]

    results = amc.find_period_mismatches(Timesheet(path, "Табель"), records)

    assert [r["roster_pib_raw"] for r in results] == ["ПЕРШИЙ Перший Перший"]
    flagged = {d.day: (actual, expected) for d, (actual, expected) in results[0]["dates"].items()
               if not amc._statuses_match(expected, actual)}
    assert flagged == {4: ("РРез", "ППД"), 5: ("РВЗ", "ППД"), 6: ("РВЗ", "ППД")}


def test_find_period_mismatches_gives_rrez_priority_and_resumes_the_parallel_period(tmp_path):
    """Рішення користувача: "РРез" (ТВО, не оплачується)
    переважає паралельний "ППД" на весь свій проміжок, навіть якщо ППД
    почався пізніше; ППД продовжується лише після кінця ТВО."""
    path = _write_roster(tmp_path / "ОБЛІК.xlsx", [
        ["Посада1", "сержант", "ПЕРШИЙ Перший Перший", "РВЗ", "Ррез", "Ррез", "Ррез", "ППД", "ППД", "РВЗ"],
    ], header=_WEEK_HEADER)
    timesheet = Timesheet(path, "Табель")
    records = [
        _period("ПЕРШИЙ Перший Перший", "РРез", date(2026, 10, 2), date(2026, 10, 4), source_sheet="Прибулі у ВД"),
        _period("ПЕРШИЙ Перший Перший", "ППД", date(2026, 10, 3), None, source_sheet="ППД"),
    ]

    results = amc.find_period_mismatches(timesheet, records)

    assert len(results) == 1
    assert results[0]["sources"] == ["book61: Прибулі у ВД", "book61: ППД"]
    assert {d.day: expected for d, (_actual, expected) in results[0]["dates"].items()} == {
        2: "РРез", 3: "РРез", 4: "РРез", 5: "ППД", 6: "ППД", 7: "ППД",
    }


def test_find_period_mismatches_never_duplicates_a_person_within_a_file(tmp_path):
    """Рішення користувача: ПІБ у mis_match.xlsx не
    дублюються, а звіт розділено за addon-файлами: записи однієї людини з
    КІЛЬКОХ аркушів і з різним написанням ПІБ (подвійний пробіл, регістр) -
    ОДИН запис на кожен файл, кожен зі своїм ланцюжком аркушів і лише своїми
    датами."""
    path = _write_roster(tmp_path / "ОБЛІК.xlsx", [
        ["Посада1", "сержант", "ПЕРШИЙ Перший Перший", "ППД", "ППД", "ШП", "ВП", "СЗЧ", "СЗЧ", "РВЗ"],
    ], header=_WEEK_HEADER)
    timesheet = Timesheet(path, "Табель")
    records = [
        _period("ПЕРШИЙ Перший Перший", "ППД", date(2026, 10, 1), date(2026, 10, 2), source_sheet="ППД"),
        _period("ПЕРШИЙ  Перший Перший", "ШП", date(2026, 10, 3), date(2026, 10, 4), source_sheet="ШП"),
        _period("перший перший перший", "СЗЧ", date(2026, 10, 5), None, source_sheet="СЗЧ", source_file="addon2"),
    ]

    results = amc.find_period_mismatches(timesheet, records)

    assert [(r["source_file"], r["sources"], sorted(d.day for d in r["dates"])) for r in results] == [
        ("book61", ["book61: ППД", "book61: ШП"], [1, 2, 3, 4]),
        ("addon2", ["addon2: СЗЧ"], [5, 6, 7]),
    ]


def test_check_addon_mismatches_file_lists_every_person_only_once(tmp_path):
    """Наскрізна перевірка самого mis_match.xlsx (рішення користувача: ПІБ не мають дублюватись у файлі): людина з
    ППД/ШП/СЗЧ-періодами, що перекриваються, й іще одним ШП-записом з іншим
    написанням ПІБ - ОДИН рядок "book61: ППД -> book61: ШП -> book61: СЗЧ";
    людина з двома рядками Алф2 "самостріл" - ОДИН раз у "Застереження".
    В обох аркушах - ПІДРОЗДІЛ людини з Табеля."""
    oblik = _write_roster(tmp_path / "ОБЛІК.xlsx", [
        ["Підрозділ1", "Посада1", "сержант", "ПЕРШИЙ Перший Перший", "ППД", "ППД", "ШП", "ШП", "ШП", "СЗЧ", "РВЗ"],
        ["Підрозділ2", "Посада2", "матрос", "ДРУГИЙ Другий Другий", "РВЗ", "ВП", "ВП", "РВЗ", "РВЗ", "РВЗ", "РВЗ"],
    ], header=["ПІДРОЗДІЛ"] + _WEEK_HEADER)
    book61 = write_full_book61(tmp_path / "book61.xlsm", {
        "ППД": [[1, "Підрозділ1", "сержант", "ПЕРШИЙ Перший Перший", "ППД", date(2026, 8, 13), "у лік з 04.10.2026"]],
        "ШП": [
            [1, "Підрозділ1", "сержант", "ПЕРШИЙ Перший Перший", "лік", date(2026, 10, 3), date(2026, 10, 6)],
            [2, "Підрозділ1", "сержант", "ПЕРШИЙ  Перший Перший", "лік", date(2026, 9, 1), date(2026, 9, 5)],
        ],
        "СЗЧ": [[1, "Підрозділ1", "сержант", "ПЕРШИЙ Перший Перший", "СЗЧ", date(2026, 10, 6)]],
        "ВП": [[1, "Підрозділ1", "матрос", "ДРУГИЙ Другий Другий", "від", date(2026, 10, 2), None, date(2026, 10, 5),
                "Перша частина щорічної відпустки"]],
        "Алф2": [
            [1, "ДРУГИЙ Другий Другий", "самостріл", None, None, None],
            [2, "ДРУГИЙ Другий Другий", "самостріл", None, None, None],
        ],
    })
    report = str(tmp_path / "mis_match.xlsx")

    result = check_addon_mismatches_file(oblik, [book61], report_file_name=report)

    assert result.report_path == report
    assert result.checked_files == [book61] and result.unrecognized_files == []
    wb = load_workbook(report)
    assert wb.sheetnames == ["book61", "Застереження"]
    ws = wb["book61"]
    assert [ws.cell(row=1, column=c).value for c in (1, 2, 3)] == ["Джерело", "Підрозділ", "ПІБ"]
    pib_rows = {
        normalize_name(ws.cell(row=r, column=3).value): r
        for r in range(2, ws.max_row + 1) if ws.cell(row=r, column=3).value not in (None, "—")
    }
    pib_counts = Counter(
        normalize_name(ws.cell(row=r, column=3).value)
        for r in range(2, ws.max_row + 1) if ws.cell(row=r, column=3).value not in (None, "—")
    )
    assert pib_counts == Counter({"ПЕРШИЙ ПЕРШИЙ ПЕРШИЙ": 1, "ДРУГИЙ ДРУГИЙ ДРУГИЙ": 1})
    first_row = pib_rows["ПЕРШИЙ ПЕРШИЙ ПЕРШИЙ"]
    assert ws.cell(row=first_row + 1, column=1).value == "book61: ППД -> book61: ШП -> book61: СЗЧ"
    assert ws.cell(row=first_row, column=2).value == "Підрозділ1"
    assert ws.cell(row=pib_rows["ДРУГИЙ ДРУГИЙ ДРУГИЙ"], column=2).value == "Підрозділ2"

    warn_ws = wb["Застереження"]
    warn_rows = [
        (warn_ws.cell(row=r, column=1).value, warn_ws.cell(row=r, column=2).value)
        for r in range(2, warn_ws.max_row + 1)
    ]
    assert warn_rows == [("Підрозділ2", "ДРУГИЙ Другий Другий")]


def test_find_period_mismatches_merges_multiple_records_for_the_same_person_and_sheet(tmp_path):
    path = _write_roster(tmp_path / "ОБЛІК.xlsx", [
        ["Посада1", "сержант", "ПЕРШИЙ Перший Перший", "РВЗ", "ВП", "РВЗ"],
    ])
    timesheet = Timesheet(path, "Табель")
    records = [
        _period("ПЕРШИЙ Перший Перший", "ВП", date(2026, 7, 1), date(2026, 7, 1)),
        _period("ПЕРШИЙ Перший Перший", "ВП", date(2026, 7, 3), date(2026, 7, 3)),
    ]

    results = amc.find_period_mismatches(timesheet, records)

    assert len(results) == 1
    assert results[0]["dates"] == {
        date(2026, 7, 1): ("РВЗ", "ВП"),
        date(2026, 7, 3): ("РВЗ", "ВП"),
    }


def test_find_period_mismatches_treats_bpshp_vpbp_as_a_non_mismatch(tmp_path):
    path = _write_roster(tmp_path / "ОБЛІК.xlsx", [
        ["Посада1", "сержант", "ПЕРШИЙ Перший Перший", "ВПБП", "РВЗ", "РВЗ"],
    ])
    timesheet = Timesheet(path, "Табель")
    records = [_period("ПЕРШИЙ Перший Перший", "БПШП", date(2026, 7, 1), date(2026, 7, 1), source_sheet="ШП")]

    assert amc.find_period_mismatches(timesheet, records) == []


def test_find_period_mismatches_treats_payment_converted_value_as_a_non_mismatch(tmp_path):
    """Реальний випадок: output/ОБЛІК.xlsx вже "для виплат" - клітинка
    каже "100_БПШП" (число 10 - аналогічно для "ВД") замість буквального
    book61-статусу "БПШП"."""
    path = _write_roster(tmp_path / "ОБЛІК.xlsx", [
        ["Посада1", "сержант", "ПЕРШИЙ Перший Перший", "100_БПШП", "РВЗ", "РВЗ"],
    ])
    timesheet = Timesheet(path, "Табель")
    records = [_period("ПЕРШИЙ Перший Перший", "БПШП", date(2026, 7, 1), date(2026, 7, 1), source_sheet="ШП")]

    assert amc.find_period_mismatches(timesheet, records) == []


def test_find_period_mismatches_still_flags_10_vs_navch_as_a_genuine_mismatch(tmp_path):
    """Застереження користувача: "10" еквівалентний "ВД", АЛЕ якщо book61
    очікує саме "Навч" - число 10 в Табелі МАЄ лишатись розбіжністю."""
    path = _write_roster(tmp_path / "ОБЛІК.xlsx", [
        ["Посада1", "сержант", "ПЕРШИЙ Перший Перший", 10, "РВЗ", "РВЗ"],
    ])
    timesheet = Timesheet(path, "Табель")
    records = [_period("ПЕРШИЙ Перший Перший", "Навч", date(2026, 7, 1), date(2026, 7, 1), source_sheet="ВД")]

    results = amc.find_period_mismatches(timesheet, records)

    assert len(results) == 1
    assert results[0]["dates"] == {date(2026, 7, 1): (10, "Навч")}


def test_find_period_mismatches_without_any_date_columns_overlapping_is_empty(tmp_path):
    path = _write_roster(tmp_path / "ОБЛІК.xlsx", [
        ["Посада1", "сержант", "ПЕРШИЙ Перший Перший", "РВЗ", "РВЗ", "РВЗ"],
    ])
    timesheet = Timesheet(path, "Табель")
    records = [_period("ПЕРШИЙ Перший Перший", "ВП", date(2020, 1, 1), date(2020, 1, 2))]

    assert amc.find_period_mismatches(timesheet, records) == []


# --- unmapped_people_in_roster -------------------------------------------------

def test_unmapped_people_in_roster_keeps_only_people_present_in_the_tabel(tmp_path):
    """Те саме правило, що й для періодів: кого немає в Табелі, той не
    потрапляє в mis_match.xlsx навіть як застереження."""
    path = _write_roster(tmp_path / "ОБЛІК.xlsx", [
        ["Посада1", "сержант", "ПЕРШИЙ Перший Перший", "РВЗ", "РВЗ", "РВЗ"],
    ])
    timesheet = Timesheet(path, "Табель")
    unmapped_people = [
        {"pib_raw": "ПЕРШИЙ Перший Перший", "alf2_status": "самостріл"},
        {"pib_raw": "ВІДСУТНІЙ Відсутній Відсутній", "alf2_status": "самостріл"},
    ]

    kept = amc.unmapped_people_in_roster(timesheet, unmapped_people)

    assert [p["pib_raw"] for p in kept] == ["ПЕРШИЙ Перший Перший"]


def test_unmapped_people_in_roster_lists_each_person_once():
    """Алф2 може мати кілька рядків з тим самим ПІБ - у "Застереження"
    людина МАЄ бути один раз."""
    class _Roster:
        label_columns = {}

        def person_rows(self):
            return {"ПЕРШИЙ ПЕРШИЙ ПЕРШИЙ": 2}

    unmapped_people = [
        {"pib_raw": "ПЕРШИЙ Перший Перший", "alf2_status": "самостріл"},
        {"pib_raw": "ПЕРШИЙ  Перший Перший", "alf2_status": "самостріл"},
    ]

    assert len(amc.unmapped_people_in_roster(_Roster(), unmapped_people)) == 1


# --- write_addon_mismatch_report ----------------------------------------------

def test_write_addon_mismatch_report_returns_none_when_nothing_to_report(tmp_path):
    assert amc.write_addon_mismatch_report([], [], str(tmp_path / "mis_match.xlsx")) is None


def test_write_addon_mismatch_report_writes_period_mismatches_and_warnings(tmp_path):
    period_mismatches = [{
        "source_file": "book61",
        "sources": ["book61: ВП"], "roster_pidrozdil": "Підрозділ1",
        "roster_pib_raw": "ПЕРШИЙ Перший Перший", "addon_pib_raw": "ПЕРШИЙ Перший Перший",
        "dates": {date(2026, 7, 1): ("РВЗ", "ВП")},
    }]
    unmapped_people = [{"pib_raw": "ТРЕТІЙ Третій Третій", "alf2_status": "самостріл", "roster_pidrozdil": "Підрозділ3"}]
    output_path = str(tmp_path / "mis_match.xlsx")

    result = amc.write_addon_mismatch_report(period_mismatches, unmapped_people, output_path)

    assert result == output_path
    wb = load_workbook(output_path)
    assert wb.sheetnames == ["book61", "Застереження"]
    ws = wb["book61"]
    assert [ws.cell(row=1, column=c).value for c in (1, 2, 3)] == ["Джерело", "Підрозділ", "ПІБ"]
    assert ws.cell(row=2, column=1).value == "Табель (ОБЛІК.xlsx)"
    assert ws.cell(row=3, column=1).value == "book61: ВП"
    assert ws.cell(row=2, column=2).value == "Підрозділ1"
    assert "B2:B3" in {str(merged) for merged in ws.merged_cells.ranges}
    assert ws.cell(row=2, column=3).value == "ПЕРШИЙ Перший Перший"

    warn_ws = wb["Застереження"]
    assert [warn_ws.cell(row=1, column=c).value for c in (1, 2)] == ["Підрозділ", "ПІБ (book61)"]
    assert (warn_ws.cell(row=2, column=1).value, warn_ws.cell(row=2, column=2).value) == (
        "Підрозділ3", "ТРЕТІЙ Третій Третій",
    )


def test_write_addon_mismatch_report_shows_a_placeholder_when_the_tabel_has_no_pidrozdil(tmp_path):
    period_mismatches = [{
        "source_file": "book61",
        "sources": ["book61: ВП"], "roster_pidrozdil": None,
        "roster_pib_raw": "ПЕРШИЙ Перший Перший", "addon_pib_raw": "ПЕРШИЙ Перший Перший",
        "dates": {date(2026, 7, 1): ("РВЗ", "ВП")},
    }]
    output_path = str(tmp_path / "mis_match.xlsx")

    amc.write_addon_mismatch_report(period_mismatches, [], output_path)

    assert load_workbook(output_path)["book61"].cell(row=2, column=2).value == "—"


def test_write_addon_mismatch_report_writes_only_warnings_when_no_period_mismatches(tmp_path):
    unmapped_people = [{"pib_raw": "ДРУГИЙ Другий Другий", "alf2_status": "самостріл", "roster_pidrozdil": None}]
    output_path = str(tmp_path / "mis_match.xlsx")

    result = amc.write_addon_mismatch_report([], unmapped_people, output_path)

    assert result == output_path
    wb = load_workbook(output_path)
    assert "Застереження" in wb.sheetnames


def test_write_addon_mismatch_report_renders_the_sheet_chain_as_the_source_label(tmp_path):
    """Рішення користувача: колонка "Джерело" показує ЛАНЦЮЖОК
    аркушів book61 у тому порядку, в якому вони змінюють одне одного на
    шкалі людини: "book61: ППД -> book61: ШП -> book61: СЗЧ"."""
    period_mismatches = [{
        "source_file": "book61",
        "sources": ["book61: ППД", "book61: ШП", "book61: СЗЧ"], "roster_pidrozdil": "Підрозділ1",
        "roster_pib_raw": "ПЕРШИЙ Перший Перший", "addon_pib_raw": "ПЕРШИЙ Перший Перший",
        "dates": {date(2026, 10, 8): ("РВЗ", "СЗЧ")},
    }]
    output_path = str(tmp_path / "mis_match.xlsx")

    amc.write_addon_mismatch_report(period_mismatches, [], output_path)

    wb = load_workbook(output_path)
    ws = wb["book61"]
    assert ws.cell(row=3, column=1).value == "book61: ППД -> book61: ШП -> book61: СЗЧ"


def test_addon_files_lists_excel_books_without_lock_files(tmp_path):
    folder = tmp_path / "addons"
    folder.mkdir()
    for name in ("b.xlsx", "a.xlsm", "~$b.xlsx", "notes.txt"):
        (folder / name).write_text("x")
    assert [os.path.basename(path) for path in addon_files(str(folder))] == ["a.xlsm", "b.xlsx"]


def test_check_addon_mismatches_file_puts_each_addon_file_on_its_own_sheet(tmp_path):
    """Рішення користувача: mis_match.xlsx розділено на аркуші з
    назвами addon-файлів: шкала людини одна (book61 + ЦВС), але кожна дата -
    на аркуші файлу, чий запис її визначив, ланцюжок "Джерело" - лише цього
    файлу; звірений файл без розбіжностей - аркуш із написом. День виписки
    (ЦВС, аркуш лікування) з РВЗ у Табелі - не розбіжність. Книга, тип якої не
    розпізнано (чи це взагалі не Excel), пропускається й повертається окремо."""
    oblik = _write_roster(tmp_path / "ОБЛІК.xlsx", [
        ["Підрозділ1", "Посада1", "сержант", "ПЕРШИЙ Перший Перший", "ППД", "РВЗ", "ШП", "ВП", "РВЗ", "РВЗ", "РВЗ"],
        ["Підрозділ2", "Посада2", "матрос", "ДРУГИЙ Другий Другий", "РВЗ", "РВЗ", "РВЗ", "РВЗ", "РВЗ", "РВЗ", "РВЗ"],
    ], header=["ПІДРОЗДІЛ"] + _WEEK_HEADER)
    book61 = write_full_book61(tmp_path / "book61.xlsm", {
        "ППД": [[1, "Підрозділ1", "сержант", "ПЕРШИЙ Перший Перший", "ППД", date(2026, 8, 13), "у лік з 03.10.2026"]],
    })
    cvs = write_cvs(tmp_path / "cvs.xlsx", {
        MEDICAL_SHEET: [[None, 1, "Підрозділ1", "сержант", "ПЕРШИЙ Перший Перший", "соматика", datetime(2026, 10, 3),
                         "Лікарня", "№1", datetime(2026, 10, 5)]],
    })
    cvs2 = write_cvs(tmp_path / "cvs2.xlsx", {
        VLK_SHEET: [[1, "ДРУГИЙ Другий Другий", "матрос", None, datetime(2026, 10, 2), None, "придатний"]],
    })
    junk = tmp_path / "junk.xlsx"
    junk_wb = Workbook()
    junk_wb.worksheets[0].append(["ПІБ", "Щось інше"])
    junk_wb.save(str(junk))
    broken = tmp_path / "broken.xlsx"
    broken.write_text("не книга")
    report = str(tmp_path / "mis_match.xlsx")

    result = check_addon_mismatches_file(oblik, [book61, cvs, cvs2, str(junk), str(broken)], report_file_name=report)

    assert result.checked_files == [book61, cvs, cvs2]
    assert result.unrecognized_files == [str(junk), str(broken)]
    # 04.10 "ВП" у Табелі - ще й на аркуші book61: у book61 немає запису "ВП"
    # (ADDON_REQUIRED_BOOK61_SHEETS).
    assert [(r["source_file"], r["roster_pib_raw"], r["sources"]) for r in result.period_mismatches] == [
        ("book61", "ПЕРШИЙ Перший Перший", ["book61: ППД", "book61: ВП"]),
        ("cvs", "ПЕРШИЙ Перший Перший", ["cvs: СОМАТИКА-ПОРАНЕНІ"]),
    ]
    bad_days = [
        [d.day for d, (actual, expected) in record["dates"].items() if not amc._statuses_match(expected, actual)]
        for record in result.period_mismatches
    ]
    assert bad_days == [[2, 4], [4]]

    wb = load_workbook(report)
    assert wb.sheetnames == ["book61", "cvs", "cvs2"]
    assert wb["book61"].cell(row=3, column=1).value == "book61: ППД -> book61: ВП"
    assert wb["cvs"].cell(row=3, column=1).value == "cvs: СОМАТИКА-ПОРАНЕНІ"
    header_dates = [wb["cvs"].cell(row=1, column=c).value for c in range(4, wb["cvs"].max_column + 1)]
    assert [value.day for value in header_dates if isinstance(value, datetime)] == [3, 4, 5]
    assert len(header_dates) == 3
    assert wb["cvs2"].cell(row=2, column=1).value == "Розбіжностей з цим файлом немає."


def test_sheet_titles_follow_excel_rules_and_stay_unique():
    used = set()
    assert amc._sheet_title("a/b:c*?[x]\\y", used) == "a_b_c___x__y"
    assert len(amc._sheet_title("Б" * 40, used)) == 31
    assert amc._sheet_title("book61", used) == "book61"
    assert amc._sheet_title("BOOK61", used) == "BOOK61 (2)"


def test_non_combat_treatment_accepts_any_of_its_statuses():
    """Рішення користувача: не бойове лікування з книги ЦВС - у
    Табелі ШП, ЛМР або ВПСЗ (і їхні форми в ОБЛІК для виплат, напр. 10_МЕД);
    бойове (БПШП) чи РВЗ - розбіжність; в обидва боки (звірка книг між собою)."""
    from constants import ADDON_NON_COMBAT_TREATMENT_STATUS as non_combat
    for actual in ("ШП", "ЛМР", "ВПСЗ"):
        assert amc._statuses_match(non_combat, actual) and amc._statuses_match(actual, non_combat)
    assert amc._statuses_match(non_combat, "10_МЕД")
    for actual in ("БПШП", "РВЗ", 100):
        assert not amc._statuses_match(non_combat, actual)
    assert not amc._statuses_match("БПШП", non_combat)


def test_open_non_combat_treatment_takes_the_end_from_another_books_hospital_stay():
    """Відкритий запис ЦВС "ШП/ЛМР/ВПСЗ" бере кінець того самого епізоду з
    іншої книги (book61 "ШП"), як раніше відкритий "ШП"."""
    from constants import ADDON_NON_COMBAT_TREATMENT_STATUS as non_combat
    cvs = {"pib_raw": "ПЕРШИЙ Перший Перший", "status": non_combat, "start_date": date(2026, 10, 10),
           "end_date": None, "source_sheet": "СОМАТИКА", "arrival_date": None}
    book61 = {"pib_raw": "ПЕРШИЙ Перший Перший", "status": "ШП", "start_date": date(2026, 10, 10),
              "end_date": date(2026, 10, 15), "source_sheet": "ШП", "arrival_date": date(2026, 10, 15)}
    assert amc._known_end_of_open_period(cvs, [cvs, book61]) == date(2026, 10, 15)


def test_leave_in_tabel_must_be_covered_by_a_book61_vp_record(tmp_path):
    """Рішення користувача (реальний випадок): у Табелі ВПСЗ до
    кінця тижня, а запис "ВП" book61 закінчився раніше - дні після нього -
    розбіжність на аркуші book61 ("немає запису на аркуші ВП"), хоча жодна
    книга в ці дні нічого не очікує. Людина з відпусткою в Табелі й без
    жодного запису в addons - теж. Без book61 серед книг - не перевіряється."""
    oblik = _write_roster(tmp_path / "ОБЛІК.xlsx", [
        ["Посада1", "сержант", "ПЕРШИЙ Перший Перший", "ВПСЗ", "ВПСЗ", "ВПСЗ", "ВПСЗ", "РВЗ", "РВЗ", "РВЗ"],
        ["Посада2", "матрос", "ДРУГИЙ Другий Другий", "РВЗ", "ВП", "ВП", "РВЗ", "РВЗ", "РВЗ", "РВЗ"],
        ["Посада3", "матрос", "ТРЕТІЙ Третій Третій", "ВП", "ВП", "РВЗ", "РВЗ", "РВЗ", "РВЗ", "РВЗ"],
    ], header=_WEEK_HEADER)
    timesheet = Timesheet(oblik, "Табель")
    records = [
        _period("ПЕРШИЙ Перший Перший", "ВПСЗ", date(2026, 9, 13), date(2026, 10, 2)),
        _period("ТРЕТІЙ Третій Третій", "ВП", date(2026, 9, 25), date(2026, 10, 2)),
    ]

    results = amc.find_period_mismatches(timesheet, records, files_by_kind={"book61": "book61"})

    no_record = "немає запису на аркуші ВП"
    assert {r["roster_pib_raw"]: {d.day: pair for d, pair in r["dates"].items() if pair[1] == no_record}
            for r in results} == {
        "ПЕРШИЙ Перший Перший": {3: ("ВПСЗ", no_record), 4: ("ВПСЗ", no_record)},
        "ДРУГИЙ Другий Другий": {2: ("ВП", no_record), 3: ("ВП", no_record)},
    }
    assert all(r["source_file"] == "book61" and "book61: ВП" in r["sources"] for r in results)
    assert amc.find_period_mismatches(timesheet, records) == []  # без book61 - не перевіряється


def test_tracked_statuses_without_any_addon_record_are_shown(tmp_path):
    """Рішення користувача: "теж їх відображати, враховуючи
    попередні правила": статус Табеля, який веде книга addons, у день, коли
    ЖОДНА книга нічого не каже, - розбіжність "немає запису на аркуші ..." на
    аркуші відповідної книги (ШП - book61 "ШП", Адап - журнал; форма для
    виплат 100_БЗ - book61 "Алф2"); запис будь-якої книги цього дня -
    звичайна звірка (ЦВС "ШП/ЛМР/ВПСЗ" і ШП у Табелі - збіг); статус, якого не
    веде жодна книга (РТГр), і книга, якої немає серед звірених, - не
    перевіряються."""
    from constants import ADDON_NON_COMBAT_TREATMENT_STATUS
    oblik = _write_roster(tmp_path / "ОБЛІК.xlsx", [
        ["Посада1", "сержант", "ПЕРШИЙ Перший Перший", "ШП", "ШП", "ШП", "РВЗ", "РВЗ", "РВЗ", "РВЗ"],
        ["Посада2", "матрос", "ДРУГИЙ Другий Другий", "Адап", "Адап", "РВЗ", "РВЗ", "РВЗ", "РВЗ", "РВЗ"],
        ["Посада3", "матрос", "ТРЕТІЙ Третій Третій", "100_БЗ", "100_БЗ", "100_БЗ", "100_БЗ", "100_БЗ", "100_БЗ", "100_БЗ"],
        ["Посада4", "матрос", "ЧЕТВЕРТИЙ Четвертий Четвертий", "РТГр", "РТГр", "РВЗ", "РВЗ", "РВЗ", "РВЗ", "РВЗ"],
    ], header=_WEEK_HEADER)
    timesheet = Timesheet(oblik, "Табель")
    records = [
        # ПЕРШИЙ: 01-02 - ЦВС (не бойове) - збіг; 03 - жодного запису.
        _period("ПЕРШИЙ Перший Перший", ADDON_NON_COMBAT_TREATMENT_STATUS, date(2026, 10, 1), date(2026, 10, 2),
                "СОМАТИКА", "cvs"),
    ]
    no_record = lambda sheet: f"немає запису на аркуші {sheet}"

    results = amc.find_period_mismatches(timesheet, records, files_by_kind={"book61": "book61", "journal": "journal"})

    assert sorted((r["source_file"], r["roster_pib_raw"], sorted((d.day, e) for d, (a, e) in r["dates"].items()))
                  for r in results) == [
        ("book61", "ПЕРШИЙ Перший Перший", [(3, no_record("ШП"))]),
        ("book61", "ТРЕТІЙ Третій Третій", [(day, no_record("Алф2")) for day in range(1, 8)]),
        ("journal", "ДРУГИЙ Другий Другий", [(1, no_record("Адаптація")), (2, no_record("Адаптація"))]),
    ]
    only_book61 = amc.find_period_mismatches(timesheet, records, files_by_kind={"book61": "book61"})
    assert "ДРУГИЙ Другий Другий" not in {r["roster_pib_raw"] for r in only_book61}  # журналу немає


def test_killed_on_alf2_allows_bz_until_the_death_certificate(tmp_path):
    """Правило книги ЦВС (рішення користувача): "Загиб" - лише з
    дати лікарського свідоцтва. Реальний випадок: Алф2 - "загибель" з 2024
    (без свідоцтва), ЦВС - Загиб з 05.10 (свідоцтво), у Табелі БЗ до 04.10,
    далі Загиб - не розбіжність. Число в Табелі (10 - і ВД, і ППД...) у
    зворотну перевірку не йде."""
    from constants import ADDON_KILLED_WITHOUT_CERTIFICATE_STATUS
    oblik = _write_roster(tmp_path / "ОБЛІК.xlsx", [
        ["Посада1", "сержант", "ПЕРШИЙ Перший Перший", "БЗ", "БЗ", "БЗ", "БЗ", "Загиб", "Загиб", "Загиб"],
        ["Посада2", "матрос", "ДРУГИЙ Другий Другий", 10, 10, "РВЗ", "РВЗ", "РВЗ", "РВЗ", "РВЗ"],
    ], header=_WEEK_HEADER)
    records = [
        _period("ПЕРШИЙ Перший Перший", ADDON_KILLED_WITHOUT_CERTIFICATE_STATUS, date(2024, 8, 6), None, "Алф2"),
        _period("ПЕРШИЙ Перший Перший", "Загиб", date(2026, 10, 5), None, "Загиблі-200", "cvs"),
    ]

    assert amc.find_period_mismatches(Timesheet(oblik, "Табель"), records, files_by_kind={"book61": "book61"}) == []


def test_report_has_no_gaps_dates_compared_elsewhere_show_the_tabel_value_in_grey(tmp_path):
    """Рішення користувача (реальний випадок: ВПСЗ | порожньо |
    ВПСЗ): дата, яку визначила інша книга (її аркуш) чи жодна, - не порожні
    клітинки, а значення Табеля в сірій об'єднаній клітинці."""
    d1, d2, d3 = date(2026, 10, 8), date(2026, 10, 9), date(2026, 10, 10)
    tabel = {d1: "ВПСЗ", d2: "10_МЕД", d3: "ВПСЗ"}
    records = [
        {"source_file": "book61", "sources": ["book61: ВП"], "roster_pidrozdil": "Підрозділ1",
         "roster_pib_raw": "ПЕРШИЙ Перший Перший", "addon_pib_raw": "ПЕРШИЙ Перший Перший",
         "dates": {d1: ("ВПСЗ", "ВПСЗ"), d3: ("ВПСЗ", "ВП")}, "tabel_values": tabel},
        {"source_file": "book61", "sources": ["book61: ВП"], "roster_pidrozdil": "Підрозділ1",
         "roster_pib_raw": "ДРУГИЙ Другий Другий", "addon_pib_raw": "ДРУГИЙ Другий Другий",
         "dates": {d2: ("РВЗ", "ВП")}, "tabel_values": {d1: None, d2: "РВЗ", d3: "РВЗ"}},
    ]
    path = str(tmp_path / "mis_match.xlsx")

    amc.write_addon_mismatch_report(records, [], path, source_files=["book61"])

    ws = load_workbook(path)["book61"]
    assert [ws.cell(row=1, column=c).value for c in (4, 5, 6)] == [datetime(2026, 10, d) for d in (8, 9, 10)]
    gap = ws.cell(row=2, column=5)
    assert (gap.value, gap.fill.fgColor.rgb) == ("10_МЕД", "FFF2F2F2")
    merged = {str(rng) for rng in ws.merged_cells.ranges}
    assert "E2:E3" in merged
    assert ws.cell(row=4, column=4).value == "—" and "D4:D5" in merged  # Табель порожній
    assert all(ws.cell(row=row, column=col).value is not None for row in (2, 4) for col in (4, 5, 6))


def _episode(status, start, end, source_file, **extra):
    return {"pib_raw": "ПЕРШИЙ Перший Перший", "status": status, "start_date": start, "end_date": end,
            "source_sheet": "аркуш", "source_file": source_file, **extra}


def test_open_period_takes_the_end_of_the_same_episode_only():
    """Рішення користувача ("продовження статусів"): відкритий
    період бере кінець того самого епізоду з іншої книги (закінчився пізніше
    чи того самого дня з документом), але НЕ попереднього епізоду, що
    закінчився в день його початку (переведення з лікарні до лікарні 16.10), і
    НЕ одноденного запису без документа (ЦВС ВЛК лише з датою проходження)."""
    open_bpshp = _episode("БПШП", date(2026, 10, 16), None, "book61")
    previous = _episode("БПШП", date(2026, 9, 24), date(2026, 10, 16), "cvs")
    assert amc._known_end_of_open_period(open_bpshp, [open_bpshp, previous]) is None

    open_vlk = _episode("ВЛК", date(2026, 10, 13), None, "book61")
    one_day = _episode("ВЛК", date(2026, 10, 13), date(2026, 10, 13), "cvs", end_assumed=True)
    assert amc._known_end_of_open_period(open_vlk, [open_vlk, one_day]) is None

    same_day_document = _episode("ВЛК", date(2026, 10, 13), date(2026, 10, 13), "cvs", end_assumed=False)
    later_document = _episode("ВЛК", date(2026, 10, 13), date(2026, 10, 20), "cvs", end_assumed=False)
    assert amc._known_end_of_open_period(open_vlk, [open_vlk, same_day_document]) == date(2026, 10, 13)
    assert amc._known_end_of_open_period(open_vlk, [open_vlk, later_document]) == date(2026, 10, 20)


def test_a_name_spelled_differently_in_an_addon_is_matched_and_reported(tmp_path):
    """Рішення користувача (реальний випадок: в ОБЛІК "В’...", у
    book61 - без апострофа; book61-записи цієї людини раніше не звірялись
    зовсім): людину зіставлено попри апостроф (дати звіряються), а різниця
    написання - окремий рядок звіту навіть без розбіжних дат (ПІБ червоним)."""
    oblik = _write_roster(tmp_path / "ОБЛІК.xlsx", [
        ["Посада1", "сержант", "П'ЯТИЙ П'ятий П'ятий", "ВП", "ВП", "РВЗ", "РВЗ", "РВЗ", "РВЗ", "РВЗ"],
    ], header=_WEEK_HEADER)
    records = [_period("ПЯТИЙ Пятий Пятий", "ВП", date(2026, 9, 25), date(2026, 10, 2))]

    results = amc.find_period_mismatches(Timesheet(oblik, "Табель"), records)

    assert [(r["source_file"], r["roster_pib_raw"], r["addon_pib_raw"]) for r in results] == [
        ("book61", "П'ЯТИЙ П'ятий П'ятий", "ПЯТИЙ Пятий Пятий"),
    ]
    assert all(amc._statuses_match(e, a) for a, e in results[0]["dates"].values())  # дати звіряються й збігаються
    report = str(tmp_path / "mis_match.xlsx")
    amc.write_addon_mismatch_report(results, [], report, source_files=["book61"])
    ws = load_workbook(report)["book61"]
    assert [ws.cell(row=row, column=3).value for row in (2, 3)] == ["П'ЯТИЙ П'ятий П'ятий", "ПЯТИЙ Пятий Пятий"]
    assert ws.cell(row=2, column=3).fill.fgColor.rgb == "FFFFC7CE"


def test_vacation_without_a_destination_is_a_yellow_mismatch_in_the_tabel_month(tmp_path):
    """Рішення користувача: відпустка book61 без "Куди вибув"
    (порожньо чи "?") - розбіжність, навіть коли статуси збігаються; у
    mis_match.xlsx - дні цієї відпустки в місяці Табеля просто жовтим
    (об'єднана клітинка зі значенням Табеля, без тексту), а в "Джерело" -
    примітка. Та сама відпустка з "Куди вибув" - не розбіжність."""
    oblik = _write_roster(tmp_path / "ОБЛІК.xlsx", [
        ["Посада1", "сержант", "ПЕРШИЙ Перший Перший", "ВП", "ВП", "РВЗ", "РВЗ", "РВЗ", "РВЗ", "РВЗ"],
        ["Посада2", "матрос", "ДРУГИЙ Другий Другий", "ВП", "ВП", "РВЗ", "РВЗ", "РВЗ", "РВЗ", "РВЗ"],
    ], header=_WEEK_HEADER)
    records = [
        {**_period("ПЕРШИЙ Перший Перший", "ВП", date(2026, 9, 25), date(2026, 10, 2)), "destination_missing": True},
        {**_period("ДРУГИЙ Другий Другий", "ВП", date(2026, 9, 25), date(2026, 10, 2)), "destination_missing": False},
    ]

    results = amc.find_period_mismatches(Timesheet(oblik, "Табель"), records)

    assert [(r["roster_pib_raw"], r["no_destination"]) for r in results] == [
        ("ПЕРШИЙ Перший Перший", [date(2026, 10, 1), date(2026, 10, 2)]),
    ]
    report = str(tmp_path / "mis_match.xlsx")
    amc.write_addon_mismatch_report(results, [], report, source_files=["book61"])
    ws = load_workbook(report)["book61"]
    assert ws.cell(row=3, column=1).value == "book61: ВП; ВП: не вказано «Куди вибув»"
    merged = {str(rng) for rng in ws.merged_cells.ranges}
    assert [ws.cell(row=2, column=col).value for col in (4, 5)] == ["ВП", "ВП"]
    assert {"D2:D3", "E2:E3"} <= merged  # одна клітинка, без тексту внизу
    assert {ws.cell(row=2, column=col).fill.fgColor.rgb for col in (4, 5)} == {"FFFFEB9C"}


def test_report_pib_column_is_wide_enough_and_left_aligned(tmp_path):
    """Рішення користувача: ПІБ у mis_match.xlsx - повністю."""
    from content.report_pib_column import pib_column_width
    name = "ВІСІМНАДЦЯТИЙ Вісімнадцятий Вісімнадцятий"
    records = [{"source_file": "book61", "sources": ["book61: ВП"], "roster_pidrozdil": "Підрозділ1",
                "roster_pib_raw": name, "addon_pib_raw": name, "dates": {date(2026, 10, 1): ("РВЗ", "ВП")},
                "tabel_values": {date(2026, 10, 1): "РВЗ"}}]
    path = str(tmp_path / "mis_match.xlsx")

    amc.write_addon_mismatch_report(records, [], path, source_files=["book61"])

    ws = load_workbook(path)["book61"]
    assert ws.column_dimensions["C"].width == pib_column_width([name]) > 30
    assert ws.cell(row=2, column=3).alignment.horizontal == "left"


def test_vacation_without_a_destination_alone_is_only_a_warning():
    """Рішення користувача: у терміналі жовтим, це не помилка:
    лише відпустка без "Куди вибув", а статуси збігаються й ПІБ те саме.
    Будь-яка розбіжність статусу чи інше написання ПІБ - вже помилка."""
    base = {"roster_pib_raw": "ПЕРШИЙ Перший Перший", "addon_pib_raw": "ПЕРШИЙ Перший Перший",
            "dates": {date(2026, 10, 1): ("ВП", "ВП")}, "no_destination": [date(2026, 10, 1)]}
    assert amc.is_warning_only(base)
    assert not amc.is_warning_only({**base, "dates": {date(2026, 10, 1): ("РВЗ", "ВП")}})
    assert not amc.is_warning_only({**base, "addon_pib_raw": "ПЕРШИЙ Перший Перши"})
    assert not amc.is_warning_only({**base, "no_destination": []})


def test_a_name_spelled_differently_is_reported_even_without_dates_in_the_tabel(tmp_path):
    """Запис книги цілком поза датами Табеля (лише вересень) - дат звіряти
    нема, але інше написання ПІБ однаково показується."""
    oblik = _write_roster(tmp_path / "ОБЛІК.xlsx", [
        ["Посада1", "сержант", "П'ЯТИЙ П'ятий П'ятий", "РВЗ", "РВЗ", "РВЗ", "РВЗ", "РВЗ", "РВЗ", "РВЗ"],
    ], header=_WEEK_HEADER)
    records = [_period("ПЯТИЙ Пятий Пятий", "ВП", date(2026, 9, 1), date(2026, 9, 5))]

    results = amc.find_period_mismatches(Timesheet(oblik, "Табель"), records)

    assert [(r["addon_pib_raw"], r["sources"], r["dates"]) for r in results] == [("ПЯТИЙ Пятий Пятий", ["book61: ВП"], {})]


def test_a_day_without_a_document_does_not_end_an_open_treatment(tmp_path):
    """Реальний випадок: БПШП з 01.10 без виписки, ЦВС - ВЛК 05.10 лише з
    датою проходження (end_assumed). Лікування після ВЛК триває, а не
    лишається без жодного запису ("немає запису на аркуші ШП")."""
    oblik = _write_roster(tmp_path / "ОБЛІК.xlsx", [
        ["Посада1", "матрос", "ПЕРШИЙ Перший Перший", "БПШП", "БПШП", "БПШП", "БПШП", "ВЛК", "БПШП", "БПШП"],
    ], header=_WEEK_HEADER)
    records = [
        _period("ПЕРШИЙ Перший Перший", "БПШП", date(2026, 10, 1), None, "СОМАТИКА-ПОРАНЕНІ", "cvs"),
        {**_period("ПЕРШИЙ Перший Перший", "ВЛК", date(2026, 10, 5), date(2026, 10, 5), "ВЛК", "cvs"), "end_assumed": True},
    ]

    assert amc.find_period_mismatches(Timesheet(oblik, "Табель"), records,
                                      files_by_kind={"book61": "book61", "cvs": "cvs"}) == []


def test_returned_from_szch_accepts_any_status_of_someone_back_on_duty():
    """Рішення користувача: "Повер" (book61, повернувся з СЗЧ) - збіг з
    будь-яким статусом Табеля, крім ADDON_NOT_RETURNED_AFTER_SZCH_STATUSES
    (і їхніх форм для виплат, напр. "100_БЗ")."""
    from constants import ADDON_NOT_RETURNED_AFTER_SZCH_STATUSES

    for actual in ("РВЗ", 100, 30, "ВП", "ШП", "Повер"):
        assert amc._statuses_match("Повер", actual), actual
    for actual in (*ADDON_NOT_RETURNED_AFTER_SZCH_STATUSES, *amc._payment_converted_forms("БЗ")):
        assert not amc._statuses_match("Повер", actual), actual


def test_a_misspelled_name_in_a_book_is_matched_and_shown_as_a_spelling_difference(tmp_path):
    """Реальний випадок: у Табелі ім'я з опискою, у book61 - інакше (прізвище
    й по батькові ті самі). Записи книги приєднуються до людини Табеля, дати
    звіряються як звичайно, а звіт показує обидва написання ПІБ. Інше ім'я
    (не описка) і людина, яку книга вже знає під табельним ПІБ, - ні."""
    oblik = _write_roster(tmp_path / "ОБЛІК.xlsx", [
        ["Посада1", "матрос", "ДРУГИЙ Другий Другий", "СЗЧ", "СЗЧ", "СЗЧ", "СЗЧ", "СЗЧ", "СЗЧ", "СЗЧ"],
        ["Посада2", "матрос", "ТРЕТІЙ Третій Третій", "РВЗ", "РВЗ", "РВЗ", "РВЗ", "РВЗ", "РВЗ", "РВЗ"],
        ["Посада3", "матрос", "ЧЕТВЕРТИЙ Четвертий Четвертий", "РВЗ", "РВЗ", "РВЗ", "РВЗ", "РВЗ", "РВЗ", "РВЗ"],
    ], header=_WEEK_HEADER)
    records = [
        _period("ДРУГИЙ Дугий Другий", "СЗЧ", date(2026, 10, 1), None, "СЗЧ"),
        _period("ТРЕТІЙ Сьомий Третій", "СЗЧ", date(2026, 10, 1), None, "СЗЧ"),  # інша людина
        _period("ЧЕТВЕРТИЙ Четвертий Четвертий", "ВП", date(2026, 9, 1), date(2026, 9, 5)),
        _period("ЧЕТВЕРТИЙ Четвртий Четвертий", "СЗЧ", date(2026, 10, 1), None, "СЗЧ"),
    ]

    results = amc.find_period_mismatches(Timesheet(oblik, "Табель"), records)

    assert [(r["roster_pib_raw"], r["addon_pib_raw"], r["dates"]) for r in results] == [
        ("ДРУГИЙ Другий Другий", "ДРУГИЙ Дугий Другий", {
            date(2026, 10, day): ("СЗЧ", "СЗЧ") for day in range(1, 8)
        }),
    ]
