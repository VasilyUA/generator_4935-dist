"""Звірка книг resources/addons між собою (content/addon_cross_checker.py,
generators/check_addon_mismatches.py::cross_check_addons)."""

from datetime import date, datetime

from book61_fixtures import write_full_book61
from cvs_fixtures import MEDICAL_SHEET, write_cvs
from openpyxl import Workbook, load_workbook

from constants import ADDON_NON_COMBAT_TREATMENT_STATUS
import content.addon_cross_checker as acc
import content.addon_mismatch_checker as amc
from content.addon_reader import ANY_STATUS
from content.oblik_timesheet import Timesheet
from generators.check_addon_mismatches import cross_check_addons

_HEADER = ["ПІДРОЗДІЛ", "ПОСАДА", "ЗВАННЯ", "ПІБ"] + [datetime(2026, 10, day) for day in range(1, 8)]


def _roster(path, names):
    wb = Workbook()
    ws = wb.worksheets[0]
    ws.title = "Табель"
    ws.append(_HEADER)
    for index, name in enumerate(names, start=1):
        ws.append([f"Підрозділ{index}", f"Посада{index}", "сержант", name] + [None] * 7)
    wb.save(str(path))
    return str(path)


def _period(pib_raw, status, start, end, source_file, source_sheet, arrival_date=None):
    return {
        "pib_raw": pib_raw, "status": status, "start_date": start, "end_date": end,
        "source_sheet": source_sheet, "source_file": source_file, "arrival_date": arrival_date,
    }


def test_statuses_of_newcomers_are_identical():
    """Рішення користувача: "НОВ" і "Адап" - не розбіжність."""
    assert amc._statuses_match("Адап", "НОВ") is True
    assert amc._statuses_match("НОВ", "Адап.") is True


def test_find_addon_conflicts_compares_each_book_on_its_own_timeline(tmp_path):
    """Кожна книга - своя шкала; звіряються лише дати, де обидві книги щось
    очікують, тими самими правилами еквівалентності (ППД=НОВ, 10 = ВД,
    "будь-який"); день виписки однієї з книг не звіряється; людина не з
    Табеля - не звіряється."""
    timesheet = Timesheet(_roster(tmp_path / "ОБЛІК.xlsx", ["ПЕРШИЙ Перший Перший", "ДРУГИЙ Другий Другий"]), "Табель")
    records = [
        # ПЕРШИЙ: 01-02 збіг ППД/НОВ, 03-04 ВП проти ШП, 05 - день виписки ЦВС.
        _period("ПЕРШИЙ Перший Перший", "ППД", date(2026, 10, 1), date(2026, 10, 2), "book61", "ППД"),
        _period("ПЕРШИЙ Перший Перший", "ВП", date(2026, 10, 3), date(2026, 10, 5), "book61", "ВП"),
        _period("ПЕРШИЙ Перший Перший", "НОВ", date(2026, 10, 1), date(2026, 10, 2), "cvs", "НОВ"),
        _period("ПЕРШИЙ Перший Перший", "ШП", date(2026, 10, 3), date(2026, 10, 5), "cvs", MEDICAL_SHEET,
                arrival_date=date(2026, 10, 5)),
        # ДРУГИЙ: усе збігається.
        _period("ДРУГИЙ Другий Другий", "ВД", date(2026, 10, 1), date(2026, 10, 3), "book61", "ВД"),
        _period("ДРУГИЙ Другий Другий", 10, date(2026, 10, 1), date(2026, 10, 2), "journal", "Фахова підготовка"),
        _period("ДРУГИЙ Другий Другий", ANY_STATUS, date(2026, 10, 3), date(2026, 10, 3), "journal", "БЗВП"),
        # Не з Табеля.
        _period("ТРЕТІЙ Третій Третій", "ВП", date(2026, 10, 1), date(2026, 10, 2), "book61", "ВП"),
        _period("ТРЕТІЙ Третій Третій", "ШП", date(2026, 10, 1), date(2026, 10, 2), "cvs", MEDICAL_SHEET),
    ]

    conflicts = acc.find_addon_conflicts(timesheet, records, ["book61", "cvs", "journal"], last_date=date(2026, 10, 7))

    assert [(c["files"], c["roster_pib_raw"], c["roster_pidrozdil"], c["days"]) for c in conflicts] == [
        (("book61", "cvs"), "ПЕРШИЙ Перший Перший", "Підрозділ1", 2),
    ]
    assert conflicts[0]["periods"] == [{
        "start": date(2026, 10, 3), "end": date(2026, 10, 4), "status_a": "ВП", "status_b": "ШП",
        "sheet_a": "ВП", "sheet_b": MEDICAL_SHEET,
        # Періоди самих записів, як у файлах (рішення користувача: для звіту).
        "period_a": (date(2026, 10, 3), date(2026, 10, 5)), "period_b": (date(2026, 10, 3), date(2026, 10, 5)),
    }]


def test_find_addon_conflicts_checks_every_period_not_only_the_tabel_month(tmp_path):
    """Рішення користувача: за ВСІ періоди: Табель за жовтень, а
    розбіжність книг у березні теж знаходиться; поспіль ідучі дні з тією
    самою парою статусів - один період."""
    timesheet = Timesheet(_roster(tmp_path / "ОБЛІК.xlsx", ["ПЕРШИЙ Перший Перший"]), "Табель")
    records = [
        _period("ПЕРШИЙ Перший Перший", "ППД", date(2026, 3, 1), date(2026, 3, 31), "book61", "ППД"),
        _period("ПЕРШИЙ Перший Перший", "ШП", date(2026, 3, 10), date(2026, 3, 12), "cvs", MEDICAL_SHEET),
        _period("ПЕРШИЙ Перший Перший", "ВЛК", date(2026, 3, 20), date(2026, 3, 20), "cvs", "ВЛК"),
    ]

    (conflict,) = acc.find_addon_conflicts(timesheet, records, ["book61", "cvs"], last_date=date(2026, 10, 7))

    assert [(p["start"], p["end"], p["status_a"], p["status_b"]) for p in conflict["periods"]] == [
        (date(2026, 3, 10), date(2026, 3, 12), "ППД", "ШП"),
        (date(2026, 3, 20), date(2026, 3, 20), "ППД", "ВЛК"),
    ]
    assert [(p["period_a"], p["period_b"]) for p in conflict["periods"]] == [
        ((date(2026, 3, 1), date(2026, 3, 31)), (date(2026, 3, 10), date(2026, 3, 12))),
        ((date(2026, 3, 1), date(2026, 3, 31)), (date(2026, 3, 20), date(2026, 3, 20))),
    ]
    assert conflict["days"] == 4


def test_hospital_stay_inside_bzvp_is_not_a_conflict(tmp_path):
    """Рішення користувача: ШП посеред БЗВП - нормально (і БЗВП з
    аркуша журналу "БЗВП" зі статусом "Навч" - навчання не у власній в/ч);
    ШП, що почалось ще до БЗВП, - не "всередині", розбіжність лишається."""
    names = ["ПЕРШИЙ Перший Перший", "ДРУГИЙ Другий Другий", "ТРЕТІЙ Третій Третій"]
    timesheet = Timesheet(_roster(tmp_path / "ОБЛІК.xlsx", names), "Табель")
    records = [
        _period(names[0], "БЗВП", date(2026, 2, 16), date(2026, 4, 7), "journal", "БЗВП", arrival_date=date(2026, 4, 7)),
        _period(names[0], "ШП", date(2026, 2, 19), date(2026, 2, 27), "cvs", MEDICAL_SHEET, arrival_date=date(2026, 2, 27)),
        _period(names[1], "Навч", date(2026, 3, 10), date(2026, 5, 8), "journal", "БЗВП", arrival_date=date(2026, 5, 8)),
        _period(names[1], "ШП", date(2026, 4, 23), date(2026, 5, 6), "cvs", MEDICAL_SHEET, arrival_date=date(2026, 5, 6)),
        _period(names[2], "БЗВП", date(2026, 2, 16), date(2026, 4, 7), "journal", "БЗВП", arrival_date=date(2026, 4, 7)),
        _period(names[2], "ШП", date(2026, 2, 10), date(2026, 2, 20), "cvs", MEDICAL_SHEET, arrival_date=date(2026, 2, 20)),
    ]

    conflicts = acc.find_addon_conflicts(timesheet, records, ["journal", "cvs"], last_date=date(2026, 10, 7))

    assert [(c["roster_pib_raw"], [(p["start"], p["end"]) for p in c["periods"]]) for c in conflicts] == [
        (names[2], [(date(2026, 2, 16), date(2026, 2, 19))]),  # 20.02 - день виписки, не звіряється
    ]


def test_conflict_periods_split_where_the_records_of_a_file_change(tmp_path):
    """Рядок звіту - одна пара записів: та сама пара статусів, але інший
    запис файлу (інший період у ньому) - окремий розбіжний період."""
    timesheet = Timesheet(_roster(tmp_path / "ОБЛІК.xlsx", ["ПЕРШИЙ Перший Перший"]), "Табель")
    records = [
        _period("ПЕРШИЙ Перший Перший", "ППД", date(2026, 3, 1), date(2026, 3, 31), "book61", "ППД"),
        _period("ПЕРШИЙ Перший Перший", "ШП", date(2026, 3, 10), date(2026, 3, 11), "cvs", MEDICAL_SHEET),
        _period("ПЕРШИЙ Перший Перший", "ШП", date(2026, 3, 12), None, "cvs", MEDICAL_SHEET),
        _period("ПЕРШИЙ Перший Перший", "ППД", date(2026, 3, 14), date(2026, 3, 31), "cvs", "ППД"),
    ]

    (conflict,) = acc.find_addon_conflicts(timesheet, records, ["book61", "cvs"], last_date=date(2026, 10, 7))

    assert [(p["start"], p["end"], p["period_b"]) for p in conflict["periods"]] == [
        (date(2026, 3, 10), date(2026, 3, 11), (date(2026, 3, 10), date(2026, 3, 11))),
        (date(2026, 3, 12), date(2026, 3, 13), (date(2026, 3, 12), None)),  # кінця у файлі немає
    ]


def _conflict(name, file_a, file_b, *periods):
    """periods - (початок, кінець розбіжності, статус A, статус B, аркуш A,
    аркуш B, період запису A, період запису B)."""
    return {
        "files": (file_a, file_b), "roster_pib_raw": name, "roster_pidrozdil": "Підрозділ1",
        "periods": [
            {"start": start, "end": end, "status_a": a, "status_b": b, "sheet_a": sheet_a, "sheet_b": sheet_b,
             "period_a": period_a, "period_b": period_b}
            for start, end, a, b, sheet_a, sheet_b, period_a, period_b in periods
        ],
        "days": sum((end - start).days + 1 for start, end, *_rest in periods),
    }


def test_write_addon_conflicts_report_puts_each_file_on_its_own_sheet_compared_with_the_others(tmp_path):
    """Рішення користувача: аркуш - назва файлу, на ньому цей
    файл порівняно з усіма іншими - одна таблиця з однією шапкою; перед
    аркушем, статусом і періодом - колонка з назвою їхнього файлу в кожному
    рядку (не "інший файл" у шапці); червоним - лише те, що не сходиться,
    решта - одного вигляду; та сама розбіжність - на аркушах обох файлів;
    файл без розбіжностей - аркуш із написом; ПІБ людини - одна об'єднана
    клітинка."""
    conflicts = [
        _conflict("ПЕРШИЙ Перший Перший", "book61", "cvs",
                  # Реальний випадок: періоди в обох файлах однакові, різний лише статус
                  # (розбіжні дні - на день коротші: день виписки не звіряється).
                  (date(2026, 8, 10), date(2026, 8, 12), "ШП", "БПШП", "ШП", MEDICAL_SHEET,
                   (date(2026, 8, 10), date(2026, 8, 13)), (date(2026, 8, 10), date(2026, 8, 13))),
                  (date(2026, 3, 10), date(2026, 3, 12), "ППД", "ШП", "ППД", MEDICAL_SHEET,
                   (date(2026, 3, 1), date(2026, 3, 31)), (date(2026, 3, 10), None))),
    ]
    output_path = str(tmp_path / "addons_mis_match.xlsx")

    assert acc.write_addon_conflicts_report([], output_path) is None
    assert acc.write_addon_conflicts_report(conflicts, output_path, source_files=["book61", "cvs", "journal"]) == output_path

    wb = load_workbook(output_path)
    assert wb.sheetnames == ["book61", "cvs", "journal"]
    book61 = wb["book61"]
    assert [book61.cell(row=3, column=c).value for c in range(1, 14)] == [
        "Підрозділ", "ПІБ",
        "Файл", "Аркуш", "Статус", "З", "По",
        "Порівняно з файлом", "Аркуш", "Статус", "З", "По",
        "У чому розбіжність",
    ]
    assert book61.cell(row=3, column=14).value is None  # колонок "розбіжні дні" немає
    assert not any(rng.min_row == 3 for rng in book61.merged_cells.ranges)  # одна шапка, без об'єднаних груп
    # Періоди ОБОХ файлів - як записані в них; кінця у файлі немає - "не вказано".
    assert [book61.cell(row=4, column=c).value for c in range(2, 14)] == [
        "ПЕРШИЙ Перший Перший",
        "book61", "ППД", "ППД", datetime(2026, 3, 1), datetime(2026, 3, 31),
        "cvs", MEDICAL_SHEET, "ШП", datetime(2026, 3, 10), "не вказано",
        "Не сходиться: статус, дата початку, дата кінця",
    ]
    assert [book61.cell(row=5, column=c).value for c in range(3, 14)] == [
        "book61", "ШП", "ШП", datetime(2026, 8, 10), datetime(2026, 8, 13),
        "cvs", MEDICAL_SHEET, "БПШП", datetime(2026, 8, 10), datetime(2026, 8, 13),
        "Не сходиться: статус",
    ]
    assert book61.cell(row=4, column=6).number_format == "dd.MM.yyyy"
    # Дата шрифтом звіту вміщається, а не "#######".
    assert all(book61.column_dimensions[letter].width >= 16 for letter in "FGKL")
    assert "B4:B5" in {str(rng) for rng in book61.merged_cells.ranges}  # одна людина - один ПІБ

    def red(row, column):
        return book61.cell(row=row, column=column).fill.fgColor.rgb == "FFFFC7CE"

    # Червоним - ЛИШЕ те, що не сходиться: у рядку з однаковими періодами - самі статуси.
    assert [c for c in range(1, 14) if red(5, c)] == [5, 10]
    assert [c for c in range(1, 14) if red(4, c)] == [5, 6, 7, 10, 11, 12]
    # Решта - одного вигляду: шапка одного кольору, клітинки без заливки.
    assert {book61.cell(row=3, column=c).fill.fgColor.rgb for c in range(1, 14)} == {"FF035C6B"}
    assert {book61.cell(row=5, column=c).fill.fill_type for c in range(1, 14) if c not in (5, 10)} == {None}
    cvs = wb["cvs"]
    assert [cvs.cell(row=4, column=c).value for c in range(3, 13)] == [
        "cvs", MEDICAL_SHEET, "ШП", datetime(2026, 3, 10), "не вказано",
        "book61", "ППД", "ППД", datetime(2026, 3, 1), datetime(2026, 3, 31),
    ]
    assert wb["journal"].cell(row=4, column=1).value == "Розбіжностей з іншими файлами немає."


def test_report_points_at_a_date_far_in_the_future_as_a_likely_year_typo(tmp_path):
    """Реальний випадок: ЦВС - лікування 30.06.2026-05.07.2027 (рік виписки
    з помилкою), тож "розбіжність" з відпусткою book61 у серпні; дата - червоним
    жирним, а в поясненні - що це, ймовірно, помилка в році."""
    conflicts = [
        _conflict("ПЕРШИЙ Перший Перший", "book61", "cvs",
                  (date(2026, 8, 13), date(2026, 9, 10), "ВПСЗ", "БПШП", "ВП", MEDICAL_SHEET,
                   (date(2026, 8, 13), date(2026, 9, 10)), (date(2026, 6, 30), date(2027, 7, 5)))),
    ]
    output_path = str(tmp_path / "addons_mis_match.xlsx")

    acc.write_addon_conflicts_report(conflicts, output_path, source_files=["book61", "cvs"], today=date(2026, 10, 26))

    book61 = load_workbook(output_path)["book61"]
    assert book61.cell(row=4, column=13).value == (
        "Не сходиться: статус, дата початку, дата кінця. Увага: у «cvs» дата 05.07.2027 - далеко в майбутньому, "
        "ймовірно помилка в році (виправте у файлі)"
    )
    end, start = book61.cell(row=4, column=12), book61.cell(row=4, column=11)
    assert (end.font.color.rgb, end.font.bold) == ("FFC00000", True)  # кінець запису cvs - 2027 рік
    assert (start.font.color.rgb, start.font.bold) == ("FF9C0006", False)  # початок - просто інший, ніж у book61


def test_file_rows_show_the_conflict_from_each_files_point_of_view():
    period_a, period_b = (date(2025, 12, 30), date(2026, 1, 5)), (date(2026, 1, 1), None)
    conflicts = [_conflict("ПЕРШИЙ Перший Перший", "a", "b",
                           (date(2026, 1, 1), date(2026, 1, 2), "ВП", "ШП", "ВП", "ШП", period_a, period_b))]
    (row_a,) = acc.file_rows(conflicts, "a")
    (row_b,) = acc.file_rows(conflicts, "b")
    assert (row_a["other_file"], row_a["my_status"], row_a["other_status"], row_a["days"]) == ("b", "ВП", "ШП", 2)
    assert (row_a["my_period"], row_a["other_period"]) == (period_a, period_b)
    assert (row_b["other_file"], row_b["my_status"], row_b["other_status"]) == ("a", "ШП", "ВП")
    assert (row_b["my_period"], row_b["other_period"]) == (period_b, period_a)
    assert acc.file_rows(conflicts, "c") == []


def test_file_rows_show_the_same_pair_of_records_once():
    """Та сама пара записів кількома проміжками днів (між ними - день, що не
    звіряється) - у звіті лише періоди записів, тож ОДИН рядок, а не кілька
    однакових."""
    period_a, period_b = (date(2026, 3, 1), date(2026, 3, 31)), (date(2026, 3, 1), date(2026, 3, 31))
    conflicts = [_conflict("ПЕРШИЙ Перший Перший", "a", "b",
                           (date(2026, 3, 1), date(2026, 3, 9), "ППД", "ШП", "ППД", "ШП", period_a, period_b),
                           (date(2026, 3, 11), date(2026, 3, 30), "ППД", "ШП", "ППД", "ШП", period_a, period_b),
                           (date(2026, 3, 10), date(2026, 3, 10), "ВЛК", "ШП", "ВЛК", "ШП", period_a, period_b))]

    rows = acc.file_rows(conflicts, "a")

    assert [(r["my_status"], r["start"], r["end"], r["days"]) for r in rows] == [
        ("ППД", date(2026, 3, 1), date(2026, 3, 30), 29),
        ("ВЛК", date(2026, 3, 10), date(2026, 3, 10), 1),
    ]


def test_cross_check_addons_reads_real_books_and_reports_conflicts(tmp_path):
    roster = _roster(tmp_path / "ОБЛІК.xlsx", ["ПЕРШИЙ Перший Перший"])
    book61 = write_full_book61(tmp_path / "book61.xlsm", {
        "ППД": [[1, "Підрозділ1", "сержант", "ПЕРШИЙ Перший Перший", "ППД", date(2026, 9, 1), None]],
    })
    cvs = write_cvs(tmp_path / "cvs.xlsx", {
        MEDICAL_SHEET: [[None, 1, "Підрозділ1", "сержант", "ПЕРШИЙ Перший Перший", "соматика", datetime(2026, 10, 3),
                         "Лікарня", "№1", datetime(2026, 10, 6)]],
    })
    report = str(tmp_path / "addons_mis_match.xlsx")

    result = cross_check_addons([book61, cvs], roster_path=roster, report_file_name=report)

    assert result.checked_files == [book61, cvs] and result.unrecognized_files == []
    assert result.implausible_date_records == []
    assert result.report_path == report
    (conflict,) = result.conflicts
    assert [(p["start"], p["end"], p["status_a"], p["status_b"]) for p in conflict["periods"]] == [
        (date(2026, 10, 3), date(2026, 10, 5), "ППД", ADDON_NON_COMBAT_TREATMENT_STATUS),  # 06.10 - день виписки, не звіряється
    ]
    assert load_workbook(report).sheetnames == ["book61", "cvs"]


def test_non_combat_treatment_inside_bzvp_is_not_a_conflict():
    """ШП посеред БЗВП - нормально (ADDON_NESTED_STATUSES); не бойове
    лікування з ЦВС ("ШП/ЛМР/ВПСЗ") - теж."""
    outer = _period("ПЕРШИЙ Перший Перший", "БЗВП", date(2026, 2, 16), date(2026, 4, 7), "journal", "БЗВП")
    inner = _period("ПЕРШИЙ Перший Перший", ADDON_NON_COMBAT_TREATMENT_STATUS, date(2026, 2, 19), date(2026, 2, 27),
                    "cvs", MEDICAL_SHEET)
    assert acc._nested(outer, inner)


def test_a_person_found_in_only_one_book_has_nothing_to_compare(tmp_path):
    timesheet = Timesheet(_roster(tmp_path / "ОБЛІК.xlsx", ["ПЕРШИЙ Перший Перший"]), "Табель")
    records = [_period("ПЕРШИЙ Перший Перший", "ВП", date(2026, 10, 1), date(2026, 10, 5), "book61", "ВП")]

    assert acc.find_addon_conflicts(timesheet, records, ["book61", "cvs"], last_date=date(2026, 10, 7)) == []
