from datetime import date, datetime

from docx_fixtures import ARRIVAL_HEADER, DEPARTURE_HEADER, write_daily_report
from openpyxl import Workbook, load_workbook

from generators.generate_timesheet import _review_line, generate_timesheet

_HEADER = ["ПОСАДА", "ЗВАННЯ", "ПІБ", datetime(2026, 7, 31), datetime(2026, 8, 1), datetime(2026, 8, 2), datetime(2026, 8, 3)]


def _write_roster(path, rows):
    wb = Workbook()
    ws = wb.active
    ws.title = "Табель"
    for col_idx, value in enumerate(_HEADER, start=1):
        ws.cell(row=1, column=col_idx, value=value)
    for row_idx, row_values in enumerate(rows, start=2):
        for col_idx, value in enumerate(row_values, start=1):
            ws.cell(row=row_idx, column=col_idx, value=value)
    wb.save(str(path))
    return str(path)


def test_generate_timesheet_end_to_end(tmp_path):
    report_dir = tmp_path / "report"
    report_dir.mkdir()
    write_daily_report(
        report_dir / "01.08.2026 - щоденний рапорт.docx",
        groups=[(
            "ВИБУЛИ зі складу сил та засобів 9 армійського корпусу:",
            [("У відпустку:", DEPARTURE_HEADER, [["1", "матрос", "ПЕРШИЙ Перший Перший", "Посада", "01.08.2026", "Вибув"]])],
        )],
    )
    write_daily_report(
        report_dir / "02.08.2026 - щоденний рапорт.docx",
        groups=[(
            "ПРИБУЛИ до складу сил та засобів 9 армійського корпусу:",
            [("З відпустки:", ARRIVAL_HEADER, [["1", "матрос", "ПЕРШИЙ Перший Перший", "Посада", "02.08.2026", "Прибув"]])],
        )],
        freeform_paragraphs=["хтось невідомий згадується вільним текстом тут."],
    )

    roster_path = _write_roster(tmp_path / "ОБЛІК.xlsx", [
        ["Посада1", "сержант", "ПЕРШИЙ Перший Перший", "РВЗ", None, None, None],
        ["Посада2", "матрос", "ДРУГИЙ Другий Другий", "РВЗ", None, None, None],
    ])
    output_path = tmp_path / "output" / "ОБЛІК.xlsx"
    review_path = tmp_path / "output" / "review.txt"

    result_path, unresolved = generate_timesheet(
        report_dir=str(report_dir),
        personel_list_file_name=roster_path,
        personel_list_sheet_name="Табель",
        output_file_name=str(output_path),
        review_file_name=str(review_path),
    )

    assert result_path == str(output_path)
    assert output_path.exists()

    saved = load_workbook(str(output_path))["Табель"]
    # ПЕРШИЙ: 31.07=РВЗ (базовий) -> 01.08=ВП (У відпустку) -> 02.08=РВЗ (З відпустки, останній рапорт);
    # 03.08 - після останнього рапорту, свідомо не займається.
    assert [saved.cell(row=2, column=c).value for c in range(4, 8)] == ["РВЗ", "ВП", "РВЗ", None]
    # ДРУГИЙ: жодних подій - просто продовжує базовий статус через дати наявних звітів (01.08-02.08).
    assert [saved.cell(row=3, column=c).value for c in range(4, 8)] == ["РВЗ", "РВЗ", "РВЗ", None]

    assert len(unresolved) == 1  # вільний текст "Поза межами" з 02.08 - на ручну перевірку
    assert review_path.exists()
    assert "Поза межами" in review_path.read_text(encoding="utf-8")


def test_confirmed_health_leave_next_day_overrides_plain_vp_from_needs_leave_report(tmp_path):
    """Реальний випадок: рапорт дня N - "потребує відпустки за висновком ВЛК
    №... від ДАТА." (без прямої причини) - дає "ВП" на день N+1
    (_discharge_vlk_leave_events); рапорт дня N+1 - ОКРЕМЕ речення "З ДАТА
    відпустка за станом здоров'я висновок ВЛК №... від ...:" - дає "ВПСЗ" на
    ТУ САМУ дату N+1. Дві РІЗНІ події на ОДНУ дату - автоматичний override
    (STATUS_CONFLICT_OVERRIDES: ВПСЗ - специфічніший статус, перемагає) - без
    жодного запису на ручну перевірку, фінальний статус - "ВПСЗ", не "ВП"."""
    report_dir = tmp_path / "report"
    report_dir.mkdir()
    write_daily_report(
        report_dir / "01.08.2026 - щоденний рапорт.docx",
        freeform_paragraphs=["матрос ПЕРШИЙ Перший Перший потребує відпустки за висновком ВЛК №123 від 01.08.2026."],
    )
    write_daily_report(
        report_dir / "02.08.2026 - щоденний рапорт.docx",
        freeform_paragraphs=["матрос ПЕРШИЙ Перший Перший з 02.08.2026 відпустка за станом здоров'я висновок ВЛК №123 від 01.08.2026:"],
    )
    roster_path = _write_roster(tmp_path / "ОБЛІК.xlsx", [["Посада1", "матрос", "ПЕРШИЙ Перший Перший", "РВЗ", None, None, None]])
    output_path = tmp_path / "output" / "ОБЛІК.xlsx"

    _result_path, unresolved = generate_timesheet(
        report_dir=str(report_dir),
        personel_list_file_name=roster_path,
        personel_list_sheet_name="Табель",
        output_file_name=str(output_path),
        review_file_name=str(tmp_path / "output" / "review.txt"),
    )

    assert unresolved == []
    saved = load_workbook(str(output_path))["Табель"]
    # 31.07=РВЗ (базовий) -> 01.08=ВЛК (якорна подія) -> 02.08=ВПСЗ (не "ВП").
    assert [saved.cell(row=2, column=c).value for c in range(4, 7)] == ["РВЗ", "ВЛК", "ВПСЗ"]


def test_review_line_includes_pib_when_not_already_in_reason():
    item = {"reason": "Невідомий заголовок пункту рапорту: \"Щось:\".", "pib_raw": "ПЕРШИЙ Перший Перший", "report_date": None}

    assert "ПІБ: ПЕРШИЙ Перший Перший" in _review_line(item)


def test_review_line_does_not_duplicate_pib_already_embedded_in_reason():
    item = {
        "reason": "ПІБ \"ПЕРШИЙ Перший Перший\" з рапорту не знайдено в ОБЛІК.xlsx.",
        "pib_raw": "ПЕРШИЙ Перший Перший", "report_date": None,
    }

    assert _review_line(item).count("ПЕРШИЙ Перший Перший") == 1


def test_review_line_includes_would_be_status_and_report_text(tmp_path):
    """Реальний випадок: подія коректно розпізнана (статус мав би стати
    "ВПСЗ"), але людину не знайдено в ОБЛІК.xlsx - повідомлення має пояснити,
    ЯКИЙ статус мав би застосуватись і З ЯКОГО речення рапорту він узятий, а
    не лише "не знайдено", без жодного контексту."""
    item = {
        "reason": "ПІБ \"ТРЕТІЙ Третій Третій\" з рапорту не знайдено в ОБЛІК.xlsx.",
        "pib_raw": "ТРЕТІЙ Третій Третій", "status": "ВПСЗ",
        "raw_text": "матрос ТРЕТІЙ Третій Третій ... вибув у відпустку за станом здоров'я...",
        "report_date": date(2026, 8, 8),
    }
    line = _review_line(item)

    assert line.startswith("[08.08.2026] ")
    assert "мав би застосуватись статус \"ВПСЗ\"" in line
    assert "текст рапорту: \"матрос ТРЕТІЙ" in line


def test_review_line_plain_reason_without_any_extra_details():
    assert _review_line({"reason": "Файл не знайдено.", "report_date": None}) == "Файл не знайдено."


def test_table_return_from_health_leave_with_wrong_month_typo_is_flagged_not_silently_applied(tmp_path):
    """Реальний випадок: відпустка за станом здоров'я (ВПСЗ) мала б
    закінчитись 20.08.2026 (людина повертається "З відпустки:" - таблична
    подія, DEFAULT_STATUS), але в самому рапорті - друкарська помилка: вказано
    "20.07.2026" (той самий день місяця, ПОПЕРЕДНІЙ місяць - типова помилка
    "не той місяць"). "20.07.2026" РАНІШЕ за найпершу колонку табеля
    (31.07.2026), АЛЕ статус події ("РВЗ") НЕ збігається з тим, що ВЖЕ стоїть
    у базовій колонці ("ВПСЗ") - сумнів НЕ знято (на відміну від випадку, де
    статус збігається б і подію можна було б мовчки пропустити - див.
    test_apply_events_reports_date_before_the_earliest_column_when_it_
    contradicts_the_baseline у test_oblik_timesheet.py) - тож подія лишається
    на ручну перевірку, а НЕ застосовується мовчки з хибною датою (людина й
    далі показує "ВПСЗ" на всі наявні дати - жодної зміни, доки хтось не
    виправить дату в рапорті вручну)."""
    header = ["ПОСАДА", "ЗВАННЯ", "ПІБ", datetime(2026, 7, 31)] + [
        datetime(2026, 8, day) for day in range(1, 21)
    ]
    wb = Workbook()
    ws = wb.active
    ws.title = "Табель"
    for col_idx, value in enumerate(header, start=1):
        ws.cell(row=1, column=col_idx, value=value)
    ws.cell(row=2, column=1, value="Посада1")
    ws.cell(row=2, column=2, value="сержант")
    ws.cell(row=2, column=3, value="ШОСТИЙ Шостий Шостий")
    ws.cell(row=2, column=4, value="ВПСЗ")
    roster_path = tmp_path / "ОБЛІК.xlsx"
    wb.save(str(roster_path))

    report_dir = tmp_path / "report"
    report_dir.mkdir()
    write_daily_report(
        report_dir / "20.08.2026 - щоденний рапорт.docx",
        groups=[(
            "ПРИБУЛИ до складу сил та засобів 9 армійського корпусу:",
            [("З відпустки:", ARRIVAL_HEADER, [["1", "сержант", "ШОСТИЙ Шостий Шостий", "Посада1", "20.07.2026", "Прибув"]])],
        )],
    )
    output_path = tmp_path / "output" / "ОБЛІК.xlsx"
    review_path = tmp_path / "output" / "review.txt"

    _result_path, unresolved = generate_timesheet(
        report_dir=str(report_dir),
        personel_list_file_name=str(roster_path),
        personel_list_sheet_name="Табель",
        output_file_name=str(output_path),
        review_file_name=str(review_path),
    )

    assert len(unresolved) == 1
    assert "відсутня серед колонок" in unresolved[0]["reason"]
    saved = load_workbook(str(output_path))["Табель"]
    # Жодна колонка не змінилась - "ВПСЗ" продовжується аж до 20.08 включно
    # (подія з хибною датою НЕ застосована).
    assert [saved.cell(row=2, column=c).value for c in range(4, 25)] == ["ВПСЗ"] * 21


def test_freeform_szch_with_out_of_range_date_is_flagged_not_silently_applied(tmp_path):
    """Реальний випадок (07.08.2026 - щоденний рапорт.docx, "Поза межами"):
    "<ПІБ> ... 06.07.2026 вибув у СЗЧ" - дата в реченні поза діапазоном
    колонок ОБЛІК.xlsx (07 замість, судячи з усього, 08 - друкарська помилка
    в самому рапорті). Подія ВИТЯГУЄТЬСЯ (статус "СЗЧ"), але НЕ застосовується
    мовчки з хибною датою - лишається на ручну перевірку з конкретною
    причиною, як і однотипні друкарські помилки в табличних подіях."""
    report_dir = tmp_path / "report"
    report_dir.mkdir()
    write_daily_report(
        report_dir / "07.08.2026 - щоденний рапорт.docx",
        freeform_paragraphs=["матрос ЧЕТВЕРТИЙ Четвертий Четвертий під час проходження ВЛК 06.07.2026 вибув у СЗЧ."],
    )
    roster_path = _write_roster(tmp_path / "ОБЛІК.xlsx", [["Посада1", "матрос", "ЧЕТВЕРТИЙ Четвертий Четвертий", "РВЗ", None, None, None]])
    output_path = tmp_path / "output" / "ОБЛІК.xlsx"

    _result_path, unresolved = generate_timesheet(
        report_dir=str(report_dir),
        personel_list_file_name=roster_path,
        personel_list_sheet_name="Табель",
        output_file_name=str(output_path),
        review_file_name=str(tmp_path / "output" / "review.txt"),
    )

    assert len(unresolved) == 1
    assert "відсутня серед колонок" in unresolved[0]["reason"]
    assert unresolved[0]["pib_raw"] == "ЧЕТВЕРТИЙ Четвертий Четвертий"
    # Базовий статус лишається незайманим - подія не застосована.
    saved = load_workbook(str(output_path))["Табель"]
    assert saved.cell(row=2, column=4).value == "РВЗ"


def test_generate_timesheet_reports_report_file_without_a_parseable_date(tmp_path):
    report_dir = tmp_path / "report"
    report_dir.mkdir()
    write_daily_report(report_dir / "щоденний рапорт без дати.docx")  # без дати в назві файлу
    write_daily_report(
        report_dir / "01.08.2026 - щоденний рапорт.docx",
        groups=[(
            "ВИБУЛИ зі складу сил та засобів 9 армійського корпусу:",
            [("У відпустку:", DEPARTURE_HEADER, [["1", "матрос", "ПЕРШИЙ Перший Перший", "Посада", "01.08.2026", "Вибув"]])],
        )],
    )
    roster_path = _write_roster(tmp_path / "ОБЛІК.xlsx", [["Посада1", "сержант", "ПЕРШИЙ Перший Перший", "РВЗ", None, None, None]])

    _result_path, unresolved = generate_timesheet(
        report_dir=str(report_dir),
        personel_list_file_name=roster_path,
        personel_list_sheet_name="Табель",
        output_file_name=str(tmp_path / "output" / "ОБЛІК.xlsx"),
        review_file_name=str(tmp_path / "output" / "review.txt"),
    )

    assert any("не вдалося визначити дату" in item["reason"].lower() for item in unresolved)


def test_generate_timesheet_ignores_office_lock_file_report(tmp_path):
    """Реальний випадок: користувач тримає рапорт СЬОГОДНІШНЬОГО дня
    відкритим у Word під час генерації - Word створює тимчасовий lock-файл
    ("~$..." - той самий принцип, що й information_unit_reader._is_office_
    lock_file), чиє МАНГЛЕНЕ ім'я (перші символи замінено на "~$") взагалі
    не містить розпізнаваної дати - раніше потрапляв у unresolved як
    "не вдалося визначити дату", хоча сам користувач нічого не переплутав -
    lock-файл просто НЕ є справжнім рапортом і має мовчки пропускатись."""
    report_dir = tmp_path / "report"
    report_dir.mkdir()
    write_daily_report(
        report_dir / "01.08.2026 - щоденний рапорт.docx",
        groups=[(
            "ВИБУЛИ зі складу сил та засобів 9 армійського корпусу:",
            [("У відпустку:", DEPARTURE_HEADER, [["1", "матрос", "ПЕРШИЙ Перший Перший", "Посада", "01.08.2026", "Вибув"]])],
        )],
    )
    (report_dir / "~$.08.2026 - щоденний рапорт.docx").write_bytes(b"")
    roster_path = _write_roster(tmp_path / "ОБЛІК.xlsx", [["Посада1", "сержант", "ПЕРШИЙ Перший Перший", "РВЗ", None, None, None]])

    _result_path, unresolved = generate_timesheet(
        report_dir=str(report_dir),
        personel_list_file_name=roster_path,
        personel_list_sheet_name="Табель",
        output_file_name=str(tmp_path / "output" / "ОБЛІК.xlsx"),
        review_file_name=str(tmp_path / "output" / "review.txt"),
    )

    assert not any("не вдалося визначити дату" in item["reason"].lower() for item in unresolved)


def test_generate_timesheet_raises_when_only_a_lock_file_is_present(tmp_path):
    """Тека з ЄДИНИМ файлом - lock-файлом Word - рахується "без жодного
    рапорту" (той самий FileNotFoundError, що й для геть порожньої теки),
    а НЕ "рапорт без розпізнаваної дати" - lock-файл ніколи не рахується
    справжнім рапортом."""
    report_dir = tmp_path / "report"
    report_dir.mkdir()
    (report_dir / "~$.08.2026 - щоденний рапорт.docx").write_bytes(b"")
    roster_path = _write_roster(tmp_path / "ОБЛІК.xlsx", [["Посада1", "сержант", "ПЕРШИЙ Перший Перший", "РВЗ", None, None, None]])

    try:
        generate_timesheet(
            report_dir=str(report_dir),
            personel_list_file_name=roster_path,
            personel_list_sheet_name="Табель",
            output_file_name=str(tmp_path / "output" / "ОБЛІК.xlsx"),
            review_file_name=str(tmp_path / "output" / "review.txt"),
        )
        assert False, "мало підняти FileNotFoundError"
    except FileNotFoundError:
        pass


def test_generate_timesheet_raises_without_any_report_files(tmp_path):
    empty_report_dir = tmp_path / "report"
    empty_report_dir.mkdir()
    roster_path = _write_roster(tmp_path / "ОБЛІК.xlsx", [["Посада1", "сержант", "ПЕРШИЙ Перший Перший", "РВЗ", None, None, None]])

    try:
        generate_timesheet(
            report_dir=str(empty_report_dir),
            personel_list_file_name=roster_path,
            personel_list_sheet_name="Табель",
            output_file_name=str(tmp_path / "output" / "ОБЛІК.xlsx"),
            review_file_name=str(tmp_path / "output" / "review.txt"),
        )
        assert False, "мало підняти FileNotFoundError"
    except FileNotFoundError:
        pass


def test_review_file_states_no_issues_when_everything_resolved(tmp_path):
    report_dir = tmp_path / "report"
    report_dir.mkdir()
    write_daily_report(
        report_dir / "01.08.2026 - щоденний рапорт.docx",
        groups=[(
            "ВИБУЛИ зі складу сил та засобів 9 армійського корпусу:",
            [("У відпустку:", DEPARTURE_HEADER, [["1", "матрос", "ПЕРШИЙ Перший Перший", "Посада", "01.08.2026", "Вибув"]])],
        )],
    )
    roster_path = _write_roster(tmp_path / "ОБЛІК.xlsx", [["Посада1", "сержант", "ПЕРШИЙ Перший Перший", "РВЗ", None, None, None]])
    review_path = tmp_path / "output" / "review.txt"

    _result_path, unresolved = generate_timesheet(
        report_dir=str(report_dir),
        personel_list_file_name=roster_path,
        personel_list_sheet_name="Табель",
        output_file_name=str(tmp_path / "output" / "ОБЛІК.xlsx"),
        review_file_name=str(review_path),
    )

    assert unresolved == []
    assert "не знайдено" in review_path.read_text(encoding="utf-8")
