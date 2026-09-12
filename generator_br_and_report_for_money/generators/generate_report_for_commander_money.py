import os
from datetime import datetime

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches

from constants import (
    OUTPUT_DIR,
    SHORT_UNIT_BATTALION,
    HIGHER_COMMANDER_TITLE,
    COMMANDER_MONEY_REPORT_CATEGORIES,
    MONTH,
    YEAR,
)
from utils.logging_utils import print_green
from formatting.docx_utils import set_margins, add_page_number_header, add_paragraph_with_style, add_blank_paragraphs, save_docx_safely
from content.br_helpers import find_higher_commander
from content.money_report_helpers import (
    month_nominative_upper,
    month_genitive_lower,
    get_date_columns,
    get_tvo_commander_periods,
    normalize_name,
    _blank_non_szch_days_for_szch_started_this_month,
)
from utils.date_utils import to_date
from content.report_changes import build_changes_entries
from generators.generate_report_for_get_money import (
    STATIK,
    NOT_PAID_TABLE_HEADERS,
    _NOT_PAID_POINT,
    _build_categories,
    _add_category_section,
    _add_changes_sections,
    _flatten_changes_rows_for_logging,
    _log_missing_legal_basis,
    _log_missing_posada,
    _add_footer,
)

# Рапорт на додаткову винагороду ДЛЯ КОМАНДИРА/ТВО - штатний командир батальйону
# (ПОСАДА == HIGHER_COMMANDER_TITLE в ОБЛІК.xlsx) не може фігурувати у ГОЛОВНОМУ
# рапорті (той пишеться командиром від третьої особи про підлеглих), тож для
# нього (і для ТВО, якщо такий був протягом місяця) - окремий, від першої особи.


def _find_regular_commander_row(rows_with_data):
    return next((row for row in rows_with_data if row.get("ПОСАДА") == HIGHER_COMMANDER_TITLE), None)


def _clip_row_to_dates(row, date_columns, allowed_dates):
    """Копія row, де значення комірок дат ПОЗА allowed_dates прибрані (мов їх узагалі
    не було в ОБЛІК.xlsx) - лише дні within allowed_dates лишаються видимими для
    build_category_rows/_build_categories (решта комірок відсутні -> resolve_day_value_and_category
    поверне (None, None) для них, так само, як для порожньої комірки)."""
    clipped = dict(row)
    for col in date_columns:
        if col not in allowed_dates:
            clipped.pop(col, None)
    return clipped


def _collect_commander_rows(rows_with_data, rows_with_tvo_data, date_columns, period_start, period_end, narrator_row):
    """ХТО ПІДПИСУЄ рапорт (narrator_row - find_higher_commander на "сьогодні",
    рахує викликач і передає сюди) визначає, ЧИ додавати ТВО-рядок окремо від
    штатного командира - підтверджено користувачем:

    - Якщо підписує ШТАТНИЙ командир (narrator_row без позначки ТВО, тобто
      наразі немає активного на "сьогодні" ТВО на цю посаду) - жоден ТВО
      НЕ додається взагалі, навіть якщо хтось був ТВО на цю посаду ЧАСТИНУ
      звітного місяця: штатний командир отримує ОДИН рядок за ВЕСЬ звітний
      період (build_category_rows/_build_categories, викликані далі, самі
      відфільтрують лише його СПРАВДІ оплачувані дні за фактичними
      значеннями комірок - "статус 100").

    - Якщо підписує ТВО (narrator_row позначений ТВО, тобто хтось активний
      на цю посаду САМЕ "сьогодні") - і цей(і) ТВО, і штатний командир (якщо
      є) потрапляють РАЗОМ, кожен своїм окремим рядком, ОБМЕЖЕНИМ лише
      днями, коли він ФАКТИЧНО виконував цю роль: штатний командир - за дні,
      НЕ покриті жодним ТВО цього місяця, кожен ТВО - за дні своєї заявленої
      заміни (Start/End, обрізані до меж звітного періоду). Якщо штатний
      командир сам одночасно значиться ТВО (рідкісний збіг) - обидва
      діапазони дат об'єднуються в ОДИН його рядок, без дублювання."""
    regular_row = _find_regular_commander_row(rows_with_data)

    if not narrator_row.get("ТВО", False):
        return [dict(regular_row)] if regular_row is not None else []

    tvo_periods = get_tvo_commander_periods(rows_with_tvo_data, HIGHER_COMMANDER_TITLE, period_start, period_end)

    tvo_dates_by_pib = {}
    for tvo in tvo_periods:
        allowed = {col for col in date_columns if tvo["start"] <= to_date(col) <= tvo["end"]}
        tvo_dates_by_pib.setdefault(tvo["pib"], set()).update(allowed)

    rows = []
    if regular_row is not None:
        regular_pib = normalize_name(regular_row.get("ПІБ", ""))
        all_tvo_dates = set().union(*tvo_dates_by_pib.values()) if tvo_dates_by_pib else set()
        own_dates = {col for col in date_columns if col not in all_tvo_dates}
        own_dates |= tvo_dates_by_pib.pop(regular_pib, set())
        if own_dates:
            rows.append(_clip_row_to_dates(regular_row, date_columns, own_dates))

    rows_by_pib = {normalize_name(row.get("ПІБ", "")): row for row in rows_with_data}
    for pib, dates in tvo_dates_by_pib.items():
        person_row = rows_by_pib.get(pib)
        if person_row is None or not dates:
            continue
        rows.append(_clip_row_to_dates(person_row, date_columns, dates))

    return rows


def _add_commander_intro(doc, month_int, narrator_row):
    add_paragraph_with_style(doc, STATIK["HIGH_COMMANDER"], alignment=WD_ALIGN_PARAGRAPH.RIGHT)
    add_blank_paragraphs(doc, 3)
    add_paragraph_with_style(doc, STATIK["REPORT_TITLE"], alignment=WD_ALIGN_PARAGRAPH.CENTER)
    add_blank_paragraphs(doc, 2)

    text = STATIK["COMMANDER_INTRO"].format(
        rank=narrator_row.get("ЗВАННЯ", ""),
        pib=narrator_row.get("ПІБ", ""),
        posada=narrator_row.get("ПОСАДА", ""),
        month=month_genitive_lower(month_int),
        year=YEAR,
    )
    add_paragraph_with_style(doc, text, first_line_indent=34)
    add_blank_paragraphs(doc, 1)


def process_generate_report_for_commander_money(rows_with_data, rows_with_tvo_data, column_names_personel, changes_file_pairs=()):
    month_int = int(MONTH)
    date_columns = get_date_columns(column_names_personel)
    if not date_columns:
        return None

    period_start, period_end = date_columns[0], date_columns[-1]
    period_str = f"{period_start.strftime('%d.%m.%Y')}-{period_end.strftime('%d.%m.%Y')}"

    # Той самий, хто підписує ЦЕЙ рапорт від першої особи (find_higher_commander на
    # дату підписання - сьогодні) - визначає ще й ЧИЮ участь враховувати нижче
    # (_collect_commander_rows) - обчислюється тут, ОДИН раз, а не окремо перед
    # самим підписом наприкінці.
    today = datetime.now()
    narrator_row = find_higher_commander(rows_with_tvo_data, rows_with_data, HIGHER_COMMANDER_TITLE, today)

    commander_rows = _collect_commander_rows(rows_with_data, rows_with_tvo_data, date_columns, period_start, period_end, narrator_row)
    if not commander_rows:
        print_green("Немає штатного командира чи ТВО за цей місяць - рапорт КБ/ТВО не сформовано.")
        return None

    # СЗЧ, що почався цього місяця - той самий принцип, що й у головному
    # рапорті (process_generate_report_for_get_money): не отримує оплату за
    # жоден звичайний пункт, АЛЕ й далі з'являється в "NOT_PAID".
    commander_rows = _blank_non_szch_days_for_szch_started_this_month(commander_rows, date_columns)

    # ОБЛІК.xlsx більше не містить колонки СТАТУС - див. пояснення в
    # process_generate_report_for_get_money (generate_report_for_get_money.py).
    status_lookup = {}

    categories = _build_categories(commander_rows, date_columns, status_lookup, categories=COMMANDER_MONEY_REPORT_CATEGORIES)
    categories = [(key, rows) for key, rows in categories if rows]

    # commander_and_tvo_only=True - на відміну від головного рапорту (де КБ/ТВО
    # прибираються з розділу "Прошу внести зміни..." повністю), тут навпаки -
    # лишаються ЛИШЕ ЇХНІ зміни за попередні місяці: якщо в самого командира/ТВО
    # є зміни - вони мають потрапити саме в ЙОГО рапорт, а не залишитись невидимими
    # лише тому, що головний рапорт про них не пише - підтверджено користувачем.
    changes_entries = build_changes_entries(
        rows_with_tvo_data, categories=COMMANDER_MONEY_REPORT_CATEGORIES,
        changes_file_pairs=changes_file_pairs, commander_and_tvo_only=True, narrator_row=narrator_row,
    )

    if not categories and not changes_entries:
        print_green("У штатного командира/ТВО немає жодного дня, що потрапляє в рапорт КБ/ТВО - не сформовано.")
        return None

    combined_for_logging = categories + _flatten_changes_rows_for_logging(changes_entries)
    _log_missing_legal_basis(combined_for_logging)
    _log_missing_posada(combined_for_logging)

    doc = Document()
    set_margins(doc, Inches(0.7875), Inches(0.7875), Inches(1.18125), Inches(0.4003))
    add_page_number_header(doc)

    _add_commander_intro(doc, month_int, narrator_row)

    # _NOT_PAID_POINT рендериться ОДРАЗУ ПІСЛЯ звичайних пунктів ОБРАНОГО місяця,
    # ПЕРЕД "Прошу внести зміни..." - як і в головному рапорті (той самий
    # принцип - див. process_generate_report_for_get_money): обраний місяць - у
    # пріоритеті, зміни за попередні місяці йдуть ПІСЛЯ нього.
    not_paid_entry = next((item for item in categories if item[0] == _NOT_PAID_POINT), None)
    categories = [item for item in categories if item[0] != _NOT_PAID_POINT]

    carried_over_cm = 0.0
    point_number = 1
    for template_key, rows in categories:
        carried_over_cm = _add_category_section(
            doc, point_number, template_key, rows, month_int, period_start.day, period_end.day, carried_over_cm,
            section_map_key="COMMANDER_SECTIONS",
        )
        point_number += 1

    if not_paid_entry:
        _, not_paid_rows = not_paid_entry
        carried_over_cm = _add_category_section(
            doc, point_number, _NOT_PAID_POINT, not_paid_rows, month_int, period_start.day, period_end.day, carried_over_cm,
            section_map_key="COMMANDER_SECTIONS", headers=NOT_PAID_TABLE_HEADERS,
        )
        point_number += 1

    point_number, carried_over_cm = _add_changes_sections(doc, point_number, changes_entries, carried_over_cm)

    _add_footer(doc, narrator_row)

    file_name = f"Рапорт ДВ КБ {SHORT_UNIT_BATTALION.upper()} {month_nominative_upper(month_int).lower()} від {period_str}.docx"
    file_path = os.path.join(OUTPUT_DIR, file_name)
    save_docx_safely(doc, file_path)
    print_green(f"- {file_name}")
    return file_path
