import os
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from constants import OUTPUT_DIR, SHORT_UNIT_BATTALION, MONTH, YEAR
from content.money_report_helpers import month_nominative_upper
from content.report_document_reader import extract_timetable_rows_from_report
from content.oblik_style_reference import load_timetable_style_reference
from utils.logging_utils import print_green
from utils.excel_writer import save_workbook_safely

_LABEL_HEADERS = ["ПОСАДА", "ЗВАННЯ", "ПІБ"]
_DATE_NUMBER_FORMAT = "DD.MM.YYYY"

# Передостання колонка табеля - текст "Примітка (підстави)" рапорту, для
# СПЕЦКОНТИНГЕНТУ/100_ВП/100_ШП/NOT_PAID (extract_timetable_rows_from_report вже
# заповнює "ПІДСТАВИ" порожнім рядком для решти людей) - підтверджено
# користувачем: саме так, як колонка "ПІДСТАВИ" у самому ОБЛІК.xlsx (одразу
# після колонок днів).
_BASIS_HEADER = "ПІДСТАВИ"
_BASIS_COL_WIDTH = 40

# Остання колонка - "Дата зникнення безвісти" з рядка(ів) СПЕЦКОНТИНГЕНТУ
# (extract_timetable_rows_from_report заповнює "ДАТА_ЗНИКНЕННЯ" порожнім рядком
# для решти людей) - як і колонка "ДАТА ЗНИКНЕННЯ"/"ДАТА В СТАТУС
# СПЕЦКОНТИНГЕНТУ" в самому ОБЛІК.xlsx, одразу після "ПІДСТАВИ".
_DISAPPEARANCE_HEADER = "ДАТА В СТАТУС СПЕЦКОНТИНГЕНТУ"
_DISAPPEARANCE_COL_WIDTH = 24

# День, який не потрапив у ЖОДНУ таблицю рапорту (тобто в цей день людина не
# отримувала додаткової винагороди взагалі) - табель не повинен лишати клітинку
# порожньою (підтверджено користувачем): такий день - якийсь вид відпустки чи
# відрядження, що рапорт (звіт лише про ВИПЛАТИ) не деталізує - тож замість
# вгадувати ЯКИЙ саме, пишемо всі три можливі варіанти одним текстом.
_UNCLAIMED_DAY_LABEL = "ВП/ВД/ШП"

# Типові (запасні) стилі - використовуються, лише якщо ОБЛІК.xlsx недоступний для
# зчитування стилю (content.oblik_style_reference.load_timetable_style_reference
# повернув None) - табель тоді все одно формується, просто виглядає простіше.
_FALLBACK_LABEL_COL_WIDTH = 22
_FALLBACK_DATE_COL_WIDTH = 6
_FALLBACK_HEADER_FONT = Font(bold=True)
_FALLBACK_ALIGNMENT = Alignment(horizontal="center", vertical="center", wrap_text=True)
_FALLBACK_BORDER = Border(left=Side(style="thin"), right=Side(style="thin"), top=Side(style="thin"), bottom=Side(style="thin"))


def _style_or_fallback(style):
    """Повертає готові до застосування атрибути клітинок - або з реального
    ОБЛІК.xlsx (style, якщо вдалось зчитати), або прості типові (style is None)."""
    if style is None:
        return {
            "header_font": _FALLBACK_HEADER_FONT, "header_fill": None, "header_border": _FALLBACK_BORDER,
            "header_alignment_label": _FALLBACK_ALIGNMENT, "header_alignment_date": _FALLBACK_ALIGNMENT,
            "body_font": None, "body_alignment": _FALLBACK_ALIGNMENT, "body_border": _FALLBACK_BORDER,
            "label_col_width": _FALLBACK_LABEL_COL_WIDTH, "pib_col_width": _FALLBACK_LABEL_COL_WIDTH,
            "date_col_width": _FALLBACK_DATE_COL_WIDTH,
            "header_row_height": None, "body_row_height": None, "keyword_colors": {},
        }
    return style


def _fill_for_value(value, keyword_colors):
    """Заливка клітинки дня за тим самим принципом, що й умовне форматування
    ОБЛІК.xlsx ("клітинка МІСТИТЬ ключове слово") - перший ключ keyword_colors
    (уже впорядкований за пріоритетом самого файлу), що трапляється в тексті
    значення, перемагає; значення без жодного відповідного ключового слова
    (напр. 70/170/10/50 - файл-джерело їх окремо не фарбував) лишається без
    заливки. value тут ніколи не None - непретензований день уже замінено на
    _UNCLAIMED_DAY_LABEL до виклику цієї функції."""
    text = str(value)
    for keyword, color in keyword_colors.items():
        if keyword in text:
            return PatternFill(start_color=color, end_color=color, fill_type="solid")
    return None


def process_generate_timetable(report_path):
    """Табель обліку робочого часу - .xlsx сітка ПОСАДА/ЗВАННЯ/ПІБ x дні РІВНО
    обраного місяця генерації (constants.MONTH/YEAR - те саме питання "Введіть
    номер місяця..."), побудована ЦІЛКОМ з готового рапорту на додаткову
    винагороду (report_path - .docx, обраний користувачем у resources), а НЕ з
    ОБЛІК.xlsx: до нього тут НІДЕ не звертаємось ЗА ДАНИМИ - хто фактично
    потрапив у рапорт (пройшов усі його виключення - штатний командир/ТВО,
    ПРВД тощо) ЗА ОБРАНИЙ МІСЯЦЬ, той і потрапляє в табель; кого нема в рапорті,
    чи є лише за ІНШИЙ місяць (напр. перехідний період пункту "перебування на
    лікуванні", що частково стосується попереднього місяця) - не потрапляє.
    Немає колонки ПІДРОЗДІЛ - рапорт її не містить.

    Вигляд (шрифти/кольори по значенню дня/ширини колонок) - з РЕАЛЬНОГО
    ОБЛІК.xlsx (content.oblik_style_reference), ЛИШЕ заради стилю - жодне
    значення звідти не потрапляє в саму таблицю табеля.

    Значення дня - див. content.report_document_reader.extract_timetable_rows_from_report
    (число пункту рапорту - 30/100/70/170/10/50 - чи "NOT_PAID"/"100_ШП"/"100_ВП"/
    "СПЕЦКОНТИНГЕНТ"/"п.N" для решти випадків) - максимально наближене до сирого
    значення ОБЛІК.xlsx, але не завжди тотожне йому (не всі деталі відновлювані з
    готового тексту рапорту). День, що не потрапив у ЖОДНУ таблицю рапорту (в
    ОБЛІК.xlsx був би, напр., кодом відпустки/відрядження) - "ВП/ВД/ШП"
    (_UNCLAIMED_DAY_LABEL), а не порожня клітинка - рапорт описує лише ВИПЛАТИ, тож
    не може підказати, яка САМЕ з трьох причин відсутності виплати мала місце.

    Передостання колонка - "ПІДСТАВИ" (_BASIS_HEADER) - заповнена для людей із
    СПЕЦКОНТИНГЕНТУ/100_ВП/100_ШП/NOT_PAID (текстом "Примітка (підстави)" з
    їхнього рядка(ів) рапорту), порожня для решти - підтверджено користувачем.
    Остання колонка - "ДАТА В СТАТУС СПЕЦКОНТИНГЕНТУ" (_DISAPPEARANCE_HEADER) -
    заповнена ЛИШЕ для людей із СПЕЦКОНТИНГЕНТУ ("Дата зникнення безвісти" з
    їхнього рядка рапорту), порожня для решти.

    У рапорті взагалі немає жодного розпізнаного пункту з датами ЗА ОБРАНИЙ
    місяць - табель не формується, повертає None."""
    rows_with_data, date_columns = extract_timetable_rows_from_report(report_path, target_month=int(MONTH), target_year=int(YEAR))
    if not rows_with_data:
        print_green(f"У {report_path} не знайдено жодного пункту з періодами участі за {MONTH}.{YEAR} - табель не сформовано.")
        return None

    style = _style_or_fallback(load_timetable_style_reference())
    label_col_widths = [style["label_col_width"], style["label_col_width"], style["pib_col_width"]]

    headers = _LABEL_HEADERS + date_columns + [_BASIS_HEADER, _DISAPPEARANCE_HEADER]

    wb = Workbook()
    ws = wb.active
    ws.title = "Табель"

    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = style["header_font"]
        if style["header_fill"] is not None:
            cell.fill = style["header_fill"]
        cell.border = style["header_border"]
        is_date = isinstance(header, datetime)
        cell.alignment = style["header_alignment_date"] if is_date else style["header_alignment_label"]
        if is_date:
            cell.number_format = _DATE_NUMBER_FORMAT
    if style["header_row_height"]:
        ws.row_dimensions[1].height = style["header_row_height"]

    for row_idx, row in enumerate(rows_with_data, start=2):
        values = (
            [row.get(name, "") for name in _LABEL_HEADERS]
            + [row.get(col) if row.get(col) is not None else _UNCLAIMED_DAY_LABEL for col in date_columns]
            + [row.get(_BASIS_HEADER, ""), row.get("ДАТА_ЗНИКНЕННЯ", "")]
        )
        for col_idx, value in enumerate(values, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            if style["body_font"] is not None:
                cell.font = style["body_font"]
            cell.alignment = style["body_alignment"]
            cell.border = style["body_border"]
            if len(_LABEL_HEADERS) < col_idx <= len(_LABEL_HEADERS) + len(date_columns):
                fill = _fill_for_value(value, style["keyword_colors"])
                if fill is not None:
                    cell.fill = fill
        if style["body_row_height"]:
            ws.row_dimensions[row_idx].height = style["body_row_height"]

    for col_idx, width in enumerate(label_col_widths, start=1):
        ws.column_dimensions[get_column_letter(col_idx)].width = width
    for col_idx in range(len(_LABEL_HEADERS) + 1, len(_LABEL_HEADERS) + len(date_columns) + 1):
        ws.column_dimensions[get_column_letter(col_idx)].width = style["date_col_width"]
    ws.column_dimensions[get_column_letter(len(headers) - 1)].width = _BASIS_COL_WIDTH
    ws.column_dimensions[get_column_letter(len(headers))].width = _DISAPPEARANCE_COL_WIDTH
    ws.freeze_panes = ws.cell(row=2, column=len(_LABEL_HEADERS) + 1).coordinate

    os.makedirs(OUTPUT_DIR, exist_ok=True)
    file_name = f"Табель {SHORT_UNIT_BATTALION.upper()} {month_nominative_upper(int(MONTH)).lower()} {YEAR}.xlsx"
    file_path = os.path.join(OUTPUT_DIR, file_name)
    save_workbook_safely(wb, file_path)
    print_green(f"- {file_name}")
    return file_path
