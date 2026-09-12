import pandas as pd
from datetime import datetime
from docx import Document
from formatting.docx_utils import create_or_clear_output_directory
from utils.date_utils import date_to_str
from content.br_helpers import find_higher_commander
from content.br_general_catalog import generate_content_br_general
from content.money_report_helpers import get_date_columns, _exclude_by_szch_started_this_month
from constants import (
    OUTPUT_DIR, OUTPUT_DIR_BR, OUTPUT_DIR_EXTRACT_BR, OUTPUT_DIR_COMBAT_LOG_EXTRACT_WAR,
    HIGHER_COMMANDER_TITLE, COMMANDER_TITLE, OUTPUT_DIR_BR_SAVE, NUMBER_OF_DOCUMENTS_BRS_EVERY_DAY,
    NUMBER_OF_DOCUMENTS_BRS_EVERY_WEEK, MONTH, YEAR,
)
from generators.generate_documents_br_every_day import generate_documents_br_general
from generators.generate_documents_br_weekly_task import generate_documents_br
from generators.generate_documents_combat_log_extract_war_every_week import generate_documents_combat_log_extract_war_every_week
from generators.generate_br_save_army import generate_documents_br_save
from generators.generate_extract_log_war_general_br_every_day import generate_extract_log_war_general_br_every_day

def process_generate_br(column_names_personel, rows_with_data, rows_with_tvo_data, rows_with_dowries_data, city, coordinates, rows_with_pridani_sheet_data=()):
    create_or_clear_output_directory(OUTPUT_DIR, OUTPUT_DIR_BR, OUTPUT_DIR_BR_SAVE, OUTPUT_DIR_EXTRACT_BR, OUTPUT_DIR_COMBAT_LOG_EXTRACT_WAR)

    # Аркуш "ПРИДАНІ" У САМОМУ ОБЛІК.xlsx (constants.PRIDANI_SHEET_NAME) - той
    # самий формат рядків, що й звичайний особовий склад (ПІДРОЗДІЛ/ПОСАДА/
    # ЗВАННЯ/ПІБ + дні з тими ж значеннями 30/70/100/170 тощо, ключовані ТИМИ Ж
    # datetime, що й column_names_personel) - тож просто ДОДАЄМО ці рядки до
    # rows_with_data ще ДО решти обробки: приданий особовий склад отримує БР/
    # витяги з ЖБД за свої дні 100/70/170 через ТУ САМУ логіку категоризації
    # (content/br_general_catalog.py), без жодних змін там. НЕ потрапляє в
    # рапорт на додаткову винагороду - money-звіти НЕ отримують ці рядки
    # взагалі (окремий, не об'єднаний список у GeneratorInputs/user_input.py) -
    # підтверджено користувачем.
    rows_with_data = [*rows_with_data, *rows_with_pridani_sheet_data]

    # В/сл, хто пішов у СЗЧ ПРОТЯГОМ поточного місяця - не потрапляє в БР
    # документи взагалі (раніше прибиралось ще в user_input.py, СПІЛЬНОЮ
    # точкою для БР і рапорту на додаткову винагороду - тепер рапорт застосовує
    # ІНШЕ правило, див. _blank_non_szch_days_for_szch_started_this_month у
    # generate_report_for_get_money.py, тож фільтрація рознесена по викликачах).
    rows_with_data = _exclude_by_szch_started_this_month(rows_with_data, get_date_columns(column_names_personel))

    # ОБЛІК.xlsx більше не містить колонки СТАТУС - див. пояснення в
    # generate_report_for_get_money.process_generate_report_for_get_money.
    status_lookup = {}

    for col_name in column_names_personel:
        if isinstance(col_name, datetime):
            higher_commander_data = find_higher_commander(rows_with_tvo_data, rows_with_data, HIGHER_COMMANDER_TITLE, col_name)
            commander_data = find_higher_commander(rows_with_tvo_data, rows_with_data, COMMANDER_TITLE, col_name)

            date = pd.to_datetime(col_name, errors='coerce')
            if pd.notna(date) and date_to_str(date) in NUMBER_OF_DOCUMENTS_BRS_EVERY_DAY: # type: ignore
                section_lists = generate_content_br_general(rows_with_data, rows_with_dowries_data, col_name, higher_commander_data, status_lookup)
                generate_documents_br_general(Document(), col_name, rows_with_data, section_lists, higher_commander_data, city, coordinates)
                generate_documents_br_general(Document(), col_name, rows_with_data, section_lists, higher_commander_data, city, coordinates, commander_data, is_extract_br=True)
                previous_date = col_name - pd.DateOffset(days=1)
                previous_section_lists = generate_content_br_general(rows_with_data, rows_with_dowries_data, previous_date, higher_commander_data, status_lookup)
                generate_extract_log_war_general_br_every_day(Document(), col_name, commander_data, section_lists, previous_section_lists, city, coordinates)

            generate_documents_br_save(Document(), col_name, higher_commander_data, city, coordinates, OUTPUT_DIR_BR_SAVE)

            # "ЗАВДАННЯ" (generate_documents_br) і не-щоденний "Витяг з ЖБД"
            # (generate_documents_combat_log_extract_war_every_week) генеруються лише для дат із
            # ФАКТИЧНИМ БР бат (NUMBER_OF_DOCUMENTS_BRS_EVERY_WEEK, поле 'бат' непусте) - а не за
            # умовною сіткою "1-ше число чи кожні 7 днів": реальні БР видаються
            # нерівномірно, тож генерація завжди прив'язана до того, коли БР дійсно
            # видано (period_text/variant у самих генераторах рахується так само -
            # див. get_bat_period_and_variant).
            if col_name.month == int(MONTH) and col_name.year == int(YEAR) and NUMBER_OF_DOCUMENTS_BRS_EVERY_WEEK.get(date_to_str(col_name), {}).get('бат'):
                generate_documents_br(Document(), col_name, higher_commander_data, city, coordinates, OUTPUT_DIR_BR)
                generate_documents_br(Document(), col_name, higher_commander_data, city, coordinates, OUTPUT_DIR_EXTRACT_BR, is_extract_br=True)
                generate_documents_combat_log_extract_war_every_week(Document(), col_name, commander_data)

    return True
