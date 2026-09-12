"""Генерує МІНІМАЛЬНИЙ приклад resources/ОБЛІК.xlsx (кілька вигаданих рядків,
щоб було видно точну структуру) - не чіпає файл, якщо він уже існує.
Запуск: python resources_example/generate_sample.py (з кореня проєкту)."""
import os
from datetime import datetime, timedelta

from openpyxl import Workbook

_MONTH_NAMES = ["СІЧЕНЬ", "ЛЮТИЙ", "БЕРЕЗЕНЬ", "КВІТЕНЬ", "ТРАВЕНЬ", "ЧЕРВЕНЬ",
                "ЛИПЕНЬ", "СЕРПЕНЬ", "ВЕРЕСЕНЬ", "ЖОВТЕНЬ", "ЛИСТОПАД", "ГРУДЕНЬ"]

RESOURCES_DIR = os.path.join(os.path.dirname(__file__), "..", "resources")
OUT_PATH = os.path.join(RESOURCES_DIR, "ОБЛІК.xlsx")

# Вигадані люди для прикладу - НІКОЛИ не вставляйте сюди реальні ПІБ.
PEOPLE = [
    ("Підрозділ 1", "Командир взводу", "лейтенант", "ПЕРШИЙ Перший Першович"),
    ("Підрозділ 2", "Стрілець", "солдат", "ДРУГИЙ Другий Другович"),
    ("Підрозділ 3", "Санітар", "старший солдат", "ТРЕТІЙ Третій Третьович"),
]


def main():
    if os.path.exists(OUT_PATH):
        print(f"{OUT_PATH} вже існує - нічого не роблю (заберіть/перейменуйте файл, якщо хочете новий приклад).")
        return

    today = datetime.now()
    sheet_name = _MONTH_NAMES[today.month - 1]
    first_of_month = today.replace(day=1)
    dates = [first_of_month + timedelta(days=i) for i in range(3)]  # лише 3 дні для прикладу

    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name
    ws.append(["ПІДРОЗДІЛ", "ПОСАДА", "ЗВАННЯ", "ПІБ", *dates, "ПІДСТАВИ"])
    for pidrozdil, posada, zvannya, pib in PEOPLE:
        ws.append([pidrozdil, posada, zvannya, pib, 100, 100, "ВП", ""])
    for col in range(5, 5 + len(dates)):
        ws.cell(row=1, column=col).number_format = "DD.MM.YYYY"

    ws_tvo = wb.create_sheet("ТВО")
    ws_tvo.append(["Підрозділ", "ПОСАДА", "Start", "End", "ЗВАННЯ", "ПІБ", "ТВО"])
    # Порожній (без рядків людей) - так само коректно, як і заповнений.

    os.makedirs(RESOURCES_DIR, exist_ok=True)
    wb.save(OUT_PATH)
    print(f"Створено {OUT_PATH} (аркуш '{sheet_name}' + порожній 'ТВО').")
    print("Замініть вигаданих людей на реальних і заповніть усі дні місяця.")


if __name__ == "__main__":
    main()
