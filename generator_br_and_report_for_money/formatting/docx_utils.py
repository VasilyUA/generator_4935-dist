import os
import re
import shutil
import time
from typing import Union
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ROW_HEIGHT_RULE
from docx.shared import Pt, Inches, RGBColor, Cm
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

from constants import (
    NUMBER_OF_DOCUMENTS_BRS_EVERY_WEEK, SHORT_UNIT_BATTALION, SHORT_UNIT_BRIGADE,
    OUTPUT_DIR_COMBAT_LOG_EXTRACT_WAR, FULL_UNIT_BUT, FULL_MILITARY_UNIT, YEAR,
    MONEY_REPORT_GENERAL_REFERENCES_LOG_WAR, SEQUENCE, SEQUENCE_HIGHER_COMMANDER,
)
from utils.logging_utils import print_red, print_green
from utils.date_utils import to_date
from content.br_helpers import convert_to_short_name, get_number_br
from content.money_report_helpers import MONTH_NAMES_GENITIVE_LOWER

_MONTH_NUMBER_BY_GENITIVE_NAME = {name: number for number, name in MONTH_NAMES_GENITIVE_LOWER.items()}


def set_line_spacing(doc, spacing):
    for paragraph in doc.paragraphs:
        paragraph.paragraph_format.line_spacing = Pt(spacing)


def _close_matching_word_documents(matches):
    """Закриває (без збереження) кожен відкритий у Word документ, для якого
    matches(абсолютний_шлях) - True. Якщо Word не запущено або pywin32
    недоступний — тихо виходить. Спільна основа для _close_word_documents_under
    (усі файли під директорією) та _close_word_document_if_open (рівно один
    файл) - різниця між ними лише в самому предикаті matches."""
    try:
        import pythoncom
        import win32com.client
    except ImportError:
        return

    try:
        pythoncom.CoInitialize()
        try:
            word = win32com.client.GetActiveObject("Word.Application")
        except Exception:
            return
        for doc in list(word.Documents):
            try:
                full_name = os.path.abspath(doc.FullName)
            except Exception:
                continue
            if matches(full_name):
                print_red(f"Закриваю відкритий у Word файл: {os.path.basename(full_name)}")
                doc.Close(SaveChanges=False)
    finally:
        pythoncom.CoUninitialize()


def _close_word_documents_under(directory):
    """Якщо серед відкритих у Word документів є файли з попереднього запуску
    генератора (лежать усередині directory) — закриває їх без збереження, щоб
    shutil.rmtree не впав з PermissionError через залишений відкритим .docx."""
    abs_dir = os.path.abspath(directory) + os.sep
    _close_matching_word_documents(lambda full_name: full_name.startswith(abs_dir))


def _rmtree_retry(path, attempts=5, delay=0.5):
    """shutil.rmtree з повторними спробами — деякі процеси (антивірус, OneDrive)
    тримають файл заблокованим лише короткий час одразу після закриття Word."""
    for attempt in range(attempts):
        try:
            shutil.rmtree(path)
            return
        except PermissionError as e:
            if attempt == attempts - 1:
                raise PermissionError(
                    f"Не вдалось видалити '{path}': файл заблоковано іншою програмою. "
                    f"Закрийте його (напр. Word) і спробуйте ще раз. ({e})"
                ) from e
            time.sleep(delay)


def _close_word_document_if_open(file_path):
    """Як _close_word_documents_under, але для ОДНОГО конкретного файлу (а не всієї
    директорії) - щоб не закривати випадково інші відкриті документи в тій самій папці
    (напр. коли перегенеровуємо лише рапорт на додаткову винагороду, а не весь output/)."""
    abs_path = os.path.abspath(file_path)
    _close_matching_word_documents(lambda full_name: full_name == abs_path)


def save_docx_safely(doc, file_path, attempts=5, delay=0.5):
    """doc.save(), але спершу закриває цей самий файл у Word, якщо він там відкритий
    (напр. користувач переглядав попередній рапорт із такою самою назвою), і повторює
    спробу, якщо файл лишається заблокованим ще коротку мить (антивірус/OneDrive)."""
    _close_word_document_if_open(file_path)
    for attempt in range(attempts):
        try:
            doc.save(file_path)
            return
        except PermissionError as e:
            if attempt == attempts - 1:
                raise PermissionError(
                    f"Не вдалось зберегти '{file_path}': файл заблоковано іншою програмою. "
                    f"Закрийте його (напр. Word) і спробуйте ще раз. ({e})"
                ) from e
            time.sleep(delay)


def create_or_clear_output_directory(directory, output_dir_br, output_dir_br_save, output_dir_extract_br, output_dir_combat_log_extract_war):
    _close_word_documents_under(directory)

    if os.path.exists(output_dir_combat_log_extract_war):
        _rmtree_retry(output_dir_combat_log_extract_war)
    if os.path.exists(output_dir_br):
        _rmtree_retry(output_dir_br)
    if os.path.exists(output_dir_br_save):
        _rmtree_retry(output_dir_br_save)
    if os.path.exists(output_dir_extract_br):
        _rmtree_retry(output_dir_extract_br)
    if os.path.exists(directory):
        _rmtree_retry(directory)

    os.makedirs(directory)
    os.makedirs(output_dir_br)
    os.makedirs(output_dir_br_save)
    os.makedirs(output_dir_extract_br)
    os.makedirs(output_dir_combat_log_extract_war)


def _set_paragraph_mark_font(paragraph, font_name, font_size):
    # Без цього порожній абзац (без жодного run-у з текстом) показує шрифт теми (Cambria 11)
    # замість заданого, бо Word бере шрифт кінцевого маркера абзацу (pPr/rPr), а не run.font.
    rPr = OxmlElement('w:rPr')
    rFonts = OxmlElement('w:rFonts')
    rFonts.set(qn('w:ascii'), font_name)
    rFonts.set(qn('w:hAnsi'), font_name)
    rFonts.set(qn('w:eastAsia'), font_name)
    rPr.append(rFonts)
    color = OxmlElement('w:color')
    color.set(qn('w:val'), '000000')
    rPr.append(color)
    sz = OxmlElement('w:sz')
    sz.set(qn('w:val'), str(int(font_size * 2)))
    rPr.append(sz)
    paragraph._p.get_or_add_pPr().append(rPr)


def add_paragraph_with_style(doc, text, font_name='Times New Roman', font_size=13.5, bold: Union[bool, list] = False, alignment=WD_ALIGN_PARAGRAPH.JUSTIFY, first_line_indent=0, format_tabs=False, heading_level=None):
    paragraph = doc.add_paragraph()
    if heading_level is not None:
        # Позначає абзац як заголовок (з'являється у структурі/навігації документа),
        # але весь вигляд (шрифт, розмір, жирність, вирівнювання) лишається таким самим,
        # бо нижче все одно виставляється напряму через прямі runs/paragraph_format,
        # які завжди переважають над стилем Heading.
        paragraph.style = f"Heading {heading_level}"
        # ВИНЯТОК - keep_with_next: стиль "Heading N" за замовчуванням має keepNext=True
        # (успадковано від базового шаблону Word), а paragraph_format сам по собі цього
        # не перевизначає (лишається None -> бере значення зі стилю). Без явного False
        # цей абзац примусово "приклеювався" б до наступного вмісту (в цьому проєкті -
        # до таблиці, разом із її власним ланцюжком "шапка+перший рядок") - і якщо весь
        # цей ланцюжок не вміщується на залишок сторінки, він ЦІЛИКОМ переноситься на
        # нову, лишаючи порожнє місце на попередній (саме такий баг був знайдений).
        paragraph.paragraph_format.keep_with_next = False
    paragraph.alignment = alignment
    paragraph.paragraph_format.line_spacing = 1.0
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.space_before = Pt(0)

    if first_line_indent > 0:
        paragraph.paragraph_format.first_line_indent = Pt(first_line_indent)

    if format_tabs:
        width = Inches(doc.sections[0].page_width.inches - (doc.sections[0].left_margin.inches + doc.sections[0].right_margin.inches))
        paragraph.paragraph_format.tab_stops.add_tab_stop(width, WD_TAB_ALIGNMENT.RIGHT)

    def add_run(t, is_bold):
        run = paragraph.add_run(t)
        run.font.name = font_name
        run._element.rPr.rFonts.set(qn('w:eastAsia'), font_name)
        run.font.size = Pt(font_size)
        run.font.color.rgb = RGBColor(0, 0, 0)
        run.font.bold = is_bold
        return run

    if isinstance(bold, list):
        # bold — список слів які треба виділити жирним
        remaining = text
        while remaining:
            earliest_word = None
            earliest_pos = len(remaining)

            for word in bold:
                pos = remaining.find(word)
                if pos != -1 and pos < earliest_pos:
                    earliest_pos = pos
                    earliest_word = word

            if earliest_word is None:
                # жирних слів більше немає — додаємо залишок як звичайний текст
                add_run(remaining, False)
                break

            # текст до жирного слова
            if earliest_pos > 0:
                add_run(remaining[:earliest_pos], False)

            # жирне слово
            add_run(earliest_word, True)

            remaining = remaining[earliest_pos + len(earliest_word):]
    else:
        # bold — True або False — весь текст однаковий
        add_run(text, bool(bold))

    _set_paragraph_mark_font(paragraph, font_name, font_size)

    return paragraph


def add_custom_heading(doc, text, font_name='Times New Roman', heading=1, color_rgb=RGBColor(0, 0, 0), font_size=16, alignment=WD_ALIGN_PARAGRAPH.JUSTIFY, first_line_indent=0):
    paragraph = doc.add_paragraph()
    paragraph.style = f"Heading {heading}"

    paragraph.alignment = alignment
    paragraph.paragraph_format.line_spacing = 1.0  # Одинарний міжрядковий
    paragraph.paragraph_format.space_before = Pt(0)  # Взагалі нуль перед абзацом
    paragraph.paragraph_format.space_after = Pt(0)  # Взагалі нуль після абзацу
    paragraph.paragraph_format.first_line_indent = Pt(first_line_indent)

    run = paragraph.add_run(text)
    run.bold = True
    run.font.name = font_name
    run._element.rPr.rFonts.set(qn('w:eastAsia'), font_name)
    run.font.color.rgb = color_rgb
    run.font.size = Pt(font_size)

    return paragraph


# Спільні "будівельні блоки" абзаців для generate_*.py (раніше дублювались у кожному файлі окремо)
def add_body_paragraph(doc, text):
    add_paragraph_with_style(doc, text, first_line_indent=34)


def add_section_heading(doc, text):
    add_custom_heading(doc, text, font_size=14, alignment=WD_ALIGN_PARAGRAPH.LEFT, first_line_indent=34)


def add_blank_paragraphs(doc, count=1):
    for _ in range(count): add_paragraph_with_style(doc, "")


def add_recipients_preamble(doc):
    """Спільний блок отримувачів БР-документів (вищий командир + "КОМАНДИРУ
    {SEQUENCE}") - однаковий для generate_documents_br_every_day.py,
    generate_documents_br_weekly_task.py, generate_br_save_army.py."""
    add_body_paragraph(doc, ', '.join(SEQUENCE_HIGHER_COMMANDER))
    add_body_paragraph(doc, f"КОМАНДИРУ {', '.join(SEQUENCE)}")
    add_blank_paragraphs(doc)


def add_signature_block(doc, role, data):
    add_paragraph_with_style(doc, f"{role} {FULL_UNIT_BUT} {FULL_MILITARY_UNIT}")
    add_paragraph_with_style(doc, f"{data.get('ЗВАННЯ', '')}\t{convert_to_short_name(data.get('ПІБ', ''))}", format_tabs=True)


def add_titled_section(doc, title, text, blanks=1):
    add_section_heading(doc, title)
    add_body_paragraph(doc, text)
    add_blank_paragraphs(doc, blanks)


# -> група 1 "359дск/6", група 2 - дата, цифрами ("23.06.2026") АБО словами
# ("23 червня 2026", родовий відмінок місяця - MONTH_NAMES_GENITIVE_LOWER) -
# реальні lines одних записів constants.py написані цифрами, інших - словами
# (так природньо пишуть у справжніх документах) - підтверджено користувачем
# (падало на "26 липня 2026" - лише числовий формат розпізнавався раніше).
# "№\s*" (а не просто "№") - деякі lines мають пробіл між № і номером
# ("№ 359дск/8"), інші - без ("№359дск/7") - обидва варіанти й далі мають
# розпізнаватись однаково.
_LOG_WAR_REFERENCE_RE = re.compile(
    r"№\s*(\S+?)\s*від\s+(\d{2}\.\d{2}\.\d{4}|\d{1,2}\s+[а-яіїєґ]+\s+\d{4})", re.IGNORECASE,
)


def _normalize_log_war_reference_date(raw_date_text):
    """Приводить дату з посилання "№... від ..." (_LOG_WAR_REFERENCE_RE - цифрами
    чи словами) до єдиного канонічного "DD.MM.YYYY", яким і так уже написані
    записи, введені цифрами - щоб результат виглядав однаково незалежно від
    того, як саме користувач написав ЦЕЙ конкретний рядок."""
    raw_date_text = raw_date_text.strip()
    if re.fullmatch(r"\d{2}\.\d{2}\.\d{4}", raw_date_text):
        return raw_date_text
    day_str, month_name, year_str = raw_date_text.split()
    month_int = _MONTH_NUMBER_BY_GENITIVE_NAME[month_name.lower()]
    return f"{int(day_str):02d}.{month_int:02d}.{year_str}"


def _resolve_log_war_reference_text(safe_col_name):
    """Знаходить у constants.MONEY_REPORT_GENERAL_REFERENCES_LOG_WAR запис, чий діапазон
    [start; end] містить дату конкретного витяга з ЖБД (safe_col_name), і дістає з його
    "lines" номер справи та дату ("№359дск/6 від 23.06.2026") - САМЕ ця, а не фіксована
    одна пара, підставляється в заголовок "за номенклатурою №... від ... року", бо різні
    періоди місяця можуть вестись за різними номерами справи (номенклатури).

    Якщо для дати не знайшлось жодного відповідного запису (напр. дата випадає поза всіма
    діапазонами) - повертає заповнювач "№_____ від ___.___.{YEAR} року", а не падає."""
    date = to_date(safe_col_name)
    for reference in MONEY_REPORT_GENERAL_REFERENCES_LOG_WAR:
        start = to_date(reference["start"])
        end = to_date(reference["end"]) if reference.get("end") else date
        if not (start <= date <= end):
            continue
        for line in reference.get("lines", []):
            match = _LOG_WAR_REFERENCE_RE.search(line)
            if match:
                return f"№{match.group(1)} від {_normalize_log_war_reference_date(match.group(2))}"

    print_red(f"Не знайдено запису MONEY_REPORT_GENERAL_REFERENCES_LOG_WAR для дати {safe_col_name} - вставлено заповнювач.")
    return f"№_____ від ___.___.{YEAR}"


def generate_combat_log_extract_title(doc, safe_col_name):
    reference_text = _resolve_log_war_reference_text(safe_col_name)
    for text in ("ВИТЯГ ІЗ ЖУРНАЛУ БОЙОВИХ ДІЙ", f"{FULL_UNIT_BUT} {FULL_MILITARY_UNIT} за номенклатурою {reference_text} року"):
        add_paragraph_with_style(doc, text, alignment=WD_ALIGN_PARAGRAPH.CENTER, font_size=14)


def bold_cell_heading_lines(cell, prefixes):
    """Робить ЖИРНИМ кожен абзац комірки, чий текст починається з одного з
    prefixes (напр. ("4.", "5.", "6.")) - для нумерованих пунктів усередині
    багаторядкової комірки таблиці (Витяг з ЖБД -
    generate_documents_combat_log_extract_war_every_week.py,
    generate_extract_log_war_general_br_every_day.py), де кожен такий пункт
    має виглядати як заголовок розділу, а не звичайний текст."""
    for paragraph in cell.paragraphs:
        if paragraph.text.startswith(tuple(prefixes)):
            for run in paragraph.runs:
                run.font.bold = True


def setup_combat_log_extract_document(doc, safe_col_name):
    """Спільне налаштування документа "Витяг з ЖБД" - однакове для щоденного і
    щотижневого варіантів (generate_extract_log_war_general_br_every_day.py,
    generate_documents_combat_log_extract_war_every_week.py): поля, альбомна
    орієнтація, міжрядковий інтервал і заголовок. На відміну від звичайних БР
    (generate_head_documents_br) - БЕЗ грифу "ДЛЯ СЛУЖБОВОГО КОРИСТУВАННЯ" в
    колонтитулах, підтверджено користувачем."""
    set_margins(doc, Inches(0.786), Inches(0.786), Inches(1.18), Inches(0.395))
    switch_landscape(doc)
    set_line_spacing(doc, 1.0)
    generate_combat_log_extract_title(doc, safe_col_name)


COMBAT_LOG_EXTRACT_TABLE_HEADERS = ["Дата, час", "Завдання військ та стисле висвітлення ходу бойових дій"]


def add_combat_log_extract_table(doc, safe_col_name, text):
    """Додає єдину таблицю (один рядок даних) документа "Витяг з ЖБД" і робить
    пункти 4./5./6. усередині неї заголовками - спільний завершальний крок
    generate_main_combat_log_extract_war_every_day/every_week."""
    create_table(doc, COMBAT_LOG_EXTRACT_TABLE_HEADERS, [[safe_col_name, text]], cant_split=False)
    bold_cell_heading_lines(doc.tables[-1].cell(1, 1), ("4.", "5.", "6."))


def check_number_file_is_exist(num__doc):
    return "" if num__doc is None else f"№{num__doc} "


def save_combat_log_extract_war(doc, safe_col_name, commander_data, number_objects=NUMBER_OF_DOCUMENTS_BRS_EVERY_WEEK, file_suffix=""):
    add_blank_paragraphs(doc, 2)
    add_paragraph_with_style(doc, "Згідно з оригіналом:")
    role = "Тимчасово виконуючий обов'язки начальника штабу - заступник командира" if commander_data.get('ТВО', "") else "Начальник штабу - заступник командира"
    add_signature_block(doc, role, commander_data)
    num_bat = check_number_file_is_exist(get_number_br(safe_col_name, number_objects).get('бат', None))
    file_name = f"Витяг з ЖБД для БР - {num_bat}{safe_col_name}{file_suffix}.docx"
    doc.save(os.path.join(OUTPUT_DIR_COMBAT_LOG_EXTRACT_WAR, file_name))
    print_green(file_name)


def set_margins(doc, top, bottom, left, right):
    sections = doc.sections

    for section in sections:
        section.page_height = Inches(11.69)  # 297 мм
        section.page_width = Inches(8.27)    # 210 мм
        section.top_margin = top
        section.bottom_margin = bottom
        section.left_margin = left
        section.right_margin = right


def add_text_for_header_and_footer_document_br(doc):
    section = doc.sections[0]
    section.different_first_page_header_footer = True  # 🔧 різний колонтитул для першої сторінки

    # Верхній колонтитул (з другої сторінки), нижній колонтитул (усі сторінки), нижній колонтитул першої сторінки
    for para in (section.header.paragraphs[0], section.footer.paragraphs[0], section.first_page_footer.paragraphs[0]):
        para.text = "ДЛЯ СЛУЖБОВОГО КОРИСТУВАННЯ"
        para.alignment = 1  # по центру
        para.style.font.size = Pt(14)
        para.style.font.name = "Times New Roman"


def generate_head_documents_br(doc):
    """Спільний "шапковий" блок БР-документів (ЗАВДАННЯ, застосування безпеки,
    щоденне/щотижневе бойове розпорядження) - поля, колонтитул, міжрядковий
    інтервал, гриф "Для службового користування"/"Прим.№1" і відступ перед
    основним текстом. Однаковий для generate_documents_br_every_day.py,
    generate_documents_br_weekly_task.py, generate_br_save_army.py."""
    set_margins(doc, Inches(0.786), Inches(0.786), Inches(1.18), Inches(0.395))
    add_text_for_header_and_footer_document_br(doc)
    set_line_spacing(doc, 1.0)
    add_paragraph_with_style(doc, "Для службового користування", font_size=12, first_line_indent=280)
    add_paragraph_with_style(doc, "Прим.№1", font_size=12, first_line_indent=280)
    add_blank_paragraphs(doc, 3)


def add_page_number_header(doc, font_name="Times New Roman", font_size=12):
    """Додає в колонтитул (спільний для всіх сторінок, включно з першою) поле номера
    сторінки Word (інструкція "PAGE" - оновлюється автоматично при перегляді/друку),
    по центру, без жодного іншого тексту - лише число, як у зразку документа."""
    header_paragraph = doc.sections[0].header.paragraphs[0]
    header_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    header_paragraph.paragraph_format.space_after = Pt(0)
    header_paragraph.paragraph_format.space_before = Pt(0)

    def _new_run():
        run = header_paragraph.add_run()
        run.font.name = font_name
        run._element.rPr.rFonts.set(qn("w:eastAsia"), font_name)
        run.font.size = Pt(font_size)
        return run

    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    _new_run()._r.append(fld_begin)

    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    _new_run()._r.append(instr_text)

    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    _new_run()._r.append(fld_end)


def switch_landscape(doc):
    # Змінюємо орієнтацію на альбомну
    section = doc.sections[-1]
    new_width, new_height = section.page_height, section.page_width
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = new_width
    section.page_height = new_height


def _set_table_cell_margins_cm(tbl, margin_cm):
    """Встановлює ОДНАКОВІ внутрішні відступи (верх/низ/ліво/право) усіх комірок
    таблиці - OOXML w:tblCellMar (стандартний API python-docx цього не має, за
    замовчуванням Word застосовує свої власні ненульові відступи)."""
    margin_dxa = str(Cm(margin_cm).twips)
    tbl_cell_mar = OxmlElement('w:tblCellMar')
    for tag in ('w:top', 'w:left', 'w:bottom', 'w:right'):
        node = OxmlElement(tag)
        node.set(qn('w:w'), margin_dxa)
        node.set(qn('w:type'), 'dxa')
        tbl_cell_mar.append(node)
    tbl._tbl.tblPr.append(tbl_cell_mar)


def create_table(doc, headers: list=[], data_rows: list=[], widths_cm: list=[3, 21], font_size=12, font_name='Times New Roman', italic=False, color_rgb=(0,0,0), cell_alignment=None, rotated_header_indices=(), header_row_height_cm=None, header_bold=True, body_first_line_bold=True, body_italic_first_cell=True, cant_split=True, cell_margin_cm=None):
    """
    Створює таблицю (1 рядок — заголовок, 1-??? рядки з даними).
    :param doc: об’єкт Document()
    :param data_rows: список з списком значень для рядків [["Дані 1", "Дані 2", "Дані 3"], ["Дані 1", "Дані 2", "Дані 3"]]
    :param headers: список з кількість заголовків кількості стовпчиків ["Заголовок 1", "Заголовок 2", "Заголовок 3"]
    :param widths_cm: список зі кількістю ширини стовпчиків в сантиметрах [2.5, 5.0, 3.0]
    :param cell_alignment: якщо задано (напр. WD_ALIGN_PARAGRAPH.CENTER) - застосовується як
        горизонтальне вирівнювання ДО всіх абзаців (заголовок + дані), а комірки (заголовок і дані)
        також центруються по вертикалі (w:vAlign); якщо None (за замовчуванням) - жодне з двох не
        чіпається, щоб не зламати вигляд уже існуючих документів, які use create_table без цього параметра.
    :param rotated_header_indices: індекси колонок (0-відлік), заголовок яких повертається знизу
        вгору (w:textDirection btLr) - для дуже вузьких колонок на кшталт "Кількість днів".
    :param header_row_height_cm: якщо задано - мінімальна висота (см) рядка заголовків
        (рядок все одно може вирости більше, якщо текст не вміщується).
    :param header_bold: жирний текст заголовків (за замовчуванням True - як раніше).
    :param body_first_line_bold: жирний перший рядок кожної багаторядкової комірки з даними
        (за замовчуванням True - як раніше).
    :param body_italic_first_cell: курсив + відступ для комірки (рядок 1, колонка 1) даних
        (за замовчуванням True - як раніше; той самий особливий випадок, що й був).
    :param cant_split: забороняти розрив рядка таблиці між сторінками (за замовчуванням True,
        як і раніше - підходить для таблиць із КОРОТКИМИ рядками, напр. рапорт на додаткову
        винагороду, де кожен рядок - один військовослужбовець і не повинен "осиротіти" на межі
        сторінки). Витяги з ЖБД (generate_documents_combat_log_extract_war_every_week.py,
        generate_extract_log_war_general_br_every_day.py) мають РІВНО один рядок даних із текстом на
        КІЛЬКА сторінок - для такого рядка "не розривати між сторінками" виконати неможливо
        (він і так завжди більший за одну сторінку), і Word замість розриву як слід зображує
        порожню сторінку та накладення тексту на колонтитул. Такі виклики мають передавати
        cant_split=False, щоб Word розривав рядок природно, як звичайний текст.
    :param cell_margin_cm: якщо задано - внутрішні відступи (см) усіх комірок таблиці
        (верх/низ/ліво/право однаково), напр. 0 для суцільного тексту без полів. Якщо
        None (за замовчуванням) - Word-івські типові відступи лишаються без змін.
    """

    n = len(headers)
    # Перевірка: ширини та кожен data_row мають правильну довжину
    if len(widths_cm) != n or any(len(row) != n for row in data_rows):
        print_red(f"❌ Некоректні параметри: columns={n}, widths={len(widths_cm)}, rows={[(i+1,len(r)) for i,r in enumerate(data_rows)]}")
        return

    tbl = doc.add_table(rows=1 + len(data_rows), cols=n)
    tbl.style = 'Table Grid'
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = tbl.allow_autofit = False

    if cell_margin_cm is not None:
        _set_table_cell_margins_cm(tbl, cell_margin_cm)

    # Забороняємо розрив рядка таблиці між сторінками - якщо рядок не вміщується
    # до кінця сторінки, Word переносить його ЦІЛИКОМ на наступну, а не розриває
    # посередині комірки (як у зразку - там теж стоїть cantSplit). Див. docstring
    # параметра cant_split - для рядків, що завідомо більші за одну сторінку, це
    # прапорець свідомо не застосовується.
    if cant_split:
        for row in tbl.rows:
            row_tr_pr = row._tr.get_or_add_trPr()
            cant_split_el = OxmlElement('w:cantSplit')
            row_tr_pr.append(cant_split_el)

    if header_row_height_cm is not None:
        tbl.rows[0].height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
        tbl.rows[0].height = Cm(header_row_height_cm)

    # Задаємо ширину колонок. cell.width сам по собі НЕ впливає на те, як Word
    # реально малює таблицю - за це відповідає tblGrid, який оновлюється лише
    # через tbl.columns[i].width. Без цього рядка Word ігнорує задані ширини й
    # ділить таблицю на рівні колонки (навіть при tblLayout=fixed).
    for i, w in enumerate(widths_cm):
        tbl.columns[i].width = Cm(w)
        for cell in tbl.columns[i].cells:
            cell.width = Cm(w)

    # Заповнення заголовків
    for i, h in enumerate(headers):
        header_paragraph = tbl.rows[0].cells[i].paragraphs[0]
        # "Тримати з наступним" - забороняє розрив сторінки МІЖ рядком заголовків і першим
        # рядком даних. Без цього Word може лишити сам заголовок один унизу сторінки,
        # а всі дані таблиці перенести на наступну - разом вони підуть на наступну сторінку.
        if data_rows:
            header_paragraph.paragraph_format.keep_with_next = True
        run = header_paragraph.add_run(h)
        run.bold = header_bold
        run.font.name = font_name
        run._element.rPr.rFonts.set(qn('w:eastAsia'), font_name)
        run.font.size = Pt(font_size)
        if cell_alignment is not None:
            header_paragraph.alignment = cell_alignment
            v_align = OxmlElement('w:vAlign')
            v_align.set(qn('w:val'), 'center')
            tbl.rows[0].cells[i]._tc.get_or_add_tcPr().append(v_align)

        if i in rotated_header_indices:
            text_direction = OxmlElement('w:textDirection')
            text_direction.set(qn('w:val'), 'btLr')
            tbl.rows[0].cells[i]._tc.get_or_add_tcPr().append(text_direction)

    # Дані
    for r, row in enumerate(data_rows, start=1):
        for c, val in enumerate(row):
            cell = tbl.rows[r].cells[c]
            cell.text = ""  # очищаємо попередній вміст

            if cell_alignment is not None:
                v_align = OxmlElement('w:vAlign')
                v_align.set(qn('w:val'), 'center')
                cell._tc.get_or_add_tcPr().append(v_align)

            lines = str(val).split('\n')
            for idx, line in enumerate(lines):
                # Для першої строки переходимо до існуючого параграфу,
                # для інших — додаємо нові
                p = cell.paragraphs[0] if idx == 0 else cell.add_paragraph()
                if cell_alignment is not None:
                    p.alignment = cell_alignment
                run = p.add_run(line)

                # Загальне форматування
                run.font.name = font_name
                run._element.rPr.rFonts.set(qn('w:eastAsia'), font_name)
                run.font.size = Pt(font_size)
                run.font.color.rgb = RGBColor(*color_rgb)
                run.font.italic = italic
                if idx == 0 and body_first_line_bold:
                    run.font.bold = True  # жирний лише перший рядок

                # 🎯 Відступ першого рядка лише для потрібного абзацу:
                if r == 1 and c == 1 and body_italic_first_cell:
                    run.font.italic = True
                    p.paragraph_format.first_line_indent = Cm(0.6)
