import sys
from datetime import datetime
from pathlib import Path

sys.path.append(str(Path(__file__).resolve().parent.parent))

from docx import Document
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, Side

import generators.generate_timetable as gt


def _write_report(path, sections):
    """sections - список (heading_text, data_rows); data_rows - [посада, звання,
    піб, період, дні, підстава]. Той самий мінімальний формат таблиці (7 колонок,
    ПІБ - колонка 3, ПЕРІОД - колонка 4), що й реальні таблиці рапорту."""
    doc = Document()
    for heading_text, data_rows in sections:
        doc.add_heading(heading_text, level=1)
        table = doc.add_table(rows=1 + len(data_rows), cols=7)
        for row_idx, values in enumerate(data_rows, start=1):
            for col_idx, value in enumerate(values, start=1):
                table.rows[row_idx].cells[col_idx].text = value
    doc.save(str(path))
    return str(path)


def test_process_generate_timetable_returns_none_without_recognized_sections(tmp_path, monkeypatch):
    monkeypatch.setattr(gt, "MONTH", "06")
    monkeypatch.setattr(gt, "YEAR", "2026")
    path = _write_report(tmp_path / "report.docx", [
        ("11.1 Виключити пункти в додатку 1:", [
            ["Оператор", "старший матрос", "ПЕРШИЙ Перший Перший", "01.06.2026-15.06.2026", "15", ""],
        ]),
    ])

    assert gt.process_generate_timetable(path) is None


def test_process_generate_timetable_filters_to_selected_month_and_builds_full_calendar_grid(tmp_path, monkeypatch):
    """Регресія: рапорт може містити пункт, що частково стосується ІНШОГО місяця
    (напр. перехідний період "перебування на лікуванні") - людина, у якої є дні
    ЛИШЕ поза обраним місяцем генерації, взагалі не повинна потрапити в табель,
    а колонки дат мають охоплювати ЦІЛИЙ обраний місяць (як і ОБЛІК.xlsx), а не
    лише дні, для яких у рапорті є дані."""
    monkeypatch.setattr(gt, "OUTPUT_DIR", str(tmp_path))
    monkeypatch.setattr(gt, "MONTH", "07")
    monkeypatch.setattr(gt, "YEAR", "2026")

    path = _write_report(tmp_path / "report.docx", [
        ("1. Виплатити додаткову винагороду у розмірі 30 000 грн. 00 коп. військовослужбовцям", [
            ["Стрілець", "сержант", "ДРУГИЙ Другий Другий", "01.07.2026-02.07.2026", "2", ""],
        ]),
        ("8. Виплатити особовому складу за період перебування на лікуванні", [
            ["старший технік", "штаб-сержант", "ТРЕТІЙ Третій Третій", "22.06.2026-30.06.2026", "9", ""],
        ]),
    ])

    file_path = gt.process_generate_timetable(path)
    ws = load_workbook(file_path).active

    header = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
    assert header[3] == datetime(2026, 7, 1)
    assert header[-3] == datetime(2026, 7, 31)
    assert header[-2] == "ПІДСТАВИ"
    assert header[-1] == "ДАТА В СТАТУС СПЕЦКОНТИНГЕНТУ"
    assert len(header) == 3 + 31 + 2

    pib_column = [row[2].value for row in ws.iter_rows(min_row=2)]
    assert pib_column == ["ДРУГИЙ Другий Другий"]


def test_process_generate_timetable_writes_raw_values_grid(tmp_path, monkeypatch):
    monkeypatch.setattr(gt, "OUTPUT_DIR", str(tmp_path))
    monkeypatch.setattr(gt, "MONTH", "07")
    monkeypatch.setattr(gt, "YEAR", "2026")

    path = _write_report(tmp_path / "report.docx", [
        ("1. Виплатити додаткову винагороду у розмірі 30 000 грн. 00 коп. військовослужбовцям", [
            ["Стрілець", "сержант", "ДРУГИЙ Другий Другий", "01.07.2026-02.07.2026", "2", ""],
        ]),
    ])

    file_path = gt.process_generate_timetable(path)
    ws = load_workbook(file_path).active

    row1 = [cell.value for cell in next(ws.iter_rows(min_row=2, max_row=2))]
    assert row1[:5] == ["Стрілець", "сержант", "ДРУГИЙ Другий Другий", 30, 30]
    # 03.07.2026 - поза періодом участі цієї людини - не порожня клітинка, а
    # позначка "ВП/ВД/ШП" (рапорт описує лише виплати, тож не підкаже, яка саме
    # з трьох причин відсутності виплати мала місце того дня).
    assert row1[5] == "ВП/ВД/ШП"


def test_process_generate_timetable_fills_pidstavy_and_disappearance_date_for_spetskontyngent(tmp_path, monkeypatch):
    """Колонка "ПІДСТАВИ" заповнюється текстом "Примітка (підстави)" для людей
    із пункту СПЕЦКОНТИНГЕНТ - для решти (звичайний пункт 1, без підстави в
    цьому фікстурі) лишається порожньою. Остання колонка "ДАТА В СТАТУС
    СПЕЦКОНТИНГЕНТУ" заповнюється ЛИШЕ для СПЕЦКОНТИНГЕНТУ ("Дата зникнення
    безвісти" з їхнього рядка рапорту)."""
    monkeypatch.setattr(gt, "OUTPUT_DIR", str(tmp_path))
    monkeypatch.setattr(gt, "MONTH", "07")
    monkeypatch.setattr(gt, "YEAR", "2026")

    doc = Document()
    doc.add_heading("1. Виплатити додаткову винагороду у розмірі 30 000 грн. 00 коп. військовослужбовцям", level=1)
    table1 = doc.add_table(rows=2, cols=7)
    for col_idx, value in enumerate(["Стрілець", "сержант", "ДРУГИЙ Другий Другий", "01.07.2026-01.07.2026", "1", ""], start=1):
        table1.rows[1].cells[col_idx].text = value

    doc.add_heading("4. Виплатити додаткову винагороду за період дії воєнного часу, які захоплені в полон", level=1)
    headers = ["№", "Посада", "Військове звання", "Прізвище, ім'я, по батькові", "Дата зникнення безвісти", "Період виплати", "Примітка (підстави)"]
    table2 = doc.add_table(rows=2, cols=7)
    for col_idx, header in enumerate(headers):
        table2.rows[0].cells[col_idx].text = header
    for col_idx, value in enumerate(["Гранатометник", "солдат", "ЧЕТВЕРТИЙ Четвертий Четвертий", "11.08.2024", "01.07.2026-31.07.2026", "Наказ №1164-ОД"], start=1):
        table2.rows[1].cells[col_idx].text = value

    path = tmp_path / "report.docx"
    doc.save(str(path))

    file_path = gt.process_generate_timetable(str(path))
    ws = load_workbook(file_path).active

    rows_by_pib = {row[2].value: row for row in ws.iter_rows(min_row=2)}
    # openpyxl зберігає порожній рядок як None після збереження/перечитування.
    assert rows_by_pib["ДРУГИЙ Другий Другий"][-2].value is None
    assert rows_by_pib["ДРУГИЙ Другий Другий"][-1].value is None
    assert rows_by_pib["ЧЕТВЕРТИЙ Четвертий Четвертий"][-2].value == "Наказ №1164-ОД"
    assert rows_by_pib["ЧЕТВЕРТИЙ Четвертий Четвертий"][-1].value == "11.08.2024"


def test_process_generate_timetable_fills_pidstavy_for_100_vp_100_shp_and_not_paid(tmp_path, monkeypatch):
    """Регресія: користувач повідомив, що "ПІДСТАВИ" мають заповнюватись не
    лише для СПЕЦКОНТИНГЕНТУ, а й для 100_ВП/100_ШП/NOT_PAID."""
    monkeypatch.setattr(gt, "OUTPUT_DIR", str(tmp_path))
    monkeypatch.setattr(gt, "MONTH", "07")
    monkeypatch.setattr(gt, "YEAR", "2026")

    path = _write_report(tmp_path / "report.docx", [
        ("8. Виплатити особовому складу за період перебування на стаціонарному лікуванні", [
            ["старший технік", "штаб-сержант", "ТРЕТІЙ Третій Третій", "01.07.2026-01.07.2026", "1", "Довідка №1 від 01.07.2026"],
        ]),
        ("9. Виплатити особовому складу, які перебувають у відпустці для лікування після тяжкого поранення", [
            ["Стрілець", "сержант", "ШОСТИЙ Шостий Шостий", "02.07.2026-02.07.2026", "1", "Довідка №2 від 02.07.2026"],
        ]),
        ("7. Не виплачувати додаткову винагороду нижчепойменованим", [
            ["Гранатометник", "матрос", "СЬОМИЙ Сьомий Сьомий", "03.07.2026-03.07.2026", "1", "Наказ №3 від 03.07.2026"],
        ]),
    ])

    file_path = gt.process_generate_timetable(path)
    ws = load_workbook(file_path).active

    rows_by_pib = {row[2].value: row for row in ws.iter_rows(min_row=2)}
    assert rows_by_pib["ТРЕТІЙ Третій Третій"][-2].value == "Довідка №1 від 01.07.2026"
    assert rows_by_pib["ШОСТИЙ Шостий Шостий"][-2].value == "Довідка №2 від 02.07.2026"
    assert rows_by_pib["СЬОМИЙ Сьомий Сьомий"][-2].value == "Наказ №3 від 03.07.2026"


def test_process_generate_timetable_file_name_uses_selected_month_and_year(tmp_path, monkeypatch):
    monkeypatch.setattr(gt, "OUTPUT_DIR", str(tmp_path))
    monkeypatch.setattr(gt, "SHORT_UNIT_BATTALION", "0 бат")
    monkeypatch.setattr(gt, "MONTH", "03")
    monkeypatch.setattr(gt, "YEAR", "2025")

    path = _write_report(tmp_path / "report.docx", [
        ("1. Виплатити додаткову винагороду у розмірі 30 000 грн. 00 коп. військовослужбовцям", [
            ["Стрілець", "сержант", "ДРУГИЙ Другий Другий", "05.03.2025-06.03.2025", "2", ""],
        ]),
    ])

    file_path = gt.process_generate_timetable(path)

    assert file_path == str(tmp_path / "Табель 0 БАТ березень 2025.xlsx")


def test_process_generate_timetable_applies_style_reference(tmp_path, monkeypatch):
    """Стиль (шрифт/заливка заголовка, заливка клітинки дня за ключовим словом) -
    з content.oblik_style_reference.load_timetable_style_reference, коли він
    доступний (тут - підмінений керованим фейком, щоб не залежати від реального
    ОБЛІК.xlsx у тестовому середовищі)."""
    monkeypatch.setattr(gt, "OUTPUT_DIR", str(tmp_path))
    monkeypatch.setattr(gt, "MONTH", "07")
    monkeypatch.setattr(gt, "YEAR", "2026")

    header_font = Font(name="Bahnschrift", bold=True)
    body_font = Font(name="Times New Roman", size=10)
    header_border = Border(top=Side(style="medium"))
    body_border = Border(top=Side(style="thin"))
    fake_style = {
        "header_font": header_font, "header_fill": None, "header_border": header_border,
        "header_alignment_label": Alignment(horizontal="left"), "header_alignment_date": Alignment(textRotation=90),
        "body_font": body_font, "body_alignment": Alignment(horizontal="center"), "body_border": body_border,
        "label_col_width": 27, "pib_col_width": 36, "date_col_width": 8,
        "header_row_height": 120, "body_row_height": 25.5,
        "keyword_colors": {"30": "BDD7EE"},
    }
    monkeypatch.setattr(gt, "load_timetable_style_reference", lambda: fake_style)

    path = _write_report(tmp_path / "report.docx", [
        ("1. Виплатити додаткову винагороду у розмірі 30 000 грн. 00 коп. військовослужбовцям", [
            ["Стрілець", "сержант", "ДРУГИЙ Другий Другий", "01.07.2026-01.07.2026", "1", ""],
        ]),
    ])

    file_path = gt.process_generate_timetable(path)
    ws = load_workbook(file_path).active

    assert ws.cell(row=1, column=3).font.name == "Bahnschrift"
    assert ws.cell(row=1, column=4).alignment.textRotation == 90
    assert ws.cell(row=1, column=3).border.top.style == "medium"
    assert ws.cell(row=2, column=3).border.top.style == "thin"
    assert ws.row_dimensions[1].height == 120
    assert ws.row_dimensions[2].height == 25.5
    assert ws.column_dimensions["C"].width == 36

    day_cell = ws.cell(row=2, column=4)
    assert day_cell.value == 30
    assert day_cell.fill.fgColor.rgb == "00BDD7EE"


def test_process_generate_timetable_falls_back_to_default_style_when_reference_unavailable(tmp_path, monkeypatch):
    monkeypatch.setattr(gt, "OUTPUT_DIR", str(tmp_path))
    monkeypatch.setattr(gt, "MONTH", "07")
    monkeypatch.setattr(gt, "YEAR", "2026")
    monkeypatch.setattr(gt, "load_timetable_style_reference", lambda: None)

    path = _write_report(tmp_path / "report.docx", [
        ("1. Виплатити додаткову винагороду у розмірі 30 000 грн. 00 коп. військовослужбовцям", [
            ["Стрілець", "сержант", "ДРУГИЙ Другий Другий", "01.07.2026-01.07.2026", "1", ""],
        ]),
    ])

    file_path = gt.process_generate_timetable(path)
    ws = load_workbook(file_path).active

    assert ws.cell(row=1, column=1).font.bold is True
    assert ws.cell(row=2, column=4).fill.fgColor.rgb in (None, "00000000")
