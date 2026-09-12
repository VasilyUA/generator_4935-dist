import sys
from pathlib import Path

sys.path.append(str(Path(__file__).resolve().parent.parent))

from datetime import datetime

from docx import Document

import content.log_war_document_reader as lwdr


def _write_log_war_doc(path, paragraphs, use_table=True):
    """Реальні документи ЖБД тримають увесь наратив ВСЕРЕДИНІ ОДНІЄЇ клітинки
    таблиці (use_table=True - за замовчуванням, матиме сенс для більшості
    тестів); use_table=False - запасний варіант (звичайні абзаци документа),
    для перевірки fallback-гілки _read_log_war_paragraphs."""
    doc = Document()
    if use_table:
        table = doc.add_table(rows=1, cols=1)
        cell = table.cell(0, 0)
        cell.paragraphs[0].text = paragraphs[0] if paragraphs else ""
        for text in paragraphs[1:]:
            cell.add_paragraph(text)
    else:
        for text in paragraphs:
            doc.add_paragraph(text)
    doc.save(path)
    return path


# -------------------------
# scan_log_war_folder
# -------------------------
def test_scan_log_war_folder_matches_expected_filename_pattern(tmp_path):
    _write_log_war_doc(tmp_path / "ЖБД 01.05.2026.docx", ["1. Загальна обстановка."])
    _write_log_war_doc(tmp_path / "ЖБД 02.05.2026.docx", ["1. Загальна обстановка."])

    result = lwdr.scan_log_war_folder(str(tmp_path))

    assert set(result.keys()) == {datetime(2026, 5, 1), datetime(2026, 5, 2)}


def test_scan_log_war_folder_case_insensitive():
    assert lwdr._LOG_WAR_FILENAME_RE.match("жбд 01.05.2026.docx")
    assert lwdr._LOG_WAR_FILENAME_RE.match("ЖБД 01.05.2026.DOCX")


def test_scan_log_war_folder_ignores_non_matching_and_doc_files(tmp_path):
    _write_log_war_doc(tmp_path / "ЖБД 01.05.2026.docx", ["текст"])
    (tmp_path / "Титулка МП 04.05.2026.doc").write_bytes(b"not a real doc file")
    (tmp_path / "№27 ЗАВДАННЯ 01.05.2026.docx").write_bytes(b"")

    result = lwdr.scan_log_war_folder(str(tmp_path))

    assert list(result.keys()) == [datetime(2026, 5, 1)]


def test_scan_log_war_folder_warns_on_duplicate_date(tmp_path, capsys):
    (tmp_path / "a").mkdir()
    (tmp_path / "b").mkdir()
    _write_log_war_doc(tmp_path / "a" / "ЖБД 01.05.2026.docx", ["текст"])
    _write_log_war_doc(tmp_path / "b" / "ЖБД 01.05.2026.docx", ["текст"])

    result = lwdr.scan_log_war_folder(str(tmp_path))

    assert len(result) == 1
    assert "Кілька файлів ЖБД" in capsys.readouterr().out


def test_scan_log_war_folder_empty_folder_returns_empty(tmp_path):
    assert lwdr.scan_log_war_folder(str(tmp_path)) == {}


# -------------------------
# _read_log_war_paragraphs
# -------------------------
def test_read_log_war_paragraphs_from_table_cell(tmp_path):
    path = _write_log_war_doc(tmp_path / "doc.docx", ["1. Раз.", "2. Два.", "3. Три."])
    assert lwdr._read_log_war_paragraphs(str(path)) == ["1. Раз.", "2. Два.", "3. Три."]


def test_read_log_war_paragraphs_fallback_to_top_level_paragraphs(tmp_path):
    path = _write_log_war_doc(tmp_path / "doc.docx", ["1. Раз.", "2. Два."], use_table=False)
    assert lwdr._read_log_war_paragraphs(str(path)) == ["1. Раз.", "2. Два."]


def test_read_log_war_paragraphs_unreadable_file_returns_empty(tmp_path, capsys):
    path = tmp_path / "broken.docx"
    path.write_bytes(b"not a real docx")
    assert lwdr._read_log_war_paragraphs(str(path)) == []
    assert "Не вдалось прочитати" in capsys.readouterr().out


# -------------------------
# _slice_after_heading / extract_log_war_sections
# -------------------------
def test_slice_after_heading_returns_content_up_to_next_numbered_heading():
    paragraphs = ["4. Щось.", "5. Рішення командира.", "Абзац А.", "Абзац Б.", "5.1 Підпункт.", "6. Наступне."]
    index, text = lwdr._slice_after_heading(paragraphs, lwdr._SECTION5_HEADING_RE)
    assert index == 1
    assert text == "Абзац А.\nАбзац Б."


def test_slice_after_heading_not_found_returns_none_and_empty():
    paragraphs = ["1. Щось.", "2. Інше."]
    assert lwdr._slice_after_heading(paragraphs, lwdr._SECTION5_HEADING_RE) == (None, "")


def test_slice_after_heading_content_runs_to_end_of_document():
    paragraphs = ["5. Рішення командира.", "Останній абзац."]
    index, text = lwdr._slice_after_heading(paragraphs, lwdr._SECTION5_HEADING_RE)
    assert index == 0
    assert text == "Останній абзац."


def test_extract_log_war_sections_full_shape(tmp_path):
    paragraphs = [
        "4. Коротке викладення.",
        "5. Рішення командира військової частини та бойові завдання.",
        "Здійснити переміщення ПЕРШИЙ Перший Перший з А в Б.",
        "6. Відомості про виконання бойового завдання.",
        "Хід виконання спланованих завдань:",
        "Залучено особовий склад: ПЕРШИЙ Перший Перший (Місце);",
        "6.1 Хід ведення бойових дій.",
        "Щось інше.",
    ]
    path = _write_log_war_doc(tmp_path / "doc.docx", paragraphs)

    sections = lwdr.extract_log_war_sections(str(path))

    assert "Здійснити переміщення" in sections["section5"]
    assert "Залучено особовий склад" in sections["section6_hid"]
    assert "Щось інше" not in sections["section6_hid"]


def test_extract_log_war_sections_unit_designator_without_dot_is_not_mistaken_for_heading(tmp_path):
    """РЕГРЕСІЯ (виявлено користувачем на реальному зразку 01.07.2026): текст
    пункту "5" рясніє позначеннями підрозділів БЕЗ крапки після числа
    ("1 бмп", "40 обрмн" тощо) - _NUMBERED_HEADING_RE раніше помилково
    визнавав голе число (без внутрішньої крапки, з ОПЦІЙНОЮ крапкою в кінці)
    за заголовок і обривав секцію "5" одразу після нього - увесь список ПІБ,
    що йшов ДАЛІ в тому ж абзаці/наступних абзацах, губився."""
    paragraphs = [
        "4. Коротке викладення.",
        "5. Рішення командира військової частини та бойові завдання.",
        "Управління силами здійснювати з КСП 1 бмп 40 обрмн в районі: "
        "майор ПЕРШИЙ Перший Перший, майор ДРУГИЙ Другий Другий.",
        "6. Відомості про виконання бойового завдання.",
        "Хід виконання спланованих завдань:",
        "Залучено особовий склад: ПЕРШИЙ Перший Перший (Місце);",
    ]
    path = _write_log_war_doc(tmp_path / "doc.docx", paragraphs)

    sections = lwdr.extract_log_war_sections(str(path))

    assert "ПЕРШИЙ Перший Перший" in sections["section5"]
    assert "ДРУГИЙ Другий Другий" in sections["section5"]


def test_extract_log_war_sections_empty_hid_subsection_is_not_an_error(tmp_path):
    """Реальний зразок (перехідний день) - "Хід виконання спланованих завдань:"
    одразу ж змінюється "6.1" без жодного вмісту - легітимний, ОЧІКУВАНИЙ
    результат (порожній рядок), а не помилка."""
    paragraphs = [
        "5. Рішення командира.",
        "Нічого особливого.",
        "6. Відомості про виконання.",
        "Хід виконання спланованих завдань:",
        "6.1 Хід ведення бойових дій.",
    ]
    path = _write_log_war_doc(tmp_path / "doc.docx", paragraphs)

    sections = lwdr.extract_log_war_sections(str(path))

    assert sections["section6_hid"] == ""


def test_extract_log_war_sections_missing_sections_return_empty_strings(tmp_path):
    path = _write_log_war_doc(tmp_path / "doc.docx", ["1. Загальна обстановка.", "2. Бойовий склад."])

    sections = lwdr.extract_log_war_sections(str(path))

    assert sections == {"section5": "", "section6_hid": ""}
