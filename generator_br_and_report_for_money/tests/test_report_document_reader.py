import sys
from datetime import datetime
from pathlib import Path

sys.path.append(str(Path(__file__).resolve().parent.parent))

from docx import Document

import content.report_document_reader as rdr

_TABLE_HEADERS = ["№", "Посада", "Військове звання", "Прізвище, ім'я, по батькові", "Період участі", "Кількість днів", "Підстава для виплати"]


def _add_section(doc, heading_text, data_rows):
    """data_rows - список [посада, звання, піб, період, дні, підстава] - будує
    Heading 1 + таблицю (7 колонок, як create_table у formatting/docx_utils.py:
    ПІБ у колонці 3, ПЕРІОД у колонці 4, 0-відлік) одразу під ним."""
    doc.add_heading(heading_text, level=1)
    table = doc.add_table(rows=1 + len(data_rows), cols=7)
    for col_idx, header in enumerate(_TABLE_HEADERS):
        table.rows[0].cells[col_idx].text = header
    for row_idx, values in enumerate(data_rows, start=1):
        table.rows[row_idx].cells[0].text = str(row_idx)
        for col_idx, value in enumerate(values, start=1):
            table.rows[row_idx].cells[col_idx].text = value


def _write_report(path, sections):
    doc = Document()
    for heading_text, data_rows in sections:
        _add_section(doc, heading_text, data_rows)
    doc.save(str(path))
    return str(path)


def _add_plain_paragraph_section(doc, heading_text, data_rows):
    """Як _add_section, але заголовок - ЗВИЧАЙНИЙ "Normal"-абзац (doc.add_paragraph),
    БЕЗ стилю Word "Heading" - так реально виглядають пункти в рапортах,
    підготовлених вручну у Word (напр. живий зразок
    resources/ДВ_зі_змінами_1бмп_ТРАВЕНЬ_v2.3.docx: усі його пункти - "Normal"-
    абзаци з номером прямо в тексті)."""
    doc.add_paragraph(heading_text)
    table = doc.add_table(rows=1 + len(data_rows), cols=7)
    for col_idx, header in enumerate(_TABLE_HEADERS):
        table.rows[0].cells[col_idx].text = header
    for row_idx, values in enumerate(data_rows, start=1):
        table.rows[row_idx].cells[0].text = str(row_idx)
        for col_idx, value in enumerate(values, start=1):
            table.rows[row_idx].cells[col_idx].text = value


def _write_report_with_plain_paragraph_headings(path, sections):
    doc = Document()
    for heading_text, data_rows in sections:
        _add_plain_paragraph_section(doc, heading_text, data_rows)
    doc.save(str(path))
    return str(path)


# -------------------------
# _raw_value_for_heading
# -------------------------
def test_raw_value_for_heading_returns_none_for_changes_subsection():
    assert rdr._raw_value_for_heading("11.1 Виключити пункти в додатку 1:") is None
    assert rdr._raw_value_for_heading("11.2 Додаток 1 доповнити наступними пунктами:") is None


def test_raw_value_for_heading_detects_not_paid_case_insensitively():
    assert rdr._raw_value_for_heading("7. Не виплачувати додаткову винагороду нижчепойменованим...") == "NOT_PAID"
    assert rdr._raw_value_for_heading("7. НЕ ВИПЛАЧУВАТИ додаткову винагороду...") == "NOT_PAID"


def test_raw_value_for_heading_extracts_money_amount_without_confusing_70_and_170():
    assert rdr._raw_value_for_heading("1. Виплатити додаткову винагороду у розмірі 30 000 грн. 00 коп. військовослужбовцям") == 30
    assert rdr._raw_value_for_heading("2. Виплатити додаткову винагороду в розмірі 100 000 грн. 00 коп. військовослужбовцям, які беруть участь") == 100
    assert rdr._raw_value_for_heading("3. Виплатити додаткову винагороду в розмірі 10 000 грн. 00 коп. за час виконання обов'язків") == 10
    assert rdr._raw_value_for_heading("5. Виплатити одноразову винагороду... у розмірі 70 000 грн. 00 коп.  (з розрахунку...") == 70
    assert rdr._raw_value_for_heading("6. Виплатити одноразову винагороду... у розмірі 170 000 грн. 00 коп.  (з розрахунку...") == 170
    assert rdr._raw_value_for_heading("9. Виплатити додаткову винагороду у розмірі 50 000 грн. 00 коп. військовослужбовцям") == 50


def test_raw_value_for_heading_detects_100_shp_for_hospitalization():
    """"стаціонарне лікування"/"лікарняний заклад" - госпіталізація ("100_ШП"),
    незалежно від точної назви пункту чи наявності суми "100 000 грн" узагалі."""
    assert rdr._raw_value_for_heading(
        "9. Виплатити додаткову винагороду в розмірі 100 000 грн. 00 коп. на підставі висновку "
        "військово-лікарської комісії про потребу в тривалому лікуванні в лікарняних закладах",
    ) == "100_ШП"
    assert rdr._raw_value_for_heading(
        "8. Виплатити особовому складу за період перебування на стаціонарному лікуванні в закладах охорони здоров'я",
    ) == "100_ШП"
    # Без "лікува..."/"поранен..." - звичайний бойовий пункт 100, не 100_ШП.
    assert rdr._raw_value_for_heading("2. Виплатити додаткову винагороду в розмірі 100 000 грн. 00 коп. військовослужбовцям") == 100


def test_raw_value_for_heading_detects_100_vp_for_medical_vacation_after_wound():
    """"відпустці для лікування"/"тяжкого поранення" - людина НЕ в лікарняному
    закладі, а у відпустці ("100_ВП") - НЕ те саме, що госпіталізація ("100_ШП"),
    навіть попри спільні слова "лікування"/"поранення" в обох формулюваннях."""
    assert rdr._raw_value_for_heading(
        "9. Виплатити додаткову винагороду у розмірі 100 000 грн. 00 коп. військовослужбовцям, які у "
        "зв'язку з пораненням (контузією, травмою, каліцтвом) перебувають у відпустці для лікування "
        "після тяжкого поранення за висновком військово-лікарської комісії",
    ) == "100_ВП"


def test_raw_value_for_heading_falls_back_to_100_shp_for_unrecognized_medical_wording():
    """Згадка лікування/поранення, що не підійшла під жодну з двох вужчих
    перевірок (стаціонар/відпустка) - все одно "100_ШП" (запасний варіант), а не
    втрачається зовсім."""
    assert rdr._raw_value_for_heading("9. Виплатити винагороду у зв'язку з пораненням, отриманим під час бою") == "100_ШП"


def test_raw_value_for_heading_detects_spetskontyngent_via_captured_missing_keywords():
    """"полон"/"безвісти"/"заручник" - позначка "СПЕЦКОНТИНГЕНТ" незалежно від
    точної назви пункту (тут - НОВИЙ "100_СПЕЦКОНТИНГЕНТ", доданий користувачем
    у resources/data.json) і від наявності суми "NNN 000 грн" у тексті (цей
    пункт її взагалі не містить)."""
    assert rdr._raw_value_for_heading(
        "4. Виплатити додаткову винагороду за період дії воєнного часу нижчепойменованим військовослужбовцям, "
        "які захоплені в полон крім тих, які добровільно здалися в полон або є заручниками",
    ) == "100_СПЕЦКОНТИНГЕНТ"
    assert rdr._raw_value_for_heading("4. ...інтерновані нейтральні держави або безвісно відсутні...") == "100_СПЕЦКОНТИНГЕНТ"


def test_raw_value_for_heading_falls_back_to_point_number_when_truly_unrecognized():
    assert rdr._raw_value_for_heading("4. Якийсь цілком новий, ще не описаний тут вид пункту") == "п.4"


def test_raw_value_for_heading_unknown_shape_returns_question_mark():
    assert rdr._raw_value_for_heading("Довільний текст без номера на початку") == "?"


# -------------------------
# _dates_from_period_text
# -------------------------
def test_dates_from_period_text_expands_ranges():
    dates = rdr._dates_from_period_text("01.07.2026-03.07.2026; 05.07.2026-05.07.2026")
    assert dates == [datetime(2026, 7, 1), datetime(2026, 7, 2), datetime(2026, 7, 3), datetime(2026, 7, 5)]


def test_dates_from_period_text_blank_returns_empty():
    assert rdr._dates_from_period_text("") == []
    assert rdr._dates_from_period_text(None) == []


# -------------------------
# _date_subranges_from_period_text - checker_accounting.report_log_war_checker
# потребує МЕЖІ кожного під-діапазону (перший день кожного - "перехідний"),
# а не лише плаский список дат.
# -------------------------
def test_date_subranges_from_period_text_single_range():
    subranges = rdr._date_subranges_from_period_text("01.07.2026-03.07.2026")
    assert subranges == [(datetime(2026, 7, 1), [datetime(2026, 7, 1), datetime(2026, 7, 2), datetime(2026, 7, 3)])]


def test_date_subranges_from_period_text_multiple_ranges():
    subranges = rdr._date_subranges_from_period_text("01.07.2026-02.07.2026; 05.07.2026-05.07.2026")
    assert subranges == [
        (datetime(2026, 7, 1), [datetime(2026, 7, 1), datetime(2026, 7, 2)]),
        (datetime(2026, 7, 5), [datetime(2026, 7, 5)]),
    ]


def test_date_subranges_from_period_text_blank_returns_empty():
    assert rdr._date_subranges_from_period_text("") == []
    assert rdr._date_subranges_from_period_text(None) == []


def test_dates_from_period_text_matches_flattened_subranges():
    """Регресія після рефакторингу _dates_from_period_text у термінах
    _date_subranges_from_period_text - поведінка для викликачів (напр.
    checker_accounting.report_checker) не повинна змінитись."""
    period_text = "01.07.2026-03.07.2026; 05.07.2026-06.07.2026"
    flattened = [date for _start, dates in rdr._date_subranges_from_period_text(period_text) for date in dates]
    assert rdr._dates_from_period_text(period_text) == flattened


# -------------------------
# extract_timetable_rows_from_report
# -------------------------
def test_extract_timetable_rows_from_report_builds_grid_from_recognized_sections(tmp_path):
    path = _write_report(tmp_path / "report.docx", [
        ("1. Виплатити додаткову винагороду у розмірі 30 000 грн. 00 коп. військовослужбовцям", [
            ["Стрілець", "сержант", "ДРУГИЙ Другий Другий", "01.07.2026-02.07.2026", "2", ""],
        ]),
        ("7. Не виплачувати додаткову винагороду нижчепойменованим", [
            ["Гранатометник", "матрос", "СЬОМИЙ Сьомий Сьомий", "03.07.2026-03.07.2026", "1", ""],
        ]),
        ("11.1 Виключити пункти в додатку 1:", [
            ["Оператор", "старший матрос", "ПРИЗРАК Examples Examples", "01.06.2026-15.06.2026", "15", ""],
        ]),
    ])

    rows, date_columns = rdr.extract_timetable_rows_from_report(path)

    assert date_columns == [datetime(2026, 7, 1), datetime(2026, 7, 2), datetime(2026, 7, 3)]
    by_pib = {row["ПІБ"]: row for row in rows}
    assert set(by_pib) == {"ДРУГИЙ Другий Другий", "СЬОМИЙ Сьомий Сьомий"}
    assert by_pib["ДРУГИЙ Другий Другий"]["ПОСАДА"] == "Стрілець"
    assert by_pib["ДРУГИЙ Другий Другий"][datetime(2026, 7, 1)] == 30
    assert by_pib["ДРУГИЙ Другий Другий"][datetime(2026, 7, 2)] == 30
    assert by_pib["ДРУГИЙ Другий Другий"][datetime(2026, 7, 3)] is None
    assert by_pib["СЬОМИЙ Сьомий Сьомий"][datetime(2026, 7, 3)] == "NOT_PAID"


def test_extract_timetable_rows_from_report_defers_100_rollup_behind_specific_point(tmp_path):
    """Пункт 100 рахує ще й дні 70/170 додатковим рядком (той самий день - у ДВОХ
    таблицях) - значення з таблиці "100" не повинно затирати справжнє "70" для
    того самого дня тієї самої людини."""
    path = _write_report(tmp_path / "report.docx", [
        ("5. Виплатити одноразову винагороду... у розмірі 70 000 грн. 00 коп.  (з розрахунку", [
            ["Навідник", "матрос", "ВОСЬМИЙ Восьмий Восьмий", "01.07.2026-01.07.2026", "1", ""],
        ]),
        ("2. Виплатити додаткову винагороду в розмірі 100 000 грн. 00 коп. військовослужбовцям", [
            ["Навідник", "матрос", "ВОСЬМИЙ Восьмий Восьмий", "01.07.2026-01.07.2026", "1", ""],
        ]),
    ])

    rows, _date_columns = rdr.extract_timetable_rows_from_report(path)

    assert rows[0][datetime(2026, 7, 1)] == 70


# -------------------------
# _detect_columns
# -------------------------
def test_detect_columns_finds_period_by_header_text_when_shifted_by_extra_column(tmp_path):
    """Регресія: таблиця пункту "СПЕЦКОНТИНГЕНТ" (полонені/зниклі безвісти) має
    ДОДАТКОВУ колонку "Дата зникнення безвісти" ПЕРЕД періодом ("Період виплати"
    замість "Період участі"), зсуваючи сам період з позиції 4 на позицію 5 - за
    фіксованою позицією (стара поведінка) ця колонка прочиталась би як період,
    хоча там дата зникнення (без діапазону "DD.MM.YYYY-DD.MM.YYYY")."""
    doc = Document()
    headers = ["№", "Посада", "Військове звання", "Прізвище, ім'я, по батькові", "Дата зникнення безвісти", "Період виплати", "Примітка"]
    table = doc.add_table(rows=1, cols=7)
    for col_idx, header in enumerate(headers):
        table.rows[0].cells[col_idx].text = header

    columns = rdr._detect_columns(table.rows[0])

    # "days"/"basis" тут - запасні позиції (5/6): жоден заголовок не містить
    # "днів"/"підстав" у цьому конкретному зразку (не про це цей тест) - лише
    # "period"/"disappearance_date" справді визначені за назвою заголовка.
    assert columns == {"posada": 1, "zvannya": 2, "pib": 3, "period": 5, "days": 5, "basis": 6, "disappearance_date": 4}


def test_extract_timetable_rows_from_report_reads_shifted_period_column(tmp_path):
    """Той самий регресійний сценарій, наскрізь: людина з таблиці "СПЕЦКОНТИНГЕНТ"
    (зсунута колонка періоду) МАЄ потрапити в результат з правильними датами, а
    не бути втраченою (як було до _detect_columns - "Дата зникнення безвісти" не
    містить діапазону дат, тож _dates_from_period_text повертала [] і людина без
    жодного дня відкидалась)."""
    doc = Document()
    doc.add_heading(
        "4. Виплатити додаткову винагороду за період дії воєнного часу нижчепойменованим військовослужбовцям, "
        "які захоплені в полон або безвісно відсутні",
        level=1,
    )
    headers = ["№", "Посада", "Військове звання", "Прізвище, ім'я, по батькові", "Дата зникнення безвісти", "Період виплати", "Примітка"]
    table = doc.add_table(rows=2, cols=7)
    for col_idx, header in enumerate(headers):
        table.rows[0].cells[col_idx].text = header
    data_row = ["1", "Колишній гранатометник", "солдат", "ЧЕТВЕРТИЙ Четвертий Четвертий", "11.08.2024", "01.07.2026-02.07.2026", "Наказ №1164-ОД"]
    for col_idx, value in enumerate(data_row):
        table.rows[1].cells[col_idx].text = value
    path = tmp_path / "report.docx"
    doc.save(str(path))

    rows, date_columns = rdr.extract_timetable_rows_from_report(str(path))

    assert date_columns == [datetime(2026, 7, 1), datetime(2026, 7, 2)]
    assert len(rows) == 1
    assert rows[0]["ПІБ"] == "ЧЕТВЕРТИЙ Четвертий Четвертий"
    assert rows[0][datetime(2026, 7, 1)] == "100_СПЕЦКОНТИНГЕНТ"
    assert rows[0][datetime(2026, 7, 2)] == "100_СПЕЦКОНТИНГЕНТ"


def test_extract_timetable_rows_from_report_skips_rows_without_pib(tmp_path):
    path = _write_report(tmp_path / "report.docx", [
        ("1. Виплатити додаткову винагороду у розмірі 30 000 грн. 00 коп. військовослужбовцям", [
            ["", "", "", "01.07.2026-01.07.2026", "1", ""],
        ]),
    ])

    rows, date_columns = rdr.extract_timetable_rows_from_report(path)

    assert rows == []
    assert date_columns == []


def test_extract_timetable_rows_from_report_ignores_document_with_only_changes_sections(tmp_path):
    path = _write_report(tmp_path / "report.docx", [
        ("11.1 Виключити пункти в додатку 1:", [
            ["Оператор", "старший матрос", "ПЕРШИЙ Перший Перший", "01.06.2026-15.06.2026", "15", ""],
        ]),
    ])

    rows, date_columns = rdr.extract_timetable_rows_from_report(path)

    assert rows == []
    assert date_columns == []


# -------------------------
# extract_timetable_rows_from_report - target_month/target_year
# -------------------------
def test_extract_timetable_rows_from_report_filters_out_other_months(tmp_path):
    """Регресія: пункт "перебування на лікуванні" може охоплювати перехідний
    період з ПОПЕРЕДНЬОГО місяця - людина, чиї дні є ЛИШЕ поза обраним місяцем,
    не повинна потрапити в результат узагалі."""
    path = _write_report(tmp_path / "report.docx", [
        ("1. Виплатити додаткову винагороду у розмірі 30 000 грн. 00 коп. військовослужбовцям", [
            ["Стрілець", "сержант", "ДРУГИЙ Другий Другий", "01.07.2026-02.07.2026", "2", ""],
        ]),
        ("8. Виплатити особовому складу за період перебування на лікуванні", [
            ["старший технік", "штаб-сержант", "ТРЕТІЙ Третій Третій", "22.06.2026-30.06.2026", "9", ""],
        ]),
    ])

    rows, date_columns = rdr.extract_timetable_rows_from_report(path, target_month=7, target_year=2026)

    assert date_columns[0] == datetime(2026, 7, 1)
    assert date_columns[-1] == datetime(2026, 7, 31)
    assert len(date_columns) == 31
    assert [row["ПІБ"] for row in rows] == ["ДРУГИЙ Другий Другий"]


def test_extract_timetable_rows_from_report_keeps_only_target_month_days_for_split_period(tmp_path):
    """Одна людина може мати дні І в обраному місяці, І поза ним (той самий
    пункт, один запис "ПЕРІОД") - лишаються лише дні обраного місяця."""
    path = _write_report(tmp_path / "report.docx", [
        ("1. Виплатити додаткову винагороду у розмірі 30 000 грн. 00 коп. військовослужбовцям", [
            ["Стрілець", "сержант", "ДРУГИЙ Другий Другий", "30.06.2026-02.07.2026", "3", ""],
        ]),
    ])

    rows, date_columns = rdr.extract_timetable_rows_from_report(path, target_month=7, target_year=2026)

    assert len(rows) == 1
    assert rows[0][datetime(2026, 7, 1)] == 30
    assert rows[0][datetime(2026, 7, 2)] == 30
    assert datetime(2026, 6, 30) not in rows[0]


def test_extract_timetable_rows_from_report_returns_empty_when_nothing_in_target_month(tmp_path):
    path = _write_report(tmp_path / "report.docx", [
        ("1. Виплатити додаткову винагороду у розмірі 30 000 грн. 00 коп. військовослужбовцям", [
            ["Стрілець", "сержант", "ДРУГИЙ Другий Другий", "01.06.2026-02.06.2026", "2", ""],
        ]),
    ])

    rows, date_columns = rdr.extract_timetable_rows_from_report(path, target_month=7, target_year=2026)

    assert rows == []
    assert len(date_columns) == 31


# -------------------------
# extract_report_entries
# -------------------------
def test_extract_report_entries_returns_period_and_days_as_written(tmp_path):
    """На відміну від extract_timetable_rows_from_report (сітка днів), тут ПЕРІОД/
    ДНІ повертаються ЯК НАПИСАНО в самому рапорті (рядки тексту) - для
    checker_accounting.report_checker, який звіряє САМЕ написане з ОБЛІК.xlsx."""
    path = _write_report(tmp_path / "report.docx", [
        ("1. Виплатити додаткову винагороду у розмірі 30 000 грн. 00 коп. військовослужбовцям", [
            ["Стрілець", "сержант", "ДРУГИЙ Другий Другий", "01.07.2026-02.07.2026; 05.07.2026-05.07.2026", "3", ""],
        ]),
        ("11.1 Виключити пункти в додатку 1:", [
            ["Оператор", "старший матрос", "ПЕРШИЙ Перший Перший", "01.06.2026-15.06.2026", "15", ""],
        ]),
    ])

    entries = rdr.extract_report_entries(path)

    assert entries == [{
        "raw_value": 30, "ПОСАДА": "Стрілець", "ЗВАННЯ": "сержант", "ПІБ": "ДРУГИЙ Другий Другий",
        "ПЕРІОД": "01.07.2026-02.07.2026; 05.07.2026-05.07.2026", "ДНІ": "3",
        "ПІДСТАВА": "", "ДАТА_ЗНИКНЕННЯ": None, "НОМЕР_ПУНКТУ": "1",
    }]


# -------------------------
# _is_heading - РЕГРЕСІЯ (виявлено на живому зразку
# resources/ДВ_зі_змінами_1бмп_ТРАВЕНЬ_v2.3.docx): реальні рапорти, підготовлені
# вручну у Word, часто мають пункти як ЗВИЧАЙНІ "Normal"-абзаци (номер прямо в
# тексті), БЕЗ стилю Word "Heading" - без розпізнавання такого абзацу як
# заголовка current_heading лишається порожнім на ввесь документ, і КОЖЕН
# рядок отримує raw_value/НОМЕР_ПУНКТУ "?" (саме на це поскаржився користувач).
# -------------------------
def test_is_heading_true_for_heading_styled_paragraph():
    doc = Document()
    doc.add_heading("1. Виплатити додаткову винагороду...", level=1)
    assert rdr._is_heading(doc.paragraphs[0])


def test_is_heading_true_for_plain_paragraph_with_leading_point_number():
    doc = Document()
    doc.add_paragraph("1. Виплатити додаткову винагороду у розмірі 30 000 грн...")
    assert rdr._is_heading(doc.paragraphs[0])


def test_is_heading_true_for_plain_paragraph_changes_subsection():
    doc = Document()
    doc.add_paragraph("11.1 Виключити пункти в додатку 1:")
    assert rdr._is_heading(doc.paragraphs[0])


def test_is_heading_false_for_unrelated_plain_paragraph():
    doc = Document()
    doc.add_paragraph("Рапорт подається на 118 (ста вісімнадцяти) аркушах.")
    assert not rdr._is_heading(doc.paragraphs[0])


def test_is_heading_false_for_empty_plain_paragraph():
    doc = Document()
    doc.add_paragraph("")
    assert not rdr._is_heading(doc.paragraphs[0])


def test_is_heading_false_for_table_block():
    """_iter_block_items проходить абзаци й таблиці впереміш - таблиця сама
    по собі ніколи не є заголовком пункту (лише перевіряти block.style/.text
    таблиці було б і безглуздо, і небезпечно - Table не має цих атрибутів)."""
    doc = Document()
    table = doc.add_table(rows=1, cols=1)
    assert not rdr._is_heading(table)


def test_extract_report_entries_recognizes_plain_paragraph_headings(tmp_path):
    """Наскрізь: рапорт, де пункти - ЗВИЧАЙНІ "Normal"-абзаци (як у живому
    зразку), а НЕ doc.add_heading() - і raw_value, і НОМЕР_ПУНКТУ мають бути
    правильно розпізнані, так само як для Heading-стильованих документів."""
    path = _write_report_with_plain_paragraph_headings(tmp_path / "report.docx", [
        ("1. Виплатити додаткову винагороду у розмірі 30 000 грн. 00 коп. військовослужбовцям", [
            ["Стрілець", "сержант", "ПЕРШИЙ Перший Перший", "01.07.2026-02.07.2026", "2", ""],
        ]),
        ("2. Виплатити додаткову винагороду у розмірі 100 000 грн. 00 коп. військовослужбовцям", [
            ["Оператор", "старший матрос", "ДРУГИЙ Другий Другий", "01.07.2026-02.07.2026", "2", ""],
        ]),
        ("2.1 Виключити пункти в додатку 1:", [
            ["Оператор", "старший матрос", "ТРЕТІЙ Третій Третій", "01.06.2026-15.06.2026", "15", ""],
        ]),
    ])

    entries = rdr.extract_report_entries(path)

    assert [(e["ПІБ"], e["raw_value"], e["НОМЕР_ПУНКТУ"]) for e in entries] == [
        ("ПЕРШИЙ Перший Перший", 30, "1"),
        ("ДРУГИЙ Другий Другий", 100, "2"),
    ]


# -------------------------
# "НОМЕР_ПУНКТУ" - літеральний номер пункту з заголовка рапорту (підтверджено
# користувачем: показувати службовий raw_value замість номера ЗАПЛУТУЄ - людина
# звіряє рапорт за ЙОГО ВЛАСНОЮ нумерацією, не за внутрішнім кодом).
# -------------------------
def test_extract_report_entries_heading_number_matches_document_numbering(tmp_path):
    path = _write_report(tmp_path / "report.docx", [
        ("1. Виплатити додаткову винагороду у розмірі 100 000 грн. 00 коп.", [
            ["Стрілець", "сержант", "ПЕРШИЙ Перший Перший", "01.07.2026-02.07.2026", "2", ""],
        ]),
    ])
    entries = rdr.extract_report_entries(path)
    assert entries[0]["НОМЕР_ПУНКТУ"] == "1"


def test_extract_report_entries_heading_number_present_even_when_raw_value_is_fallback_code(tmp_path):
    """РЕГРЕСІЯ (виявлено користувачем): заголовок, що НЕ підійшов під жоден
    відомий орієнтир (raw_value - запасний "п.N" - _raw_value_for_heading),
    однаково має правильний "НОМЕР_ПУНКТУ" - обидва поля обчислюються
    НЕЗАЛЕЖНО одне від одного, тож "raw_value" (службовий код, незрозумілий
    людині) і "НОМЕР_ПУНКТУ" (номер САМЕ як у тексті рапорту) можуть
    відрізнятись."""
    path = _write_report(tmp_path / "report.docx", [
        ("7. Щось геть незрозуміле без жодного відомого орієнтиру.", [
            ["Стрілець", "сержант", "ПЕРШИЙ Перший Перший", "01.07.2026-02.07.2026", "2", ""],
        ]),
    ])
    entries = rdr.extract_report_entries(path)
    assert entries[0]["raw_value"] == "п.7"
    assert entries[0]["НОМЕР_ПУНКТУ"] == "7"


def test_extract_report_entries_heading_number_blank_when_no_leading_number(tmp_path):
    """Заголовок без жодного провідного числа - вкрай рідкісний випадок, але
    НЕ має падати - порожній рядок, а не помилка."""
    path = _write_report(tmp_path / "report.docx", [
        ("Не виплачувати додаткову винагороду нижчепойменованим.", [
            ["Стрілець", "сержант", "ПЕРШИЙ Перший Перший", "01.07.2026-02.07.2026", "2", ""],
        ]),
    ])
    entries = rdr.extract_report_entries(path)
    assert entries[0]["raw_value"] == "NOT_PAID"
    assert entries[0]["НОМЕР_ПУНКТУ"] == ""


def test_extract_report_entries_reads_days_column_by_header_text(tmp_path):
    """Регресія - позиція колонки "ДНІ" визначається за назвою заголовка
    ("К-ь днів"/"Кількість днів"/"К-сть днів" тощо реально трапляються в
    resources/РАПОРТ.docx), а не фіксованою позицією 5."""
    doc = Document()
    doc.add_heading("1. Виплатити додаткову винагороду у розмірі 30 000 грн. 00 коп. військовослужбовцям", level=1)
    headers = ["№", "Посада", "Військове звання", "Прізвище, ім'я, по батькові", "Період участі", "К-ь днів", "Примітка"]
    table = doc.add_table(rows=2, cols=7)
    for col_idx, header in enumerate(headers):
        table.rows[0].cells[col_idx].text = header
    data_row = ["1", "Стрілець", "сержант", "ДРУГИЙ Другий Другий", "01.07.2026-02.07.2026", "2", ""]
    for col_idx, value in enumerate(data_row):
        table.rows[1].cells[col_idx].text = value
    path = tmp_path / "report.docx"
    doc.save(str(path))

    entries = rdr.extract_report_entries(str(path))

    assert entries[0]["ДНІ"] == "2"
