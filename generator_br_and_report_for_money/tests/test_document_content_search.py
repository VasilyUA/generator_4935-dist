import sys
from datetime import datetime
from pathlib import Path

sys.path.append(str(Path(__file__).resolve().parent.parent))

import pytest
from docx import Document

from constants import SHORT_UNIT_BATTALION, SHORT_UNIT_BRIGADE
from content import document_content_search as dcs

_BR = f"БР {SHORT_UNIT_BATTALION} {SHORT_UNIT_BRIGADE}"


def _write_docx(path, paragraphs=(), table_rows=None):
    doc = Document()
    for text in paragraphs:
        doc.add_paragraph(text)
    if table_rows:
        table = doc.add_table(rows=len(table_rows), cols=len(table_rows[0]))
        for r, row in enumerate(table_rows):
            for c, value in enumerate(row):
                table.cell(r, c).text = value
    doc.save(path)
    return str(path)


# -------------------------
# extract_document_text
# -------------------------
def test_extract_document_text_reads_docx_paragraphs_and_tables(tmp_path):
    path = _write_docx(tmp_path / "file.docx", paragraphs=["Абзац з текстом"], table_rows=[["Клітинка таблиці"]])
    text = dcs.extract_document_text(path)
    assert "Абзац з текстом" in text
    assert "Клітинка таблиці" in text


def test_extract_document_text_unreadable_docx_returns_empty_and_warns(tmp_path, capsys):
    path = tmp_path / "broken.docx"
    path.write_text("not a real docx", encoding="utf-8")
    assert dcs.extract_document_text(str(path)) == ""
    assert "Не вдалось прочитати" in capsys.readouterr().out


def test_extract_document_text_unknown_extension_returns_empty(tmp_path):
    path = tmp_path / "file.txt"
    path.write_text("щось", encoding="utf-8")
    assert dcs.extract_document_text(str(path)) == ""


def test_extract_document_text_doc_missing_pywin32_returns_empty_and_warns(tmp_path, monkeypatch, capsys):
    monkeypatch.setitem(sys.modules, "pythoncom", None)
    monkeypatch.setitem(sys.modules, "win32com.client", None)
    path = tmp_path / "file.doc"
    path.write_bytes(b"")
    assert dcs.extract_document_text(str(path)) == ""
    assert "pywin32 недоступний" in capsys.readouterr().out


class _FakeDocument:
    def __init__(self, text):
        self.Content = type("Content", (), {"Text": text})()
        self.closed_with = None

    def Close(self, SaveChanges=None):
        self.closed_with = SaveChanges


class _FakeWord:
    def __init__(self, document=None, open_error=None):
        self._document = document
        self._open_error = open_error
        self.Visible = None
        self.Documents = self

    def Open(self, path, ReadOnly=None):
        if self._open_error:
            raise self._open_error
        return self._document


def test_extract_document_text_doc_happy_path_via_word_com(tmp_path, monkeypatch):
    """pywin32 - реальна (вже встановлена в проєкті) залежність, тож import
    pythoncom/win32com.client тут відбувається по-справжньому - підміняється лише
    сам Dispatch("Word.Application"), а не sys.modules цілком (заміна sys.modules
    фейковими об'єктами провокує внутрішні lazy-імпорти pywin32 (gencache) і падає
    непередбачувано - monkeypatch.setattr на реальний атрибут модуля безпечніший)."""
    pytest.importorskip("win32com.client")
    fake_document = _FakeDocument("Текст зі старого .doc файлу")
    fake_word = _FakeWord(document=fake_document)
    monkeypatch.setattr("win32com.client.Dispatch", lambda name: fake_word)

    path = tmp_path / "file.doc"
    path.write_bytes(b"")

    assert dcs.extract_document_text(str(path)) == "Текст зі старого .doc файлу"
    assert fake_document.closed_with is False


def test_extract_document_text_doc_open_failure_returns_empty_and_warns(tmp_path, monkeypatch, capsys):
    pytest.importorskip("win32com.client")
    fake_word = _FakeWord(open_error=RuntimeError("Word не відкрив файл"))
    monkeypatch.setattr("win32com.client.Dispatch", lambda name: fake_word)

    path = tmp_path / "file.doc"
    path.write_bytes(b"")

    assert dcs.extract_document_text(str(path)) == ""
    assert "Не вдалось відкрити" in capsys.readouterr().out


def test_extract_document_text_doc_dispatch_failure_returns_empty_and_warns(tmp_path, monkeypatch, capsys):
    """Dispatch("Word.Application") сам падає (напр. Word не встановлено взагалі) -
    зовнішній except (охоплює й саме створення word, а не лише Documents.Open)."""
    pytest.importorskip("win32com.client")

    def _raise(name):
        raise RuntimeError("Word недоступний")

    monkeypatch.setattr("win32com.client.Dispatch", _raise)

    path = tmp_path / "file.doc"
    path.write_bytes(b"")

    assert dcs.extract_document_text(str(path)) == ""
    assert "Не вдалось прочитати" in capsys.readouterr().out


# -------------------------
# find_person_document_references
# -------------------------
def test_find_person_document_references_matches_name_in_paragraph(tmp_path):
    _write_docx(tmp_path / "№79 ЗАВДАННЯ 05.06.2026.docx", paragraphs=["Особовий склад: ВОСЬМИЙ Восьмий Восьмий"])
    roster = {"ВОСЬМИЙ ВОСЬМИЙ ВОСЬМИЙ"}

    result = dcs.find_person_document_references(str(tmp_path), roster, "all")

    assert result == {"ВОСЬМИЙ ВОСЬМИЙ ВОСЬМИЙ": [f"{_BR} №79 від 05.06.2026"]}


def test_find_person_document_references_matches_name_in_table_cell(tmp_path):
    _write_docx(tmp_path / "№27 ЩОДЕННА 01.06.2026.docx", table_rows=[["ВОСЬМИЙ Восьмий Восьмий", "медик"]])
    roster = {"ВОСЬМИЙ ВОСЬМИЙ ВОСЬМИЙ"}

    result = dcs.find_person_document_references(str(tmp_path), roster, "all")

    assert result == {"ВОСЬМИЙ ВОСЬМИЙ ВОСЬМИЙ": [f"{_BR} №27 від 01.06.2026"]}


def test_find_person_document_references_only_task_scope_skips_daily_files(tmp_path):
    _write_docx(tmp_path / "№27 ЩОДЕННА 01.06.2026.docx", paragraphs=["ВОСЬМИЙ Восьмий Восьмий"])
    _write_docx(tmp_path / "№79 ЗАВДАННЯ 05.06.2026.docx", paragraphs=["ВОСЬМИЙ Восьмий Восьмий"])
    roster = {"ВОСЬМИЙ ВОСЬМИЙ ВОСЬМИЙ"}

    result = dcs.find_person_document_references(str(tmp_path), roster, "only_task")

    assert result == {"ВОСЬМИЙ ВОСЬМИЙ ВОСЬМИЙ": [f"{_BR} №79 від 05.06.2026"]}


def test_find_person_document_references_skips_file_with_empty_extracted_text(tmp_path, monkeypatch):
    """Файл відповідає імені ЩОДЕННА/ЗАВДАННЯ, але extract_document_text повернув "" (не
    вдалось прочитати чи файл справді порожній) - пропускається без падіння."""
    path = tmp_path / "№79 ЗАВДАННЯ 05.06.2026.docx"
    path.write_bytes(b"")
    monkeypatch.setattr(dcs, "extract_document_text", lambda file_path: "")

    result = dcs.find_person_document_references(str(tmp_path), {"ВОСЬМИЙ ВОСЬМИЙ ВОСЬМИЙ"}, "all")

    assert result == {}


def test_find_person_document_references_ignores_non_matching_filenames(tmp_path):
    _write_docx(tmp_path / "прев_06_ОБЛІК.docx", paragraphs=["ВОСЬМИЙ Восьмий Восьмий"])
    roster = {"ВОСЬМИЙ ВОСЬМИЙ ВОСЬМИЙ"}

    result = dcs.find_person_document_references(str(tmp_path), roster, "all")

    assert result == {}


def test_find_person_document_references_no_match_absent_from_result(tmp_path):
    _write_docx(tmp_path / "№79 ЗАВДАННЯ 05.06.2026.docx", paragraphs=["ДЕСЯТИЙ Десятий Десятий"])
    roster = {"ВОСЬМИЙ ВОСЬМИЙ ВОСЬМИЙ"}

    result = dcs.find_person_document_references(str(tmp_path), roster, "all")

    assert result == {}


def test_find_person_document_references_accumulates_multiple_documents_for_same_person(tmp_path):
    _write_docx(tmp_path / "№72 ЗАВДАННЯ 01.06.2026.docx", paragraphs=["ВОСЬМИЙ Восьмий Восьмий"])
    _write_docx(tmp_path / "№79 ЗАВДАННЯ 05.06.2026.docx", paragraphs=["ВОСЬМИЙ Восьмий Восьмий"])
    roster = {"ВОСЬМИЙ ВОСЬМИЙ ВОСЬМИЙ"}

    result = dcs.find_person_document_references(str(tmp_path), roster, "all")

    assert result == {"ВОСЬМИЙ ВОСЬМИЙ ВОСЬМИЙ": [
        f"{_BR} №72 від 01.06.2026",
        f"{_BR} №79 від 05.06.2026",
    ]}


# -------------------------
# _format_display_pib
# -------------------------
def test_format_display_pib_surname_upper_rest_titlecase():
    assert dcs._format_display_pib("ТРИНАДЦЯТИЙ ТРИНАДЦЯТИЙ ТРИНАДЦЯТИЙ") == "ТРИНАДЦЯТИЙ Тринадцятий Тринадцятий"
    assert dcs._format_display_pib("Тринадцятий Тринадцятий тринадцятий") == "ТРИНАДЦЯТИЙ Тринадцятий Тринадцятий"


def test_format_display_pib_keeps_apostrophe_casing_correct():
    assert dcs._format_display_pib("ТРИНАДЦЯТИЙ В'ОСЬМИЙ ЧОТИРНАДЦЯТИЙ") == "ТРИНАДЦЯТИЙ В'осьмий Чотирнадцятий"


def test_format_display_pib_empty_string_returns_as_is():
    assert dcs._format_display_pib("") == ""


# -------------------------
# warn_missing_document_coverage
# -------------------------
def test_collect_missing_document_coverage_warnings_includes_rank_when_provided(tmp_path):
    _write_docx(tmp_path / "№79 ЗАВДАННЯ 01.06.2026.docx", paragraphs=["ТРИНАДЦЯТИЙ Тринадцятий Тринадцятий"])
    person_dates = {"ТРИНАДЦЯТИЙ ТРИНАДЦЯТИЙ ТРИНАДЦЯТИЙ": [datetime(2026, 6, 5)]}
    rank_by_pib = {"ТРИНАДЦЯТИЙ ТРИНАДЦЯТИЙ ТРИНАДЦЯТИЙ": "сержант"}

    warnings = dcs.collect_missing_document_coverage_warnings(str(tmp_path), person_dates, "all", rank_by_pib=rank_by_pib)

    assert warnings == ["Немає документа (ЩОДЕННА/ЗАВДАННЯ) за 05.06.2026 для сержант ТРИНАДЦЯТИЙ Тринадцятий Тринадцятий - підстава на цей день відсутня."]


def test_collect_missing_document_coverage_warnings_includes_subdivision_and_rank_in_order(tmp_path):
    person_dates = {"ТРИНАДЦЯТИЙ ТРИНАДЦЯТИЙ ТРИНАДЦЯТИЙ": [datetime(2026, 6, 5)]}
    rank_by_pib = {"ТРИНАДЦЯТИЙ ТРИНАДЦЯТИЙ ТРИНАДЦЯТИЙ": "сержант"}
    subdivision_by_pib = {"ТРИНАДЦЯТИЙ ТРИНАДЦЯТИЙ ТРИНАДЦЯТИЙ": "Підрозділ 1"}

    warnings = dcs.collect_missing_document_coverage_warnings(
        str(tmp_path), person_dates, "all", rank_by_pib=rank_by_pib, subdivision_by_pib=subdivision_by_pib,
    )

    assert warnings == ["Немає документа (ЩОДЕННА/ЗАВДАННЯ) за 05.06.2026 для Підрозділ 1 сержант ТРИНАДЦЯТИЙ Тринадцятий Тринадцятий - підстава на цей день відсутня."]


def test_collect_missing_document_coverage_warnings_subdivision_without_rank(tmp_path):
    person_dates = {"ТРИНАДЦЯТИЙ ТРИНАДЦЯТИЙ ТРИНАДЦЯТИЙ": [datetime(2026, 6, 5)]}
    subdivision_by_pib = {"ТРИНАДЦЯТИЙ ТРИНАДЦЯТИЙ ТРИНАДЦЯТИЙ": "Підрозділ 1"}

    warnings = dcs.collect_missing_document_coverage_warnings(
        str(tmp_path), person_dates, "all", subdivision_by_pib=subdivision_by_pib,
    )

    assert warnings == ["Немає документа (ЩОДЕННА/ЗАВДАННЯ) за 05.06.2026 для Підрозділ 1 ТРИНАДЦЯТИЙ Тринадцятий Тринадцятий - підстава на цей день відсутня."]


def test_collect_missing_document_coverage_warnings_no_rank_omits_leading_space(tmp_path):
    person_dates = {"ТРИНАДЦЯТИЙ ТРИНАДЦЯТИЙ ТРИНАДЦЯТИЙ": [datetime(2026, 6, 5)]}

    warnings = dcs.collect_missing_document_coverage_warnings(str(tmp_path), person_dates, "all")

    assert warnings[0].startswith("Немає документа (ЩОДЕННА/ЗАВДАННЯ) за 05.06.2026 для ТРИНАДЦЯТИЙ Тринадцятий Тринадцятий")


def test_collect_missing_document_coverage_warnings_no_warning_when_every_date_is_covered(tmp_path):
    _write_docx(tmp_path / "№79 ЗАВДАННЯ 05.06.2026.docx", paragraphs=["ВОСЬМИЙ Восьмий Восьмий"])
    person_dates = {"ВОСЬМИЙ ВОСЬМИЙ ВОСЬМИЙ": [datetime(2026, 6, 5)]}

    assert dcs.collect_missing_document_coverage_warnings(str(tmp_path), person_dates, "all") == []


def test_collect_missing_document_coverage_warnings_warns_for_date_without_matching_document(tmp_path):
    _write_docx(tmp_path / "№79 ЗАВДАННЯ 05.06.2026.docx", paragraphs=["ВОСЬМИЙ Восьмий Восьмий"])
    person_dates = {"ВОСЬМИЙ ВОСЬМИЙ ВОСЬМИЙ": [datetime(2026, 6, 5), datetime(2026, 6, 6)]}

    warnings = dcs.collect_missing_document_coverage_warnings(str(tmp_path), person_dates, "all")

    assert len(warnings) == 1
    assert "06.06.2026" in warnings[0]
    assert "ВОСЬМИЙ Восьмий Восьмий" in warnings[0]
    assert "05.06.2026" not in warnings[0]


def test_collect_missing_document_coverage_warnings_document_on_that_date_but_wrong_person_still_warns(tmp_path):
    """Документ датований потрібним днем, але ПІБ людини в ньому НЕ згадується -
    все одно попередження (сам факт наявності файлу на цю дату не є підставою,
    якщо людину в ньому не названо)."""
    _write_docx(tmp_path / "№79 ЗАВДАННЯ 05.06.2026.docx", paragraphs=["ІНШИЙ Інший Інший"])
    person_dates = {"ВОСЬМИЙ ВОСЬМИЙ ВОСЬМИЙ": [datetime(2026, 6, 5)]}

    warnings = dcs.collect_missing_document_coverage_warnings(str(tmp_path), person_dates, "all")

    assert "05.06.2026" in warnings[0]


def test_collect_missing_document_coverage_warnings_only_task_scope_ignores_daily_files(tmp_path):
    _write_docx(tmp_path / "№27 ЩОДЕННА 05.06.2026.docx", paragraphs=["ВОСЬМИЙ Восьмий Восьмий"])
    person_dates = {"ВОСЬМИЙ ВОСЬМИЙ ВОСЬМИЙ": [datetime(2026, 6, 5)]}

    warnings = dcs.collect_missing_document_coverage_warnings(str(tmp_path), person_dates, "only_task")

    assert "05.06.2026" in warnings[0]


def test_collect_missing_document_coverage_warnings_checks_each_person_independently(tmp_path):
    _write_docx(tmp_path / "№79 ЗАВДАННЯ 05.06.2026.docx", paragraphs=["ВОСЬМИЙ Восьмий Восьмий"])
    person_dates = {
        "ВОСЬМИЙ ВОСЬМИЙ ВОСЬМИЙ": [datetime(2026, 6, 5)],
        "ДЕСЯТИЙ ДЕСЯТИЙ ДЕСЯТИЙ": [datetime(2026, 6, 5)],
    }

    warnings = dcs.collect_missing_document_coverage_warnings(str(tmp_path), person_dates, "all")

    assert len(warnings) == 1
    assert "ДЕСЯТИЙ Десятий Десятий" in warnings[0]
