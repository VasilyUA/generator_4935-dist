import os
import sys
from datetime import datetime
from pathlib import Path

sys.path.append(str(Path(__file__).resolve().parent.parent))

import numpy as np
import pytest
from docx import Document
from openpyxl import Workbook, load_workbook

import checker_accounting.report_checker as rc

D1, D2, D3 = datetime(2026, 7, 1), datetime(2026, 7, 2), datetime(2026, 7, 3)


@pytest.fixture(autouse=True)
def _isolate_ignored_basis_lines(monkeypatch):
    """Ізолює ВСІ тести цього файлу від РЕАЛЬНОГО constants.BR_HIGHT_UNIT/
    NUMBER_OF_DOCUMENTS_BRS_EVERY_WEEK/_EVERY_DAY (той самий підхід, що й
    _no_general_money_references у tests/test_money_report.py для
    NUMBER_OF_DOCUMENTS_BRS_EVERY_WEEK) - ці словники, як і решта "динамічних"
    об'єктів constants.py (MONEY_REPORT_CATEGORIES/_CHANGES_ORDER_REFERENCES/
    _GENERAL_REFERENCES_*/COMMANDER_MONEY_REPORT_CATEGORIES), ЩОМІСЯЦЯ
    наповнюються РЕАЛЬНИМИ даними користувача - тест НЕ МАЄ покладатись на їхній
    ПОТОЧНИЙ вміст (інакше зламається наступного місяця, коли вміст зміниться,
    хоча сам код лишиться правильним). Тести, яким справді потрібен виняток
    (вищий штаб / тижневий БР / 'брг' щоденного БР), підміняють
    rc._IGNORED_BASIS_CITATIONS/rc._IGNORED_WEEKLY_BR_NUMBERS власним,
    синтетичним значенням локально (monkeypatch дозволяє це поверх фікстури)."""
    monkeypatch.setattr(rc, "_IGNORED_BASIS_CITATIONS", set())
    monkeypatch.setattr(rc, "_IGNORED_WEEKLY_BR_NUMBERS", set())


def _write_oblik_xlsx(path, headers, rows):
    wb = Workbook()
    ws = wb.active
    ws.title = "Аркуш1"
    ws.append(headers)
    for row in rows:
        ws.append(row)
    wb.save(str(path))
    return str(path)


def _write_report(path, sections):
    """sections - список (heading_text, data_rows); data_rows - [посада, звання,
    піб, період, дні, підстава] - той самий формат, що й tests/test_report_document_reader.py."""
    doc = Document()
    for heading_text, data_rows in sections:
        doc.add_heading(heading_text, level=1)
        table = doc.add_table(rows=1 + len(data_rows), cols=7)
        for row_idx, values in enumerate(data_rows, start=1):
            for col_idx, value in enumerate(values, start=1):
                table.rows[row_idx].cells[col_idx].text = value
    doc.save(str(path))
    return str(path)


def _use_fixture(monkeypatch, oblik_path, columns_letters):
    monkeypatch.setattr(rc, "PERSONEL_LIST_FILE_NAME", oblik_path)
    monkeypatch.setattr(rc, "PERSONEL_LIST_SHEET_NAME", "Аркуш1")
    monkeypatch.setattr(rc, "PERSONEL_LIST_COLUMNS_LETTERS", columns_letters)


def _read_result_rows(output_path):
    """Текст КОЖНОГО рядка (крім заголовка) аркуша .xlsx, збереженого
    check_report_against_oblik - усі колонки (rc._RESULT_HEADERS - ПІБ/
    Повідомлення/Ресурс/Шлях/Дата періоду/Документ) одного рядка об'єднані в
    один рядок тексту (щоб тестам не важливо було, в якій САМЕ колонці опинився
    шуканий фрагмент) - список, а не множина: порядок findings інколи важливий
    для тестів."""
    ws = load_workbook(output_path).active
    return [
        " | ".join(str(cell.value) for cell in row if cell.value is not None)
        for row in ws.iter_rows(min_row=2)
        if any(cell.value is not None for cell in row)
    ]


# -------------------------
# _blank
# -------------------------
def test_blank_recognizes_none_empty_string_and_nan():
    assert rc._blank(None) is True
    assert rc._blank("") is True
    assert rc._blank(np.nan) is True
    assert rc._blank(0) is False
    assert rc._blank("30") is False
    assert rc._blank(30) is False


# -------------------------
# _check_pib_spelling
# -------------------------
def test_check_pib_spelling_flags_missing_person():
    entries = [{"ПІБ": "ПЕРШИЙ Перший Перший", "raw_value": 30, "ПОСАДА": "", "ЗВАННЯ": "", "ПЕРІОД": "", "ДНІ": "0"}]
    findings = rc._check_pib_spelling(entries, {})
    assert len(findings) == 1
    assert findings[0]["ПІБ"] == "ПЕРШИЙ Перший Перший"
    assert "не знайдено" in findings[0]["Повідомлення"]
    assert findings[0]["Ресурс"] == rc.PERSONEL_LIST_FILE_NAME


def test_check_pib_spelling_flags_mismatched_spelling():
    entries = [{"ПІБ": "Другий Другий", "raw_value": 30, "ПОСАДА": "", "ЗВАННЯ": "", "ПЕРІОД": "", "ДНІ": "0"}]
    oblik_by_pib = {"ДРУГИЙ ДРУГИЙ": {"ПІБ": "Другий  Другий"}}  # зайвий пробіл в ОБЛІК

    findings = rc._check_pib_spelling(entries, oblik_by_pib)

    assert len(findings) == 1
    assert "по-різному" in findings[0]["Повідомлення"]


def test_check_pib_spelling_no_finding_when_exact_match():
    entries = [{"ПІБ": "Другий Другий", "raw_value": 30, "ПОСАДА": "", "ЗВАННЯ": "", "ПЕРІОД": "", "ДНІ": "0"}]
    oblik_by_pib = {"ДРУГИЙ ДРУГИЙ": {"ПІБ": "Другий Другий"}}
    assert rc._check_pib_spelling(entries, oblik_by_pib) == []


def test_check_pib_spelling_deduplicates_same_person_across_entries():
    """Та сама людина в ДВОХ рядках рапорту (напр. РУДЕНКО - госпіталізація й
    відпустка для лікування різними періодами) - лише ОДНЕ повідомлення, не два."""
    entries = [
        {"ПІБ": "Третій Третій", "raw_value": "100_ШП", "ПОСАДА": "", "ЗВАННЯ": "", "ПЕРІОД": "", "ДНІ": "0"},
        {"ПІБ": "Третій Третій", "raw_value": "100_ВП", "ПОСАДА": "", "ЗВАННЯ": "", "ПЕРІОД": "", "ДНІ": "0"},
    ]
    assert len(rc._check_pib_spelling(entries, {})) == 1


# -------------------------
# _check_periods_and_days
# -------------------------
def test_check_periods_and_days_no_findings_when_everything_matches():
    entries = [{"ПІБ": "Другий Другий", "raw_value": 30, "ПЕРІОД": "01.07.2026-02.07.2026", "ДНІ": "2"}]
    oblik_by_pib = {"ДРУГИЙ ДРУГИЙ": {"ПІБ": "Другий Другий", D1: 30, D2: 30}}
    assert rc._check_periods_and_days(entries, oblik_by_pib) == []


def test_check_periods_and_days_flags_day_count_mismatch():
    entries = [{"ПІБ": "Другий Другий", "raw_value": 30, "ПЕРІОД": "01.07.2026-02.07.2026", "ДНІ": "5"}]
    oblik_by_pib = {"ДРУГИЙ ДРУГИЙ": {"ПІБ": "Другий Другий", D1: 30, D2: 30}}

    findings = rc._check_periods_and_days(entries, oblik_by_pib)

    assert len(findings) == 1
    assert "2 дн." in findings[0]["Повідомлення"] and "5" in findings[0]["Повідомлення"]


def test_check_periods_and_days_flags_non_numeric_days_text():
    entries = [{"ПІБ": "Другий Другий", "raw_value": 30, "ПЕРІОД": "01.07.2026-02.07.2026", "ДНІ": "два"}]
    oblik_by_pib = {"ДРУГИЙ ДРУГИЙ": {"ПІБ": "Другий Другий", D1: 30, D2: 30}}

    findings = rc._check_periods_and_days(entries, oblik_by_pib)

    assert any("не число" in f["Повідомлення"] for f in findings)


def test_check_periods_and_days_flags_blank_oblik_cell():
    entries = [{"ПІБ": "Другий Другий", "raw_value": 30, "ПЕРІОД": "01.07.2026-02.07.2026", "ДНІ": "2"}]
    oblik_by_pib = {"ДРУГИЙ ДРУГИЙ": {"ПІБ": "Другий Другий", D1: 30, D2: None}}

    findings = rc._check_periods_and_days(entries, oblik_by_pib)

    assert len(findings) == 1
    assert "02.07.2026" in findings[0]["Повідомлення"] and "порожня" in findings[0]["Повідомлення"]


def test_check_periods_and_days_skips_oblik_lookup_when_person_not_found():
    """Людину не знайдено в ОБЛІК.xlsx - _check_pib_spelling вже про це повідомила,
    тут не дублюємо ще й "порожня клітинка" для кожного дня."""
    entries = [{"ПІБ": "Невідомий Хтось", "raw_value": 30, "ПЕРІОД": "01.07.2026-02.07.2026", "ДНІ": "2"}]
    assert rc._check_periods_and_days(entries, {}) == []


def test_check_periods_and_days_skips_spetskontyngent_entries_entirely():
    """СПЕЦКОНТИНГЕНТ має ЗОВСІМ іншу таблицю ("Дата зникнення безвісти"/"Період
    виплати") - "ДНІ" там не про кількість днів (тут - навмисно "сміттєвий"
    текст, який зламав би int()), і день-за-днем перевіряється ОКРЕМО, в
    _check_spetskontyngent, а не тут."""
    entries = [{
        "ПІБ": "Четвертий Четвертий", "raw_value": rc.SPETSKONTYNGENT_VALUE,
        "ПЕРІОД": "01.07.2026-31.07.2026", "ДНІ": "01.07.2026-31.07.2026",
    }]
    assert rc._check_periods_and_days(entries, {"ЧЕТВЕРТИЙ ЧЕТВЕРТИЙ": {"ПІБ": "Четвертий Четвертий"}}) == []


# -------------------------
# _check_periods_self_consistent / check_report_periods (легкий інструмент -
# лише сам рапорт, без ОБЛІК.xlsx чи будь-якого іншого зовнішнього джерела)
# -------------------------
def test_check_periods_self_consistent_no_findings_when_days_match():
    entries = [{"ПІБ": "Другий Другий", "raw_value": 30, "ПЕРІОД": "01.07.2026-02.07.2026", "ДНІ": "2"}]
    assert rc._check_periods_self_consistent(entries, "report.docx") == []


def test_check_periods_self_consistent_flags_day_count_mismatch():
    entries = [{"ПІБ": "Другий Другий", "raw_value": 30, "ПЕРІОД": "01.07.2026-02.07.2026", "ДНІ": "5"}]

    findings = rc._check_periods_self_consistent(entries, "report.docx")

    assert len(findings) == 1
    assert "2 дн." in findings[0]["Повідомлення"] and "5" in findings[0]["Повідомлення"]
    assert findings[0]["Ресурс"] == "report.docx"


def test_check_periods_self_consistent_flags_non_numeric_days_text():
    entries = [{"ПІБ": "Другий Другий", "raw_value": 30, "ПЕРІОД": "01.07.2026-02.07.2026", "ДНІ": "два"}]

    findings = rc._check_periods_self_consistent(entries, "report.docx")

    assert len(findings) == 1 and "не число" in findings[0]["Повідомлення"]


def test_check_periods_self_consistent_skips_spetskontyngent_entries():
    entries = [{
        "ПІБ": "Четвертий Четвертий", "raw_value": rc.SPETSKONTYNGENT_VALUE,
        "ПЕРІОД": "01.07.2026-31.07.2026", "ДНІ": "01.07.2026-31.07.2026",
    }]
    assert rc._check_periods_self_consistent(entries, "report.docx") == []


def test_check_report_periods_end_to_end_finds_discrepancy(tmp_path, monkeypatch):
    """Ключова відмінність від check_report_against_oblik: жодного ОБЛІК.xlsx
    тут немає взагалі - findings лише за самоузгодженістю самого рапорту."""
    monkeypatch.setattr(rc, "OUTPUT_DIR", str(tmp_path))
    report_path = _write_report(tmp_path / "report.docx", [
        ("1. Виплатити додаткову винагороду у розмірі 30 000 грн. 00 коп. військовослужбовцям", [
            ["Стрілець", "сержант", "СЬОМИЙ Сьомий Сьомий", "01.07.2026-03.07.2026", "5", ""],
        ]),
    ])

    output_path = rc.check_report_periods(report_path)

    assert output_path == str(tmp_path / "Перевірка рапорту (періоди-дні).xlsx")
    content = _read_result_rows(output_path)
    assert any("3 дн." in row and "5" in row for row in content)


def test_check_report_periods_writes_clean_message_when_consistent(tmp_path, monkeypatch):
    monkeypatch.setattr(rc, "OUTPUT_DIR", str(tmp_path))
    report_path = _write_report(tmp_path / "report.docx", [
        ("1. Виплатити додаткову винагороду у розмірі 30 000 грн. 00 коп. військовослужбовцям", [
            ["Стрілець", "сержант", "СЬОМИЙ Сьомий Сьомий", "01.07.2026-02.07.2026", "2", ""],
        ]),
    ])

    output_path = rc.check_report_periods(report_path)

    assert _read_result_rows(output_path) == ["Невідповідностей між періодом і кількістю днів не знайдено."]


def test_check_report_periods_uses_custom_output_path(tmp_path):
    report_path = _write_report(tmp_path / "report.docx", [
        ("1. Виплатити додаткову винагороду у розмірі 30 000 грн. 00 коп. військовослужбовцям", [
            ["Стрілець", "сержант", "СЬОМИЙ Сьомий Сьомий", "01.07.2026-01.07.2026", "1", ""],
        ]),
    ])
    custom_output = tmp_path / "custom" / "result.xlsx"

    output_path = rc.check_report_periods(report_path, output_path=str(custom_output))

    assert output_path == str(custom_output)
    assert custom_output.exists()


# -------------------------
# _is_death_marker / _oblik_disappearance_date
# -------------------------
def test_is_death_marker_recognizes_int_and_string_200():
    assert rc._is_death_marker(200) is True
    assert rc._is_death_marker("200") is True
    assert rc._is_death_marker(" 200 ") is True
    assert rc._is_death_marker(30) is False
    assert rc._is_death_marker(rc.SPETSKONTYNGENT_VALUE) is False


def test_is_death_marker_recognizes_real_zagybel_text():
    """Реальні дані (реальний випадок, повідомлений користувачем) показують, що
    ОБЛІК.xlsx позначає загибель буквальним текстом "загибель", а не числом "200"."""
    assert rc._is_death_marker("загибель") is True
    assert rc._is_death_marker("Загибель") is True
    assert rc._is_death_marker(" загибель ") is True
    assert rc._is_death_marker("полон") is False


# oblik_disappearance_date - перенесено в content.money_report_helpers (спільне
# з build_category_rows/"Очк.БЗ") - тест тепер у tests/test_money_report.py.


# -------------------------
# _check_spetskontyngent
# -------------------------
def _spk_entry(pib, period, disappearance_date):
    return {"ПІБ": pib, "raw_value": rc.SPETSKONTYNGENT_VALUE, "ПЕРІОД": period, "ДНІ": period, "ДАТА_ЗНИКНЕННЯ": disappearance_date}


def _use_fixed_month(monkeypatch):
    """Усі тести _check_spetskontyngent мають знати ТОЧНИЙ "кінець місяця" -
    фіксуємо MONTH/YEAR (а не покладаємось на те, який місяць зараз реально
    веде користувач у справжньому ОБЛІК.xlsx)."""
    monkeypatch.setattr(rc, "MONTH", "07")
    monkeypatch.setattr(rc, "YEAR", "2026")


def test_check_spetskontyngent_no_findings_when_period_ends_the_day_before_death(monkeypatch):
    """"Повний період" - НЕ обов'язково до кінця місяця: реальні дані (реальний
    випадок, повідомлений користувачем) показують, що маркер загибелі в ОБЛІК.xlsx
    з'являється ПОЧИНАЮЧИ з дня, наступного за останнім активним статусом
    (22.07 - ще статус, 23.07 - вже "загибель") - тож період рапорту, що
    закінчується РІВНО днем ПЕРЕД першим таким маркером (тут - D2), вже
    ПОВНИЙ (очікуваний) період, не розбіжність."""
    _use_fixed_month(monkeypatch)
    entries = [_spk_entry("Четвертий Четвертий", "01.07.2026-01.07.2026", "11.08.2024")]
    oblik_by_pib = {"ЧЕТВЕРТИЙ ЧЕТВЕРТИЙ": {"ПІБ": "Четвертий Четвертий", D1: "полон", D2: "загибель", "ДАТА ЗНИКНЕННЯ": "11.08.2024"}}
    assert rc._check_spetskontyngent(entries, oblik_by_pib) == []


def test_check_spetskontyngent_flags_disappearance_date_mismatch(monkeypatch):
    _use_fixed_month(monkeypatch)
    entries = [_spk_entry("Четвертий Четвертий", "01.07.2026-01.07.2026", "11.08.2024")]
    oblik_by_pib = {"ЧЕТВЕРТИЙ ЧЕТВЕРТИЙ": {"ПІБ": "Четвертий Четвертий", D1: "полон", D2: 200, "ДАТА ЗНИКНЕННЯ": "12.08.2024"}}

    findings = rc._check_spetskontyngent(entries, oblik_by_pib)

    assert len(findings) == 1
    assert "дата зникнення" in findings[0]["Повідомлення"]


def test_check_spetskontyngent_flags_unparseable_disappearance_date(monkeypatch):
    """Дата "зникнення" в ОБЛІК.xlsx (чи в рапорті) у форматі, який to_date не
    розпізнає (ValueError) - теж розбіжність, а не падіння звірки."""
    _use_fixed_month(monkeypatch)
    entries = [_spk_entry("Четвертий Четвертий", "01.07.2026-01.07.2026", "11.08.2024")]
    oblik_by_pib = {"ЧЕТВЕРТИЙ ЧЕТВЕРТИЙ": {"ПІБ": "Четвертий Четвертий", D1: "полон", D2: 200, "ДАТА ЗНИКНЕННЯ": "не дата"}}

    findings = rc._check_spetskontyngent(entries, oblik_by_pib)

    assert len(findings) == 1
    assert "дата зникнення" in findings[0]["Повідомлення"]


def test_check_spetskontyngent_flags_missing_disappearance_date_in_oblik(monkeypatch):
    _use_fixed_month(monkeypatch)
    entries = [_spk_entry("Четвертий Четвертий", "01.07.2026-01.07.2026", "11.08.2024")]
    oblik_by_pib = {"ЧЕТВЕРТИЙ ЧЕТВЕРТИЙ": {"ПІБ": "Четвертий Четвертий", D1: "полон", D2: 200}}  # немає жодної колонки дати

    findings = rc._check_spetskontyngent(entries, oblik_by_pib)

    assert len(findings) == 1
    assert "дата зникнення" in findings[0]["Повідомлення"]


def test_check_spetskontyngent_skips_date_check_when_report_has_no_date(monkeypatch):
    _use_fixed_month(monkeypatch)
    entries = [_spk_entry("Четвертий Четвертий", "01.07.2026-01.07.2026", "")]
    oblik_by_pib = {"ЧЕТВЕРТИЙ ЧЕТВЕРТИЙ": {"ПІБ": "Четвертий Четвертий", D1: "полон", D2: 200}}
    assert rc._check_spetskontyngent(entries, oblik_by_pib) == []


def test_check_spetskontyngent_flags_blank_participation_day(monkeypatch):
    """Смерть трапляється лише на день4 (04.07) - за межами періоду блоку, тож
    единственна розбіжність тут - порожня комірка на D2, без сторонніх
    "період триває довше" повідомлень."""
    _use_fixed_month(monkeypatch)
    entries = [_spk_entry("Четвертий Четвертий", "01.07.2026-03.07.2026", "11.08.2024")]
    oblik_by_pib = {
        "ЧЕТВЕРТИЙ ЧЕТВЕРТИЙ": {
            "ПІБ": "Четвертий Четвертий", D1: "полон", D2: None, D3: "полон",
            datetime(2026, 7, 4): "загибель", "ДАТА ЗНИКНЕННЯ": "11.08.2024",
        },
    }

    findings = rc._check_spetskontyngent(entries, oblik_by_pib)

    assert len(findings) == 1
    assert "02.07.2026" in findings[0]["Повідомлення"] and "порожня" in findings[0]["Повідомлення"]


def test_check_spetskontyngent_stops_checking_after_death_marker(monkeypatch):
    """Регресія: якщо в ОБЛІК.xlsx трапляється маркер загибелі на якийсь день
    періоду, дні ПІСЛЯ нього не повинні позначатись як "порожні" - людина
    вибула зі статусу, а не просто пропущена (про сам факт, що рапорт помилково
    продовжує рахувати участь ПІСЛЯ загибелі, вже сигналізує ОКРЕМЕ повідомлення
    про довжину періоду)."""
    _use_fixed_month(monkeypatch)
    entries = [_spk_entry("Четвертий Четвертий", "01.07.2026-03.07.2026", "11.08.2024")]
    oblik_by_pib = {"ЧЕТВЕРТИЙ ЧЕТВЕРТИЙ": {"ПІБ": "Четвертий Четвертий", D1: "полон", D2: "загибель", D3: None, "ДАТА ЗНИКНЕННЯ": "11.08.2024"}}

    findings = rc._check_spetskontyngent(entries, oblik_by_pib)

    assert not any("порожня" in f["Повідомлення"] for f in findings)
    assert any("період у рапорті" in f["Повідомлення"] for f in findings)


def test_check_spetskontyngent_flags_period_continuing_past_death(monkeypatch):
    """Головний сценарій, про який повідомив користувач: людина зі статусом
    СПЕЦКОНТИНГЕНТ по 22.07 включно, з 23.07 -
    "загибель" в ОБЛІК.xlsx, а рапорт помилково продовжує рахувати період аж
    до 30.07 - розбіжність (мав закінчитись 22.07)."""
    _use_fixed_month(monkeypatch)
    entries = [_spk_entry("ВІСІМНАДЦЯТИЙ Вісімнадцятий Вісімнадцятий", "01.07.2026-30.07.2026", "")]
    oblik_row = {"ПІБ": "ВІСІМНАДЦЯТИЙ Вісімнадцятий Вісімнадцятий"}
    for day in range(1, 23):
        oblik_row[datetime(2026, 7, day)] = "100_СПЕЦКОНТИНГЕНТ"
    for day in range(23, 32):
        oblik_row[datetime(2026, 7, day)] = "загибель"
    oblik_by_pib = {"ВІСІМНАДЦЯТИЙ ВІСІМНАДЦЯТИЙ ВІСІМНАДЦЯТИЙ": oblik_row}

    findings = rc._check_spetskontyngent(entries, oblik_by_pib)

    period_findings = [f["Повідомлення"] for f in findings if "період у рапорті" in f["Повідомлення"]]
    assert len(period_findings) == 1
    assert "30.07.2026" in period_findings[0]
    assert "22.07.2026" in period_findings[0]
    assert "триває до" in period_findings[0]


# -------------------------
# _month_end_or_death_date
# -------------------------
def test_month_end_or_death_date_returns_last_day_of_month_without_death_marker():
    oblik_row = {D1: "полон", D2: "полон"}
    assert rc._month_end_or_death_date(oblik_row, 7, 2026) == datetime(2026, 7, 31)


def test_month_end_or_death_date_returns_day_before_first_death_marker():
    oblik_row = {D1: "полон", D2: "загибель", D3: None}
    assert rc._month_end_or_death_date(oblik_row, 7, 2026) == D1


# -------------------------
# _check_spetskontyngent - повний період (до кінця місяця чи до загибелі)
# -------------------------
def test_check_spetskontyngent_flags_period_shorter_than_full_month(monkeypatch):
    """Головний сценарій, про який повідомив користувач: рапорт вказує період,
    що закінчується РАНІШЕ кінця місяця, хоча в ОБЛІК.xlsx немає жодної "200"
    (загибелі), яка виправдовувала б коротший період - людина мала б і далі
    рахуватись СПЕЦКОНТИНГЕНТОМ до кінця місяця."""
    _use_fixed_month(monkeypatch)
    entries = [_spk_entry("Четвертий Четвертий", "01.07.2026-15.07.2026", "11.08.2024")]
    oblik_by_pib = {"ЧЕТВЕРТИЙ ЧЕТВЕРТИЙ": {"ПІБ": "Четвертий Четвертий", D1: "полон", "ДАТА ЗНИКНЕННЯ": "11.08.2024"}}

    findings = rc._check_spetskontyngent(entries, oblik_by_pib)

    period_findings = [f["Повідомлення"] for f in findings if "період у рапорті закінчується" in f["Повідомлення"]]
    assert len(period_findings) == 1
    assert "15.07.2026" in period_findings[0]
    assert "31.07.2026" in period_findings[0]
    assert "до кінця місяця" in period_findings[0]


def test_check_spetskontyngent_does_not_flag_period_shortened_by_death(monkeypatch):
    """Той самий "короткий" (відносно кінця місяця) період - але цього разу
    ОБЛІК.xlsx підтверджує, що загибель настає одразу НАСТУПНОГО дня після
    закінчення періоду рапорту (15.07) - тобто період рапорту закінчується
    РІВНО в очікувану дату (14.07 - день перед загибеллю) - НЕ розбіжність
    (підтверджено користувачем). Проміжні дні (02-14.07) тут не заповнені
    навмисно - це поза межами цього тесту (день-за-днем перевіряє окремий
    тест), тож дивимось САМЕ на відсутність повідомлення про довжину періоду."""
    _use_fixed_month(monkeypatch)
    entries = [_spk_entry("Четвертий Четвертий", "01.07.2026-14.07.2026", "11.08.2024")]
    oblik_by_pib = {
        "ЧЕТВЕРТИЙ ЧЕТВЕРТИЙ": {"ПІБ": "Четвертий Четвертий", D1: "полон", datetime(2026, 7, 15): "загибель", "ДАТА ЗНИКНЕННЯ": "11.08.2024"},
    }

    findings = rc._check_spetskontyngent(entries, oblik_by_pib)

    assert not any("період у рапорті" in f["Повідомлення"] for f in findings)


def test_check_spetskontyngent_flags_no_dates_at_all(monkeypatch):
    _use_fixed_month(monkeypatch)
    entries = [_spk_entry("Четвертий Четвертий", "", "11.08.2024")]
    oblik_by_pib = {"ЧЕТВЕРТИЙ ЧЕТВЕРТИЙ": {"ПІБ": "Четвертий Четвертий", "ДАТА ЗНИКНЕННЯ": "11.08.2024"}}

    findings = rc._check_spetskontyngent(entries, oblik_by_pib)

    period_findings = [f["Повідомлення"] for f in findings if "період у рапорті закінчується" in f["Повідомлення"]]
    assert len(period_findings) == 1
    assert "немає жодної дати" in period_findings[0]


def test_check_spetskontyngent_merges_multiple_entries_for_same_person(monkeypatch):
    """Та сама людина в ДВОХ рядках СПЕЦКОНТИНГЕНТУ (напр. окремі записи за
    різні періоди) - дати об'єднуються перед перевіркою "повного місяця", а не
    перевіряються (і не звітуються) окремо на кожен рядок."""
    _use_fixed_month(monkeypatch)
    entries = [
        _spk_entry("Четвертий Четвертий", "01.07.2026-15.07.2026", "11.08.2024"),
        _spk_entry("Четвертий Четвертий", "16.07.2026-31.07.2026", ""),
    ]
    oblik_by_pib = {"ЧЕТВЕРТИЙ ЧЕТВЕРТИЙ": {"ПІБ": "Четвертий Четвертий", D1: "полон", "ДАТА ЗНИКНЕННЯ": "11.08.2024"}}

    findings = rc._check_spetskontyngent(entries, oblik_by_pib)

    assert not any("період у рапорті закінчується" in f["Повідомлення"] for f in findings)


def test_check_spetskontyngent_skips_date_and_blank_day_checks_when_person_not_found_in_oblik(monkeypatch):
    """Людину не знайдено в ОБЛІК.xlsx узагалі (напр. прибрано з поточного
    ОБЛІК.xlsx після визнання зниклою безвісти) - перевірки, які ПОТРЕБУЮТЬ
    рядка ОБЛІК.xlsx (дата зникнення, день-за-днем) пропускаються (вже є
    окреме "не знайдено" від _check_pib_spelling), але період все одно
    ПОВНИЙ (до кінця місяця) - жодної розбіжності немає взагалі."""
    _use_fixed_month(monkeypatch)
    entries = [_spk_entry("Невідомий Хтось", "01.07.2026-31.07.2026", "11.08.2024")]
    assert rc._check_spetskontyngent(entries, {}) == []


def test_check_spetskontyngent_flags_short_period_even_when_person_not_found_in_oblik(monkeypatch):
    """Регресія на реальний випадок, повідомлений користувачем: троє людей
    прибрані з поточного ОБЛІК.xlsx (напр. після
    визнання зниклими безвісти), тому "не знайдено в ОБЛІК.xlsx" - але
    перевірка "період до кінця місяця" НЕ повинна через це мовчати: вона
    самодостатня (звіряє лише ПЕРІОД рапорту з календарем), і саме тому
    користувач очікував побачити цю розбіжність, а бачив тільки "не знайдено"."""
    _use_fixed_month(monkeypatch)
    entries = [_spk_entry("ШІСТНАДЦЯТИЙ Шістнадцятий Шістнадцятий", "01.07.2026-30.07.2026", "11.08.2024")]

    findings = rc._check_spetskontyngent(entries, {})

    period_findings = [f["Повідомлення"] for f in findings if "період у рапорті закінчується" in f["Повідомлення"]]
    assert len(period_findings) == 1
    assert "30.07.2026" in period_findings[0] and "31.07.2026" in period_findings[0]
    assert "до кінця місяця" in period_findings[0]
    # Дата зникнення й день-за-днем ПОТРЕБУЮТЬ рядка ОБЛІК.xlsx - без нього не звітуються.
    assert not any("дата зникнення" in f["Повідомлення"] for f in findings)
    assert not any("порожня" in f["Повідомлення"] for f in findings)


def test_check_spetskontyngent_ignores_non_spetskontyngent_entries():
    entries = [{"ПІБ": "Другий Другий", "raw_value": 30, "ПЕРІОД": "01.07.2026-01.07.2026", "ДНІ": "1"}]
    assert rc._check_spetskontyngent(entries, {}) == []


# -------------------------
# check_report_against_oblik
# -------------------------
def test_check_report_against_oblik_end_to_end_finds_discrepancies(tmp_path, monkeypatch):
    monkeypatch.setattr(rc, "OUTPUT_DIR", str(tmp_path))
    oblik_path = _write_oblik_xlsx(tmp_path / "ОБЛІК.xlsx", ["ПІДРОЗДІЛ", "ПОСАДА", "ЗВАННЯ", "ПІБ", D1, D2, D3], [
        ["Підрозділ 1", "Стрілець", "сержант", "СЬОМИЙ Сьомий Сьомий", 30, 30, None],
    ])
    _use_fixture(monkeypatch, oblik_path, ["A", "B", "C", "D", "E", "F", "G"])

    report_path = _write_report(tmp_path / "report.docx", [
        ("1. Виплатити додаткову винагороду у розмірі 30 000 грн. 00 коп. військовослужбовцям", [
            ["Стрілець", "сержант", "СЬОМИЙ Сьомий Сьомий", "01.07.2026-03.07.2026", "3", ""],
        ]),
        ("2. Виплатити додаткову винагороду в розмірі 100 000 грн. 00 коп. військовослужбовцям", [
            ["Стрілець", "сержант", "Шостий Шостий", "01.07.2026-01.07.2026", "1", ""],
        ]),
    ])

    output_path = rc.check_report_against_oblik(report_path)

    assert output_path == str(tmp_path / "Звірка рапорту з ОБЛІК.xlsx")
    content = _read_result_rows(output_path)
    assert any("Шостий Шостий" in row and "не знайдено" in row for row in content)
    assert any("03.07.2026" in row and "порожня" in row for row in content)


def test_check_report_against_oblik_writes_clean_message_when_no_discrepancies(tmp_path, monkeypatch):
    monkeypatch.setattr(rc, "OUTPUT_DIR", str(tmp_path))
    oblik_path = _write_oblik_xlsx(tmp_path / "ОБЛІК.xlsx", ["ПІДРОЗДІЛ", "ПОСАДА", "ЗВАННЯ", "ПІБ", D1, D2], [
        ["Підрозділ 1", "Стрілець", "сержант", "СЬОМИЙ Сьомий Сьомий", 30, 30],
    ])
    _use_fixture(monkeypatch, oblik_path, ["A", "B", "C", "D", "E", "F"])

    report_path = _write_report(tmp_path / "report.docx", [
        ("1. Виплатити додаткову винагороду у розмірі 30 000 грн. 00 коп. військовослужбовцям", [
            ["Стрілець", "сержант", "СЬОМИЙ Сьомий Сьомий", "01.07.2026-02.07.2026", "2", ""],
        ]),
    ])

    output_path = rc.check_report_against_oblik(report_path)

    assert _read_result_rows(output_path) == ["Розбіжностей між рапортом і ОБЛІК.xlsx не знайдено."]


def test_check_report_against_oblik_uses_custom_output_path(tmp_path, monkeypatch):
    oblik_path = _write_oblik_xlsx(tmp_path / "ОБЛІК.xlsx", ["ПІДРОЗДІЛ", "ПОСАДА", "ЗВАННЯ", "ПІБ", D1], [
        ["Підрозділ 1", "Стрілець", "сержант", "СЬОМИЙ Сьомий Сьомий", 30],
    ])
    _use_fixture(monkeypatch, oblik_path, ["A", "B", "C", "D", "E"])
    report_path = _write_report(tmp_path / "report.docx", [
        ("1. Виплатити додаткову винагороду у розмірі 30 000 грн. 00 коп. військовослужбовцям", [
            ["Стрілець", "сержант", "СЬОМИЙ Сьомий Сьомий", "01.07.2026-01.07.2026", "1", ""],
        ]),
    ])
    custom_output = tmp_path / "custom" / "result.xlsx"

    output_path = rc.check_report_against_oblik(report_path, output_path=str(custom_output))

    assert output_path == str(custom_output)
    assert custom_output.exists()


# -------------------------
# _number_and_date
# -------------------------
def test_number_and_date_extracts_both_when_present():
    assert rc._number_and_date("БР 2169 №2217 від 03.07.2026") == ("2217", "03.07.2026")


def test_number_and_date_strips_leading_punctuation_from_number():
    """Той самий номер записаний "№-1622" (constants.py, вручну) чи "№ 1622"/
    "№1622" (як його фактично передрукували в готовий рапорт) - має нормалізуватись
    до ОДНАКОВОГО значення, інакше звірка з _IGNORED_BASIS_CITATIONS ніколи б
    не збіглася через суто косметичну різницю запису."""
    assert rc._number_and_date("БР 2169 №-1622 від-16.05.2026") == ("1622", "16.05.2026")
    assert rc._number_and_date("БР 2169 № 1622від 16.05.2026") == ("1622", "16.05.2026")
    assert rc._number_and_date("БР 2169 №1622 від 16.05.2026") == ("1622", "16.05.2026")


def test_number_and_date_none_none_without_date():
    assert rc._number_and_date("БР 2169 №2217") == (None, None)


def test_number_and_date_none_number_without_number_sign():
    assert rc._number_and_date("БР 2169 від 03.07.2026") == (None, "03.07.2026")


# -------------------------
# _reference_line_citation / _reference_citations / _all_reference_lines - на
# СИНТЕТИЧНИХ даних, незалежно від ПОТОЧНОГО (щомісяця змінюваного) реального
# вмісту constants.BR_HIGHT_UNIT.
# -------------------------
def test_reference_line_citation_extracts_br_line():
    assert rc._reference_line_citation("БР 0 БАТ 2169 №2217 від 03.07.2026") == ("БР", "2217", "03.07.2026")


def test_reference_line_citation_none_for_bn_or_pozbd_line():
    assert rc._reference_line_citation("БН 0 БАТ 2169 №1 від 08.07.2026") is None
    assert rc._reference_line_citation("ПозБД 2169 №242 від 18.06.2026") is None


def test_reference_line_citation_none_without_date():
    assert rc._reference_line_citation("ЖБД (Справа №2т)") is None


def test_reference_line_citation_normalizes_embedded_newlines_and_spacing():
    assert rc._reference_line_citation("ЖБД 0 БАТ 2169 №1\n від   01.07.2026") == ("ЖБД", "1", "01.07.2026")


def test_reference_citations_collects_from_multiple_lines_skipping_untracked():
    lines = ["БР 0 БАТ 2169 №1 від 01.07.2026", "БН 0 БАТ 2169 №1 від 02.07.2026", "ЖБД (Справа №2т)"]
    assert rc._reference_citations(lines) == {("БР", "1", "01.07.2026")}


def test_reference_citations_empty_for_empty_input():
    assert rc._reference_citations([]) == set()


def test_all_reference_lines_collects_from_all_groups_and_references():
    references_by_group = {
        "general": [{"start": "01.07.2026", "end": None, "lines": ["ЖБД (Справа №1)"]}],
        "ЖИТТЄДІЯЛЬНІСТЬ": [
            {"start": "01.07.2026", "end": None, "lines": ["БР 0 БАТ 2169 №1 від 01.07.2026"]},
            {"start": "02.07.2026", "end": None, "lines": ["БР 0 БАТ 2169 №2 від 02.07.2026", "БН 0 БАТ 2169 №1 від 02.07.2026"]},
        ],
    }
    assert rc._all_reference_lines(references_by_group) == [
        "ЖБД (Справа №1)",
        "БР 0 БАТ 2169 №1 від 01.07.2026",
        "БР 0 БАТ 2169 №2 від 02.07.2026",
        "БН 0 БАТ 2169 №1 від 02.07.2026",
    ]


def test_all_reference_lines_empty_for_empty_input():
    assert rc._all_reference_lines({}) == []


# -------------------------
# _dates_from_date_keyed_dict / _every_day_brg_lines - на СИНТЕТИЧНИХ словниках
# форми NUMBER_OF_DOCUMENTS_BRS_EVERY_WEEK/_EVERY_DAY, незалежно від їхнього
# ПОТОЧНОГО (щомісяця змінюваного) реального вмісту.
# -------------------------
def test_dates_from_date_keyed_dict_parses_valid_keys():
    result = rc._dates_from_date_keyed_dict({"01.07.2026": {}, "02.07.2026": {}})
    assert sorted(result) == [datetime(2026, 7, 1), datetime(2026, 7, 2)]


def test_dates_from_date_keyed_dict_skips_unparseable_keys():
    result = rc._dates_from_date_keyed_dict({"не дата": {}, "01.07.2026": {}})
    assert result == [datetime(2026, 7, 1)]


def test_dates_from_date_keyed_dict_empty_for_empty_input():
    assert rc._dates_from_date_keyed_dict({}) == []


def test_every_day_brg_lines_builds_line_per_nonblank_brg_entry(monkeypatch):
    monkeypatch.setattr(rc, "NUMBER_OF_DOCUMENTS_BRS_EVERY_DAY", {
        "01.07.2026": {"бат": "88", "брг": "2169"},
        "02.07.2026": {"бат": "90", "брг": ""},
        "03.07.2026": {"бат": "92"},
    })
    assert rc._every_day_brg_lines() == {f"БР {rc.SHORT_UNIT_BRIGADE} №2169 від 01.07.2026"}


def test_every_day_brg_lines_empty_when_no_brg_field_present(monkeypatch):
    monkeypatch.setattr(rc, "NUMBER_OF_DOCUMENTS_BRS_EVERY_DAY", {"01.07.2026": {"бат": "88"}})
    assert rc._every_day_brg_lines() == set()


# -------------------------
# _ignored_weekly_br_numbers - на СИНТЕТИЧНОМУ NUMBER_OF_DOCUMENTS_BRS_EVERY_WEEK,
# незалежно від його ПОТОЧНОГО (щомісяця змінюваного) реального вмісту.
# -------------------------
def test_ignored_weekly_br_numbers_collects_bat_and_reference_numbers(monkeypatch):
    """_ignored_weekly_br_numbers() передає дати в build_brs_chain_lines
    (content.money_report_helpers), яка сама читає СВІЙ ВЛАСНИЙ NUMBER_OF_DOCUMENTS_BRS_EVERY_WEEK
    з __globals__ ЦІЄЇ функції - підміна лише rc.NUMBER_OF_DOCUMENTS_BRS_EVERY_WEEK
    (report_checker) не зачіпає цей другий, окремий простір імен (а в
    багатопроєктній pytest-сесії навіть підміна content.money_report_helpers
    напряму як модуля не гарантовано влучає в ТОЙ САМИЙ об'єкт модуля, що його
    вже захопила rc.build_brs_chain_lines - див. кореневий conftest.py,
    _COLLIDING_MODULES/_switch_to). monkeypatch.setitem напряму в
    rc.build_brs_chain_lines.__globals__ - єдиний спосіб гарантовано влучити
    САМЕ туди, звідки ця функція реально читає, незалежно від того, скільки
    разів і де ще міг перезавантажитись сам модуль. Без цього тест "працює"
    лише випадково, поки реальний MONTH/DATA з constants.py збігається з
    датою тут (і ламається щоразу, як користувач переходить у наступний
    місяць) - саме так і сталось (2026-08-20: MONTH перейшов на серпень)."""
    fake_week_dict = {
        "27.07.2026": {"бат": "154", "посилання_брг": "2400 від 13.07.2026", "посилання_бат": "131 від 19.07.2026"},
    }
    monkeypatch.setattr(rc, "NUMBER_OF_DOCUMENTS_BRS_EVERY_WEEK", fake_week_dict)
    monkeypatch.setitem(rc.build_brs_chain_lines.__globals__, "NUMBER_OF_DOCUMENTS_BRS_EVERY_WEEK", fake_week_dict)
    assert rc._ignored_weekly_br_numbers() == {"154", "2400", "131"}


def test_ignored_weekly_br_numbers_empty_for_empty_dict(monkeypatch):
    monkeypatch.setattr(rc, "NUMBER_OF_DOCUMENTS_BRS_EVERY_WEEK", {})
    assert rc._ignored_weekly_br_numbers() == set()


# -------------------------
# _basis_citations - виняток для тижневого БР (NUMBER_OF_DOCUMENTS_BRS_EVERY_WEEK,
# ЦІЛКОМ) і 'брг' щоденного БР (NUMBER_OF_DOCUMENTS_BRS_EVERY_DAY) - підтверджено
# користувачем: жодне з них не відповідає локальному файлу ЩОДЕННОЇ цього проєкту.
# _IGNORED_BASIS_CITATIONS тут - синтетичне значення (через фікстуру
# _isolate_ignored_basis_lines), не реальний, щомісяця змінюваний вміст.
# -------------------------
def test_basis_citations_excludes_weekly_br_number_regardless_of_cited_date(monkeypatch):
    """Регресія на реальний випадок, повідомлений користувачем: тижневий БР
    №154 виданий 27.07.2026 (constants.NUMBER_OF_DOCUMENTS_BRS_EVERY_WEEK -
    'бат': '154' саме на цю дату), але лишається чинним і НАСТУПНІ дні, тож у
    ГОТОВОМУ рапорті цитується з ІНШОЮ датою (датою участі людини, тут -
    28.07.2026, а не датою видання) - раніше (звірка тижневого БР за (тип,
    номер, ДАТА) - _IGNORED_BASIS_CITATIONS) ця розбіжність НЕ ловилась: номер
    збігався, а дата - ні, тож "154" помилково НЕ виключався. Виправлено:
    тижневий БР виключається ЛИШЕ за номером (_IGNORED_WEEKLY_BR_NUMBERS),
    дата взагалі не звіряється."""
    monkeypatch.setattr(rc, "_IGNORED_WEEKLY_BR_NUMBERS", {"154"})
    assert rc._basis_citations("БР 0 БАТ 2169 №154 від 28.07.2026") == []


def test_basis_citations_keeps_br_number_not_in_ignored_weekly_set(monkeypatch):
    monkeypatch.setattr(rc, "_IGNORED_WEEKLY_BR_NUMBERS", {"154"})
    assert rc._basis_citations("БР 0 БАТ 2169 №156 від 28.07.2026") == [("БР", "156", "28.07.2026")]


def test_basis_citations_excludes_daily_brg_citation_but_keeps_daily_bat_citation(monkeypatch):
    """'брг' щоденного БР (NUMBER_OF_DOCUMENTS_BRS_EVERY_DAY) - ігнорується;
    'бат' ЦЬОГО Ж словника - НІ, навіть коли обидва згадані в тому самому тексті."""
    monkeypatch.setattr(rc, "_IGNORED_BASIS_CITATIONS", {("БР", "2217", "01.07.2026")})
    text = "БР 0 БАТ 2169 №88 від 01.07.2026\nБР 2169 №2217 від 01.07.2026"
    assert rc._basis_citations(text) == [("БР", "88", "01.07.2026")]


# -------------------------
# _basis_citations
# -------------------------
def test_basis_citations_extracts_br_line_with_standard_format():
    text = "БР 0 БАТ 2169 №2217 від 03.07.2026"
    assert rc._basis_citations(text) == [("БР", "2217", "03.07.2026")]


def test_basis_citations_extracts_zhbd_line():
    text = "ЖБД 0 БАТ 2169 №359дск/6 від 23.06.2026"
    assert rc._basis_citations(text) == [("ЖБД", "359дск/6", "23.06.2026")]


def test_basis_citations_ignores_bn_and_pozbd_markers():
    text = "БН 0 БАТ 2169 №1 від 08.07.2026\nПозБД 2169 №242 від 18.06.2026"
    assert rc._basis_citations(text) == []


def test_basis_citations_extracts_multiple_lines():
    text = "БР 0 БАТ 2169 №2217 від 03.07.2026\nЖБД 0 БАТ 2169 №359дск/6 від 23.06.2026"
    assert rc._basis_citations(text) == [("БР", "2217", "03.07.2026"), ("ЖБД", "359дск/6", "23.06.2026")]


def test_basis_citations_handles_date_split_onto_next_line_within_one_citation():
    """Реальний запис (constants.MONEY_REPORT_GENERAL_REFERENCES_LOG_WAR) сам
    переносить дату на наступний рядок формулювання - "ЖБД ... №359дск/N" і
    "від {дата}" опиняються на РІЗНИХ рядках готового тексту, хоч це ОДНА
    підстава - зведення пробілів/переносів в один рядок перед аналізом має це
    виправити."""
    text = "ЖБД 0 БАТ 2169 №359дск/6\n від 23.06.2026"
    assert rc._basis_citations(text) == [("ЖБД", "359дск/6", "23.06.2026")]


def test_basis_citations_handles_number_after_date_without_vid_keyword():
    text = "БР 2169 №114 11.07.2026"
    assert rc._basis_citations(text) == [("БР", "114", "11.07.2026")]


def test_basis_citations_handles_dash_and_middot_separators():
    """Той самий "покручений" формат (дефіси/крапки замість пробілів), що й у
    реальних constants.BR_HIGHT_UNIT, але з ІНШИМ номером (не точним записом з
    константи), щоб перевірити САМ розбір, а не потрапити під виняток нижче."""
    text = "БР-9·АК №119/1/2500т/99-від-12.05.2026"
    assert rc._basis_citations(text) == [("БР", "119/1/2500т/99", "12.05.2026")]


def test_basis_citations_excludes_exact_higher_unit_reference_citation(monkeypatch):
    """constants.BR_HIGHT_UNIT - підстави ВИЩОГО штабу (9 АК, командир бригади
    тощо) - НЕ перевіряються взагалі, підтверджено користувачем, навіть попри
    те, що формально мають "БР"/"ЖБД" і дату. Синтетична (тип, номер, дата)
    (не з РЕАЛЬНОГО constants.BR_HIGHT_UNIT - _isolate_ignored_basis_lines) -
    перевіряється сам МЕХАНІЗМ винятку, а не поточний, щомісяця змінюваний
    вміст константи."""
    monkeypatch.setattr(rc, "_IGNORED_BASIS_CITATIONS", {("БР", "119/1/2500т/9", "12.05.2026")})
    assert rc._basis_citations("БР-9·АК №119/1/2500т/9-від-12.05.2026") == []


def test_basis_citations_excludes_higher_unit_reference_among_other_citations(monkeypatch):
    monkeypatch.setattr(rc, "_IGNORED_BASIS_CITATIONS", {("БР", "119/1/2500т/9", "12.05.2026")})
    text = "БР 0 БАТ 2169 №2217 від 03.07.2026\nБР-9·АК №119/1/2500т/9-від-12.05.2026"
    assert rc._basis_citations(text) == [("БР", "2217", "03.07.2026")]


def test_basis_citations_excludes_reformatted_higher_unit_reference_regardless_of_wording(monkeypatch):
    """Регресія на реальний випадок: той самий запис BR_HIGHT_UNIT з'являється в
    ГОТОВОМУ рапорті записаним ЗОВСІМ інакше, ніж він написаний у ПОТОЧНОМУ
    constants.py (рапорт зберігає текст ТАКИМ, яким BR_HIGHT_UNIT був на момент
    генерації рапорту, а сам BR_HIGHT_UNIT редагується/переформатовується далі) -
    напр. рапорт: "БР 9 АК №119/1/2500т/9 від 12.05.2026;" (пробіли, крапка з
    комою, номер БЕЗ ведучого дефіса), а BR_HIGHT_UNIT СЬОГОДНІ:
    "БР-9·АК №119/1/2500т/9-від-12.05.2026" (дефіси/крапки, номер З ведучим
    дефісом "№-..." - хоча тут дефіс належить самому запису вище рангом, тест
    моделює саме "номер без дефіса в рапорті" =="номер з дефісом в constants.py"
    - обидва мають нормалізуватись до ОДНАКОВОЇ (тип, номер, дата)). Раніше
    (дослівне порівняння рядків) це НЕ співпадало й хибно НЕ виключалось."""
    monkeypatch.setattr(rc, "_IGNORED_BASIS_CITATIONS", rc._reference_citations([
        "БР·9 бр ·№-1622 від-16.05.2026",
    ]))
    text = "БР 9 бр № 1622від 16.05.2026;"
    assert rc._basis_citations(text) == []


def test_basis_citations_does_not_leak_date_across_adjacent_citations():
    """"БР ... 11.07.2026" (без "від") одразу за ним "БР ... №2400 13.07.2026" в
    ОДНОМУ тексті (без переносу рядка) - МАРКЕР другої підстави має зупинити
    пошук дати для першої, інакше перша хибно підхопила б 13.07.2026."""
    text = "БР 2169 №114 11.07.2026 БР 2169 №2400 13.07.2026"
    assert rc._basis_citations(text) == [("БР", "114", "11.07.2026"), ("БР", "2400", "13.07.2026")]


def test_basis_citations_returns_empty_for_reference_without_date():
    assert rc._basis_citations("ЖБД (Справа №2т)") == []


def test_basis_citations_returns_empty_for_blank_text():
    assert rc._basis_citations("") == []
    assert rc._basis_citations(None) == []


def test_basis_citations_number_is_none_when_no_number_present():
    text = "БР 2169 від 03.07.2026"
    assert rc._basis_citations(text) == [("БР", None, "03.07.2026")]


# -------------------------
# _br_files_for_date
# -------------------------
def test_br_files_for_date_matches_daily_file(tmp_path):
    (tmp_path / "№117 ЩОДЕННА 13.07.2026.docx").write_bytes(b"")
    assert rc._br_files_for_date(str(tmp_path), "13.07.2026") == [str(tmp_path / "№117 ЩОДЕННА 13.07.2026.docx")]


def test_br_files_for_date_matches_task_order_pdf(tmp_path):
    (tmp_path / "№87 ЗАВДАННЯ 01.07.2026.pdf").write_bytes(b"")
    assert rc._br_files_for_date(str(tmp_path), "01.07.2026") == [str(tmp_path / "№87 ЗАВДАННЯ 01.07.2026.pdf")]


def test_br_files_for_date_recurses_into_subfolders(tmp_path):
    (tmp_path / "br").mkdir()
    (tmp_path / "br" / "№117 ЩОДЕННА 13.07.2026.docx").write_bytes(b"")
    assert rc._br_files_for_date(str(tmp_path), "13.07.2026") == [str(tmp_path / "br" / "№117 ЩОДЕННА 13.07.2026.docx")]


def test_br_files_for_date_ignores_wrong_date(tmp_path):
    (tmp_path / "№117 ЩОДЕННА 14.07.2026.docx").write_bytes(b"")
    assert rc._br_files_for_date(str(tmp_path), "13.07.2026") == []


def test_br_files_for_date_empty_for_folder_without_matches(tmp_path):
    assert rc._br_files_for_date(str(tmp_path), "13.07.2026") == []


# -------------------------
# _log_war_files_for_number - за НОМЕРОМ справи в заголовку файлу (а не за
# датою з filename, яка не пов'язана з датою "від DATE" самої підстави рапорту).
# -------------------------
def _write_docx_with_text(path, text):
    from docx import Document
    doc = Document()
    doc.add_paragraph(text)
    doc.save(str(path))
    return str(path)


def test_log_war_files_for_number_matches_by_content_not_filename_date(tmp_path):
    """Реальний сценарій: filename датований 16.07.2026, а заголовок усередині
    файлу - "за номенклатурою №359дск/7 від 12.07.2026" (ІНША дата, спільна для
    ВСЬОГО періоду дії цього запису) - пошук за номером справи в змісті файлу
    все одно знаходить цей файл, повертаючи його ВЛАСНУ дату з filename."""
    path = _write_docx_with_text(
        tmp_path / "Витяг з ЖБД для БР - №124 16.07.2026 ЩОДЕННА.docx",
        "... за номенклатурою №359дск/7 від 12.07.2026 року ...",
    )
    result = rc._log_war_files_for_number(str(tmp_path), "359дск/7")
    assert result == [(path, datetime(2026, 7, 16))]


def test_log_war_files_for_number_matches_weekly_suffix_too(tmp_path):
    path = _write_docx_with_text(
        tmp_path / "Витяг з ЖБД для БР - №121 14.07.2026 ЩОТИЖНЕВА.docx",
        "за номенклатурою №359дск/7 від 12.07.2026",
    )
    assert rc._log_war_files_for_number(str(tmp_path), "359дск/7") == [(path, datetime(2026, 7, 14))]


def test_log_war_files_for_number_ignores_files_with_different_number(tmp_path):
    _write_docx_with_text(tmp_path / "Витяг з ЖБД для БР - №124 16.07.2026 ЩОДЕННА.docx", "за номенклатурою №359дск/6 від 23.06.2026")
    assert rc._log_war_files_for_number(str(tmp_path), "359дск/7") == []


def test_log_war_files_for_number_recurses_into_subfolders(tmp_path):
    (tmp_path / "logwar").mkdir()
    path = _write_docx_with_text(tmp_path / "logwar" / "Витяг з ЖБД для БР - №124 16.07.2026 ЩОДЕННА.docx", "№359дск/7 від 12.07.2026")
    assert rc._log_war_files_for_number(str(tmp_path), "359дск/7") == [(path, datetime(2026, 7, 16))]


def test_log_war_files_for_number_empty_when_number_is_none():
    assert rc._log_war_files_for_number("будь-яка-неіснуюча-папка", None) == []


def test_log_war_files_for_number_ignores_non_matching_filenames(tmp_path):
    _write_docx_with_text(tmp_path / "№117 ЩОДЕННА 13.07.2026.docx", "№359дск/7 від 12.07.2026")
    assert rc._log_war_files_for_number(str(tmp_path), "359дск/7") == []


# -------------------------
# _check_basis_documents
# -------------------------
def test_check_basis_documents_no_finding_when_document_found_and_person_named(tmp_path):
    _write_docx_with_text(tmp_path / "№117 ЩОДЕННА 13.07.2026.docx", "Особовий склад: СЬОМИЙ Сьомий Сьомий")
    entries = [{"ПІБ": "СЬОМИЙ Сьомий Сьомий", "ПІДСТАВА": "БР 0 БАТ 2169 №117 від 13.07.2026"}]

    assert rc._check_basis_documents(entries, str(tmp_path)) == []


def test_check_basis_documents_flags_missing_document(tmp_path):
    entries = [{"ПІБ": "СЬОМИЙ Сьомий Сьомий", "ПІДСТАВА": "БР 0 БАТ 2169 №117 від 13.07.2026"}]

    findings = rc._check_basis_documents(entries, str(tmp_path))

    assert len(findings) == 1
    assert findings[0]["ПІБ"] == "СЬОМИЙ Сьомий Сьомий"
    assert "не знайдено" in findings[0]["Повідомлення"].lower()
    assert findings[0]["Шлях"] == str(tmp_path)
    # Регресія: користувач попросив дописати номер БР ("Ресурс" - окрема колонка).
    assert findings[0]["Ресурс"] == "БР №117 від 13.07.2026"
    # Регресія: користувач попросив дописати ЗА ЯКЕ ЧИСЛО ПЕРІОДУ й У ЯКОМУ
    # ДОКУМЕНТІ саме не підтвердилось - для "БР" це рівно дата самої підстави;
    # "Документ" порожній, бо жодного файлу за цю дату взагалі не знайдено.
    assert findings[0]["Дата періоду"] == "13.07.2026"
    assert findings[0]["Документ"] == ""


def test_check_basis_documents_flags_document_found_but_person_not_named(tmp_path):
    _write_docx_with_text(tmp_path / "№117 ЩОДЕННА 13.07.2026.docx", "Особовий склад: ВОСЬМИЙ Восьмий Восьмий")
    entries = [{"ПІБ": "СЬОМИЙ Сьомий Сьомий", "ПІДСТАВА": "БР 0 БАТ 2169 №117 від 13.07.2026"}]

    findings = rc._check_basis_documents(entries, str(tmp_path))

    assert len(findings) == 1
    assert findings[0]["ПІБ"] == "СЬОМИЙ Сьомий Сьомий"
    assert "не згадано" in findings[0]["Повідомлення"]
    assert findings[0]["Ресурс"] == "БР №117 від 13.07.2026"
    assert findings[0]["Дата періоду"] == "13.07.2026"
    assert findings[0]["Документ"] == "№117 ЩОДЕННА 13.07.2026.docx"


def test_check_basis_documents_label_omits_number_when_not_recognized(tmp_path):
    """Підстава без розпізнаваного номера (напр. "№" узагалі відсутнє) - у
    "Ресурс" лишається лише тип+дата, без "№None" чи подібного сміття."""
    entries = [{"ПІБ": "СЬОМИЙ Сьомий Сьомий", "ПІДСТАВА": "БР 0 БАТ 2169 від 13.07.2026"}]

    findings = rc._check_basis_documents(entries, str(tmp_path))

    assert len(findings) == 1
    assert "№" not in findings[0]["Ресурс"]
    assert findings[0]["Ресурс"] == "БР від 13.07.2026"
    assert findings[0]["Дата періоду"] == "13.07.2026"


def test_check_basis_documents_skips_entries_without_br_or_zhbd_citation():
    entries = [{"ПІБ": "СЬОМИЙ Сьомий Сьомий", "ПІДСТАВА": "Особиста примітка без дати"}]
    assert rc._check_basis_documents(entries, "будь-яка-неіснуюча-папка") == []


def test_check_basis_documents_deduplicates_same_person_type_date_across_entries(tmp_path):
    entries = [
        {"ПІБ": "СЬОМИЙ Сьомий Сьомий", "ПІДСТАВА": "БР 0 БАТ 2169 №117 від 13.07.2026"},
        {"ПІБ": "СЬОМИЙ Сьомий Сьомий", "ПІДСТАВА": "БР 0 БАТ 2169 №117 від 13.07.2026"},
    ]

    findings = rc._check_basis_documents(entries, str(tmp_path))

    assert len(findings) == 1


def test_check_basis_documents_passes_when_any_candidate_file_names_person(tmp_path):
    """Кілька файлів відповідають ТИПУ+ДАТІ (напр. ЩОДЕННА і застосування безпеки
    того самого дня) - досить, щоб ПІБ згадувався ХОЧ В ОДНОМУ з них."""
    _write_docx_with_text(tmp_path / "№117 ЩОДЕННА 13.07.2026.docx", "ВОСЬМИЙ Восьмий Восьмий")
    _write_docx_with_text(tmp_path / "№118 ЗАВДАННЯ 13.07.2026.docx", "СЬОМИЙ Сьомий Сьомий")
    entries = [{"ПІБ": "СЬОМИЙ Сьомий Сьомий", "ПІДСТАВА": "БР 0 БАТ 2169 №117 від 13.07.2026"}]

    assert rc._check_basis_documents(entries, str(tmp_path)) == []


def test_check_basis_documents_zhbd_citation_never_matches_daily_br_files(tmp_path):
    """"ЖБД" шукається за НОМЕРОМ справи у ЗМІСТІ файлу (_log_war_files_for_number),
    а не за типом/датою filename - файл ЩОДЕННА, що просто МІСТИТЬ ПІБ, але не
    сам номер справи, не рахується. Жодного файлу з цим номером узагалі -
    ОДНЕ загальне "не знайдено" (без "Дата періоду" - немає конкретної дати,
    яку можна було б звинуватити, підтверджено користувачем)."""
    _write_docx_with_text(tmp_path / "№117 ЩОДЕННА 13.07.2026.docx", "СЬОМИЙ Сьомий Сьомий")
    entries = [{"ПІБ": "СЬОМИЙ Сьомий Сьомий", "ПІДСТАВА": "ЖБД 0 БАТ 2169 №359дск/6 від 13.07.2026", "ПЕРІОД": "13.07.2026-13.07.2026"}]

    findings = rc._check_basis_documents(entries, str(tmp_path))

    assert len(findings) == 1
    assert "не знайдено" in findings[0]["Повідомлення"].lower()
    assert findings[0]["Дата періоду"] == ""
    assert findings[0]["Документ"] == ""


def test_check_basis_documents_zhbd_finds_person_via_case_number_despite_mismatched_citation_date(tmp_path):
    """Регресія на реальний випадок, повідомлений користувачем: підстава
    рапорту каже "ЖБД ... №359дск/7 від 12.07.2026" (дата РЕЄСТРАЦІЇ запису,
    СПІЛЬНА для багатьох витягів), а сам файл витяга - за 16.07.2026 (ДЕНЬ
    участі людини, з її "ПЕРІОД") - заголовок файлу теж каже "...№359дск/7 від
    12.07.2026" (та сама дата реєстрації, а не 16.07.2026). Раніше (пошук за
    датою з filename == дата з підстави) це НІКОЛИ не знаходилось - тепер
    шукається за НОМЕРОМ справи в змісті файлу і ВЛАСНОЮ датою файлу серед
    ПЕРІОДУ людини."""
    path = _write_docx_with_text(
        tmp_path / "Витяг з ЖБД для БР - №124 16.07.2026 ЩОДЕННА.docx",
        "за номенклатурою №359дск/7 від 12.07.2026 року\nОсобовий склад: СЬОМИЙ Сьомий Сьомий",
    )
    entries = [{"ПІБ": "СЬОМИЙ Сьомий Сьомий", "ПІДСТАВА": "ЖБД 0 БАТ 2169 №359дск/7 від 12.07.2026", "ПЕРІОД": "16.07.2026-16.07.2026"}]

    assert rc._check_basis_documents(entries, str(tmp_path)) == []
    assert "СЬОМИЙ" in rc._cached_document_text(path)


def test_check_basis_documents_zhbd_flags_missing_when_no_file_covers_persons_period(tmp_path):
    """Файл з ПРАВИЛЬНИМ номером справи існує, але його ВЛАСНА дата (17.07.2026)
    - поза межами періоду, який людина заявляє (16.07.2026) - не рахується.
    Немає ЖОДНОГО файлу цього номера серед періоду людини - ОДНЕ загальне
    "не знайдено" (без конкретної "Дата періоду" - див. попередній тест)."""
    _write_docx_with_text(
        tmp_path / "Витяг з ЖБД для БР - №127 17.07.2026 ЩОДЕННА.docx",
        "за номенклатурою №359дск/7 від 12.07.2026\nСЬОМИЙ Сьомий Сьомий",
    )
    entries = [{"ПІБ": "СЬОМИЙ Сьомий Сьомий", "ПІДСТАВА": "ЖБД 0 БАТ 2169 №359дск/7 від 12.07.2026", "ПЕРІОД": "16.07.2026-16.07.2026"}]

    findings = rc._check_basis_documents(entries, str(tmp_path))

    assert len(findings) == 1
    assert "не знайдено" in findings[0]["Повідомлення"].lower()
    assert findings[0]["Дата періоду"] == ""


def test_check_basis_documents_zhbd_flags_only_the_specific_dates_missing_confirmation(tmp_path):
    """Регресія на прохання користувача: людина заявляє ЖБД за ДВА дні періоду
    (16.07 і 17.07) - обидва мають витяг із ПРАВИЛЬНИМ номером справи, але ПІБ
    згадано лише в ОДНОМУ з них (16.07) - результат має вказати РІВНО на
    17.07.2026 (і саме той файл), а не мовчати (бо "хоч десь підтвердилось") і
    не помилково звинувачувати 16.07 (де все гаразд)."""
    confirmed_path = _write_docx_with_text(
        tmp_path / "Витяг з ЖБД для БР - №124 16.07.2026 ЩОДЕННА.docx",
        "за номенклатурою №359дск/7 від 12.07.2026\nСЬОМИЙ Сьомий Сьомий",
    )
    missing_path = _write_docx_with_text(
        tmp_path / "Витяг з ЖБД для БР - №127 17.07.2026 ЩОДЕННА.docx",
        "за номенклатурою №359дск/7 від 12.07.2026\nВОСЬМИЙ Восьмий Восьмий",
    )
    entries = [{"ПІБ": "СЬОМИЙ Сьомий Сьомий", "ПІДСТАВА": "ЖБД 0 БАТ 2169 №359дск/7 від 12.07.2026", "ПЕРІОД": "16.07.2026-17.07.2026"}]

    findings = rc._check_basis_documents(entries, str(tmp_path))

    assert len(findings) == 1
    assert findings[0]["Дата періоду"] == "17.07.2026"
    assert findings[0]["Документ"] == os.path.basename(missing_path)
    assert "не згадано" in findings[0]["Повідомлення"]
    assert "16.07.2026" not in findings[0]["Повідомлення"]
    assert os.path.basename(confirmed_path) != findings[0]["Документ"]


def test_check_basis_documents_zhbd_only_checks_dates_that_have_a_matching_numbered_file(tmp_path):
    """Один із двох днів періоду (18.07) НЕ має ЖОДНОГО файлу з цим номером
    справи в усьому folder (хоча номер справи реально трапляється - лише за
    ІНШИЙ день, 16.07) - цей день НЕ перевіряється взагалі під ЦІЄЮ підставою:
    у самого рапорту могла бути ЩЕ ОДНА, ІНША ЖБД-підстава саме на 18.07 (з
    іншим номером справи), яку ЦЯ перевірка вже не бачить (у фікстурі її просто
    нема) - тож 18.07 мовчки пропускається, а не хибно звинувачується під ЦИМ
    номером. 16.07 - має файл і в ньому справді згадано ПІБ - findings порожні."""
    _write_docx_with_text(
        tmp_path / "Витяг з ЖБД для БР - №124 16.07.2026 ЩОДЕННА.docx",
        "за номенклатурою №359дск/7 від 12.07.2026\nСЬОМИЙ Сьомий Сьомий",
    )
    entries = [{"ПІБ": "СЬОМИЙ Сьомий Сьомий", "ПІДСТАВА": "ЖБД 0 БАТ 2169 №359дск/7 від 12.07.2026", "ПЕРІОД": "16.07.2026-16.07.2026;18.07.2026-18.07.2026"}]

    findings = rc._check_basis_documents(entries, str(tmp_path))

    assert findings == []


def test_check_basis_documents_zhbd_flags_person_not_named_in_covering_file(tmp_path):
    _write_docx_with_text(
        tmp_path / "Витяг з ЖБД для БР - №124 16.07.2026 ЩОДЕННА.docx",
        "за номенклатурою №359дск/7 від 12.07.2026\nВОСЬМИЙ Восьмий Восьмий",
    )
    entries = [{"ПІБ": "СЬОМИЙ Сьомий Сьомий", "ПІДСТАВА": "ЖБД 0 БАТ 2169 №359дск/7 від 12.07.2026", "ПЕРІОД": "16.07.2026-16.07.2026"}]

    findings = rc._check_basis_documents(entries, str(tmp_path))

    assert len(findings) == 1
    assert "не згадано" in findings[0]["Повідомлення"]
    assert findings[0]["Дата періоду"] == "16.07.2026"
    assert findings[0]["Документ"] == "Витяг з ЖБД для БР - №124 16.07.2026 ЩОДЕННА.docx"


# -------------------------
# check_report_against_oblik - basis_folder
# -------------------------
def test_check_report_against_oblik_skips_basis_check_when_folder_not_given(tmp_path, monkeypatch):
    oblik_path = _write_oblik_xlsx(tmp_path / "ОБЛІК.xlsx", ["ПІДРОЗДІЛ", "ПОСАДА", "ЗВАННЯ", "ПІБ", D1], [
        ["Підрозділ 1", "Стрілець", "сержант", "СЬОМИЙ Сьомий Сьомий", 30],
    ])
    _use_fixture(monkeypatch, oblik_path, ["A", "B", "C", "D", "E"])
    report_path = _write_report(tmp_path / "report.docx", [
        ("1. Виплатити додаткову винагороду у розмірі 30 000 грн. 00 коп. військовослужбовцям", [
            ["Стрілець", "сержант", "СЬОМИЙ Сьомий Сьомий", "01.07.2026-01.07.2026", "1", "БР 0 БАТ 2169 №117 від 13.07.2026"],
        ]),
    ])

    output_path = rc.check_report_against_oblik(report_path, output_path=str(tmp_path / "result.xlsx"))

    assert _read_result_rows(output_path) == ["Розбіжностей між рапортом і ОБЛІК.xlsx не знайдено."]


def test_check_report_against_oblik_includes_basis_findings_when_folder_given(tmp_path, monkeypatch):
    oblik_path = _write_oblik_xlsx(tmp_path / "ОБЛІК.xlsx", ["ПІДРОЗДІЛ", "ПОСАДА", "ЗВАННЯ", "ПІБ", D1], [
        ["Підрозділ 1", "Стрілець", "сержант", "СЬОМИЙ Сьомий Сьомий", 30],
    ])
    _use_fixture(monkeypatch, oblik_path, ["A", "B", "C", "D", "E"])
    report_path = _write_report(tmp_path / "report.docx", [
        ("1. Виплатити додаткову винагороду у розмірі 30 000 грн. 00 коп. військовослужбовцям", [
            ["Стрілець", "сержант", "СЬОМИЙ Сьомий Сьомий", "01.07.2026-01.07.2026", "1", "БР 0 БАТ 2169 №117 від 13.07.2026"],
        ]),
    ])
    empty_basis_folder = tmp_path / "output"
    empty_basis_folder.mkdir()

    output_path = rc.check_report_against_oblik(
        report_path, output_path=str(tmp_path / "result.xlsx"), basis_folder=str(empty_basis_folder),
    )

    content = _read_result_rows(output_path)
    assert any("не знайдено" in row.lower() and "БР №117 від 13.07.2026" in row for row in content)
    assert any("СЬОМИЙ Сьомий Сьомий" in row for row in content)
