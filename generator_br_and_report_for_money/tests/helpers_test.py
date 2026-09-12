import random
import re
import sys
from datetime import datetime, timedelta
from pathlib import Path

sys.path.append(str(Path(__file__).resolve().parent.parent))

import pytest
from colorama import Fore, Style
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, RGBColor

import constants
import helpers

# -------------------------
# Фікстури
# -------------------------
@pytest.fixture(autouse=True)
def fixed_random_seed():
    random.seed(0)
    yield
    random.seed()


@pytest.fixture
def doc():
    return Document()


# -------------------------
# Прості утилітарні функції
# -------------------------
def test_print_colored_caps(capsys):
    helpers.print_red("err")
    helpers.print_green("ok")
    out = capsys.readouterr().out
    assert "err" in out
    assert "ok" in out


def test_print_red_output(capsys):
    helpers.print_red("Hello")
    out = capsys.readouterr().out
    assert out == Fore.RED + "Hello" + Style.RESET_ALL + "\n"


def test_get_data_with_dates_frame():
    import pandas as pd

    df = pd.DataFrame({'A': [1], 'B': [pd.Timestamp('2025-11-01')], 'C': ['text']})
    rows = helpers.get_data(df, ['A', 'B', 'C'])
    assert rows['rows'][0]['B'] == '01.11.2025'
    assert rows['rows'][0]['A'] == 1
    assert rows['rows'][0]['C'] == 'text'


def test_get_data_date_can_be_empty_true_with_nat():
    import pandas as pd

    df = pd.DataFrame([{"A": pd.NaT, "B": "x"}])
    result = helpers.get_data(df, ["A", "B"], date_can_be_empty=True)
    assert result["rows"] == [{"A": None, "B": "x"}]


def test_get_data_multiple_rows_mixed():
    import pandas as pd

    df = pd.DataFrame([
        {"A": pd.Timestamp(2025, 2, 2), "B": None},
        {"A": None, "B": "text"},
        {"A": 50, "B": pd.Timestamp(2020, 10, 10)},
    ])
    result = helpers.get_data(df, ["A", "B"], date_can_be_empty=True)
    assert result["rows"] == [
        {"A": "02.02.2025", "B": None},
        {"A": None, "B": "text"},
        {"A": 50, "B": "10.10.2020"},
    ]


def test_excel_col_to_index_basic():
    assert helpers.excel_col_to_index('A') == 0
    assert helpers.excel_col_to_index('Z') == 25
    assert helpers.excel_col_to_index('AA') == 26
    assert helpers.excel_col_to_index('AD') == 29


def test_to_date_parsing_and_errors():
    assert helpers.to_date(datetime(2025, 11, 1)) == datetime(2025, 11, 1).date()
    assert helpers.to_date("2025-11-01") == datetime(2025, 11, 1).date()
    assert helpers.to_date("01.11.2025") == datetime(2025, 11, 1).date()
    with pytest.raises(ValueError):
        helpers.to_date(12345)


def test_to_date_strips_stray_whitespace():
    """constants.py (MONEY_REPORT_GENERAL_REFERENCES_LOG_WAR тощо) редагується вручну -
    випадковий пробіл на кінці дати ("08.08.2026 ") - реальна помилка, знайдена
    користувачем (падало з "Невідомий формат або тип дати"), а не привід відхилити
    дату як "невідомий формат"."""
    assert helpers.to_date(" 01.11.2025 ") == datetime(2025, 11, 1).date()
    assert helpers.to_date("01.11.2025 ") == datetime(2025, 11, 1).date()
    assert helpers.to_date(" 2025-11-01") == datetime(2025, 11, 1).date()


def test_convert_to_short_name_basic_and_errors():
    assert helpers.convert_to_short_name("Перший Другий Третій") == "Другий Перший"
    assert isinstance(helpers.convert_to_short_name("В' осьмий  Четвертий  Шостий"), str)
    assert helpers.convert_to_short_name(None) is None

    with pytest.raises(ValueError):
        helpers.convert_to_short_name("OnlyTwoWords")

    assert helpers.convert_to_short_name("Сьомий ВосьмийДесятий") == "Восьмий Сьомий"
    assert helpers.convert_to_short_name("Одинадцятий ДванадцятийТринадцятий") == "Дванадцятий Одинадцятий"
    assert helpers.convert_to_short_name("Чотирнадцятий ШістнадцятийСімнадцятий  ") == "Шістнадцятий Чотирнадцятий"


def test_check_number_file_is_exist():
    assert helpers.check_number_file_is_exist(None) == ""
    assert helpers.check_number_file_is_exist("123") == "№123 "


def _expected_log_war_reference_text(docx_utils_module, reference):
    match = docx_utils_module._LOG_WAR_REFERENCE_RE.search(reference["lines"][0])
    return f"№{match.group(1)} від {docx_utils_module._normalize_log_war_reference_date(match.group(2))}"


def test_resolve_log_war_reference_text_picks_entry_matching_date():
    """Дата на самому початку діапазону ПЕРШОГО реального запису
    constants.MONEY_REPORT_GENERAL_REFERENCES_LOG_WAR резолвиться саме в номер
    справи й дату з ЦЬОГО запису - а не з якогось іншого чи фіксованого значення.
    Очікуваний текст береться напряму з самого constants.py (не дублюється
    літералом), щоб тест не застарівав щоразу, як користувач редагує ці записи
    (нова точка ЖБД щомісяця - звичайна практика)."""
    import formatting.docx_utils as docx_utils_module

    reference = constants.MONEY_REPORT_GENERAL_REFERENCES_LOG_WAR[0]
    expected = _expected_log_war_reference_text(docx_utils_module, reference)

    result = docx_utils_module._resolve_log_war_reference_text(reference["start"])
    assert result == expected


def test_resolve_log_war_reference_text_picks_different_entry_for_later_date():
    """Аналогічно, але бере ДРУГИЙ реальний запис - підтверджує, що функція справді
    вибирає ВІДПОВІДНИЙ запис за датою (а не завжди перший чи останній)."""
    import formatting.docx_utils as docx_utils_module

    if len(constants.MONEY_REPORT_GENERAL_REFERENCES_LOG_WAR) < 2:
        pytest.skip("Потрібно щонайменше 2 реальні записи MONEY_REPORT_GENERAL_REFERENCES_LOG_WAR")

    reference = constants.MONEY_REPORT_GENERAL_REFERENCES_LOG_WAR[1]
    expected = _expected_log_war_reference_text(docx_utils_module, reference)

    result = docx_utils_module._resolve_log_war_reference_text(reference["start"])
    assert result == expected


def test_log_war_reference_recognizes_word_based_genitive_month_date():
    """"lines" деяких реальних записів написані природньою мовою ("26 липня
    2026", родовий відмінок місяця), а не лише цифрами ("26.07.2026") - обидва
    формати мають розпізнаватись однаково, і результат завжди приводиться до
    "DD.MM.YYYY" - падало на реальних даних до цього фікса (AttributeError:
    'NoneType' object has no attribute 'group' - регекс розумів лише цифри)."""
    import formatting.docx_utils as docx_utils_module

    cases = [
        ("№359дск/11 від 3 січня 2027", "№359дск/11 від 03.01.2027"),
        ("№359дск/12 від 23 грудня 2026", "№359дск/12 від 23.12.2026"),
    ]
    for line, expected in cases:
        match = docx_utils_module._LOG_WAR_REFERENCE_RE.search(line)
        assert match is not None, f"регекс не розпізнав рядок: {line!r}"
        result = f"№{match.group(1)} від {docx_utils_module._normalize_log_war_reference_date(match.group(2))}"
        assert result == expected


def test_log_war_reference_recognizes_number_with_space_after_hash():
    """"№ 359дск/8" (з пробілом) розпізнається так само, як "№359дск/7" (без) -
    обидва написання трапляються в реальних записах."""
    import formatting.docx_utils as docx_utils_module

    match = docx_utils_module._LOG_WAR_REFERENCE_RE.search("ЖБД № 359дск/8 від 26.07.2026")
    assert match is not None
    assert match.group(1) == "359дск/8"


def test_resolve_log_war_reference_text_falls_back_when_no_entry_matches(monkeypatch, capsys):
    """Якщо для дати не знайшлось жодного запису (напр. увесь список порожній) -
    повертається заповнювач, а не падіння, з попередженням у консоль."""
    import formatting.docx_utils as docx_utils_module

    monkeypatch.setattr(docx_utils_module, "MONEY_REPORT_GENERAL_REFERENCES_LOG_WAR", [])

    result = docx_utils_module._resolve_log_war_reference_text("01.07.2026")

    assert result == f"№_____ від ___.___.{constants.YEAR}"
    assert "Не знайдено запису MONEY_REPORT_GENERAL_REFERENCES_LOG_WAR" in capsys.readouterr().out


def test_generate_combat_log_extract_title_embeds_resolved_reference_text(doc):
    import formatting.docx_utils as docx_utils_module

    reference = constants.MONEY_REPORT_GENERAL_REFERENCES_LOG_WAR[0]
    expected_reference_text = _expected_log_war_reference_text(docx_utils_module, reference)

    docx_utils_module.generate_combat_log_extract_title(doc, reference["start"])

    texts = [p.text for p in doc.paragraphs]
    assert "ВИТЯГ ІЗ ЖУРНАЛУ БОЙОВИХ ДІЙ" in texts
    assert any(f"за номенклатурою {expected_reference_text} року" in text for text in texts)


def test_show_coordinates_masking():
    coords = "36T TT 12345 67890"
    assert "***" in helpers.show_coordinates(True, coords)
    assert helpers.show_coordinates(False, coords) == coords


def test_get_number_br_using_real_constants():
    key = next(iter(constants.NUMBER_OF_DOCUMENTS_BRS_EVERY_WEEK.keys()))
    res = helpers.get_number_br(key)
    assert 'number_documents' in res

    res2 = helpers.get_number_br("not a date")
    assert res2['брг'] == '___' and res2['бат'] == '___'


# -------------------------
# Робота з docx.Document() (реальний, без Dummy-мокапів)
# -------------------------
def test_set_line_spacing_applies_to_all(doc):
    doc.add_paragraph("перший")
    doc.add_paragraph("другий")
    helpers.set_line_spacing(doc, 12)
    for p in doc.paragraphs:
        assert p.paragraph_format.line_spacing.pt == 12


def test_create_or_clear_output_directory(tmp_path):
    base = tmp_path / "out"
    br = base / "br"
    save = base / "br_save"
    extract = base / "extractbr"
    logwar = tmp_path / "logwar"

    for d, f in [(base, "base_file.txt"), (br, "br_file.txt"), (save, "save_file.txt"),
                 (extract, "extract_file.txt"), (logwar, "log_file.txt")]:
        d.mkdir(parents=True, exist_ok=True)
        (d / f).write_text("old data")

    for d in [base, br, save, extract, logwar]:
        assert any(d.rglob("*"))

    helpers.create_or_clear_output_directory(str(base), str(br), str(save), str(extract), str(logwar))

    for d in [base, br, save, extract, logwar]:
        assert d.exists() and d.is_dir()
        for f in d.iterdir():
            if f.is_file():
                pytest.fail(f"Файл {f} не був видалений")


def test_add_paragraph_with_style(doc):
    paragraph = helpers.add_paragraph_with_style(
        doc, "Test text", font_name="Arial", font_size=16, bold=True,
        alignment=WD_ALIGN_PARAGRAPH.CENTER, first_line_indent=20, format_tabs=True
    )

    run = paragraph.runs[0]
    assert run.text == "Test text"
    assert run.font.name == "Arial"
    assert run._element.rPr.rFonts.get(qn('w:eastAsia')) == "Arial"
    assert run.font.size.pt == 16
    assert run.font.bold is True
    assert run.font.color.rgb == RGBColor(0, 0, 0)
    assert paragraph.alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert paragraph.paragraph_format.first_line_indent.pt == 20
    assert paragraph.paragraph_format.line_spacing == 1.0
    assert paragraph.paragraph_format.space_after.pt == 0
    assert paragraph.paragraph_format.space_before.pt == 0

    tab_stops = paragraph.paragraph_format.tab_stops
    sec = doc.sections[0]
    width = Inches(sec.page_width.inches - (sec.left_margin.inches + sec.right_margin.inches))
    tabs = [t.position.inches for t in tab_stops]
    assert any(abs(t - width.inches) < 0.01 for t in tabs)


def test_add_paragraph_with_style_bold_word_list(doc):
    paragraph = helpers.add_paragraph_with_style(doc, "Підрозділ 1 забезпечити готовність резервів", bold=["Підрозділ 1"])
    bold_runs = [r for r in paragraph.runs if r.font.bold]
    assert any(r.text == "Підрозділ 1" for r in bold_runs)


def test_add_paragraph_with_style_heading_level_overrides_style_keep_with_next(doc):
    """Стиль "Heading N" за замовчуванням має keepNext=True (успадковано від базового
    шаблону Word) - без явного False цей абзац "приклеювався" б до наступного вмісту
    (в рапорті - до таблиці), і весь ланцюжок "заголовок+шапка+перший рядок" міг би
    цілком переноситись на нову сторінку, лишаючи порожнє місце на попередній."""
    paragraph = helpers.add_paragraph_with_style(doc, "1. Пункт", heading_level=1)
    assert paragraph.style.name == "Heading 1"
    assert paragraph.style.paragraph_format.keep_with_next is True  # сам стиль - як і був
    assert paragraph.paragraph_format.keep_with_next is False  # але абзац - явно перевизначено


def test_add_custom_heading(doc):
    heading = helpers.add_custom_heading(doc, "head", heading=2, font_size=13)
    assert heading.style.name == "Heading 2"
    assert any(run.text == "head" for run in heading.runs)


def test_set_margins(doc):
    helpers.set_margins(doc, Inches(1), Inches(2), Inches(3), Inches(4))
    for s in doc.sections:
        assert s.top_margin == Inches(1)
        assert s.bottom_margin == Inches(2)
        assert s.left_margin == Inches(3)
        assert s.right_margin == Inches(4)


def test_add_text_for_header_footer(doc):
    helpers.add_text_for_header_and_footer_document_br(doc)
    sec = doc.sections[0]
    assert sec.different_first_page_header_footer is True
    assert sec.header.paragraphs[0].text.strip() == "ДЛЯ СЛУЖБОВОГО КОРИСТУВАННЯ"
    assert sec.footer.paragraphs[0].text.strip() == "ДЛЯ СЛУЖБОВОГО КОРИСТУВАННЯ"
    assert sec.first_page_footer.paragraphs[0].text.strip() == "ДЛЯ СЛУЖБОВОГО КОРИСТУВАННЯ"


def test_add_page_number_header_inserts_centered_page_field(doc):
    helpers.add_page_number_header(doc)
    header_paragraph = doc.sections[0].header.paragraphs[0]
    assert header_paragraph.alignment == WD_ALIGN_PARAGRAPH.CENTER

    xml = header_paragraph._p.xml
    assert 'w:fldCharType="begin"' in xml
    assert 'w:fldCharType="end"' in xml
    assert "PAGE" in xml


def test_switch_landscape(doc):
    section = doc.sections[-1]
    assert section.orientation == WD_ORIENT.PORTRAIT
    original_width, original_height = section.page_width, section.page_height

    helpers.switch_landscape(doc)

    section = doc.sections[-1]
    assert section.orientation == WD_ORIENT.LANDSCAPE
    assert section.page_width == original_height
    assert section.page_height == original_width


def test_create_table_full_coverage(doc):
    headers = ["H1", "H2"]

    # Некоректні параметри (довжина widths_cm не збігається з кількістю заголовків)
    result = helpers.create_table(doc, headers=headers, data_rows=[["a", "b"]], widths_cm=[1])
    assert result is None
    assert len(doc.tables) == 0

    doc2 = Document()
    widths_cm = [2, 3]
    data_rows = [["a\nline2", "b"], ["c", "d"]]
    helpers.create_table(doc2, headers=headers, data_rows=data_rows, widths_cm=widths_cm,
                          font_size=10, font_name="Arial", italic=True, color_rgb=(1, 2, 3))

    tbl = doc2.tables[-1]
    for i, h in enumerate(headers):
        assert tbl.rows[0].cells[i].text == h

    for i, w in enumerate(widths_cm):
        for cell in tbl.columns[i].cells:
            assert abs(cell.width.cm - w) < 0.01

    for r, row in enumerate(data_rows, start=1):
        for c, val in enumerate(row):
            cell = tbl.rows[r].cells[c]
            expected_lines = str(val).split('\n')
            assert len(cell.paragraphs) == len(expected_lines)
            for p, expected_line in zip(cell.paragraphs, expected_lines):
                assert p.text == expected_line
                # cell.text = "" (усередині create_table) лишає порожній перший run
                # без шрифту — реальний вміст завжди в останньому доданому run.
                run = p.runs[-1]
                assert run.font.name == "Arial"
                assert run.font.size.pt == 10
                assert run.font.color.rgb == RGBColor(1, 2, 3)
                assert run.font.italic is True


def test_create_table_cell_alignment_centers_all_cells_vertically_too(doc):
    """Коли задано cell_alignment - усі комірки (заголовок і дані, будь-яка колонка)
    центруються і по горизонталі (алгоритм абзацу), і по вертикалі (w:vAlign)."""
    headers = ["H1", "H2"]
    data_rows = [["a", "b\nc"]]
    helpers.create_table(doc, headers=headers, data_rows=data_rows, widths_cm=[2, 3], cell_alignment=WD_ALIGN_PARAGRAPH.CENTER)

    tbl = doc.tables[-1]
    for r in range(2):
        for c in range(2):
            assert 'w:vAlign w:val="center"' in tbl.cell(r, c)._tc.xml


def test_create_table_no_cell_alignment_leaves_vertical_alignment_untouched(doc):
    headers = ["H1", "H2"]
    data_rows = [["a", "b"]]
    helpers.create_table(doc, headers=headers, data_rows=data_rows, widths_cm=[2, 3])

    tbl = doc.tables[-1]
    for r in range(2):
        for c in range(2):
            assert "w:vAlign" not in tbl.cell(r, c)._tc.xml


def test_create_table_cant_split_defaults_to_true(doc):
    """За замовчуванням (як і раніше) кожен рядок таблиці має w:cantSplit - підходить
    для таблиць із короткими рядками (напр. рапорт на додаткову винагороду), де рядок
    не повинен розриватись між сторінками."""
    headers = ["H1", "H2"]
    data_rows = [["a", "b"]]
    helpers.create_table(doc, headers=headers, data_rows=data_rows, widths_cm=[2, 3])

    tbl = doc.tables[-1]
    for row in tbl.rows:
        assert "w:cantSplit" in row._tr.xml


def test_create_table_cant_split_false_allows_row_to_split_across_pages():
    """cant_split=False (Витяги з ЖБД - одна дата, текст на кілька сторінок) не додає
    w:cantSplit - інакше Word не може ні розірвати рядок, ні вмістити його на одну
    сторінку, і замість цього малює порожню сторінку та накладає текст на колонтитул."""
    doc = Document()
    headers = ["H1", "H2"]
    data_rows = [["a", "b"]]
    helpers.create_table(doc, headers=headers, data_rows=data_rows, widths_cm=[2, 3], cant_split=False)

    tbl = doc.tables[-1]
    for row in tbl.rows:
        assert "w:cantSplit" not in row._tr.xml


# -------------------------
# get_validated_ksp_data (InquirerPy, через conftest FakePrompt)
# -------------------------
def test_get_validated_ksp_data_first_try_valid(inquirer_inputs):
    inquirer_inputs.extend(["Львів", "36T TT 54321 12345"])
    city, coords = helpers.get_validated_ksp_data()
    assert city == "Львів"
    assert coords == "36T TT 54321 12345"


def test_get_validated_ksp_data_retries_on_invalid(inquirer_inputs, capsys):
    inquirer_inputs.extend(["123", "Київ", "BAD", "Київ", "36T TT 12345 67890"])
    city, coords = helpers.get_validated_ksp_data()
    assert city == "Київ"
    assert coords == "36T TT 12345 67890"

    out = capsys.readouterr().out
    assert "Некоректна назва" in out
    assert "Некоректні координати" in out


def test_get_validated_ksp_data_city_with_dash_and_apostrophe(inquirer_inputs):
    inquirer_inputs.extend(["Кам'янець-Подільський", "36T TT 00001 99999"])
    city, coords = helpers.get_validated_ksp_data()
    assert city == "Кам'янець-Подільський"
    assert coords == "36T TT 00001 99999"


def test_get_validated_ksp_data_uses_defaults_on_enter(inquirer_inputs):
    # Порожня черга -> FakePrompt повертає default, як і реальний Enter користувача
    city, coords = helpers.get_validated_ksp_data()
    assert city == "КИЇВ"
    assert coords == "36T TT 12345 67890"


# -------------------------
# find_higher_commander (реальний to_date, без моків)
# -------------------------
def test_find_higher_commander_prefers_tvo_range():
    col = "01.11.2025"
    rows_tvo = [{'ПОСАДА': constants.HIGHER_COMMANDER_TITLE, 'Start': '01.11.2025', 'End': '05.11.2025', 'ТВО': True}]
    rows = [{'ПОСАДА': constants.HIGHER_COMMANDER_TITLE}]
    found = helpers.find_higher_commander(rows_tvo, rows, constants.HIGHER_COMMANDER_TITLE, col)
    assert found is not None
    assert found.get('ТВО') is True


def test_find_higher_commander_outside_range_fallback_to_rows():
    date_col = datetime(2025, 1, 10)
    rows_tvo = [{"ПОСАДА": "Командир", "ТВО": True, "Start": datetime(2025, 1, 11), "End": datetime(2025, 1, 20)}]
    rows = [{"ПОСАДА": "Командир", "Name": "Backup"}]
    result = helpers.find_higher_commander(rows_tvo, rows, "Командир", date_col)
    assert result is rows[0]


def test_find_higher_commander_no_tvo_rows_found_in_rows():
    date_col = datetime(2025, 1, 10)
    rows = [{"ПОСАДА": "Командир", "Name": "Real"}]
    result = helpers.find_higher_commander([], rows, "Командир", date_col)
    assert result is rows[0]


def test_find_higher_commander_not_found_anywhere():
    date_col = datetime(2025, 1, 10)
    rows_tvo = [{"ПОСАДА": "Командир", "ТВО": True, "Start": datetime(2025, 2, 1), "End": datetime(2025, 2, 5)}]
    rows = [{"ПОСАДА": "Інша посада"}]
    result = helpers.find_higher_commander(rows_tvo, rows, "Командир", date_col)
    assert result == {}


def test_find_higher_commander_exact_date_edges():
    date_col = datetime(2025, 3, 10)
    rows_tvo = [{"ПОСАДА": "Командир", "ТВО": True, "Start": datetime(2025, 3, 10), "End": datetime(2025, 3, 15)}]
    result = helpers.find_higher_commander(rows_tvo, [], "Командир", date_col)
    assert result is rows_tvo[0]


# -------------------------
# generate_content_br_general (SECTION_LISTS монкіпатчено фейковою секцією -
# раніше тест покладався на те, що тестове значення ПІДРОЗДІЛ реально збігається
# з ключем constants.SECTION_LISTS у ЛОКАЛЬНИХ resources/data.json конкретного
# користувача, що й було справжнім прихованим hardcode'ом реального підрозділу
# в трекованому коді; тепер секція - повністю синтетична, підставлена
# монкіпатчем, тож тест не залежить від того, які підрозділи реально є в
# resources/data.json)
#
# Монкіпатч іде напряму в helpers.generate_content_br_general.__globals__, А
# НЕ через "import content.br_general_catalog as br_catalog" +
# monkeypatch.setattr на цей СВІЖИЙ модуль: кореневий conftest.py евіктить
# content/helpers із sys.modules між проєктами (_COLLIDING_MODULES), тож
# свіжий import тут може дати ІНШИЙ об'єкт модуля, ніж той, на який уже
# посилається helpers.generate_content_br_general.__globals__ (прив'язаний ще
# при колекції файлу helpers_test.py) - монкіпатч на "не той" модуль тоді
# мовчки НЕ впливає на функцію під тестом (KeyError на нібито підставленому
# ключі, підтверджено падінням у повному прогоні всіх 5 проєктів разом,
# невидимо при ізольованому запуску лише цього проєкту). __globals__ - це
# буквально той самий namespace dict, який функція читає, незалежно від
# того, який об'єкт модуля наразі лежить у sys.modules.
# -------------------------
def test_generate_content_br_general_appends_row_to_real_section(monkeypatch):
    section_lists = {"підрозділ1": []}
    monkeypatch.setitem(helpers.generate_content_br_general.__globals__, "SECTION_LISTS", section_lists)
    col = datetime(2025, 11, 1)
    row = {"ПІДРОЗДІЛ": "Підрозділ 1", "ПОСАДА": "Командир роти", "ПІБ": "Перший Перший", col: 100}

    result = helpers.generate_content_br_general(
        rows_with_data=[row],
        rows_with_dowries_data=[],
        col_name=col,
        higher_commander_data={"ПОСАДА": constants.HIGHER_COMMANDER_TITLE},
    )

    assert row in result["підрозділ1"]
    # get_unique_sections_list будує новий словник, але наповнює його з того самого
    # (очищеного й перезаповненого) SECTION_LISTS
    assert row in section_lists["підрозділ1"]


def test_generate_content_br_general_skips_higher_commander(monkeypatch):
    monkeypatch.setitem(helpers.generate_content_br_general.__globals__, "SECTION_LISTS", {"підрозділ1": []})
    col = datetime(2025, 11, 1)
    row = {"ПІДРОЗДІЛ": "Підрозділ 1", "ПОСАДА": constants.HIGHER_COMMANDER_TITLE, "ПІБ": "Комбат", col: 100}

    result = helpers.generate_content_br_general(
        rows_with_data=[row],
        rows_with_dowries_data=[],
        col_name=col,
        higher_commander_data={"ПОСАДА": constants.HIGHER_COMMANDER_TITLE},
    )

    assert row not in result["підрозділ1"]


def test_generate_content_br_general_excludes_enemy_territory_status(monkeypatch):
    monkeypatch.setitem(helpers.generate_content_br_general.__globals__, "SECTION_LISTS", {"підрозділ1": []})
    col = datetime(2025, 11, 1)
    row = {"ПІДРОЗДІЛ": "Підрозділ 1", "ПОСАДА": "Стрілець", "ПІБ": "Перший Перший", col: 100}
    status_lookup = {"ПЕРШИЙ ПЕРШИЙ": "РТГр"}

    result = helpers.generate_content_br_general(
        rows_with_data=[row],
        rows_with_dowries_data=[],
        col_name=col,
        higher_commander_data={"ПОСАДА": constants.HIGHER_COMMANDER_TITLE},
        status_lookup=status_lookup,
    )

    assert row not in result["підрозділ1"]


def test_generate_content_br_general_keeps_non_enemy_territory_status_with_lookup(monkeypatch):
    monkeypatch.setitem(helpers.generate_content_br_general.__globals__, "SECTION_LISTS", {"підрозділ1": []})
    col = datetime(2025, 11, 1)
    row = {"ПІДРОЗДІЛ": "Підрозділ 1", "ПОСАДА": "Стрілець", "ПІБ": "Перший Перший", col: 100}
    status_lookup = {"ПЕРШИЙ ПЕРШИЙ": "БД(СЗ)"}

    result = helpers.generate_content_br_general(
        rows_with_data=[row],
        rows_with_dowries_data=[],
        col_name=col,
        higher_commander_data={"ПОСАДА": constants.HIGHER_COMMANDER_TITLE},
        status_lookup=status_lookup,
    )

    assert row in result["підрозділ1"]


@pytest.mark.parametrize("cell_value", [70, 170])
def test_generate_content_br_general_includes_70_and_170_cell_values(cell_value, monkeypatch):
    """70 і 170 - так само "бойові" дні для щоденного БР, як і 100 (нові точки
    MONEY_REPORT_CATEGORIES) - мають потрапляти в розділ так само, як 100."""
    monkeypatch.setitem(helpers.generate_content_br_general.__globals__, "SECTION_LISTS", {"підрозділ1": []})
    col = datetime(2025, 11, 1)
    row = {"ПІДРОЗДІЛ": "Підрозділ 1", "ПОСАДА": "Стрілець", "ПІБ": "Перший Перший", col: cell_value}

    result = helpers.generate_content_br_general(
        rows_with_data=[row],
        rows_with_dowries_data=[],
        col_name=col,
        higher_commander_data={"ПОСАДА": constants.HIGHER_COMMANDER_TITLE},
    )

    assert row in result["підрозділ1"]


@pytest.mark.parametrize("cell_value,status", [(70, "70_РТГр"), (170, "170_РТГр")])
def test_generate_content_br_general_excludes_70_170_enemy_territory_status(cell_value, status, monkeypatch):
    """Аналог test_generate_content_br_general_excludes_enemy_territory_status для
    нових точок 70/170 - їхні власні "*_РТГр" категорії теж не отримують БР."""
    monkeypatch.setitem(helpers.generate_content_br_general.__globals__, "SECTION_LISTS", {"підрозділ1": []})
    col = datetime(2025, 11, 1)
    row = {"ПІДРОЗДІЛ": "Підрозділ 1", "ПОСАДА": "Стрілець", "ПІБ": "Перший Перший", col: cell_value}
    status_lookup = {"ПЕРШИЙ ПЕРШИЙ": status}

    result = helpers.generate_content_br_general(
        rows_with_data=[row],
        rows_with_dowries_data=[],
        col_name=col,
        higher_commander_data={"ПОСАДА": constants.HIGHER_COMMANDER_TITLE},
        status_lookup=status_lookup,
    )

    assert row not in result["підрозділ1"]


def test_generate_content_br_general_enemy_territory_exclusion_is_dynamic(monkeypatch):
    """"Дії на території противника" вираховуються ДИНАМІЧНО з "generate_br": False у
    MONEY_REPORT_CATEGORIES (content/br_general_catalog.py::_enemy_territory_categories),
    а не за захардкодженим списком назв - категорія без цього прапорця (чи взагалі
    видалена) за замовчуванням НЕ виключається з БР."""
    import content.br_general_catalog as br_catalog
    import content.money_report_helpers as mrh

    monkeypatch.setattr(br_catalog, "SECTION_LISTS", {"підрозділ1": []})

    def _patch_categories(categories):
        monkeypatch.setattr(br_catalog, "MONEY_REPORT_CATEGORIES", categories)
        monkeypatch.setattr(mrh, "MONEY_REPORT_CATEGORIES", categories)

    _patch_categories({
        100: {
            "general": [],
            "ВОРОЖА": {"grounds": [], "use_brs": False, "generate_br": False},
            "БД(СЗ)": {"grounds": [], "use_brs": True},
        },
    })
    col = datetime(2025, 11, 1)
    excluded_row = {"ПІДРОЗДІЛ": "Підрозділ 1", "ПОСАДА": "Стрілець", "ПІБ": "Другий Другий", col: 100}
    included_row = {"ПІДРОЗДІЛ": "Підрозділ 1", "ПОСАДА": "Стрілець", "ПІБ": "Третій Третій", col: 100}
    status_lookup = {"ДРУГИЙ ДРУГИЙ": "ВОРОЖА", "ТРЕТІЙ ТРЕТІЙ": "БД(СЗ)"}

    result = br_catalog.generate_content_br_general(
        rows_with_data=[excluded_row, included_row],
        rows_with_dowries_data=[],
        col_name=col,
        higher_commander_data={"ПОСАДА": constants.HIGHER_COMMANDER_TITLE},
        status_lookup=status_lookup,
    )

    assert excluded_row not in result["підрозділ1"]
    assert included_row in result["підрозділ1"]

    # Той самий сценарій, але БЕЗ "generate_br" (прапорець прибрано) -
    # "ВОРОЖА" за замовчуванням більше не виключається з БР.
    _patch_categories({
        100: {
            "general": [],
            "ВОРОЖА": {"grounds": [], "use_brs": False},
            "БД(СЗ)": {"grounds": [], "use_brs": True},
        },
    })
    result_without_flag = br_catalog.generate_content_br_general(
        rows_with_data=[excluded_row, included_row],
        rows_with_dowries_data=[],
        col_name=col,
        higher_commander_data={"ПОСАДА": constants.HIGHER_COMMANDER_TITLE},
        status_lookup=status_lookup,
    )

    assert excluded_row in result_without_flag["підрозділ1"]




# -------------------------
# validate_battalion_commander_exists
# -------------------------
def test_validate_battalion_commander_exists_raises_when_missing():
    with pytest.raises(ValueError, match=constants.HIGHER_COMMANDER_TITLE):
        helpers.validate_battalion_commander_exists([{"ПОСАДА": "Щось інше"}], file_path="ОБЛІК.xlsx")


def test_validate_battalion_commander_exists_passes_when_present():
    helpers.validate_battalion_commander_exists([{"ПОСАДА": constants.HIGHER_COMMANDER_TITLE}])


def test_get_prev_general_returns_none_for_unparseable_date():
    assert helpers.get_prev_general("не дата", {"01.11.2025": {}}) is None


def test_get_prev_general_never_returns_the_date_itself():
    """Навіть якщо date_str теж є ключем d (як завжди для реального виклику -
    сьогоднішня дата вже має власний запис) - не має повертати саму себе, лише
    СТРОГО раніший запис."""
    d = {"01.07.2026": {"бат": "87"}, "02.07.2026": {"бат": "90"}}
    assert helpers.get_prev_general("02.07.2026", d) == "01.07.2026"


def test_get_prev_general_skips_entries_with_empty_bat():
    """Записи з порожнім 'бат' - лише заглушки сітки дат (напр. кінець попереднього
    місяця до першого реального щоденного розпорядження) - не рахуються "попереднім"
    розпорядженням, навіть якщо хронологічно раніші."""
    d = {"29.06.2026": {"бат": ""}, "30.06.2026": {"бат": ""}, "01.07.2026": {"бат": "88"}}
    assert helpers.get_prev_general("01.07.2026", d) is None
    assert helpers.get_prev_general("02.07.2026", d) == "01.07.2026"


def test_get_prev_general_picks_most_recent_of_several_earlier_entries():
    d = {"01.07.2026": {"бат": "87"}, "02.07.2026": {"бат": "90"}, "03.07.2026": {"бат": "92"}}
    assert helpers.get_prev_general("03.07.2026", d) == "02.07.2026"


# -------------------------
# get_bat_period_and_variant
# -------------------------
_BAT_MONTH = {
    "01.07.2026": {"бат": "87"},
    "02.07.2026": {"бат": ""},
    "07.07.2026": {"бат": "102"},
    "14.07.2026": {"бат": "121"},
    "21.07.2026": {"бат": "138"},
    "27.07.2026": {"бат": "154"},
    "28.07.2026": {"бат": ""},
    "01.08.2026": {"бат": "1"},
}


def test_get_bat_period_and_variant_first_of_month_starts_same_day():
    """Розпорядження від 1-го числа описує період, що починається того ж дня -
    закінчується датою НАСТУПНОГО фактичного БР бат (07.07), а не умовною межею тижня."""
    period_text, variant_index = helpers.get_bat_period_and_variant("01.07.2026", _BAT_MONTH)
    assert period_text == "з 01.07.2026 по 07.07.2026"
    assert variant_index == 0


def test_get_bat_period_and_variant_mid_month_starts_next_day():
    """Розпорядження не від 1-го числа описує НАСТУПНИЙ період (від наступного дня) -
    до дати НАСТУПНОГО фактичного БР бат."""
    period_text, variant_index = helpers.get_bat_period_and_variant("07.07.2026", _BAT_MONTH)
    assert period_text == "з 08.07.2026 по 14.07.2026"
    assert variant_index == 1


def test_get_bat_period_and_variant_uneven_real_gap_not_forced_to_seven_days():
    """Реальні БР видаються нерівномірно - 21.07 і 27.07 (лише 6 днів), а не 21-28,
    як було б за фіксованою тижневою сіткою."""
    period_text, variant_index = helpers.get_bat_period_and_variant("21.07.2026", _BAT_MONTH)
    assert period_text == "з 22.07.2026 по 27.07.2026"
    assert variant_index == 0


def test_get_bat_period_and_variant_last_bat_of_month_ends_on_last_day():
    """Якщо непустий 'бат' - останній у місяці, період закінчується останнім днем
    місяця (наступного БР бат цього місяця вже немає)."""
    period_text, variant_index = helpers.get_bat_period_and_variant("27.07.2026", _BAT_MONTH)
    assert period_text == "з 28.07.2026 по 31.07.2026"
    assert variant_index == 1


def test_get_bat_period_and_variant_variant_rotates_by_bat_order_not_by_week():
    """Індекс варіанту - порядковий номер серед фактичних БР бат цього місяця (0,1,2,0),
    а не номер тижня - навіть коли між БР менше/більше 7 днів."""
    assert helpers.get_bat_period_and_variant("01.07.2026", _BAT_MONTH)[1] == 0
    assert helpers.get_bat_period_and_variant("07.07.2026", _BAT_MONTH)[1] == 1
    assert helpers.get_bat_period_and_variant("14.07.2026", _BAT_MONTH)[1] == 2
    assert helpers.get_bat_period_and_variant("21.07.2026", _BAT_MONTH)[1] == 0
    assert helpers.get_bat_period_and_variant("27.07.2026", _BAT_MONTH)[1] == 1


def test_get_bat_period_and_variant_ignores_other_months_bat_dates():
    """Пошук наступного/попереднього БР бат обмежений МІСЯЦЕМ col_name - бат з іншого
    місяця (01.08) не впливає на розрахунок для 27.07."""
    period_text, _ = helpers.get_bat_period_and_variant("27.07.2026", _BAT_MONTH)
    assert "08.2026" not in period_text


# -------------------------
# excel_reader - файлові фолбеки, відсутній аркуш
# -------------------------
def test_find_file_with_any_extension_returns_original_if_exists(tmp_path):
    f = tmp_path / "data.xlsx"
    f.write_text("x")
    assert helpers.find_file_with_any_extension(str(f)) == str(f)


def test_find_file_with_any_extension_falls_back_to_other_extension(tmp_path, capsys):
    xlsm_path = tmp_path / "data.xlsm"
    xlsm_path.write_text("x")

    result = helpers.find_file_with_any_extension(str(tmp_path / "data.xlsx"))

    assert result == str(xlsm_path)
    assert "не знайдено" in capsys.readouterr().out


def test_find_file_with_any_extension_raises_when_no_extension_matches(tmp_path):
    with pytest.raises(FileNotFoundError):
        helpers.find_file_with_any_extension(str(tmp_path / "nope.xlsx"))


def test_resolve_sheet_name_raises_when_no_sheet_matches(tmp_path):
    import utils.excel_reader as excel_reader_module
    from openpyxl import Workbook

    path = tmp_path / "book.xlsx"
    wb = Workbook()
    wb.active.title = "СтороннійАркуш"
    wb.save(str(path))

    with pytest.raises(ValueError, match="не знайдено жодного"):
        excel_reader_module._resolve_sheet_name(str(path), "НеІснуючийАркуш", "openpyxl")


def test_resolve_sheet_name_finds_tvo_data_in_month_named_sheet_fallback(tmp_path):
    """Окремого аркуша "ТВО" вже немає (підтверджено користувачем) - таблиця
    Start/End/ПОСАДА/ПІБ/ТВО тепер ведеться прямо в аркуші поточного місяця
    (напр. "СЕРПЕНЬ") - TVO_LIST_FALLBACK_SHEET_NAMES дозволяє її знайти й
    там. Падало на реальних даних до цього фікса (ТВО завжди виходив
    порожнім, як для файлів, де аркуша "ТВО" справді нема - другий, ТВО-
    призначений в/сл тихо зникав із рапорту на командира)."""
    import utils.excel_reader as excel_reader_module
    from openpyxl import Workbook

    path = tmp_path / "book.xlsx"
    wb = Workbook()
    wb.active.title = "СЕРПЕНЬ"
    wb.save(str(path))

    resolved = excel_reader_module._resolve_sheet_name(
        str(path), "ТВО", "openpyxl", excel_reader_module.TVO_LIST_FALLBACK_SHEET_NAMES,
    )
    assert resolved == "СЕРПЕНЬ"


def test_tvo_fallback_names_do_not_include_tabel():
    """На відміну від PERSONEL_LIST_FALLBACK_SHEET_NAMES, тут НАВМИСНО немає
    "Табель" - той аркуш тепер аркуш ОСОБОВОГО СКЛАДУ (інша структура
    колонок), а не ТВО-таблиці - якби він тут теж був запасним варіантом,
    TVO_LIST_COLUMNS_LETTERS прочитав би з нього геть не ті колонки."""
    import utils.excel_reader as excel_reader_module

    assert "Табель" not in excel_reader_module.TVO_LIST_FALLBACK_SHEET_NAMES
    assert "СЕРПЕНЬ" in excel_reader_module.TVO_LIST_FALLBACK_SHEET_NAMES


def test_read_optional_datafile_returns_empty_list_when_file_missing(tmp_path, capsys):
    result = helpers.read_optional_datafile(str(tmp_path / "missing.xlsx"), "Sheet1", ["A"])
    assert result == []


def test_read_optional_datafile_returns_empty_list_when_sheet_missing_and_allowed(tmp_path, capsys):
    """sheet_can_be_missing=True: файл є, але шуканого аркуша (і жодного запасного)
    у ньому немає - ValueError від _resolve_sheet_name перехоплюється, а не падає
    назовні (сама відсутність аркуша - очікуваний стан, напр. "ТВО" у файлі, який
    зберіг сусідній проєкт, що про ТВО не знає)."""
    from openpyxl import Workbook

    path = tmp_path / "book.xlsx"
    wb = Workbook()
    wb.active.title = "СтороннійАркуш"
    wb.save(str(path))

    result = helpers.read_optional_datafile(str(path), "НеІснуючийАркуш", ["A"], sheet_can_be_missing=True)

    assert result == []
    assert "не знайдено" in capsys.readouterr().out


def test_read_optional_datafile_reraises_when_sheet_missing_and_not_allowed(tmp_path):
    """sheet_can_be_missing=False (за замовчуванням) - та сама відсутність аркуша
    ЛИШАЄТЬСЯ помилкою, а не мовчки ігнорується (напр. ПРИДАНІ.xlsx - відсутність
    очікуваного аркуша тут якраз і означає зіпсований/не той файл)."""
    from openpyxl import Workbook

    path = tmp_path / "book.xlsx"
    wb = Workbook()
    wb.active.title = "СтороннійАркуш"
    wb.save(str(path))

    with pytest.raises(ValueError, match="не знайдено жодного"):
        helpers.read_optional_datafile(str(path), "НеІснуючийАркуш", ["A"])


def test_detect_personel_list_columns_letters_raises_when_no_date_columns_for_month(tmp_path):
    import utils.excel_reader as excel_reader_module
    from openpyxl import Workbook
    from datetime import datetime as dt

    path = tmp_path / "oblik.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Аркуш1"
    ws.append(["ПІДРОЗДІЛ", "ПОСАДА", "ЗВАННЯ", "ПІБ", dt(2026, 5, 1)])
    wb.save(str(path))

    with pytest.raises(ValueError, match="не знайдено колонок з датами"):
        excel_reader_module.detect_personel_list_columns_letters(str(path), "Аркуш1", "07", "2026")


def test_detect_personel_list_columns_letters_detects_singular_pidstava_header(tmp_path):
    """Реальний ОБЛІК.xlsx називає опційну колонку в однині - "ПІДСТАВА" (не
    "ПІДСТАВИ") - detect_personel_list_columns_letters має підхоплювати ОБИДВА
    написання (utils.excel_reader.OPTIONAL_PERSONEL_COLUMN_NAMES)."""
    import utils.excel_reader as excel_reader_module
    from openpyxl import Workbook
    from datetime import datetime as dt

    path = tmp_path / "oblik.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Аркуш1"
    ws.append(["ПІДРОЗДІЛ", "ПОСАДА", "ЗВАННЯ", "ПІБ", dt(2026, 7, 1), "ПІДСТАВА"])
    wb.save(str(path))

    letters = excel_reader_module.detect_personel_list_columns_letters(str(path), "Аркуш1", "07", "2026")

    assert letters == ["A", "B", "C", "D", "E", "F"]


# -------------------------
# br_general_catalog - додані підрозділи (dowries) та пропуск секції без шаблону
# -------------------------
def test_generate_content_br_general_appends_matching_dowry_within_date_range(monkeypatch):
    monkeypatch.setitem(helpers.generate_content_br_general.__globals__, "SECTION_LISTS", {"підрозділ1": []})
    col = datetime(2025, 11, 1)
    dowry_row = {
        "ПРИДАНИЙ ДО": "Підрозділ 1",
        "військове звання": "сержант",
        "Прізвище та ініціали": "Перший П.П.",
        "ДАТА ПРИБУВ (З)": "25.10.2025",
        "ДАТА ВІДБУТТЯ (ПО)": None,
    }

    result = helpers.generate_content_br_general(
        rows_with_data=[],
        rows_with_dowries_data=[dowry_row],
        col_name=col,
        higher_commander_data={"ПОСАДА": constants.HIGHER_COMMANDER_TITLE},
    )

    assert any(d.get("ПІБ") == "Перший П.П." for d in result["підрозділ1"])


def test_generate_content_br_general_excludes_dowry_outside_date_range(monkeypatch):
    monkeypatch.setitem(helpers.generate_content_br_general.__globals__, "SECTION_LISTS", {"підрозділ1": []})
    col = datetime(2025, 11, 1)
    dowry_row = {
        "ПРИДАНИЙ ДО": "Підрозділ 1",
        "військове звання": "сержант",
        "Прізвище та ініціали": "Другий Д.Д.",
        "ДАТА ПРИБУВ (З)": "01.09.2025",
        "ДАТА ВІДБУТТЯ (ПО)": "01.10.2025",
    }

    result = helpers.generate_content_br_general(
        rows_with_data=[],
        rows_with_dowries_data=[dowry_row],
        col_name=col,
        higher_commander_data={"ПОСАДА": constants.HIGHER_COMMANDER_TITLE},
    )

    assert all(d.get("ПІБ") != "Другий Д.Д." for d in result["підрозділ1"])


def test_get_content_log_general_extract_log_war_skips_section_without_data():
    result = helpers.get_content_log_general_extract_log_war(
        section_lists={"Підрозділ 1": []},
        paragraph_map={"Підрозділ 1": "Шаблон {unit} {city} ({coordinates}):"},
        city="КИЇВ",
        coordinates="36T TT 12345 67890",
    )
    assert result == ""


# -------------------------
# excel_reader.read_optional_pridani_sheet - опційний аркуш "ПРИДАНІ"
# -------------------------
def _write_pridani_xlsx(path, header, rows, sheet_name="ПРИДАНІ"):
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name
    ws.append(header)
    for row in rows:
        ws.append(row)
    wb.save(str(path))
    return str(path)


def test_read_optional_pridani_sheet_returns_empty_when_sheet_missing(tmp_path):
    import utils.excel_reader as excel_reader_module

    path = _write_pridani_xlsx(
        tmp_path / "ПРИДАНІ.xlsx",
        ["ПРИДАНИЙ ДО", "ВІЙСЬКОВЕ ЗВАННЯ", "ПІБ", datetime(2026, 7, 1)],
        [["Підрозділ 1", "сержант", "Перший Перший", 100]],
        sheet_name="СтороннійАркуш",
    )

    result = excel_reader_module.read_optional_pridani_sheet(path, "ПРИДАНІ", "07", "2026")

    assert result == []


def test_read_optional_pridani_sheet_returns_empty_when_no_pib_column(tmp_path):
    import utils.excel_reader as excel_reader_module

    path = _write_pridani_xlsx(
        tmp_path / "ПРИДАНІ.xlsx",
        ["ПРИДАНИЙ ДО", "ВІЙСЬКОВЕ ЗВАННЯ", datetime(2026, 7, 1)],
        [["Підрозділ 1", "сержант", 100]],
    )

    result = excel_reader_module.read_optional_pridani_sheet(path, "ПРИДАНІ", "07", "2026")

    assert result == []


def test_read_optional_pridani_sheet_returns_empty_when_no_matching_date_columns(tmp_path):
    import utils.excel_reader as excel_reader_module

    path = _write_pridani_xlsx(
        tmp_path / "ПРИДАНІ.xlsx",
        ["ПРИДАНИЙ ДО", "ВІЙСЬКОВЕ ЗВАННЯ", "ПІБ", datetime(2026, 6, 1)],
        [["Підрозділ 1", "сержант", "Перший Перший", 100]],
    )

    result = excel_reader_module.read_optional_pridani_sheet(path, "ПРИДАНІ", "07", "2026")

    assert result == []


def test_read_optional_pridani_sheet_skips_rows_with_blank_pib(tmp_path):
    import utils.excel_reader as excel_reader_module

    path = _write_pridani_xlsx(
        tmp_path / "ПРИДАНІ.xlsx",
        ["ПРИДАНИЙ ДО", "ВІЙСЬКОВЕ ЗВАННЯ", "ПІБ", datetime(2026, 7, 1)],
        [
            ["Підрозділ 1", "сержант", "Перший Перший", 100],
            ["Підрозділ 1", "матрос", None, 100],
        ],
    )

    result = excel_reader_module.read_optional_pridani_sheet(path, "ПРИДАНІ", "07", "2026")

    assert len(result) == 1
    assert result[0]["ПІБ"] == "Перший Перший"


def test_get_content_log_general_extract_log_war_skips_section_without_template():
    result = helpers.get_content_log_general_extract_log_war(
        section_lists={"невідома_секція": [{"ЗВАННЯ": "сержант", "ПІБ": "Перший"}]},
        paragraph_map={},
        city="КИЇВ",
        coordinates="36T TT 12345 67890",
    )
    assert result == ""


def test_get_catalog_log_general_extract_log_war_builds_from_store():
    store = {"Підрозділ 1": "Шаблон {unit} {city} ({coordinates}):"}
    result = helpers.get_catalog_log_general_extract_log_war(
        section_lists={"Підрозділ 1": [{"ЗВАННЯ": "сержант", "ПІБ": "Перший Перший"}]},
        store=store,
        city="КИЇВ",
        coordinates="36T TT 12345 67890",
    )
    assert "Перший Перший" in result
    assert "КИЇВ" in result


# -------------------------
# docx_utils - взаємодія з відкритим Word (_close_word_documents_under) та
# повторні спроби видалення (_rmtree_retry), через публічний
# create_or_clear_output_directory
# -------------------------
def _br_output_dirs(base):
    return str(base), str(base / "br"), str(base / "br_save"), str(base / "extractbr"), str(base / "logwar")


def test_create_or_clear_output_directory_word_com_unavailable(tmp_path, monkeypatch):
    monkeypatch.setitem(sys.modules, "win32com.client", None)
    base = tmp_path / "out"
    base.mkdir()

    helpers.create_or_clear_output_directory(*_br_output_dirs(base))

    assert base.exists()


def test_create_or_clear_output_directory_word_not_running(tmp_path, monkeypatch):
    import pythoncom
    import win32com.client

    monkeypatch.setattr(pythoncom, "CoInitialize", lambda: None)
    monkeypatch.setattr(pythoncom, "CoUninitialize", lambda: None)

    def _raise(name):
        raise Exception("Word.Application не запущено")

    monkeypatch.setattr(win32com.client, "GetActiveObject", _raise)

    base = tmp_path / "out"
    base.mkdir()

    helpers.create_or_clear_output_directory(*_br_output_dirs(base))

    assert base.exists()


class _FakeBRWordDoc:
    def __init__(self, full_name):
        self.FullName = full_name
        self.closed_with = None

    def Close(self, SaveChanges=False):
        self.closed_with = SaveChanges


class _BrokenBRWordDoc:
    @property
    def FullName(self):
        raise RuntimeError("помилка COM при читанні шляху")


def test_create_or_clear_output_directory_closes_matching_word_docs(tmp_path, monkeypatch, capsys):
    import pythoncom
    import win32com.client

    base = tmp_path / "out"
    base.mkdir()
    target_file = base / "report.docx"
    target_file.write_text("x")

    matching_doc = _FakeBRWordDoc(str(target_file))
    unrelated_doc = _FakeBRWordDoc(str(tmp_path / "other" / "unrelated.docx"))
    broken_doc = _BrokenBRWordDoc()

    class FakeWordApp:
        Documents = [matching_doc, broken_doc, unrelated_doc]

    monkeypatch.setattr(pythoncom, "CoInitialize", lambda: None)
    monkeypatch.setattr(pythoncom, "CoUninitialize", lambda: None)
    monkeypatch.setattr(win32com.client, "GetActiveObject", lambda name: FakeWordApp())

    helpers.create_or_clear_output_directory(*_br_output_dirs(base))

    assert matching_doc.closed_with is False
    assert unrelated_doc.closed_with is None
    assert "Закриваю відкритий у Word файл" in capsys.readouterr().out


def test_create_or_clear_output_directory_retries_then_succeeds(tmp_path, monkeypatch):
    import formatting.docx_utils as docx_utils_module

    base = tmp_path / "out"
    base.mkdir()
    (base / "old.txt").write_text("x")

    real_rmtree = docx_utils_module.shutil.rmtree
    call_count = {"n": 0}

    def flaky_rmtree(path):
        call_count["n"] += 1
        if call_count["n"] == 1:
            raise PermissionError("заблоковано")
        return real_rmtree(path)

    monkeypatch.setattr(docx_utils_module.shutil, "rmtree", flaky_rmtree)
    monkeypatch.setattr(docx_utils_module.time, "sleep", lambda s: None)

    helpers.create_or_clear_output_directory(*_br_output_dirs(base))

    assert base.exists()
    assert call_count["n"] >= 2


def test_create_or_clear_output_directory_raises_after_all_retries_fail(tmp_path, monkeypatch):
    import formatting.docx_utils as docx_utils_module

    base = tmp_path / "out"
    base.mkdir()
    (base / "old.txt").write_text("x")

    def always_fails(path):
        raise PermissionError("завжди заблоковано")

    monkeypatch.setattr(docx_utils_module.shutil, "rmtree", always_fails)
    monkeypatch.setattr(docx_utils_module.time, "sleep", lambda s: None)

    with pytest.raises(PermissionError, match="Не вдалось видалити"):
        helpers.create_or_clear_output_directory(*_br_output_dirs(base))


# -------------------------
# save_docx_safely / _close_word_document_if_open (рапорт на дод. винагороду -
# ізольовано від _close_word_documents_under, бо закриває лише ОДИН точний файл)
# -------------------------
def test_close_word_document_if_open_word_com_unavailable(tmp_path, monkeypatch):
    import formatting.docx_utils as docx_utils_module

    monkeypatch.setitem(sys.modules, "win32com.client", None)
    docx_utils_module._close_word_document_if_open(str(tmp_path / "report.docx"))


def test_close_word_document_if_open_word_not_running(monkeypatch):
    import pythoncom
    import win32com.client
    import formatting.docx_utils as docx_utils_module

    monkeypatch.setattr(pythoncom, "CoInitialize", lambda: None)
    monkeypatch.setattr(pythoncom, "CoUninitialize", lambda: None)

    def _raise(name):
        raise Exception("Word.Application не запущено")

    monkeypatch.setattr(win32com.client, "GetActiveObject", _raise)

    docx_utils_module._close_word_document_if_open("report.docx")


def test_close_word_document_if_open_closes_only_exact_match(tmp_path, monkeypatch, capsys):
    import pythoncom
    import win32com.client
    import formatting.docx_utils as docx_utils_module

    target_file = tmp_path / "report.docx"
    target_file.write_text("x")

    matching_doc = _FakeBRWordDoc(str(target_file))
    unrelated_doc = _FakeBRWordDoc(str(tmp_path / "other.docx"))
    broken_doc = _BrokenBRWordDoc()

    class FakeWordApp:
        Documents = [matching_doc, broken_doc, unrelated_doc]

    monkeypatch.setattr(pythoncom, "CoInitialize", lambda: None)
    monkeypatch.setattr(pythoncom, "CoUninitialize", lambda: None)
    monkeypatch.setattr(win32com.client, "GetActiveObject", lambda name: FakeWordApp())

    docx_utils_module._close_word_document_if_open(str(target_file))

    assert matching_doc.closed_with is False
    assert unrelated_doc.closed_with is None
    assert "Закриваю відкритий у Word файл" in capsys.readouterr().out


def test_save_docx_safely_succeeds_first_try(tmp_path, monkeypatch):
    import formatting.docx_utils as docx_utils_module

    monkeypatch.setattr(docx_utils_module, "_close_word_document_if_open", lambda path: None)

    class FakeDoc:
        def __init__(self):
            self.saved_to = None

        def save(self, path):
            self.saved_to = path

    doc = FakeDoc()
    target = str(tmp_path / "report.docx")
    docx_utils_module.save_docx_safely(doc, target)

    assert doc.saved_to == target


def test_save_docx_safely_retries_then_succeeds(monkeypatch):
    import formatting.docx_utils as docx_utils_module

    monkeypatch.setattr(docx_utils_module, "_close_word_document_if_open", lambda path: None)
    monkeypatch.setattr(docx_utils_module.time, "sleep", lambda s: None)

    call_count = {"n": 0}

    class FlakyDoc:
        def save(self, path):
            call_count["n"] += 1
            if call_count["n"] == 1:
                raise PermissionError("заблоковано")

    docx_utils_module.save_docx_safely(FlakyDoc(), "report.docx")

    assert call_count["n"] >= 2


def test_save_docx_safely_raises_after_all_retries_fail(monkeypatch):
    import formatting.docx_utils as docx_utils_module

    monkeypatch.setattr(docx_utils_module, "_close_word_document_if_open", lambda path: None)
    monkeypatch.setattr(docx_utils_module.time, "sleep", lambda s: None)

    class AlwaysFailsDoc:
        def save(self, path):
            raise PermissionError("завжди заблоковано")

    with pytest.raises(PermissionError, match="Не вдалось зберегти"):
        docx_utils_module.save_docx_safely(AlwaysFailsDoc(), "report.docx")


# -------------------------
# generate_documents_br_general - секція без запису в PARAGRAPH_MAP
# -------------------------
def test_generate_content_br_returns_early_for_unknown_section(doc):
    from generator_br_and_report_for_money.generators.generate_documents_br_every_day import generate_content_br

    before = len(doc.paragraphs)
    generate_content_br(doc, "невідома_секція_якої_немає", [{"ЗВАННЯ": "сержант", "ПІБ": "Перший"}])

    assert len(doc.paragraphs) == before


# -------------------------
# constants - гілка "робота з папкою етапу" (виконується лише на імпорті)
# -------------------------
def test_constants_stage_folder_branch_sets_paths_when_stage_selected(tmp_path, monkeypatch):
    import importlib.util
    import os as os_module
    import sync.stage_folder as stage_folder_module

    month_dir = tmp_path / "07.2026"
    stage_dir = month_dir / "1 етап"
    stage_dir.mkdir(parents=True)

    monkeypatch.setattr(stage_folder_module, "prompt_should_open_stage_folder", lambda: True)
    monkeypatch.setattr(
        stage_folder_module, "prompt_stage_selection",
        lambda project_dir: (str(month_dir), str(stage_dir), True),
    )
    monkeypatch.setattr(
        stage_folder_module, "sync_stage_constants_backup",
        lambda *args, **kwargs: str(stage_dir / "constants.py.bak"),
    )

    # Глобальне ім'я `constants` у цьому файлі теоретично може (через колізію
    # однакових імен модулів між generator_br_and_report_for_money і final_combat_report в одному
    # pytest-процесі - див. коментар у кореневому conftest.py) опинитись
    # прив'язаним не до того файлу, а importlib.reload() довіряє вже
    # встановленому module.__spec__, тож не рятує від цього. Щоб перевірка
    # гілки "робота з папкою етапу" не залежала від цього взагалі, виконуємо
    # САМЕ generator_br_and_report_for_money/constants.py як окремий, повністю ізольований модуль
    # напряму за шляхом до файлу - без sys.modules і без впливу на інші тести.
    # Тестова тека етапу порожня (немає ОБЛІК.xlsx) - автоматичне визначення колонок
    # (нижче за виконанням, ніж усі перевірені тут атрибути) тому зупинить модуль
    # через SystemExit, не дійшовши до кінця файлу; самі перевірені нижче атрибути
    # (STAGE_DIR, PERSONEL_LIST_FILE_NAME тощо) уже встановлені до цього моменту.
    constants_path = str(Path(__file__).resolve().parent.parent / "constants.py")
    spec = importlib.util.spec_from_file_location("generator_br_and_report_for_money_constants_for_stage_test", constants_path)
    br_constants = importlib.util.module_from_spec(spec)
    with pytest.raises(SystemExit):
        spec.loader.exec_module(br_constants)

    assert br_constants.USING_STAGE_FOLDER is True
    assert br_constants.STAGE_DIR == str(stage_dir)
    assert br_constants.PERSONEL_LIST_FILE_NAME == os_module.path.join(str(stage_dir), "resources", "ОБЛІК.xlsx")
    assert br_constants.DOWRIES_LIST_FILE_NAME == os_module.path.join(str(stage_dir), "resources", "ПРИДАНІ.xlsx")
    assert br_constants.OUTPUT_DIR == os_module.path.join(str(stage_dir), "output")
    assert br_constants.OUTPUT_DIR_BR == os_module.path.join(str(stage_dir), "output", "br")
    assert br_constants.STAGE_CONSTANTS_BACKUP_PATH == str(stage_dir / "constants.py.bak")


def test_constants_stage_folder_branch_not_used_when_no_stage_selected():
    # sanity-перевірка стандартного (немодифікованого) стану, з яким працюють усі інші тести
    assert constants.USING_STAGE_FOLDER is False
    assert constants.STAGE_DIR is None


def test_constants_month_uses_current_month_when_interactive_prompts_skipped(monkeypatch):
    """SKIP_CONSTANTS_INTERACTIVE_PROMPTS=1 (index.py виставляє це ДО імпорту
    для деяких режимів - run_mode_prompt.MODES_WITHOUT_CONSTANTS_PROMPTS,
    напр. "Перевірка рапорту та ЖБД", де питання місяця геть не по темі) -
    MONTH береться напряму з поточної дати, без InquirerPy-запиту. Той самий
    ізольований спосіб завантаження модуля, що й test_constants_stage_folder_
    branch_sets_paths_when_stage_selected вище (і той самий прапорець ТАКОЖ
    вимикає гілку "робота з папкою етапу" - short-circuit на "not
    _SKIP_INTERACTIVE_PROMPTS and ..." - тож жодного InquirerPy-виклику тут
    не відбувається взагалі, монкіпатчити нема чого)."""
    import importlib.util
    from datetime import datetime as dt

    monkeypatch.setenv("SKIP_CONSTANTS_INTERACTIVE_PROMPTS", "1")

    constants_path = str(Path(__file__).resolve().parent.parent / "constants.py")
    spec = importlib.util.spec_from_file_location("generator_br_and_report_for_money_constants_for_skip_test", constants_path)
    br_constants = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(br_constants)

    assert br_constants.MONTH == dt.now().strftime("%m")
