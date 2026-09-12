import os
import sys
from datetime import datetime
from pathlib import Path

sys.path.append(str(Path(__file__).resolve().parent.parent))

from docx import Document

import generators.generate_documents_br_weekly_task as weekly_task_module

_WEEK = {"01.07.2026": {"бат": "87"}, "07.07.2026": {"бат": "102"}}


def test_generate_documents_br_saves_full_document(tmp_path, monkeypatch):
    """Наскрізна перевірка (не окремих текстових шматків, як
    test_combat_log_extract_generators.py, а самого факту генерації документа
    "ЗАВДАННЯ" з правильною назвою файлу) - раніше покривалось лише через
    реальні дані ОБЛІК.xlsx (test_integration.py), тепер - синтетичними, що не
    залежать від того, який місяць зараз веде користувач у реальному файлі."""
    monkeypatch.setattr(weekly_task_module, "NUMBER_OF_DOCUMENTS_BRS_EVERY_WEEK", _WEEK)
    monkeypatch.setattr(weekly_task_module, "TASK_ORDER_VARIANTS", [[("Підрозділ 1", "Перший варіант")], [("Підрозділ 1", "Другий варіант")]])

    higher_commander_data = {"ЗВАННЯ": "підполковник", "ПІБ": "ПЕРШИЙ Перший Перший", "ТВО": False}
    doc = Document()

    weekly_task_module.generate_documents_br(
        doc, datetime(2026, 7, 1), higher_commander_data, "КИЇВ", "36T TT 12345 67890", str(tmp_path),
    )

    files = os.listdir(tmp_path)
    assert any(f.endswith("ЗАВДАННЯ 01.07.2026.docx") for f in files)


def test_generate_documents_br_extract_variant_uses_different_intro_text(tmp_path, monkeypatch):
    monkeypatch.setattr(weekly_task_module, "NUMBER_OF_DOCUMENTS_BRS_EVERY_WEEK", _WEEK)
    monkeypatch.setattr(weekly_task_module, "TASK_ORDER_VARIANTS", [[("Підрозділ 1", "Варіант")]])

    higher_commander_data = {"ЗВАННЯ": "підполковник", "ПІБ": "ПЕРШИЙ Перший Перший", "ТВО": True}
    doc = Document()

    weekly_task_module.generate_documents_br(
        doc, datetime(2026, 7, 1), higher_commander_data, "КИЇВ", "36T TT 12345 67890", str(tmp_path), is_extract_br=True,
    )

    text = "\n".join(p.text for p in doc.paragraphs)
    assert "ВИТЯГ з БОЙОВОГО" in text
    assert "ТВО командира" in text


def test_generate_order_section_falls_back_to_first_variant_when_not_filled_in(capsys, monkeypatch):
    """Регресія: варіант переліку завдань, на який вказує get_bat_period_and_variant,
    ще порожній у constants.py (TASK_ORDER_VARIANTS[variant_index] == []) -
    використовується варіант №1 (TASK_ORDER_VARIANTS[0]), а не порожній розділ
    "НАКАЗАВ:". 07.07 - ДРУГА дата місяця (index=1) -> variant_index=1 (порожній
    у цьому тесті), тоді як TASK_ORDER_VARIANTS[0] (перша дата, 01.07) - справжній
    заповнений варіант, який і має стати запасним."""
    week = {"01.07.2026": {"бат": "87"}, "07.07.2026": {"бат": "102"}}
    monkeypatch.setattr(weekly_task_module, "NUMBER_OF_DOCUMENTS_BRS_EVERY_WEEK", week)
    monkeypatch.setattr(weekly_task_module, "TASK_ORDER_VARIANTS", [[("Підрозділ 1", "Запасний варіант")], []])

    doc = Document()
    weekly_task_module.generate_order_section(doc, datetime(2026, 7, 7))

    text = "\n".join(p.text for p in doc.paragraphs)
    assert "Підрозділ 1. Запасний варіант" in text
    assert "ще не заповнено" in capsys.readouterr().out
