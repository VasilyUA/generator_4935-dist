import sys
from pathlib import Path

sys.path.append(str(Path(__file__).resolve().parent.parent))

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import Rule
from openpyxl.styles import Alignment, Border, Color, Font, PatternFill, Side
from openpyxl.styles.differential import DifferentialStyle
from openpyxl.utils import get_column_letter

import content.oblik_style_reference as osr


def _build_oblik_fixture(path, extra_rules=()):
    """Мінімальний файл на кшталт ОБЛІК.xlsx - заголовок (ПІДРОЗДІЛ/ПОСАДА/ЗВАННЯ/
    ПІБ/дата), тілесний рядок, ширини колонок, висоти рядків і умовне форматування
    (containsText) з ОБОМА видами заливки - прямий RGB ("СЗЧ") і тема+tint ("30") -
    той самий формат, що й реальний ОБЛІК.xlsx (перевірено вручну на реальному файлі)."""
    wb = Workbook()
    ws = wb.active
    ws.title = "ТРАВЕНЬ"

    ws["A1"], ws["B1"], ws["C1"], ws["D1"], ws["E1"] = "ПІДРОЗДІЛ", "ПОСАДА", "ЗВАННЯ", "ПІБ", "01.07.2026"
    ws.cell(row=1, column=2).font = Font(name="Bahnschrift", bold=True, size=14)
    ws.cell(row=1, column=4).alignment = Alignment(horizontal="center", wrap_text=True)
    ws.cell(row=1, column=5).alignment = Alignment(textRotation=90)

    ws.cell(row=1, column=2).border = Border(top=Side(style="medium"), bottom=Side(style="medium"), left=Side(style="medium"), right=Side(style="medium"))

    ws["B2"] = "Стрілець"
    ws.cell(row=2, column=2).font = Font(name="Times New Roman", size=10)
    ws.cell(row=2, column=2).alignment = Alignment(horizontal="center")
    ws.cell(row=2, column=2).border = Border(top=Side(style="thin"), bottom=Side(style="thin"), left=Side(style="thin"), right=Side(style="thin"))

    ws.column_dimensions["B"].width = 27.0
    ws.column_dimensions["D"].width = 36.0
    ws.column_dimensions["E"].width = 8.0
    ws.row_dimensions[1].height = 120.0
    ws.row_dimensions[2].height = 25.5

    dxf_rgb = DifferentialStyle(fill=PatternFill(bgColor=Color(rgb="FFC00000")))
    ws.conditional_formatting.add("E2:E100", Rule(
        type="containsText", operator="containsText", text="СЗЧ",
        formula=['NOT(ISERROR(SEARCH("СЗЧ",E2)))'], dxf=dxf_rgb,
    ))
    dxf_theme = DifferentialStyle(fill=PatternFill(bgColor=Color(theme=4, tint=0.6)))
    ws.conditional_formatting.add("E2:E100", Rule(
        type="containsText", operator="containsText", text="30",
        formula=['NOT(ISERROR(SEARCH("30",E2)))'], dxf=dxf_theme,
    ))
    for rule in extra_rules:
        ws.conditional_formatting.add("E2:E100", rule)

    wb.save(str(path))
    return str(path)


# -------------------------
# _apply_tint
# -------------------------
def test_apply_tint_lightens_for_positive_tint():
    # 5B9BD5 - accent1 з реальної теми Office, tint 0.6 -> помітно світліший відтінок.
    lightened = osr._apply_tint("5B9BD5", 0.6)
    assert lightened != "5B9BD5"
    r, g, b = (int(lightened[i:i + 2], 16) for i in (0, 2, 4))
    assert (r, g, b) > (0x5B, 0x9B, 0xD5)


def test_apply_tint_darkens_for_negative_tint():
    darkened = osr._apply_tint("5B9BD5", -0.5)
    r, g, b = (int(darkened[i:i + 2], 16) for i in (0, 2, 4))
    assert (r, g, b) < (0x5B, 0x9B, 0xD5)


# -------------------------
# _resolve_fill_color
# -------------------------
def test_resolve_fill_color_returns_none_for_missing_color():
    assert osr._resolve_fill_color(None, {}) is None


def test_resolve_fill_color_resolves_direct_rgb():
    assert osr._resolve_fill_color(Color(rgb="FFC00000"), {}) == "C00000"


def test_resolve_fill_color_resolves_theme_with_tint():
    theme_colors = {"accent1": "5B9BD5"}
    resolved = osr._resolve_fill_color(Color(theme=4, tint=0.6), theme_colors)
    assert resolved == osr._apply_tint("5B9BD5", 0.6)


def test_resolve_fill_color_returns_none_for_unresolvable_color():
    # Індексований колір (палітра Excel 2003) - тут свідомо не підтримується.
    assert osr._resolve_fill_color(Color(indexed=5), {}) is None
    # Ім'я кольору теми відсутнє серед прочитаних з theme1.xml.
    assert osr._resolve_fill_color(Color(theme=4, tint=0.0), {}) is None


# -------------------------
# load_timetable_style_reference
# -------------------------
def test_load_timetable_style_reference_extracts_fonts_widths_and_keyword_colors(tmp_path, monkeypatch):
    path = _build_oblik_fixture(tmp_path / "ОБЛІК.xlsx")
    monkeypatch.setattr(osr, "PERSONEL_LIST_FILE_NAME", path)
    monkeypatch.setattr(osr, "PERSONEL_LIST_SHEET_NAME", "ТРАВЕНЬ")

    style = osr.load_timetable_style_reference()

    assert style is not None
    assert style["header_font"].name == "Bahnschrift"
    assert style["header_alignment_date"].textRotation == 90
    assert style["body_font"].name == "Times New Roman"
    assert style["label_col_width"] == 27.0
    assert style["pib_col_width"] == 36.0
    assert style["date_col_width"] == 8.0
    assert style["header_row_height"] == 120.0
    assert style["body_row_height"] == 25.5
    assert style["header_border"].top.style == "medium"
    assert style["body_border"].top.style == "thin"
    assert style["keyword_colors"]["СЗЧ"] == "C00000"
    # accent1 типової теми openpyxl (Office) - "4F81BD".
    assert style["keyword_colors"]["30"] == osr._apply_tint("4F81BD", 0.6)
    # Пріоритет - "СЗЧ" (priority 1) перед "30" (priority 2), як і в самому файлі.
    assert list(style["keyword_colors"]) == ["СЗЧ", "30"]


def test_load_timetable_style_reference_falls_back_to_first_sheet_when_expected_sheet_missing(tmp_path, monkeypatch):
    path = _build_oblik_fixture(tmp_path / "ОБЛІК.xlsx")
    monkeypatch.setattr(osr, "PERSONEL_LIST_FILE_NAME", path)
    monkeypatch.setattr(osr, "PERSONEL_LIST_SHEET_NAME", "НЕІСНУЮЧИЙ")

    assert osr.load_timetable_style_reference() is not None


def test_load_timetable_style_reference_returns_none_when_file_missing(tmp_path, monkeypatch):
    monkeypatch.setattr(osr, "PERSONEL_LIST_FILE_NAME", str(tmp_path / "не_існує.xlsx"))

    assert osr.load_timetable_style_reference() is None


def test_load_timetable_style_reference_ignores_non_containstext_and_unmatched_rules(tmp_path, monkeypatch):
    """duplicateValues (без dxf-заливки, яку тут розпізнати можна) і containsText
    без "SEARCH(...)" у формулі - жодне з двох не повинно потрапити в keyword_colors,
    і не повинно ламати решту зчитування."""
    from openpyxl.formatting.rule import Rule as RuleCls

    extra_rules = [
        RuleCls(type="duplicateValues"),
        RuleCls(type="containsText", operator="containsText", text="ВП", formula=["SOME_OTHER_FORMULA()"], dxf=DifferentialStyle(fill=PatternFill(bgColor=Color(rgb="FFFFFF00")))),
    ]
    path = _build_oblik_fixture(tmp_path / "ОБЛІК.xlsx", extra_rules=extra_rules)
    monkeypatch.setattr(osr, "PERSONEL_LIST_FILE_NAME", path)
    monkeypatch.setattr(osr, "PERSONEL_LIST_SHEET_NAME", "ТРАВЕНЬ")

    style = osr.load_timetable_style_reference()

    assert "ВП" not in style["keyword_colors"]
    assert set(style["keyword_colors"]) == {"СЗЧ", "30"}


def test_load_timetable_style_reference_keeps_first_priority_color_for_duplicate_keyword(tmp_path, monkeypatch):
    """Два правила з ОДНАКОВИМ ключовим словом (можливо в різних діапазонах) -
    перемагає те, що має ВИЩИЙ пріоритет (менше число), як і в самому Excel."""
    from openpyxl.formatting.rule import Rule as RuleCls

    extra_rules = [
        RuleCls(type="containsText", operator="containsText", text="30", formula=['NOT(ISERROR(SEARCH("30",E2)))'], dxf=DifferentialStyle(fill=PatternFill(bgColor=Color(rgb="FF000000")))),
    ]
    path = _build_oblik_fixture(tmp_path / "ОБЛІК.xlsx", extra_rules=extra_rules)
    monkeypatch.setattr(osr, "PERSONEL_LIST_FILE_NAME", path)
    monkeypatch.setattr(osr, "PERSONEL_LIST_SHEET_NAME", "ТРАВЕНЬ")

    style = osr.load_timetable_style_reference()

    # Перше (найвищий пріоритет) правило для "30" - тема+tint, а НЕ додане
    # пізніше (нижчий пріоритет) чорне FF000000.
    # accent1 типової теми openpyxl (Office) - "4F81BD".
    assert style["keyword_colors"]["30"] == osr._apply_tint("4F81BD", 0.6)
