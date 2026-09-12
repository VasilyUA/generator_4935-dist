import os
import sys
from pathlib import Path

sys.path.append(str(Path(__file__).resolve().parent.parent))

from docx import Document

import generators.generate_documents_combat_log_extract_war_every_week as weekly_module
import generators.generate_extract_log_war_general_br_every_day as daily_module
import formatting.docx_utils as docx_utils_module


# -------------------------
# generate_main_combat_log_extract_war_every_week
# -------------------------
def _fake_get_number_br(number_objects):
    def _get(col_name):
        return {"брг": "1", "бат": number_objects.get(col_name, {}).get("бат", "")}
    return _get


def test_weekly_extract_point_6_has_no_reference_on_first_bat_of_month(monkeypatch):
    """01.07 - перший тижневий БР місяця, попереднього просто не існує - пункт 6
    лишається без посилання ("..."), а пункт 5 (сьогоднішнє рішення) і пункт 4
    (бригадне розпорядження) заповнені як завжди."""
    week = {"01.07.2026": {"бат": "87"}, "07.07.2026": {"бат": "102"}}
    monkeypatch.setattr(weekly_module, "get_number_br", _fake_get_number_br(week))
    monkeypatch.setattr(weekly_module, "NUMBER_OF_DOCUMENTS_BRS_EVERY_WEEK", week)
    monkeypatch.setattr(weekly_module, "TASK_ORDER_VARIANTS", [[("Підрозділ 1", "Перший варіант")], [("Підрозділ 1", "Другий варіант")]])

    doc = Document()
    weekly_module.generate_main_combat_log_extract_war_every_week(doc, "01.07.2026")

    text = doc.tables[-1].cell(1, 1).text
    assert "4. Коротке викладення" in text
    assert "5. Рішення командира" in text
    assert "№87 від 01.07.2026, для виконання завдань" in text
    assert "Підрозділ 1. Перший варіант" in text
    assert "6. Відомості про виконання" in text
    assert "підрозділи здійснили" not in text


def test_weekly_extract_point_6_references_previous_week_order_when_new_bat_appears(monkeypatch):
    """07.07 - з'явився НОВИЙ тижневий БР - пункт 5 показує СЬОГОДНІШНЄ (07.07)
    рішення (свій варіант переліку завдань), пункт 6 - ВЖЕ ЗАВЕРШЕНЕ попереднє
    (01.07) розпорядження, зі СВОЇМ варіантом переліку завдань (а не сьогоднішнім) -
    так само, як і в щоденному витягу."""
    week = {"01.07.2026": {"бат": "87"}, "07.07.2026": {"бат": "102"}}
    monkeypatch.setattr(weekly_module, "get_number_br", _fake_get_number_br(week))
    monkeypatch.setattr(weekly_module, "NUMBER_OF_DOCUMENTS_BRS_EVERY_WEEK", week)
    monkeypatch.setattr(weekly_module, "TASK_ORDER_VARIANTS", [[("Підрозділ 1", "Перший варіант")], [("Підрозділ 1", "Другий варіант")]])

    doc = Document()
    weekly_module.generate_main_combat_log_extract_war_every_week(doc, "07.07.2026")

    text = doc.tables[-1].cell(1, 1).text
    assert "5. Рішення командира" in text
    assert "№102 від 07.07.2026, для виконання завдань" in text
    assert "Підрозділ 1. Другий варіант" in text
    assert "6. Відомості про виконання" in text
    assert "07.07.2026 відповідно до бойового розпорядження командира" in text
    assert "№87 від 01.07.2026" in text
    assert "Підрозділ 1. Перший варіант" in text


# -------------------------
# generate_documents_combat_log_extract_war_every_week (обгортка) /
# _tasks_text_for fallback
# -------------------------
def test_generate_documents_combat_log_extract_war_every_week_saves_full_document(tmp_path, monkeypatch):
    """Наскрізна перевірка самої обгортки (setup + основний вміст + збереження
    файлу) - раніше покривалась лише через реальні дані ОБЛІК.xlsx
    (test_integration.py), тепер - синтетичними, незалежно від того, який
    місяць зараз веде користувач у реальному файлі."""
    week = {"01.07.2026": {"бат": "87"}}
    monkeypatch.setattr(weekly_module, "get_number_br", _fake_get_number_br(week))
    monkeypatch.setattr(weekly_module, "NUMBER_OF_DOCUMENTS_BRS_EVERY_WEEK", week)
    monkeypatch.setattr(weekly_module, "TASK_ORDER_VARIANTS", [[("Підрозділ 1", "Варіант")]])
    monkeypatch.setattr(docx_utils_module, "OUTPUT_DIR_COMBAT_LOG_EXTRACT_WAR", str(tmp_path))

    commander_data = {"ЗВАННЯ": "майор", "ПІБ": "ПЕРШИЙ Перший Перший", "ТВО": False}
    weekly_module.generate_documents_combat_log_extract_war_every_week(Document(), "01.07.2026", commander_data)

    files = os.listdir(tmp_path)
    assert any("01.07.2026" in f for f in files)


def test_tasks_text_for_falls_back_to_first_variant_when_empty(monkeypatch):
    """Регресія: варіант переліку завдань, на який вказує певний variant_index,
    ще порожній у constants.py - використовується варіант №1
    (TASK_ORDER_VARIANTS[0]), а не порожній текст."""
    monkeypatch.setattr(weekly_module, "TASK_ORDER_VARIANTS", [[("Підрозділ 1", "Запасний варіант")], []])

    assert weekly_module._tasks_text_for(1) == "Підрозділ 1. Запасний варіант"


# -------------------------
# generate_main_combat_log_extract_war_every_day
# -------------------------
def test_daily_extract_point_6_references_real_previous_day_order(monkeypatch):
    """02.07 - точка 5 показує СЬОГОДНІШНЄ (02.07) розпорядження, точка 6 -
    ВЧОРАШНЄ (01.07), уже у виконаному часі."""
    monkeypatch.setattr(daily_module, "NUMBER_OF_DOCUMENTS_BRS_EVERY_DAY", {
        "01.07.2026": {"бат": "88", "брг": "1"},
        "02.07.2026": {"бат": "90", "брг": "1"},
    })

    doc = Document()
    daily_module.generate_main_combat_log_extract_war_every_day(doc, "02.07.2026", "СЬОГОДНІШНІЙ ЗМІСТ", "ВЧОРАШНІЙ ЗВІТ")

    text = doc.tables[-1].cell(1, 1).text
    assert "5. Рішення командира" in text
    assert "№90 від 02.07.2026, здійснити" in text
    assert "СЬОГОДНІШНІЙ ЗМІСТ" in text
    assert "6. Відомості про виконання" in text
    assert "№88 від 01.07.2026 підрозділи батальйону здійснили" in text
    assert "ВЧОРАШНІЙ ЗВІТ" in text


def test_daily_extract_point_6_has_no_reference_on_first_tracked_day(monkeypatch):
    """01.07 - немає жодного попереднього щоденного розпорядження (це перший
    день обліку) - пункт 6 лишається без конкретного посилання/змісту, а не
    показує помилково "None" чи посилання саме на себе."""
    monkeypatch.setattr(daily_module, "NUMBER_OF_DOCUMENTS_BRS_EVERY_DAY", {
        "01.07.2026": {"бат": "88", "брг": "1"},
    })

    doc = Document()
    daily_module.generate_main_combat_log_extract_war_every_day(doc, "01.07.2026", "СЬОГОДНІШНІЙ ЗМІСТ", "ЗВІТ ЯКИЙ НЕ МАЄ З'ЯВИТИСЬ")

    text = doc.tables[-1].cell(1, 1).text
    assert "№88 від 01.07.2026, здійснити" in text
    assert "6. Відомості про виконання" in text
    assert "ЗВІТ ЯКИЙ НЕ МАЄ З'ЯВИТИСЬ" not in text
    assert "None" not in text


def test_daily_extract_point_6_skips_placeholder_entries_without_bat(monkeypatch):
    """Записи з порожнім 'бат' (заглушки сітки дат, напр. кінець попереднього
    місяця) не рахуються "попереднім" розпорядженням для пункту 6."""
    monkeypatch.setattr(daily_module, "NUMBER_OF_DOCUMENTS_BRS_EVERY_DAY", {
        "29.06.2026": {"бат": "", "брг": ""},
        "30.06.2026": {"бат": "", "брг": ""},
        "01.07.2026": {"бат": "88", "брг": "1"},
    })

    doc = Document()
    daily_module.generate_main_combat_log_extract_war_every_day(doc, "01.07.2026", "СЬОГОДНІШНІЙ ЗМІСТ", "ЗВІТ ЯКИЙ НЕ МАЄ З'ЯВИТИСЬ")

    text = doc.tables[-1].cell(1, 1).text
    assert "ЗВІТ ЯКИЙ НЕ МАЄ З'ЯВИТИСЬ" not in text
    assert "від 30.06.2026 підрозділи батальйону здійснили" not in text
    assert "від 29.06.2026 підрозділи батальйону здійснили" not in text


# -------------------------
# bold_cell_heading_lines (пункти 4./5./6. як заголовки)
# -------------------------
def test_bold_cell_heading_lines_bolds_only_matching_paragraphs():
    from formatting.docx_utils import bold_cell_heading_lines

    doc = Document()
    table = doc.add_table(rows=1, cols=1)
    cell = table.cell(0, 0)
    cell.text = "4. Заголовок"
    cell.add_paragraph("Звичайний текст")
    cell.add_paragraph("5. Інший заголовок")

    bold_cell_heading_lines(cell, ("4.", "5.", "6."))

    assert cell.paragraphs[0].runs[0].font.bold is True
    assert cell.paragraphs[1].runs[0].font.bold is not True
    assert cell.paragraphs[2].runs[0].font.bold is True


def test_daily_extract_bolds_points_4_5_6_headings(monkeypatch):
    monkeypatch.setattr(daily_module, "NUMBER_OF_DOCUMENTS_BRS_EVERY_DAY", {
        "01.07.2026": {"бат": "88", "брг": "1"},
    })

    doc = Document()
    daily_module.generate_main_combat_log_extract_war_every_day(doc, "01.07.2026", "ЗМІСТ", "ЗВІТ")

    cell = doc.tables[-1].cell(1, 1)
    heading_paragraphs = [p for p in cell.paragraphs if p.text.startswith(("4.", "5.", "6."))]
    assert len(heading_paragraphs) == 3
    for paragraph in heading_paragraphs:
        assert all(run.font.bold for run in paragraph.runs)


def test_weekly_extract_bolds_points_4_5_6_headings(monkeypatch):
    week = {"01.07.2026": {"бат": "87"}}
    monkeypatch.setattr(weekly_module, "get_number_br", _fake_get_number_br(week))
    monkeypatch.setattr(weekly_module, "NUMBER_OF_DOCUMENTS_BRS_EVERY_WEEK", week)
    monkeypatch.setattr(weekly_module, "TASK_ORDER_VARIANTS", [[("Підрозділ 1", "Варіант")]])

    doc = Document()
    weekly_module.generate_main_combat_log_extract_war_every_week(doc, "01.07.2026")

    cell = doc.tables[-1].cell(1, 1)
    heading_paragraphs = [p for p in cell.paragraphs if p.text.startswith(("4.", "5.", "6."))]
    assert len(heading_paragraphs) == 3
    for paragraph in heading_paragraphs:
        assert all(run.font.bold for run in paragraph.runs)
