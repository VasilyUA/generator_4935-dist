from docx import Document

def read_docx_text(file_path: str) -> str:
    """Зчитує весь текст з Word-документу."""
    doc = Document(file_path)
    full_text = []
    for paragraph in doc.paragraphs:
        full_text.append(paragraph.text)
    # якщо у тебе таблиці — теж додаємо їхній текст
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                full_text.append(cell.text)
    return "\n".join(full_text)
