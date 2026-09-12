import copy
import glob
import os
import re
from datetime import datetime

from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from utils.logging_utils import print_green, print_red
from utils.excel_writer import save_workbook_safely

# Тека з файлами для звірки (ОБЛІК1.xlsx, ОБЛІК2.xlsx, ... - будь-яка кількість) і
# шлях до результату - ті самі "resources"/"output", що й в решті проєкту.
CHECK_DIR = os.path.join("resources", "check")
RESULT_FILE_NAME = os.path.join("output", "RESULT_accounting.xlsx")

GREEN_FILL = PatternFill(start_color="FFC6EFCE", end_color="FFC6EFCE", fill_type="solid")
RED_FILL = PatternFill(start_color="FFFFC7CE", end_color="FFFFC7CE", fill_type="solid")

# Порядок "іменованих" колонок у результаті - ОБЛІК.xlsx-подібні файли можуть мати
# їх у РІЗНОМУ порядку/позиції (напр. з додатковою колонкою на початку) - визначаються
# за НАЗВОЮ заголовка, а не за фіксованою літерою.
_LABEL_COLUMNS = ("ПІДРОЗДІЛ", "ПОСАДА", "ЗВАННЯ", "ПІБ")

# Розширення, які openpyxl.load_workbook вміє читати напряму (той самий ZIP+XML
# формат, macro чи ні) - .xls (старий бінарний формат) свідомо НЕ підтримується тут,
# бо потребував би окремого читання через xlrd, як у utils/excel_reader.py.
_SUPPORTED_EXTENSIONS = (".xlsx", ".xlsm")


def _find_check_files(check_dir):
    """Відсортований (за іменем файлу) список .xlsx/.xlsm у check_dir - порядок
    визначає, яка з N-рядкової групи людини буде "першою" (файл-джерело стилю та
    порядку появи нових людей у результаті)."""
    file_paths = sorted(
        path for path in glob.glob(os.path.join(check_dir, "*"))
        if os.path.splitext(path)[1].lower() in _SUPPORTED_EXTENSIONS
    )
    if not file_paths:
        raise FileNotFoundError(f"У теці {check_dir} не знайдено жодного .xlsx/.xlsm файлу для звірки.")
    return file_paths


def _read_check_file(file_path, extra_label_columns=()):
    """Читає ОДИН файл (перший аркуш) - колонки ПІДРОЗДІЛ/ПОСАДА/ЗВАННЯ/ПІБ визначаються
    за НАЗВОЮ заголовка, дати - за ТИПОМ значення заголовка (datetime) - працює навіть
    якщо порядок/кількість колонок (чи порожні "технічні" колонки між ними) різняться
    між файлами check_dir.

    extra_label_columns - додаткові назви заголовків (у ВЕРХНЬОМУ регістрі), що
    зберігаються в "label" ТАК САМО, як _LABEL_COLUMNS, але НЕ входять у сам
    _LABEL_COLUMNS - той список визначає ще й структуру результату "Звірка"
    (_build_check_result нижче: об'єднані колонки/заголовки), який має лишитись
    БЕЗ ЗМІН для звичайного виклику без цього параметра. content.report_changes
    (той самий парсер для resources/changes/prev_MM_ОБЛІК.xlsx/actual_MM_ОБЛІК.xlsx)
    передає сюди ("ПІДСТАВИ", "ПІДСТАВА") - інакше ця колонка мовчки губилась би
    (не дата й не в _LABEL_COLUMNS), і build_pidstavy_extra_grounds_by_person
    (content.money_report_helpers) завжди отримувала б порожній результат.

    Повертає (people, date_columns): people - список {"pib_raw", "label": {назва: значення},
    "days": {date: значення}} (рядки з порожнім ПІБ пропускаються); date_columns - множина
    дат (datetime.date), знайдених у ЦЬОМУ файлі."""
    wb = load_workbook(file_path, data_only=True)
    ws = wb[wb.sheetnames[0]]
    header = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))

    label_columns = _LABEL_COLUMNS + tuple(extra_label_columns)
    label_col_index = {}
    date_col_index = {}
    for idx, value in enumerate(header):
        if isinstance(value, datetime):
            date_col_index[value.date()] = idx
        elif isinstance(value, str):
            # " ".join(value.split()) - схлопує БУДЬ-ЯКИЙ пробільний символ
            # (не лише краї, як .strip()) в один пробіл - реальний ОБЛІК.xlsx
            # переносить довгі заголовки на 2 рядки (напр. комірка буквально
            # містить "ПІДСТАВИ\nДЛЯ ІНШИХ СИТУАЦІЙ") - без цього порівняння з
            # label_columns (усі - однорядкові літерали) мовчки не збігалось би.
            normalized = " ".join(value.split()).upper()
            if normalized in label_columns:
                label_col_index[normalized] = idx

    if "ПІБ" not in label_col_index:
        raise ValueError(f"У файлі {file_path} не знайдено колонку 'ПІБ' у заголовку.")

    people = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        pib_col = label_col_index["ПІБ"]
        pib_raw = row[pib_col] if pib_col < len(row) else None
        if not pib_raw:
            continue
        people.append({
            "pib_raw": pib_raw,
            "label": {name: (row[idx] if idx < len(row) else None) for name, idx in label_col_index.items()},
            "days": {date: (row[idx] if idx < len(row) else None) for date, idx in date_col_index.items()},
        })

    return people, set(date_col_index)


def _extract_style_templates(file_path):
    """Бере СТИЛЬ (шрифт/заливка/рамка/вирівнювання, ширина колонок, висота рядків)
    заголовка й рядка даних із file_path (перший файл check_dir) - як шаблон для
    результату, щоб RESULT_accounting.xlsx виглядав так само, як вихідні ОБЛІК*.xlsx."""
    wb = load_workbook(file_path)
    ws = wb[wb.sheetnames[0]]
    header = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))

    label_col = next((i for i, v in enumerate(header) if isinstance(v, str) and v.strip().upper() in _LABEL_COLUMNS), None)
    date_col = next((i for i, v in enumerate(header) if isinstance(v, datetime)), None)
    if label_col is None or date_col is None:
        raise ValueError(f"Не вдалось визначити стиль заголовка (ПІБ чи дати) з {file_path}.")

    return {
        "header_label": ws.cell(row=1, column=label_col + 1),
        "header_date": ws.cell(row=1, column=date_col + 1),
        "body_label": ws.cell(row=2, column=label_col + 1),
        "body_date": ws.cell(row=2, column=date_col + 1),
        "header_row_height": ws.row_dimensions[1].height,
        "body_row_height": ws.row_dimensions[2].height,
        "label_col_width": ws.column_dimensions[get_column_letter(label_col + 1)].width,
        "date_col_width": ws.column_dimensions[get_column_letter(date_col + 1)].width,
    }


def _normalize_name(name):
    """Нормалізує ПІБ для зіставлення між файлами check_dir (пробіли, апостроф,
    регістр) - той самий підхід, що й content.money_report_helpers.normalize_name,
    але СВОЯ, локальна копія: імпорт того модуля тягне за собою constants.py, а
    той - інтерактивні запити (вибір місяця тощо) вже при самому імпорті, що
    зайве для цього окремого, самодостатнього інструменту звірки."""
    if not isinstance(name, str):
        return ""
    cleaned = re.sub(r"(В|в)'[\s]+([а-яА-ЯІіЇїЄєҐґ])", r"\1'\2", name.strip())
    return " ".join(cleaned.split()).upper()


def _style_bundle(source_cell):
    """Знімає копію стилю source_cell ОДИН РАЗ (викликається лише для 4 шаблонних
    клітинок - header_label/header_date/body_label/body_date - а не для кожної з
    тисяч клітинок результату): Font/Fill/Border/Alignment openpyxl трактує як
    незмінні значення, тож той самий об'єкт можна безпечно призначити багатьом
    клітинкам - єдина подальша зміна (заливка при розбіжності, нижче) ЗАМІНЮЄ
    посилання на клітинці, а не мутує сам об'єкт стилю."""
    return {
        "font": copy.copy(source_cell.font),
        "fill": copy.copy(source_cell.fill),
        "border": copy.copy(source_cell.border),
        "alignment": copy.copy(source_cell.alignment),
        "number_format": source_cell.number_format,
    }


def _apply_style(target_cell, bundle):
    target_cell.font = bundle["font"]
    target_cell.fill = bundle["fill"]
    target_cell.border = bundle["border"]
    target_cell.alignment = bundle["alignment"]
    target_cell.number_format = bundle["number_format"]


def _normalize_for_compare(value):
    """Порожня клітинка (None) і рядок зводяться до одного вигляду (обрізані пробіли,
    верхній регістр), щоб "100" в одному файлі й "100 " чи "Вп"/"ВП" в іншому не
    вважались розбіжністю там, де її фактично немає. Числа лишаються як є."""
    if value is None:
        return ""
    if isinstance(value, str):
        return value.strip().upper()
    return value


def _values_match(values):
    return len({_normalize_for_compare(v) for v in values}) <= 1


def _format_day_value(value):
    """Нормалізоване текстове представлення значення комірки дати для колонки
    "Підсумок" (нижче) - число БЕЗ зайвого ".0" (openpyxl часто зберігає цілі
    числа як float), текст - у ВЕРХНЬОМУ регістрі й без зайвих пробілів (той
    самий підхід, що й _normalize_for_compare вище) - щоб "Вп"/"ВП" рахувались
    як ОДИН статус, а не два окремих рядки підсумку."""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, str):
        return " ".join(value.split()).upper()
    return str(value)


def _row_day_value_summary(day_values):
    """"100 - 15 днів; 30 - 5 днів; ВП - 5 днів; ..." - скільки днів МАЄ КОЖНЕ
    фактично наявне значення серед day_values ЦЬОГО РЯДКА (одного файлу-
    джерела для однієї людини - рахується ОКРЕМО для кожного рядка, а не
    сумарно по всіх файлах людини відразу) - підтверджено користувачем:
    "трапа кількість днів для кожного статусу... наприклад ВП - 5 днів" -
    ДИНАМІЧНИЙ перелік УСІХ значень, що трапляються (не фіксований список
    100/30/70/170 - це були лише приклади). Порожні комірки НЕ рахуються
    взагалі (ні як окрема категорія, ні в жодному лічильнику). Порядок - за
    ПЕРШОЮ появою значення серед дат (зліва направо, тобто хронологічно)."""
    counts = {}
    order = []
    for raw_value in day_values:
        if raw_value is None or raw_value == "":
            continue
        key = _format_day_value(raw_value)
        if key not in counts:
            counts[key] = 0
            order.append(key)
        counts[key] += 1
    return "; ".join(f"{key} - {counts[key]} днів" for key in order)


def _people_order_and_index(per_file_people):
    """people_order - список нормалізованих ПІБ у порядку першої появи (файл за
    файлом, рядок за рядком); per_file_by_pib - для кожного файлу {normalized_pib:
    person}, щоб потім легко перевірити, чи є (і які) дані про людину в кожному файлі."""
    people_order = []
    seen = set()
    per_file_by_pib = []
    for people in per_file_people:
        by_pib = {}
        for person in people:
            normalized = _normalize_name(person["pib_raw"])
            if normalized not in seen:
                seen.add(normalized)
                people_order.append(normalized)
            by_pib[normalized] = person
        per_file_by_pib.append(by_pib)
    return people_order, per_file_by_pib


def check_accounting(check_dir=CHECK_DIR, output_path=RESULT_FILE_NAME, file_paths=None):
    """Звіряє всі .xlsx/.xlsm файли з check_dir (кожен - структура на кшталт ОБЛІК.xlsx,
    можливо з іншим порядком/кількістю колонок) і зберігає результат у output_path.

    file_paths - якщо передано (не None), звіряються РІВНО ЦІ файли (у цьому порядку),
    а check_dir взагалі не сканується - потрібно, коли в одній директорії лежить кілька
    НЕЗАЛЕЖНИХ груп файлів для звірки (напр. resources/changes/ з кількома парами
    prev_MM_ОБЛІК/actual_MM_ОБЛІК за різні місяці) і звичайний глоб усієї директорії
    змішав би їх в одну звірку.

    Для кожної людини - N рядків (N = кількість файлів у check_dir), по одному на
    файл (перша колонка результату - "Файл", ім'я джерела цього рядка). Кожна колонка
    з _LABEL_COLUMNS (ПІДРОЗДІЛ/ПОСАДА/ЗВАННЯ/ПІБ) - ОДНА клітинка, об'єднана на всі
    N рядків цієї людини (не дублюється рядок у рядок) - показує канонічне значення з
    ПЕРШОГО файлу, де ця КОЛОНКА фактично має значення (не просто де людину знайдено -
    файл може містити людину, але взагалі не мати такої колонки в заголовку, як інший
    файл з ІНШИМ набором колонок). Кожна дата - зелена заливка, якщо
    значення в УСІХ N рядках людини однакові, інакше - червона (в усіх рядках цієї
    колонки для цієї людини); так само, якщо людина відсутня в одному з файлів, це
    теж проявляється як розбіжність (порожнє значення проти заповненого). ПІДРОЗДІЛ/
    ПОСАДА/ЗВАННЯ фарбуються так само - за ФАКТИЧНИМИ (за кожним файлом) значеннями,
    а не за вже об'єднаним канонічним, інакше об'єднання саме по собі приховало б
    розбіжність. ПІБ (сам ключ зіставлення) не фарбується. Стиль заголовка й клітинок
    копіюється з ПЕРШОГО (за іменем) файлу check_dir.

    "Підсумок" - ОСТАННЯ колонка, ПІСЛЯ всіх дат (_row_day_value_summary,
    підтверджено користувачем: "в кінці таблички після кінця місяця") - для
    КОЖНОГО рядка (файлу-джерела) окремо, текст на кшталт "100 - 15 днів; 30 -
    5 днів; ВП - 5 днів" - скільки днів МАЄ КОЖНЕ фактично наявне значення в
    ЦЬОМУ рядку (ДИНАМІЧНИЙ перелік, не фіксований список - будь-яке значення,
    що трапляється, отримує свій запис). Фарбується ТІЄЮ САМОЮ логікою, що й
    дати (зелена/червона за збігом між рядками людини) - природно узгоджено з
    рештою: якщо самі дати розходяться між файлами, розійдеться й підсумок.

    Повертає шлях до збереженого файлу, або None, якщо file_paths не передано і в
    check_dir немає жодного .xlsx/.xlsm файлу (не падає - лише друкує попередження,
    як і решта опціональних перевірок у проєкті)."""
    if file_paths is None:
        try:
            file_paths = _find_check_files(check_dir)
        except FileNotFoundError as error:
            print_red(str(error))
            return None

    file_names = [os.path.basename(path) for path in file_paths]
    parsed_files = [_read_check_file(path) for path in file_paths]
    per_file_people = [people for people, _ in parsed_files]
    all_dates = sorted({date for _, date_columns in parsed_files for date in date_columns})

    people_order, per_file_by_pib = _people_order_and_index(per_file_people)
    styles = _extract_style_templates(file_paths[0])
    # Стиль знімається один раз на роль клітинки (4 рази), а не для кожної з тисяч
    # клітинок результату - див. _style_bundle.
    header_label_style = _style_bundle(styles["header_label"])
    header_date_style = _style_bundle(styles["header_date"])
    body_label_style = _style_bundle(styles["body_label"])
    body_date_style = _style_bundle(styles["body_date"])

    wb = Workbook()
    ws = wb.active
    assert isinstance(ws, Worksheet)  # завжди так для щойно створеного Workbook()
    ws.title = "Звірка"

    # "Підсумок" - ОСТАННЯ колонка, ПІСЛЯ всіх дат (_row_day_value_summary) -
    # підтверджено користувачем: "в кінці таблички після кінця місяця".
    headers = ["Файл"] + list(_LABEL_COLUMNS) + [datetime.combine(date, datetime.min.time()) for date in all_dates] + ["Підсумок"]
    pib_col = 1 + len(_LABEL_COLUMNS)
    summary_col = len(headers)
    # {"ПІДРОЗДІЛ": 2, "ПОСАДА": 3, "ЗВАННЯ": 4, "ПІБ": 5} - колонка кожної з
    # _LABEL_COLUMNS у результаті (усі вони об'єднуються на N рядків людини).
    label_col_by_name = {name: idx for idx, name in enumerate(_LABEL_COLUMNS, start=2)}

    for col_idx, value in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx, value=value)
        _apply_style(cell, header_date_style if isinstance(value, datetime) else header_label_style)
    ws.row_dimensions[1].height = styles["header_row_height"]
    ws.freeze_panes = "A2"

    current_row = 2
    for pib in people_order:
        start_row = current_row
        # Канонічне значення ПІДРОЗДІЛ/ПОСАДА/ЗВАННЯ - з ПЕРШОГО файлу, де САМЕ ЦЯ
        # колонка має значення (не просто де людину знайдено) - інакше файл без такої
        # колонки в заголовку (напр. коротший "prev"-зріз) "затер" би канонічне
        # значення пусткою, навіть коли інший файл цю людину й колонку має.
        canonical = {
            name: next(
                (
                    per_file_by_pib[i][pib]["label"].get(name)
                    for i in range(len(file_paths))
                    if pib in per_file_by_pib[i] and per_file_by_pib[i][pib]["label"].get(name)
                ),
                None,
            )
            for name in _LABEL_COLUMNS if name != "ПІБ"
        }
        canonical["ПІБ"] = next((per_file_by_pib[i][pib]["pib_raw"] for i in range(len(file_paths)) if pib in per_file_by_pib[i]), pib)

        # ФАКТИЧНІ (за кожним файлом) значення ПІДРОЗДІЛ/ПОСАДА/ЗВАННЯ - для порівняння;
        # об'єднана клітинка показує лише канонічне значення, тож розбіжність треба
        # зафіксувати ДО того, як рядки перетворяться на однакові.
        actual_labels_by_file = []

        for file_index, file_name in enumerate(file_names):
            person = per_file_by_pib[file_index].get(pib)
            actual_labels_by_file.append({name: (person["label"].get(name) if person else None) for name in _LABEL_COLUMNS if name != "ПІБ"})

            day_values = [(person["days"].get(date) if person else None) for date in all_dates]
            row_values = (
                [file_name]
                + [canonical[name] for name in _LABEL_COLUMNS]
                + day_values
                + [_row_day_value_summary(day_values)]
            )

            for col_idx, value in enumerate(row_values, start=1):
                cell = ws.cell(row=current_row, column=col_idx, value=value)
                is_date_cell = pib_col < col_idx < summary_col
                _apply_style(cell, body_date_style if is_date_cell else body_label_style)
            ws.row_dimensions[current_row].height = styles["body_row_height"]
            current_row += 1

        end_row = current_row - 1
        if end_row > start_row:
            for col_idx in range(2, pib_col + 1):
                ws.merge_cells(start_row=start_row, start_column=col_idx, end_row=end_row, end_column=col_idx)

        for name, col_idx in label_col_by_name.items():
            if name == "ПІБ":
                continue
            values = [labels[name] for labels in actual_labels_by_file]
            fill = GREEN_FILL if _values_match(values) else RED_FILL
            # Клітинка об'єднана (start_row..end_row) - Excel (і сам openpyxl після
            # збереження/перевідкриття) бере стиль лише з "якірної" (верхньої лівої)
            # клітинки діапазону; стиль інших клітинок діапазону не зберігається.
            ws.cell(row=start_row, column=col_idx).fill = fill

        for col_idx in range(pib_col + 1, len(headers) + 1):
            values = [ws.cell(row=r, column=col_idx).value for r in range(start_row, end_row + 1)]
            fill = GREEN_FILL if _values_match(values) else RED_FILL
            for r in range(start_row, end_row + 1):
                ws.cell(row=r, column=col_idx).fill = fill

    file_col_width = max(len(name) for name in file_names) + 2
    # "Підсумок" - текстовий рядок на кшталт "100 - 15 днів; 30 - 5 днів; ..." -
    # значно ДОВШИЙ за звичайну дату, тож ширина колонки-дати (styles["date_col_width"])
    # тут була б замалою - беремо ЩОНАЙМЕНШЕ 40 символів (чи ширину дати, якщо
    # вона й так більша - реальний шаблон-файл міг мати нетипово широкі колонки).
    summary_col_width = max(styles["date_col_width"] or 0, 40)
    for col_idx in range(1, len(headers) + 1):
        letter = get_column_letter(col_idx)
        if col_idx == 1:
            ws.column_dimensions[letter].width = file_col_width
        elif col_idx <= pib_col:
            ws.column_dimensions[letter].width = styles["label_col_width"]
        elif col_idx == summary_col:
            ws.column_dimensions[letter].width = summary_col_width
        else:
            ws.column_dimensions[letter].width = styles["date_col_width"]

    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    save_workbook_safely(wb, output_path)
    print_green(os.path.basename(output_path))
    return output_path


if __name__ == "__main__":  # pragma: no cover
    check_accounting()
