import calendar
import re
from datetime import datetime, timedelta

from docx import Document
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P
from docx.table import Table
from docx.text.paragraph import Paragraph

from content.money_report_helpers import normalize_name

# Позиції колонок за замовчуванням (0-відлік) - ПОСАДА/ЗВАННЯ/ПІБ/ПЕРІОД/ДНІ/
# ПІДСТАВА, як у звичайній таблиці рапорту (formatting/docx_utils.create_table,
# TABLE_HEADERS/NOT_PAID_TABLE_HEADERS у generators/generate_report_for_get_money.py).
# Не всі пункти рапорту мають РІВНО такий набір колонок - напр. "СПЕЦКОНТИНГЕНТ"
# (полонені/зниклі безвісти) має ДОДАТКОВУ колонку "Дата зникнення безвісти" ПЕРЕД
# періодом, зсуваючи сам період на позицію 5 замість 4 (і решту колонок далі) -
# тож ці позиції лише ЗАПАСНИЙ варіант, якщо _detect_columns (за назвою
# заголовка, нижче) нічого не знайшла.
_POSADA_COL, _ZVANNYA_COL, _PIB_COL, _PERIOD_COL, _DAYS_COL, _BASIS_COL = 1, 2, 3, 4, 5, 6

# Ключові слова заголовків колонок (без урахування регістру) - для визначення
# ФАКТИЧНОЇ позиції ПІБ/ПЕРІОД/ДНІВ/ПІДСТАВИ у КОЖНІЙ конкретній таблиці окремо
# (а не одного фіксованого номера на весь документ), бо позиція періоду
# відрізняється для "СПЕЦКОНТИНГЕНТ" (див. вище). ПОСАДА/ЗВАННЯ так само за
# назвою - для узгодженості й на випадок майбутніх пунктів із власними
# додатковими колонками. "disappearance_date" ("Дата зникнення безвісти") - БЕЗ
# запасної позиції в _detect_columns нижче: ця колонка є ЛИШЕ в таблиці
# "СПЕЦКОНТИНГЕНТ", немає сенсу вигадувати для неї позицію "на весь випадок" в
# звичайній таблиці, де такої колонки просто нема.
_COLUMN_HEADER_KEYWORDS = {
    "posada": "посада", "zvannya": "звання", "pib": "прізвище", "period": "період",
    "days": "днів", "basis": "підстав", "disappearance_date": "зникненн",
}

# "СПЕЦКОНТИНГЕНТ" (полон/зниклі безвісти/заручники) - єдиний пункт, чия таблиця
# рапорту має інакший набір колонок (див. вище) і чию "ДНІ"/"Дата зникнення
# безвісти" content/report_document_reader.py та checker_accounting/report_checker.py
# обробляють окремо від решти пунктів (підтверджено користувачем) - назва тут
# ОДНА, а не дублюється рядковим літералом у кожному місці використання.
SPETSKONTYNGENT_VALUE = "100_СПЕЦКОНТИНГЕНТ"

# Пункти, чия "Примітка (підстави)" відтворюється в табелі (extract_timetable_rows_
# from_report - колонка "ПІДСТАВИ"): СПЕЦКОНТИНГЕНТ (як і раніше), а також 100_ВП/
# 100_ШП (лікування/поранення) і NOT_PAID (СЗЧ/Задув.) - підтверджено користувачем.
# Решта пунктів (звичайні 30/70/100/170/...) підстави в табелі не мають взагалі.
_BASIS_TRACKED_VALUES = (SPETSKONTYNGENT_VALUE, "100_ВП", "100_ШП", "NOT_PAID")

# Підпункти розділу "Прошу внести зміни..." нумеруються "N.M текст" (без пробілу
# між номером пункту й підпункту - _add_changes_subsection), звичайні ж пункти
# рапорту - "N. текст" (з пробілом одразу після крапки - _add_category_section).
# Ця різниця у форматуванні номера - єдина СТАБІЛЬНА ознака, яка відрізняє
# "зміни за попередні місяці" (не стосуються поточного місяця - табель на них
# не зважає) від звичайних пунктів, незалежна від того, як саме сформульовано
# сам текст пункту (текст пунктів роками редагувався, тож зіставляти рапорт
# довільного (можливо, старого) місяця з ТОЧНИМ поточним текстом STATIK -
# ненадійно).
_CHANGES_SUBSECTION_RE = re.compile(r"^\s*\d+\.\d+\s")
_LEADING_NUMBER_RE = re.compile(r"^\s*(\d+)\.\s")
_MONEY_AMOUNT_RE = re.compile(r"(\d[\d\s]*\d|\d)\s*000\s*грн")
_NOT_PAID_KEYWORD = "не виплачувати"

# Ключові слова замість зіставлення з ТОЧНИМ поточним текстом
# resources/data.json/constants.py - новий пункт (напр. "100_ВП"/"100_СПЕЦКОНТИНГЕНТ",
# додані користувачем) підхоплюється табелем АВТОМАТИЧНО, без жодних змін тут, доки
# його юридичний текст згадує ЦІ Ж стабільні (законодавчо усталені, не зафромульовані
# в кожному пункті наново) поняття - незалежно від того, як САМЕ пункт названо чи
# структуровано в MONEY_REPORT_CATEGORIES/SECTIONS (звідти цей модуль НІЧОГО не
# читає взагалі - лише сам текст рапорту).
# "стаціонарне лікування"/"лікарняний заклад" - госпіталізація (напр. "100_ШП") -
# на відміну від "відпустці ДЛЯ ЛІКУВАННЯ" (людина НЕ в закладі, а у відпустці,
# напр. "100_ВП") - обидва варіанти рясно згадують спільні слова "лікування"/
# "поранення" (звідси й окремі, вужчі перевірки нижче ПЕРЕД загальним "лікува|поранен"
# fallback, а не одна спільна перевірка - інакше "100_ВП" завжди хибно збігався б
# із "100_ШП").
_HOSPITAL_KEYWORD_RE = re.compile(r"стаціонарн|лікарня", re.IGNORECASE)
_MEDICAL_VACATION_KEYWORD_RE = re.compile(r"відпустці для лікування|тяжкого поранення", re.IGNORECASE)
# Загальний fallback - будь-яка ІНША згадка лікування/поранення (напр. майбутній
# пункт із формулюванням, ще не описаним двома вужчими перевірками вище) - все одно
# позначається як "100_ШП", а не губиться зовсім.
_MEDICAL_LEAVE_KEYWORD_RE = re.compile(r"лікува|поранен", re.IGNORECASE)
# "полон"/"безвісн"/"заручник" - полонені, зниклі безвісти, заручники (спецконтингент).
_MISSING_CAPTURED_KEYWORD_RE = re.compile(r"полон|безвісн|заручник", re.IGNORECASE)

_PERIOD_RANGE_RE = re.compile(r"(\d{2}\.\d{2}\.\d{4})\s*-\s*(\d{2}\.\d{2}\.\d{4})")

# "100" рахує ще й дні 70/170 (_COMBINED_TARGET_CELL_VALUES у
# content/money_report_helpers.py) - той самий день людини тому може з'явитись
# ОДНОЧАСНО і в її власній таблиці "70"/"170" (справжнє значення), і, ЗАЙВИМ
# рядком, у таблиці "100" (де підстава вже про пункт 100, а не 70/170). Щоб не
# затерти справжнє 70/170 значення дня генеричним "100", значення з таблиці
# "100" застосовуються ОСТАННІМИ і лише туди, де ще нічого не записано.
_DEFERRED_RAW_VALUE = 100


def _iter_block_items(doc):
    """Абзаци й таблиці документа В ПОРЯДКУ ЇХ ПОЯВИ (doc.paragraphs/doc.tables
    окремо не зберігають взаємний порядок - потрібен спільний прохід по
    doc.element.body, стандартний рецепт python-docx)."""
    for child in doc.element.body.iterchildren():
        if isinstance(child, CT_P):
            yield Paragraph(child, doc)
        elif isinstance(child, CT_Tbl):
            yield Table(child, doc)


def _is_heading(block):
    """Абзац - "заголовок пункту" (оновлює current_heading), якщо АБО стиль
    явно містить "heading" (документи, згенеровані doc.add_heading() -
    напр. тестові фікстури tests/test_report_document_reader.py), АБО текст
    сам починається з провідного номера пункту "N. "/"N.M " - _LEADING_NUMBER_RE/
    _CHANGES_SUBSECTION_RE. Другий варіант потрібен для РЕАЛЬНИХ рапортів,
    підготовлених вручну у Word: підтверджено на живому зразку
    resources/ДВ_зі_змінами_1бмп_ТРАВЕНЬ_v2.3.docx - усі пункти рапорту
    (включно з підпунктами "Прошу внести зміни...") там - звичайні
    "Normal"-абзаци з номером прямо в тексті, БЕЗ стилю Word "Заголовок" - без
    цієї гілки current_heading лишався б порожнім на весь документ, і КОЖЕН
    рядок отримував би raw_value/НОМЕР_ПУНКТУ "?" (саме такий "?" і поскаржився
    користувач у колонці "Пункт рапорту")."""
    if not isinstance(block, Paragraph):
        return False
    if block.style is not None and "heading" in block.style.name.lower():
        return True
    text = block.text
    return bool(_LEADING_NUMBER_RE.match(text) or _CHANGES_SUBSECTION_RE.match(text))


def _raw_value_for_heading(heading_text):
    """Визначає "сире" значення (як у комірці дня ОБЛІК.xlsx) для таблиці під
    заголовком heading_text - None, якщо це підпункт "Прошу внести зміни..."
    (не звичайний пункт поточного місяця, ігнорується табелем).

    Текст самих пунктів роками редагувався (порівняно зі старим зразком
    resources/РАПОРТ.docx: "у розмірі"/"в розмірі", "К-ь днів"/"Кількість днів"
    тощо), і НОВІ пункти (напр. "100_ВП"/"100_СПЕЦКОНТИНГЕНТ") можуть додаватись у
    constants.py/resources/data.json в будь-який момент - тож зіставлення НІЯК не
    звіряється з ЦИМИ файлами (цей модуль їх узагалі не читає), а спирається лише
    на стабільні, законодавчо усталені орієнтири в самому тексті: "не виплачувати"
    (NOT_PAID); "стаціонарне лікування"/"лікарняний заклад" - госпіталізація
    ("100_ШП"); "відпустці для лікування"/"тяжкого поранення" - людина НЕ в
    закладі, а у відпустці для лікування ("100_ВП"); будь-яка ІНША згадка
    лікування/поранення, що не підійшла під ці дві вужчі перевірки - також
    "100_ШП" (запасний варіант, не губити рядок зовсім); "полон"/"безвісти"/
    "заручник" (спецконтингент); і, нарешті, сума "NNN 000 грн" (сам номер
    пункту). Новий пункт, чий текст згадує один з ЦИХ орієнтирів, підхоплюється
    автоматично, без жодних змін коду.

    Пункт без жодного з цих орієнтирів (напр. дійсно НЕВІДОМИЙ, ще не описаний тут
    вид пункту) - лишає ПОЗНАЧКУ "п.N" (N - номер пункту з самого заголовка), а не
    помилку - людина в табелі краще з приблизною позначкою, ніж не потрапить у
    нього взагалі."""
    if _CHANGES_SUBSECTION_RE.match(heading_text):
        return None

    lowered = heading_text.lower()
    if _NOT_PAID_KEYWORD in lowered:
        return "NOT_PAID"
    if _HOSPITAL_KEYWORD_RE.search(heading_text):
        return "100_ШП"
    if _MEDICAL_VACATION_KEYWORD_RE.search(heading_text):
        return "100_ВП"
    if _MISSING_CAPTURED_KEYWORD_RE.search(heading_text):
        return SPETSKONTYNGENT_VALUE
    if _MEDICAL_LEAVE_KEYWORD_RE.search(heading_text):
        return "100_ШП"

    money_match = _MONEY_AMOUNT_RE.search(heading_text)
    if money_match:
        return int(money_match.group(1).replace(" ", ""))

    number_match = _LEADING_NUMBER_RE.match(heading_text)
    return f"п.{number_match.group(1)}" if number_match else "?"


def _detect_columns(header_row):
    """{posada,zvannya,pib,period,days,basis: індекс колонки; disappearance_date:
    індекс колонки чи None} для ОДНІЄЇ конкретної таблиці - за назвою заголовка
    (_COLUMN_HEADER_KEYWORDS), а не фіксованою позицією: пункт "СПЕЦКОНТИНГЕНТ"
    (полонені/зниклі безвісти) має додаткову колонку "Дата зникнення безвісти"
    ПЕРЕД періодом, зсуваючи сам період (і решту колонок) на іншу позицію -
    фіксовані _POSADA_COL/etc. підійшли б лише для "звичайних" таблиць.
    Колонку з відповідним ключовим словом не знайдено (напр. дуже старий чи
    пошкоджений заголовок) - лишається запасна фіксована позиція (крім
    "disappearance_date" - для неї запасної позиції нема, None)."""
    header_texts = [cell.text.strip().lower() for cell in header_row.cells]
    indices = {
        "posada": _POSADA_COL, "zvannya": _ZVANNYA_COL, "pib": _PIB_COL, "period": _PERIOD_COL,
        "days": _DAYS_COL, "basis": _BASIS_COL, "disappearance_date": None,
    }
    for key, keyword in _COLUMN_HEADER_KEYWORDS.items():
        match = next((idx for idx, text in enumerate(header_texts) if keyword in text), None)
        if match is not None:
            indices[key] = match
    return indices


def _date_subranges_from_period_text(period_text):
    """Як _dates_from_period_text (нижче), але зберігає МЕЖІ кожного під-діапазону
    (розділеного "; " у самому тексті ПЕРІОДУ, напр. "08.05.2026-24.05.2026;
    28.05.2026-31.05.2026" - _PERIOD_RANGE_RE шукає пари дат незалежно від
    роздільника, тож "; " саме по собі тут не парситься окремо, лише природно
    розділяє сусідні пари) - [(start_date, [дата, дата, ...]), ...]. Потрібно
    checker_accounting.report_log_war_checker, де ПЕРШИЙ день КОЖНОГО під-
    діапазону (а не лише перший день усього періоду) - "перехідний" (людина
    щойно почала цю ділянку участі), решта днів під-діапазону - "звичайні"."""
    subranges = []
    for start_str, end_str in _PERIOD_RANGE_RE.findall(period_text or ""):
        start, end = datetime.strptime(start_str, "%d.%m.%Y"), datetime.strptime(end_str, "%d.%m.%Y")
        dates, current = [], start
        while current <= end:
            dates.append(current)
            current += timedelta(days=1)
        subranges.append((start, dates))
    return subranges


def _dates_from_period_text(period_text):
    return [date for _start, dates in _date_subranges_from_period_text(period_text) for date in dates]


def _iter_recognized_rows(doc):
    """(raw_value, posada, zvannya, pib_raw, period_text, days_text, basis_text,
    disappearance_text, heading_number) для КОЖНОГО рядка КОЖНОЇ таблиці рапорту,
    чий заголовок - звичайний пункт поточного місяця (_raw_value_for_heading не
    None - тобто НЕ підпункт "Прошу внести зміни...") і чий ПІБ непорожній.
    disappearance_text - None, якщо в таблиці немає колонки "Дата зникнення
    безвісти" (звичайні пункти, усі КРІМ "СПЕЦКОНТИНГЕНТ"). heading_number -
    ЛІТЕРАЛЬНИЙ номер пункту, ЯК НАПИСАНО в самому заголовку рапорту (напр. "1"
    для "1. Виплатити...") - НЕЗАЛЕЖНО від raw_value (сума 30/70/100/170 тощо чи
    службовий код "NOT_PAID"/"100_ШП" - зрозумілий коду, але НЕ людині, яка
    звіряє рапорт за його ВЛАСНОЮ нумерацією пунктів); порожній рядок, якщо в
    заголовку взагалі немає провідного числа (украй рідкісний випадок). Спільна
    основа для extract_timetable_rows_from_report (табель - будує сітку днів із
    periodів) і extract_report_entries (checker_accounting.report_checker/
    report_log_war_checker - звіряє ПЕРІОД/ДНІ/ПІБ/номер пункту, ЯК НАПИСАНО в
    самому рапорті) - обидва йдуть по ТИХ САМИХ рядках, лише по-різному
    використовують ці поля."""
    current_heading = ""
    for block in _iter_block_items(doc):
        if isinstance(block, Paragraph):
            if _is_heading(block):
                current_heading = block.text
            continue

        raw_value = _raw_value_for_heading(current_heading)
        if raw_value is None:
            continue

        heading_number_match = _LEADING_NUMBER_RE.match(current_heading)
        heading_number = heading_number_match.group(1) if heading_number_match else ""

        columns = _detect_columns(block.rows[0])
        for row in block.rows[1:]:
            cells = row.cells
            pib_raw = cells[columns["pib"]].text.strip()
            if not pib_raw:
                continue
            disappearance_col = columns["disappearance_date"]
            yield (
                raw_value, cells[columns["posada"]].text.strip(), cells[columns["zvannya"]].text.strip(),
                pib_raw, cells[columns["period"]].text, cells[columns["days"]].text.strip(),
                cells[columns["basis"]].text.strip(),
                cells[disappearance_col].text.strip() if disappearance_col is not None else None,
                heading_number,
            )


def extract_report_entries(file_path):
    """Повертає список {"raw_value", "ПОСАДА", "ЗВАННЯ", "ПІБ", "ПЕРІОД", "ДНІ",
    "ПІДСТАВА", "ДАТА_ЗНИКНЕННЯ", "НОМЕР_ПУНКТУ"} - ОДИН запис на КОЖЕН рядок
    КОЖНОЇ визнаної таблиці рапорту, з ПЕРІОД/ДНІ ЯК НАПИСАНО в самому рапорті
    (рядки тексту, без відновлення сітки днів) - призначено для
    checker_accounting.report_checker (звірка рапорту з ОБЛІК.xlsx за
    періодами/кількістю днів/написанням ПІБ), а не для табеля (див.
    _iter_recognized_rows). "ДАТА_ЗНИКНЕННЯ" - None для будь-якого пункту, КРІМ
    SPETSKONTYNGENT_VALUE ("СПЕЦКОНТИНГЕНТ") - лише його таблиця має цю
    колонку. "НОМЕР_ПУНКТУ" - літеральний номер пункту з самого заголовка
    рапорту (напр. "1") - підтверджено користувачем: показувати "raw_value"
    (30/70/100/170/службовий код) у звірках замість цього номера незрозуміло
    людині, яка звіряє рапорт за ЙОГО ВЛАСНОЮ нумерацією."""
    doc = Document(file_path)
    return [
        {
            "raw_value": raw_value, "ПОСАДА": posada, "ЗВАННЯ": zvannya, "ПІБ": pib_raw,
            "ПЕРІОД": period_text, "ДНІ": days_text, "ПІДСТАВА": basis_text, "ДАТА_ЗНИКНЕННЯ": disappearance_text,
            "НОМЕР_ПУНКТУ": heading_number,
        }
        for raw_value, posada, zvannya, pib_raw, period_text, days_text, basis_text, disappearance_text, heading_number
        in _iter_recognized_rows(doc)
    ]


def extract_timetable_rows_from_report(file_path, target_month=None, target_year=None):
    """Читає готовий рапорт на додаткову винагороду (.docx - будь-який, обраний
    користувачем з resources, за будь-який минулий чи поточний місяць) і
    відтворює з нього сітку табеля - {"ПОСАДА", "ЗВАННЯ", "ПІБ", <дата>: значення}
    на кожну людину, ЯКА РЕАЛЬНО Є В РАПОРТІ (а не в ОБЛІК.xlsx - цей модуль
    жодного разу не звертається до ОБЛІК.xlsx чи будь-яких інших джерел, лише
    до самого файлу рапорту).

    target_month/target_year (обидва задано, чи обидва None) - обмежує табель
    РІВНО обраним місяцем генерації (питання "Введіть номер місяця...",
    constants.MONTH/YEAR): дати з ІНШИХ місяців, знайдені в рапорті (напр. пункт
    "перебування на лікуванні" може охоплювати перехідний період із попереднього
    місяця в тому самому реченні), відкидаються, а колонки дат - ЦІЛИЙ
    календарний місяць (1-ше по останній день, calendar.monthrange), а не лише
    дні, для яких у рапорті справді є дані - як і колонки самого ОБЛІК.xlsx.
    Людина, у якої після цього фільтра не лишилось жодного дня, у результаті
    відсутня взагалі. target_month=None (за замовчуванням) - без фільтра, колонки
    дат - від найранішої до найпізнішої дати, знайденої в БУДЬ-ЯКОМУ періоді
    документа (використовується лише в тестах цього модуля - generate_timetable.py
    завжди передає обраний місяць).

    Обмеження (неминучі при відтворенні з готового ТЕКСТУ рапорту, а не з
    вихідних даних): (1) один пункт рапорту може об'єднувати кілька НАЗВАНИХ
    категорій ОБЛІК.xlsx (напр. "РТГр" і звичайний "БД(СЗ)" в пункті 100) в
    ОДНУ таблицю без будь-якої видимої позначки, звідки саме взявся кожен
    рядок - тому такі дні відтворюються під ЧИСЛОВИМ кодом пункту (100/70/...),
    а не назвою вихідної категорії; (2) так само "СЗЧ" і "Задув." в межах
    пункту "Не виплачувати..." невідрізнювані одне від одного в самому тексті
    рапорту - обидва відтворюються як "NOT_PAID".

    Кожен рядок результату додатково має "ПІДСТАВИ" - текст(и) колонки "Примітка
    (підстави)" з рядка(ів) цієї людини, чий пункт - один із _BASIS_TRACKED_VALUES
    (СПЕЦКОНТИНГЕНТ/100_ВП/100_ШП/NOT_PAID - підтверджено користувачем; порожній
    рядок для решти пунктів), і "ДАТА_ЗНИКНЕННЯ" - текст колонки "Дата зникнення
    безвісти" САМЕ з рядка(ів) SPETSKONTYNGENT_VALUE цієї людини (порожній рядок
    для решти, ця колонка є ЛИШЕ в таблиці СПЕЦКОНТИНГЕНТУ)."""
    doc = Document(file_path)

    rows_by_pib = {}
    deferred_updates = []

    for raw_value, posada, zvannya, pib_raw, period_text, _days_text, basis_text, disappearance_text, _heading_number in _iter_recognized_rows(doc):
        normalized = normalize_name(pib_raw)
        entry = rows_by_pib.setdefault(
            normalized,
            {"ПОСАДА": "", "ЗВАННЯ": "", "ПІБ": pib_raw, "ДАТА_ЗНИКНЕННЯ": "", "_basis_parts": [], "_days": {}},
        )
        entry["ПОСАДА"] = entry["ПОСАДА"] or posada
        entry["ЗВАННЯ"] = entry["ЗВАННЯ"] or zvannya
        if raw_value in _BASIS_TRACKED_VALUES and basis_text and basis_text not in entry["_basis_parts"]:
            entry["_basis_parts"].append(basis_text)
        if raw_value == SPETSKONTYNGENT_VALUE:
            entry["ДАТА_ЗНИКНЕННЯ"] = entry["ДАТА_ЗНИКНЕННЯ"] or disappearance_text or ""

        dates = _dates_from_period_text(period_text)
        if target_month is not None:
            dates = [date for date in dates if date.month == target_month and date.year == target_year]
        if raw_value == _DEFERRED_RAW_VALUE:
            deferred_updates.extend((normalized, date, raw_value) for date in dates)
        else:
            entry["_days"].update((date, raw_value) for date in dates)

    for normalized, date, raw_value in deferred_updates:
        rows_by_pib[normalized]["_days"].setdefault(date, raw_value)

    if target_month is not None:
        last_day = calendar.monthrange(target_year, target_month)[1]
        date_columns = [datetime(target_year, target_month, day) for day in range(1, last_day + 1)]
    else:
        all_dates = [date for entry in rows_by_pib.values() for date in entry["_days"]]
        if not all_dates:
            return [], []
        start, end = min(all_dates), max(all_dates)
        date_columns = [start + timedelta(days=i) for i in range((end - start).days + 1)]

    rows_with_data = [
        {
            "ПОСАДА": entry["ПОСАДА"], "ЗВАННЯ": entry["ЗВАННЯ"], "ПІБ": entry["ПІБ"],
            "ПІДСТАВИ": "; ".join(entry["_basis_parts"]), "ДАТА_ЗНИКНЕННЯ": entry["ДАТА_ЗНИКНЕННЯ"],
            **{date: entry["_days"].get(date) for date in date_columns},
        }
        for entry in rows_by_pib.values() if entry["_days"]
    ]
    return rows_with_data, date_columns
