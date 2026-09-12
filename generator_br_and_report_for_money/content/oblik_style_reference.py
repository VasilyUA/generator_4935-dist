import colorsys
import copy
import re
import zipfile
from xml.etree import ElementTree as ET

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from constants import PERSONEL_LIST_FILE_NAME, PERSONEL_LIST_SHEET_NAME
from utils.excel_reader import find_file_with_any_extension

# Табель має ВИГЛЯДАТИ як ОБЛІК.xlsx (кольори/шрифти/межі клітинок/ширини колонок/
# умовне форматування по значенню дня), хоча самі ДАНІ табеля беруться ЦІЛКОМ з
# обраного рапорту (content.report_document_reader) - ОБЛІК.xlsx тут відкривається
# ЛИШЕ заради стилю, жодне значення з нього не потрапляє в саму таблицю табеля.

_THEME_NS = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
# Порядок у самому <a:clrScheme> theme1.xml - dk1,lt1,dk2,lt2,accent1..6; але
# посилання theme=N у стилях клітинок Excel використовують ІНШИЙ порядок (0 і 1,
# 2 і 3 поміняні місцями відносно clrScheme) - задокументована особливість
# формату OOXML, не помилка тут.
_EXCEL_THEME_ORDER = ["lt1", "dk1", "lt2", "dk2", "accent1", "accent2", "accent3", "accent4", "accent5", "accent6"]

_SEARCH_KEYWORD_RE = re.compile(r'SEARCH\("([^"]+)"')

# Колонки самого ОБЛІК.xlsx (fixed_columns_count=4: ПІДРОЗДІЛ/ПОСАДА/ЗВАННЯ/ПІБ),
# звідки береться стиль - ПОСАДА (2) для звичайних "лейбл"-колонок табеля
# (табель не має власної ПІДРОЗДІЛ - рапорт її не містить), ПІБ (4) для колонки
# ПІБ (у ОБЛІК вона зазвичай ширша й вирівняна ліворуч), перша дата (5) для
# колонок днів.
_LABEL_STYLE_COL, _PIB_STYLE_COL, _DATE_STYLE_COL = 2, 4, 5


def _read_theme_colors(file_path):
    """{ім'я кольору теми: "RRGGBB"} з xl/theme/theme1.xml (.xlsx/.xlsm - звичайний
    zip-архів) - частина кольорів умовного форматування ОБЛІК.xlsx задана через
    тему (theme+tint), а не прямим RGB, тож без цього їх не розв'язати."""
    with zipfile.ZipFile(file_path) as archive:
        theme_xml = archive.read("xl/theme/theme1.xml")
    scheme = ET.fromstring(theme_xml).find(".//a:clrScheme", _THEME_NS)
    colors = {}
    for child in scheme:
        name = child.tag.split("}")[-1]
        srgb = child.find("a:srgbClr", _THEME_NS)
        sys_clr = child.find("a:sysClr", _THEME_NS)
        if srgb is not None:
            colors[name] = srgb.get("val")
        elif sys_clr is not None:
            colors[name] = sys_clr.get("lastClr")
    return colors


def _apply_tint(rgb_hex, tint):
    """Той самий алгоритм, що й сам Excel (ECMA-376, HSL) - освітлює (tint>0) чи
    затемнює (tint<0) базовий колір теми, як показує колірна палітра Excel."""
    r, g, b = (int(rgb_hex[i:i + 2], 16) / 255 for i in (0, 2, 4))
    h, l, s = colorsys.rgb_to_hls(r, g, b)
    l = l * (1.0 + tint) if tint < 0 else l * (1.0 - tint) + tint
    r, g, b = colorsys.hls_to_rgb(h, min(max(l, 0.0), 1.0), s)
    return f"{round(r * 255):02X}{round(g * 255):02X}{round(b * 255):02X}"


def _resolve_fill_color(color, theme_colors):
    """openpyxl Color (rgb чи theme+tint) -> "RRGGBB", або None, якщо колір
    індексований чи іншого рідкісного типу, який тут не розпізнається - клітинка
    тоді лишається без заливки, а не падає."""
    if color is None:
        return None
    if color.type == "rgb" and isinstance(color.rgb, str) and len(color.rgb) == 8:
        return color.rgb[2:]
    if color.type == "theme" and isinstance(color.theme, int) and color.theme < len(_EXCEL_THEME_ORDER):
        base_rgb = theme_colors.get(_EXCEL_THEME_ORDER[color.theme])
        if base_rgb:
            return _apply_tint(base_rgb, color.tint or 0.0)
    return None


def _keyword_colors(ws, dxf_styles, theme_colors):
    """{ключове_слово: "RRGGBB"} з правил умовного форматування ОБЛІК.xlsx виду
    "клітинка МІСТИТЬ ключове_слово" (напр. "30"/"100"/"СЗЧ"/"ВП"/...) - у
    пріоритеті самого файлу (rule.priority - менше число застосовується першим),
    перше знайдене правило для кожного слова перемагає, як і в самому Excel."""
    colors = {}
    for cf_range in ws.conditional_formatting:
        rules = sorted(cf_range.rules, key=lambda rule: rule.priority if rule.priority is not None else 0)
        for rule in rules:
            if rule.type != "containsText" or not rule.formula or rule.dxfId is None:
                continue
            match = _SEARCH_KEYWORD_RE.search(rule.formula[0])
            if not match or match.group(1) in colors:
                continue
            dxf = dxf_styles[rule.dxfId] if rule.dxfId < len(dxf_styles) else None
            resolved = _resolve_fill_color(dxf.fill.bgColor, theme_colors) if dxf and dxf.fill else None
            if resolved:
                colors[match.group(1)] = resolved
    return colors


def load_timetable_style_reference():
    """Стиль (шрифти/кольори/ширини колонок/висоти рядків/умовне форматування по
    значенню дня) з РЕАЛЬНОГО ОБЛІК.xlsx - виключно для оформлення табеля.

    Файл недоступний, пошкоджений чи іншої несподіваної структури - повертає
    None (не піднімає виключення) - generate_timetable.py тоді використовує
    прості типові стилі замість падіння: табель важливіший за його оформлення."""
    try:
        file_path = find_file_with_any_extension(PERSONEL_LIST_FILE_NAME)
        wb = load_workbook(file_path)
        ws = wb[PERSONEL_LIST_SHEET_NAME] if PERSONEL_LIST_SHEET_NAME in wb.sheetnames else wb[wb.sheetnames[0]]
        theme_colors = _read_theme_colors(file_path)

        header_label_cell = ws.cell(row=1, column=_LABEL_STYLE_COL)
        header_pib_cell = ws.cell(row=1, column=_PIB_STYLE_COL)
        header_date_cell = ws.cell(row=1, column=_DATE_STYLE_COL)
        body_label_cell = ws.cell(row=2, column=_LABEL_STYLE_COL)

        return {
            "header_font": copy.copy(header_label_cell.font),
            "header_fill": copy.copy(header_label_cell.fill),
            "header_border": copy.copy(header_label_cell.border),
            "header_alignment_label": copy.copy(header_pib_cell.alignment),
            "header_alignment_date": copy.copy(header_date_cell.alignment),
            "body_font": copy.copy(body_label_cell.font),
            "body_alignment": copy.copy(body_label_cell.alignment),
            "body_border": copy.copy(body_label_cell.border),
            "label_col_width": ws.column_dimensions[get_column_letter(_LABEL_STYLE_COL)].width,
            "pib_col_width": ws.column_dimensions[get_column_letter(_PIB_STYLE_COL)].width,
            "date_col_width": ws.column_dimensions[get_column_letter(_DATE_STYLE_COL)].width,
            "header_row_height": ws.row_dimensions[1].height,
            "body_row_height": ws.row_dimensions[2].height,
            "keyword_colors": _keyword_colors(ws, wb._differential_styles.styles, theme_colors),
        }
    except Exception:
        return None
