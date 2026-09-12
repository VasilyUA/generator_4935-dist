import os
import re

from docx import Document

from constants import SHORT_UNIT_BATTALION, SHORT_UNIT_BRIGADE
from content.money_report_helpers import normalize_name
from utils.date_utils import to_date
from utils.logging_utils import print_red

# "№79 ЗАВДАННЯ 05.06.2026.docx"/"№27 ЩОДЕННА 01.06.2026.doc" - номер, тип
# (ЩОДЕННА/ЗАВДАННЯ), дата, розширення (.docx чи ЛЕГАСІ-бінарний .doc - на відміну
# від sync/constants_sync.py, тут .doc теж підтримується, бо файли реально
# трапляються в обох форматах - extract_document_text відкриває .doc через Word COM).
_DAILY_OR_TASK_FILENAME_RE = re.compile(r"№(\d+)\s*(ЩОДЕННА|ЗАВДАННЯ)\s+(\d{2}\.\d{2}\.\d{4})\.(docx|doc)$", re.IGNORECASE)


def extract_document_text(file_path):
    """Текст .docx (python-docx, абзаци+таблиці) чи .doc (Word COM - python-docx
    НЕ вміє читати старий бінарний формат узагалі). Будь-яка помилка (немає
    pywin32, Word не встановлено/не відкрився, файл пошкоджено) - червоне
    попередження і порожній рядок, а не падіння: пошук по решті файлів триває."""
    ext = os.path.splitext(file_path)[1].lower()
    if ext == ".docx":
        try:
            doc = Document(file_path)
        except Exception as error:
            print_red(f"Не вдалось прочитати {file_path}: {error}")
            return ""
        paragraphs_text = "\n".join(p.text for p in doc.paragraphs)
        tables_text = "\n".join(cell.text for table in doc.tables for row in table.rows for cell in row.cells)
        return f"{paragraphs_text}\n{tables_text}"
    if ext == ".doc":
        return extract_document_text_via_word_com(file_path)
    return ""


def extract_document_text_via_word_com(file_path):
    try:
        import pythoncom
        import win32com.client
    except ImportError:
        print_red(f"pywin32 недоступний - неможливо прочитати {file_path} (.doc).")
        return ""

    pythoncom.CoInitialize()
    try:
        word = win32com.client.Dispatch("Word.Application")
        word.Visible = False
        try:
            document = word.Documents.Open(os.path.abspath(file_path), ReadOnly=True)
        except Exception as error:
            print_red(f"Не вдалось відкрити {file_path} у Word: {error}")
            return ""
        try:
            return document.Content.Text
        finally:
            document.Close(SaveChanges=False)
    except Exception as error:
        print_red(f"Не вдалось прочитати {file_path}: {error}")
        return ""
    finally:
        pythoncom.CoUninitialize()


def _scan_daily_or_task_files(folder, scope):
    """Спільне сканування ЩОДЕННА/ЗАВДАННЯ файлів folder (рекурсивно, os.walk, за
    іменем у межах директорії) - спільна основа для find_person_document_references
    (підстави за ПІБ) і collect_missing_document_coverage_warnings (перевірка покриття по датах):
    для КОЖНОГО файлу, чий текст вдалось прочитати, повертає (дата документа - date,
    готовий рядок підстави, НОРМАЛІЗОВАНИЙ - через .upper() - текст файлу).

    scope="only_task" - пропускає файли ЩОДЕННА, лишає лише ЗАВДАННЯ (менша,
    швидша вибірка); будь-яке інше значення (зокрема "all") - обидва типи."""
    for dirpath, _dirnames, filenames in os.walk(folder):
        for filename in sorted(filenames):
            match = _DAILY_OR_TASK_FILENAME_RE.search(filename)
            if not match:
                continue
            number, doc_type, date_str, _ext = match.groups()
            if scope == "only_task" and doc_type.upper() != "ЗАВДАННЯ":
                continue

            text = extract_document_text(os.path.join(dirpath, filename))
            if not text:
                continue
            normalized_text = " ".join(text.split()).upper()

            line = f"БР {SHORT_UNIT_BATTALION} {SHORT_UNIT_BRIGADE} №{number} від {date_str}"
            yield to_date(date_str), line, normalized_text


def find_person_document_references(folder, roster_names, scope):
    """Шукає ПОВНЕ ПІБ (roster_names - вже normalize_name'ені) у тексті КОЖНОГО
    файлу ЩОДЕННА/ЗАВДАННЯ в folder - на відміну від
    sync.constants_sync.sync_document_numbers_from_folder (один номер на ДАТУ,
    застосовується до всіх однаково), тут номер того самого документа
    підставляється ЛИШЕ тим, чиє ПІБ дійсно згадане в його тексті.

    Повертає {normalized_pib: [рядок, ...]} - лише для людей, кого дійсно
    знайдено хоча б в одному файлі; порядок рядків - за порядком сканування
    файлів (_scan_daily_or_task_files)."""
    results = {}
    for _date, line, normalized_text in _scan_daily_or_task_files(folder, scope):
        for name in roster_names:
            if name and name in normalized_text:
                results.setdefault(name, []).append(line)
    return results


def _format_display_pib(pib):
    """"Прізвище Ім'я По батькові" для читабельного показу в повідомленнях -
    перше слово (прізвище) ВЕЛИКИМИ літерами, решта слів (ім'я, по батькові) -
    лише перша літера велика, решта маленькі - підтверджено користувачем.
    Регістр самого аргументу не важливий (працює однаково і з normalize_name'еним
    "ПЕРШИЙ ПЕРШИЙ ПЕРШИЙ", і зі звичайним "Перший Перший Перший") - результат
    завжди нормалізується до цього єдиного вигляду."""
    words = pib.split()
    if not words:
        return pib
    surname, *rest = words
    return " ".join([surname.upper()] + [word[:1].upper() + word[1:].lower() for word in rest])


def collect_missing_document_coverage_warnings(folder, person_dates, scope, rank_by_pib=None, subdivision_by_pib=None):
    """Для КОЖНОЇ дати з person_dates[normalized_pib] (дні фактичної участі людини
    в пункті 100 - єдиному, для якого застосовується, підтверджено користувачем,
    лише для рапорту на поправки в наказі за попередні місяці) - перевіряє, чи є
    ХОЧ ОДИН файл ЩОДЕННА/ЗАВДАННЯ, датований САМЕ цим днем (за назвою файлу), чий
    текст дійсно згадує ПІБ цієї людини.

    Повертає список готових рядків попередження (не друкує їх сама, на відміну
    від старої версії) - підтверджено користувачем: ці попередження мають
    потрапити в САМ WORD-документ (окремий список у кінці), а не в термінал -
    викликач (content.report_changes) збирає їх і передає далі в рендер
    (generators.generate_report_for_get_money._add_missing_coverage_section).

    ПІБ у самому тексті попередження - ЗАВЖДИ у вигляді "ПРІЗВИЩЕ Ім'я По
    батькові" (_format_display_pib), незалежно від регістру ключа person_dates.
    rank_by_pib/subdivision_by_pib - опційно, {normalized_pib: ЗВАННЯ/ПІДРОЗДІЛ} -
    додають звання й підрозділ ПЕРЕД ПІБ (у порядку "підрозділ звання ПІБ", як у
    колонках звичайної таблиці рапорту - ПІДРОЗДІЛ, ПОСАДА, ЗВАННЯ, ПІБ), кожен -
    лише якщо передано і непорожнє для цієї людини - підтверджено користувачем.

    person_dates - {normalized_pib: [дата, ...]} (date чи datetime - порівнюється
    лише за .date() через to_date). find_person_document_references (вище)
    окремо й незалежно будує самі рядки підстави."""
    covered_dates_by_pib = {}
    for date, _line, normalized_text in _scan_daily_or_task_files(folder, scope):
        for pib in person_dates:
            if pib and pib in normalized_text:
                covered_dates_by_pib.setdefault(pib, set()).add(date)

    warnings = []
    for pib, dates in person_dates.items():
        covered = covered_dates_by_pib.get(pib, set())
        subdivision = (subdivision_by_pib or {}).get(pib, "")
        rank = (rank_by_pib or {}).get(pib, "")
        name_parts = [part for part in (subdivision, rank, _format_display_pib(pib)) if part]
        display_name = " ".join(name_parts)
        for raw_date in sorted(dates):
            date = to_date(raw_date)
            if date not in covered:
                warnings.append(
                    f"Немає документа (ЩОДЕННА/ЗАВДАННЯ) за {date.strftime('%d.%m.%Y')} "
                    f"для {display_name} - підстава на цей день відсутня."
                )
    return warnings
